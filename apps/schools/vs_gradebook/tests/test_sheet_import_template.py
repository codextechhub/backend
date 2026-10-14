"""The template download: who gets it, what it holds, and what it leaves out."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from vs_rbac.tests.helpers import install_declared_fields, set_field_access

from .import_base import ImportFixture

MARK = "gradebook.sheet.mark"


def read(response):
    book = load_workbook(io.BytesIO(response.content))
    return book, list(book["Scores"].iter_rows(values_only=True))


class TemplateAccessTests(ImportFixture):
    def test_a_main_teacher_without_the_import_key_is_refused(self):
        self.assertEqual(self.template(self.keyless, self.jss1b, self.maths).status_code, 403)

    def test_a_teacher_who_is_not_the_main_teacher_is_refused(self):
        for user in (self.bello, self.musa):
            with self.subTest(user=user.email):
                response = self.template(user, self.jss1a, self.maths)
                self.assertEqual(response.status_code, 403)
                self.assertEqual(response.data["error"]["code"], "NOT_MAIN_TEACHER")

    def test_enter_any_without_the_import_key_is_refused(self):
        self.assertEqual(self.template(self.enter_only, self.jss1a, self.maths).status_code, 403)

    def test_a_user_with_no_key_is_refused(self):
        self.assertEqual(self.template(self.nobody, self.jss1a, self.maths).status_code, 403)

    def test_another_schools_class_is_a_404(self):
        self.assertEqual(self.template(self.eze, self.solo_class, self.solo_maths).status_code, 404)

    def test_a_class_in_a_branch_the_caller_cannot_see_is_a_404(self):
        self.assertEqual(self.template(self.ikeja_importer, self.jss1a, self.maths).status_code, 404)

    def test_enter_any_with_the_import_key_may_download_any_sheet(self):
        self.assertEqual(self.template(self.importer, self.jss1a, self.maths).status_code, 200)
        self.assertEqual(self.template(self.lekki_importer, self.jss1a, self.maths).status_code, 200)

    def test_a_locked_sheet_can_still_be_downloaded(self):
        self.make_sheet(self.jss1a, self.maths, status="SUBMITTED")
        self.assertEqual(self.template(self.eze, self.jss1a, self.maths).status_code, 200)


class TemplateContentTests(ImportFixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.make_sheet(
            cls.jss1a, cls.maths,
            rows={cls.chiamaka: (18, 15, 46), cls.tunde: ("ABS", "EXC", None)},
        )

    def test_it_is_an_xlsx_named_for_the_class_subject_and_term(self):
        response = self.template(self.eze, self.jss1a, self.maths)
        self.assertIn("spreadsheetml", response["Content-Type"])
        disposition = response["Content-Disposition"]
        self.assertIn("attachment", disposition)
        for word in ("mathematics", "jss1-a", "first-term"):
            self.assertIn(word, disposition)
        self.assertTrue(disposition.endswith('.xlsx"'))

    def test_the_header_names_each_component_with_its_maximum(self):
        _, rows = read(self.template(self.eze, self.jss1a, self.maths))
        self.assertEqual(
            list(rows[0]),
            ["Admission number", "Student", "1st CA (max 20)", "2nd CA (max 20)", "Examination (max 60)"],
        )

    def test_the_rows_follow_the_sheets_roll_order(self):
        sheet = self.data(self.sheet(self.eze, self.jss1a, self.maths))
        _, rows = read(self.template(self.eze, self.jss1a, self.maths))
        self.assertEqual(
            [r[0] for r in rows[1:]], [r["student"]["admission_number"] for r in sheet["rows"]],
        )
        self.assertEqual([r[1] for r in rows[1:]], [r["student"]["name"] for r in sheet["rows"]])

    def test_current_marks_are_filled_in_with_abs_and_exc_as_text(self):
        _, rows = read(self.template(self.eze, self.jss1a, self.maths))
        by_number = {r[0]: r[2:] for r in rows[1:]}
        self.assertEqual(by_number["BFS/001"], (18, 15, 46))
        self.assertEqual(by_number["BFS/002"], ("ABS", "EXC", None))
        self.assertEqual(by_number["BFS/003"], (None, None, None))

    def test_a_notes_sheet_explains_blanks_abs_and_exc(self):
        book, _ = read(self.template(self.eze, self.jss1a, self.maths))
        text = " ".join(str(r[0]) for r in book["Notes"].iter_rows(values_only=True) if r[0])
        for word in ("ABS", "EXC", "empty"):
            self.assertIn(word, text)

    def test_an_elective_sheet_lists_only_the_students_who_take_it(self):
        _, rows = read(self.template(self.importer, self.jss1a, self.french))
        self.assertEqual({r[0] for r in rows[1:]}, {"BFS/001", "BFS/002"})
        self.assertEqual(self.template(self.importer, self.jss1b, self.french).status_code, 404)

    def test_an_unstarted_sheet_has_the_roll_and_no_marks(self):
        _, rows = read(self.template(self.adewale, self.sss1a, self.maths))
        self.assertEqual({r[0] for r in rows[1:]}, {"BFS/008", "BFS/009"})
        self.assertTrue(all(cell is None for r in rows[1:] for cell in r[2:]))

    def test_a_school_with_one_branch_gets_the_same_shape(self):
        _, rows = read(self.template(self.solo_teacher, self.solo_class, self.solo_maths))
        self.assertEqual({r[0] for r in rows[1:]}, {"SUN/001", "SUN/002"})
        self.assertEqual(len(rows[0]), 5)


class TemplateFieldAccessTests(ImportFixture):
    def test_a_role_that_may_not_read_marks_gets_the_roll_with_no_marks(self):
        self.make_sheet(self.jss1a, self.maths, rows={self.chiamaka: (18, 15, 46)})
        install_declared_fields("gradebook.sheet")
        set_field_access(self.teacher_role, MARK, read=False, write=False)
        _, rows = read(self.template(self.eze, self.jss1a, self.maths))
        self.assertIn("BFS/001", [r[0] for r in rows[1:]])
        self.assertTrue(all(cell is None for r in rows[1:] for cell in r[2:]))
