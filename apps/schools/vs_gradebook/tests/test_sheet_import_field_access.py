"""Which parts of an upload a role reads and writes: the template, the preview, the confirm."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from schools.vs_gradebook.models import Score
from vs_rbac.models import FieldDefinition
from vs_rbac.tests.helpers import install_declared_fields, set_field_access

from .import_base import ImportFixture

MARK = "gradebook.sheet.mark"
ROWS = [["BFS/001", "Chiamaka", 18, 16, None], ["BFS/002", "Tunde", 10, 12, 50]]


class UploadFieldAccessTests(ImportFixture):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        install_declared_fields("gradebook.sheet")
        cls.make_sheet(cls.jss1a, cls.maths, rows={cls.chiamaka: (18, 15, 46)})

    def close(self, *, read=False, write=False):
        set_field_access(self.teacher_role, MARK, read=read, write=write)

    def test_the_preview_names_are_carried_by_the_mark_switch(self):
        self.assertEqual(
            set(FieldDefinition.objects.get(key=MARK).api_names),
            {"mark", "marker", "cells", "value", "existing", "outcome", "replaces"},
        )

    def test_by_default_the_preview_shows_values_outcomes_and_the_replacement_count(self):
        found = self.preview(self.eze, self.jss1a, self.maths, ROWS)
        self.assertEqual(found["summary"]["replaces"], 1)
        cell = self.cell_of(found["rows"][0], self.ca2)
        self.assertEqual((cell["value"], cell["existing"], cell["outcome"]), (16, 15, "REPLACES"))

    def test_a_role_that_may_not_read_marks_is_refused_the_upload_and_nothing_is_kept(self):
        self.close()
        response = self.upload(self.eze, self.jss1a, self.maths, self._book())
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.data["error"]["code"], "field_write_denied")

    def test_a_role_that_may_read_but_not_write_marks_is_refused_the_upload(self):
        self.close(read=True, write=False)
        self.assertEqual(self.upload(self.eze, self.jss1a, self.maths, self._book()).status_code, 403)

    def test_it_may_still_download_the_template_with_the_marks_when_it_can_read_them(self):
        self.close(read=True, write=False)
        book = load_workbook(io.BytesIO(self.template(self.eze, self.jss1a, self.maths).content))
        rows = list(book["Scores"].iter_rows(values_only=True))
        self.assertIn((18, 15, 46), [r[2:] for r in rows[1:]])

    def test_a_confirm_is_refused_when_the_role_lost_the_right_to_write_marks_since_the_preview(self):
        found = self.preview(self.eze, self.jss1a, self.maths, ROWS)
        self.close(read=True, write=False)
        response = self.confirm(self.eze, self.jss1a, self.maths, found["token"], [3])
        self.assertEqual(response.status_code, 403)
        self.assertEqual(Score.objects.count(), 3)

    def test_the_confirm_answer_follows_the_read_switches_like_any_sheet_read(self):
        found = self.preview(self.eze, self.jss1a, self.maths, ROWS)
        set_field_access(self.teacher_role, "gradebook.sheet.row_result", read=False, write=False)
        body = self.data(self.confirm(self.eze, self.jss1a, self.maths, found["token"], [3]))
        self.assertNotIn("total", body["rows"][0])
        self.assertIn("cells", body["rows"][0])

    @staticmethod
    def _book():
        from .import_base import workbook

        return workbook(ROWS)
