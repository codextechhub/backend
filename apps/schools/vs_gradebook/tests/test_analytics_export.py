"""The two analytics tables as files: the screen's figures, cut to the caller.

Every file is compared with the JSON the analytics endpoint serves the same
caller for the same term and branch, so a file and the screen cannot disagree.
"""
from __future__ import annotations

import csv
import datetime as dt
import io

from openpyxl import load_workbook

from schools.vs_academics.models import AcademicTerm
from vs_audit.models import AuditEvent

from .analytics_base import AnalyticsFixture, snapshot


def read_csv(response):
    """``(title, subtitle, headers, rows)`` of a downloaded CSV."""
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8"))))
    return rows[0][0], rows[1][0], rows[3], rows[4:]


def read_xlsx(response):
    """The same four things from a downloaded workbook."""
    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = [[("" if c is None else c) for c in row] for row in sheet.iter_rows(values_only=True)]
    return rows[0][0], rows[1][0], [str(c) for c in rows[3]], rows[4:]


def cells(rows):
    """*rows* with every value as text, so a CSV and a workbook compare."""
    return [[str(value) for value in row] for row in rows]


class RequestTests(AnalyticsFixture):
    def test_an_unknown_or_missing_table_is_a_400(self):
        for params in ({}, {"table": "teachers"}, {"table": ""}, {"table": "tops"}):
            with self.subTest(params=params):
                self.assertEqual(self.export(self.reader, **params).status_code, 400)

    def test_an_unknown_file_type_is_a_400(self):
        for kind in ("pdf", "docx", "json"):
            with self.subTest(kind=kind):
                self.assertEqual(
                    self.export(self.reader, table="subjects", export=kind).status_code, 400,
                )

    def test_a_bad_request_is_not_recorded(self):
        self.export(self.reader, table="bogus")
        self.export(self.nobody, table="subjects")
        self.assertFalse(AuditEvent.objects.filter(entity_type="GradebookAnalytics").exists())

    def test_the_file_is_an_attachment_named_for_the_table_and_the_term(self):
        response = self.export(self.reader, table="subjects", export="csv")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        self.assertEqual(
            response["Content-Disposition"],
            'attachment; filename="gradebook_subjects_first_term.csv"',
        )
        response = self.export(self.reader, table="movers", export="xlsx")
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn('filename="gradebook_movers_first_term.xlsx"', response["Content-Disposition"])

    def test_csv_is_the_default_file_type(self):
        self.assertEqual(self.export(self.reader, table="subjects")["Content-Type"], "text/csv")

    def test_the_term_names_the_file_when_one_is_asked_for(self):
        response = self.export(self.reader, table="subjects", term=self.second_term.pk)
        self.assertIn("gradebook_subjects_second_term.csv", response["Content-Disposition"])


class SubjectsFileTests(AnalyticsFixture):
    def expected(self, body, with_branch=True):
        homes = {c["id"]: c["branch"]["name"] for c in body["scope"]["classes"] if c["branch"]}
        rows = []
        for row in body["subjects"]:
            line = [row["subject"]["name"], row["class"]["name"]]
            if with_branch:
                line.append(homes[row["class"]["id"]])
            spread = ", ".join(f"{e['grade']} {e['count']}" for e in row["spread"] if e["count"])
            line += [row["average"], f"{row['pass_rate']}%", spread, row["marked"]]
            rows.append(line)
        return cells(rows)

    def test_the_csv_holds_the_screens_rows_in_the_screens_order(self):
        body = self.data(self.analytics(self.reader))
        title, subtitle, headers, rows = read_csv(self.export(self.reader, table="subjects"))
        self.assertEqual(title, "Performance by subject")
        self.assertEqual(subtitle, "First Term, 2025/2026 · All branches")
        self.assertEqual(
            headers,
            ["Subject", "Class", "Branch", "Average", "Pass rate", "Grade spread", "Marked"],
        )
        self.assertEqual(rows, self.expected(body))
        self.assertEqual(rows[1], [
            "Mathematics", "JSS1 A", "Lekki Branch", "60.00", "75%", "A 2, B 1, F 1", "4",
        ])

    def test_the_workbook_holds_the_same_rows(self):
        body = self.data(self.analytics(self.reader))
        title, _, headers, rows = read_xlsx(self.export(self.reader, table="subjects", export="xlsx"))
        self.assertEqual(title, "Performance by subject")
        self.assertEqual(headers[0], "Subject")
        self.assertEqual(cells(rows), self.expected(body))

    def test_a_branch_pinned_readers_file_holds_only_their_branchs_classes(self):
        _, subtitle, _, rows = read_csv(self.export(self.lekki_reader, table="subjects"))
        self.assertEqual({row[1] for row in rows}, {"JSS1 A", "JSS1 B"})
        self.assertIn("the reader's branches only", subtitle)
        _, _, _, rows = read_csv(self.export(self.ikeja_reader, table="subjects"))
        self.assertEqual({row[1] for row in rows}, {"JSS1 C"})

    def test_a_named_branch_narrows_the_file_and_the_subtitle(self):
        _, subtitle, _, rows = read_csv(
            self.export(self.reader, table="subjects", branch=self.ikeja.pk),
        )
        self.assertEqual({row[1] for row in rows}, {"JSS1 C"})
        self.assertEqual(subtitle, "First Term, 2025/2026 · Ikeja Branch")

    def test_a_school_with_one_branch_has_no_branch_column(self):
        self.put(self.solo_class, self.solo_term, [
            (self.solo_pupil, snapshot(
                self.solo_pupil, self.solo_class, self.solo_term, [(self.solo_maths, 77)],
                average=77, position=1,
            )),
        ], branch=self.solo_branch)
        _, subtitle, headers, rows = read_csv(self.export(self.solo_reader, table="subjects"))
        self.assertEqual(
            headers, ["Subject", "Class", "Average", "Pass rate", "Grade spread", "Marked"],
        )
        self.assertEqual(rows, [["Mathematics", "Basic 3", "77.00", "100%", "A 1", "1"]])
        self.assertEqual(subtitle, "First Term, 2025/2026")

    def test_an_empty_term_is_the_headings_alone(self):
        empty = AcademicTerm.all_objects.create(
            tenant=self.tenant, session=self.year, name="Third Term", order_index=3,
            start_date=dt.date(2026, 4, 20), end_date=dt.date(2026, 7, 17),
        )
        _, _, headers, rows = read_csv(self.export(self.reader, table="subjects", term=empty.pk))
        self.assertEqual(rows, [])
        self.assertEqual(headers[0], "Subject")

    def test_a_cell_that_would_run_as_a_formula_is_written_as_text(self):
        body = snapshot(self.chiamaka, self.jss1a, self.first_term, [(self.maths, 80)], average=80)
        body["subjects"][0]["subject"] = {"id": 999999, "name": "=1+1"}
        from ..models import StudentTermResult

        StudentTermResult.objects.filter(
            student=self.chiamaka, class_term__term=self.first_term,
        ).update(snapshot=body)
        _, _, _, rows = read_csv(self.export(self.reader, table="subjects"))
        self.assertIn("'=1+1", {row[0] for row in rows})


class MoversFileTests(AnalyticsFixture):
    def test_the_csv_holds_the_screens_movers_biggest_fall_first(self):
        body = self.data(self.analytics(self.reader))
        title, subtitle, headers, rows = read_csv(self.export(self.reader, table="movers"))
        self.assertEqual(title, "Change in overall average")
        self.assertEqual(subtitle, "First Term, 2025/2026 · All branches")
        self.assertEqual(headers, [
            "Student", "Admission number", "Class", "Branch", "Compared with",
            "Previous average", "Average this term", "Change",
        ])
        homes = {c["id"]: c["branch"]["name"] for c in body["scope"]["classes"]}
        self.assertEqual(rows, cells([
            [m["student"]["name"], m["student"]["admission_number"], m["class"]["name"],
             homes[m["class"]["id"]], m["from_term"]["name"], m["from"], m["to"], m["change"]]
            for m in body["movers"]
        ]))
        self.assertEqual([row[0] for row in rows], ["Sade Ojo", "Tunde Bello", "Chiamaka Nwosu"])
        self.assertEqual(rows[0][5:], ["99.50", "90.00", "-9.50"])

    def test_the_workbook_holds_the_same_rows(self):
        _, _, _, csv_rows = read_csv(self.export(self.reader, table="movers"))
        _, _, _, xlsx_rows = read_xlsx(self.export(self.reader, table="movers", export="xlsx"))
        self.assertEqual(cells(xlsx_rows), csv_rows)

    def test_a_branch_pinned_reader_gets_only_their_own_students(self):
        _, _, _, rows = read_csv(self.export(self.ikeja_reader, table="movers"))
        self.assertEqual([row[0] for row in rows], ["Sade Ojo"])


class AuditTests(AnalyticsFixture):
    def test_each_export_is_recorded_with_who_which_table_and_which_term(self):
        self.export(self.reader, table="movers", export="xlsx")
        event = AuditEvent.objects.get(entity_type="GradebookAnalytics")
        self.assertEqual(event.action_type, "EXPORT_COMPLETED")
        self.assertEqual(event.module_key, "GRADEBOOK")
        self.assertEqual(event.tenant_id, self.tenant.pk)
        self.assertEqual(event.actor_user_id, self.reader.pk)
        self.assertEqual(event.metadata["table"], "movers")
        self.assertEqual(event.metadata["term_id"], self.first_term.pk)
        self.assertEqual(event.metadata["file_type"], "xlsx")
        self.assertEqual(event.metadata["rows"], 3)
        self.assertIn("First Term", event.summary)

    def test_the_trail_holds_no_figure_and_no_name(self):
        self.export(self.reader, table="movers")
        event = AuditEvent.objects.get(entity_type="GradebookAnalytics")
        text = f"{event.summary} {event.metadata}"
        for needle in ("Sade", "Tunde", "99.50", "-9.50"):
            self.assertNotIn(needle, text)
