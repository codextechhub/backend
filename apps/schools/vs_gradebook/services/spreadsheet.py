"""Reading an uploaded score sheet, and writing the template it is filled in on.

**Reading** turns an upload into a :class:`Grid`: the heading row and every
data row that holds something, each with the row number a person sees in the
spreadsheet, so a problem found in the file can be pointed at. The format is
judged by its extension and, for a workbook, by its signature. Size and row
count are bounded before anything is parsed, because the file is untrusted. A
workbook is read cell by cell with its stored types (a number stays a number,
and a date stays a date, so it can be refused as the mark it is not); a CSV is
read as text. Blank rows are skipped and keep their numbers.

**Writing** is the template a teacher downloads: the scores on one worksheet,
with the headings the reader expects, and the rules for filling it on another,
so the file explains itself offline. Admission numbers are written as text, so
a spreadsheet does not turn ``00123`` into ``123``.
"""
from __future__ import annotations

import csv
import io
import os
from itertools import islice
from dataclasses import dataclass, field

from ..constants import (
    IMPORT_ADMISSION_HEADING,
    IMPORT_FORMATS,
    IMPORT_MAX_BYTES,
    IMPORT_MAX_ROWS,
    IMPORT_NOTES_SHEET,
    IMPORT_SCORES_SHEET,
    IMPORT_STUDENT_HEADING,
)
from ..exceptions import UnreadableUpload, UnsupportedUpload, UploadTooLarge

_ZIP_SIGNATURE = b"PK\x03\x04"
_CSV_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
_MAX_COLUMNS = 60


@dataclass
class Grid:
    """The heading row of an upload and its data rows.

    ``rows`` is ``[(row_number, values)]`` with *values* padded to the heading's
    width. A value is ``None`` for an empty cell, else the cell as stored: a
    stripped string, a number, or a date for a workbook date.
    """

    headings: list
    rows: list = field(default_factory=list)


def _extension(name: str) -> str:
    return os.path.splitext(name or "")[1].lower().lstrip(".")


def _clean(value):
    """An empty cell as None, text stripped, anything else as stored."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def read_grid(upload) -> Grid:
    """The grid of an uploaded ``.xlsx`` or ``.csv``, or a refusal saying what is wrong with the file."""
    extension = _extension(getattr(upload, "name", ""))
    if extension not in IMPORT_FORMATS:
        raise UnsupportedUpload()
    if getattr(upload, "size", 0) and upload.size > IMPORT_MAX_BYTES:
        raise UploadTooLarge()
    data = upload.read(IMPORT_MAX_BYTES + 1)
    if len(data) > IMPORT_MAX_BYTES:
        raise UploadTooLarge()
    if not data.strip():
        raise UnreadableUpload("That file is empty. Download the template and fill it in.")
    if extension == "xlsx":
        if not data.startswith(_ZIP_SIGNATURE):
            raise UnreadableUpload()
        records = _workbook_records(data)
    else:
        records = _csv_records(data)
    return _grid(records)


def _grid(records) -> Grid:
    """The heading row (the first with a value) and the data rows after it."""
    headings, rows = None, []
    for number, values in records:
        cleaned = [_clean(value) for value in values[:_MAX_COLUMNS]]
        if not any(value is not None for value in cleaned):
            continue
        if headings is None:
            headings = cleaned
            continue
        if len(rows) >= IMPORT_MAX_ROWS:
            raise UploadTooLarge(
                f"That file has more than {IMPORT_MAX_ROWS} rows, which is more than "
                f"a class has. Upload the file downloaded from this sheet.",
            )
        cleaned += [None] * (len(headings) - len(cleaned))
        rows.append((number, cleaned[: len(headings)]))
    if headings is None:
        raise UnreadableUpload("That file is empty. Download the template and fill it in.")
    return Grid(headings=headings, rows=rows)


def _bounded(records) -> list:
    """The first records of a file, refusing one that runs on past any class's size.

    Blank rows count, so a sheet padded with a million empty rows is refused
    after a bounded read and not walked to its end.
    """
    cap = IMPORT_MAX_ROWS * 3
    found = list(islice(records, cap + 1))
    if len(found) > cap:
        raise UploadTooLarge(
            f"That file has more than {cap} rows, which is more than a class has. "
            f"Upload the file downloaded from this sheet.",
        )
    return found


def _workbook_records(data: bytes):
    """``(row_number, values)`` for the scores sheet of a workbook, or the first sheet.

    A workbook from outside is untrusted, and the library signals a damaged one
    in several ways, so any failure to open or walk it is one refusal.
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = (
                workbook[IMPORT_SCORES_SHEET]
                if IMPORT_SCORES_SHEET in workbook.sheetnames else workbook.worksheets[0]
            )
            return _bounded(
                (number, list(values))
                for number, values in enumerate(sheet.iter_rows(values_only=True), start=1)
            )
        finally:
            workbook.close()
    except Exception as exc:  # noqa: BLE001 - see the docstring
        raise UnreadableUpload() from exc


def _csv_records(data: bytes):
    """``(record_number, values)`` for a CSV, decoded as UTF-8 or the Windows code page."""
    text = None
    for encoding in _CSV_ENCODINGS:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None or "\x00" in text:
        raise UnreadableUpload()
    try:
        return _bounded(enumerate(csv.reader(io.StringIO(text)), start=1))
    except csv.Error as exc:
        raise UnreadableUpload() from exc


# ── The template ───────────────────────────────────────────────────────────

NOTES = (
    "How to fill in this sheet",
    "",
    "Enter each student's mark in the column for that assessment, as a whole number "
    "from 0 up to the maximum shown in the column heading.",
    "Type ABS if the student was absent and EXC if they were excused. These are not marks and count as neither zero nor blank.",
    "Leave a cell empty to leave it as it is on the sheet. An empty cell never clears a mark that is already entered.",
    "Do not change the Admission number column or the column headings. Rows are matched by admission number.",
    "Upload the file on the sheet's Upload screen. You will see what each row would do before anything is saved.",
)


def template_workbook(heading_columns, rows) -> bytes:
    """The template as ``.xlsx`` bytes.

    *heading_columns* are the component headings in order. *rows* is
    ``[(admission_number, name, [cell, ...])]``, a cell being a whole mark, the
    text ``ABS`` or ``EXC``, or None.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    scores = workbook.active
    scores.title = IMPORT_SCORES_SHEET
    bold = Font(bold=True)
    headings = [IMPORT_ADMISSION_HEADING, IMPORT_STUDENT_HEADING, *heading_columns]
    scores.append(headings)
    for cell in scores[1]:
        cell.font = bold
    for admission, name, cells in rows:
        scores.append([admission, name, *cells])
        scores.cell(row=scores.max_row, column=1).number_format = "@"
        scores.cell(row=scores.max_row, column=1).data_type = "s"
    scores.freeze_panes = "C2"
    widths = [len(h) for h in headings]
    for admission, name, cells in rows:
        widths[0] = max(widths[0], len(str(admission)))
        widths[1] = max(widths[1], len(str(name)))
    for index, width in enumerate(widths, start=1):
        scores.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 40)

    notes = workbook.create_sheet(IMPORT_NOTES_SHEET)
    for line in NOTES:
        notes.append([line])
    notes["A1"].font = bold
    notes.column_dimensions["A"].width = 110
    for row in notes.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
