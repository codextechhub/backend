#!/usr/bin/env python3
"""Cut M18 v1.13: the Fake provider gate, webhook provider binding and rate limit,
the stageless payout route, and the payment fields under Field Access.

What changed in the backend, and therefore in the document:

* 786596d5. The in-memory Fake provider exists only where
  PAYMENTS_FAKE_PROVIDER_ENABLED is on (the development, test and CI settings).
  Off, which is what the staging settings Render runs inherit from base, FAKE is
  refused like an unknown provider when a collection, virtual account or payout
  batch is created; its webhook answers 404 before the body is read; a stored
  FAKE event ends FAILED on replay; old FAKE rows still list and read, while
  ?verify=1 and payout confirmation on them are refused. The signing secret
  comes from settings, and an empty one verifies nothing.
* c916e0b6. A webhook event settles only records of the provider that signed
  it, and the public receiver is throttled per client address
  (payments_webhook, 120 a minute) before the signature is read.
* 15d9ac3d. That address is decided once per request by ClientIPMiddleware,
  from Cloudflare's CF-Connecting-IP on Render, never from X-Forwarded-For.
* 3aab9000. New books publish the tenant's payout route with no steps and no
  approver groups; the default ladder is published only on request, and
  migration 0007 removes the groups earlier provisioning created unasked.
* 4767f637 and c5549965. A virtual account's number and name and a payout's
  beneficiary fields are Field Access fields decided by role switches; the
  view_sensitive keys that used to hide them are retired.

* 966e5770 (the finance pass, merged at 32194627). The deposit account a
  collection or virtual account books to, and the source account a payout or
  payout batch is paid from, are named through finance's account resolver
  under the caller's bank reach, and this module's own entity lookup refuses to
  resolve a ledger account at all.

* D62 (0c59e23d, merged 864468ba). A payment request's and a virtual account's
  deposit account follows the customer's branch: a ledger account behind
  another branch's bank account is refused 400, "Deposit it into ...".
  Payouts and payout batches are school-wide and unaffected.
* D63 (112fd7ef, merged 864468ba). Collection create, virtual-account create,
  a single payout and each payout-batch line resolve their customer, invoice
  and vendor within the caller's branches, answering 400 as for an unknown one.
* D67 (6c4853b0 and a992a8e6, merged 0ea6a97c). PaymentsReach is the one way a
  payments view reaches the gateway tables; lists and summaries count only rows
  in reach, a row outside it is a 404 on its detail and actions, and the Export
  Centre's collections and payouts datasets narrow the same way.

The MRD is reconciled in a separate step and is not touched here.

    python tools/patch_payments_provider_gate_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    edit_box,
    insert_row_after,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    normalise_change_log,
    repair_ooxml,
    row,
    set_value,
)

import patch_record_history_docs

REVIEW_DATE = "28 September 2026"
CODE_BASELINE = (
    "Backend main at 7860ea86, 28 September 2026, carrying the payments reach pass (6c4853b0 "
    "and a992a8e6, merged at 0ea6a97c), the branch parties and bank ledgers of the second "
    "finance pass (0c59e23d and 112fd7ef, merged at 864468ba), 966e5770 from the finance pass "
    "(ledger accounts named under the caller's bank reach, merged at 32194627), 3aab9000 (stageless payout "
    "route), 4767f637 and c5549965 (payment fields under Field Access), 786596d5 (Fake "
    "provider gate), c916e0b6 (webhook provider binding and rate limit) and 15d9ac3d "
    "(trusted client address)"
)
TEST_EVIDENCE = (
    "Verified by the vs_payments suite (242 tests) and core.tests_client_ip (12), "
    "each run on its own and all passing; the ledger reach rests on the tests the finance "
    "pass added, read at 6d3251c6 and not re-run here. The branch reach rests on "
    "tests_branch_reach.py in vs_payments and tests_ledger_reach.py in vs_finance, read at "
    "7860ea86 and not re-run here. Backend evidence only; nothing here claims deployment."
)

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

M18_DIR = "18-payments-and-collections"
M18_STEM = "XVS_M18_Payments_and_Collections_Functional_Requirements_Document"
M18_SOURCE, M18_TARGET = "1.12", "1.13"
SOURCE_MRD = "XVS Module Requirements Document v2.93"


# ── editing helpers ──────────────────────────────────────────────────────────


def add_fr_after(doc, anchor: str, heading: str, header_cell: str,
                 rows: list[tuple[str, str]]) -> None:
    """Add a requirement directly after ``anchor``'s table, cloned from it.

    The heading and the table are copies of the anchor's, so the new requirement
    carries the anchor's heading style, status colour and row formatting.
    """
    template = fr_table(doc, anchor)
    heading_p = template._tbl.getprevious()
    while heading_p is not None and heading_p.tag.split("}")[1] != "p":
        heading_p = heading_p.getprevious()
    if heading_p is None or not Paragraph(heading_p, doc).text.strip().startswith(anchor):
        raise ValueError(f"{anchor} has no heading directly above its table")
    new_heading = copy.deepcopy(heading_p)
    new_table = copy.deepcopy(template._tbl)
    template._tbl.addnext(new_heading)
    new_heading.addnext(new_table)
    mrd_tools.retitle(Paragraph(new_heading, doc), heading)
    table = next(t for t in doc.tables if t._tbl is new_table)
    if len(table.rows) - 1 != len(rows):
        raise ValueError(f"{anchor} has {len(table.rows) - 1} rows, {len(rows)} given")
    for cell in table.rows[0].cells:
        keep_format(cell, header_cell)
    for r, (label, value) in zip(table.rows[1:], rows):
        if r.cells[0].text.strip() != label:
            raise ValueError(f"Row {r.cells[0].text!r} is not {label!r}")
        keep_format(r.cells[-1], value)


def retitle_heading(doc, start: str, text: str) -> None:
    hits = [p for p in doc.paragraphs if p.text.strip().startswith(start)]
    if len(hits) != 1:
        raise ValueError(f"{start!r} starts {len(hits)} paragraphs")
    mrd_tools.retitle(hits[0], text)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
#: CT_TrPr children that must follow cantSplit in the transitional schema.
_AFTER_CANT_SPLIT = {"trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden", "ins",
                     "del", "trPrChange"}


def keep_rows_together(table) -> None:
    """Stop each row of ``table`` breaking across a page, in schema order.

    This document sets no row to stay whole, so a row added near the foot of a
    page is drawn half on one page and half on the next.
    """
    from docx.oxml import OxmlElement

    for r in table.rows:
        tr_pr = r._tr.get_or_add_trPr()
        if tr_pr.find(_W + "cantSplit") is not None:
            continue
        element = OxmlElement("w:cantSplit")
        later = next((c for c in tr_pr if c.tag.replace(_W, "") in _AFTER_CANT_SPLIT), None)
        if later is None:
            tr_pr.append(element)
        else:
            later.addprevious(element)


def keep_requirement_together(table) -> None:
    """A requirement never leaves its heading and first line, or its last row, alone on a page.

    The status bar and the Requirement row keep with the evidence below them, the
    second-last row keeps with the last, and a short requirement moves whole.
    """
    short = sum(len(c.text) for r in table.rows for c in r.cells) < 1500
    rows = table.rows[:-1] if short else list(table.rows[:2]) + list(table.rows[-2:-1])
    for table_row in rows:
        for cell in table_row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def box(doc, heading: str):
    hits = [t for t in doc.tables
            if t.rows[0].cells[0].paragraphs[0].text.strip() == heading]
    if len(hits) != 1:
        raise ValueError(f"{heading!r} heads {len(hits)} boxes")
    return hits[0].rows[0].cells[0]


# ── the revision ─────────────────────────────────────────────────────────────

FR023 = [
    ("Requirement",
     "The in-memory Fake provider moves no money, so it must not exist where real money "
     "moves. No collection, virtual account or payout may be started against it there, and "
     "no event signed with its secret may be stored or booked."),
    ("Current evidence",
     "PAYMENTS_FAKE_PROVIDER_ENABLED decides whether Fake exists at all. The development, "
     "test and CI settings turn it on; the base settings turn it off and do not read it from "
     "the environment, and the staging settings Render runs inherit that. The provider "
     "registry is the one gate every path resolves through: with the setting off, "
     "available_providers omits FAKE and get_provider refuses it exactly as it refuses a "
     "name that was never registered, even when a test override is registered under it. "
     "Creating a collection, a virtual account or a payout batch resolves the provider "
     "through resolve_provider_name before any row is written. The webhook receiver answers "
     "404 for FAKE before reading the body, and processing a stored FAKE event marks it "
     "FAILED with the reason. The signing secret comes from PAYMENTS_FAKE_WEBHOOK_SECRET: an "
     "enabled Fake with no secret is not configured, and an empty secret verifies no "
     "signature, because an HMAC keyed with the empty string is one anybody can compute. "
     "The shared finance screens (FinPro v0.7.24, pinned by the school app and the console) "
     "offer Fake only in a development build."),
    ("Acceptance",
     "With the setting off: naming FAKE when creating a collection, virtual account or "
     "payout batch is a 400 on provider that names the providers offered, and nothing is "
     "written; a configured default provider that is not offered is 503 "
     "PAYMENT_PROVIDER_NOT_CONFIGURED, a configuration fault rather than the caller's; a "
     "forged Fake webhook is 404 and stores, books and audits nothing; a stored FAKE event "
     "ends FAILED on replay and books nothing; Fake rows already on file still list and "
     "read; ?verify=1 on a Fake collection is 503 and confirming a Fake payout is refused, "
     "each booking nothing and leaving the row's status as it was. With the setting on, a "
     "Fake collection is created and a signed webhook books it, and an unknown provider's "
     "webhook is still 404. FakeProviderDisabledTests, FakeProviderSecretTests and "
     "FakeProviderEnabledTests cover each case."),
    ("Current limit",
     "Rows recorded against Fake before a deployment switched it off stay in the state they "
     "reached. See Section 9."),
]

FR024 = [
    ("Requirement",
     "A valid signature proves which provider sent an event and nothing more. An event must "
     "never link, re-check or book a collection or payout that belongs to another provider."),
    ("Current evidence",
     "_find_record, which every lookup passes through on ingestion and on processing, "
     "refuses a record whose provider is not the one that signed the event. The platform's "
     "own reference is unique across providers, so it is looked up across all of them and a "
     "hit on another provider's record is refused. The provider's reference can repeat "
     "across providers, so it is looked up within the signing provider only. A "
     "virtual-account deposit resolves the account within the signing provider, and a "
     "reference that collides with another provider's intent is refused. A refused event is "
     "marked FAILED with the reason and linked to nothing, so its received audit row names "
     "no entity and it is listed on the platform view of unattributed events (FR-010), not "
     "on any tenant's."),
    ("Acceptance",
     "An event from one provider naming another provider's collection or payout ends "
     "FAILED, and nothing is linked, re-checked or booked; replaying it refuses it again. "
     "The same event from the record's own provider books. A provider reference matches "
     "only within its own provider, and a deposit into a Fake account cannot claim a "
     "Paystack reference. WebhookProviderBindingTests covers each case."),
    ("Current limit", "None."),
]

FR025 = [
    ("Requirement",
     "A payments write must name its customer, invoice and payout vendor only from the "
     "caller's branches and the school-wide rows, and a branch's payment request or virtual "
     "account must deposit only into that branch's bank or a school-wide one."),
    ("Current evidence",
     "Collection create resolves its customer and its invoice, virtual-account create its "
     "customer, and a single payout and every payout-batch line their vendor, within the "
     "entity and the caller's branches plus the school-wide rows, the reading finance's own "
     "resolvers use (_entity_obj and _payout_vendor). A row another branch keeps answers "
     "exactly as one that does not exist: 400 \"No customer '<ref>' in this entity.\", \"No "
     "invoice '<ref>' in this entity.\", or \"No such vendor in this entity.\", on "
     "items[n].vendor for a batch line. The deposit account then obeys the document's branch, "
     "Module 19's same-branch rule (its FR-020): a payment request and a virtual account take "
     "the branch of the customer they name, which the receipt they produce carries, so a "
     "ledger account behind another branch's bank account is refused 400 on deposit_account, "
     "naming the branch: \"This payment request belongs to Ikeja Branch. Deposit it into an "
     "Ikeja Branch account or a school-wide one.\", and \"This virtual account belongs to "
     "...\" for an account. A school-wide customer may deposit into any account in reach, and "
     "a ledger account behind no bank account is unaffected. A payout and a payout batch pay "
     "through a vendor payment that carries no branch, so they are school-wide documents and "
     "their source account may be any account in reach."),
    ("Acceptance",
     "Ikeja's clerk naming a Lekki customer, a Lekki invoice or a vendor Lekki keeps to itself "
     "gets the 400 of an unknown one and nothing is written: "
     "PaymentsNameOnlyWhatTheClerkReachesTests in tests_branch_reach.py covers the payment "
     "request's customer and invoice, the virtual account's customer, the payout's vendor and "
     "the batch line's vendor. Mrs Okafor, who covers Ikeja and Lekki, cannot deposit an Ikeja "
     "family's payment request or virtual account into Lekki's bank ledger "
     "(test_a_payment_request_and_a_virtual_account_deposited_into_it in Module 19's "
     "tests_ledger_reach.py)."),
    ("Current limit",
     "A payment request that names an invoice and no customer is held to no branch for its "
     "deposit account. See Section 9."),
]

FR026 = [
    ("Requirement",
     "Inside the entity, a caller must read, count and act on only the gateway records of "
     "their own branches and the school-wide ones."),
    ("Current evidence",
     "No gateway table carries a branch. Each record takes its reach from the finance or "
     "procurement row it hangs on, read inclusively as finance reads it, the caller's branches "
     "plus the school-wide rows, and PaymentsReach (vs_payments/reach.py) is the one way a "
     "payments view reaches these tables: a collection by its customer and by the invoice it "
     "settles, either of which may be absent and then restricts nothing; a virtual account by "
     "its customer; a payout by the vendor it pays, a payout naming no vendor being "
     "school-wide; a payout batch by every line in it, so one line paying a vendor outside "
     "reach withholds the whole batch; a transactions-log entry by the record its reference "
     "names, and a virtual account's entries by metadata.virtual_account_id as well; a webhook "
     "event by the collection or payout it matched; and a settlement-reconciliation bank line "
     "by its bank account. The collection, virtual-account, payout and batch lists, the "
     "transactions log, the movements feed, the webhook queue, settlement reconciliation and "
     "every summary hold and count only rows in reach, so a branch reader sees smaller totals "
     "than a whole-school reader. A caller who covers the whole school is not narrowed and "
     "gets the same querysets as the entity filter alone. The staff-only unattributed-webhooks "
     "view reads outside it, because its events belong to no entity (FR-010). The Export "
     "Centre's collections and payouts datasets narrow by the same rule (Module 26)."),
    ("Acceptance",
     "A row outside reach is absent from its list, uncounted in its summary, and answers 404 "
     "as an unknown one does on collection detail, ?verify=1 included (\"No such collection "
     "in this entity.\"), virtual-account detail and status change (\"No virtual account "
     "matches this id for the entity.\"), payout-batch detail, its direct-submit refusal and "
     "submit-for-approval (\"No such payout batch in this entity.\"), and webhook replay "
     "(\"Webhook event not found for this entity.\"); nothing is verified, changed or "
     "submitted. PaymentsShowOnlyWhatTheClerkReachesTests in tests_branch_reach.py proves "
     "each list, summary, settlement report, detail, status change and export holds to the "
     "reach, and PaymentsViewsStartFromTheReachTests fails if views.py reads a gateway table "
     "other than through PaymentsReach."),
    ("Current limit",
     "A transactions-log entry whose reference names no gateway record stays visible to every "
     "branch: a virtual account's creation logged before entries carried "
     "metadata.virtual_account_id, whose reference is the one-off request sent to the "
     "provider, and a rejected initiation that never wrote a record. See Section 9."),
]

SUMMARY = (
    "Minor revision. Records nine backend changes. (1) The in-memory Fake provider exists "
    "only where PAYMENTS_FAKE_PROVIDER_ENABLED is on, which the development, test and CI "
    "settings set and the staging settings Render runs do not. Off, FAKE is refused like an "
    "unknown provider when a collection, virtual account or payout batch is created, its "
    "webhook is 404 before the body is read, a stored FAKE event ends FAILED on replay, "
    "and old FAKE rows still list and read while verifying or confirming them is refused. "
    "The signing secret comes from settings and an empty one verifies nothing. New FR-023. "
    "(2) A webhook event settles only records of the provider that signed it; a "
    "cross-provider event ends FAILED with the reason and nothing is linked, re-checked or "
    "booked. New FR-024. (3) The public receiver is throttled per client address "
    "(payments_webhook, 120 a minute) before the signature is read, and that address is "
    "Cloudflare's CF-Connecting-IP on Render, never X-Forwarded-For, which the caller "
    "writes. FR-005 records both. (4) New books publish the tenant's payout route with no "
    "steps and no approver groups, because who approves is the tenant's own answer; the "
    "default two-stage ladder is published only on request by seed_payout_approvals, and "
    "migration 0007 removes the unused groups earlier provisioning created. FR-015, "
    "FR-020, the actor table, the dependencies and the fail-closed box follow. (5) A "
    "virtual account's number and name and a payout's beneficiary fields are Field Access "
    "fields, decided by role switches and absent for a caller who cannot read them; the "
    "view_sensitive keys that hid them are retired. The sensitive-field box, FR-003, "
    "FR-012, FR-014 and FR-017 follow. Section 5.3, the webhook route, the typed errors and "
    "traceability are reconciled, and Section 9 records the Fake rows left from before the "
    "gate. (6) The deposit account of a collection or virtual account and the source account "
    "of a payout or payout batch are named through finance's resolver under the caller's "
    "bank reach, so another branch's bank ledger is refused as an unknown account, and this "
    "module's own entity lookup no longer resolves a ledger account (966e5770, the finance "
    "pass; FR-019 and the typed errors). (7) The deposit account of a payment request or a "
    "virtual account follows the customer's branch, so a ledger account behind another "
    "branch's bank account is refused 400 naming the branch, \"Deposit it into\" that "
    "branch's account or a school-wide one, while payouts and payout batches stay school-wide "
    "(0c59e23d). (8) Collection create, virtual-account create, a single payout and each "
    "payout-batch line find the customer, invoice and vendor they name only within the "
    "caller's branches and the school-wide rows, and answer 400 for any other as for an "
    "unknown one (112fd7ef). New FR-025. (9) Every payments read and status change reaches "
    "the gateway tables through PaymentsReach: each record takes its reach from the customer, "
    "invoice, vendor, matched record or bank account it hangs on; lists and summaries count "
    "only rows in reach, so a branch reader sees smaller totals; a row outside it answers 404 "
    "on collection detail, virtual-account detail and status change, batch detail and submit, "
    "and webhook replay; and the Export Centre's collections and payouts datasets narrow the "
    "same way, a scope with no caller narrowing nothing (6c4853b0, a992a8e6). New FR-026. "
    "FR-008, FR-016, FR-017, FR-019, the actor table, the data model, the typed errors and "
    "the dependencies follow, and Section 9 records the log entries no branch can be traced "
    "for and a payment request naming only an invoice, whose deposit account is held to no "
    "branch. Traceability is retraced against MRD v2.93, whose 24 "
    "Module 18 entries are unchanged. " + TEST_EVIDENCE
)


def patch_branch_reach(doc) -> None:
    """D62, D63 and D67: the parties, deposit accounts and records a caller reaches.

    A write names its customer, invoice, vendor and deposit account within the
    caller's branches (FR-025), and every read and status change reaches only the
    gateway records of those branches and the school-wide ones (FR-026).
    """
    add_fr_after(doc, "FR-019",
                 "FR-025 Name Parties and Deposit Accounts Within the Caller's Branches",
                 "FR-025 | Implemented", FR025)
    add_fr_after(doc, "FR-025",
                 "FR-026 Read and Act Only on the Caller's Branches' Records",
                 "FR-026 | Implemented", FR026)

    actors = table_headed(doc, "Actor", "May do", "Governed by")
    edit_cell(row(actors, "Finance operator").cells[1],
              "and the movement feeds.",
              "and the movement feeds, of their own branches and the school-wide records "
              "(FR-026).")

    append_to(fr_table(doc, "FR-008"), "Acceptance",
              " Inside the entity it is listed only to a caller who reaches the collection or "
              "payout it matched (FR-026).")
    append_to(fr_table(doc, "FR-016"), "Current evidence",
              " Both sides are narrowed to the caller's branches: the gateway records in reach, "
              "and the statement lines of the bank accounts in reach (FR-026).")
    edit_value(fr_table(doc, "FR-017"), "Current evidence",
               "are entity-scoped read endpoints over the gateway records.",
               "are entity-scoped read endpoints over the gateway records, narrowed to the "
               "caller's branches (FR-026).")
    append_to(fr_table(doc, "FR-019"), "Current evidence",
              " Inside the entity the parties a write names and the records a caller reads "
              "answer to the caller's branches as well (FR-025, FR-026).")

    model = table_headed(doc, "Model", "Holds", "Links to")
    keep_format(row(model, "PaymentEvent").cells[-1],
                "The gateway record it concerns, by that record's reference. A virtual "
                "account's entries also carry metadata.virtual_account_id.")

    errors = table_headed(doc, "Condition", "Answer")
    anchor = row(errors, "A deposit or source account that backs a bank account outside the "
                         "caller's branches")
    for values in [
        ["A customer, invoice or payout vendor outside the caller's branches",
         "400 \"No customer '<ref>' in this entity.\", \"No invoice '<ref>' in this entity.\" "
         "or \"No such vendor in this entity.\", as for an unknown one; nothing is written."],
        ["A payment request or virtual account depositing into a ledger account behind "
         "another branch's bank account",
         "400 on deposit_account, naming the branch: \"This payment request belongs to Ikeja "
         "Branch. Deposit it into an Ikeja Branch account or a school-wide one.\"; nothing is "
         "written."],
        ["Reading or acting on a collection, virtual account, payout batch or webhook event "
         "outside the caller's branches",
         "404 with the route's own message, as for an unknown id; nothing is verified, changed "
         "or submitted."],
    ]:
        anchor = insert_row_after(anchor, values)

    deps = table_headed(doc, "Dependency", "Contract")
    append_to(deps, "Module 19, Finance & Accounting",
              " Its account resolver holds a payment request's and a virtual account's deposit "
              "account to the customer's branch (FR-025).")
    insert_row_after(row(deps, "Module 5, Audit"), [
        "Module 26, Reporting & Exports",
        "Publishes the gateway collections and payout instructions datasets this module "
        "registers. Each narrows to the exporting caller's branches by the rule of FR-026, and "
        "a scope with no caller, a system-triggered estimate or the catalogue's own checks, "
        "narrows nothing.",
    ])

    edit_box(box(doc, "FURTHER GAPS"), [
        ("append",
         "• A transactions-log entry whose reference names no gateway record stays visible to "
         "every branch. A virtual account's creation logged before entries carried "
         "metadata.virtual_account_id names only the one-off request sent to the provider, so "
         "Ikeja's clerk still reads the line recording a Lekki family's account; a rejected "
         "initiation that never wrote a record reads the same way (FR-026)."),
        ("append",
         "• A payment request that names an invoice and no customer is held to no branch for "
         "its deposit account. The route reads the branch from the customer it was given, finds "
         "none and treats the request as school-wide, and the service then takes the customer "
         "from the invoice, whose branch the receipt books to. Mrs Okafor, who covers Ikeja and "
         "Lekki, can raise a request against a Lekki invoice alone and have it deposit into "
         "Ikeja's bank ledger (FR-025)."),
    ])

    trace = table_headed(doc, "MRD capability", "Requirements")
    for capability, requirements in [
        ("Payment collection intents and status", "FR-001, FR-002, FR-025, FR-026"),
        ("Virtual-account operations", "FR-003, FR-025, FR-026"),
        ("Payout creation and lifecycle", "FR-012, FR-025, FR-026"),
        ("Payout batches", "FR-013, FR-025, FR-026"),
        ("Settlement import and reconciliation", "FR-016, FR-026"),
        ("Transaction log", "FR-017, FR-026"),
        ("Payment movements and movement summary", "FR-017, FR-026"),
        ("Failed-webhook operator queue and replay", "FR-008, FR-009, FR-026"),
    ]:
        set_value(trace, capability, requirements)


def patch_m18() -> None:
    require_newest(str(ROOT / "functional-requirements" / M18_DIR / f"{M18_STEM}_v*.docx"),
                   M18_SOURCE)
    doc = Document(str(frd_path(M18_DIR, M18_STEM, M18_SOURCE)))
    set_cover_version(doc, M18_SOURCE, M18_TARGET)
    set_control(doc, "Version", M18_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)

    # 1.1 In scope
    scope = table_headed(doc, "Area", "Responsibility")
    append_to(scope, "Webhook ingestion",
              " An event settles only records of the provider that signed it, and the public "
              "receiver is rate limited per client address.")

    # 3. Actors
    actors = table_headed(doc, "Actor", "May do", "Governed by")
    edit_cell(row(actors, "Invoice payer").cells[-1],
              "plus a rate limit.", "plus a rate limit per client address.")
    keep_format(
        row(actors, "Payout checker and senior checker").cells[-1],
        "Whoever the stages the tenant publishes name: role holders, an approver group it "
        "composes, a dynamic rule, or the requester's organogram. New books name nobody. The "
        "default ladder, published on request, names the payout-approver and "
        "payout-senior-approver groups, created empty. No permission key confers approval.")
    keep_format(
        row(actors, "Provider").cells[-1],
        "Signature verification against the provider secret, for a provider this deployment "
        "offers, behind a rate limit per client address. An event reaches only its own "
        "provider's records.")
    edit_box(box(doc, "SENSITIVE FIELD ACCESS"), [
        ("replace", "• Beneficiary name",
         "• A virtual account's number and name, and a payout's beneficiary name, account "
         "number and bank code, are Field Access fields: a caller reads each only where their "
         "roles switch it on, and a field they cannot read is absent from the response, with "
         "no placeholder. None is writable: the provider issues the account, and the "
         "beneficiary is copied from the verified vendor record."),
        ("after", "• A virtual account's number",
         "• The movements feed drops a payout row's party and beneficiary_account for a caller "
         "who cannot read them, and the payout approval document drops its Beneficiary and "
         "Account columns for such an approver. The account there is cut to its last four "
         "digits whoever reads it."),
    ])

    # 4. Functional requirements
    edit_value(fr_table(doc, "FR-003"), "Acceptance",
               "Provider identifiers are treated as sensitive and withheld from callers "
               "without the sensitive key.",
               "The account number and account name are Field Access fields, absent for a "
               "caller whose roles cannot read them.")

    fr005 = fr_table(doc, "FR-005")
    edit_value(fr005, "Current evidence",
               "ingest_webhook verifies",
               "The receiver answers 404 before reading the body when the path names a "
               "provider this deployment does not offer (FR-023). It is throttled per client "
               "address under the payments_webhook scope, 120 requests a minute, and the "
               "throttle runs before the signature is read. That address is decided once per "
               "request by ClientIPMiddleware: on Render, which Cloudflare fronts, it is "
               "CF-Connecting-IP, then True-Client-IP, both of which Cloudflare overwrites; "
               "elsewhere it is the connecting peer. X-Forwarded-For, which the caller writes, "
               "is never read. ingest_webhook then verifies")
    append_to(fr005, "Acceptance",
              " A request past the limit is answered 429 and stores and audits nothing, so a "
              "flood of badly signed bodies cannot fill the audit log, and a new "
              "X-Forwarded-For on every request buys no new allowance "
              "(WebhookReceiverThrottleTests).")
    set_value(fr005, "Current limit",
              "The allowance is one fixed rate per client address, shared by every provider. "
              "A provider delivering more than 120 events a minute from one address has the "
              "excess refused with 429 and relies on its own retries.")
    add_fr_after(doc, "FR-005",
                 "FR-023 Offer the Fake Provider Only Where a Deployment Switches It On",
                 "FR-023 | Implemented", FR023)

    add_fr_after(doc, "FR-007", "FR-024 Settle Only the Signing Provider's Records",
                 "FR-024 | Implemented", FR024)

    append_to(fr_table(doc, "FR-008"), "Current evidence",
              " An event whose provider this deployment no longer offers, and one naming "
              "another provider's record, are marked FAILED with the reason in the same way "
              "(FR-023, FR-024).")
    append_to(fr_table(doc, "FR-010"), "Current evidence",
              " An event refused for naming another provider's record is linked to nothing "
              "and is listed here (FR-024).")

    edit_value(fr_table(doc, "FR-012"), "Acceptance",
               "and protected bank fields remain withheld.",
               "and the beneficiary's bank fields are absent for a caller whose roles cannot "
               "read them.")
    append_to(fr_table(doc, "FR-014"), "Acceptance",
              " An approver whose roles cannot read the beneficiary fields is shown the batch "
              "without those columns.")

    retitle_heading(doc, "FR-015 Seed the Two-Stage Payout Ladder Blocked",
                    "FR-015 Publish the Two-Stage Payout Ladder Blocked, on Request")
    fr015 = fr_table(doc, "FR-015")
    set_value(fr015, "Requirement",
              "The default ladder must arrive blocked rather than self-skipping, and only "
              "when a tenant asks for it.")
    set_value(fr015, "Current evidence",
              "New books publish a route with no steps (FR-020). seed_payout_approvals "
              "publishes the default ladder for a tenant that asks: an always-on payout "
              "checker and a conditional senior checker at 50,000,000 kobo, both with "
              "skip_if_no_approvers=False, naming the payout-approver and "
              "payout-senior-approver approver groups, which it creates empty. It fills in the "
              "empty route provisioning published, and leaves a route that already carries a "
              "live step exactly as it is. Migration 0007 removes the two groups, and the "
              "steps naming them, that earlier provisioning created unasked; it skips a group "
              "with members, one whose steps a batch has run through, and one something is "
              "waiting on.")
    set_value(fr015, "Acceptance",
              "An unstaffed applicable stage parks. Adding somebody to the approver group a "
              "stage names makes it actionable. Seeding again, or over a tenant's own ladder, "
              "leaves every step as it was, and filling the empty route reports it as "
              "created. Payout handlers advertise that continue-without-approval is "
              "unavailable and reject the release endpoint with a typed 409. "
              "tests_stageless_provisioning covers provisioning, seeding and each of the "
              "migration's skips.")

    fr020 = fr_table(doc, "FR-020")
    set_value(fr020, "Current evidence",
              "New entity provisioning publishes the tenant's own payout route carrying no "
              "steps, and creates no approver group. Who approves money leaving a tenant is "
              "that tenant's own answer, read from the organogram it builds; a ladder "
              "published at creation would be a guess at the people and at the amount that "
              "needs a second pair of eyes. The empty route still stands in front of the "
              "shared platform row, so a change to that row can never begin governing the "
              "tenant's cash-out. Migration 0006 published a platform fallback for entities "
              "that predate the provisioner; the corrective vs_rbac migration then removed "
              "its stages, because a shared row cannot name a tenant's approver group. The "
              "platform row itself remains, so no entity is unroutable. payout_approval_health "
              "asks the same workflow resolver used by submission and reports every active "
              "entity that still has no route.")
    set_value(fr020, "Acceptance",
              "After migration, every active ledger entity resolves either its tenant route "
              "or the platform row. Provisioning skips a tenant that already has a route, so "
              "a second entity in the same tenant, or a tenant with a ladder of its own, keeps "
              "every step. Re-running migration 0006 is idempotent, and a platform policy an "
              "administrator owns is left alone. The release check exits successfully only "
              "when no active entity is uncovered.")
    set_value(fr020, "Current limit",
              "Resolving is not approving. Neither the empty tenant route nor the platform row "
              "carries steps, so a batch from an entity with no stages of its own is refused "
              "with 409 APPROVAL_NOT_CONFIGURED until the tenant publishes some, through the "
              "workflow template screens or seed_payout_approvals. The payout routes accept "
              "confirm_without_approval with a reason, as every submit route does, but a "
              "confirmation cannot release a payout: it ends the approval with no human vote, "
              "and the approval-time check that requires one distinct checker, or two at the "
              "high-value threshold, then refuses the batch with PAYOUT_APPROVAL_REQUIRED, so "
              "nothing is sent. That outcome is read from the code; no test submits a "
              "confirmed batch against a route with no steps. A published stage supplies no "
              "people: it stays blocked until somebody it names can act.")

    append_to(fr_table(doc, "FR-017"), "Acceptance",
              " A payout row's beneficiary name and account are absent, not masked, for a "
              "caller whose roles cannot read them.")

    # 5.3 Webhook lifecycle
    lifecycle = table_headed(doc, "Step", "Rule")
    set_value(lifecycle, "1",
              "The path must name a provider this deployment offers, or the answer is 404 "
              "before the body is read. Each client address may send 120 requests a minute; "
              "past that the answer is 429 before the signature is read, and nothing is "
              "stored or audited. The raw body's signature is then verified against the "
              "provider secret. A failure is a 401 and nothing is stored as actionable.")
    set_value(lifecycle, "4",
              "The worker matches the event only to a record of the provider that signed it; "
              "a record of another provider ends the event FAILED with the reason. It then "
              "re-verifies with the provider, and books through the ledger under a row lock "
              "with a terminal-state short-circuit.")

    # 7.3 and 7.4 API contracts
    route = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Method and path"
             and any(r.cells[0].text.strip() == "POST /webhooks/{provider}/" for r in t.rows)]
    if len(route) != 1:
        raise ValueError("The webhook route table is not unique")
    set_value(route[0], "POST /webhooks/{provider}/",
              "The provider receiver. Rate limited per client address (payments_webhook, 120 "
              "a minute); a provider this deployment does not offer is 404. "
              "Signature-verified, deduplicated, stored, acknowledged.")

    errors = table_headed(doc, "Condition", "Answer")
    anchor = row(errors, "Unverified or wrongly signed webhook body")
    for values in [
        ["Webhook path naming a provider this deployment does not offer (an unknown name, "
         "or FAKE where the Fake provider is off)",
         "404 before the body is read; nothing is verified, stored or audited."],
        ["More than 120 webhook requests in a minute from one client address",
         "429 before the signature is read; nothing is stored or audited."],
        ["Webhook event naming another provider's collection or payout",
         "Acknowledged, then marked FAILED with the reason; nothing is linked, re-checked or "
         "booked."],
        ["Creating a collection, virtual account or payout batch with a provider this "
         "deployment does not offer",
         "400 on provider, naming the providers offered; nothing is written. A configured "
         "default that is not offered is 503 PAYMENT_PROVIDER_NOT_CONFIGURED."],
        ["Verifying a collection or confirming a payout recorded against a provider this "
         "deployment no longer offers",
         "503 PAYMENT_PROVIDER_NOT_CONFIGURED; nothing is booked and the row keeps its "
         "status."],
    ]:
        anchor = insert_row_after(anchor, values)

    # 8. Dependencies
    deps = table_headed(doc, "Dependency", "Contract")
    set_value(deps, "Module 4, Roles & Permissions",
              "Supplies the permission keys the views enforce, and the Field Access switches "
              "that decide who reads a virtual account's number and name and a payout's "
              "beneficiary fields. It supplies no approving role and no approve key: who may "
              "approve a payout is decided by the workflow stages the tenant publishes.")
    set_value(deps, "Providers",
              "Paystack, plus an in-memory Fake provider that moves no money. Fake exists only "
              "where PAYMENTS_FAKE_PROVIDER_ENABLED is on: the development, test and CI "
              "settings turn it on, and the staging settings Render runs inherit the base "
              "settings' off, so there it is refused as an unknown provider (FR-023). Its "
              "signing secret comes from PAYMENTS_FAKE_WEBHOOK_SECRET. OPay was removed; the "
              "provider enumeration carries exactly those two, and a registry override seam "
              "lets a test or an operator substitute a client, though never for a provider "
              "the deployment does not offer.")
    set_value(deps, "Seeding",
              "seed_payments_permissions registers the keys. New books publish a tenant payout "
              "route with no steps; seed_payout_approvals publishes the default two-stage "
              "ladder and its empty groups for a tenant that asks; migration 0006 supplies the "
              "non-destructive platform fallback for older entities, and migration 0007 "
              "removes the unused groups earlier provisioning created. payout_approval_health "
              "is the release gate.")
    insert_row_after(row(deps, "Public application"), [
        "Client address",
        "The webhook rate limit keys on the address ClientIPMiddleware decides once per "
        "request, first in the middleware. On Render, which Cloudflare fronts, the staging "
        "settings name CF-Connecting-IP, then True-Client-IP, both of which Cloudflare "
        "overwrites with the connecting address; a value that is not an IP address is "
        "ignored. Elsewhere the connecting peer stands. X-Forwarded-For is never read, and "
        "the throttles trust no forwarding proxy.",
    ])

    # 9. Needs attention
    edit_box(box(doc, "CURRENT PAYOUT CONTROL - FAIL CLOSED"), [
        ("replace", "• New books receive tenant policy",
         "• New books receive their own payout route with no steps and no approver groups, "
         "and the shared platform row carries no steps either, so an entity that has "
         "published no ladder of its own is refused rather than approved, a confirmation "
         "cannot stand in for the missing human vote, and payouts forbid continuing without "
         "approval outright."),
    ])
    edit_box(box(doc, "FURTHER GAPS"), [
        ("append",
         "• Collections, virtual accounts and payouts recorded against the Fake provider "
         "before a deployment switched it off stay as they were. They list and read, but "
         "verifying a collection and confirming a payout are refused and no route closes "
         "either, so an open one stays open. A Fake virtual account can still be "
         "deactivated."),
    ])

    # FR-019 and the refusals: ledger accounts are named under the caller's bank reach.
    fr019 = fr_table(doc, "FR-019")
    append_to(fr019, "Current evidence",
              " Inside the entity, the ledger accounts a write names answer to the caller's "
              "branches. The deposit account a collection or a virtual account books to and "
              "the source account a payout or a payout batch is paid from are resolved "
              "through Module 19's account resolver, which refuses a ledger account behind "
              "another branch's bank account exactly as it refuses an unknown one; this "
              "module's own entity lookup raises rather than resolve a ledger account, so no "
              "route here can name one around that rule.")
    append_to(fr019, "Acceptance",
              " test_a_payout_sourced_from_it and "
              "test_the_payments_lookup_refuses_to_resolve_a_ledger_account in Module 19's "
              "tests_ledger_reach.py prove another branch's bank ledger is refused as a "
              "payout's source and that the payments lookup will not resolve a ledger "
              "account.")
    insert_row_after(row(errors, "Caller lacks the view or action key"), [
        "A deposit or source account that backs a bank account outside the caller's branches",
        "400 \"No account '<ref>' in this entity.\", as for an unknown account; nothing is "
        "written.",
    ])

    # 10. Traceability, retraced against the current MRD: its 24 Module 18 entries are unchanged.
    set_control(doc, "Source MRD", SOURCE_MRD)
    edit_paragraph(doc, "Module 18 carries 24 capability entries", "MRD v2.76", "MRD v2.93")
    trace = table_headed(doc, "MRD capability", "Requirements")
    set_value(trace, "Provider registry with per-provider overrides", "FR-001, FR-007, FR-023")
    set_value(trace, "Fake provider for controlled testing", "FR-005, FR-023")
    set_value(trace, "Signed and idempotent webhook ingestion", "FR-005, FR-024")
    set_value(trace, "Provider-side webhook re-verification", "FR-007, FR-024")
    # Section 10 starts a page, so its heading never sits apart from its table.
    heading = [p for p in doc.paragraphs if p.text.strip() == "10. MRD Traceability"]
    if len(heading) != 1:
        raise ValueError("Section 10's heading is not unique")
    heading[0].paragraph_format.page_break_before = True

    patch_branch_reach(doc)

    log_change(doc, M18_TARGET, SUMMARY)

    # Every table this revision changed keeps its rows whole.
    for table in (scope, actors, errors, deps, lifecycle, route[0], trace,
                  fr005, fr_table(doc, "FR-023"), fr_table(doc, "FR-024"), fr015, fr020,
                  patch_record_history_docs.change_log_table(doc)):
        keep_rows_together(table)
    # The sensitive-field box sits against the actor table, so the two are laid out as
    # one; a keep-with-next inside it ties both to the first requirement and pushes the
    # box alone onto the next page under a repeated actor header.
    for paragraph in box(doc, "SENSITIVE FIELD ACCESS").paragraphs:
        paragraph.paragraph_format.keep_with_next = None
    # The dependency table never ends on a lone row under a repeated header.
    for table in (deps,):
        for table_row in table.rows[-3:-1]:
            for cell in table_row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    # Every requirement keeps its rows whole too, so none is drawn across a page.
    for table in doc.tables:
        if table.rows[0].cells[0].text.strip().startswith("FR-"):
            keep_rows_together(table)
            keep_requirement_together(table)
    assert_absent_outside_log(doc, "view_sensitive", "sensitive key", "MRD v2.76",
                              "Seed the Two-Stage Payout Ladder Blocked",
                              "whose stages name approver groups created empty",
                              "creates the payout-approver and payout-senior-approver")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M18_DIR, M18_STEM, M18_TARGET),
           f"{M18_STEM.replace('_', ' ')} v{M18_TARGET}", M18_TARGET)


def main() -> None:
    patch_m18()


if __name__ == "__main__":
    main()
