#!/usr/bin/env python3
"""Version the Module 11, 13 and 14 FRDs against the code at commit 13675db.

Module 11, Student Management, v2.8 to v2.9. Version 2.8 was a narrow key
revision and version 2.7 a photograph revision, so neither read the module
whole. This one does: a guardians import (FR-024) and a guardian palette search
(FR-025) are new; the shared people matcher reaches the student and guardian
directories and not the student palette, which is recorded as a gap; promotion
moves to school.students.promote in the requirement, its acceptance and the API
table, holds an unwired level under LEVEL_NOT_WIRED instead of graduating it,
narrows to one branch when asked, and reads a batch under the inclusive branch
rule; the student import refuses what the enrol form refuses; the year lens,
the seat counts route and the next-number suggestion are recorded; and the
dependencies that have since been built are marked discharged.

Module 13, Academic Structure, v2.9 to v2.10. SchoolClass carries a class
teacher, read on every class row and written by Module 12; the academic
structure and subjects imports are FR-017; decision 10 closes on the platform's
band map; and the version 2.9 change-log row, which was written into the
section 6.8 table, is moved into the change log.

Module 14, Academic Calendar and Timetables, v3.2 to v3.3. The exam keys reach
FR-016, FR-017 and the API table; an exam period must be one the caller may see;
the clash previews, the calendar import (FR-018) and the branch lens (FR-019)
are recorded; five refusals gain codes of their own; three statements version
3.1 left describing user_type are corrected; decision 11 closes.

Earlier patch scripts rewrote some cells through ``write_cell``, which sets
Aptos Narrow, in documents set in Calibri. The helpers here keep each cell's
own run formatting, clone only Calibri rows, and restore the inherited Aptos
cells from a Calibri neighbour in the same column.

    python tools/patch_students_academics_calendar_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
)

REVIEW_DATE = "11 September 2026"
SHORT_DATE = "11 Sep 2026"
MRD_VERSION = "2.76"
CODE_BASELINE = "Backend worktree at 13675db, 11 September 2026"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ── formatting-preserving docx helpers ───────────────────────────────────────

def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError("Paragraph carries no run to inherit formatting from")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell(cell, text: str) -> None:
    """Replace a cell's text, keeping its first run's formatting."""
    for paragraph in cell.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    set_run_text(cell.paragraphs[0], text)


def replace_text(paragraphs, old: str, new: str) -> None:
    """Replace ``old`` inside the single run that holds it."""
    hits = [run for p in paragraphs for run in p.runs if old in run.text]
    if len(hits) != 1 or hits[0].text.count(old) != 1:
        raise ValueError(f"Expected one run holding {old[:70]!r}, found {len(hits)}")
    hits[0].text = hits[0].text.replace(old, new)


def edit_cell(cell, old: str, new: str) -> None:
    replace_text(cell.paragraphs, old, new)


def append_cell(cell, tail: str) -> None:
    cell.paragraphs[-1].runs[-1].text += tail


def append_paragraph(paragraph, tail: str) -> None:
    paragraph.runs[-1].text += tail


def row_by_label(table, label: str):
    hits = [row for row in table.rows if row.cells[0].text.strip() == label]
    if len(hits) != 1:
        raise ValueError(f"Row {label!r}: {len(hits)} matches")
    return hits[0]


def find_row(table, prefix: str, col: int = 0):
    hits = [row for row in table.rows if row.cells[col].text.strip().startswith(prefix)]
    if len(hits) != 1:
        raise ValueError(f"Row {prefix!r}: {len(hits)} matches")
    return hits[0]


def find_api_row(table, method: str, endpoint: str):
    hits = [
        row for row in table.rows
        if row.cells[0].text.strip() == method and row.cells[1].text.strip() == endpoint
    ]
    if len(hits) != 1:
        raise ValueError(f"API row {method} {endpoint}: {len(hits)} matches")
    return hits[0]


def add_row_after(anchor, template, values: list[str]):
    """Insert a copy of ``template`` after ``anchor`` and fill it."""
    new_tr = copy.deepcopy(template._tr)
    anchor._tr.addnext(new_tr)
    row = _Row(new_tr, anchor._parent)
    for cell, value in zip(row.cells, values):
        set_cell(cell, value)
    return row


def append_row(table, template, values: list[str]):
    return add_row_after(table.rows[-1], template, values)


def delete_row(row) -> None:
    row._tr.getparent().remove(row._tr)


def restore_cell_format(cell, template) -> None:
    """Give ``cell`` the cell, paragraph and run formatting of ``template``."""
    if cell._tc is template._tc:
        return
    if template._tc.tcPr is not None:
        if cell._tc.tcPr is not None:
            cell._tc.remove(cell._tc.tcPr)
        cell._tc.insert(0, copy.deepcopy(template._tc.tcPr))
    source = template.paragraphs[0]
    run_format = source.runs[0]._r.rPr if source.runs else None
    for paragraph in cell.paragraphs:
        if paragraph._p.pPr is not None:
            paragraph._p.remove(paragraph._p.pPr)
        if source._p.pPr is not None:
            paragraph._p.insert(0, copy.deepcopy(source._p.pPr))
        for run in paragraph.runs:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            if run_format is not None:
                run._r.insert(0, copy.deepcopy(run_format))


def restore_row_format(row, template_row) -> None:
    for cell, template in zip(row.cells, template_row.cells):
        restore_cell_format(cell, template)


def body_paragraph(doc, prefix: str):
    wanted = " ".join(prefix.split())
    hits = [p for p in doc.paragraphs if " ".join(p.text.split()).startswith(wanted)]
    if len(hits) != 1:
        raise ValueError(f"Paragraph {prefix!r}: {len(hits)} matches")
    return hits[0]


def set_callout(paragraph, text: str) -> None:
    """Rewrite a callout, keeping its marker run and its text run's format."""
    runs = paragraph.runs
    marker = runs[0].text.strip() if runs else ""
    if len(runs) < 2 or not marker or marker[0] not in ("\U0001f4cc", "\u2139"):
        raise ValueError(f"Not a callout: {paragraph.text[:60]!r}")
    runs[1].text = text
    for run in runs[2:]:
        run.text = ""


def last_text_paragraph_before(paragraph):
    element = paragraph._p.getprevious()
    while element is not None:
        if element.tag == W + "p":
            candidate = Paragraph(element, paragraph._parent)
            if candidate.text.strip():
                return candidate
        element = element.getprevious()
    raise ValueError("No text paragraph before the anchor")


def add_bullets_before(doc, heading_prefix: str, texts: list[str]) -> None:
    """Append bullets to the list that ends just before a heading."""
    anchor = last_text_paragraph_before(body_paragraph(doc, heading_prefix))
    if anchor.style is None or anchor.style.name != "List Paragraph":
        raise ValueError(f"Expected a bullet before {heading_prefix!r}")
    element = anchor._p
    for text in texts:
        clone = copy.deepcopy(anchor._p)
        element.addnext(clone)
        set_run_text(Paragraph(clone, anchor._parent), text)
        element = clone


def clone_fr_block(after_element, heading, table, heading_text, header_text, rows):
    """Copy an FR heading and table after ``after_element`` and fill them."""
    new_heading = copy.deepcopy(heading._p)
    new_table = copy.deepcopy(table._tbl)
    after_element.addnext(new_heading)
    new_heading.addnext(new_table)
    set_run_text(Paragraph(new_heading, heading._parent), heading_text)
    body = new_table.findall(W + "tr")[1:]
    while len(body) > len(rows):
        new_table.remove(body.pop())
    while len(body) < len(rows):
        extra = copy.deepcopy(body[-1])
        body[-1].addnext(extra)
        body.append(extra)
    filled = Table(new_table, table._parent)
    set_cell(filled.rows[0].cells[0], header_text)
    for row, (label, value) in zip(filled.rows[1:], rows):
        set_cell(row.cells[0], label)
        set_cell(row.cells[1], value)
    return new_table


def fr_heading(table):
    heading = Paragraph(table._tbl.getprevious(), table._parent)
    if not heading.text.startswith("FR-"):
        raise ValueError(f"No FR heading before table: {heading.text[:40]!r}")
    return heading


def require_fr(table, label: str):
    if not table.rows[0].cells[0].text.strip().startswith(label):
        raise ValueError(f"Table is not {label}")
    return table


def require_header(table, first: str):
    if table.rows[0].cells[0].text.strip() != first:
        raise ValueError(f"Table header is not {first!r}")
    return table


def set_cover(cover, values: dict[str, str], value_template) -> None:
    for label, value in values.items():
        cell = row_by_label(cover, label).cells[1]
        set_cell(cell, value)
        restore_cell_format(cell, value_template)


def finish(doc, output: Path, title: str, version: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


# ═════════════════════════════════════════════════════════════════════════════
# Module 11 - Student Management
# ═════════════════════════════════════════════════════════════════════════════

M11_DIR = "11-student-management"
M11_STEM = "XVS_M11_Student_Management_Functional_Requirements_Document"
M11_SOURCE, M11_TARGET = "2.8", "2.9"

M11_STATUS = (
    "BUILT. The module described here exists: apps/schools/vs_students, mounted "
    "at /v1/students/ and /v1/guardians/. Version 2.5 reconciled the "
    "specification with what shipped. Version 2.9 does so again: every change "
    "to the module's code from version 2.6 to commit 13675db was read against "
    "this document, including those versions 2.7 and 2.8 did not reach, and "
    "each section it touched is written as the code now stands."
)

M11_VERIFIED = (
    f"{CODE_BASELINE}. The seven models and the thirty-four routes of section "
    "10 are present, with three permission keys of this module's own beside the "
    "five it inherits, three configuration definitions, two import datasets and "
    "one export dataset. schools.vs_students ran 279 tests, all passing, at "
    "commit 6611856, which contains every change this version records. The "
    "tests named below were read from the test files at 13675db."
)

M11_SECTION_21_TAIL = (
    " Version 2.9 adds seven rows after those, and all seven correct this "
    "document against its own module's code: sentences that stayed as they were "
    "while the code beneath them moved, most of them in places versions 2.7 and "
    "2.8 did not reach."
)

M11_CORRECTIONS = [
    [
        "v2.8 (FR-010, its acceptance, and section 10): the promotion preview, "
        "run and batch read are gated by school.students.manage.",
        "school.students.promote gates all three, and the run also asserts "
        "academics.classes.assign. Version 2.8 added the key to sections 3.7 and "
        "8.1 and left the requirement, its first acceptance criterion and three "
        "rows of section 10 naming the key it replaced.",
        "apps/schools/vs_students/views/promotion.py, PromotionPreviewView, "
        "PromotionRunView and PromotionBatchView; tests/test_behaviour.py, "
        "PromotionTests.test_running_needs_the_assign_key_as_well_as_promote.",
    ],
    [
        "v2.7 (FR-010 rule 3): a level with no next_level is terminal and "
        "resolves to GRADUATED.",
        "Only a level marked is_terminal graduates. A null next_level without "
        "the flag means nobody has wired the level, and the run holds its "
        "students under the cause LEVEL_NOT_WIRED and names it on the exception "
        "list. Read as terminal, the bare null graduated every year group of a "
        "school that had not finished wiring its ladder, which is every school "
        "that has just been set up.",
        "apps/schools/vs_students/services/promotion.py, _target_class; "
        "apps/schools/vs_students/constants.py, EXC_LEVEL_NOT_WIRED; "
        "tests/test_behaviour.py, UnwiredLevelTests.",
    ],
    [
        "v2.8 (FR-001 and FR-021): the student search is over first_name, "
        "last_name and student_number, and the guardian search over name, phone "
        "and email, each field matched against the whole query.",
        "Both read one shared matcher. The fields are joined into one string and "
        "matched three ways, ranked in this order: the query contained anywhere, "
        "every typed word starting a word in any order, and the typed letters "
        "appearing in order. Matched field by field, a full name typed into the "
        "box found nobody, because no single column holds it. The student search "
        "gains middle_name. The command palette's student search, FR-020, does "
        "not use the matcher yet.",
        "apps/core/search.py, search and loose_match; "
        "apps/schools/vs_students/views/students.py, STUDENT_SEARCH_FIELDS; "
        "apps/schools/vs_students/services/guardians.py, GUARDIAN_SEARCH_FIELDS; "
        "apps/core/test_search.py, LooseSearchTests.",
    ],
    [
        "v2.8 (cover and section 10): thirty-three routes.",
        "Thirty-four. GET /v1/students/classes/seats/, which gives every class "
        "picker its seat counts, was mounted before version 2.7 and never "
        "listed, and GET /v1/guardians/search/ is new. Both are in section 10 "
        "now, under FR-018 and FR-025.",
        "apps/schools/vs_students/urls.py; apps/schools/vs_students/views/"
        "records.py, ClassSeatsView; apps/schools/vs_students/views/"
        "guardians.py, GuardianSearchView.",
    ],
    [
        "v2.4 (FR-005 and section 10): a guardian is unlinked from a student "
        "with DELETE /v1/guardians/<id>/.",
        "The unlink sits under the student, DELETE /v1/students/<id>/guardians/"
        "<guardian_id>/, because a guardian is linked to several children and "
        "deleting \"this guardian\" would not say which of them the school meant. "
        "PATCH on the same route changes the relationship or makes the guardian "
        "primary. /v1/guardians/<id>/ answers GET and PATCH only.",
        "apps/schools/vs_students/views/guardians.py, StudentGuardianDetailView "
        "and GuardianDetailView; apps/schools/vs_students/urls.py.",
    ],
    [
        "v2.4 to v2.8 (FR-012 keys, and section 13 rows 14 and 18): "
        "school.students.import cannot stand in for the generic import keys, "
        "because the import permission bridge is hard-coded to bank statements.",
        "It can, and has since version 2.5 built the registry section 8.4 "
        "describes: a domain app writes its dataset-to-key pair from "
        "AppConfig.ready and the engine imports nothing. The students and "
        "guardians datasets both register school.students.import.",
        "apps/vs_import_data/permissions.py, register_dataset_import_key and "
        "HasImportBatchRBACPermission; apps/schools/vs_students/services/"
        "import_registry.py.",
    ],
    [
        "v2.8 (row R31 above, section 13 row 10 and section 14 decision 10): no "
        "view in the platform enforces a capability.",
        "Every permission key carries a depth band, and HasRBACPermission "
        "refuses PLAN_UPGRADE_REQUIRED after the role check when the school's "
        "plan does not reach it; the seeded catalogue ships "
        "platform.entitlements.enforce on. school.students sits in the students "
        "module at Core, promote at Plus, import and export under the platform's "
        "bulk import and data export features, and view_sensitive carries no "
        "band at all, so no plan buys a child's medical record.",
        "apps/vs_rbac/permission_bands.py, CAPABILITY_MODULE, ACTION_CAPABILITY, "
        "RESOURCE_BANDS, ACTION_BANDS and NEVER_BAND; apps/vs_rbac/plan_gate.py; "
        "apps/vs_config/management/commands/seed_config_catalogue.py.",
    ],
]

M11_MANAGE_USE = (
    "FR-007 withdraw, FR-008 suspend, FR-009 reactivate, FR-017 transfer out "
    "and FR-022's status action. Promotion no longer rides on it: "
    "school.students.promote, below."
)

M11_GUARDIAN_DATASET = [
    "Guardians dataset",
    "A second dataset, guardians, carries the households a school calls: one "
    "row per guardian-and-child link. It registers school.students.import "
    "through the same registry, is classified in vs_import_data.datasets as a "
    "dataset a school may import for itself, and is published as one atomic "
    "job rather than row by row, because its primary-contact rules are read "
    "across rows. FR-024.",
]

M11_FR001_FILTERS = (
    "status, branch (only where the school has more than one branch), class, "
    "level, an unassigned flag, session, and a search. The search matches "
    "first_name, middle_name, last_name and student_number as one string, three "
    "ways, and ranks the hits: the query contained anywhere, every typed word "
    "starting a word in any order, and the typed letters appearing in order. So "
    "\"Tobi Okafor\", \"okafor tobi\" and \"to ok\" all find Tobi Okafor, a "
    "child with no middle name or number is still found, and an exact hit sorts "
    "above a coincidence. Filters are applied on top of the scoped queryset and "
    "can only narrow it, never widen it. session reads the roll as it was in "
    "that year: the students with a placement in it, each row's class read from "
    "that year's placement, including one a promotion has since closed; an "
    "unknown year is refused rather than ignored. The unassigned flag is a "
    "class filter of \"none\" and not the same query as /v1/students/unplaced/, "
    "which takes no session because every row on it is a child to place now, "
    "and level narrows through the class's level rather than through a column "
    "on the student."
)

M11_FR001_ACCEPTANCE_TAIL = (
    " (10) A full name typed in either order finds the child, and a stray "
    "regular-expression character in the query is matched literally rather "
    "than answering 500 (apps/core/test_search.py, LooseSearchTests). (11) "
    "Asked about a past year, the list shows that year's roll and each child's "
    "class in that year, and an unknown year is refused (tests/test_shape.py, "
    "SessionLensTests)."
)

M11_FR005_TRIGGER = (
    "GET and POST /v1/students/<id>/guardians/; PATCH and DELETE "
    "/v1/students/<id>/guardians/<guardian_id>/, which change one link's "
    "relationship or primary flag and unlink it, with ?promote= naming the "
    "guardian who becomes primary; PATCH /v1/guardians/<id>/, which corrects "
    "the guardian's own details; GET /v1/guardians/<id>/students/. GET "
    "/v1/guardians/ and GET /v1/guardians/<id>/ are FR-021, and GET "
    "/v1/guardians/search/ is FR-025."
)

M11_FR010_KEYS = (
    "school.students.promote to preview, run and read a batch; "
    "academics.classes.assign as well on the run, because it writes "
    "placements. The run checks the second key explicitly, because "
    "rbac_permission is any-of (section 10)."
)

M11_FR010_RULE_3 = (
    "(3) A level marked is_terminal resolves to GRADUATE. A level with no "
    "next_level that is not so marked has not been wired, and resolves to HOLD "
    "with the cause LEVEL_NOT_WIRED: read as terminal, the bare null graduated "
    "every unwired year group at once, and a school that has just been set up "
    "has wired nothing."
)

M11_FR010_RULE_8 = (
    "(8) The default outcome is GRADUATE where the level is marked terminal, "
    "HOLD where the level is not wired, where the next level exists but no "
    "class does, and where the student is only ENROLLED, and PROMOTE otherwise."
)

M11_FR010_RULES_TAIL = (
    " (12) A preview or run may be narrowed to one branch with ?branch=, and "
    "the branch reaches the classification itself: the candidates are that "
    "branch's students and the targets are that branch's classes plus the "
    "school-wide ones. The preview and the run read the branch through the same "
    "property, so the set a registrar reviews is the set that moves. (13) A "
    "batch is read under the inclusive branch rule. A run pinned to one branch "
    "opens only to a caller who covers that branch and answers 404 to anybody "
    "else, because what it records is which children repeated the year or were "
    "held back, by name; a run with no branch covered the whole school and "
    "opens from every branch."
)

M11_FR010_CAUSES_OLD = (
    "The exception list has a fixed vocabulary of five causes, because the "
    "screen prints the sentence: TERMINAL_LEVEL, NO_CLASS_AT_NEXT_LEVEL, "
    "NO_CLASS_TO_REPEAT, STUDENT_SUSPENDED and NO_CLASS_ASSIGNED. The first "
    "three are facts about a class"
)

M11_FR010_CAUSES_NEW = (
    "The exception list has a fixed vocabulary of six causes, because the "
    "screen prints the sentence: TERMINAL_LEVEL, LEVEL_NOT_WIRED, "
    "NO_CLASS_AT_NEXT_LEVEL, NO_CLASS_TO_REPEAT, STUDENT_SUSPENDED and "
    "NO_CLASS_ASSIGNED. The first four are facts about a class"
)

M11_FR010_PREVIEW_TAIL = (
    " LEVEL_NOT_WIRED says whose problem it is: set the level's promotion "
    "target in Academic Structure, or mark the level as one pupils leave after. "
    "The level map marks a class terminal only where its level says so, so a "
    "cohort nobody has wired is not shown as leaving."
)

M11_FR010_DEPENDS_OLD = (
    "is_terminal shipped in M13 v2.8 and the fallback is now live rather than "
    "defensive."
)

M11_FR010_DEPENDS_NEW = (
    "is_terminal shipped in M13 v2.8, and the classifier now reads the three "
    "states apart: a target set promotes, the flag graduates, and a bare null "
    "holds under LEVEL_NOT_WIRED rather than graduating."
)

M11_FR010_ACCEPTANCE_1_OLD = (
    "(1) A caller without academics.classes.assign answers 403 on the run and "
    "succeeds on the preview only if they hold manage."
)

M11_FR010_ACCEPTANCE_1_NEW = (
    "(1) A caller without school.students.promote answers 403 on the preview, "
    "the run and the batch read, and one holding promote without "
    "academics.classes.assign answers 403 on the run alone "
    "(tests/test_behaviour.py, "
    "PromotionTests.test_running_needs_the_assign_key_as_well_as_promote)."
)

M11_FR010_ACCEPTANCE_TAIL = (
    " (14) A level with no promotion target and no terminal flag holds its "
    "students, names LEVEL_NOT_WIRED rather than TERMINAL_LEVEL, and is not "
    "shown as leaving on the level map, while a level marked terminal still "
    "graduates them (tests/test_behaviour.py, UnwiredLevelTests). (15) Another "
    "branch's run answers 404 to a head pinned to one branch, their own branch's "
    "run opens, and a whole-school run opens from any branch "
    "(tests/test_security.py, PromotionRunReachTests). (16) A preview or run "
    "narrowed to one branch classifies only that branch's students, and offers "
    "only that branch's classes and the school-wide ones as targets."
)

M11_FR012_KEYS = (
    "school.students.import, which stands in for the engine's generic "
    "import.batches.* keys on a students batch through the dataset registry in "
    "section 8.4. A caller holding the generic keys reaches the wizard as well."
)

M11_FR012_RULES_TAIL = (
    " (9) A file is refused everything the enrolment form refuses and more, "
    "before any row is written: a date of birth that makes the child under 2 "
    "or over 25, an admission date in the future or before the child was born, "
    "a guardian email that is not an address, a guardian phone with fewer than "
    "seven digits, and any value longer than its column, which would otherwise "
    "fail halfway through the write with earlier rows committed. (10) Three "
    "faults only a file can carry are warned about rather than refused, each "
    "named once: a class the file fills past its capacity, imported all the "
    "same because a school loading a roll it already has is describing what is "
    "true; a guardian contact another row or an existing guardian holds under a "
    "different name, because the child joins that household and the row's name "
    "is discarded; and a digit inside a name, which is usually a shifted "
    "column. The template's own guidance states these rules, so a school reads "
    "them before it builds the file."
)

M11_FR012_ACCEPTANCE_TAIL = (
    " (9) A birth year that makes the child an adult or a baby, a future "
    "admission date, an admission date before the birth, an over-long value, a "
    "malformed guardian email and a phone too short to ring are each refused "
    "before anything is written, and a file that overfills a class warns once "
    "with the count (tests/test_import_export.py, "
    "ImportCatchesWhatTheFormCatchesTests)."
)

M11_FR018_RULES_TAIL = (
    " (6) seats returns every class of the year being read with its capacity, "
    "the seats used and the seats remaining, in one request, because the "
    "enrolment form, the transfer drawer and the assign bar each render every "
    "class's load at once. A class with no capacity is returned with a null "
    "capacity rather than dropped. The rows come in ladder order, programme "
    "then level then name, and carry level_order as the [programme, level] "
    "pair a client sorts its derived level list by, because each programme "
    "numbers its levels from one and level order alone interleaves JSS1 with "
    "SSS1. It is not paginated, answers an empty list rather than "
    "NO_ACTIVE_SESSION between years, counts each student's last placement in "
    "the year asked about so that a year a promotion has closed still shows its "
    "loads, and follows the branch filter. (7) summary and seats read the year "
    "?session= names. Under a session the summary's status counts describe "
    "that year's students as they stand today, because a status carries no "
    "year, and the response says so with status_is_current."
)

M11_FR018_ACCEPTANCE_TAIL = (
    " (8) The seat counts offer only the year being read, narrow with the "
    "branch filter, and keep a past year's loads (tests/test_shape.py, "
    "ClassSeatsTests and SessionLensReachesEveryListTests)."
)

M11_FR019_RULES_TAIL = (
    " (7) GET also returns a suggestion for the next admission number, read "
    "from the number the school most recently issued rather than from its "
    "pattern: the trailing digits are incremented and their zero padding kept. "
    "It returns an empty string rather than guessing when the school has "
    "issued none, when the latest number does not end in digits, or when the "
    "successor fails the school's own pattern, and it skips a successor already "
    "taken. A suggestion reserves nothing: two registrars can be handed the "
    "same one, and the uniqueness constraint in section 7.1 refuses the second."
)

M11_FR019_ACCEPTANCE_TAIL = (
    " (8) The suggestion continues the school's own format, keeps zero "
    "padding, suggests nothing before the first number, after a number ending "
    "in no digits or where the successor breaks the school's rule, and skips a "
    "taken successor (tests/test_shape.py, AdmissionSuggestionTests)."
)

M11_FR020_GAP = [
    "Current gap",
    "The palette does not read the shared matcher that FR-001 and FR-021 use. "
    "It still matches first_name, last_name and student_number each against "
    "the whole query, and never middle_name, so a child's full name typed into "
    "the palette finds nobody while the directory beneath it finds them: rule "
    "1's promise that the two agree holds for what the palette returns and not "
    "for what it misses. The guardian palette search, FR-025, reads the "
    "directory's matcher and does not share the gap. "
    "apps/schools/vs_students/views/students.py, StudentSearchView.",
]

M11_FR021_RULE_1_OLD = (
    "(1) The list is searchable by name and paginated, and each row carries "
    "the ward count and the ward names, which is what makes siblings visible "
    "at a glance."
)

M11_FR021_RULE_1_NEW = (
    "(1) The list is searchable and paginated, and each row carries the ward "
    "count and the ward names, which is what makes siblings visible at a "
    "glance. The search matches full_name, phone and email as one string, "
    "three ways and ranked, exactly as FR-001's does; the phone is in it "
    "because \"who do we call about this family\" is often asked from a missed "
    "call. The list is ordered by name and then id, by rank first where a "
    "search ran, so a page never repeats a guardian from the page before or "
    "drops one."
)

M11_FR021_RULES_TAIL = (
    " (5) ?session= narrows the list to the guardians of the children on that "
    "year's roll: a guardian carries no year any more than a branch, so the "
    "narrowing is on the wards, as the branch narrowing is."
)

M11_FR021_ACCEPTANCE_TAIL = (
    " (7) Asked about a past year, the directory lists the guardians of that "
    "year's roll, and its pages are stable (tests/test_shape.py, "
    "SessionLensReachesEveryListTests)."
)

M11_FR024 = [
    ("Description",
     "Load the households a school calls, every father, grandmother and legal "
     "guardian after the one the student import carries, through the "
     "platform's import engine, validated before anything is written."),
    ("Keys",
     "school.students.import, the roll's own key, registered for the guardians "
     "dataset beside the students one: a guardian is part of the roll, and the "
     "student import already creates one per child."),
    ("Trigger",
     "The existing /v1/import/ surface, with dataset_type=guardians. This "
     "module adds no upload endpoint of its own."),
    ("Columns",
     "guardian_full_name, guardian_phone, guardian_email, student_number, "
     "student_first_name, student_last_name, student_date_of_birth, "
     "relationship, is_primary, occupation, address. One row is one link, not "
     "one parent: the relationship and the primary-contact flag belong to the "
     "pair, so a parent with three children at the school appears three times."),
    ("Business rules",
     "(1) The child is reached by admission number, which is exact, or by first "
     "name, last name and date of birth. A name that matches two children is "
     "refused and the message asks for the number, because a resolver that "
     "took the first would attach a father to the wrong family in silence; a "
     "number no student holds, and a row naming no child at all, are refused "
     "too. (2) A guardian is matched on email first and phone second, the rule "
     "upsert_guardian applies everywhere else, so a parent the school already "
     "holds is reused and keeps the name on file. (3) Two rows claiming primary "
     "contact for one child are refused, because the second would silently "
     "win. (4) A row that moves a child's primary contact from somebody else is "
     "warned about and names who loses it, and a child this file gives "
     "contacts to while marking none primary is warned about, because the "
     "school would hold numbers and no answer to who it calls first. Both "
     "checks read the whole file, so row order does not change the answer. (5) "
     "A pair already linked is skipped rather than linked twice, and the same "
     "pair twice in one file is refused. (6) A contact the file or the school "
     "already holds under a different name is warned about, as in FR-012 rule "
     "10. (7) A phone with fewer than seven digits and a malformed email are "
     "refused. (8) The file is published as one atomic job, because the "
     "primary-contact rules are read across rows and half a file could move a "
     "primary and then fail before the row meant to replace it."),
    ("Pending tenant",
     "The import surface declares pending_tenant_surface, and guardians is "
     "classified in vs_import_data.datasets as a dataset a school may import "
     "for itself, so a school still onboarding can load its households with its "
     "roll."),
    ("Acceptance",
     "(1) A clean row reports no error, and executing it creates the guardian "
     "and the link. (2) An admission number no student holds, a child named "
     "ambiguously and a row naming no child are each refused. (3) Moving the "
     "primary contact is warned about and names who loses it; two rows claiming "
     "primary for one child are refused; a child given contacts and no primary "
     "is warned about, whatever order the rows come in. (4) A pair already "
     "linked is skipped, and the same pair twice in one file is refused. (5) A "
     "contact held under another name is warned about, and a phone too short to "
     "ring is refused. (6) The template exists with all eleven columns, and a "
     "school may import its own households. tests/test_import_export.py, "
     "GuardianImportTests."),
]

M11_FR025 = [
    ("Description",
     "The command palette's guardian hits: a small, capped lookup by name, "
     "phone or email that answers who, and whose children, and nothing more."),
    ("Key", "school.students.view."),
    ("Trigger", "GET /v1/guardians/search/?q=."),
    ("Business rules",
     "(1) It reads the same scoped queryset and the same matcher as FR-021's "
     "directory, so the palette can never find a guardian the directory cannot "
     "open, and never misses one it would have found. (2) A query shorter than "
     "two characters returns an empty list rather than the whole school. (3) "
     "The result is capped at ten and is not paginated. (4) Each hit carries "
     "id, full_name and ward_names and nothing else: no phone, no email, no "
     "photograph. The palette opens on the second keystroke over whatever page "
     "the reader is on, in a staffroom where somebody else can see the screen; "
     "a name and the children it belongs to tell two guardians of one surname "
     "apart, and the rest stays on the record for whoever opens it. (5) "
     "ward_names are narrowed to the branches the caller covers, exactly as the "
     "directory's are, so a branch-bound head sees a school-level guardian with "
     "their own branch's children beside them and no others. The palette must "
     "not become the way round that narrowing."),
    ("Acceptance",
     "(1) 403 without school.students.view (tests/test_security.py, "
     "PermissionDeniedTests, whose sweep includes guardian-search). (2) Every "
     "hit carries exactly id, full_name and ward_names "
     "(test_the_guardian_search_carries_no_contact_details). (3) A head pinned "
     "to one branch sees the guardian with their own branch's child and not the "
     "child at the other branch "
     "(test_the_guardian_search_names_only_wards_the_caller_may_see). (4) A "
     "one-character query returns an empty list. (5) No more than ten hits come "
     "back, and the response carries no pagination block."),
]

M11_SECTION_10_CALLOUT = (
    "PATCH /v1/guardians/<id>/, and every other route under /v1/guardians/, "
    "address a guardian reachable from more than one student, so the tenant "
    "check on them cannot be inherited from a student in the URL. Each resolves "
    "the guardian against request.tenant explicitly and answers 404 for one "
    "belonging to another tenant. Version 1.0 declares these routes and never "
    "mentions the scoping problem they create. The unlink sits under the "
    "student, at /v1/students/<id>/guardians/<guardian_id>/, because deleting "
    "\"this guardian\" would not say which of the guardian's children the school "
    "meant."
)

M11_M13_DEPENDENCIES = {
    "A Staff model": (
        "Partly discharged. schools.vs_staff exists (M12), and "
        "SchoolClass.class_teacher names a class's teacher through its "
        "StaffProfile, so the recipient the student.enrolled template addresses "
        "can now be found. Nothing here reads it yet: FR-002 still writes its "
        "audit event and dispatches nothing, the four student.* events are still "
        "inactive (row 4), and the teacher prebuilt role still sees every student "
        "its branch grants reach, which is section 14, decision 6."
    ),
    "AuditModuleKey.STUDENT": (
        "DISCHARGED. AuditModuleKey.STUDENT and the action types of section 8.2 "
        "are registered in vs_audit, so every event this module writes is kept."
    ),
    "school.students.import and school.students.export": (
        "DISCHARGED. Both are seeded, with school.students.promote beside them, "
        "and all three sit in the Student Bulk Data bundle (section 8.1)."
    ),
    "A students dataset in the Export Centre": (
        "DISCHARGED. apps/schools/vs_students/export_datasets.py registers the "
        "students dataset from AppConfig.ready, fenced to the tenant (FR-014)."
    ),
    "A students member of DatasetTypeChoices": (
        "DISCHARGED. DatasetTypeChoices carries students and guardians, each "
        "with its handler and template, and both are classified in "
        "vs_import_data.datasets as datasets a school may import for itself "
        "(FR-012, FR-024)."
    ),
    "A main branch on every school": (
        "DISCHARGED. School creation refuses a school with no branch: branches "
        "is required and may not be empty, and a lone branch is made the main "
        "branch. apps/schools/vs_schools/serializers.py, the school create "
        "serializer's branches field and its main-branch rule."
    ),
    "A capability gate": (
        "ANSWERED by the platform. HasRBACPermission runs a plan gate after the "
        "role check and refuses PLAN_UPGRADE_REQUIRED where the school's plan "
        "does not reach the key's band, and the seeded catalogue ships "
        "platform.entitlements.enforce on. Section 14, decision 10."
    ),
    "A generalised import permission bridge": (
        "DISCHARGED at version 2.5 and recorded here at version 2.9. "
        "vs_import_data.permissions carries a dataset-to-key registry that a "
        "domain app writes into from AppConfig.ready, and school.students.import "
        "is registered for the students and guardians datasets (section 8.4)."
    ),
    "A dataset-aware import permission bridge": (
        "DISCHARGED with row 14: the registry is built, so "
        "school.students.import reaches the wizard for a students or a guardians "
        "batch."
    ),
}

M11_SECTION_13_TAIL = (
    " Rows that have since been built stay in the table marked DISCHARGED, as "
    "row 1 does, so a reader who remembers the gap meets its answer."
)

M11_DECISION_10_WHY = (
    "Answered by the platform rather than here. Every permission key now "
    "carries a depth band, and HasRBACPermission refuses "
    "PLAN_UPGRADE_REQUIRED, naming the module and the depth the school would "
    "need, after the role check has passed; the seeded catalogue ships "
    "platform.entitlements.enforce on. school.students is in the students "
    "module at Core, so every school's plan reaches the record itself. promote "
    "is sold at Plus, import and export with the platform's bulk import and "
    "data export features, and view_sensitive carries no band at all, so a "
    "school cannot buy its way into a child's medical record. A school with no "
    "package grants at all is unprovisioned rather than unentitled, and is not "
    "refused."
)

M11_DECISION_10_COST = (
    "Nothing left to defer here. Which depth each key is sold at is decided in "
    "the band map, apps/vs_rbac/permission_bands.py, and moving one is a "
    "pricing decision rather than a change to this module."
)

M11_SECTION_12_BULLETS = {
    "12.2 Behaviour": [
        "GET /v1/guardians/search/ returns only id, full_name and ward_names for "
        "every hit, and names only the wards of the branches the caller covers "
        "(FR-025).",
        "A head pinned to one branch cannot open another branch's promotion "
        "run, which answers 404, and can open a whole-school run (FR-010).",
    ],
    "12.3 Shape and scale": [
        "A level with no promotion target and no terminal flag holds its "
        "students under LEVEL_NOT_WIRED and graduates nobody (FR-010).",
        "A child's full name typed in either order finds them in the directory, "
        "and a guardian's finds them in the guardian directory and its palette "
        "search (FR-001, FR-021, FR-025).",
        "A guardians file that claims two primary contacts for one child is "
        "refused before anything is written, and one that moves a primary "
        "contact names who loses it (FR-024).",
    ],
    "12.4 Multi-shape tenancy": [
        "The list, the summary, the guardian directory and the seat counts all "
        "read the year ?session= names, and an unknown year is refused (FR-001, "
        "FR-018, FR-021).",
    ],
}

M11_CHANGE_SUMMARY = (
    "Minor revision, and the first full reconciliation since version 2.5: every "
    "change to the module's code from version 2.6 to commit 13675db was read "
    "against this document, including the ones versions 2.7 and 2.8 did not "
    "reach. FR-024 is new and carries the guardians import, the second parent "
    "nothing else could add in bulk: one row per guardian-and-child link, the "
    "child reached by admission number or refused where a name is ambiguous, "
    "primary-contact moves warned about by name, and the whole file published "
    "as one job. FR-025 is new and carries the guardian palette search, which "
    "reads the directory's scoping and matcher, returns ten hits at most and "
    "never paginates, carries id, full name and ward names and no contact "
    "details, and names only the wards the caller's branches reach. FR-001 and "
    "FR-021 record the shared matcher that finds a person from a full name "
    "typed in either order, and FR-020 records that the student palette does "
    "not use it yet, so a full name finds nobody there. FR-010 moves to "
    "school.students.promote, which version 2.8 added to the key tables and not "
    "to the requirement, the API table or the acceptance; records "
    "LEVEL_NOT_WIRED, the sixth exception cause, under which an unwired level "
    "holds its students instead of graduating them; records that a run "
    "narrowed to one branch classifies only that branch; and records that a "
    "batch is read under the inclusive branch rule, so another branch's run "
    "answers 404. FR-012 records the refusals and warnings the student import "
    "gained. FR-018 gains the class seat counts route, never listed before, and "
    "the year lens on the summary; FR-019 gains the next-number suggestion; "
    "FR-021 gains the year lens and stable pages. Section 10 lists thirty-four "
    "routes and corrects where a guardian is unlinked, section 11 corrects the "
    "status of an unknown branch and adds an unknown year, section 13 marks the "
    "dependencies that have since been built, and decision 10 is answered by "
    "the platform's plan gate. Section 2.2 gains seven rows. FR-023 gains the "
    "heading it was missing, and cells an earlier revision set in another "
    "typeface are returned to the document's own. The note on what version 2.3 "
    "could not check is removed, because this version was rendered and read "
    "page by page. schools.vs_students ran 279 tests, all passing, at commit "
    "6611856, which contains every change recorded here. Backend evidence only; "
    "nothing here is deployed."
)


def patch_m11(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M11 Student Management Functional Requirements Document v{M11_TARGET}"
    tables = doc.tables
    cover, corrections = tables[0], tables[1]
    keys37, student_model, guardian_model = tables[5], tables[15], tables[17]
    keys81, dataset84 = tables[23], tables[25]
    fr001 = require_fr(tables[26], "FR-001")
    fr005 = require_fr(tables[30], "FR-005")
    fr010 = require_fr(tables[35], "FR-010")
    fr012 = require_fr(tables[38], "FR-012")
    fr015 = require_fr(tables[41], "FR-015")
    fr023 = require_fr(tables[42], "FR-023")
    fr018 = require_fr(tables[45], "FR-018")
    fr019 = require_fr(tables[46], "FR-019")
    fr020 = require_fr(tables[47], "FR-020")
    fr021 = require_fr(tables[48], "FR-021")
    fr022 = require_fr(tables[49], "FR-022")
    api = require_header(tables[50], "Method")
    refusals = require_header(tables[51], "Condition")
    dependencies = require_header(tables[52], "Dependency")
    decisions = require_header(tables[53], "Question")
    log = require_header(tables[54], "Version")
    require_header(corrections, "What an earlier version said")

    label_template = fr001.rows[1].cells[0]
    value_template = fr001.rows[1].cells[1]

    # Cover.
    status_cell = row_by_label(cover, "Status").cells[1]
    set_cover(cover, {
        "Version": M11_TARGET,
        "Date": REVIEW_DATE,
        "Supersedes": f"v{M11_SOURCE}",
        "Status": M11_STATUS,
        "Verified against": M11_VERIFIED,
    }, status_cell)
    append_row(cover, row_by_label(cover, "Status"), [
        "Source MRD", f"XVS Module Requirements Document v{MRD_VERSION} | Module 11",
    ])

    # Cells an earlier revision set in Aptos Narrow, returned to Calibri.
    for index in range(62, 68):
        restore_row_format(corrections.rows[index], corrections.rows[61])
    restore_row_format(keys37.rows[7], keys37.rows[6])
    restore_cell_format(student_model.rows[13].cells[1], student_model.rows[12].cells[1])
    restore_row_format(guardian_model.rows[7], guardian_model.rows[6])
    restore_row_format(keys81.rows[3], keys81.rows[2])
    for index in (1, 4, 5, 6):
        restore_cell_format(fr015.rows[index].cells[1], value_template)
    restore_cell_format(fr023.rows[0].cells[0], fr015.rows[0].cells[0])
    for row in fr023.rows[1:]:
        restore_cell_format(row.cells[0], label_template)
        restore_cell_format(row.cells[1], value_template)
    restore_row_format(api.rows[35], api.rows[34])
    restore_row_format(api.rows[36], api.rows[34])
    for index in (9, 10, 11):
        restore_row_format(log.rows[index], log.rows[8])

    # Section 2: the corrections this version records.
    append_paragraph(
        body_paragraph(doc, "The table below lists every claim a builder must not trust."),
        M11_SECTION_21_TAIL,
    )
    append_cell(corrections.rows[31].cells[1], " Superseded by row R74.")
    for values in M11_CORRECTIONS:
        append_row(corrections, corrections.rows[61], values)

    # Sections 3.7, 8.1 and 8.4.
    set_cell(find_row(keys37, "school.students.manage").cells[3], M11_MANAGE_USE)
    replace_text(
[body_paragraph(doc, "The school.students resource already exists with five keys")],
        "Two new keys are needed", "Three new keys are needed",
    )
    append_row(dataset84, dataset84.rows[-1], M11_GUARDIAN_DATASET)

    # Section 9.
    replace_text(
[body_paragraph(doc, "Twenty-two requirements.")],
        "Twenty-two requirements.", "Twenty-five requirements.",
    )
    set_cell(find_row(fr001, "Filters").cells[1], M11_FR001_FILTERS)
    append_cell(
        find_row(fr001, "Field exposure").cells[1],
        " The list row also carries applied_on, which the applicant board sorts by.",
    )
    append_cell(find_row(fr001, "Acceptance").cells[1], M11_FR001_ACCEPTANCE_TAIL)
    set_cell(find_row(fr005, "Trigger").cells[1], M11_FR005_TRIGGER)

    set_cell(find_row(fr010, "Keys").cells[1], M11_FR010_KEYS)
    rules = find_row(fr010, "Business rules").cells[1]
    edit_cell(
        rules, "(3) A level with no next_level is terminal and resolves to GRADUATED.",
        M11_FR010_RULE_3,
    )
    edit_cell(
        rules,
        "(8) The default outcome is GRADUATE where the level is terminal, HOLD "
        "where the next level exists but no class does and where the student is "
        "only ENROLLED, and PROMOTE otherwise.",
        M11_FR010_RULE_8,
    )
    append_cell(rules, M11_FR010_RULES_TAIL)
    preview = find_row(fr010, "Preview").cells[1]
    edit_cell(preview, M11_FR010_CAUSES_OLD, M11_FR010_CAUSES_NEW)
    append_cell(preview, M11_FR010_PREVIEW_TAIL)
    edit_cell(find_row(fr010, "Depends on").cells[1], M11_FR010_DEPENDS_OLD, M11_FR010_DEPENDS_NEW)
    acceptance = find_row(fr010, "Acceptance").cells[1]
    edit_cell(acceptance, M11_FR010_ACCEPTANCE_1_OLD, M11_FR010_ACCEPTANCE_1_NEW)
    append_cell(acceptance, M11_FR010_ACCEPTANCE_TAIL)

    set_cell(find_row(fr012, "Keys").cells[1], M11_FR012_KEYS)
    append_cell(find_row(fr012, "Business rules").cells[1], M11_FR012_RULES_TAIL)
    append_cell(find_row(fr012, "Acceptance").cells[1], M11_FR012_ACCEPTANCE_TAIL)

    edit_cell(
        find_row(fr018, "Description").cells[1],
        "Five reads the directory, the profile and the class screens cannot be drawn without",
        "Six reads the directory, the profile, the class screens and the class "
        "pickers cannot be drawn without",
    )
    edit_cell(
        find_row(fr018, "Key").cells[1],
        "school.students.view on all five; the roster additionally requires "
        "academics.classes.view, because it is a fact about a class",
        "school.students.view on all six; the roster and the seat counts "
        "additionally require academics.classes.view, because each is a fact "
        "about a class",
    )
    edit_cell(
        find_row(fr018, "Trigger").cells[1],
        "GET /v1/students/classes/<class_id>/roster/.",
        "GET /v1/students/classes/<class_id>/roster/, GET /v1/students/classes/seats/.",
    )
    append_cell(find_row(fr018, "Business rules").cells[1], M11_FR018_RULES_TAIL)
    fr018_acceptance = find_row(fr018, "Acceptance").cells[1]
    edit_cell(fr018_acceptance, "(1) Each of the five answers 403", "(1) Each of the six answers 403")
    append_cell(fr018_acceptance, M11_FR018_ACCEPTANCE_TAIL)

    append_cell(find_row(fr019, "Business rules").cells[1], M11_FR019_RULES_TAIL)
    append_cell(find_row(fr019, "Acceptance").cells[1], M11_FR019_ACCEPTANCE_TAIL)

    fr020_rules = find_row(fr020, "Business rules")
    add_row_after(fr020_rules, fr020_rules, M11_FR020_GAP)

    fr021_rules = find_row(fr021, "Business rules").cells[1]
    edit_cell(fr021_rules, M11_FR021_RULE_1_OLD, M11_FR021_RULE_1_NEW)
    append_cell(fr021_rules, M11_FR021_RULES_TAIL)
    append_cell(find_row(fr021, "Acceptance").cells[1], M11_FR021_ACCEPTANCE_TAIL)

    # FR-023 had no heading, so Word ran it into FR-015's table.
    heading = copy.deepcopy(fr_heading(fr015)._p)
    fr023._tbl.addprevious(heading)
    set_run_text(Paragraph(heading, fr015._parent), "FR-023  Guardian photograph")

    # FR-024 and FR-025 after FR-022, the last requirement in section 9.
    after = clone_fr_block(
        fr022._tbl, fr_heading(fr012), fr012,
        "FR-024  Bulk import guardians", "FR-024  Bulk Import Guardians", M11_FR024,
    )
    clone_fr_block(
        after, fr_heading(fr020), fr020,
        "FR-025  Search guardians", "FR-025  Search Guardians", M11_FR025,
    )

    # Section 10.
    intro = body_paragraph(doc, "All under /v1/students/ except")
    replace_text(
        [intro],
        "except the two guardian routes and the guardian directory, which are under",
        "except the guardian routes, which are under",
    )
    replace_text(
        [intro],
        "on every list except FR-020's capped search.",
        "on every list except the two capped palette searches, FR-020 and "
        "FR-025, and FR-018's seat counts.",
    )
    api_template = api.rows[33]
    set_cell(find_api_row(api, "PATCH", "/v1/guardians/<id>/").cells[3],
             "FR-005. Correct the guardian's own details.")
    unlink = find_api_row(api, "DELETE", "/v1/guardians/<id>/")
    set_cell(unlink.cells[1], "/v1/students/<id>/guardians/<guardian_id>/")
    set_cell(unlink.cells[3],
             "FR-005. Unlink from this student; ?promote= names the guardian who "
             "becomes primary.")
    add_row_after(unlink, api_template, [
        "PATCH", "/v1/students/<id>/guardians/<guardian_id>/", "school.students.update",
        "FR-005. Change the relationship, or make this guardian primary.",
    ])
    set_cell(find_api_row(api, "POST", "/v1/students/promotions/preview/").cells[2],
             "school.students.promote")
    set_cell(find_api_row(api, "POST", "/v1/students/promotions/").cells[2],
             "school.students.promote + academics.classes.assign")
    batch = find_api_row(api, "GET", "/v1/students/promotions/<id>/")
    set_cell(batch.cells[2], "school.students.promote")
    set_cell(batch.cells[3],
             "FR-010. Summary and counters, read under the inclusive branch rule.")
    add_row_after(find_api_row(api, "GET", "/v1/guardians/<id>/"), api_template, [
        "GET", "/v1/guardians/search/", "school.students.view",
        "FR-025. Command palette. Ten hits at most, never paginated, and no "
        "contact details.",
    ])
    add_row_after(
        find_api_row(api, "GET", "/v1/students/classes/<class_id>/roster/"), api_template, [
            "GET", "/v1/students/classes/seats/",
            "academics.classes.view + school.students.view",
            "FR-018. Every class of the year with its seats used and remaining, "
            "for the pickers. Not paginated.",
        ],
    )
    add_row_after(
        find_api_row(api, "(engine)", "/v1/import/ with dataset_type=students"), api_template, [
            "(engine)", "/v1/import/ with dataset_type=guardians", "school.students.import",
            "FR-024. Declares pending_tenant_surface.",
        ],
    )
    set_callout(body_paragraph(doc, "\U0001f4cc PATCH and DELETE on /v1/guardians/<id>/"),
                M11_SECTION_10_CALLOUT)

    # Section 11.
    branch_row = find_row(refusals, "Naming a branch that is unknown, malformed or another tenant's")
    set_cell(branch_row.cells[1], "400")
    set_cell(branch_row.cells[2],
             "Standard validation error on the field, from resolve_branch_reference. "
             "One answer for all three, so the parameter cannot be used to "
             "enumerate. A list's ?branch= filter at a school with one branch is "
             "ignored rather than refused.")
    append_row(refusals, refusals.rows[-1], [
        "?session= naming a year this school does not have", "400",
        "Standard validation error on the session field, refused rather than "
        "ignored, because a filter that quietly does nothing reads as a year with "
        "nobody in it (FR-001, FR-018, FR-021).",
    ])

    # Section 12.
    for heading_text, bullets in M11_SECTION_12_BULLETS.items():
        add_bullets_before(doc, heading_text, bullets)

    # Section 13.
    append_paragraph(
        body_paragraph(doc, "Each item below is named by version 1.0 as if it were present"),
        M11_SECTION_13_TAIL,
    )
    for prefix, state in M11_M13_DEPENDENCIES.items():
        set_cell(find_row(dependencies, prefix).cells[2], state)

    # Section 14.
    decision = find_row(decisions, "10. Is Student Management entitlement-gated?")
    set_cell(decision.cells[1], M11_DECISION_10_WHY)
    set_cell(decision.cells[2], M11_DECISION_10_COST)

    # Section 15: this family lists its revisions oldest first.
    append_row(log, log.rows[8], [M11_TARGET, SHORT_DATE, M11_CHANGE_SUMMARY])
    stale = body_paragraph(doc, "\U0001f4cc What this version could not check")
    stale._p.getparent().remove(stale._p)

    finish(doc, output, title, M11_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 13 - Academic Structure
# ═════════════════════════════════════════════════════════════════════════════

M13_DIR = "13-academic-structure"
M13_STEM = "XVS_M13_Academic_Structure_Functional_Requirements_Document"
M13_SOURCE, M13_TARGET = "2.9", "2.10"

M13_STATUS = (
    "Built. Version 2.10 records three things that arrived beside the module's "
    "own routes: SchoolClass carries a class teacher, read on every class row "
    "and written by Module 12 under its own key; the academic structure and the "
    "subject list can each be loaded from a spreadsheet through the import "
    "engine, under academics.structure.import; and the platform's plan gate, "
    "under which academic structure is core for every school. The finance "
    "abstraction layer link version 2.9 recorded stands."
)

M13_VERIFIED = (
    f"{CODE_BASELINE}. Every claim this version adds was read from the code at "
    "that commit, and every test it names from the test files there."
)

M13_CLASS_TEACHER_ROW = [
    "class_teacher",
    "ForeignKey to vs_staff.StaffProfile, on_delete=SET_NULL, null=True, "
    "blank=True, related_name=\"classes_led\". The member of staff who owns the "
    "class, if one has been designated. Declared here because the designation "
    "is unique per class by construction, and a copy on an assignment row would "
    "let two people claim it. It points at the staff record rather than the "
    "login, so a class teacher who has left stops holding the class when their "
    "employment record says so, and it is SET_NULL because a class outlives "
    "whoever taught it. Written by Module 12, never by this module.",
]

M13_CLASS_TEACHER_CALLOUT = (
    "class_teacher is read here and written in Module 12. Every class row "
    "carries it as {staff_id, name}, or null where nobody holds the class, and "
    "never an email address, because the class list is read by the whole "
    "school. The list joins the staff record and its account in its own query, "
    "so a page of classes asks nothing per class. The write is PUT "
    "/v1/i/me/staff/teaching/class-teacher/ under school.teachers.assign: it "
    "resolves the class under the inclusive branch rule, so a branch "
    "administrator cannot name another branch's class and a school-wide class "
    "stays every branch's, and it refuses a class of a closed year. "
    "schools/vs_staff/tests/test_teaching.py, ClassTeacherTests and "
    "BranchReachTests."
)

M13_ASSIGNMENT_LAYER = (
    "Module 12, where the teacher half is built as TeachingAssignment under "
    "school.teachers.assign: who teaches which subject in which class, as the "
    "lead or an assistant, and nothing about when. This module neither declares "
    "nor reads it. The student half is academics.classes.assign, which gates "
    "M11's enrolment writes."
)

M13_IMPORT_KEY_ROW = [
    "academics.structure.import", "SENSITIVE", "school_admin",
    "Run the academic structure and subjects imports (FR-017). Its own key "
    "rather than .create, because one upload builds a school's whole spine, "
    "which is not the same act as adding one class through the form. It is sold "
    "with the platform's bulk import feature rather than with the structure "
    "band, which is core for every school.",
]

M13_FR006_RULE_8 = (
    " (8) A class carries its class teacher, as {staff_id, name} on the list "
    "and the detail and as null where nobody holds it, never as an email "
    "address, because the class list is one a whole school reads. This module "
    "declares the column and reads it and never writes it; Module 12 writes it, "
    "under school.teachers.assign (section 6.6)."
)

M13_FR017 = [
    ("Description",
     "Build a school's academic spine, and the subject list that hangs off it, "
     "from two spreadsheets through the platform's import engine rather than "
     "one form at a time. A secondary school is sixty classes and thirty year "
     "groups, and M11's student import refuses any row naming a class the school "
     "has not built, so this is what stands between a new school and its roll."),
    ("Actors",
     "academics.structure.import, SENSITIVE, held by school_admin by default. "
     "It is its own key rather than .create for the reason school.students.import "
     "is: one upload builds a school's whole spine. It stands in for the "
     "engine's generic import keys on both datasets through the dataset "
     "registry, and both datasets are classified as ones a school may import for "
     "itself."),
    ("Trigger",
     "The existing /v1/import/ surface, with dataset_type=academic_structure or "
     "dataset_type=subjects. This module adds no upload endpoint."),
    ("Business rules",
     "(1) Structure: one row is one class, naming its programme, level, level "
     "order, what the level promotes to, the class name, arm, capacity, branch "
     "and department. The programme and the level are created the first time a "
     "row names them, and codes are generated by the helper the screens use, so "
     "a school never invents one. (2) Because a typo creates a phantom year "
     "group rather than failing, the file is read whole and refused before "
     "anything is written for: a year group under two programmes, in the file "
     "or against what the school runs; a promotion chain that loops, reported "
     "once per loop; a year group promoting into itself, or into one the file "
     "does not contain; two rows creating one class; two year groups given one "
     "position; a missing programme, level or class; a capacity that is not a "
     "number or is larger than 200; and a branch the school does not have. (3) "
     "A year group that says nothing about where its pupils go is warned about, "
     "because a null next_level reads as unwired and M11's promotion holds it; "
     "a class that does not look like its level is warned about; and a class the "
     "school already holds is warned about and not duplicated. (4) Subjects: one "
     "row is one subject, with the year groups it is taught at in one column "
     "separated by semicolons, not commas, because \"Primary 4, 5 and 6\" is a "
     "thing a school writes. A subject is catalogue and is created once; its "
     "offerings are attached to the year groups of the year the school is "
     "running. (5) A year group the school does not run is refused rather than "
     "skipped, and so is the same subject on two rows. A subject naming no year "
     "group, a year group listed twice on one row, and a Core or Elective value "
     "the importer does not know, which becomes Core, are warned about. So is a "
     "year group the file would leave with no subject at all, which is the check "
     "no form can make: the subject screen shows what each subject covers and "
     "never what a year group is missing. (6) Each file is published whole in "
     "one transaction, and re-uploading a corrected file builds onto what is "
     "there: no second programme, level, class, subject or offering, and a held "
     "subject has year groups added rather than replaced. (7) A school with no "
     "running year is told so rather than failing. (8) Validation and execution "
     "read every row through one function, so the two passes cannot read a row "
     "differently."),
    ("Acceptance",
     "(1) One structure file builds the programme, the levels, the classes and "
     "the promotion chain, with generated codes and the capacity the file "
     "gives. (2) Each whole-file fault in rule 2 is refused, and a loop is "
     "reported once. (3) A level that says nothing about promotion is warned "
     "about, and one that says pupils leave is not. (4) Re-uploading a corrected "
     "file builds no second copy, and a class the school holds is warned about "
     "rather than duplicated. (5) A subject file creates each subject and its "
     "offerings once, refuses an unknown year group and a repeated subject, and "
     "warns about a year group left with nothing to teach. (6) A school may "
     "import its own structure and subjects, the module key reaches the wizard, "
     "and the templates carry nine and six columns. "
     "schools/vs_academics/tests/test_structure_import.py, WholeFileFaultsTests, "
     "UnwiredLevelTests, RowFaultsTests, BuildTests, OwnershipTests and "
     "SubjectImportTests."),
]

M13_SECTION_9_TAIL = (
    " The two rows at the foot of the table are the import engine's surface "
    "rather than this module's views, listed for completeness: they carry "
    "academics.structure.import and declare nothing here."
)

M13_BULLETS = [
    "A class row carries its class teacher as a staff id and a display name, "
    "or null where nobody holds it, and never an email address; clearing the "
    "designation returns null (FR-006).",
    "An academic structure file whose promotion chain loops, or that puts one "
    "year group under two programmes, is refused before anything is written, "
    "and re-uploading a corrected file builds no second copy (FR-017).",
]

M13_SECTION_12_INTRO = (
    "Each item below was named by version 1.0 as if it were present. Most now "
    "are, and those rows say so rather than disappearing, so that a reader who "
    "remembers the gap meets its answer. None of what remains is described in "
    "this document as present, and none of it should be stubbed silently."
)

M13_DEPENDENCIES = {
    "A Staff model": (
        "DISCHARGED. schools.vs_staff exists, built by Module 12, and "
        "SchoolClass.class_teacher points at its StaffProfile (section 6.6). The "
        "teacher side of the assignment layer is built there too, as "
        "TeachingAssignment under school.teachers.assign. "
        "academics.classes.assign stays unused here and is enforced by M11."
    ),
    "A Student model": (
        "DISCHARGED. schools.vs_students exists, built by M11, and owns "
        "ClassEnrolment, which points here with PROTECT and never the other way. "
        "The two readings this row asked M11 to take are taken: a null capacity "
        "means no limit, and a null next_level without is_terminal means "
        "promotion is unconfigured, which M11 holds as LEVEL_NOT_WIRED rather "
        "than graduating (M11 FR-010)."
    ),
    "A timetable module": (
        "DISCHARGED. schools.vs_calendar exists, built by M14, and points at "
        "this module's sessions, classes and subjects with PROTECT. It carries no "
        "foreign key to AcademicTerm and derives the term an event falls in from "
        "its dates, so nothing on that side protects a term."
    ),
    "AuditModuleKey.ACADEMICS": (
        "DISCHARGED. AuditModuleKey.ACADEMICS and the action types of section "
        "7.2 are registered, so the trail is kept."
    ),
    "academics.structure.* and academics.subject.*": (
        "DISCHARGED. seed_school_permissions registers academics.structure.* "
        "and academics.subject.*, with academics.structure.import beside them."
    ),
    "An academics capability": (
        "ANSWERED. Every permission key now carries a depth band, and the band "
        "map places academics.session, .classes, .subject and .structure in no "
        "capability module, which means core for every school: no plan refuses "
        "academic structure. academics.structure.import is the exception, sold "
        "with the platform's bulk import feature. Section 13, decision 10 closes "
        "on it."
    ),
}

M13_SECTION_13_OLD = (
    "The six that remain keep the numbers they have always carried, which is "
    "why the list now starts at 3 and skips 4, 5, 6 and 12."
)

M13_SECTION_13_NEW = (
    "Version 2.10 closes decision 10, and not by a ruling here: the platform's "
    "band map leaves academic structure core for every school, which is the "
    "first answer the question offered. The five that remain keep the numbers "
    "they have always carried, which is why the list now starts at 3 and skips "
    "4, 5, 6, 10 and 12."
)

M13_CHANGE_SUMMARY = (
    "Minor revision. Three things arrived beside the module's own routes and "
    "are recorded. SchoolClass carries class_teacher, a nullable foreign key to "
    "Module 12's staff record with SET_NULL, read on every class row as a staff "
    "id and a display name or null, never an email address, and written only "
    "by Module 12's class-teacher route under school.teachers.assign, which "
    "resolves the class under the inclusive branch rule and refuses a closed "
    "year; section 6.6, section 6.8 and FR-006 rule 8 carry it. FR-017 is new "
    "and carries the academic structure and subjects imports under "
    "academics.structure.import: one row per class building its programme and "
    "level as it goes, one row per subject with its year groups beside it, "
    "each file read whole and published in one transaction, refusing a looping "
    "promotion chain, a year group under two programmes and a year group the "
    "school does not run, and warning about a level whose promotion is unset "
    "and a year group left with nothing to teach. Section 13, decision 10 "
    "closes, because the platform's band map leaves academic structure core for "
    "every school, and section 12 marks the dependencies that have since been "
    "built: the staff, student and timetable modules, the audit key and the "
    "permission keys. The version 2.9 row, which had been written into the "
    "section 6.8 table rather than into this log, is moved here unchanged. No "
    "route, refusal code or audit action of this module changed, and the one "
    "model change is the class_teacher column. Backend evidence only; nothing "
    "here is deployed."
)


def patch_m13(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M13 Academic Structure Functional Requirements Document v{M13_TARGET}"
    tables = doc.tables
    cover, school_class, not_built, keys71 = tables[0], tables[20], tables[24], tables[25]
    fr006 = require_fr(tables[32], "FR-006")
    fr016 = require_fr(tables[42], "FR-016")
    api = require_header(tables[43], "Method")
    dependencies = require_header(tables[46], "Dependency")
    decisions = require_header(tables[47], "Question")
    log = require_header(tables[48], "Version")
    require_header(not_built, "Version 1.0 model")
    require_header(school_class, "SchoolClass")

    status_cell = row_by_label(cover, "Status").cells[1]
    supersedes = row_by_label(cover, "Supersedes").cells[1].text.strip()
    set_cover(cover, {
        "Version": M13_TARGET,
        "Date": REVIEW_DATE,
        "Supersedes": f"v{M13_SOURCE} (August 2026), {supersedes}",
        "Status": M13_STATUS,
        "Verified against": M13_VERIFIED,
    }, status_cell)
    append_row(cover, row_by_label(cover, "Status"), [
        "Source MRD", f"XVS Module Requirements Document v{MRD_VERSION} | Module 13",
    ])

    # Section 6.6 and the callout under its constraint block.
    active = row_by_label(school_class, "is_active")
    add_row_after(active, active, M13_CLASS_TEACHER_ROW)
    set_callout(body_paragraph(doc, "ℹ class_teacher is deliberately absent"),
                M13_CLASS_TEACHER_CALLOUT)

    # Section 6.8, including the change-log row an earlier revision put there.
    set_cell(find_row(not_built, "ClassSubjectAssignment").cells[2], M13_ASSIGNMENT_LAYER)
    teacher_row = find_row(not_built, "SchoolClass.class_teacher")
    set_cell(teacher_row.cells[1], "A ForeignKey to Staff, in version 1.0.")
    set_cell(teacher_row.cells[2],
             "Built, as a nullable foreign key to vs_staff.StaffProfile with "
             "SET_NULL, added when Module 12 shipped: section 6.6.")
    stray = find_row(not_built, M13_SOURCE)
    moved = [cell.text.strip() for cell in stray.cells]
    if moved[0] != M13_SOURCE or not moved[2].startswith("Minor revision."):
        raise ValueError("The stray version 2.9 row is not where it was expected")
    delete_row(stray)

    append_row(keys71, keys71.rows[-1], M13_IMPORT_KEY_ROW)
    append_cell(find_row(fr006, "Business rules").cells[1], M13_FR006_RULE_8)

    clone_fr_block(
        fr016._tbl, fr_heading(fr016), fr016,
        "FR-017  Import the structure and the subject list",
        "FR-017  Import the Structure and the Subject List", M13_FR017,
    )

    append_paragraph(body_paragraph(doc, "All endpoints are under /v1/academics/"),
                     M13_SECTION_9_TAIL)
    for dataset in ("academic_structure", "subjects"):
        append_row(api, api.rows[-1], [
            "(engine)", f"/v1/import/ with dataset_type={dataset}",
            "academics.structure.import", "FR-017",
        ])

    add_bullets_before(doc, "11.3 Shape and scale", M13_BULLETS)

    set_run_text(body_paragraph(doc, "Each item below is named by version 1.0 as if it were present."),
                 M13_SECTION_12_INTRO)
    for prefix, state in M13_DEPENDENCIES.items():
        set_cell(find_row(dependencies, prefix).cells[2], state)

    replace_text(
[body_paragraph(doc, "This document deliberately does not decide the following.")],
        M13_SECTION_13_OLD, M13_SECTION_13_NEW,
    )
    delete_row(find_row(decisions, "10. Is academic structure entitlement-gated?"))

    template = log.rows[-1]
    append_row(log, template, moved)
    append_row(log, template, [M13_TARGET, SHORT_DATE, M13_CHANGE_SUMMARY])

    finish(doc, output, title, M13_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 14 - Academic Calendar and Timetables
# ═════════════════════════════════════════════════════════════════════════════

M14_DIR = "14-timetable-and-calendar"
M14_STEM = "XVS_M14_Academic_Calendar_and_Timetables_Functional_Requirements_Document"
M14_SOURCE, M14_TARGET = "3.2", "3.3"

M14_STATUS = (
    "Built and verified. Version 3.3 is a reconciliation: it records every "
    "change to schools.vs_calendar between version 3.1 and commit 13675db, "
    "including the six that version 3.2's key revision did not reach. The "
    "school year and the term lifecycle remain owned upstream by M13."
)

M14_VERIFIED = (
    f"{CODE_BASELINE}. schools.vs_calendar is mounted at three prefixes under "
    "/v1/academics/ with eight models and twenty-five routes, and every claim "
    "this version adds was read from the code and the test files at that "
    "commit. The staff record earlier versions said was missing now exists in "
    "Module 12, and this module reads none of it: a teacher is still an ACTIVE "
    "grant of the tenant role keyed teacher (section 4.8). Line references "
    "inherited from earlier versions have not been re-numbered, so treat one as "
    "a pointer to a file rather than to a line."
)

M14_UPSTREAM = (
    "M13 Academic Structure FRD v2.10, which owns AcademicSession, "
    "AcademicTerm, Level, SchoolClass and Subject, and whose FR-002, FR-003 and "
    "FR-011 are the authority on the session lifecycle, the term lifecycle and "
    "the branch scope of everything this module schedules. Citations in the "
    "body name the M13 version that settled each point, which is often an "
    "earlier one."
)

M14_KEYS_INTRO_OLD = (
    "The restored half adds one resource of its own, academics.timetable, "
    "with five keys."
)

M14_KEYS_INTRO_NEW = (
    "The restored half adds two resources of its own: academics.timetable, "
    "with five keys, and academics.exam, with five more, split from the "
    "timetable at version 3.2 so that exams can be sold and granted apart from "
    "the weekly timetable."
)

M14_IMPORT_KEY_ROW = [
    "import.batches.view, .create, .run and .import", "school_admin",
    "The calendar import (FR-018). These are the import engine's own keys; no "
    "calendar key stands in for them, as school.students.import does for a "
    "students batch. Rolling a batch back is not among them, because unwinding "
    "live data is a support action.",
]

M14_FR003_CONTAINMENT = (
    "(5) An event bound to a branch may only name levels and classes that are "
    "school-wide or belong to that same branch. A school-wide event may name "
    "anything of the tenant. The refusal is 422 EVENT_AUDIENCE_OUT_OF_SCOPE and "
    "not M13's BRANCH_SCOPE_CONFLICT, because an event is not a class's parent "
    "and a school reading the message needs to know which rule it met (section "
    "10)."
)

M14_FR013_PICKER = (
    "The list of people who may fill a slot is every teacher of the tenant, "
    "and it is not narrowed by the branch of the room being filled. Narrowing "
    "it would be the obvious thing to do and would be wrong: a teacher's grants "
    "record where they may work, not where they will be needed, so a picker "
    "filtered by the room's branch would make somebody who teaches at two "
    "branches unschedulable at one of them. What makes the wide picker safe is "
    "FR-014, whose teacher clash query is wide for the same reason. FR-019's "
    "lens narrows the teacher list a screen shows, never who may be put in a "
    "slot."
)

M14_FR014_PREVIEW = [
    "Asking before saving",
    "POST /v1/academics/timetable/slots/preview/ and POST "
    "/v1/academics/exams/<id>/slots/preview/ answer what a lesson or a paper "
    "would clash with, and write nothing and audit nothing. Each builds the "
    "same unsaved row its create path builds and hands it to the same warning "
    "function, so what a school is shown before saving and what it is told "
    "after cannot disagree. Both take exclude, so an edited cell is not told it "
    "clashes with itself, and both apply rule 7's redaction. The exam preview "
    "returns a refusal separately from its warnings, because a class sitting "
    "twice or a date outside the exam period is refused outright and a form "
    "must not offer to add it anyway. Each carries the create key, "
    "academics.timetable.create and academics.exam.create, not the view key: "
    "a preview answers who is where, and under the view key it would let "
    "read-only access walk the school's staffing one request at a time.",
]

M14_FR016_ACTORS = (
    "academics.exam.create, .update, .manage and .view, the exam keys rather "
    "than the timetable's; .publish is FR-017. Exams took keys of their own at "
    "version 3.2 so they can be sold and granted apart from the weekly "
    "timetable (section 7.1). school_admin holds all four; branch_admin holds "
    "create, update and view; teacher holds view."
)

M14_FR016_RULES_TAIL = (
    " (7) The exam period named in the body must be one the caller may see "
    "under FR-002's inclusive read. Another branch's period answers 404, No "
    "such calendar entry in this year, and nothing is written: every date on "
    "an exam timetable is read from the period it hangs off, so building on "
    "another branch's period would hand that branch control of them. A "
    "school-wide period stays every branch's to build on. (8) A paper whose "
    "end time is before its start time is refused 422 EXAM_TIMES_INVALID, and "
    "a class already sitting a paper in that sitting is refused 409 "
    "CLASS_ALREADY_SITTING, with a message naming the class, the paper and the "
    "sitting; the constraints stay behind both."
)

M14_FR017_ACTORS = (
    "academics.timetable.publish for a class timetable and "
    "academics.exam.publish for an exam timetable. school_admin and "
    "branch_admin hold both, and both are SENSITIVE. Section 7.1 records that a "
    "school wanting exam publication reserved to the head office withholds the "
    "exam key in a role template of its own rather than asking for a code "
    "change."
)

M14_FR018 = [
    ("Description",
     "Load a school year's calendar from a spreadsheet through the platform's "
     "import engine: the thirty to sixty dated entries a school already keeps. "
     "It is the first dataset a school may import for itself."),
    ("Actors",
     "The engine's own import.batches.view, .create, .run and .import, which "
     "the school_admin template holds; no calendar key stands in for them. The "
     "engine's keys are sold with the platform's bulk import feature, at Plus, "
     "so a school whose plan stops at Core is refused them with "
     "PLAN_UPGRADE_REQUIRED, while a school with no package grants at all is "
     "unprovisioned rather than unentitled and is not refused. Rolling a batch "
     "back is not a school's key: unwinding live data is a support action."),
    ("Trigger",
     "The existing /v1/import/ surface, with dataset_type=calendar_events. The "
     "template's columns are name, event_type, start_date, end_date, branch, "
     "closes_school, description and applies_to."),
    ("Business rules",
     "(1) Validation and execution read every row through one resolver, so a "
     "file cannot pass with one reading and import with another. (2) "
     "References are names, resolved case-insensitively inside the school's "
     "own tenant and its running year: the branch by name, the audience as "
     "level and class names separated by semicolons. A name that resolves to "
     "nothing is refused, never defaulted to everybody, because shrugging at "
     "\"Primary 4 (Lekki)\" turns an afternoon off for one level into a "
     "school-wide closure. A branch name at a single-branch school is refused "
     "rather than ignored. (3) The event type is taken as its label or its "
     "stored code. A date outside the year, an end before the start, a date "
     "that is not a date, a missing name and a closes_school value other than "
     "yes or no are refused. (4) A batch uploaded for one branch cannot write "
     "another branch's calendar, and a row naming no branch takes the batch's. "
     "(5) The same entry twice in one file is refused; an entry already on the "
     "calendar is warned about and skipped on import rather than repeated; a "
     "date between terms and an overlap with an entry of the same kind are "
     "warned about and kept, exactly as FR-001's warnings. (6) An archived year "
     "refuses the whole file. (7) Each row is written into the uploading "
     "school's own year with its audience, and the tenant comes from the batch "
     "and nowhere else. (8) The import surface declares pending_tenant_surface "
     "and calendar_events is classified as a dataset a school may import for "
     "itself, so a school still onboarding may load its calendar."),
    ("Rollback",
     "A batch can be reversed through the engine's rollback record, which "
     "removes an entry and its audience. The reverser refuses an id that now "
     "names something else, an entry of another school, and an exam period "
     "that has since had an exam timetable built on it, because the events "
     "route refuses that delete and a rollback that went ahead would be a way "
     "round it. A skipped row is not reversed."),
    ("Acceptance",
     "(1) A good file passes, and each fault in rules 2, 3 and 6 is refused "
     "with its row and column. (2) A level and a class both resolve as an "
     "audience, one bad name among good ones is still refused, and a blank "
     "audience means everybody. (3) A branch-scoped upload cannot write another "
     "branch's calendar and fills blank rows with its own. (4) An entry already "
     "on the calendar is warned about and skipped, not repeated. (5) The "
     "uploading school is the only school a batch can write to. (6) Rollback "
     "removes the entry and its audience and refuses in each case of the "
     "rollback row. (7) The template's columns match the handler. "
     "schools/vs_calendar/tests/test_imports.py, TemplateAgreementTests, "
     "ValidationTests, SingleBranchValidationTests, ExecutionTests and "
     "RollbackTests."),
]

M14_FR019 = [
    ("Description",
     "Let a caller who covers several branches read one of them on every "
     "screen, through the branch switcher, without hiding what the whole school "
     "shares."),
    ("Actors",
     "Every list read in the module. It adds no key: which branch a caller is "
     "looking at is a choice of view, and which branches they may see is "
     "FR-002's rule, which still applies first."),
    ("Trigger",
     "?branch=<id> on GET /v1/academics/calendar/events/, "
     "/v1/academics/calendar/overview/, /v1/academics/timetable/periods/, "
     "/v1/academics/timetable/classes/, /v1/academics/timetable/teachers/, "
     "/v1/academics/timetable/rooms/ and /v1/academics/exams/."),
    ("The rule",
     "(1) The lens narrows inclusively, to the branch's own rows and the "
     "school-wide ones, so a branch's calendar still carries the school's "
     "public holidays and the everyday bell schedule stays on the list. (2) "
     "Rooms are the exception, because a room always has a branch and there is "
     "no shared row to keep. (3) An exam has no branch of its own and follows "
     "the branch of the exam period it hangs off; a school-wide period shows at "
     "every branch. (4) The teacher list narrows to the people who teach at "
     "that branch or whose account is tied to it; a teacher with neither, not "
     "yet timetabled, appears under every branch, because hiding them makes a "
     "new teacher unreachable from all of them. A teacher's week never narrows: "
     "a teacher at two branches seen through one would show empty days the "
     "other branch fills, and would be booked twice. (5) The overview's next-up "
     "list and its counts follow the lens. (6) A branch that is not this "
     "school's is refused with a validation error rather than ignored, because "
     "a filter nobody reads looks exactly like one that found nothing. (7) At a "
     "single-branch school the parameter is ignored rather than refused, "
     "because a stale tab can still carry it."),
    ("Acceptance",
     "(1) With no lens every visible branch shows; the lens drops the other "
     "branch, keeps the school-wide entries, and composes with the event list's "
     "scope facet. (2) The class picker narrows and keeps a school-wide class; "
     "exams follow their period, and school-wide periods show everywhere. (3) "
     "The bell schedule list narrows and keeps the everyday schedule. (4) The "
     "teacher list narrows to who teaches there, keeps somebody teaching at "
     "both, and keeps a teacher with no lessons yet, and a teacher's week is "
     "never narrowed. (5) The overview's next-up list, room count and event "
     "count follow the lens. (6) A foreign branch answers 400. (7) A "
     "single-branch school's parameter is ignored. "
     "schools/vs_calendar/tests/test_scoping_lens.py, one class per surface."),
]

M14_SECTION_12_INTRO = (
    "Each item below was named by version 1.0 as if it were present, or was "
    "needed by a later version and was not there. Several now are, and those "
    "rows say so rather than disappearing, so that a reader who remembers the "
    "gap meets its answer. None of what remains is described in this document "
    "as present, and none of it should be stubbed silently. The Staff row has "
    "changed character again at this version and is worth reading closely: the "
    "staff record exists, and this module still reads none of it."
)

M14_DEPENDENCIES = {
    "M13's AcademicSession, AcademicTerm, Level, SchoolClass and Subject": (
        "Built and mounted at /v1/academics/, specified by M13 v2.10. The "
        "dependency is met and stays hard: CalendarEvent.session, "
        "TimetableSlot.school_class and TimetableSlot.subject are non-null, so "
        "this module's migrations run after M13's."
    ),
    "M13's Level and SchoolClass": (
        "Built (M13 v2.10). FR-003's audience reads them."
    ),
    "AuditModuleKey.ACADEMICS": (
        "Registered, with ACADEMIC_TIMETABLE_PUBLISHED, so the trail either "
        "module writes is kept."
    ),
    "A Staff model": (
        "Built by Module 12 as schools.vs_staff: a staff record, "
        "qualifications, documents, leave and teaching assignments. This module "
        "reads none of it. A teacher is still identified by an ACTIVE grant of "
        "the tenant's teacher role (section 4.8), and no rule in section 3.5 "
        "reads a staff fact, so each is still a limit of this module. Whether "
        "those rules should now read the staff record, and refuse or warn when "
        "they do, is section 13, decision 15."
    ),
    "A Student model": (
        "Built by M11 as schools.vs_students. Nothing in this version reads it, "
        "and FR-003 stays written so that it needs none: an audience narrows a "
        "display, never a set of people."
    ),
    "An academics capability": (
        "Answered. The calendar is sold as the Calendar module: "
        "academics.calendar at Core, academics.timetable at Plus and "
        "academics.exam at Advanced. HasRBACPermission refuses "
        "PLAN_UPGRADE_REQUIRED after the role check when a school's plan does "
        "not reach the band, and the seeded catalogue ships enforcement on. "
        "Section 13, decision 11 closes."
    ),
    "The finance abstraction layer (FAL)": (
        "The FAL exists, at apps/schools/core/fal, and ties a fee structure to a "
        "session and a term (M13 FR-003). It ties no school year to a fiscal "
        "year. vs_finance.FiscalYear is a separate, entity-scoped, "
        "domain-neutral year with its own start and end dates and its own OPEN "
        "and CLOSED status, and the two calendars need not align. Nothing in "
        "this module may reference it, and nothing in vs_finance may reference "
        "a session or an event."
    ),
    "A subject-and-teacher assignment layer": (
        "Built by Module 12 as TeachingAssignment, under "
        "school.teachers.assign: who teaches which subject in which class, as "
        "the lead or an assistant, and nothing about when. TimetableSlot.teacher "
        "still points at a user and does not read it, so a slot can name a "
        "teacher who holds no assignment for that class and subject, and an "
        "assignment can name a teacher no grid schedules. The fact now lives in "
        "two tables, which is what this row warned against, and which one a slot "
        "should defer to is not decided."
    ),
    "A record of which people teach at more than one branch": (
        "Answered at version 3.1: a role assignment carries a branch, so the "
        "same teacher role at two branches is two active grants (section 13, "
        "decision 20). FR-013's picker stays tenant-wide and FR-014's clash query "
        "stays wide, for the reasons each gives."
    ),
}

M14_DECISION_11_WHY = (
    "Answered by the platform. The calendar is a module a school is sold, cut "
    "at three depths: academics.calendar at Core, academics.timetable at Plus "
    "and academics.exam at Advanced, which is why exams took keys of their own "
    "at version 3.2. The plan gate refuses a key the school's plan does not "
    "reach with PLAN_UPGRADE_REQUIRED, after the role check, and the seeded "
    "catalogue ships it on. M13's half, academic structure, is core for every "
    "school and carries no capability."
)

M14_BULLETS = {
    "11.2 Behaviour": [
        "A branch administrator cannot hang an exam off another branch's exam "
        "period: the create answers 404 and writes nothing (FR-016).",
        "A clash preview is refused to a caller holding only the view key, and "
        "where the other side of a clash is at a branch the caller cannot see, "
        "the preview names neither its class nor its room (FR-014).",
    ],
    "11.3 Shape and scale": [
        "A duplicate room name or code, a filled cell and a class sitting twice "
        "are each refused with a code and a message of their own, never the "
        "platform's generic duplicate sentence (FR-011, FR-013, FR-016).",
        "A calendar file naming a level, class or branch that resolves to "
        "nothing is refused on that row, and an entry already on the calendar "
        "is skipped rather than repeated (FR-018).",
    ],
    "12. Dependencies That Do Not Exist Yet": [
        "Read through one branch's lens, every list keeps the school-wide rows "
        "beside that branch's own, except rooms, which have none, and a "
        "teacher's week is never narrowed (FR-019).",
    ],
}

M14_CHANGE_SUMMARY = (
    "Minor revision, and a reconciliation: every change to schools.vs_calendar "
    "between version 3.1 and commit 13675db is recorded, including the six "
    "that version 3.2's key revision did not reach. FR-016, FR-017 and the API "
    "table now name the academics.exam keys that version 3.2 added to the key "
    "tables only, and FR-016 records that the exam period named in the body "
    "must be one the caller may see: another branch's period answers 404, and a "
    "school-wide period stays every branch's. FR-014 gains the two clash "
    "previews, which write nothing, reuse the warning functions the saves use, "
    "redact a branch the caller cannot see and carry the create keys. FR-018 is "
    "new and carries the calendar import, the first dataset a school may import "
    "for itself, whose names resolve inside the school's own year and are "
    "refused rather than defaulted when they resolve to nothing. FR-019 is new "
    "and carries the branch lens every list now reads, inclusive of school-wide "
    "rows except for rooms, with a teacher's week never narrowed. Five refusals "
    "that answered with the platform's generic duplicate sentence, or with a "
    "500, now carry codes and words of their own: DUPLICATE_NAME and "
    "DUPLICATE_CODE for rooms, CELL_ALREADY_FILLED, CLASS_ALREADY_SITTING and "
    "EXAM_TIMES_INVALID; naming one level twice in an audience is one narrowing "
    "rather than a refusal. Three statements version 3.1 left behind are "
    "corrected: a teacher is a role grant in FR-013, FR-015 and section 10 as "
    "well as in section 4.8, an audience refusal is EVENT_AUDIENCE_OUT_OF_SCOPE "
    "in FR-003, and a period's time refusal is PERIOD_TIME_INVALID. An unknown "
    "branch on a list answers 400, not 404. Section 13, decision 11 closes on "
    "the platform's plan gate, and section 12 records that the staff, student "
    "and structure modules exist, and that who teaches what now lives in two "
    "tables. Cells an earlier revision set in another typeface are returned to "
    "the document's own. Backend evidence only; nothing here is deployed."
)


def patch_m14(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M14 Academic Calendar and Timetables Functional Requirements "
        f"Document v{M14_TARGET}"
    )
    tables = doc.tables
    cover, existing_keys, keys71 = tables[0], tables[7], tables[28]
    fr001 = require_fr(tables[30], "FR-001")
    fr003 = require_fr(tables[32], "FR-003")
    fr011 = require_fr(tables[40], "FR-011")
    fr012 = require_fr(tables[41], "FR-012")
    fr013 = require_fr(tables[42], "FR-013")
    fr014 = require_fr(tables[43], "FR-014")
    fr015 = require_fr(tables[44], "FR-015")
    fr016 = require_fr(tables[45], "FR-016")
    fr017 = require_fr(tables[46], "FR-017")
    api = require_header(tables[47], "Method")
    refusals = require_header(tables[48], "Condition")
    dependencies = require_header(tables[50], "Dependency")
    decisions = require_header(tables[51], "Question")
    log = require_header(tables[52], "Version")
    require_header(keys71, "Key")

    status_cell = row_by_label(cover, "Status").cells[1]
    set_cover(cover, {
        "Version": M14_TARGET,
        "Date": REVIEW_DATE,
        "Supersedes": f"v{M14_SOURCE}",
        "Status": M14_STATUS,
        "Verified against": M14_VERIFIED,
        "Upstream document": M14_UPSTREAM,
    }, status_cell)
    append_row(cover, row_by_label(cover, "Status"), [
        "Source MRD", f"XVS Module Requirements Document v{MRD_VERSION} | Module 14",
    ])

    # Cells an earlier revision set in Aptos Narrow, returned to Calibri.
    restore_row_format(existing_keys.rows[7], existing_keys.rows[6])
    restore_row_format(keys71.rows[11], keys71.rows[10])
    restore_row_format(log.rows[9], log.rows[8])

    # Section 7.1.
    replace_text(
[body_paragraph(doc, "Version 2.3 said this module registers no permission key")],
                 M14_KEYS_INTRO_OLD, M14_KEYS_INTRO_NEW)
    edit_cell(find_row(keys71, "academics.timetable.view").cells[2],
              "Every read in the restored half: rooms and periods (FR-011, FR-012), the class "
              "grid (FR-013), the teacher grid (FR-015) and the exam schedule (FR-016).",
              "Every read in the timetable half: rooms and periods (FR-011, FR-012), the "
              "class grid (FR-013) and the teacher grid (FR-015). The exam schedule reads "
              "academics.exam.view.")
    edit_cell(find_row(keys71, "academics.timetable.create").cells[2],
              "Create a room, a period, a slot, an exam or an exam slot.",
              "Create a room, a period or a slot, and ask what a slot would clash with "
              "before saving it (FR-014).")
    edit_cell(find_row(keys71, "academics.timetable.manage").cells[2],
              "Delete a room, a period, a slot, an exam or an exam slot, and clear a whole "
              "class's grid.",
              "Delete a room, a period or a slot, and clear a whole class's grid.")
    set_cell(find_row(keys71, "academics.timetable.publish").cells[2],
             "Publish a class timetable (FR-017). The verb is already seeded and its "
             "seeded description names timetables by name. branch_admin holds it so "
             "that a branch's grid does not wait on the head office. An exam timetable "
             "is published under academics.exam.publish instead.")
    set_cell(find_row(keys71, "academics.exam.view").cells[2],
             "Every exam route: exams, exam papers, the paper preview (FR-014, FR-016) "
             "and exam publication (FR-017), in an Exams bundle of their own. Sold at "
             "Advanced while the class timetable is sold at Plus.")
    append_row(keys71, keys71.rows[10], M14_IMPORT_KEY_ROW)

    # Section 8.
    intro = body_paragraph(doc, "Seventeen requirements, every one of them this module's own.")
    replace_text([intro], "Seventeen requirements, every one of them",
                 "Nineteen requirements, every one of them")
    append_paragraph(intro, " Version 3.3 adds FR-018, the calendar import, and FR-019, "
                            "the branch lens every list reads through.")

    append_cell(find_row(fr001, "List and filters").cells[1],
                " The branch filter is FR-019's lens: it keeps the school-wide entries "
                "beside the branch's own, and composes with the scope facet, which "
                "narrows the screen to the school-wide entries or to one branch's.")

    append_cell(find_row(fr003, "Business rules").cells[1],
                " Naming the same level or class twice narrows it once: \"the whole of "
                "JSS1, and JSS1\" is redundant rather than invalid.")
    set_cell(find_row(fr003, "Containment").cells[1], M14_FR003_CONTAINMENT)
    edit_cell(find_row(fr003, "Acceptance").cells[1],
              "(5) An event at branch A naming a class at branch B is refused 422 "
              "BRANCH_SCOPE_CONFLICT;",
              "(5) An event at branch A naming a class at branch B is refused 422 "
              "EVENT_AUDIENCE_OUT_OF_SCOPE;")

    edit_cell(find_row(fr011, "Business rules").cells[1],
              "(4) name is unique within a branch, case-insensitively; the same name at "
              "another branch is ordinary and must not be refused. (5) code, when given, "
              "is unique per tenant.",
              "(4) name is unique within a branch, case-insensitively, refused as 409 "
              "DUPLICATE_NAME with a message naming the field, the branch and the rule; "
              "the same name at another branch is ordinary and must not be refused. (5) "
              "code, when given, is unique per tenant, refused as 409 DUPLICATE_CODE "
              "naming the room that holds it.")
    edit_cell(find_row(fr011, "Acceptance").cells[1],
              "(2) A second \"Block A Room 1\" at the same branch is refused as 400 "
              "DUPLICATE, and so is \"block a room 1\".",
              "(2) A second \"Block A Room 1\" at the same branch is refused as 409 "
              "DUPLICATE_NAME, and so is \"block a room 1\", and the refusal names the "
              "field, the branch and the rule; a code already used in the school is "
              "refused as 409 DUPLICATE_CODE naming its room (tests/test_rooms.py, "
              "RoomRuleTests).")

    edit_cell(find_row(fr012, "Business rules").cells[1],
              "refused as 422 INVALID_TIME_RANGE", "refused as 422 PERIOD_TIME_INVALID")

    fr013_rules = find_row(fr013, "Business rules").cells[1]
    edit_cell(fr013_rules,
              "(5) The teacher, when given, must be a User of the asserted tenant with "
              "user_type STAFF and status ACTIVE, refused as 422 NOT_A_TEACHING_USER;",
              "(5) The teacher, when given, must be an ACTIVE user of the asserted tenant "
              "holding an ACTIVE grant of the tenant's teacher role, refused as 422 "
              "NOT_A_TEACHING_USER;")
    edit_cell(fr013_rules,
              "enforced by the unique constraint in section 6.9 and reported as 400 "
              "DUPLICATE.",
              "refused as 409 CELL_ALREADY_FILLED, whose message names the lesson already "
              "in the cell, with the unique constraint in section 6.9 behind it.")
    set_cell(find_row(fr013, "The teacher picker is tenant-wide, deliberately").cells[1],
             M14_FR013_PICKER)
    fr013_acceptance = find_row(fr013, "Acceptance").cells[1]
    edit_cell(fr013_acceptance,
              "(3) A second slot for the same class, day and period is refused 400 "
              "DUPLICATE.",
              "(3) A second slot for the same class, day and period is refused 409 "
              "CELL_ALREADY_FILLED, and the message names the lesson already there.")
    edit_cell(fr013_acceptance,
              "(4) A teacher who is a SCHOOL_ADMIN, a PARENT or a STUDENT, or whose status "
              "is not ACTIVE, or who belongs to another tenant, is refused 422 "
              "NOT_A_TEACHING_USER, each case asserted separately.",
              "(4) A user without the teacher role and a user of another tenant are each "
              "refused 422 NOT_A_TEACHING_USER, and an administrator who also holds the "
              "teacher role can be scheduled, each asserted separately "
              "(tests/test_timetable.py, TeacherIdentityTests).")
    edit_cell(fr013_acceptance,
              "(7) A teacher pinned to Lekki by User.branch can be scheduled",
              "(7) A teacher pinned to Lekki can be scheduled")

    append_cell(find_row(fr014, "Actors").cells[1],
                " The two previews in the row below carry the create keys of the "
                "surfaces they preview.")
    cost = find_row(fr014, "Cost")
    add_row_after(cost, cost, M14_FR014_PREVIEW)
    append_cell(find_row(fr014, "Acceptance").cells[1],
                " (9) Each preview names the busy teacher or the taken room, writes "
                "nothing, agrees with the save, ignores the cell being edited and "
                "redacts a branch the caller cannot see; a reader without the create key "
                "may not preview; the exam preview reports a date outside the period and "
                "a class sitting twice as refusals rather than warnings "
                "(tests/test_previews.py, SlotPreviewTests and ExamSlotPreviewTests).")

    edit_cell(find_row(fr015, "Acceptance").cells[1],
              "(7) A user who exists in this tenant but whose user_type is not STAFF "
              "answers an empty grid rather than 404, because they may hold slots "
              "created before their type changed and hiding those would hide a real "
              "booking.",
              "(7) A user of this tenant who does not hold the teacher role answers their "
              "grid rather than 404, because they may hold slots written before the role "
              "was withdrawn, and hiding those would hide a real booking.")

    set_cell(find_row(fr016, "Actors").cells[1], M14_FR016_ACTORS)
    append_cell(find_row(fr016, "Trigger").cells[1],
                " POST /v1/academics/exams/<id>/slots/preview/ asks what a paper would "
                "clash with before saving it (FR-014).")
    append_cell(find_row(fr016, "Business rules").cells[1], M14_FR016_RULES_TAIL)
    fr016_acceptance = find_row(fr016, "Acceptance").cells[1]
    edit_cell(fr016_acceptance,
              "(3) A class given two papers in the same date and sitting is refused 400 "
              "DUPLICATE.",
              "(3) A class given two papers in the same date and sitting is refused 409 "
              "CLASS_ALREADY_SITTING, naming the class, the paper and the sitting.")
    append_cell(fr016_acceptance,
                " (9) Another branch's exam period cannot be built on: the create answers "
                "404 and writes nothing, and a school-wide period still can be "
                "(tests/test_exams.py, ExamSecurityTests). (10) End before start is "
                "refused 422 EXAM_TIMES_INVALID rather than answering 500 "
                "(ExamPaperRefusalTests).")

    set_cell(find_row(fr017, "Actors").cells[1], M14_FR017_ACTORS)

    after = clone_fr_block(
        fr017._tbl, fr_heading(fr017), fr017,
        "FR-018  Import a school calendar", "FR-018  Import a School Calendar", M14_FR018,
    )
    clone_fr_block(
        after, fr_heading(fr017), fr017,
        "FR-019  Read one branch at a time", "FR-019  Read One Branch at a Time", M14_FR019,
    )

    # Section 9.
    for method, endpoint, key in (
        ("GET, POST", "/v1/academics/exams/", "academics.exam.view / .create"),
        ("GET, PATCH, DELETE", "/v1/academics/exams/<id>/", "academics.exam.view / .update / .manage"),
        ("GET, POST", "/v1/academics/exams/<id>/slots/", "academics.exam.view / .create"),
        ("GET, PATCH, DELETE", "/v1/academics/exams/<id>/slots/<slot_id>/",
         "academics.exam.view / .update / .manage"),
        ("POST", "/v1/academics/exams/<id>/publish/", "academics.exam.publish"),
    ):
        set_cell(find_api_row(api, method, endpoint).cells[2], key)
    api_template = api.rows[3]
    add_row_after(find_api_row(api, "GET, POST", "/v1/academics/timetable/slots/"), api_template, [
        "POST", "/v1/academics/timetable/slots/preview/", "academics.timetable.create",
        "FR-014. New in 3.3. Writes nothing.",
    ])
    add_row_after(find_api_row(api, "GET, POST", "/v1/academics/exams/<id>/slots/"), api_template, [
        "POST", "/v1/academics/exams/<id>/slots/preview/", "academics.exam.create",
        "FR-014, FR-016. New in 3.3. Writes nothing.",
    ])
    append_row(api, api_template, [
        "(engine)", "/v1/import/ with dataset_type=calendar_events", "import.batches.*",
        "FR-018. The import engine's surface, listed for completeness.",
    ])
    replace_text(
[body_paragraph(doc, "\U0001f4cc Twenty-three endpoints")],
        "Twenty-three endpoints, four of them new in version 3.1, and the three",
        "Twenty-five endpoints, four of them new in version 3.1 and the two clash "
        "previews new in version 3.3, and the three",
    )

    # Section 10.
    branch_row = find_row(refusals, "An unknown or malformed ?branch=")
    set_cell(branch_row.cells[1], "400")
    set_cell(branch_row.cells[2],
             "Validation error on the branch field, from resolve_branch_reference, "
             "which collapses unknown, malformed and foreign branches into one answer so "
             "the three cannot be told apart. A single-branch school's lens ignores the "
             "parameter instead (FR-019).")
    edit_cell(find_row(refusals, "Period end_time not strictly after start_time").cells[2],
              "INVALID_TIME_RANGE", "PERIOD_TIME_INVALID")
    cell_row = find_row(refusals, "A second slot for the same class, day and period")
    set_cell(cell_row.cells[1], "409")
    set_cell(cell_row.cells[2],
             "CELL_ALREADY_FILLED, naming the lesson already in the cell. The unique "
             "constraint stays behind it, so two concurrent writes still race to the "
             "database rather than past it (FR-013).")
    edit_cell(find_row(refusals, "A teacher or invigilator who is not an ACTIVE STAFF user").cells[0],
              "A teacher or invigilator who is not an ACTIVE STAFF user of this tenant",
              "A teacher or invigilator who is not an ACTIVE user of this tenant holding "
              "an ACTIVE grant of its teacher role")
    paper_row = find_row(refusals, "A second paper for the same class, date and sitting")
    set_cell(paper_row.cells[1], "409")
    set_cell(paper_row.cells[2],
             "CLASS_ALREADY_SITTING, naming the class, the paper and the sitting, with the "
             "unique constraint behind it. This is version 1.0's student-level exam "
             "clash, re-expressed against classes (FR-016).")
    refusal_template = refusals.rows[3]
    for values in (
        ["A room name already used at the same branch", "409",
         "DUPLICATE_NAME, naming the field, the branch and the rule. The same name at "
         "another branch is ordinary (FR-011)."],
        ["A room code already used in this school", "409",
         "DUPLICATE_CODE, naming the room that holds it (FR-011)."],
        ["An exam paper whose end time is before its start time", "422",
         "EXAM_TIMES_INVALID (FR-016). Left to the check constraint, the same typo "
         "answered 500."],
        ["An exam anchored to an exam period the caller cannot see", "404",
         "Standard not-found body, No such calendar entry in this year, and nothing "
         "written (FR-016)."],
    ):
        append_row(refusals, refusal_template, values)

    # Section 11.
    for heading_text, bullets in M14_BULLETS.items():
        add_bullets_before(doc, heading_text, bullets)

    # Section 12.
    set_run_text(
        body_paragraph(doc, "Each item below is named by version 1.0 as if it were present, "
                            "or is needed by this version"),
        M14_SECTION_12_INTRO,
    )
    for prefix, state in M14_DEPENDENCIES.items():
        set_cell(find_row(dependencies, prefix).cells[2], state)

    # Section 13.
    replace_text(
[body_paragraph(doc, "Each of the following is a product question")],
        "The eleven questions version 2.3 carried are all still open and all keep their numbers",
        "The eleven questions version 2.3 carried all keep their numbers, and all but "
        "decision 11, which the platform's plan gate answers at version 3.3, are still open",
    )
    decision = find_row(decisions, "11. Is the academic calendar entitlement-gated?")
    set_cell(decision.cells[0], "11. Is the academic calendar entitlement-gated? CLOSED in version 3.3.")
    set_cell(decision.cells[1], M14_DECISION_11_WHY)

    # Section 14: this family lists its revisions oldest first.
    append_row(log, log.rows[8], [M14_TARGET, SHORT_DATE, M14_CHANGE_SUMMARY])

    finish(doc, output, title, M14_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for folder, stem, source, target, patch in (
        (M11_DIR, M11_STEM, M11_SOURCE, M11_TARGET, patch_m11),
        (M13_DIR, M13_STEM, M13_SOURCE, M13_TARGET, patch_m13),
        (M14_DIR, M14_STEM, M14_SOURCE, M14_TARGET, patch_m14),
    ):
        directory = root / folder
        patch(directory / f"{stem}_v{source}.docx", directory / f"{stem}_v{target}.docx")
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
