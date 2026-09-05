#!/usr/bin/env python3
"""Version the MRD and six FRDs for school-owned approval configuration.

A school used to arrive holding approval rules it never chose: provisioning
minted an approver role in every tenant, the seeded ladder pointed at it, and
neither the role nor the step could be removed. This records the change that
ends that, across every module it touches.
"""

from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)

REVIEW_DATE = "5 September 2026"
CHANGE_DATE = "5 Sep 2026"

MRD_SOURCE_VERSION = "2.63"
MRD_TARGET_VERSION = "2.64"

MRD_SOURCE_SCOPE = (
    "Backend change handing a school ownership of its own approval "
    "configuration. Provisioning no longer creates approver roles or shared "
    "approval steps; a tenant ladder names approver groups it composes, and an "
    "unconfigured document is refused with an audited confirmation rather than "
    "approved (5 September 2026)"
)

MRD_CHANGE_SUMMARY = (
    "Handed each school ownership of the approval rules and roles CodeX set up "
    "for it. Provisioning created an approver role in every tenant so a shared "
    "ladder could resolve, and neither that role nor the step naming it could be "
    "removed. Provisioning now creates approver groups instead, which a school "
    "composes from its own people, roles or organogram seats, and the shared "
    "platform template carries no steps at all: it still resolves, so no "
    "document is unroutable, but a school that has not built its ladder is "
    "refused with a named reason and an audited confirm-anyway rather than "
    "approved by rules it never chose. Role-sourced stages remain a first-class "
    "choice for any stage a school builds. Baseline roles are no longer locked, "
    "so a school may edit School Admin, Branch Admin and Teacher instead of "
    "creating a second role to carry one extra permission. A migration removed "
    "the unused seeded steps and the roles behind them, keeping any step a real "
    "document had already run through. Two reporting defects were closed with "
    "it: approval coverage counted role holders only and so reported nobody "
    "behind every group step, and it called a school configured whenever a "
    "template existed, which was true of exactly the schools whose every submit "
    "was refused. Backend evidence only; nothing here is deployed."
)

MRD_DELTA_SUBTITLE = "School-owned approval rules, with nothing seeded that cannot be removed"

MRD_DELTA_INTRO = (
    "This revision changes who owns a school's approval configuration across "
    "Modules 4, 7, 18, 19, 20 and 22. It adds no route and no permission, and it "
    "removes one capability entry in favour of another rather than raising the "
    "count. What a school is given at provisioning is now a starting point it can "
    "change: approver groups it composes instead of roles it did not ask for, and "
    "a shared route with no steps instead of a ladder chosen on its behalf. A "
    "document that reaches no step is refused and can be posted only on a recorded "
    "confirmation, so an unreviewed post is a decision somebody made."
)

MRD_DELTA_ROWS = [
    [
        "Seeded approver roles",
        "Withdrawn",
        "Provisioning creates approver groups, not roles. A school no longer "
        "finds roles on its roles screen that it never asked for and cannot "
        "delete, and a migration removed the ones already created where no "
        "document had used them.",
    ],
    [
        "Shared platform ladder",
        "Carries no steps",
        "One shared row cannot name a tenant's approver group, so it names "
        "nobody. It still resolves, so no document is unroutable.",
    ],
    [
        "Unconfigured approval",
        "Refused, not assumed",
        "A document resolving to a stageless ladder is refused with "
        "ApprovalNotConfiguredError and may proceed only on an explicit "
        "confirmation recorded against whoever gives it.",
    ],
    [
        "Role-sourced stages",
        "Unchanged",
        "A school may still build a stage approved by the holders of a named "
        "role. What was withdrawn is CodeX choosing that role on its behalf.",
    ],
    [
        "Baseline school roles",
        "Editable",
        "School Admin, Branch Admin and Teacher are no longer locked. Which "
        "permissions a school may hold is still bounded by tenant scope and "
        "enforced at the grant.",
    ],
    [
        "Approval coverage report",
        "Two defects closed",
        "It now resolves approver groups through the engine's own lookup "
        "instead of reporting nobody behind every group step, and counts a "
        "document type as configured only when a step will actually route it.",
    ],
]


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_cover_version(table, source: str, target: str) -> None:
    """Set the cover's version label to ``target``, whatever it currently reads.

    Matching on the outgoing version is not enough. A cover can carry a stale
    label - M18 v1.7 shipped with "Version: 1.6" on its cover - and a
    search-and-replace for the version we are leaving would find nothing and
    silently ship the staleness forward. So the label is found by its own word
    and rewritten, and its absence is an error rather than a no-op.
    """
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if "Version:" in run.text:
                run.text = re.sub(r"Version:\s*[\d.]+", f"Version: {target}", run.text)
                return
    raise ValueError(f"No cover version label found (expected to set {target}).")


def update_control_table(table, *, version, source_scope, review_date=REVIEW_DATE,
                         mrd_baseline=None) -> None:
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], review_date, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], source_scope, size=9)
        elif label == "MRD baseline" and mrd_baseline:
            replace_cell(row.cells[1], mrd_baseline, size=9)


def prepend_change_log(table, version: str, date: str, summary: str) -> None:
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    new_row = table.rows[1]
    replace_cell(new_row.cells[0], version, size=8)
    replace_cell(new_row.cells[1], date, size=8)
    replace_cell(new_row.cells[2], summary, size=8)


def rewrite_capability(table, old_fragment: str, new_text: str) -> bool:
    """Rewrite the one capability bullet containing ``old_fragment``.

    A capability is found by its text rather than by a coordinate, because a
    coordinate shifts whenever a row is added above it. The rewrite lands on the
    matching *paragraph*: a single cell often carries several bullets, and
    replacing the cell would silently take its siblings with it.
    """
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if old_fragment in paragraph.text:
                    write_paragraph(paragraph, new_text, size=8.2)
                    return True
    return False


def rewrite_block(table, old_fragment: str, new_text: str) -> bool:
    """Rewrite a NEEDS ATTENTION or CURRENT DECISION cell found by its text."""
    for row in table.rows:
        # Merged cells repeat across ``row.cells``, so the element itself is the
        # key. ``id()`` is not: python-docx builds a fresh proxy per access and
        # CPython recycles the id of the one just collected.
        seen = set()
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            if old_fragment in cell.text:
                replace_cell(cell, new_text, size=8.2)
                return True
    return False


M04_DECISION = (
    "CURRENT DECISION\n"
    "• No material capability gap was identified within this module's stated backend scope.\n"
    "• A school owns the roles CodeX set up for it. The baseline roles arrived locked, so a "
    "school wanting its administrator to hold one extra permission had to create a second role "
    "and assign it alongside. Which permissions a school may hold at all is already bounded by "
    "tenant scope and enforced on the grant models, so locking the role on top of that decided "
    "on the school's behalf that the first guess was final. The system-role flag still records "
    "where a role came from and still groups the roles screen; it no longer means read-only.\n"
    "• Two role-grant escalations are closed at their source. Scope refuses a tenant carrying "
    "a platform-only key. Separately, a restricted key is refused on direct role writes, in "
    "permission groups, and as an ALLOW override; it may enter a custom role only through a "
    "change request decided by somebody other than the requester, and that reviewer may grant "
    "only restricted keys they already hold. Existing unverified restricted grants are removed "
    "from custom roles by migration, while trusted seeded system roles remain as the approval "
    "bootstrap.\n"
    "• A role granted for the whole school now reaches the whole school. The branch narrowing "
    "and the permission gate are resolved from the same grants, so a Finance Officer for the "
    "whole school is no longer narrowed to the one branch her staff record happens to name while "
    "the gate lets her act at the others. Writes follow reads: what she raises without naming a "
    "branch is filed school-wide.\n"
    "• The same role may now be granted at two branches. The schema always allowed it and both "
    "write paths refused it, so a teacher working at two branches could hold the role at one. "
    "Every count that reads assignments counts people rather than grant rows.\n"
    "• Deactivation is a runtime kill switch. A permission works only while its own row and its "
    "module, resource and action are active; a group grant additionally requires an active group. "
    "Evaluation, overrides and permission-based routing share that rule, and a durable registry "
    "revision invalidates warm permission snapshots across application workers.\n"
    "• Role access writes now have one authoritative boundary. Direct edits, approved requests, "
    "prebuilt provisioning and role suggestions all lock the role and write one reasoned "
    "before-and-after RBAC audit in the same transaction, so a missing audit rolls the access "
    "change back.\n"
    "• Fresh school installations carry the school_admin and branch_admin prebuilt templates "
    "from a school-app data migration. Existing template rows, including deliberately deactivated "
    "ones, are preserved. Missing or unusable required templates now refuse school or branch "
    "creation instead of leaving a parent record with no administrator."
)

M07_NEEDS = (
    "NEEDS ATTENTION\n"
    "• An approval route with no steps is refused rather than treated as approval. A tenant "
    "that has not built its ladder resolves to the shared platform row, which carries no steps, "
    "and submission raises a named refusal that the submitter may confirm past. The confirmation "
    "is recorded against whoever gives it, so an unreviewed post is a decision somebody made "
    "rather than a silence.\n"
    "• Continue-without-approval remains available only where a document handler permits it. "
    "Payouts forbid it; permitted document types still rely on audit rather than a second "
    "reviewer.\n"
    "• Timeouts, escalation timers, and template versioning remain future enhancements, not "
    "completed features."
)

M22_NEEDS = (
    "NEEDS ATTENTION\n"
    "• Approval fails closed for controlled spend and nothing is imposed to achieve it. A "
    "tenant ladder is published when its books are created, tenant-scoped and branch-aware, and "
    "each step names an approver group created empty, so the first document parks rather than "
    "skipping. A repair pass frees parked work once somebody joins the group, with no "
    "resubmission. A tenant with no ladder of its own resolves to the shared platform row, which "
    "carries no steps, and is refused with a named reason rather than approved.\n"
    "• Approval coverage reports the gaps it exists to report. It resolves approver groups "
    "through the engine's own lookup rather than counting role holders only, and counts a "
    "document type as configured only when a step will actually route it, so a tenant with no "
    "ladder reads as unconfigured instead of as configured with no gaps.\n"
    "• A requisition raised by a holder of a school-wide role who names no branch is filed "
    "school-wide, and school-wide spend is deliberately hidden from branch-pinned staff. That is "
    "the intended reading of a purchase for the school as a whole, and it means such a "
    "requisition does not appear on the raiser's own branch queue."
)


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = f"XVS Module Requirements Document v{MRD_TARGET_VERSION}"
    doc.core_properties.version = MRD_TARGET_VERSION

    replace_cover_version(doc.tables[0], MRD_SOURCE_VERSION, MRD_TARGET_VERSION)
    update_control_table(
        doc.tables[1], version=MRD_TARGET_VERSION, source_scope=MRD_SOURCE_SCOPE,
    )

    # Contents: the delta section is named for the version it belongs to.
    for row in doc.tables[2].rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(row.cells[0], f"5. v{MRD_TARGET_VERSION} Capability Delta",
                         size=9, bold=True, color=BLUE)
            replace_cell(row.cells[1], MRD_DELTA_SUBTITLE, size=9)

    # Module 4: the roles a school starts with are now its own to shape.
    assert rewrite_capability(
        doc.tables[14], "Locked and protected roles",
        "▸  Platform-locked roles, with school roles school-owned",
    ), "M04 locked-roles capability not found"
    assert rewrite_block(doc.tables[14], "Two role-grant escalations", M04_DECISION), \
        "M04 decision block not found"

    # Module 7: resolution is no longer limited to roles provisioning invented.
    assert rewrite_capability(
        doc.tables[20], "Approver resolution limited to provisioned approval roles",
        "▸  Tenant-composed approver groups, with no role provisioned to fill them",
    ), "M07 provisioned-roles capability not found"
    assert rewrite_capability(
        doc.tables[20], "Fail-closed handling when no approver resolves",
        "▸  Fail-closed handling when no approver resolves, and when no step exists",
    ), "M07 fail-closed capability not found"
    assert rewrite_block(doc.tables[20], "Continue-without-approval remains available",
                         M07_NEEDS), "M07 needs block not found"

    # Module 22: the seeded ladder no longer names anything the school did not choose.
    assert rewrite_block(doc.tables[54], "the seeded stages park rather than skip",
                         M22_NEEDS), "M22 needs block not found"

    heading_seen = False
    intro_written = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == f"5. v{MRD_SOURCE_VERSION} Capability Delta":
            write_paragraph(paragraph, f"5. v{MRD_TARGET_VERSION} Capability Delta",
                            size=13, bold=True, color=BLUE)
            heading_seen = True
            continue
        # The paragraph under the heading introduces the delta, so it describes the
        # previous release until it is rewritten and then contradicts the table
        # directly beneath it.
        if heading_seen and not intro_written and text:
            write_paragraph(paragraph, MRD_DELTA_INTRO, size=9)
            intro_written = True
    assert heading_seen and intro_written, "delta heading or intro not found"

    delta = doc.tables[76]
    replace_cell(delta.rows[0].cells[0], f"v{MRD_TARGET_VERSION} capability delta",
                 size=8.2, bold=True, color="FFFFFF", fill=BLUE)
    while len(delta.rows) > 1:
        delta._tbl.remove(delta.rows[-1]._tr)
    template_missing = True
    for values in MRD_DELTA_ROWS:
        row = delta.add_row()
        for idx, value in enumerate(values):
            replace_cell(row.cells[idx], value, size=8.2)
        template_missing = False
    assert not template_missing, "delta rows not written"

    prepend_change_log(doc.tables[78], MRD_TARGET_VERSION, CHANGE_DATE, MRD_CHANGE_SUMMARY)

    doc.save(str(output))
    update_extended_title(output, f"XVS Module Requirements Document v{MRD_TARGET_VERSION}")
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    print(f"MRD written: {output}")



# --------------------------------------------------------------------------- #
# Module 7: Workflow & Approval Engine                                         #
# --------------------------------------------------------------------------- #

M07_SOURCE_VERSION = "1.7"
M07_TARGET_VERSION = "1.8"

M07_SOURCE_SCOPE = (
    "Working tree, 5 September 2026 (pending commit)"
)

M07_FR005_EVIDENCE = (
    "A ROLE stage names a role by key and the engine reads that role's active "
    "assignments in the tenant that raised the request, so one definition serves "
    "every tenant that holds the key. Only a role carrying the provisioning mark "
    "is resolved: ensure_approver_role sets it, and the resolver requires it. "
    "Nothing seeds through this path any more. Provisioning gives a tenant "
    "approver groups instead, so a ROLE stage exists only where a tenant built "
    "one, and the mark then guards a key that tenant chose."
)

M07_FR005_LIMIT = (
    "The mark is required because a role key is derived from the name its "
    "creator typed, so an unmarked look-alike must confer nothing. A tenant that "
    "creates a role by hand and points a stage at it therefore has to obtain the "
    "mark through ensure_approver_role rather than by naming the role well. "
    "manage.py audit_approver_roles lists roles the mark was never applied to, "
    "so each can be reviewed rather than silently bleeding authority. A group's "
    "role member needs no mark: it is stored as a reference to one chosen role, "
    "not matched by a name anything could share."
)

M07_FR031 = {
    "header": "FR-031 | Implemented",
    "requirement": (
        "An approval route carrying no steps must never produce an approved "
        "outcome on its own. The tenant is asked to decide rather than answered "
        "for."
    ),
    "evidence": (
        "submit_for_approval raises ApprovalNotConfiguredError (409, "
        "APPROVAL_NOT_CONFIGURED) when the resolved template has no stages, "
        "after rolling back the pending flip so a refused submit writes nothing. "
        "A caller may retry with an explicit confirmation, which terminates the "
        "instance as approved and records who confirmed and why through the "
        "POSTED_WITHOUT_APPROVAL audit action. The finance direct-post gate "
        "applies the same rule at its own boundary."
    ),
    "acceptance": (
        "A tenant that has published no ladder of its own resolves to the shared "
        "platform template, which carries no steps, and its first submission is "
        "refused with a named reason rather than approved. The same submission "
        "sent with a confirmation succeeds and leaves an audit event naming the "
        "confirmer. An empty ladder is never treated as consent."
    ),
    "limit": (
        "The confirmation is a single person's decision recorded after the fact, "
        "not a second reviewer. Document handlers that forbid release, payouts "
        "among them, refuse it outright."
    ),
}

M07_GAPS = (
    "FURTHER GAPS\n"
    "• Continue-without-approval remains a submitter-owned release for handlers that permit "
    "it. Payouts forbid it, but the generic default is permissive and other document types "
    "still rely on audit rather than a second reviewer. The same holds for confirming past an "
    "unconfigured route.\n"
    "• The parked-document repair runs on read, not on a schedule, so a parked document nobody "
    "opens stays parked.\n"
    "• No timeouts, escalation timers, or reminders for a stage that has been waiting.\n"
    "• No template versioning. A republish changes the definition future activations use.\n"
    "• Six of the ten defined notification event keys are not wired.\n"
    "• Route graphs are not checked for cycles at publish.\n"
    "• Condition field paths are not validated against the document type."
)

M07_CHANGE = (
    "Recorded the move of seeded approval configuration off roles and onto "
    "approver groups. Provisioning no longer creates a role in every tenant for "
    "a shared ladder to resolve through, so FR-005 now describes a role-sourced "
    "stage as something a tenant builds rather than something it is given, and "
    "its limit is stated in those terms. New FR-031 covers the case that "
    "created: a resolved template with no steps is refused with "
    "APPROVAL_NOT_CONFIGURED and proceeds only on a confirmation recorded "
    "against whoever gives it, which the error contract and the audit action "
    "both carry. The parked-approval repair memoises group-sourced stages on the "
    "same terms as role-sourced ones, so a page of parked documents costs one "
    "lookup per stage configuration rather than one per document."
)


def add_fr_table(doc, after_table, fr: dict, *, heading: str):
    """Clone an existing requirement section and fill it.

    Requirement sections are visually identical, so the new one is a deep copy of
    its neighbour rather than a fresh build: that inherits the borders, shading,
    column widths and heading style without restating any of them here. The
    heading paragraph is copied too. Without it the new requirement renders glued
    to the bottom of its predecessor, which reads as more of that requirement
    rather than as a new one.

    The header cell keeps the shading it inherited. Writing a fill here would
    override the pale green every other requirement header uses and single this
    one out for no reason.
    """
    from docx.table import Table

    body = doc.element.body
    children = list(body.iterchildren())
    anchor = children.index(after_table._tbl)
    heading_element = None
    for index in range(anchor - 1, max(anchor - 4, -1), -1):
        candidate = children[index]
        if candidate.tag.endswith("}p"):
            from docx.text.paragraph import Paragraph

            if Paragraph(candidate, doc).style.name.startswith("Heading"):
                heading_element = candidate
                break
    assert heading_element is not None, "no heading found above the requirement table"

    new_tbl = copy.deepcopy(after_table._tbl)
    after_table._tbl.addnext(new_tbl)
    new_heading = copy.deepcopy(heading_element)
    after_table._tbl.addnext(new_heading)

    from docx.text.paragraph import Paragraph

    # Retitle the cloned heading by editing its first run and dropping the rest.
    # Rebuilding the paragraph would discard what makes it look like the headings
    # around it: the Heading 2 colour and the rule underneath.
    cloned = Paragraph(new_heading, doc)
    assert cloned.runs, "cloned heading carries no run to retitle"
    cloned.runs[0].text = heading
    for run in cloned.runs[1:]:
        run._r.getparent().remove(run._r)

    table = Table(new_tbl, after_table._parent)
    replace_cell(table.rows[0].cells[0], fr["header"], size=9, bold=True)
    labels = {
        "Requirement": fr["requirement"],
        "Current evidence": fr["evidence"],
        "Acceptance": fr["acceptance"],
        "Current limit": fr["limit"],
    }
    for row in table.rows[1:]:
        label = row.cells[0].text.strip()
        if label in labels:
            replace_cell(row.cells[1], labels[label], size=8.2)
    return table


def patch_m07(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = (
        f"XVS M07 Workflow and Approval Engine Functional Requirements Document "
        f"v{M07_TARGET_VERSION}"
    )
    doc.core_properties.version = M07_TARGET_VERSION

    replace_cover_version(doc.tables[0], M07_SOURCE_VERSION, M07_TARGET_VERSION)
    for row in doc.tables[1].rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], M07_TARGET_VERSION, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Code baseline":
            replace_cell(row.cells[1], M07_SOURCE_SCOPE, size=9)
        elif label == "Source MRD":
            replace_cell(
                row.cells[1],
                f"XVS Module Requirements Document v{MRD_TARGET_VERSION} | Module 7",
                size=9,
            )

    # FR-005: a role-sourced stage is now something a tenant builds, not one it is given.
    fr005 = doc.tables[13]
    for row in fr005.rows:
        label = row.cells[0].text.strip()
        if label == "Current evidence":
            replace_cell(row.cells[1], M07_FR005_EVIDENCE, size=8.2)
        elif label == "Current limit":
            replace_cell(row.cells[1], M07_FR005_LIMIT, size=8.2)

    # The refusal earns its own error-contract row.
    errors = doc.tables[44]
    template_row = errors.rows[-1]
    new_tr = copy.deepcopy(template_row._tr)
    template_row._tr.addnext(new_tr)
    added = errors.rows[-1]
    replace_cell(added.cells[0], "APPROVAL_NOT_CONFIGURED", size=8.2)
    replace_cell(added.cells[1],
                 "The resolved template has no steps and no confirmation was given.",
                 size=8.2)
    replace_cell(added.cells[2], "409", size=8.2)

    # FR-031, immediately after the last requirement table.
    add_fr_table(doc, doc.tables[38], M07_FR031,
                 heading="FR-031 Refuse an Approval Route That Has No Steps")

    assert rewrite_block(doc.tables[47], "Continue-without-approval remains a submitter-owned",
                         M07_GAPS), "M07 gaps block not found"

    # Traceability: the renamed capability, and the new requirement it maps to.
    trace = doc.tables[48]
    for row in trace.rows:
        if row.cells[0].text.strip() == "Approver resolution limited to provisioned approval roles":
            replace_cell(row.cells[0],
                         "Tenant-composed approver groups, with no role provisioned to fill them",
                         size=8.2)
            replace_cell(row.cells[1], "FR-005, FR-006, FR-022", size=8.2)
        elif row.cells[0].text.strip() == "Fail-closed handling when no approver resolves":
            replace_cell(row.cells[0],
                         "Fail-closed handling when no approver resolves, and when no step exists",
                         size=8.2)
            replace_cell(row.cells[1], "FR-015, FR-031", size=8.2)

    prepend_change_log(doc.tables[49], M07_TARGET_VERSION, CHANGE_DATE, M07_CHANGE)

    doc.save(str(output))
    update_extended_title(output, doc.core_properties.title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    print(f"M07 FRD written: {output}")




# --------------------------------------------------------------------------- #
# Module 4: Roles & Permissions (RBAC)                                         #
# --------------------------------------------------------------------------- #

M04_SOURCE_VERSION = "1.12"
M04_TARGET_VERSION = "1.13"

M04_SOURCE_SCOPE = (
    "Backend change handing a school the roles CodeX provisioned for it: the "
    "baseline roles arrive unlocked and may be edited within tenant scope, while "
    "deletion stays refused and platform roles stay locked (5 September 2026)"
)

M04_CODE_INSPECTED = (
    "apps/vs_rbac/services.provision_role_from_prebuilt and "
    "create_role_from_suggestion, apps/vs_rbac/views TenantRoleTemplate detail "
    "update and delete, and the vs_rbac data migration unlocking school-owned "
    "roles, against the RBAC suite"
)

M04_FR021_EVIDENCE = (
    "PrebuiltRoleTemplate rows are read-only library definitions. "
    "provision_role_from_prebuilt copies one into a tenant-owned, unlocked "
    "system role and records access through set_role_access. Owning a role means "
    "being able to change it: the copy arrives editable, and the tenant may "
    "rename it and adjust its permissions within tenant scope without creating a "
    "second role to carry the difference. The school-app migration now guarantees "
    "school_admin and branch_admin template rows on fresh installations because "
    "School and Branch creation require them. It uses get_or_create and does not "
    "reactivate or overwrite an existing row. The broader seed command remains "
    "responsible for converging descriptions, tier, scope, teacher, and "
    "permission defaults."
)

M04_FR021_ACCEPTANCE = (
    "A fresh migrated school installation contains school_admin and branch_admin "
    "without an operator seed step. The seed command converges the complete "
    "three-template catalogue, including teacher and access defaults. Adoption "
    "remains idempotent, tenant-owned, editable, undeletable, and durably "
    "audited: an update by the owning tenant succeeds, and a delete is still "
    "refused."
)

M04_TRACE_EVIDENCE = (
    "Implemented. is_system_role and is_locked block deletion of provisioned "
    "roles, and platform roles stay locked. A school's own roles are unlocked and "
    "editable: the flags record provenance and protect the role from removal, not "
    "from change. What a tenant may grant is bounded by permission scope and "
    "enforced at the grant."
)

M04_CHANGE = (
    "Stopped a school's own baseline roles being read-only. School Admin, Branch "
    "Admin and Teacher were provisioned locked and the update endpoint refused "
    "them, so a school wanting its administrator to hold one further permission "
    "had to create a second role and assign it alongside. Which permissions a "
    "school may hold is already bounded by tenant scope and enforced on the grant "
    "models, so the lock added nothing except the workaround. Provisioning now "
    "creates the copy unlocked, a migration unlocked the school roles already "
    "created, and the update refusals are gone. Deletion is unchanged and still "
    "refused, and platform roles remain locked. FR-021 and the Module 4 "
    "traceability entry are restated in those terms."
)


def patch_m04(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = (
        f"XVS M04 Roles and Permissions RBAC Functional Requirements Document "
        f"v{M04_TARGET_VERSION}"
    )
    doc.core_properties.version = M04_TARGET_VERSION

    replace_cover_version(doc.tables[0], M04_SOURCE_VERSION, M04_TARGET_VERSION)
    for row in doc.tables[1].rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], M04_TARGET_VERSION, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], M04_SOURCE_SCOPE, size=9)
        elif label == "Code inspected":
            replace_cell(row.cells[1], M04_CODE_INSPECTED, size=9)
        elif label == "MRD baseline":
            replace_cell(
                row.cells[1],
                f"XVS Module Requirements Document v{MRD_TARGET_VERSION}, Module 4, "
                "eighteen capability entries",
                size=9,
            )

    fr021 = doc.tables[29]
    for row in fr021.rows:
        label = row.cells[0].text.strip()
        if label == "Current evidence":
            replace_cell(row.cells[1], M04_FR021_EVIDENCE, size=8.2)
        elif label == "Acceptance":
            replace_cell(row.cells[1], M04_FR021_ACCEPTANCE, size=8.2)

    for row in doc.tables[44].rows:
        if row.cells[0].text.strip() == "Locked and protected roles":
            replace_cell(row.cells[0],
                         "Platform-locked roles, with school roles school-owned", size=8.2)
            replace_cell(row.cells[2], M04_TRACE_EVIDENCE, size=8.2)

    prepend_change_log(doc.tables[46], M04_TARGET_VERSION, CHANGE_DATE, M04_CHANGE)

    doc.save(str(output))
    update_extended_title(output, doc.core_properties.title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    print(f"M04 FRD written: {output}")




# --------------------------------------------------------------------------- #
# The four domain FRDs                                                         #
# --------------------------------------------------------------------------- #

def patch_domain_frd(source: Path, output: Path, *, title: str, module_number: int,
                     source_version: str, target_version: str, code_baseline: str,
                     edits, change_summary: str, changelog_index: int) -> None:
    """Version one consumer FRD and apply its cell edits.

    ``edits`` is a sequence of ``(table_index, target, text)``. ``target`` is a
    row label for a requirement table, ``None`` for a whole block cell, or a
    ``(row, column)`` pair for a grid cell that carries no label. Every edit is
    applied before anything is inserted, so the indices stay the ones that were
    read off the source.
    """
    doc = Document(str(source))
    doc.core_properties.title = title
    doc.core_properties.version = target_version

    replace_cover_version(doc.tables[0], source_version, target_version)
    for row in doc.tables[1].rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], target_version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Code baseline":
            replace_cell(row.cells[1], code_baseline, size=9)
        elif label == "Source MRD":
            replace_cell(
                row.cells[1],
                f"XVS Module Requirements Document v{MRD_TARGET_VERSION} | "
                f"Module {module_number}",
                size=9,
            )

    for table_index, row_label, text in edits:
        table = doc.tables[table_index]
        if row_label is None:
            replace_cell(table.rows[0].cells[0], text, size=8.2)
            continue
        if isinstance(row_label, tuple):
            row_index, column_index = row_label
            replace_cell(table.rows[row_index].cells[column_index], text, size=8.2)
            continue
        applied = False
        for row in table.rows:
            if row.cells[0].text.strip() == row_label:
                replace_cell(row.cells[1], text, size=8.2)
                applied = True
                break
        assert applied, f"{row_label} not found in T{table_index} of {source.name}"

    prepend_change_log(doc.tables[changelog_index], target_version, CHANGE_DATE,
                       change_summary)

    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    print(f"FRD written: {output}")


M18_EDITS = [
    # The stages name groups now, so the actor row must not send an administrator
    # looking for a role to assign.
    (6, (5, 2), "Membership of the payout-approvers and payout-senior-approvers "
                "approver groups, which the tenant composes"),
    (24, "Current evidence",
     "New entity provisioning publishes a tenant ladder whose stages name approver "
     "groups it creates empty. Migration 0006 published a two-stage platform "
     "fallback for entities that predate that provisioner; the corrective vs_rbac "
     "migration then removed those stages, because a shared row cannot name a "
     "tenant's approver group and the roles it used to name are gone. The platform "
     "row itself remains, so no entity is unroutable. payout_approval_health asks "
     "the same workflow resolver used by submission and reports every active entity "
     "that still has no route."),
    (24, "Acceptance",
     "After migration, every active ledger entity resolves either its tenant policy "
     "or the platform row. Re-running the migration is idempotent, and a platform "
     "policy an administrator owns is left alone. The release check exits "
     "successfully only when no active entity is uncovered."),
    (24, "Current limit",
     "Resolving is not approving. The platform row supplies routing and no steps, so "
     "an entity relying on it is refused with APPROVAL_NOT_CONFIGURED until it "
     "publishes its own ladder, and may proceed only on a recorded confirmation. A "
     "tenant ladder supplies steps but no people: each stage stays blocked until "
     "somebody joins the approver group it names."),
    (37, None,
     "CURRENT PAYOUT CONTROL - FAIL CLOSED\n"
     "• Approval is not optional by route. Every single or bulk payout enters Module 7 and "
     "only the terminal callback may ask the provider to transfer.\n"
     "• New books receive tenant policy, whose stages name approver groups created empty. "
     "The shared platform row carries no steps, so an entity that has published no ladder of "
     "its own is refused rather than approved, and payouts forbid continuing without "
     "approval outright.\n"
     "• payout_approval_health names any active entity that still cannot resolve a standard "
     "ladder and fails the release check. Missing approver appointments are a separate "
     "parked-state health concern.\n"
     "• The default threshold is N500,000: one checker below it, and a distinct senior "
     "checker at or above it.\n"
     "• Approval and the transfer are separate transactions. The approving vote commits "
     "before the provider is called, so a rollback during approval cannot leave money sent, "
     "and a guard refuses any provider call made with a transaction open.\n"
     "• An instruction is claimed before its own send and only a clean provider rejection "
     "marks it failed, so no retry re-sends it and no timeout records moved money as a "
     "failure."),
]

M18_CHANGE = (
    "Moved payout approval off roles CodeX created and onto approver groups the "
    "tenant composes. A tenant ladder still arrives blocked, but it is blocked by "
    "an empty group rather than by an unheld role nobody asked for, and no "
    "payout-approver or payout-senior-approver role is created in any tenant. The "
    "shared platform row no longer carries the two-stage ladder migration 0006 "
    "published: a shared row cannot name a tenant's group, so it carries the "
    "document type only. It still resolves, so no entity is unroutable, and an "
    "entity relying on it is refused with APPROVAL_NOT_CONFIGURED rather than "
    "approved. FR-020 is restated accordingly, and its limit now separates "
    "routing from steps from people."
)


M22_EDITS = [
    (6, (4, 2), "Publish this tenant's approval ladders, compose the approver groups "
                "their stages name, and read the coverage report"),
    (11, "Current evidence",
     "ensure_tenant_approval_templates publishes a tenant-scoped ladder per document "
     "type, non-destructively. Each stage names an approver group the same call "
     "creates empty, so the ladder arrives closed without inventing a role the tenant "
     "never asked for and cannot delete. ensure_default_approval_templates publishes "
     "the shared platform row, which carries no stages at all: one row every tenant "
     "runs cannot name a group that belongs to one of them. seed_procurement_approvals "
     "is the operational half, and an endpoint offers the same to an administrator."),
    (11, "Acceptance",
     "Re-running after an administrator customised a ladder reports it and leaves it "
     "alone; only the platform row upserts. A tenant's template wins over the platform "
     "row through the engine's own cascade, and a tenant with no template of its own "
     "resolves to a row with no steps and is refused rather than approved. The seed "
     "runs inside the transaction that creates a tenant's books, registered through "
     "finance's entity provisioning, so the gate arrives with the chart of accounts "
     "rather than with a remembered command."),
    (11, "Current limit",
     "Only tenants created after the change are covered automatically; "
     "seed_procurement_approvals remains the route for the ones before it. A tenant "
     "that runs neither is not silently open: it resolves to the stageless platform "
     "row and is refused with APPROVAL_NOT_CONFIGURED."),
    (13, "Acceptance",
     "Adding somebody to the approver group the stage names releases the document with "
     "no resubmission. The repair leaves a staffed stage untouched, and it never "
     "crosses a tenant boundary."),
    (15, "Current evidence",
     "The coverage report projects two things the engine already owns: the template the "
     "engine would resolve for that scope through its own cascade, and the people, "
     "resolved through the engine's own lookups - role holders for a role-sourced "
     "stage, and live group membership for a group-sourced one. It counts a document "
     "type as configured only when a step will actually route it, so a tenant "
     "resolving to the stageless platform row reads as unconfigured rather than as "
     "configured with no gaps."),
    (15, "Acceptance",
     "Because it calls the same lookups routing calls, the report and live routing "
     "cannot disagree about who is eligible, whether the stage names a role or a "
     "group. Gaps are named as gaps, per branch and per stage. Only display-safe "
     "fields are returned; a person's email is never exposed."),
    (30, (3, 1),
     "Supplies the keys the views enforce, including the two CRITICAL overrides. It "
     "supplies no approving roles: the seeded stages name approver groups, and who is "
     "in them is the tenant's decision."),
]

M22_CHANGE = (
    "Moved seeded procurement approval off roles CodeX created and onto approver "
    "groups the tenant composes, and closed two defects in the coverage report "
    "with it. A tenant ladder still arrives closed, but through an empty group "
    "rather than an unheld procurement-approver role no tenant asked for and none "
    "could delete. The shared platform row carries no stages, because one row "
    "every tenant runs cannot name a group belonging to one of them; it still "
    "resolves, so nothing is unroutable, and a tenant relying on it is refused "
    "with APPROVAL_NOT_CONFIGURED rather than approved. FR-008 gained the two "
    "corrections: coverage counted role holders only and so reported nobody behind "
    "every group-sourced stage, and it called a document type configured whenever "
    "a template existed, which was true of exactly the tenants whose every submit "
    "was refused."
)

M19_EDITS = [
    (26, "Current evidence",
     "register_entity_provisioner runs registered callables inside the transaction "
     "that creates a ledger entity, alongside the currencies, chart of accounts and "
     "fiscal periods. Finance, procurement and payments each register their own "
     "approval ladders, whose stages name approver groups the same call creates "
     "empty, so a tenant's books arrive with closed rules rather than with roles it "
     "never asked for. The finance abstraction layer provisions a school's books "
     "through this path rather than creating a ledger entity itself, which is the "
     "difference between books a school can post to and a row that fails at the "
     "first posting."),
]

M19_CHANGE = (
    "Moved the seeded expense-claim ladder off a role CodeX created and onto an "
    "approver group the tenant composes. Provisioning creates the group empty, so "
    "the first claim still parks rather than approving itself, and no "
    "finance-expense-claim-approver role is created in any tenant. A tenant that "
    "has published no ladder resolves to a shared platform row carrying no steps "
    "and is refused with APPROVAL_NOT_CONFIGURED, which the existing direct-post "
    "gate already recorded against whoever confirms past it."
)

M20_EDITS = [
    (7, (3, 2), "finance.refund.* plus membership of the approver group the stage names"),
    (7, (4, 2), "finance.writeoff.* plus membership of the approver group the stage names"),
    (16, "Current evidence",
     "Both types carry a workflow document type, a handler, a submit endpoint and a "
     "submit permission. The seeded ladder puts the threshold on both of its stages, "
     "so no stage applies below it, and each stage names an approver group created "
     "empty rather than a role provisioning invented. The direct-post gate asks the "
     "engine's own question - would any stage of the resolved template actually apply "
     "to this document - by walking the template with the same route resolution and "
     "condition evaluation the router uses. Posting is refused only while some stage "
     "would run, and separately when the resolved template has no stages at all, "
     "which is approval undecided rather than approval-free and is cleared only by a "
     "recorded confirmation."),
]

M20_CHANGE = (
    "Moved the seeded adjustment ladder off roles CodeX created and onto approver "
    "groups the tenant composes. Both stages of the concession, credit-note, "
    "refund and write-off ladder now name a group, created empty, so a ladder "
    "still arrives closed without a finance-adjustment-approver or "
    "finance-senior-adjustment-approver role appearing in a tenant that never "
    "asked for one. The threshold behaviour recorded at v1.2 is unchanged. A "
    "tenant with no ladder of its own resolves to a shared row with no steps and "
    "is refused rather than posted, unless somebody confirms and is recorded."
)


FR_ROOT = Path(__file__).resolve().parent.parent / "functional-requirements"

DOMAIN_FRDS = [
    {
        "folder": "18-payments-and-collections",
        "stem": "XVS_M18_Payments_and_Collections_Functional_Requirements_Document",
        "title_words": "XVS M18 Payments and Collections Functional Requirements Document",
        "module": 18, "source": "1.7", "target": "1.8",
        "edits": M18_EDITS, "change": M18_CHANGE, "changelog": 40,
    },
    {
        "folder": "19-finance-and-accounting",
        "stem": "XVS_M19_Finance_and_Accounting_Functional_Requirements_Document",
        "title_words": "XVS M19 Finance and Accounting Functional Requirements Document",
        "module": 19, "source": "1.7", "target": "1.8",
        "edits": M19_EDITS, "change": M19_CHANGE, "changelog": 43,
    },
    {
        "folder": "20-adjustments-and-concessions",
        "stem": "XVS_M20_Adjustments_and_Concessions_Functional_Requirements_Document",
        "title_words": "XVS M20 Adjustments and Concessions Functional Requirements Document",
        "module": 20, "source": "1.2", "target": "1.3",
        "edits": M20_EDITS, "change": M20_CHANGE, "changelog": 29,
    },
    {
        "folder": "22-procurement-and-requisitions",
        "stem": "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document",
        "title_words": "XVS M22 Procurement and Requisitions Functional Requirements Document",
        "module": 22, "source": "1.5", "target": "1.6",
        "edits": M22_EDITS, "change": M22_CHANGE, "changelog": 34,
    },
]

DOMAIN_BASELINE = "Working tree, 5 September 2026 (pending commit)"


def patch_domain_frds() -> None:
    for spec in DOMAIN_FRDS:
        folder = FR_ROOT / spec["folder"]
        source = folder / f"{spec['stem']}_v{spec['source']}.docx"
        output = folder / f"{spec['stem']}_v{spec['target']}.docx"
        patch_domain_frd(
            source, output,
            title=f"{spec['title_words']} v{spec['target']}",
            module_number=spec["module"],
            source_version=spec["source"], target_version=spec["target"],
            code_baseline=DOMAIN_BASELINE,
            edits=spec["edits"], change_summary=spec["change"],
            changelog_index=spec["changelog"],
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mrd-source", type=Path)
    parser.add_argument("--mrd-output", type=Path)
    parser.add_argument("--m07-source", type=Path)
    parser.add_argument("--m07-output", type=Path)
    parser.add_argument("--m04-source", type=Path)
    parser.add_argument("--m04-output", type=Path)
    parser.add_argument("--domain-frds", action="store_true")
    args = parser.parse_args()
    if args.mrd_source:
        patch_mrd(args.mrd_source, args.mrd_output)
    if args.m07_source:
        patch_m07(args.m07_source, args.m07_output)
    if args.m04_source:
        patch_m04(args.m04_source, args.m04_output)
    if args.domain_frds:
        patch_domain_frds()


if __name__ == "__main__":
    main()
