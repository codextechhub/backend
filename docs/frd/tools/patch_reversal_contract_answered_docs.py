#!/usr/bin/env python3
"""Record that every registered procurement type answers the reversal check.

Module 7 still described the engine's reversal contract as guarded in one place:
"Procurement guards purchase orders alone, so a vendor invoice or vendor payment
that has already posted can still have its approval reversed." Each procurement
type now answers for what its own approval released, through an overridable hook
on procurement's shared handler, so FR-019, the domain-handler dependency, the
Needs Attention lead and its reversal bullet are restated as current state. What
stays is the contract itself: the engine can withdraw its own record of a vote,
never what the vote released, which is why it asks before it writes.

Module 23 gives that refusal a capability entry of its own, as Module 18 has for
a payout that has reached the provider, taking Module 23 from 23 to 24 entries.

Both documents take one minor version. Source MRD references move to v2.77.

    python tools/patch_reversal_contract_answered_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.table import _Row

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
)

REVIEW_DATE = "12 September 2026"
SHORT_DATE = "12 Sep 2026"
MRD_VERSION = "2.77"
CODE_BASELINE = "Backend worktree at 19a50b3, 12 September 2026"
DEPLOYMENT = "Backend evidence only; nothing here is deployed."


# ── keep-format docx helpers ────────────────────────────────────────────────

def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError(f"Paragraph carries no run to inherit formatting from: {paragraph.text!r}")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    """Rewrite a one-paragraph cell, keeping the formatting it already has."""
    for paragraph in cell.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    set_run_text(cell.paragraphs[0], text)


def append_cell_text(cell, tail: str) -> None:
    set_cell_text(cell, cell.text.strip() + tail)


def rewrite_box(cell, lines: list[str]) -> None:
    """Rewrite a callout cell line by line.

    A callout is one bold coloured heading paragraph followed by regular bullet
    paragraphs, so line i reuses paragraph i, overflow clones the last bullet,
    and surplus paragraphs are removed. Cloning the heading for a new bullet
    would render every added line as another heading.
    """
    original = list(cell.paragraphs)
    if len(original) < 2:
        raise ValueError("A callout needs a heading and at least one bullet to reuse")
    for index, line in enumerate(lines):
        if index < len(original):
            set_run_text(original[index], line)
        else:
            clone = copy.deepcopy(original[-1]._p)
            cell.paragraphs[-1]._p.addnext(clone)
            set_run_text(cell.paragraphs[-1], line)
    for paragraph in original[len(lines):]:
        paragraph._p.getparent().remove(paragraph._p)


def replace_bullet(cell, prefix: str, text: str) -> None:
    """Rewrite exactly one callout bullet in place, identified by how it starts."""
    lines = [paragraph.text for paragraph in cell.paragraphs]
    matches = [index for index, line in enumerate(lines) if line.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one bullet starting {prefix!r}, found {len(matches)}")
    lines[matches[0]] = text
    rewrite_box(cell, lines)


def unique_cells(row) -> list:
    seen, cells = set(), []
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        cells.append(cell)
    return cells


def find_table(doc, header: str, *, contains: str | None = None):
    """The first table whose first cell starts with ``header``."""
    for table in doc.tables:
        if not table.rows[0].cells[0].text.strip().startswith(header):
            continue
        if contains is None or any(
            row.cells[0].text.strip().startswith(contains) for row in table.rows
        ):
            return table
    raise ValueError(f"Table not found: {header!r} containing {contains!r}")


def find_fr(doc, label: str):
    return find_table(doc, f"{label} |")


def find_row(table, prefix: str, *, column: int = 0):
    for row in table.rows:
        if row.cells[column].text.strip().startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix!r}")


def field(table, label: str):
    """The value cell of a requirement or control table row named ``label``."""
    return find_row(table, label).cells[1]


def clone_row(table, anchor, values: list[str], *, before: bool = False):
    new_tr = copy.deepcopy(anchor._tr)
    if before:
        anchor._tr.addprevious(new_tr)
    else:
        anchor._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)
    return row


def hold_together(table, *, lead_rows: int = 2) -> None:
    """Keep every row whole, and the first ``lead_rows`` with the row after them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
    for row in table.rows[:lead_rows]:
        for cell in unique_cells(row):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def paginate(doc) -> None:
    """Stop a requirement's heading and banner, or a table's lead, being stranded."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    for table in doc.tables:
        first = table.rows[0].cells[0].text.strip()
        if not ((first.startswith("FR-") and " | " in first) or first == "MRD capability"):
            continue
        hold_together(table)
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            Paragraph(previous, table._parent).paragraph_format.keep_with_next = True
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("All routes are mounted under /v1/procurement/"):
            paragraph.paragraph_format.keep_with_next = True


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def update_control(doc, *, version: str, module: int) -> None:
    control = find_table(doc, "Document", contains="Code baseline")
    set_cell_text(field(control, "Version"), version)
    set_cell_text(field(control, "Review date"), REVIEW_DATE)
    set_cell_text(field(control, "Code baseline"), CODE_BASELINE)
    set_cell_text(
        field(control, "Source MRD"),
        f"XVS Module Requirements Document v{MRD_VERSION} | Module {module}",
    )


def body_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Body paragraph not found: {prefix!r}")


def prepend_change_log(doc, version: str, summary: str) -> None:
    table = find_table(doc, "Version", contains="1.0")
    clone_row(table, table.rows[1], [version, SHORT_DATE, summary], before=True)


def finish(doc, output: Path, title: str, version: str) -> None:
    paginate(doc)
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


# ═════════════════════════════════════════════════════════════════════════════
# Module 7 - Workflow & Approval Engine
# ═════════════════════════════════════════════════════════════════════════════

M07_DIR = "07-workflow-and-approval-engine"
M07_STEM = "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document"
M07_SOURCE, M07_TARGET = "1.9", "1.10"

M07_FR019_EVIDENCE = (
    "reverse_action asks the document's handler through validate_reversal before "
    "writing anything, then voids the vote: the original gains reversed_at, "
    "reversed_by, and a reason, and a reversal row is appended. Where the reversed "
    "vote is what decided its stage, the stage reopens, every stage activated "
    "after it returns to PENDING with its live votes voided and its snapshot "
    "cleared, a decided instance returns to IN_PROGRESS, and the handler's "
    "on_action_reversed puts the document back inside the same transaction. A vote "
    "the stage did not need is voided and nothing else moves. What each module "
    "refuses is its own. Payouts refuse once any instruction is claimed for the "
    "provider, and dispatch re-checks the approval under the same row lock; "
    "finance refuses once the document has moved past draft or pending approval; "
    "procurement answers for all four of its types through an overridable hook on "
    "its shared handler, refusing a requisition once a purchase order has been "
    "raised from it, a purchase order once its vendor email is pending or sent or "
    "goods have been received against it, and a vendor invoice or vendor payment "
    "once its own status has left draft or pending approval, which is what posting "
    "does to it; user creation refuses once the account has left pending approval; "
    "role changes refuse once the request is decided. Leave returns a decided "
    "request to pending. A document type with no registered handler cannot be "
    "reversed. The view returns 404 for an action belonging to another tenant."
)
M07_FR019_ACCEPTANCE_TAIL = (
    " ProcurementApprovalReversalTests covers procurement's four answers and "
    "enumerates the registered procurement document types, so a fifth cannot be "
    "added there without one."
)
M07_FR019_LIMIT = (
    "The base handler allows every reversal, so a document is protected only where "
    "its module implements the check. Every registered finance, procurement, "
    "payout, user-creation and role-change type answers it; a leave request does "
    "not, and a document type registered later inherits the permissive default "
    "rather than a refusal. A rejected, withdrawn, or cancelled instance cannot be "
    "reversed."
)

M07_DOMAIN_HANDLERS = (
    "Register document handlers and consume the approval outcome. Each handler "
    "also says whether a decision may still be reversed and puts its document back "
    "after one: payouts refuse once an instruction is claimed for the provider, "
    "finance once the document has posted, and procurement once what its own "
    "approval released has happened, which is an order raised from a requisition, "
    "a purchase order that reached its vendor or was received against, and a "
    "posted vendor bill or vendor payment. Procurement's central ladder is "
    "published by this engine's publish service."
)

M07_ATTENTION_LEAD = (
    "These are current risks and gaps, not history. The controlled-spend hole v1.0 "
    "recorded here is closed, and so is the unsafe default that could reopen it. "
    "What is left is the release the submitter holds, the one document type whose "
    "requester may decide it, and the operational gaps below."
)

M07_REVERSAL_BULLET_PREFIX = "• Reversal is guarded only where the owning module implements the check"
M07_REVERSAL_BULLET = (
    "• Reversal is answered by the module that owns the document, and every "
    "registered finance, procurement, payment, user-creation and role-change type "
    "now answers it: procurement refuses a requisition an order was raised from, "
    "an order that reached its vendor or was received against, and a bill or a "
    "payment that has posted. The base handler still allows every reversal, so a "
    "leave request, and any document type registered later, is guarded by nothing "
    "until its own handler answers."
)

M07_TRACEABILITY_LEAD = (
    f"Module 7 carries 27 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Every registered procurement type answering the "
    "reversal check strengthens action reversal and the procurement workflow "
    "handlers without changing the count."
)

M07_CHANGE = (
    "Corrects what this document says about who answers a reversal. v1.9 recorded "
    "that procurement guarded purchase orders alone, so a vendor invoice or vendor "
    "payment that had been approved and then posted could still have its approval "
    "reversed. That is no longer true: procurement answers for all four of its "
    "types through an overridable hook on its shared handler, refusing a "
    "requisition once a purchase order has been raised from it, an order once it "
    "reached the vendor or goods were received against it, and a bill or a payment "
    "once its own status has left draft or pending approval. FR-019's evidence, "
    "acceptance and limit, the domain-handler dependency, the Needs Attention lead "
    "and its reversal bullet are restated as current state. The contract itself is "
    "unchanged and the document still says so: the engine can withdraw its own "
    "record of a vote, never what the vote released, which is why it asks the "
    "owning module before it writes. What remains is narrower and is stated as "
    "such: the base handler still permits, so a leave request, and any document "
    "type registered later, is guarded by nothing until its own handler answers. "
    "No engine code changed; the procurement work this records ran 553 procurement "
    f"tests OK at 5d45f39. Module 7 stays at 27 capability entries. {DEPLOYMENT}"
)


def patch_m07(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M07 Workflow and Approval Engine Functional Requirements "
        f"Document v{M07_TARGET}"
    )

    fr019 = find_fr(doc, "FR-019")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M07_SOURCE, M07_TARGET)
    update_control(doc, version=M07_TARGET, module=7)

    set_cell_text(field(fr019, "Current evidence"), M07_FR019_EVIDENCE)
    append_cell_text(field(fr019, "Acceptance"), M07_FR019_ACCEPTANCE_TAIL)
    set_cell_text(field(fr019, "Current limit"), M07_FR019_LIMIT)

    set_cell_text(find_row(dependencies, "Modules 17 to 24").cells[1], M07_DOMAIN_HANDLERS)

    set_run_text(
        body_paragraph(doc, "These are current risks and gaps, not history."),
        M07_ATTENTION_LEAD,
    )
    replace_bullet(gaps.rows[0].cells[0], M07_REVERSAL_BULLET_PREFIX, M07_REVERSAL_BULLET)

    set_run_text(body_paragraph(doc, "Module 7 carries"), M07_TRACEABILITY_LEAD)
    prepend_change_log(doc, M07_TARGET, M07_CHANGE)
    finish(doc, output, title, M07_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 23 - Purchase Orders, Delivery & AP
# ═════════════════════════════════════════════════════════════════════════════

M23_DIR = "23-purchase-orders-delivery-and-ap"
M23_STEM = "XVS_M23_Purchase_Orders_Delivery_and_AP_Functional_Requirements_Document"
M23_SOURCE, M23_TARGET = "1.10", "1.11"

M23_NEW_CAPABILITY = [
    "Approval reversal refused once an order, bill or payment has been acted on",
    "FR-001, FR-009",
]

M23_TRACEABILITY_LEAD = (
    f"Module 23 carries 24 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Refusing to reverse the approval behind a released "
    "order, or behind a posted bill or payment, is an entry of its own, as the "
    "same refusal is for a dispatched payout in Module 18."
)

M23_CHANGE = (
    "Gives the reversal refusal a capability entry of its own, taking Module 23 "
    "from 23 to 24. Module 18 carries the same entry for a payout that has reached "
    "the provider, and a tracker that answers the question for payouts while "
    "staying silent for supplier payments reads as though procurement still had "
    "the hole this module closed in v1.10. No behaviour, requirement, route or "
    "refusal code changes: the entry names what FR-001 and FR-009 already record. "
    f"The traceability lead and table are reconciled to MRD v{MRD_VERSION} at 24 "
    f"entries. {DEPLOYMENT}"
)


def patch_m23(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M23 Purchase Orders Delivery and AP Functional Requirements "
        f"Document v{M23_TARGET}"
    )

    traceability = find_table(doc, "MRD capability")

    replace_cover_version(doc.tables[0], M23_SOURCE, M23_TARGET)
    update_control(doc, version=M23_TARGET, module=23)

    anchor = find_row(traceability, "Concurrency-safe settlement of supplier bills")
    clone_row(traceability, anchor, M23_NEW_CAPABILITY)

    set_run_text(body_paragraph(doc, "Module 23 carries"), M23_TRACEABILITY_LEAD)
    prepend_change_log(doc, M23_TARGET, M23_CHANGE)
    finish(doc, output, title, M23_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    parser.add_argument("--only", nargs="*", choices=["m07", "m23"])
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for tag, folder, stem, source, target, patch in (
        ("m07", M07_DIR, M07_STEM, M07_SOURCE, M07_TARGET, patch_m07),
        ("m23", M23_DIR, M23_STEM, M23_SOURCE, M23_TARGET, patch_m23),
    ):
        if args.only and tag not in args.only:
            continue
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists():
            raise SystemExit(f"Refusing to overwrite {output.name}")
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
