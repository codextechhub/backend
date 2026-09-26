#!/usr/bin/env python3
"""Cut M01 FRD v1.27 and MRD v2.90: school creation becomes a watchable background job.

What changed in the backend, and therefore in the documents:

* POST /v1/i/create/ still validates the payload field by field (400) and now
  also checks the prebuilt role library in the request (503 through
  require_prebuilt_roles) before anything is queued. It then creates the school
  as a tracked Celery job (vs_schools.create_school) and answers 202 with the
  job's state: job_id, status, steps, done, current, school and message. A
  client-supplied job_id (a UUID) becomes the job's id, so the console can poll
  from the moment it posts; a reused id or a non-UUID is refused with 400.
* GET /v1/i/create/<uuid:job_id>/ reads that state for the owner alone, else 404.
* The job reports each step as it starts (school, roles, school_admin, branches,
  plan, books, onboarding, invitations) on a separate database connection
  through core.job_progress, so progress shows while the single creation
  transaction is still open. Creation stays all-or-nothing for the records a
  school needs to operate; books and onboarding remain best effort inside their
  own savepoints. The job re-validates the payload before it writes, because the
  door check and the write are moments apart, so a slug or an administrator's
  address taken in the gap fails the job rather than colliding with the first
  school. A failed job records the step it stopped on and a message fit to show,
  never a stack trace. The run appears in the platform task monitor as a
  school-kind job with no completion bell.
* The same change cut one creation from about 1,490 queries to about 900: the
  vs_rbac tenant-scope guard on bulk role writes is answered once per tenant
  rather than once per row (and a refused batch names every offending key), the
  vs_finance chart of accounts and fiscal year seed only the rows they are
  missing, and a vs_audit row no longer reads itself back to confirm it is new
  because the database primary key already does.

M04 and M05 were inspected and left unversioned: the scope-guard batching and
the audit-save change are performance refactors that keep the behaviour each
document already records. A new operational dependency is recorded: creation now
needs the Celery worker running the web service's code, or CELERY_EAGER set so
the task runs inline in the request.

    python tools/patch_school_creation_job_docs.py
"""
from __future__ import annotations

import copy

from docx import Document

import patch_mrd_v2_79_docs as mrd_tools
from patch_record_history_docs import (
    ROOT,
    append_rows,
    finish,
    frd_path,
    keep_format,
    log_change,
    replace_cell,
    set_control,
    set_cover_version,
)
from patch_staff_id_and_auth_events_docs import normalise_change_log, repair_ooxml

SHORT_DATE = "26 Sep 2026"

M01_CODE_BASELINE = (
    "Backend working tree on main at a98204de with the school-creation "
    "background job, 26 September 2026"
)
MRD_CODE_BASELINE = (
    "Backend main at a98204de with the school-creation background job in its "
    "working tree, 26 September 2026. The school and console integrations are "
    "implemented and pending release"
)


# ── local helpers ─────────────────────────────────────────────────────────────


def insert_row_after(table, index: int, values: list[str]) -> None:
    """Clone the row at ``index`` and fill the clone, keeping the row's format."""
    template = table.rows[index]
    template._tr.addnext(copy.deepcopy(template._tr))
    row = table.rows[index + 1]
    for cell, value in zip(row.cells, values):
        keep_format(cell, value)


def edit_box(cell, edits) -> None:
    """Edit a heading-and-bullets box line by line, whichever shape it has.

    Edits are ("sub", start, old, new) inside the one line starting with
    ``start``, ("after", start, new) to add a line after it, ("replace", start,
    new) and ("append", new). A one-run box keeps its lines as breaks in that
    run; a paragraph-per-line box reuses its paragraphs positionally.
    """
    lines = [line.strip() for p in cell.paragraphs for line in p.text.split("\n") if line.strip()]

    def at(start):
        hits = [i for i, line in enumerate(lines) if line.startswith(start)]
        if len(hits) != 1:
            raise ValueError(f"{start!r} starts {len(hits)} box lines")
        return hits[0]

    for edit in edits:
        kind = edit[0]
        if kind == "sub":
            _, start, old, new = edit
            i = at(start)
            if lines[i].count(old) != 1:
                raise ValueError(f"{old[:60]!r} occurs {lines[i].count(old)} times in the line")
            lines[i] = lines[i].replace(old, new)
        elif kind == "after":
            lines.insert(at(edit[1]) + 1, edit[2])
        elif kind == "replace":
            lines[at(edit[1])] = edit[2]
        elif kind == "append":
            lines.append(edit[1])
        else:
            raise ValueError(f"Unknown box edit {kind!r}")
    if len(cell.paragraphs) == 1 and len(cell.paragraphs[0].runs) == 1:
        cell.paragraphs[0].runs[0].text = "\n".join(lines)
    else:
        mrd_tools.set_lines(cell, lines)


def append_bullet(cell, bullet: str, *, size: float) -> None:
    """Add one ▸ bullet to a flowing capability cell, keeping its look."""
    replace_cell(cell, cell.text.rstrip() + "\n" + bullet, size=size)


# ── M01 School & Branch Management ────────────────────────────────────────────

M01_DIR = "01-school-and-branch-management"
M01_STEM = "XVS_M01_School_and_Branch_Management_Functional_Requirements_Document"
M01_SOURCE, M01_TARGET = "1.26", "1.27"

M01_DECISION_BULLET = (
    "• Creating a school is answered before it is built. The request still "
    "refuses a bad payload field by field and a short role library with 503, but "
    "the creation itself now runs as a tracked background job that reports each "
    "step, so the console shows it progressing and names the step a failure "
    "stopped on rather than holding one request open on a spinner. A "
    "client-supplied job id lets the console poll from the first moment, which "
    "is the only way to watch the steps where the job runs inline. The job "
    "re-validates before it writes, because the door check and the write are "
    "moments apart, so a slug or an administrator's address taken in the gap "
    "fails the job rather than colliding with the first school."
)

M01_FR001_EVIDENCE_TAIL = (
    " The create endpoint no longer holds the request open for the work. It "
    "validates the payload and calls require_prebuilt_roles in the request, so a "
    "bad field still answers 400 and a short library still answers 503 before "
    "anything is queued, then hands the creation to the vs_schools.create_school "
    "Celery task and answers 202 with the job's state. run_school_creation runs "
    "this same serializer on the worker inside one transaction and re-validates "
    "before it writes, because the request's check and the write are moments "
    "apart. A client-supplied job_id (a UUID) becomes the job id, so the console "
    "can poll GET /v1/i/create/<job_id>/ from the moment it posts."
)

M01_FR001_ACCEPTANCE = (
    "A 202 response starts a tracked background job and returns its id, status "
    "and steps. When the job reaches SUCCEEDED it has persisted one School, one "
    "protected Tenant, at least one Branch with exactly one main, the four "
    "whole-School roles and a Branch Admin role for each Branch, and every "
    "required administrator account, scoped role assignment, invitation record, "
    "and admin link. A missing Branch, an invalid payload, or a library that "
    "cannot supply one of the five roles is refused in the request, before "
    "anything is queued, with 400 or 503. A failure once the job is running, "
    "including an administrator that cannot be provisioned or a slug or address "
    "taken in the gap between the two checks, leaves none of those creation "
    "records behind and marks the job FAILED with the step it stopped on and a "
    "message fit to show. ANewSchoolGetsTheRolesCodeXShipsTests proves the full "
    "set arrives without anybody running a command afterwards; "
    "ASchoolIsNeverCreatedShortOfRolesTests proves a missing template refuses "
    "the School and names what is missing; and tests_creation_job proves the 202 "
    "handoff, the progress read by job id for the owner alone, the reused and "
    "non-UUID job id refusals, and the re-validation before the write."
)

M01_FR001_LIMIT_TAIL = (
    " Creation is now queued and answered before it completes, so a deployment "
    "must run the Celery worker on the web service's code, or keep CELERY_EAGER "
    "set so the task runs inline in the request; that operational dependency is "
    "tracked in the priority gaps below."
)

M01_OPS_ROW = [
    "Watch a school being created",
    "platform.schools.create",
    "The caller's own creation job, read by its id",
]

M01_API_CREATE_PURPOSE = (
    "Validate the payload and the prebuilt role library, then create the School, "
    "Tenant and nested records as a tracked background job; 202 with the job "
    "state. An optional client job_id (a UUID) lets the console poll at once."
)

M01_API_JOB_ROW = [
    "GET",
    "/v1/i/create/{job_id}/",
    "platform.schools.create",
    "The state of a creation the caller started: its steps, the one running, "
    "and on success the School. The owner only, else 404.",
]

M01_REFUSAL_ROWS = [
    ["Creation job id",
     "A client-supplied job id must be a UUID and unused. A reused id is refused "
     "so a second creation cannot ride an existing job's row.",
     "400 on a non-UUID or a reused job id"],
    ["Creation could not be queued",
     "When the worker cannot be reached and no job row is written, the request "
     "is refused rather than reported as started.",
     "503"],
    ["Slug or admin email taken after validation",
     "The job re-validates before it writes, because the door check and the "
     "write are moments apart. A slug or address taken in the gap fails the job "
     "instead of colliding with the first school.",
     "Job FAILED, with the field message; nothing saved"],
]

M01_DEP_ROW = [
    "core background jobs",
    "Outbound tracking and async run",
    "Run school creation as a Celery task tracked on core.BackgroundJob, and "
    "publish each step on a separate connection through core.job_progress so "
    "progress shows while the one creation transaction is still open. With "
    "CELERY_EAGER set the task runs inline in the request.",
]

M01_GAP_ROW = [
    "P1",
    "School creation depends on the worker running the same code",
    "Creation is queued as a Celery task and answered before it finishes, so a "
    "worker on stale code, or none at all with CELERY_EAGER off, would leave a "
    "school queued and never built. Confirm the worker runs the web service's "
    "code on every deploy, or keep CELERY_EAGER set so the task runs inline.",
    "FR-001",
]

M01_TRACE_ROW = [
    "School creation runs as a watchable background job",
    "FR-001",
    "Implemented; the payload and the role library are checked in the request, "
    "the creation runs as a tracked job that reports each step, and the console "
    "reads its progress by job id. The job re-validates before it writes, so a "
    "slug or address taken in the gap fails the job rather than the first "
    "school.",
]

M01_SUMMARY = (
    "Records school creation as a watchable background job. The create endpoint "
    "validates the payload and the prebuilt role library in the request, "
    "refusing with 400 or 503 before anything is queued, then runs the creation "
    "as a tracked Celery job and answers 202 with the job's state; a "
    "client-supplied job_id lets the console poll at once, a reused id or a "
    "non-UUID is refused, and the new GET /v1/i/create/{job_id}/ reads the state "
    "for the owner alone. Each step is reported as it starts (school, roles, "
    "school admin, branches, plan, books, onboarding, invitations) on a separate "
    "connection, so progress shows while the one creation transaction is still "
    "open; the creation stays all-or-nothing for the records a school needs to "
    "operate, with books and onboarding still best effort. The job re-validates "
    "before it writes, because the check and the write are moments apart, so a "
    "slug or address taken in the gap fails the job rather than colliding with "
    "the first school, and a failure names the step it stopped on with a message "
    "fit to show, never a stack trace. A new dependency is recorded on core "
    "background jobs and the Celery worker: creation is answered before it "
    "finishes, so the worker must run the web service's code, or CELERY_EAGER "
    "must keep the task inline. The supporting performance work (a tenant-scope "
    "guard answered once per tenant, finance seeds writing only missing rows, an "
    "audit row that no longer reads itself back) cut a creation from about 1,490 "
    "queries to about 900 without changing this module's contract. MRD baseline "
    "v2.90. Backend evidence only; the console integration is implemented and "
    "pending release, and nothing here claims deployment."
)


def patch_m01() -> None:
    doc = Document(str(frd_path(M01_DIR, M01_STEM, M01_SOURCE)))
    set_cover_version(doc, M01_SOURCE, M01_TARGET)
    set_control(doc, "Version", M01_TARGET)
    set_control(doc, "Review date", "26 September 2026")
    set_control(doc, "Code baseline", M01_CODE_BASELINE)
    set_control(doc, "Source MRD", "XVS Module Requirements Document v2.90 | Module 1")

    decision = doc.tables[6].rows[0].cells[0]
    edit_box(decision, [
        ("sub", "• Module 1 remains", "MRD v2.77, with twenty-three capability entries",
         "MRD v2.90, with twenty-four capability entries"),
        ("append", M01_DECISION_BULLET),
    ])

    fr001 = doc.tables[9]
    for row in fr001.rows:
        label = row.cells[0].text.strip()
        if label == "Current evidence":
            keep_format(row.cells[1], row.cells[1].text.rstrip() + M01_FR001_EVIDENCE_TAIL)
        elif label == "Acceptance":
            keep_format(row.cells[1], M01_FR001_ACCEPTANCE)
        elif label == "Current limit":
            keep_format(row.cells[1], row.cells[1].text.rstrip() + M01_FR001_LIMIT_TAIL)

    ops = doc.tables[8]
    if ops.rows[2].cells[0].text.strip() != "Create School":
        raise ValueError("Operations row 2 is not Create School")
    insert_row_after(ops, 2, M01_OPS_ROW)

    api = doc.tables[35]
    if api.rows[2].cells[1].text.strip() != "/v1/i/create/":
        raise ValueError("Schools API row 2 is not the create endpoint")
    keep_format(api.rows[2].cells[3], M01_API_CREATE_PURPOSE)
    insert_row_after(api, 2, M01_API_JOB_ROW)

    append_rows(doc.tables[38], M01_REFUSAL_ROWS)
    append_rows(doc.tables[39], [M01_DEP_ROW])
    append_rows(doc.tables[40], [M01_GAP_ROW])
    append_rows(doc.tables[42], [M01_TRACE_ROW])

    log_change(doc, M01_TARGET, M01_SUMMARY)

    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M01_DIR, M01_STEM, M01_TARGET),
           f"{M01_STEM.replace('_', ' ')} v{M01_TARGET}", M01_TARGET)


# ── MRD Module Requirements Document ──────────────────────────────────────────

MRD_SOURCE, MRD_TARGET = "2.89", "2.90"

MRD_CONTENTS_NOTE = "School creation becomes a watchable background job"

MRD_INTRO = (
    "This revision records that creating a school is answered before it is "
    "built: the request validates the payload and the prebuilt role library, "
    "then a tracked background job creates the school and reports each step, and "
    "the console reads its progress by job id. One capability is added, so the "
    "platform moves to 511 entries across 31 modules. The same change made a "
    "creation far cheaper, from about 1,490 queries to about 900, by batching "
    "the tenant-scope guard, seeding only the finance rows a school is missing "
    "and dropping an audit self-check; those are performance refactors that "
    "leave Module 4's and Module 5's documented behaviour unchanged, so neither "
    "FRD is versioned."
)

MRD_M01_BULLET = "▸  School creation runs as a tracked background job, watchable step by step"

MRD_M01_NEEDS_ATTENTION = (
    "• School creation runs on the worker. Validation and the role-library check "
    "answer in the request, then the creation is queued as a Celery task and the "
    "response returns before it finishes, so the worker must run the web "
    "service's code, or CELERY_EAGER must keep the task inline. A worker on "
    "stale code, or none at all with CELERY_EAGER off, would leave a school "
    "queued and never built."
)

MRD_M01_BLURB_TAIL = (
    " School creation runs as a tracked background job: the request validates "
    "the payload and the role library, then a worker builds the school in one "
    "transaction and reports each step, and the console reads its progress by "
    "job id and is told the step a failure stopped on."
)

MRD_DELTA_ROWS = [
    ["School creation", "Runs as a watchable background job",
     "POST /v1/i/create/ validates the payload and the prebuilt role library in "
     "the request, refusing with 400 or 503 before anything is queued, then "
     "creates the school on a worker and answers 202 with the job's steps; a "
     "client job_id lets the console poll from the first moment. GET "
     "/v1/i/create/{job_id}/ reads the state for the owner alone, else 404."],
    ["The steps", "Reported as each starts",
     "School, roles, school admin, branches, plan, books, onboarding and "
     "invitations, published on a separate connection so they show while the one "
     "creation transaction is still open. A failure names the step it stopped on "
     "and a message fit to show, never a stack trace. The run appears in the "
     "platform task monitor as a school-kind job with no completion bell."],
    ["The same slug taken in the gap", "Fails the job, saves nothing",
     "The request's check and the write are moments apart, so the job "
     "re-validates before it writes; a slug or an administrator's address taken "
     "meanwhile fails the job rather than colliding with the first school. The "
     "creation stays all-or-nothing for the records a school needs to operate, "
     "with books and onboarding still best effort."],
    ["Provisioning cost", "A creation makes far fewer queries",
     "One creation fell from about 1,490 queries to about 900: the vs_rbac "
     "tenant-scope guard checks a bulk role write once per tenant rather than "
     "once per row and names every offending key when it refuses, the vs_finance "
     "chart of accounts and fiscal year write only the rows a school is missing, "
     "and a vs_audit row no longer reads itself back to confirm it is new "
     "because the database key already does."],
    ["Module FRDs", "One revised",
     "M01 v1.27. M04 and M05 were checked and left: the scope-guard batching and "
     "the audit-save change are performance work that keeps their documented "
     "behaviour."],
]

MRD_CHANGE_SUMMARY = (
    "Records school creation as a watchable background job. POST /v1/i/create/ "
    "still validates the payload field by field and now also checks the prebuilt "
    "role library in the request, refusing with 400 or 503 before anything is "
    "queued; it then creates the school as a tracked Celery job and answers 202 "
    "with the job's state, and a client-supplied job_id (a UUID) lets the "
    "console poll from the moment it posts, a reused id or a non-UUID refused. "
    "The new GET /v1/i/create/{job_id}/ reads that state for the owner alone, "
    "else 404. The job reports each step as it starts (school, roles, school "
    "admin, branches, plan, books, onboarding, invitations) on a separate "
    "database connection, so progress shows while the single creation "
    "transaction is still open; creation stays all-or-nothing for the records "
    "that make a school usable, while books and onboarding remain best effort "
    "inside their own savepoints. Because the door check and the write are "
    "moments apart, the job re-validates before it writes: a slug or an "
    "administrator's address taken in the gap fails the job rather than "
    "colliding with the first school, and a failed job records the step it "
    "stopped on and a message fit to show, never a stack trace. The run appears "
    "in the platform task monitor as a school-kind job with no completion bell. "
    "The same change cut a creation from about 1,490 queries to about 900: the "
    "tenant-scope guard on bulk role writes is answered once per tenant rather "
    "than once per row and a refused batch names every offending key, the chart "
    "of accounts and fiscal year seed only the rows a school is missing, and an "
    "audit row no longer reads itself back to confirm it is new because the "
    "database key already does. One capability is added, so the platform moves "
    "to 511 entries across 31 modules; Module 1 stays Backend Partial and In "
    "use Complete. M01 advances to v1.27. Module 4 and Module 5 were checked and "
    "not versioned: the scope-guard batching and the audit-save change are "
    "performance refactors that keep their documented behaviour. A new "
    "operational dependency is recorded: creation now needs the Celery worker "
    "running the web service's code, or CELERY_EAGER set so the task runs "
    "inline. Backend evidence only: nothing here claims frontend delivery, "
    "deployment or production adoption, though the console integration is "
    "implemented and pending release."
)


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    doc = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta, log = tables[76], tables[78]
    assert delta.rows[0].cells[0].text.strip().endswith("capability delta")
    assert index.rows[0].cells[5].text.strip() == "Entries"

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for r in control.rows:
        label = r.cells[0].text.strip()
        if label == "Version":
            replace_cell(r.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(r.cells[1], "26 September 2026", size=9)
        elif label == "Source scope":
            replace_cell(r.cells[1], MRD_CODE_BASELINE, size=9)
        elif label == "Capability entries":
            replace_cell(r.cells[1], "511", size=9)

    for r in contents.rows:
        if r.cells[0].text.strip().startswith("5."):
            keep_format(r.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(r.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (text.startswith("5. ") and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    total = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total != 510:
        raise ValueError(f"Capability total is {total}, expected 510")
    if index.rows[1].cells[1].text.strip() != "School & Branch Management":
        raise ValueError("Index row 1 is not Module 1")
    if index.rows[1].cells[5].text.strip() != "23":
        raise ValueError("Module 1 no longer has 23 entries")
    keep_format(index.rows[1].cells[5], "24")

    grid = tables[mrd_tools.CAPABILITIES[1]]
    overflow = grid.rows[7].cells[2]
    if "Every way a school's reach shrinks" not in overflow.text:
        raise ValueError("Module 1 capability overflow cell moved")
    append_bullet(overflow, MRD_M01_BULLET, size=8.5)
    needs = grid.rows[8].cells[0]
    if not needs.text.startswith("NEEDS ATTENTION"):
        raise ValueError("Module 1 closing box is not NEEDS ATTENTION")
    edit_box(needs, [("append", MRD_M01_NEEDS_ATTENTION)])

    blurb = mrd_tools.module_blurbs(doc)[1]
    if "background job" in blurb.text:
        raise ValueError("Module 1's description already names the creation job")
    mrd_tools.retitle(blurb, blurb.text.rstrip() + MRD_M01_BLURB_TAIL)

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)

    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if len(intro) != 1:
        raise ValueError(f"{len(intro)} delta introductions found")
    mrd_tools.retitle(intro[0], MRD_INTRO)

    total_after = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total_after != 511:
        raise ValueError(f"Capability total after edit is {total_after}, expected 511")

    log.rows[1]._tr.addprevious(copy.deepcopy(log.rows[1]._tr))
    for cell, text in zip(log.rows[1].cells, (MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)):
        keep_format(cell, text)

    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m01()
    patch_mrd()


if __name__ == "__main__":
    main()
