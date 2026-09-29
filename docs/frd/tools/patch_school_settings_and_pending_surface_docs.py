#!/usr/bin/env python3
"""Cut M01 FRD v1.29 and M09 FRD v2.13: a School's own settings, and what it reaches before go-live.

What changed in the backend, and therefore in the documents:

* bf517f8b. A live School reads and writes its own sign-in security and its
  payroll scope through ``/v1/i/me/settings/security/`` and
  ``/v1/i/me/settings/payroll-scope/``, on ``school.settings.view`` and
  ``school.settings.update``, which nothing checked before. Security may only
  be tightened below the platform (or below the School, for a Branch). Once a
  School has been live, ``/v1/i/me/profile/`` refuses a change to its currency
  or term structure; the platform endpoint still changes both.
* 9030348e. Provisioning a role from the prebuilt library drops a key the
  School may not hold, and logs it, rather than refusing the copy and rolling
  the whole School back.
* a2634100 and 4767f637. A provisioned role also receives the library's
  default field switches, which the conversion of field-guarding keys wrote.
* 3aab9000. A School's books arrive holding one approval route per approvable
  document type, published with no steps, and no approver groups.
* bf518f2c. Recording a how-to guide event is open to a School that is not yet
  live.
* 6cf38070. A School that is not yet live works its own support desk: the
  list, a ticket, its thread and replies, attachments, triage, escalation,
  its audit and the dashboard counts. Assignment stays the platform desk's. M09's list of what a pending School reaches had fallen well behind the
  code, and is rewritten from each view's ``pending_tenant_surface``.

M09 also corrects three statements that the School's own profile route had
made false: School metadata is not platform-staff only, and a pending School
edits its own profile apart from its name, sign-in address and code.

The MRD is not revised here. M01 is traced to MRD v2.94, which adds the Module 1
entry FR-022 maps to, a school's own sign-in security and payroll scope, taking
Module 1 to twenty-five entries; M09 stays traced to v2.93, whose Module 9
entries v2.94 leaves as they were. M01's query-scoping entry and FR-012's limit
record that the financial statements are narrowed (4e3f3456), which leaves
vs_rbac's role list the one unnarrowed list (M04 v1.29). FR-022's limit names
the other settings a School sets at main (its display time zone, staff-profile
visibility and approval-email switch), and the timezone entry says the time zone
is the School's display.timezone value.

Each document starts from the newest version in its folder and is written at
the next free number; the script refuses to run if a newer version has
appeared, and ``finish`` refuses to overwrite one.

    python tools/patch_school_settings_and_pending_surface_docs.py
"""
from __future__ import annotations

from docx import Document

import patch_mrd_v2_79_docs as mrd_tools
from patch_backend_owned_permission_registry_docs import keep_rows_whole
from patch_document_type_labels_docs import require_newest
from patch_finance_audit_reconciliation_docs import keep_requirement_headers_with_body
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    edit_box,
    insert_row_after,
    row_labelled,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    fr_table,
    insert_after,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
)

import patch_record_history_docs

MRD_VERSION = "2.93"
#: M01 traces to the MRD version that lists FR-022; M09's entries did not move.
M01_MRD_VERSION = "2.94"
BACKEND_HEAD = "0f43b3e9"


def unique_cells(r):
    """A row's cells once each, however many grid columns a merged cell spans."""
    return list({id(c._tc): c for c in r.cells}.values())


def row_with_cell(doc, column: int, text: str):
    """The one row, in any table, whose ``column`` cell reads exactly ``text``."""
    hits = [r for t in doc.tables for r in t.rows
            if len(unique_cells(r)) > column
            and unique_cells(r)[column].text.strip() == text]
    if len(hits) != 1:
        raise ValueError(f"{text!r} in column {column} matches {len(hits)} rows")
    return hits[0]


def drop_break_before_heading(doc, heading: str) -> None:
    """Remove an empty page-break paragraph standing before a heading that breaks itself.

    The heading already starts a page. When the table above it ends at the foot of
    a page, the empty paragraph carrying the second break spills onto a page of
    its own, and the heading's break then leaves that page blank.
    """
    from docx.oxml.ns import qn

    target = paragraph_starting(doc, heading)
    if target._p.find(qn("w:pPr")) is None or target._p.pPr.find(qn("w:pageBreakBefore")) is None:
        raise ValueError(f"{heading!r} does not start its own page")
    previous = target._p.getprevious()
    if (previous is None or previous.tag != qn("w:p") or "".join(previous.itertext()).strip()
            or not previous.findall(".//" + qn("w:br"))):
        raise ValueError(f"No empty page-break paragraph stands before {heading!r}")
    previous.getparent().remove(previous)


def drop_hard_break_before(doc, heading: str) -> None:
    """Remove the empty page-break paragraph that stands before ``heading``.

    A break placed where an earlier layout needed one strands the foot of the
    table above on a page of its own once that table grows, and the heading
    that follows keeps with its own table without it.
    """
    from docx.oxml.ns import qn

    target = paragraph_starting(doc, heading)
    previous = target._p.getprevious()
    if (previous is None or previous.tag != qn("w:p") or "".join(previous.itertext()).strip()
            or not any(br.get(qn("w:type")) == "page" for br in previous.iter(qn("w:br")))):
        raise ValueError(f"No empty page-break paragraph stands before {heading!r}")
    previous.getparent().remove(previous)


def keep_m09_pages_whole(doc) -> None:
    """No table row breaks across a page, and no requirement ends a page on its first row.

    A requirement's status bar and Requirement row keep with the evidence below
    them; a boxed note, one cell wide, may still run over a page.
    """
    import re

    for table in doc.tables:
        if len(table.rows[0].cells) > 1:
            keep_rows_whole(table)
        if re.match(r"^FR-\d{3}\s*\|", table.rows[0].cells[0].text.strip()):
            keep_rows_whole(table)
            for table_row in table.rows[:2]:
                for cell in table_row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True


def box_starting(doc, heading: str):
    """The single-cell box whose text starts with ``heading``."""
    hits = [t.rows[0].cells[0] for t in doc.tables
            if len(t.rows) == 1 and t.rows[0].cells[0].text.strip().startswith(heading)]
    if len(hits) != 1:
        raise ValueError(f"{heading!r} starts {len(hits)} boxes")
    return hits[0]


# ── M01 School & Branch Management ───────────────────────────────────────────

M01_DIR = "01-school-and-branch-management"
M01_STEM = "XVS_M01_School_and_Branch_Management_Functional_Requirements_Document"
M01_SOURCE, M01_TARGET = "1.28", "1.29"
M01_DATE = "28 September 2026"

M01_BASELINE = (
    f"Backend main at {BACKEND_HEAD}, which carries a School's own security and payroll-scope "
    "settings and the go-live lock on its currency and term structure (bf517f8b), 28 September "
    "2026. The school app's settings console is committed and pending release"
)

M01_SCOPE = (
    "• School and paired Tenant creation, identifiers, metadata, status, and branding, and the "
    "sign-in security and payroll scope a School sets for itself."
)

M01_DECISION = (
    "• A School sets its own sign-in security and payroll scope. Two School-facing settings, on "
    "the School's own settings keys and open only once it is live, let it tighten sign-in "
    "security for the whole School or for one Branch, and choose one payroll or one per Branch. "
    "It can never loosen the platform's security, and no configuration key that decides what a "
    "School has bought is one a School can hold. Its currency and term structure are its own to "
    "set while it onboards and fixed once it has been live."
)

M01_FR001_EVIDENCE = (
    " The copy takes nothing the School may not hold. A key the library carries that a School "
    "tenant may not hold, because it has since been made platform-only, is dropped and logged "
    "rather than refused: the School never chose it and it would confer nothing inside a "
    "tenant, while the refusal rolled back the whole School, which is how finance.entity.create, "
    "left in the Finance Admin template after it was made platform-only, once stopped every "
    "School being created. vs_rbac withdraw_from_tenants, run by vs_rbac 0021, takes such keys "
    "back from the library and from every tenant surface, and audit_permission_scope --strict "
    "fails while any remain. Each role also receives the library's default field switches, "
    "which decide who reads and writes each restrictable field and were written into the "
    "library by the conversion of field-guarding keys (vs_rbac 0025). A default on an inactive "
    "field, or on a field that is not the School's to hold, is skipped rather than refused, and "
    "a write switch on a field that can no longer be written arrives as read only."
)

M01_FR001_ACCEPTANCE = (
    " LibraryLeftoverBreaksTenantProvisioningTests proves a School's role is provisioned over a "
    "library still carrying a platform-only key, without that key, and that the library keeps "
    "every key a tenant may still hold; PrebuiltRoleConversionTests proves a role provisioned "
    "from the library carries its default field switches."
)

M01_FR009_EVIDENCE = (
    " A School's currency and term structure are its own to set while it onboards and fixed "
    "once it has been live, by the test that freezes its sign-in address "
    "(School.has_ever_been_live). /v1/i/me/profile/ then refuses a change to either with 400 on "
    "that field, telling the School it is fixed and to contact XVS, and drops both from "
    "editable_fields; sending the stored value back is not a change. Every invoice, fee "
    "schedule and academic session the School has created is laid out in them, so a later "
    "change is a migration CodeX makes through the platform update endpoint, which still "
    "accepts both."
)

M01_FR009_ACCEPTANCE = (
    " LiveSchoolProfileEndpointTests proves a live School still reads its profile and changes "
    "its street address, that currency and term structure leave editable_fields and are refused "
    "field by field, that re-sending the stored value is accepted, that a School suspended after "
    "being live stays locked, and that CodeX can still change both through the platform "
    "endpoint."
)

M01_FR012_OLD = "and every surface but onboarding is closed to it by a gate that is not a permission key."
M01_FR012_NEW = (
    "and only the surfaces a School needs to set itself up are open to it, grouped in Module 9 "
    "FR-012, by a gate that is not a permission key."
)

M01_FR022 = (
    "FR-022  Let a School Set Its Own Security and Payroll Scope",
    "FR-022 | Implemented",
    [
        ("Requirement",
         "A live School can set, for itself, the two runtime settings that are genuinely its "
         "own: how strictly its people's sign-in is guarded, and whether it runs one payroll or "
         "one per Branch. It must never be able to loosen the platform's security, write another "
         "School's values, or reach the configuration keys that decide what a School has "
         "bought."),
        ("Current evidence",
         "GET and PATCH /v1/i/me/settings/security/ and /v1/i/me/settings/payroll-scope/ are the "
         "School's door, in vs_schools views/settings.py, gated on school.settings.view to read "
         "and school.settings.update to write. Those two keys were registered and checked by "
         "nothing before these endpoints; the view key is held by School Admin and Branch Admin, "
         "the update key by School Admin. Every config.* key stays platform-only, so no School "
         "holds one. Both endpoints act on the caller's own tenant and on no identifier in the "
         "request: they do not accept a platform operator's tenant assertion, and a caller whose "
         "tenant is not a School, which is a platform operator acting as itself, receives 404, "
         "so the door cannot write the platform baseline. Neither is on the pending-tenant "
         "surface, so a School still onboarding is refused with 403 TENANT_NOT_LIVE. The "
         "security body is exactly the console's /v1/config/security-settings/ body: six values, "
         "the failed sign-in threshold, the lockout minutes, the self-service and administrator "
         "password-reset link lifetimes, the invitation lifetime and the impersonation idle "
         "timeout, each effective, as stored and with its source. The write runs through "
         "vs_config save_security_settings, the function the console uses, so a School or Branch "
         "value may only be as strict as its parent or stricter, a refused field saves nothing "
         "from the same form, null removes the layer's own value, and each change is audited as "
         "a configuration change. ?branch= reads and writes one Branch's layer, which must belong "
         "to the School and be one the caller can see, or the answer is 404. The payroll body "
         "carries the scope, CENTRAL or PER_BRANCH, whether the School, the platform or the "
         "default decided it, and the two choices worded for the person choosing. It is written "
         "through vs_config set_value, so the finance write guard runs as it does for the "
         "console: PER_BRANCH is refused while any active member of staff has no Branch, with "
         "400 INVALID_CONFIGURATION_VALUE keyed on scope and a sentence naming up to ten of "
         "them. Payroll scope is a School-level setting only, so ?branch= is not read."),
        ("Acceptance",
         "SchoolSecuritySettingsAccessTests proves a teacher is refused both verbs, a Branch "
         "administrator may read but not write and cannot read a Branch they are not posted to, "
         "one School writing never moves another's values, another School's Branch or tenant "
         "answers 404, a School that is not live is refused, a platform caller acting as itself "
         "cannot write the platform baseline, and the console endpoint stays closed to a School "
         "administrator. SchoolSecuritySettingsBehaviourTests proves the body is the console's, "
         "a School can tighten and is audited, cannot loosen past the platform, a refused field "
         "saves nothing from the same form, null returns to the platform value, an empty form is "
         "refused, and a Branch can hold its own values. SchoolPayrollScopeTests proves the same "
         "access rules, the CENTRAL default with its options, a switch where everybody has a "
         "Branch, the guard's refusal as a 400 keyed on scope, an unknown scope refused, a "
         "platform value read as the platform's, and the console's write and this read "
         "agreeing."),
        ("Current limit",
         "These two are not the only settings a School sets for itself; others reach it "
         "through their own routes, among them its display time zone, display.timezone, at "
         "GET and PATCH /v1/i/me/settings/display/ under school.settings.view and "
         "school.settings.update; who reads how much of a staff profile, at GET and PUT "
         "/v1/i/me/settings/staff-profiles/, read under school.settings.view and written under "
         "school.field_access.update; and whether its approvals send emails, at GET and PATCH "
         "/v1/workflow/notification-settings/ under workflow.template.view and "
         "workflow.template.update (Module 7). Integrations, the plan and every configuration "
         "value that decides what a School has bought remain CodeX's. The School cannot change "
         "its payroll scope, or its security, before it goes live."),
    ],
)

M01_GATE_ROW = [
    "School's own security and payroll-scope settings",
    "school.settings.view / school.settings.update",
    "The asserted tenant's School only, once live; one Branch's security layer with ?branch=",
]

M01_PROFILE_PURPOSE_OLD = "The School's own profile, the fields it still has to fill, and those it may edit"
M01_PROFILE_PURPOSE_NEW = (
    "The School's own profile, the fields it still has to fill, and those it may edit; currency "
    "and term structure are fixed once the School has been live"
)

M01_ENDPOINT_ROWS = [
    ["GET, PATCH", "/v1/i/me/settings/security/", "school.settings.view / school.settings.update",
     "The School's own sign-in security, in the console's body. A School, or one Branch with "
     "?branch=, may only tighten its parent's values, and null clears its own. Live Schools "
     "only"],
    ["GET, PATCH", "/v1/i/me/settings/payroll-scope/",
     "school.settings.view / school.settings.update",
     "One payroll for the School or one per Branch. PER_BRANCH is refused while any active "
     "member of staff has no Branch. Live Schools only"],
]

M01_VALIDATION_ROWS = [
    ["Currency and term structure on the School's own profile",
     "The School's to change until it has ever been live, then refused; the stored value sent "
     "back is not a change. The platform update endpoint is not subject to the lock",
     "400 on the field"],
    ["School or Branch security setting",
     "Only as strict as its parent or stricter, within the console's bounds; null clears the "
     "layer's own value; an empty form is refused",
     "400 keyed on the field; nothing from the form is saved"],
    ["Payroll scope",
     "CENTRAL or PER_BRANCH; PER_BRANCH needs every active member of staff assigned to a Branch",
     "400 INVALID_CONFIGURATION_VALUE keyed on scope, naming up to ten unassigned people"],
]

M01_VS_CONFIG_TAIL = (
    " Carry the two settings a School sets for itself through the same functions the console "
    "uses: sign-in security through save_security_settings, and payroll scope through set_value, "
    "so finance's write guard runs (FR-022)."
)

M01_VS_RBAC_TAIL = (
    " A library key the School may not hold is dropped from the copy and logged rather than "
    "refusing the School, and the library's default field switches are copied with the "
    "permissions."
)

M01_VERIFICATION = (
    "• A School's own security and payroll-scope settings refused to a teacher, read-only to a "
    "Branch administrator, closed before go-live and to another School, unable to loosen the "
    "platform's security, and payroll scope refused while staff have no Branch; currency and "
    "term structure refused on the School's own profile once it has been live while the "
    "platform endpoint still changes them; and a role provisioned over a library carrying a "
    "platform-only key, without it."
)

M01_TRACE = (
    f"This table reconciles every Module 1 capability entry in MRD v{M01_MRD_VERSION} to the "
    "controlling functional requirement in this FRD. All twenty-five entries are represented. "
    "Two are restated rather than added: the roles CodeX ships now records that a School is "
    "refused rather than created short, that a Branch role follows its Branch, and that a "
    "library key a School may not hold is dropped rather than refusing it; and the role-grant "
    "entry covers every way a School's reach shrinks rather than a tier move alone. A refused "
    "logo saying why in the response message strengthens the branding and logo entry, and the "
    "go-live lock on currency and term structure the create, list, retrieve and update entry, "
    "without changing the count. FR-022, a School setting its own sign-in security and payroll "
    f"scope, maps to the entry MRD v{M01_MRD_VERSION} adds for it, which takes the count to "
    "twenty-five."
)

#: The Module 1 entry MRD v2.94 adds, worded as the MRD words it.
M01_SETTINGS_TRACE = [
    "A school's own sign-in security and payroll scope, set from its settings",
    "FR-022",
    "Implemented; a live School tightens its sign-in security, for the whole School or one "
    "Branch, and never loosens the platform's, and chooses one payroll or one per Branch, "
    "through /v1/i/me/settings/security/ and /v1/i/me/settings/payroll-scope/ under "
    "school.settings.view and school.settings.update; no config.* key is one a School can hold",
]

M01_TIMEZONE_TRACE_OLD = "Partial; no School timezone field, and lifecycle"
M01_TIMEZONE_TRACE_NEW = (
    "Partial; the time zone is not a School column but the School's display.timezone "
    "configuration value, Africa/Lagos by default, which a live School sets at "
    "/v1/i/me/settings/display/, and lifecycle"
)

M01_SCOPING_TRACE_OLD = (
    "applied by every dependent domain whose rows carry a Branch, with the financial "
    "statements outstanding."
)
M01_SCOPING_TRACE_NEW = (
    "applied by every dependent domain whose rows carry a Branch, the financial statements "
    "included, which read a branch-bound reader's own journals and the School-wide entries "
    "(Module 19 FR-006). vs_rbac's role list is the one list left unnarrowed (Module 4)."
)

M01_FR012_LIMIT_OLD = (
    "Two surfaces still return rows outside the caller's Branches, both outside this module: "
    "the financial statements, which aggregate ledger lines through several relation paths, "
    "and vs_rbac's own role and assignment lists."
)
M01_FR012_LIMIT_NEW = (
    "One list still returns rows outside the caller's Branches, outside this module: vs_rbac's "
    "own role list (Module 4 Needs Attention). The financial statements read a branch-bound "
    "reader's own journals and the School-wide entries (Module 19 FR-006), and vs_rbac's "
    "assignment list shows only the holders at the caller's Branches or School-wide."
)

M01_SUMMARY = (
    "Minor revision. A School sets its own sign-in security and payroll scope, and its currency "
    "and term structure lock at go-live. FR-022 is new for GET and PATCH "
    "/v1/i/me/settings/security/ and /v1/i/me/settings/payroll-scope/, on school.settings.view "
    "and school.settings.update, which were registered and checked by nothing before; no "
    "config.* key is one a School can hold. Both act on the caller's own School only, are closed "
    "before go-live, let a School or a Branch tighten security and never loosen it, and write "
    "through the same vs_config functions and audit as the console; payroll scope runs "
    "finance's guard, which refuses PER_BRANCH while staff have no Branch. FR-009 records that "
    "/v1/i/me/profile/ refuses currency and term structure once the School has been live and "
    "drops them from editable_fields, while the platform endpoint still changes both. FR-001 "
    "records two provisioning rules this document had not stated: a library key the School may "
    "not hold is dropped and logged rather than refusing the School, which is how a key made "
    "platform-only once stopped every School being created, and each role receives the "
    "library's default field switches. FR-012 no longer says every surface but onboarding is "
    "closed to a pending School. The scope, the module decision, the endpoint gates, the School "
    "endpoints, the validation rules, the vs_config and vs_rbac dependencies, the verification "
    f"list and the traceability follow, against MRD v{M01_MRD_VERSION}, whose new Module 1 "
    "entry, a school's own sign-in security and payroll scope, FR-022 maps to, taking the count "
    "to twenty-five. FR-012's limit and the query-scoping entry record that the financial "
    "statements are narrowed to a branch-bound reader's own journals (4e3f3456), leaving "
    "vs_rbac's role list the one unnarrowed list. FR-022's limit names the settings a School "
    "sets through other routes, its display time zone, staff-profile visibility and "
    "approval-email switch, and the timezone entry says the time zone is the School's "
    "display.timezone value rather than a missing field. Covered by the tests named in each "
    "requirement, which "
    "passed in schools.vs_schools (412 tests) and vs_config (117) at bf517f8b and in vs_rbac "
    "(559) at 9030348e; the suites were not re-run for this revision. The school app's settings "
    "console is committed and pending release; nothing here claims deployment."
)


def patch_m01() -> None:
    require_newest(str(ROOT / "functional-requirements" / M01_DIR / f"{M01_STEM}_v*.docx"),
                   M01_SOURCE)
    patch_record_history_docs.REVIEW_DATE = M01_DATE
    doc = Document(str(frd_path(M01_DIR, M01_STEM, M01_SOURCE)))
    mrd_tools.replace_cover_version(doc.tables[0], M01_SOURCE, M01_TARGET)
    set_control(doc, "Version", M01_TARGET)
    set_control(doc, "Review date", M01_DATE)
    set_control(doc, "Code baseline", M01_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{M01_MRD_VERSION} | Module 1")

    mrd_tools.retitle(paragraph_starting(doc, "• School and paired Tenant creation"), M01_SCOPE)
    edit_box(box_starting(doc, "CURRENT MODULE DECISION"), [
        ("sub", "• Module 1 remains Backend Partial", "MRD v2.90", f"MRD v{M01_MRD_VERSION}"),
        ("sub", "• Module 1 remains Backend Partial", "twenty-four capability entries",
         "twenty-five capability entries"),
        ("append", M01_DECISION),
    ])

    anchor = row_labelled(doc, "School's own Branches")
    insert_row_after(anchor, M01_GATE_ROW)

    fr001 = fr_table(doc, "FR-001")
    append_to(fr001, "Current evidence", M01_FR001_EVIDENCE)
    append_to(fr001, "Acceptance", M01_FR001_ACCEPTANCE)

    fr009 = fr_table(doc, "FR-009")
    append_to(fr009, "Current evidence", M01_FR009_EVIDENCE)
    append_to(fr009, "Acceptance", M01_FR009_ACCEPTANCE)

    edit_cell(row(fr_table(doc, "FR-012"), "Current evidence").cells[-1],
              M01_FR012_OLD, M01_FR012_NEW)
    edit_cell(row(fr_table(doc, "FR-012"), "Current limit").cells[-1],
              M01_FR012_LIMIT_OLD, M01_FR012_LIMIT_NEW)

    add_fr(doc, *M01_FR022)

    profile = row_with_cell(doc, 1, "/v1/i/me/profile/")
    edit_cell(unique_cells(profile)[-1], M01_PROFILE_PURPOSE_OLD, M01_PROFILE_PURPOSE_NEW)
    logo = row_with_cell(doc, 1, "/v1/i/me/profile/logo/")
    anchor = logo
    for values in M01_ENDPOINT_ROWS:
        anchor = insert_row_after(anchor, values)

    anchor = row_labelled(doc, "Partial update scope")
    for values in M01_VALIDATION_ROWS:
        anchor = insert_row_after(anchor, values)

    for label, tail in (("vs_config", M01_VS_CONFIG_TAIL), ("vs_rbac", M01_VS_RBAC_TAIL)):
        cell = unique_cells(row_labelled(doc, label))[-1]
        keep_format(cell, cell.text.rstrip() + tail)

    insert_after(doc, "• The roles CodeX ships arriving with a new School", M01_VERIFICATION)
    mrd_tools.retitle(paragraph_starting(doc, "This table reconciles every Module 1"), M01_TRACE)
    zone = row_labelled(doc, "School status, contact, location, and timezone data")
    edit_cell(unique_cells(zone)[-1], M01_TIMEZONE_TRACE_OLD, M01_TIMEZONE_TRACE_NEW)
    scoping = row_labelled(doc, "Tenant-aware School and Branch query scoping")
    edit_cell(unique_cells(scoping)[-1], M01_SCOPING_TRACE_OLD, M01_SCOPING_TRACE_NEW)
    insert_row_after(row_labelled(doc, "School creation runs as a watchable background job"),
                     M01_SETTINGS_TRACE)
    drop_break_before_heading(doc, "11. Change Log")

    log_change(doc, M01_TARGET, M01_SUMMARY)
    # A change-log row never breaks across a page.
    keep_rows_whole(patch_record_history_docs.change_log_table(doc))
    assert_absent_outside_log(doc, "MRD v2.92", "MRD v2.90", "MRD v2.93", "every surface but onboarding",
                              "twenty-four capability entries", "is not yet traced",
                              "with the financial statements outstanding",
                              "role and assignment lists", "These are the only runtime settings",
                              "no School timezone field")
    keep_requirement_headers_with_body(doc)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M01_DIR, M01_STEM, M01_TARGET),
           f"{M01_STEM.replace('_', ' ')} v{M01_TARGET}", M01_TARGET)


# ── M09 School Onboarding ────────────────────────────────────────────────────

M09_DIR = "09-school-onboarding"
M09_STEM = "XVS_M09_School_Onboarding_Functional_Requirements_Document"
M09_SOURCE, M09_TARGET = "2.12", "2.13"
M09_DATE, M09_LOG_DATE = "28 September 2026", "28 Sep 2026"

M09_BASELINE = (
    f"Backend main at {BACKEND_HEAD}, 28 September 2026. The school app's how-to guides are "
    "committed and pending release"
)

M09_PURPOSE = (
    "Version 2.13 records what a new school's books now hold, one empty approval route per "
    "approvable document type and no approver groups, and brings the list of what a school may "
    "reach before go-live back to the code, grouped by area. No behaviour in this module "
    "changed."
)

M09_METADATA_OWNER = (
    "Module 1. Branch creation is a platform-staff responsibility under platform keys, by "
    "decision, not a missing school-facing endpoint. A school edits its own profile through "
    "Module 1's /v1/i/me/profile/ on school.profile.update, open before go-live, apart from its "
    "name, sign-in address and code, which stay CodeX's; its currency and term structure are "
    "fixed once it has been live."
)

M09_BOOKS_TAIL = (
    " The approval routes that arrive with the books are empty and name nobody, and the "
    "workflow screens where a school builds its own steps open at go-live."
)

M09_FR010_OLD = (
    "Escalation uses the existing POST /v1/support/tickets/, which is keyless by design and is "
    "the only ticket action open to a school that is not yet live."
)
M09_FR010_NEW = (
    "Escalation uses the existing POST /v1/support/tickets/, which is keyless by design. A "
    "school that is not yet live works its own support desk, as it will once live: its ticket "
    "list, a ticket, the thread and replies, attachments and their download, updating, "
    "following, moving and escalating a ticket, its audit and the dashboard counts, each "
    "scoped to the tickets it may see exactly as after go-live. Assigning a ticket, and the "
    "list of people it may be assigned to, stay the platform desk's and are closed before "
    "go-live, and a ticket cannot be deleted by anyone."
)
M09_FR010_EVIDENCE = (
    " Self-help is open before go-live as well. Recording a how-to guide event is on the "
    "pending surface, because an event carries no tenant or user, while the platform-wide guide "
    "summary is not. The school app opens its how-to guides, articles, help panel and "
    "walkthroughs before go-live, and a ticket raised from a guide carries its guide_id; that "
    "integration is committed in the school app and pending release."
)
M09_FR010_ACCEPTANCE = (
    " PendingSchoolSupportDeskTests proves a pending school lists only its own tickets, opens "
    "one, reads support's reply and answers it, attaches and downloads a file, lets its desk "
    "staff triage, reads its dashboard counts, cannot see another school's ticket, and cannot "
    "assign; PendingSchoolSurfaceTests pins the open and closed actions. A pending school "
    "records a guide event and the guide summary stays closed "
    "(test_a_school_being_set_up_can_record_guide_events)."
)

M09_FR012_OLD = (
    "The surface is the ten onboarding routes, the self-scoped sign-in and profile endpoints, "
    "the personal notification inbox, and creating a support ticket. Notification settings, the "
    "event catalogue, delivery history and all of vs_config stay closed."
)
M09_FR012_NEW = (
    "Each view's own pending_tenant_surface declaration is the definitive list: absence means "
    "closed, and a declaration may name the actions it opens, so one view can be open to read "
    "and closed to write. Grouped by area, the surface is: onboarding, the ten routes in section "
    "7.1; the person's own session, which is the current user, their security summary and "
    "password resets, the password policy, changing a password, refreshing a token and signing "
    "out; the personal notification inbox; help, which is the school's own support desk, every "
    "ticket action but assignment and the dashboard counts included, and recording a how-to "
    "guide event; the school's own "
    "records, which are stored media, the School profile and logo, and the list of its own "
    "Branches; roles, which are the tenant's role list, creating, reading and editing a role, "
    "the permission and access catalogues, and each role's field switches; the capabilities "
    "read that draws the school's navigation; academic structure and the calendar, every view in "
    "both; staff, which is the directory and search, adding a person, reading and changing their "
    "record, their qualifications and documents, postings, the roster, bulk role assignment, a "
    "person's roles and resending an invitation; and bulk import, which is the templates, "
    "batches, validation and its issues, starting and cancelling a batch, and the job list. "
    "Adding a person while the school is pending asks for School Admin or Branch Admin and "
    "refuses anything else, because onboarding is how a school's first administrators arrive; "
    "once it is live, a new person starts as Teacher and the adder chooses no role (Module 12). "
    "Everything else stays closed until go-live, among it assigning a support ticket, the guide "
    "analytics summary, notification settings, the event catalogue and delivery "
    "history, every other configuration route, one-person field exceptions and permission "
    "overrides, the organogram, teaching duties, leave and the other staff lifecycle actions, "
    "student records, the approval workflow screens, and the school's own security and payroll "
    "settings."
)
M09_FR012_ACCEPTANCE = (
    " A pending school shapes its roles' field switches and is refused one-person field "
    "exceptions and permission overrides (test_a_school_that_has_not_gone_live_can_shape_field_"
    "access, PendingSchoolExceptionTests), records a how-to guide event "
    "(test_a_school_being_set_up_can_record_guide_events), and is refused its own security and "
    "payroll settings (SchoolSecuritySettingsAccessTests and SchoolPayrollScopeTests, "
    "test_a_school_that_is_not_live_is_refused)."
)

M09_FR016_EVIDENCE = (
    " The books arrive holding one approval route for each approvable document type across "
    "finance, procurement and payments, each published with no steps, and no approver groups at "
    "all: who approves a school's money is its own answer, built from its organogram through the "
    "workflow screens, and no onboarding step asks for it. An empty route is not the same as "
    "none. For finance it is the gate itself, and for procurement and payments it stands in "
    "front of the shared platform route; either way a refund, a requisition or a payout batch "
    "submitted before the school has set any steps is refused as unconfigured, and goes through "
    "only when somebody confirms it without approval, which is recorded against them."
)
M09_FR016_ACCEPTANCE = (
    " A new school's books hold one empty route per approvable document type and no approver "
    "groups, and a later provisioning run leaves a school's own steps untouched (the "
    "tests_stageless_provisioning suites in vs_finance and vs_payments, and "
    "tests_approval_confirmation in vs_procurement)."
)

M09_METADATA_STEP = (
    "Name, slug, code, ownership type, term structure and currency are set on the school's own "
    "record. The school sets the ones it may through its own profile while it onboards; currency "
    "and term structure are fixed once it has been live."
)

M09_DECISION_4_TAIL = (
    " Superseded in part for school metadata: a school edits its own profile before go-live "
    "through Module 1's /v1/i/me/profile/, apart from its name, sign-in address and code. Branch "
    "creation is unchanged."
)

M09_DEP_M19_TAIL = (
    " Each module registers its approval route with that provisioning, so the books arrive "
    "holding one empty route per approvable document type and no approver groups."
)
M09_DEP_M31_TAIL = (
    " The school's own support desk is open before go-live, apart from assigning a ticket, and "
    "so is recording a how-to guide event; the guide summary is not."
)


M09_TRACE = (
    f"Module 9 carries 21 capability entries in MRD v{MRD_VERSION}. Each maps to the "
    "requirements below. The empty approval routes and the fuller pending surface fall inside "
    "existing entries, so no entry or count changes."
)

M09_SUMMARY = (
    "Minor revision. Records what a new school's books now hold and brings the list of what a "
    "school may reach before go-live back to the code. Books arrive holding one approval route "
    "per approvable document type across finance, procurement and payments, each published with "
    "no steps, and no approver groups: who approves is the school's own answer, and the workflow "
    "screens where it is given open at go-live, so until then such a document is refused as "
    "unconfigured and goes through only when somebody confirms it without approval (FR-016, "
    "section 1.2, the Module 19 dependency). FR-012 listed the pending surface as the onboarding "
    "routes, sign-in and profile, the inbox and filing a ticket, and said all of vs_config stays "
    "closed; the code opens far more, and the capabilities read has been open since 9 September. "
    "The surface is now grouped by area, with each view's own pending_tenant_surface declaration "
    "named as the definitive list, and it gains recording a how-to guide event, each role's "
    "field switches, and the rule that adding a person while pending asks for School Admin or "
    "Branch Admin while a live school starts everyone as Teacher. FR-010 said filing a ticket "
    "was the only ticket action open before go-live; since 6cf38070 a pending school works its "
    "whole support desk, the list, thread, replies, attachments, triage, escalation and "
    "dashboard counts, while assignment stays the platform desk's. It also records that the "
    "school app's how-to guides open before go-live, committed and pending release. Section 1.2, the "
    "SCHOOL_METADATA step and decision 4 are corrected: a school edits its own profile before "
    "go-live, apart from its name, sign-in address and code, and its currency and term structure "
    "are fixed once it has been live. Module 9 stays Backend Complete and In use Partial with "
    f"twenty-one capability entries against MRD v{MRD_VERSION}. No code in this module changed. "
    "Covered by the tests named in each requirement; the suites were not re-run for this "
    "revision. Nothing here claims deployment."
)


def patch_m09() -> None:
    require_newest(str(ROOT / "functional-requirements" / M09_DIR / f"{M09_STEM}_v*.docx"),
                   M09_SOURCE)
    patch_record_history_docs.REVIEW_DATE = M09_LOG_DATE
    doc = Document(str(frd_path(M09_DIR, M09_STEM, M09_SOURCE)))
    mrd_tools.replace_cover_version(doc.tables[0], M09_SOURCE, M09_TARGET)
    set_control(doc, "Version", M09_TARGET)
    set_control(doc, "Review date", M09_DATE)
    set_control(doc, "Code baseline", M09_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_VERSION} | Module 9")
    set_control(doc, "Supersedes", f"v{M09_SOURCE} and all earlier versions, retained unchanged")

    mrd_tools.retitle(paragraph_starting(doc, "Version 2.12 records"), M09_PURPOSE)

    keep_format(unique_cells(row_labelled(doc, "Branch creation and school metadata edits"))[-1],
                M09_METADATA_OWNER)
    books = unique_cells(row_labelled(doc, "Books, the chart of accounts and approval ladders"))[-1]
    keep_format(books, books.text.rstrip() + M09_BOOKS_TAIL)

    fr010 = fr_table(doc, "FR-010")
    edit_cell(row(fr010, "Current evidence").cells[-1], M09_FR010_OLD, M09_FR010_NEW)
    append_to(fr010, "Current evidence", M09_FR010_EVIDENCE)
    edit_cell(row(fr010, "Acceptance").cells[-1],
              "A pending tenant can create a ticket and cannot reach the rest of the desk.",
              "A pending tenant works its own support desk and cannot assign a ticket.")
    append_to(fr010, "Acceptance", M09_FR010_ACCEPTANCE)

    fr012 = fr_table(doc, "FR-012")
    edit_cell(row(fr012, "Current evidence").cells[-1], M09_FR012_OLD, M09_FR012_NEW)
    append_to(fr012, "Acceptance", M09_FR012_ACCEPTANCE)

    fr016 = fr_table(doc, "FR-016")
    append_to(fr016, "Current evidence", M09_FR016_EVIDENCE)
    append_to(fr016, "Acceptance", M09_FR016_ACCEPTANCE)

    keep_format(unique_cells(row_labelled(doc, "SCHOOL_METADATA"))[-1], M09_METADATA_STEP)

    decision = unique_cells(row_labelled(doc, "4. Who may run onboarding?"))[1]
    keep_format(decision, decision.text.rstrip() + M09_DECISION_4_TAIL)

    for label, tail in (("Module 19, Finance and Accounting", M09_DEP_M19_TAIL),
                        ("Module 31, Support Tickets", M09_DEP_M31_TAIL)):
        cell = unique_cells(row_labelled(doc, label))[-1]
        keep_format(cell, cell.text.rstrip() + tail)

    mrd_tools.retitle(paragraph_starting(doc, "Module 9 carries 21 capability entries"), M09_TRACE)

    log_change(doc, M09_TARGET, M09_SUMMARY)
    assert_absent_outside_log(
        doc, "all of vs_config stay closed", "is the only ticket action open",
        "MRD v2.77", "Version 2.12 records", "cannot read an answer",
    )
    drop_hard_break_before(doc, "FR-004  Recompute Go-Live Readiness")
    drop_hard_break_before(doc, "FR-015  Retain Go-Live History for a Year")
    keep_m09_pages_whole(doc)
    keep_requirement_headers_with_body(doc)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M09_DIR, M09_STEM, M09_TARGET),
           f"{M09_STEM.replace('_', ' ')} v{M09_TARGET}", M09_TARGET)


def main() -> None:
    patch_m01()
    patch_m09()


if __name__ == "__main__":
    main()
