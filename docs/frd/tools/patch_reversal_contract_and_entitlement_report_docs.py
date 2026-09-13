#!/usr/bin/env python3
"""Record two contracts the platform now enforces rather than assumes.

A document type says whether its approval can be undone, or it does not
register: Module 7 gains the registry refusal, the base handler's refusal by
default and the declaration a type makes when its approval genuinely releases
nothing, and Module 12 gains the rule leave never had, that an absence already
under way refuses the reversal of its approval and cancelling the request is the
remedy.

Every write that narrows a school's reach settles its roles and says what it
could not save: Module 6 records both entitlement routes reporting, the
settlement that clearing a layer now performs, and the honest limit that the
entitlement POST carries no depth field, and Module 4 withdraws its statement
that a settlement from that endpoint reaches nobody reading the response.

Each document takes one minor version. Source MRD references move to v2.79.

    python tools/patch_reversal_contract_and_entitlement_report_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.table import _Row

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
)

REVIEW_DATE = "12 September 2026"
SHORT_DATE = "12 Sep 2026"
MRD_VERSION = "2.79"
CODE_BASELINE = "Backend worktree at 5678e86, 12 September 2026"
DEPLOYMENT = "Backend evidence only; nothing here is deployed."


# ── keep-format docx helpers ────────────────────────────────────────────────

def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError(f"Paragraph carries no run to inherit formatting from: {paragraph.text!r}")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    """Rewrite a one-paragraph cell, keeping the formatting it already has."""
    for paragraph in cell.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    set_run_text(cell.paragraphs[0], text)


def append_cell_text(cell, tail: str) -> None:
    set_cell_text(cell, cell.text.strip() + tail)


def rewrite_box(cell, lines: list[str]) -> None:
    """Rewrite a callout cell line by line.

    A callout is one bold coloured heading paragraph followed by regular bullet
    paragraphs, so line i reuses paragraph i, overflow clones the last bullet,
    and surplus paragraphs are removed. Cloning the heading for a new bullet
    would render every added line as another heading.
    """
    original = list(cell.paragraphs)
    if len(original) < 2:
        raise ValueError("A callout needs a heading and at least one bullet to reuse")
    for index, line in enumerate(lines):
        if index < len(original):
            set_run_text(original[index], line)
        else:
            clone = copy.deepcopy(original[-1]._p)
            cell.paragraphs[-1]._p.addnext(clone)
            set_run_text(cell.paragraphs[-1], line)
    for paragraph in original[len(lines):]:
        paragraph._p.getparent().remove(paragraph._p)


def replace_bullet(cell, prefix: str, text: str) -> None:
    """Rewrite exactly one callout bullet in place, identified by how it starts."""
    lines = [paragraph.text for paragraph in cell.paragraphs]
    matches = [index for index, line in enumerate(lines) if line.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one bullet starting {prefix!r}, found {len(matches)}")
    lines[matches[0]] = text
    rewrite_box(cell, lines)


def replace_line(cell, prefix: str, text: str) -> None:
    """Rewrite one line of a callout written as a single paragraph of line breaks.

    Some callouts hold their heading and every bullet in one paragraph separated
    by breaks rather than in a paragraph each, so the line is found by how it
    starts and the paragraph is written back whole. A break survives the round
    trip because a newline in a run becomes one again.
    """
    paragraph = cell.paragraphs[0]
    lines = paragraph.text.split("\n")
    matches = [index for index, line in enumerate(lines) if line.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one line starting {prefix!r}, found {len(matches)}")
    lines[matches[0]] = text
    set_run_text(paragraph, "\n".join(lines))


def unique_cells(row) -> list:
    seen, cells = set(), []
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        cells.append(cell)
    return cells


def find_table(doc, header: str, *, contains: str | None = None):
    """The first table whose first cell starts with ``header``."""
    for table in doc.tables:
        if not table.rows[0].cells[0].text.strip().startswith(header):
            continue
        if contains is None or any(
            row.cells[0].text.strip().startswith(contains) for row in table.rows
        ):
            return table
    raise ValueError(f"Table not found: {header!r} containing {contains!r}")


def find_fr(doc, label: str):
    return find_table(doc, label)


def find_row(table, prefix: str, *, column: int = 0):
    for row in table.rows:
        if row.cells[column].text.strip().startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix!r}")


def field(table, label: str):
    """The value cell of a requirement or control table row named ``label``."""
    return find_row(table, label).cells[1]


def clone_row(table, anchor, values: list[str], *, before: bool = False):
    new_tr = copy.deepcopy(anchor._tr)
    if before:
        anchor._tr.addprevious(new_tr)
    else:
        anchor._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)
    return row


def set_row_text(row, values: list[str]) -> None:
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)


def hold_together(table, *, lead_rows: int = 2) -> None:
    """Keep every row whole, and the first ``lead_rows`` with the row after them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
    for row in table.rows[:lead_rows]:
        for cell in unique_cells(row):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def paginate(doc) -> None:
    """Stop a requirement's heading and banner, or a table's lead, being stranded."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    for table in doc.tables:
        first = table.rows[0].cells[0].text.strip()
        if not (first.startswith("FR-") or first == "MRD capability"):
            continue
        hold_together(table)
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            Paragraph(previous, table._parent).paragraph_format.keep_with_next = True


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def body_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Body paragraph not found: {prefix!r}")


def body_paragraph_containing(doc, fragment: str):
    """The one paragraph carrying ``fragment``, for text an icon or a space precedes."""
    matches = [p for p in doc.paragraphs if fragment in p.text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph containing {fragment!r}, found {len(matches)}")
    return matches[0]


def find_change_log(doc):
    """The change log, whose header row is exactly a version, a date and a summary.

    Matched on the whole header rather than on its first cell, because a document
    may hold other tables whose first cell opens with the word Version.
    """
    for table in doc.tables:
        cells = [cell.text.strip() for cell in unique_cells(table.rows[0])]
        if len(cells) == 3 and cells[0] == "Version" and cells[1].startswith("Date"):
            return table
    raise ValueError("Change log not found")


def prepend_change_log(doc, version: str, summary: str) -> None:
    table = find_change_log(doc)
    clone_row(table, table.rows[1], [version, SHORT_DATE, summary], before=True)


def finish(doc, output: Path, title: str, version: str) -> None:
    paginate(doc)
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


# ═════════════════════════════════════════════════════════════════════════════
# Module 7 - Workflow & Approval Engine
# ═════════════════════════════════════════════════════════════════════════════

M07_DIR = "07-workflow-and-approval-engine"
M07_STEM = "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document"
M07_SOURCE, M07_TARGET = "1.11", "1.12"

M07_FR019_REQUIREMENT = (
    "An administrator can reverse a decision without erasing it, and only where "
    "the module that owns the document agrees the decision has not already been "
    "acted on. A document type that has not said what its approval releases does "
    "not reach a running system at all."
)

M07_FR019_EVIDENCE = (
    "reverse_action asks the document's handler through validate_reversal before "
    "writing anything, then voids the vote: the original gains reversed_at, "
    "reversed_by, and a reason, and a reversal row is appended. Where the reversed "
    "vote is what decided its stage, the stage reopens, every stage activated "
    "after it returns to PENDING with its live votes voided and its snapshot "
    "cleared, a decided instance returns to IN_PROGRESS, and the handler's "
    "on_action_reversed puts the document back inside the same transaction. A vote "
    "the stage did not need is voided and nothing else moves. "
    "The question is asked of every type, because a type that does not answer it "
    "cannot register. register_handler refuses a handler that has neither answered "
    "reversal_block_reason, overridden validate_reversal, nor set "
    "approval_releases_nothing, raising REVERSAL_CONTRACT_NOT_DECLARED while the "
    "apps load, so silence stops the process that would serve a school rather than "
    "surfacing on the day an administrator undoes an approval whose effect is "
    "still standing. BaseWorkflowHandler.reversal_block_reason refuses by default "
    "behind it, for a handler that reaches the engine without passing through the "
    "registry, and the flag counts only when it is True, because saying something "
    "is released without saying what blocks a reversal is the same silence in a "
    "different spelling. A type whose approval genuinely releases nothing says so "
    "in as many words. "
    "All fourteen registered types declare an answer. Procurement's four each "
    "answer on their own handler, with no shared permissive hook above them, so a "
    "fifth added there inherits refusal rather than permission: a requisition "
    "refuses while an order raised from it still stands or has taken delivery, an "
    "order once its vendor email is pending or sent or goods have been received "
    "against it, and a vendor invoice or vendor payment once its own status has "
    "left draft or pending approval. Finance's six answer through one shared check "
    "that holds the document's row while it reads that status. Payouts refuse once "
    "any instruction is claimed for the provider, and dispatch re-checks the "
    "approval under the same row lock; user creation refuses once the account has "
    "left pending approval; role changes refuse once the request is decided. Leave "
    "answers for the absence itself, since nobody is stored as On Leave and the "
    "directory derives it from an approved request covering today: leave that has "
    "not started is reversible, and leave that is running or already taken is "
    "refused with cancelling the request named as the remedy. A document type with "
    "no registered handler cannot be reversed. The view returns 404 for an action "
    "belonging to another tenant."
)

M07_FR019_ACCEPTANCE = (
    "The original decision remains readable, tallies count only unreversed "
    "actions, and an approver whose vote was reversed can vote again. A refusal "
    "leaves the approval, its stages, and its votes as they were, and a handler "
    "that fails while putting its document back rolls the whole reversal back. A "
    "vote from an attempt its stage has since run past cannot be reversed. The "
    "ACTION_REVERSED audit entry names the reopened stage and every stage rolled "
    "back. ReverseActionTests, ReverseActionUnknownHandlerTests, and "
    "ReversalUnwindTests cover the engine. UndeclaredTypeTests proves a handler "
    "that declares nothing cannot register, that the default answer refuses rather "
    "than permits, and that approval_releases_nothing set False is silence rather "
    "than a declaration; DeclaredTypeTests covers the three ways a type declares, "
    "including a module base answering once for every type it carries; and "
    "RegisteredTypeTests.test_every_registered_document_type_answers_the_reversal_question "
    "enumerates the registry rather than sampling it, so a fifteenth type cannot "
    "arrive silently. PayoutBatchApprovalTests, JournalApprovalWorkflowTests, "
    "WorkflowApprovalTests, ProcurementApprovalReversalTests, "
    "UserCreationReversalTests and the leave ReversalTests and DecisionTests cover "
    "the refusals and the documents put back."
)

M07_FR019_LIMIT = (
    "A declaration is taken on the declaring module's word. "
    "approval_releases_nothing states that withdrawing the engine's record of the "
    "vote is the whole of the change, and nothing here checks that claim against "
    "what the module's own callbacks do. A rejected, withdrawn, or cancelled "
    "instance cannot be reversed, and a document type with no registered handler "
    "cannot be reversed at all."
)

M07_REVERSAL_BULLET_PREFIX = "• Reversal is answered by the module that owns the document"
M07_REVERSAL_BULLET = (
    "• Reversal is answered by the module that owns the document, and a document "
    "type that has not said what its approval releases does not register: the "
    "refusal fires while the apps load, and the base handler refuses by default "
    "behind it. All fourteen registered types answer, procurement's four each on "
    "their own handler, and leave refuses once an absence has begun. What is taken "
    "on trust is a type's own word that its approval releases nothing, which "
    "nothing here checks against what its callbacks do."
)

M07_DOMAIN_HANDLERS = (
    "Register document handlers and consume the approval outcome. Each handler "
    "also says whether a decision may still be reversed and puts its document back "
    "after one, and a handler that says neither is refused registration: payouts "
    "refuse once an instruction is claimed for the provider, finance once the "
    "document has posted, and procurement once what its own approval released has "
    "happened, which is an order still standing against a requisition or goods "
    "received against one, a purchase order that reached its vendor or was "
    "received against, and a posted vendor bill or vendor payment. Each of "
    "procurement's four types answers on its own handler, so a fifth added there "
    "inherits refusal. Procurement's central ladder is published by this engine's "
    "publish service."
)

M07_STAFF_APP_TAIL = (
    " The staff app also answers the reversal contract for leave: approval "
    "releases the absence itself rather than a row, so leave that has not started "
    "may be undone, and leave that is running or already taken is refused."
)

M07_TRACEABILITY_LEAD = (
    f"Module 7 carries 27 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Requiring a document type to declare what its "
    "approval releases before it may register strengthens action reversal and the "
    "domain handler entries without changing the count."
)

M07_CHANGE = (
    "Records that a document type says whether its approval can be undone, or it "
    "does not register. v1.11 recorded the base handler as allowing every "
    "reversal, so a leave request, and any type registered later, was guarded by "
    "nothing: silence read exactly like a type that genuinely releases nothing. "
    "register_handler now refuses a handler that has neither answered "
    "reversal_block_reason, overridden validate_reversal nor set "
    "approval_releases_nothing, raising REVERSAL_CONTRACT_NOT_DECLARED while the "
    "apps load, and the base handler refuses by default behind it for anything "
    "that reaches the engine without passing through the registry. All fourteen "
    "registered types declare an answer, procurement's four each on their own "
    "handler with no shared permissive hook above them, so a fifth inherits "
    "refusal. Leave gains the rule it never had: approval releases the absence "
    "itself, so leave that has not started is reversible while leave that is "
    "running or already taken is refused, with cancelling the request as the "
    "remedy. FR-019's requirement, evidence, acceptance and limit, the reversal "
    "bullet in Needs Attention, and the domain-handler and staff-app dependencies "
    "are restated as current state. What remains is narrower and is stated as "
    "such: a type's declaration that its approval releases nothing is taken on "
    "that module's word. Module 7 stays at 27 capability entries. Verified by 406 "
    "workflow tests and 232 staff tests, with finance 759, procurement 567, "
    f"payments 194, RBAC 547 and user 381 beside them. {DEPLOYMENT}"
)


def update_control_banner(doc, *, version: str, module: int, mrd_suffix: str = "") -> None:
    control = find_table(doc, "Document", contains="Code baseline")
    set_cell_text(field(control, "Version"), version)
    set_cell_text(field(control, "Review date"), REVIEW_DATE)
    set_cell_text(field(control, "Code baseline"), CODE_BASELINE)
    set_cell_text(
        field(control, "Source MRD"),
        f"XVS Module Requirements Document v{MRD_VERSION} | Module {module}{mrd_suffix}",
    )


def patch_m07(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M07 Workflow and Approval Engine Functional Requirements "
        f"Document v{M07_TARGET}"
    )

    fr019 = find_fr(doc, "FR-019")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M07_SOURCE, M07_TARGET)
    update_control_banner(doc, version=M07_TARGET, module=7)

    set_cell_text(field(fr019, "Requirement"), M07_FR019_REQUIREMENT)
    set_cell_text(field(fr019, "Current evidence"), M07_FR019_EVIDENCE)
    set_cell_text(field(fr019, "Acceptance"), M07_FR019_ACCEPTANCE)
    set_cell_text(field(fr019, "Current limit"), M07_FR019_LIMIT)

    set_cell_text(find_row(dependencies, "Modules 17 to 24").cells[1], M07_DOMAIN_HANDLERS)
    append_cell_text(find_row(dependencies, "Domain apps").cells[1], M07_STAFF_APP_TAIL)
    replace_bullet(gaps.rows[0].cells[0], M07_REVERSAL_BULLET_PREFIX, M07_REVERSAL_BULLET)

    set_run_text(body_paragraph(doc, "Module 7 carries"), M07_TRACEABILITY_LEAD)
    prepend_change_log(doc, M07_TARGET, M07_CHANGE)
    finish(doc, output, title, M07_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 12 - Staff Management
# ═════════════════════════════════════════════════════════════════════════════

M12_DIR = "12-staff-management"
M12_STEM = "XVS_M12_Staff_Management_Functional_Requirements_Document"
M12_SOURCE, M12_TARGET = "2.4", "2.5"

M12_VERIFIED_AGAINST = (
    f"{CODE_BASELINE}, and docs/designs/Staff_Management.html, which is nine "
    "screens, seven drawers and seventy-nine collections. What version 2.5 "
    "changes was checked against the code at that baseline. The rest keeps the "
    "evidence it was written with, so a statement in sections 2 to 4 that M14 or "
    "this module is unbuilt describes the platform the specification was first "
    "written against."
)

M12_FR013_RULES = (
    "(1) end_date on or after start_date, refused as 422 INVALID_DATE_RANGE. (2) "
    "Days defaults to the inclusive calendar span and is editable, because working "
    "days are not calendar days and nothing records a school's teaching week; M14 "
    "has the same question open for its own teaching-day count. (3) A request is "
    "created at PENDING and submitted to the workflow engine in the same "
    "transaction, under document type schools.leave_request, with the person's "
    "branch as the scope so the engine's branch to tenant to platform cascade "
    "resolves the right template. (4) The status is written only by the workflow "
    "handler's callbacks and by the cancel service. A reversed decision returns an "
    "APPROVED or REJECTED request to PENDING and clears decided_at, because the "
    "status follows the instance; a CANCELLED request stays cancelled. This module "
    "writes no audit event of its own for a reversal, and the engine's reversal "
    "entry is its record. Completed is derived from the end date at read time and "
    "is never stored. (5) Overlapping leave for one person WARNS and does not "
    "refuse, in data.warnings, against their PENDING and APPROVED requests: a sick "
    "day recorded inside booked annual leave is a correction, not a mistake. (6) "
    "Correcting a request is allowed while it is PENDING and refused once it is "
    "decided, because a request whose dates change after approval is a different "
    "request. (7) Cancelling sets status CANCELLED and never deletes, because "
    "leave taken is part of the employment history. (8) The report is days taken "
    "per type for a session, counted from APPROVED requests, and there is NO "
    "balance: an entitlement exists nowhere. (9) An approved request is what puts "
    "somebody on leave, and it writes nothing to their employment record. An "
    "ACTIVE person whose approved request covers today reads as ON_LEAVE, with the "
    "latest end date among such requests as on_leave_until, and reads ACTIVE again "
    "the day after it ends; somebody suspended, resigned or terminated reads as "
    "what they are. Reversing the approval removes the reading at once, because "
    "the request is PENDING again. (10) The leave document declares two fields a "
    "school's own Dynamic Role may test, the leave type and the days requested, "
    "and the staff record supplies two facts about the requester, their job title "
    "and their contract type, read from their record at this school and absent for "
    "anybody without one. The ladder this module provisions tests none of them and "
    "sends every request to the Leave Approvers group; a school that wants sick "
    "leave over five days to go to its head teacher builds that rule itself. A job "
    "title is compared exactly as the school typed it, so a rule on Head Teacher "
    "does not match Head teacher. (11) An absence that has begun REFUSES the "
    "reversal of its approval, which is this module's answer to the engine's "
    "reversal contract. Approval releases the absence itself rather than a row: "
    "nobody is stored as On Leave, the directory derives it from an approved "
    "request covering today, and the cover a school arranges is arranged from what "
    "it reads. Leave that has not started has released nothing and may have its "
    "approval undone. Leave running today is refused, because withdrawing the "
    "approval would make somebody read as present without bringing them back, and "
    "leave already taken is refused for the same fact after the event, since the "
    "days taken are counted from approved requests. The refusal begins on the "
    "first day and never lifts, and it names the remedy: cancel the request, which "
    "is what ends the absence and records the cancellation as what it is. A "
    "request that is PENDING, REJECTED or CANCELLED released nothing whatever its "
    "dates, and neither did one whose row has since gone."
)

M12_ACCEPTANCE_PREFIX = "Reversing an approval returns the request to PENDING"
M12_ACCEPTANCE = (
    "Reversing an approval returns the request to PENDING and clears decided_at "
    "(DecisionTests.test_reversing_the_decision_puts_the_request_back_to_pending), "
    "and is refused once the absence has begun: leave that has not started may be "
    "undone, leave running today and leave already taken are both refused with "
    "cancelling the request named as the remedy, and a pending, rejected or "
    "cancelled request blocks nothing whatever its dates (ReversalTests)."
)

M12_TRACEABILITY_LEAD = (
    f"MRD v{MRD_VERSION} records Module 12 as Staff Management, Phase V1, Backend "
    "Partial, In use Partial, code schools.vs_staff mounted at /v1/i/me/staff/, "
    "with ten capability entries. Each maps below to the requirements that deliver "
    "it and its state at the code baseline. Three requirements trace to no entry, "
    "bulk import (FR-016), bulk role grant (FR-019) and staff search (FR-021), "
    "because the tracker lists no search or bulk capability for this module."
)

M12_LEAVE_CAPABILITY = (
    "Partial. Leave is applied for and decided on the workflow engine, and an "
    "approved request covering today makes the row read On Leave with the date the "
    "person is back. Approval releases that absence rather than a row, so its "
    "reversal is refused from the first day of the leave onwards and cancelling "
    "the request is the remedy. Whether somebody is free at a given hour is not "
    "evidenced, because no contract records hours."
)

M12_MINOR_NOTE = (
    "ℹ  Why version 2.5 is minor. The functional baseline is intact and so "
    "are its contracts in every other place: the same six models, the same two "
    "status vocabularies, the same posting-and-reach separation, the same routes, "
    "keys, refusal codes and audit actions. What changes is one rule, whether an "
    "approved absence may have its approval undone, and the answer is that it may "
    "until the first day of the leave and not afterwards."
)

M12_CHANGE = (
    "Records the rule that decides whether an approved absence may have its "
    "approval undone. The workflow engine asks the module owning a document "
    "whether its approval has already taken effect, and leave had no answer, so an "
    "absence already under way could have its approval withdrawn: the person "
    "stopped reading as On Leave without coming back, and the cover the school had "
    "arranged was answering a question nobody was asking any more. Approval "
    "releases the absence itself rather than a row, because nobody is stored as On "
    "Leave and the directory derives it from an approved request covering today. "
    "Leave that has not started is reversible; leave running today or already "
    "taken is refused, and the refusal names the remedy, which is to cancel the "
    "request, where the absence ends and the cancellation is recorded as what it "
    "is. A request that is pending, rejected or cancelled released nothing "
    "whatever its dates, and neither did one whose row has since gone. FR-013 "
    "gains the rule as its eleventh, section 12's leave acceptance and section "
    "15's traceability record it, and the control page moves to 5678e86. No model, "
    "route, permission key or audit action changed. Covered by ReversalTests; "
    f"schools.vs_staff ran 232 tests OK. {DEPLOYMENT}"
)


def patch_m12(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M12 Staff Management Functional Requirements Document v{M12_TARGET}"

    control = doc.tables[0]
    set_cell_text(field(control, "Version"), M12_TARGET)
    set_cell_text(field(control, "Date"), REVIEW_DATE)
    set_cell_text(field(control, "Supersedes"), f"v{M12_SOURCE}")
    set_cell_text(
        field(control, "Source MRD"),
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 12",
    )
    set_cell_text(field(control, "Verified against"), M12_VERIFIED_AGAINST)

    fr013 = find_fr(doc, "FR-013")
    set_cell_text(field(fr013, "Business rules"), M12_FR013_RULES)

    set_run_text(body_paragraph(doc, M12_ACCEPTANCE_PREFIX), M12_ACCEPTANCE)
    set_run_text(body_paragraph(doc, "MRD v2.77 records Module 12"), M12_TRACEABILITY_LEAD)

    traceability = find_table(doc, "MRD capability")
    set_cell_text(
        find_row(traceability, "Leave and availability indicators").cells[2],
        M12_LEAVE_CAPABILITY,
    )

    # The change log here runs oldest first, so a new version is appended.
    change_log = find_change_log(doc)
    clone_row(change_log, change_log.rows[-1], [M12_TARGET, SHORT_DATE, M12_CHANGE])
    set_run_text(
        body_paragraph_containing(doc, "Why version 2.4 is minor"), M12_MINOR_NOTE,
    )

    finish(doc, output, title, M12_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 6 - Configuration & Capability Management
# ═════════════════════════════════════════════════════════════════════════════

M06_DIR = "06-configuration-and-capability"
M06_STEM = (
    "XVS_M06_Configuration_and_Capability_Management_Functional_Requirements_Document"
)
M06_SOURCE, M06_TARGET = "1.2", "1.3"

M06_FR010_EVIDENCE_TAIL = (
    " Clearing a layer settles the same way. A tenant's own row carries its own "
    "depth and wins over the platform row while it exists, so deleting it can "
    "return the tenant to a shallower reach at once, and clear_entitlement settles "
    "the role grants afterwards like every other write that narrows one. The "
    "report of any role the settlement could not save is an argument of the "
    "settlement itself, with no default, at all four sites that settle, so a write "
    "cannot narrow a tenant's reach without saying where those names are to go."
)

M06_FR010_LIMIT = (
    "A null depth means no depth limit at all rather than the shallowest one. "
    "Grants written before depth existed therefore reach every band, which is "
    "deliberate and is discussed under Needs Attention. Both entitlement routes "
    "return the settlement's report as data.roles_needing_attention and name the "
    "count in their message, but the POST cannot shallow a module at all, because "
    "SetEntitlementSerializer carries no depth field: its report is empty by "
    "construction rather than by evidence, and on this pair of routes the report "
    "is reachable only through the DELETE."
)

M06_RBAC_DEPENDENCY = (
    "Enforces the nineteen PLATFORM-scoped keys, and consumes this module's "
    "evaluation in the plan gate, reading platform.entitlements.enforce through "
    "get_config on every gated request. The permission-to-band map lives there, "
    "not here. It also receives the one call that runs the other way: every write "
    "here that can narrow a tenant's depth asks vs_rbac to take back the role "
    "grants beyond it, and vs_rbac owns that service, its audit trail and the "
    "report of any role it could not settle. That report is a parameter of the "
    "settlement rather than something read off its return value, because these "
    "services answer with the row they wrote, so a write here cannot settle a "
    "tenant's roles without naming where the report goes, and both entitlement "
    "routes carry it back to the operator who made the change."
)

M06_VERIFICATION = (
    "The requirements above name the tests behind each behaviour; this revision "
    "was traced from the code at the baseline named in Document Control and did "
    "not re-run them. This module's own suite ran 117 tests OK at that baseline, "
    "and the consumers of the settlement these writes perform ran green beside it: "
    "vs_rbac 547 and schools.vs_schools 358, the latter a full run holding the "
    "tests that exercise every way a tenant's reach can shrink. "
    f"{DEPLOYMENT}"
)

M06_GAP_SIX = [
    "6",
    "An entitlement written from the console cannot shallow a module",
    "Both entitlement routes settle the tenant's role grants and return what the "
    "settlement could not save, as data.roles_needing_attention with the count "
    "named in the message, and clearing a layer settles like every other write "
    "that narrows a reach. The POST cannot narrow a depth at all, because "
    "SetEntitlementSerializer carries no depth field: it writes state, source and "
    "window and nothing else. Its report is therefore empty by construction rather "
    "than by evidence, and the only way to shallow a tenant from this pair of "
    "routes is to delete the layer that was carrying the deeper grant.",
    "Put a depth field on the entitlement write, so an operator can shallow a "
    "module deliberately from the console and the report that route already "
    "carries has something to carry.",
]

M06_TRACEABILITY_LEAD = (
    f"Module 6 of XVS Module Requirements Document v{MRD_VERSION} lists twenty-six "
    "capabilities. Each maps to the requirements above; the role grants that "
    "follow a narrowing of a tenant's depth are a consequence this module's writes "
    "trigger and vs_rbac performs, so they strengthen FR-010 and FR-012 without "
    "adding a capability or changing the count."
)

M06_CHANGE = (
    "Records that every write here which narrows a tenant's reach settles that "
    "tenant's roles and says what it could not save. v1.2 recorded the entitlement "
    "endpoint as settling with nowhere to report a role it left standing, and "
    "clearing an entitlement settled nothing at all, although a tenant's own row "
    "wins over the platform row while it exists, so deleting it returns the tenant "
    "to the platform depth at once and leaves its roles holding keys the tenant no "
    "longer reaches. clear_entitlement now settles like every other write that "
    "narrows a reach, and the report is an argument of the settlement itself, with "
    "no default, at all four sites: set_entitlement, set_depth_grant, "
    "clear_depth_grant and clear_entitlement. Both entitlement routes return "
    "data.roles_needing_attention and name the count in their message through the "
    "same sentence the plan endpoints use, so the configuration screen and the "
    "plan screen cannot come to describe the same event differently. No revocation "
    "rule changed: closing a module rather than shallowing it still takes no key, "
    "and a tenant holding no package grant still reads as unprovisioned and loses "
    "nothing. FR-010's evidence and limit, the vs_rbac dependency, the "
    "verification evidence and MRD traceability are restated, and the sixth "
    "current gap is rewritten to the risk that remains: the POST carries no depth "
    "field, so it cannot shallow a module and its report is empty by construction. "
    "Module 6 remains Backend Complete and In use Complete with twenty-six "
    f"capability entries against MRD v{MRD_VERSION}. Verified by 117 configuration "
    f"tests, with vs_rbac 547 and schools.vs_schools 358 beside them. {DEPLOYMENT}"
)


def patch_m06(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M06 Configuration and Capability Management Functional Requirements "
        f"Document v{M06_TARGET}"
    )

    replace_cover_version(doc.tables[0], M06_SOURCE, M06_TARGET)
    update_control_banner(
        doc, version=M06_TARGET, module=6, mrd_suffix=", twenty-six capability entries",
    )

    fr010 = find_fr(doc, "FR-010")
    append_cell_text(field(fr010, "Current evidence"), M06_FR010_EVIDENCE_TAIL)
    set_cell_text(field(fr010, "Current limit"), M06_FR010_LIMIT)

    dependencies = find_table(doc, "Dependency")
    set_cell_text(find_row(dependencies, "vs_rbac").cells[1], M06_RBAC_DEPENDENCY)

    set_run_text(body_paragraph(doc, "The requirements above name the tests"), M06_VERIFICATION)

    gaps = find_table(doc, "Pri.", contains="6")
    set_row_text(find_row(gaps, "6"), M06_GAP_SIX)

    set_run_text(body_paragraph(doc, "Module 6 of XVS Module Requirements"), M06_TRACEABILITY_LEAD)
    prepend_change_log(doc, M06_TARGET, M06_CHANGE)
    finish(doc, output, title, M06_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 4 - Roles & Permissions (RBAC)
# ═════════════════════════════════════════════════════════════════════════════

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.19", "1.20"

M04_SOURCE_SCOPE = (
    "Backend changes making the report of a role the settlement could not save an "
    "argument of the settlement itself, so no write that narrows a tenant's reach "
    "can settle its roles without saying where that report goes, and settling a "
    "tenant's roles when its own entitlement layer is cleared (12 September 2026)"
)

M04_CODE_INSPECTED = (
    f"{CODE_BASELINE}: apps/vs_rbac (plan_grants.py, plan_gate.py, services.py), "
    "vs_config's capability, depth and entitlement services with the views above "
    "them, and the vs_schools plan endpoints and plan command they meet"
)

M04_FR025_EVIDENCE_TAIL = (
    " The report reaches the operator through a list the caller owns and the "
    "settlement fills, taken as an argument of the settlement rather than read off "
    "its return value, because the writes that settle a tenant's reach answer with "
    "the row they wrote and their callers depend on that. It carries no default, "
    "so a site cannot settle a tenant's roles without saying where the names of "
    "the ones it left standing are to go, and every site passes one: the plan "
    "endpoints, the apply_plans command, and vs_config's set_entitlement, "
    "set_depth_grant, clear_depth_grant and clear_entitlement. One shared "
    "sentence, unsettled_roles_note, is appended to the message by every surface "
    "that settles, so the plan screen and the configuration screen describe the "
    "same event in the same words."
)

M04_FR025_LIMIT = (
    "Only rbac_permission is gated: a view gated solely by rbac_group_permission "
    "or by HasAnyModuleAccess passes untouched. Enforcement is one platform switch "
    "rather than one per school, so a staged rollout would need a "
    "platform-writable per-tenant value. The catalogue deliberately does not read "
    "that switch: it reports what the school bought whether or not refusals are "
    "being issued. The settlement revokes direct role grants only: a key carried "
    "by a group or a personal ALLOW override stays and is refused at the door, and "
    "moving back up restores nothing. Nothing retries a role the settlement could "
    "not save. The refusal is deterministic, no scheduler runs in this deployment, "
    "and the role waits for the next settlement of that tenant's reach or for "
    "somebody to edit it; its keys are refused at the door in the meantime, so "
    "what is left is a stale row rather than live access (Needs Attention)."
)

M04_GAP_ROW = [
    "P2",
    "A role the settlement could not save waits for a person. The refusal is "
    "deterministic and nothing re-runs it: this deployment has no scheduler, so "
    "the role keeps grants beyond the tenant's depth until the reach is next "
    "settled or somebody edits the role. The keys are refused at the door in the "
    "meantime, so this is a stale row rather than live access. Every surface that "
    "settles a tenant's reach names the role to whoever made the change, the plan "
    "endpoints, the apply_plans command and both vs_config entitlement routes "
    "alike, so nobody is left reading a success message that mentions nothing of "
    "what it left behind.",
    "Decide whether a report of roles standing outside their tenant's depth is "
    "worth a screen of its own, or record that the next settlement of that "
    "tenant's reach trying the role again is the accepted answer.",
    "FR-025",
]

M04_REMOVAL_PREFIX = "• Against v1.18 no item leaves"
M04_REMOVAL_BULLET = (
    "• Against v1.19 no item leaves. One is rewritten to the risk that remains: "
    "every write that settles a tenant's reach now reports the role it could not "
    "save to the operator who made the change, so what is left of that item is "
    "that nothing retries one."
)

M04_TRACEABILITY_LEAD = (
    f"MRD v{MRD_VERSION} records Module 4 as Roles & Permissions (RBAC), Phase V1, "
    "Backend Complete, In use Complete, code vs_rbac, with twenty-two capability "
    "entries. The module number, name, phase, states and ownership agree with this "
    "revision. Each entry maps below; the plan-gate and role-template entries are "
    "restated rather than added to."
)

M04_PLAN_GATE_ENTRY = (
    "Implemented and in service. The catalogue registers the switch ON, a school "
    "is refused under PLAN_UPGRADE_REQUIRED once its plan is applied, and every "
    "narrowing of what the school reaches revokes the direct role grants beyond "
    "it, naming any role whose own dependencies refuse the change in the response "
    "of the write that settled rather than failing that write."
)

M04_VERIFICATION_PREFIX = "•  The reach settlement:"
M04_VERIFICATION_BULLET = (
    "•  The reach settlement: every way a tenant's reach shrinks takes back the "
    "grants beyond it and leaves the ones it still reaches; a closed module and a "
    "first plan application take nothing; a key denied on purpose is still denied; "
    "and a role whose own dependencies refuse the change is left whole, named in "
    "the response of the write that settled it and recorded under "
    "plan_downgrade_blocked, while every other role is settled."
)

M04_CHANGE = (
    "Records that every write which narrows a school's reach names the roles it "
    "could not settle, in the response of the write that settled them. v1.19 "
    "recorded a settlement performed by the vs_config entitlement endpoint as "
    "reaching the audit trail and the logs and nothing else, because that endpoint "
    "passed no collector. The report is now an argument of the settlement itself, "
    "with no default, so a site cannot settle a tenant's roles without saying "
    "where the names go, and both entitlement routes return "
    "data.roles_needing_attention and name the count in their message through one "
    "shared sentence, so the configuration screen and the plan screen describe the "
    "same event in the same words. Clearing a tenant's own entitlement layer "
    "settles too, which it did not before: that row wins over the platform row "
    "while it exists, so deleting it returns the school to the shallower platform "
    "depth at once. No revocation rule changed. FR-025's evidence and limit, the "
    "reach-settlement item in Minimum Verification, the plan-gate traceability "
    "entry and the Needs Attention item are restated; what is left of that item is "
    "that nothing retries a role the settlement could not save. Module 4 remains "
    "Backend Complete and In use Complete, with twenty-two capability entries "
    f"against MRD v{MRD_VERSION}. Verified by vs_rbac 547 tests, with vs_config "
    f"117 and schools.vs_schools 358 beside them. {DEPLOYMENT}"
)


def patch_m04(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{M04_TARGET}"
    )

    replace_cover_version(doc.tables[0], M04_SOURCE, M04_TARGET)
    control = find_table(doc, "Document", contains="Code inspected")
    set_cell_text(field(control, "Version"), M04_TARGET)
    set_cell_text(field(control, "Review date"), REVIEW_DATE)
    set_cell_text(field(control, "Source scope"), M04_SOURCE_SCOPE)
    set_cell_text(field(control, "Code inspected"), M04_CODE_INSPECTED)
    set_cell_text(
        field(control, "MRD baseline"),
        f"XVS Module Requirements Document v{MRD_VERSION}, Module 4, twenty-two "
        "capability entries",
    )

    fr025 = find_fr(doc, "FR-025")
    append_cell_text(field(fr025, "Current evidence"), M04_FR025_EVIDENCE_TAIL)
    set_cell_text(field(fr025, "Limit"), M04_FR025_LIMIT)

    set_run_text(body_paragraph(doc, M04_VERIFICATION_PREFIX), M04_VERIFICATION_BULLET)

    gaps = find_table(doc, "Pri.", contains="P2")
    set_row_text(_gap_row(gaps), M04_GAP_ROW)

    removal = find_table(doc, "REMOVAL RULE")
    replace_line(removal.rows[0].cells[0], M04_REMOVAL_PREFIX, M04_REMOVAL_BULLET)

    set_run_text(body_paragraph(doc, "MRD v2.77 records Module 4"), M04_TRACEABILITY_LEAD)
    traceability = find_table(doc, "MRD Module 4 capability")
    set_cell_text(
        find_row(traceability, "Plan depth gate with its own refusal").cells[2],
        M04_PLAN_GATE_ENTRY,
    )

    prepend_change_log(doc, M04_TARGET, M04_CHANGE)
    finish(doc, output, title, M04_TARGET)


def _gap_row(table):
    """The Needs Attention row about the settlement, found by its own opening words."""
    for row in table.rows:
        if row.cells[1].text.strip().startswith("A role the settlement could not save"):
            return row
    raise ValueError("The settlement gap row was not found")


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    parser.add_argument("--only", nargs="*", choices=["m04", "m06", "m07", "m12"])
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for tag, folder, stem, source, target, patch in (
        ("m07", M07_DIR, M07_STEM, M07_SOURCE, M07_TARGET, patch_m07),
        ("m12", M12_DIR, M12_STEM, M12_SOURCE, M12_TARGET, patch_m12),
        ("m06", M06_DIR, M06_STEM, M06_SOURCE, M06_TARGET, patch_m06),
        ("m04", M04_DIR, M04_STEM, M04_SOURCE, M04_TARGET, patch_m04),
    ):
        if args.only and tag not in args.only:
            continue
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists():
            raise SystemExit(f"Refusing to overwrite {output.name}")
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
