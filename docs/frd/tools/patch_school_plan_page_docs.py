#!/usr/bin/env python3
"""Version the MRD, M01 and M04 for the school plan page.

The console could not answer "what is this school on?", could not change a
plan at all, and could not configure a school that had not gone live - which
is every school it exists to set up.

    python tools/patch_school_plan_page_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    BLACK,
    BLUE,
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)

REVIEW_DATE = "7 September 2026"
CHANGE_DATE = "7 Sep 2026"

MRD_SOURCE, MRD_TARGET = "2.67", "2.68"
M01_SOURCE, M01_TARGET = "1.19", "1.20"
M04_SOURCE, M04_TARGET = "1.15", "1.16"

MRD_SOURCE_SCOPE = (
    "Backend change giving the console a school's plan: what it is on, why each "
    "module sits at its depth, and the two ways to change it. A platform actor "
    "can now configure a school that has not gone live (7 September 2026)"
)

MRD_CHANGE_SUMMARY = (
    "Gave the console the school plan page, and unblocked the schools it "
    "exists to set up. An operator taking a call about a refusal could answer "
    "none of the questions it raises: the effective-capability read returns on "
    "or off per capability and says nothing about depth, so the console could "
    "see that a school had been refused without seeing why, and nothing "
    "anywhere changed a plan, which was set once at school creation and never "
    "again. Each module now reports the depth it reaches, the bands that depth "
    "does and does not buy, and which of three things decided it: the plan's "
    "default, an exception the plan makes for that module, or an uplift ending "
    "on a date. Changing a plan writes the row and re-grants every module in "
    "one transaction, because writing the row alone leaves the record saying "
    "Standard while the product behaves like Basic, a state indistinguishable "
    "from the gate misbehaving. An uplift requires a reason and survives a "
    "plan change. All four routes are the platform's: what a school pays for "
    "is a decision between CodeX and the school. Rebanding is deliberately "
    "absent and stays a code change, because moving a feature between depths "
    "reprices the whole book at once, including schools invoiced last week. "
    "Two defects closed with it: a platform actor asserting a school that had "
    "not gone live was refused by the tenant surface gate and told to complete "
    "that school's onboarding, which is advice for the school given to the "
    "people whose job is to set the school up; and the platform module's "
    "generic bands held no keys and never would, because its keys live in "
    "named siblings at the same depth, so a school's plan page listed "
    "‘Core, Core’. Verified by 499 RBAC, 337 school, 177 console, "
    "171 onboarding and 94 configuration tests. Backend evidence only; nothing "
    "here is deployed."
)

M01_NEW_CAPABILITY = "▸  A school's plan, read and changed from the console"

M01_DECISION = (
    "\n• The plan is readable and changeable, and the two halves of a "
    "change are inseparable. Writing the plan row without re-granting leaves "
    "the school holding the old depths, so the record says one thing and the "
    "product does another; there is no path that writes one without the other. "
    "The read carries the depth per module and which of the plan default, a "
    "plan exception or an uplift decided it, because an operator asked to "
    "change something has to know which of the three to change."
    "\n• Rebanding is not here and should not arrive here. Moving a "
    "feature between depths reprices every school at once, including those "
    "invoiced last week, so it stays a code change that goes through review "
    "and deploy. What the console holds is per-school: one customer, one "
    "decision, audited."
)

M04_DECISION = (
    "\n• A platform actor working on a school is not the school. The "
    "tenant surface gate read the tenant being operated on, so an operator who "
    "asserted a school that had not gone live in order to configure it was "
    "refused and told to complete that school's onboarding. It now exempts a "
    "caller whose own tenant is the platform tenant. The exemption reads the "
    "real actor and does not apply while an impersonation session rides, "
    "because proxied into a school account the caller is the school for that "
    "request; an exemption that can be laundered through a proxy session is "
    "not a boundary. It grants nothing on its own: the caller still needs the "
    "key for the view and the cross-tenant opt-in to assert another tenant."
)

M06_DECISION = (
    "\n• Three more capability rows are retired. The platform module's "
    "generic bands held no permission keys and structurally never would, "
    "because that module's keys live in named siblings at the same depth: "
    "email alerts at Core, bulk import and data export at Plus. An empty band "
    "is not by itself a reason to retire one, and the seeder now records the "
    "difference: Attendance, Gradebook and the portals have empty bands "
    "because nobody has built them, and those are the price list's shape for "
    "what is coming."
)

M01_SOURCE_SCOPE = (
    "Backend change adding the school plan page: the depth per module and its "
    "source, a plan change that re-grants in the same transaction, and "
    "time-boxed uplifts (7 September 2026)"
)

M01_CHANGE_SUMMARY = (
    "Adds the school plan page. GET /v1/i/{slug}/plan/ reports the plan, the "
    "subscription expiry, whether the school is provisioned at all, and for "
    "each module the depth it reaches, the bands that depth buys and does not "
    "buy, and whether the plan's default, a plan exception or an uplift "
    "decided it. PATCH moves the school onto another plan and re-grants every "
    "module in the same transaction, because writing the plan row alone leaves "
    "the school holding the old depths and that state is indistinguishable "
    "from the gate misbehaving; moving to the plan a school is already on is "
    "refused, since a no-op that re-grants eleven modules and audits each one "
    "reads later as a change somebody made. POST and DELETE on "
    "plan/uplifts/ record and withdraw a time-boxed depth uplift for one "
    "module, with a reason required because somebody asks about it months "
    "later, usually when it expires. All four routes are platform-owned: a "
    "school administrator cannot change their own plan however their roles are "
    "composed. Rebanding is deliberately absent and stays a code change. "
    "Verified by 18 focused plan tests and 337 school tests. Backend evidence "
    "only; nothing here is deployed."
)

M04_SOURCE_SCOPE = (
    "Backend change exempting a platform actor, acting as itself, from the "
    "tenant surface gate (7 September 2026)"
)

M04_CHANGE_SUMMARY = (
    "Corrects FR-013 for the one caller it answered wrongly. The tenant "
    "surface gate reads the tenant being operated on, which is right for a "
    "school's own users and backwards for CodeX: an operator who asserted a "
    "school that had not gone live in order to configure it was refused with "
    "the school's own message, telling the platform team to complete that "
    "school's onboarding. Every school spends its first days in that state, "
    "and configuring it is the console's whole job. A caller whose own tenant "
    "is the platform tenant is now exempt. The exemption reads the real actor "
    "rather than the asserted tenant, and does not apply while an "
    "impersonation session rides: proxied into a school account the caller is "
    "the school for that request, and a school that has not gone live reaches "
    "only its onboarding surface however senior the person behind the proxy "
    "is. That half is pinned by a test predating this change. The exemption "
    "grants nothing on its own; the caller still needs the key for the view "
    "and the cross-tenant opt-in to assert another tenant at all. The "
    "reproduction was written first and watched to fail. Verified by 499 RBAC "
    "tests, including 11 pending-tenant surface tests. Backend evidence only; "
    "nothing here is deployed."
)

M01_API_ROWS = [
    ["GET", "/v1/i/{slug}/plan/", "platform.schools.view",
     "The plan, its depth per module, the bands each depth reaches, and "
     "whether the plan default, a plan exception or an uplift decided it."],
    ["PATCH", "/v1/i/{slug}/plan/", "platform.schools.manage",
     "Move the school onto another plan and re-grant every module in the same "
     "transaction. The plan it is already on is refused."],
    ["POST", "/v1/i/{slug}/plan/uplifts/", "platform.schools.manage",
     "Record a time-boxed depth uplift for one module. A reason is required."],
    ["DELETE", "/v1/i/{slug}/plan/uplifts/{capability}/", "platform.schools.manage",
     "Withdraw an uplift, returning the module to the depth its plan pays for."],
]

M01_REFUSAL_ROWS = [
    ["Plan change to the plan already held",
     "Refused rather than applied. A no-op that re-grants every module and "
     "writes an audit event for each reads later as a change somebody made.",
     "400"],
    ["Plan change by a school's own administrator",
     "Refused whatever roles the caller holds. What a school pays for is a "
     "decision between CodeX and the school.",
     "403"],
    ["Uplift naming a band rather than a module",
     "Refused, naming the module to use instead. Depth is a property of a "
     "module; a band is what a depth reaches.",
     "400"],
    ["Uplift with no reason, or an expiry in the past",
     "Refused. The reason is what answers the question somebody asks when the "
     "uplift expires and a school notices something has gone.",
     "400"],
]

M01_TRACE_ROW = [
    "A school's plan, read and changed from the console",
    "FR-007",
    "Implemented. The read carries depth and its source per module; the change "
    "re-grants in the same transaction; uplifts are time-boxed and survive a "
    "plan change. Rebanding stays a code change.",
]

M04_FR013_UPDATES = {
    "Current evidence":
        "TenantSurfaceAllowed refuses a PENDING tenant any view that does not "
        "declare pending_tenant_surface. Absence means closed, so a view added "
        "later is shut until somebody opens it deliberately. One caller is "
        "exempt: a platform actor acting as itself, because working on a "
        "school is not being one and every school spends its first days "
        "pending while CodeX configures it. The exemption reads the real "
        "actor's own tenant, not the asserted one, and is withdrawn while an "
        "impersonation session rides.",
    "Limit":
        "It reads request.tenant first, so under impersonation the tenant "
        "being operated on is the one governed, which is the intended reading "
        "and is what keeps the platform exemption from being laundered through "
        "a proxy session. The exemption is all-or-nothing for platform staff: "
        "there is no per-view narrowing of which pending-tenant surfaces an "
        "operator may reach.",
}


def replace_cell(cell, text, **kwargs):
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def retitle(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_cover_version(table, source, target):
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def update_control_table(table, *, version, source_scope, entries=None,
                         mrd_baseline=None):
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], source_scope, size=9)
        elif label == "Capability entries" and entries:
            replace_cell(row.cells[1], entries, size=9)
        elif label == "MRD baseline" and mrd_baseline:
            replace_cell(row.cells[1], mrd_baseline, size=9)


def prepend_change_log(table, version, date, summary):
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], date, size=8)
    replace_cell(row.cells[2], summary, size=8)


def append_to_cell(cell, addition, *, size):
    replace_cell(cell, cell.text.rstrip() + addition, size=size)


def insert_row_after(table, index, values, *, size=8.2, styled=()):
    template = table.rows[index]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addnext(new_tr)
    row = table.rows[index + 1]
    for idx, value in enumerate(values):
        emphasised = idx in styled
        replace_cell(
            row.cells[idx], value, size=size,
            bold=emphasised, color=BLUE if emphasised else BLACK,
        )
    return row


def keep_rows_whole(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))


def append_capability(table, entry, *, size=8.5):
    cell = table.rows[-2].cells[1]
    replace_cell(cell, cell.text.rstrip() + "\n" + entry, size=size)


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.core_properties.version = MRD_TARGET

    replace_cover_version(doc.tables[0], MRD_SOURCE, MRD_TARGET)
    update_control_table(
        doc.tables[1], version=MRD_TARGET, source_scope=MRD_SOURCE_SCOPE,
        entries="489",
    )
    for row in doc.tables[2].rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(
                row.cells[0], f"5. v{MRD_TARGET} Capability Delta",
                size=9, bold=True, color=BLUE,
            )
            replace_cell(
                row.cells[1],
                "A school's plan on a page, and the console unblocked",
                size=9,
            )
    for row in doc.tables[5].rows:
        if row.cells[0].text.strip() == "1":
            replace_cell(row.cells[5], "21", size=8.5)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == f"5. v{MRD_SOURCE} Capability Delta":
            retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")
        elif text.startswith("This revision records the gate coming into service"):
            retitle(
                paragraph,
                "This revision records the console catching up with the gate. "
                "A school's plan can be read and changed, an operator can "
                "configure a school before it goes live, and three capability "
                "rows that could never hold a key are retired.",
            )

    append_capability(doc.tables[8], M01_NEW_CAPABILITY)
    append_to_cell(doc.tables[8].rows[-1].cells[0], M01_DECISION, size=8.2)
    append_to_cell(doc.tables[14].rows[-1].cells[0], M04_DECISION, size=8.2)
    append_to_cell(doc.tables[18].rows[-1].cells[0], M06_DECISION, size=8.2)

    prepend_change_log(doc.tables[78], MRD_TARGET, CHANGE_DATE, MRD_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, f"XVS Module Requirements Document v{MRD_TARGET}")
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_m01(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M01 School and Branch Management Functional Requirements "
        f"Document v{M01_TARGET}"
    )
    doc.core_properties.title = title
    doc.core_properties.version = M01_TARGET

    schools_api = doc.tables[35]
    refusals = doc.tables[38]
    trace = doc.tables[42]
    change_log = doc.tables[43]

    replace_cover_version(doc.tables[0], M01_SOURCE, M01_TARGET)
    update_control_table(
        doc.tables[1], version=M01_TARGET, source_scope=M01_SOURCE_SCOPE,
        mrd_baseline=(
            f"XVS Module Requirements Document v{MRD_TARGET}, Module 1, "
            "twenty-one capability entries"
        ),
    )
    append_to_cell(doc.tables[6].rows[0].cells[0], M01_DECISION, size=8.2)

    anchor = len(schools_api.rows) - 1
    for values in M01_API_ROWS:
        insert_row_after(schools_api, anchor, values)
        anchor += 1
    keep_rows_whole(schools_api)

    anchor = len(refusals.rows) - 1
    for values in M01_REFUSAL_ROWS:
        insert_row_after(refusals, anchor, values)
        anchor += 1
    keep_rows_whole(refusals)

    insert_row_after(trace, len(trace.rows) - 1, M01_TRACE_ROW)
    keep_rows_whole(trace)

    prepend_change_log(change_log, M01_TARGET, CHANGE_DATE, M01_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_m04(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{M04_TARGET}"
    )
    doc.core_properties.title = title
    doc.core_properties.version = M04_TARGET

    fr013 = doc.tables[21]
    reconciliation = doc.tables[46]
    change_log = doc.tables[47]

    replace_cover_version(doc.tables[0], M04_SOURCE, M04_TARGET)
    update_control_table(
        doc.tables[1], version=M04_TARGET, source_scope=M04_SOURCE_SCOPE,
        mrd_baseline=(
            f"XVS Module Requirements Document v{MRD_TARGET}, Module 4, "
            "twenty capability entries"
        ),
    )

    for row in fr013.rows:
        label = row.cells[0].text.strip()
        if label in M04_FR013_UPDATES:
            replace_cell(row.cells[1], M04_FR013_UPDATES[label], size=8.2)

    replace_cell(
        reconciliation.rows[0].cells[0],
        reconciliation.rows[0].cells[0].text.rstrip().replace(
            f"MRD v{MRD_SOURCE}", f"MRD v{MRD_TARGET}"
        ) + (
            "\n• The count is unchanged. Exempting a platform actor from "
            "the tenant surface gate corrects FR-013 for the one caller it "
            "answered wrongly; it neither adds a product surface nor widens "
            "what anybody may do."
        ),
        size=8.2,
    )
    prepend_change_log(change_log, M04_TARGET, CHANGE_DATE, M04_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root)

    mrd = root / "module-requirements"
    patch_mrd(
        mrd / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx",
        mrd / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
    )
    print(f"Wrote MRD v{MRD_TARGET}")

    fr = root / "functional-requirements"
    m01 = fr / "01-school-and-branch-management"
    stem01 = "XVS_M01_School_and_Branch_Management_Functional_Requirements_Document"
    patch_m01(m01 / f"{stem01}_v{M01_SOURCE}.docx", m01 / f"{stem01}_v{M01_TARGET}.docx")
    print(f"Wrote M01 FRD v{M01_TARGET}")

    m04 = fr / "04-roles-and-permissions-rbac"
    stem04 = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
    patch_m04(m04 / f"{stem04}_v{M04_SOURCE}.docx", m04 / f"{stem04}_v{M04_TARGET}.docx")
    print(f"Wrote M04 FRD v{M04_TARGET}")


if __name__ == "__main__":
    main()
