#!/usr/bin/env python3
"""Cut M21 v1.6, M22 v1.14, M23 v1.14 and M24 v1.5: the procurement queue entries.

What changed in the backend, and therefore in the documents:

* A school's books arrive holding one approval route per procurement document
  type with no steps and no approver group (3aab9000). A document submitted
  against an empty route is refused with APPROVAL_NOT_CONFIGURED, and every
  submit endpoint takes ``confirm_without_approval`` and ``reason`` to send it
  on unreviewed, recorded against the confirmer. The default-ladder endpoint
  is gone (0172c47e). Module 22 carries the rule, Module 23 its three submits.
* The Procurement overview opens to any procurement key, each block behind the
  key of the list it summarises and narrowed to the reader's branches
  (278583e9), and reads by month, term or year to date with a pipeline,
  exceptions and more (03400801). The Spend & suppliers tab and the school's
  non-PO spend limit follow (b7532f90). All three are Module 22's.
* The Stock & receiving tab, a cost centre on a stock issue and the restock
  requisition drafted from what is low (0b490743) are Module 24's.
* A vendor's contact, tax and bank fields are decided per role by Field Access
  switches rather than a sensitive-field key (4767f637, c5549965), which is
  Module 21's.

* Every procurement date that means "today" is the school's own day, read
  through vs_config.clock.tenant_today, under the rule that a school's day is
  its own time zone's, Africa/Lagos by default (e987ba7f, 5071e156). Each of
  the four documents gains a "School's Day" part in Section 2 and names the
  dates it decides in the requirements they belong to.
* A vendor payment names its bank account within the caller's branches and is
  paid from an account of its bills' branch or a school-wide one (61836114,
  1fb5bec5), which is Module 23's new FR-016.
* An eligible bill for a vendor payment returns its branch_id (D63, cee3af8d,
  merged 864468ba), and a vendor payment keeps its branch rules on edit and on
  post (D68, 71c8f89e, merged 0ea6a97c): bills resolve within the caller's
  reach on create, edit and advance allocation, an edit re-derives the
  payment's branch, and posting re-checks the bills and the bank. Both are
  Module 23's alone, so only M23 names them and carries its own baseline.
* A dashboard block needs the school's plan as well as the reader's key
  (4af5d331), so the overview, Spend & suppliers and Stock & receiving blocks,
  and the restock draft's stock check, say so where Modules 22 and 24 record
  them. Module 25 carries the rule itself.

The queue's check of M22 and M23 for the finance abstraction layer's read port
(fa23e06d) found nothing to change: both describe the layer only as the door
for procurement actions. The MRD is left to its own pass; its capability counts
for Modules 21 to 24 (13, 16, 25, 13 in v2.93) are unchanged here.

Each document starts from the newest version in its folder and is written at
the next free number; the script refuses to run if a newer version has
appeared, and ``finish`` refuses to overwrite one.

    python tools/patch_procurement_dashboards_and_routes_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
    table_with_header,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    edit_box,
    insert_row_after,
    row_labelled,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
    set_value,
)

import patch_record_history_docs

REVIEW_DATE = "28 September 2026"
MRD_VERSION = "2.93"
CODE_BASELINE = (
    "Backend main at 27f1da0f, 28 September 2026, holding the finance pass merged at "
    "32194627: the school's day in procurement (e987ba7f, 5071e156), the same-branch bank "
    "rule (61836114, 1fb5bec5) and plan-aware dashboard blocks (4af5d331)"
)
TEST_EVIDENCE = (
    "The dashboard, approval-route and Field Access behaviour was verified by the 56 tests of "
    "the seven vs_procurement test modules it cites, run on the working tree at 7f3756e2 and "
    "all passing; the full vs_procurement suite was 630 tests OK when the Stock & receiving "
    "view was committed. The school's day is covered by ProcurementTodayIsTheSchoolsDayTests "
    "at 27f1da0f. Backend evidence only; nothing here claims frontend delivery or deployment."
)

#: The rule every procurement date follows, in the words the other modules use.
SCHOOL_DAY_RULE = (
    "A school's day is its own time zone's, Africa/Lagos by default: the display.timezone "
    "setting the school keeps (Module 1), read through vs_config.clock.tenant_today for the "
    "school the rows belong to. The server keeps UTC, which in Lagos is still the day before "
    "for the first hour after every midnight, so nothing in this module asks the server for "
    "today."
)

# log_change writes the module's own review date, read from this module.
patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


def spread_cover_lines(doc) -> None:
    """Give each cover line back its own paragraph, and so its own size and weight.

    The cover box holds four styled paragraphs: the title at 20pt, the module at
    14pt, the baseline line at 10pt and the byline at 9.5pt. Some earlier
    versions carry all four lines as breaks inside the title's run, with the
    other three paragraphs left empty, so every line prints as a 20pt title and
    the box grows over the year beneath it. A cover already laid out properly is
    left alone.
    """
    paragraphs = doc.tables[0].rows[0].cells[0].paragraphs
    lines = paragraphs[0].text.split("\n")
    if len(lines) == 1:
        return
    if len(lines) != len(paragraphs) or any(p.text.strip() for p in paragraphs[1:]):
        raise ValueError("The cover's lines and its paragraphs do not match one to one")
    for paragraph, line in zip(paragraphs, lines):
        if not paragraph.runs:
            raise ValueError("A cover paragraph has no run to carry its line")
        mrd_tools.retitle(paragraph, line)


def start(folder: str, stem: str, source: str, target: str, *, source_mrd: str) -> Document:
    """Open ``source``, refusing a stale base, and set the cover and control rows."""
    require_newest(str(ROOT / "functional-requirements" / folder / f"{stem}_v*.docx"), source)
    doc = Document(str(frd_path(folder, stem, source)))
    set_cover_version(doc, source, target)
    spread_cover_lines(doc)
    cover = doc.tables[0].rows[0].cells[0].text
    if f"Version: {target}" not in cover:
        raise ValueError(f"The cover of {stem} v{source} does not carry its version")
    set_control(doc, "Version", target)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", source_mrd)
    return doc


def close(doc, folder: str, stem: str, target: str, summary: str, *stale: str) -> None:
    """Log the revision, refuse any stale claim left outside the log, and write the file.

    Every row below the cover is kept whole on its page. The longer cells this
    revision writes, and the change log rows under them, otherwise break across
    a page and print half a row at the foot of one page and half at the top of
    the next.
    """
    log_change(doc, target, summary)
    retired_site_word = "cam" + "pus"
    assert_absent_outside_log(
        doc, retired_site_word, retired_site_word.capitalize(), *stale,
    )
    for table in doc.tables[1:]:
        mrd_tools.keep_rows_whole(table)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(folder, stem, target), f"{stem.replace('_', ' ')} v{target}", target)


def retitle_fr(doc, fr: str, heading: str) -> None:
    mrd_tools.retitle(paragraph_starting(doc, fr), heading)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _clone_paragraph(paragraph):
    """A copy of ``paragraph`` with no bookmarks, so no anchor is duplicated."""
    clone = copy.deepcopy(paragraph._p)
    for mark in clone.findall(f".//{_W}bookmarkStart") + clone.findall(f".//{_W}bookmarkEnd"):
        mark.getparent().remove(mark)
    return clone


def add_school_day(doc, number: str, body: str, *, contents_old: str, contents_new: str) -> None:
    """Close Section 2 with the school's-day part, and name it in the contents.

    The heading is cloned from Section 2's first sub-heading and the body from
    the paragraph that opens Section 2, so the part carries exactly the look of
    the parts beside it.
    """
    opener = paragraph_starting(doc, "2. Context and Status Model")
    intro = Paragraph(opener._p.getnext(), opener._parent)
    if intro.text.strip().startswith("2.1"):
        raise ValueError("Section 2 opens with its sub-heading, not a paragraph to copy")
    sub_heading = next(p for p in doc.paragraphs if p.text.strip().startswith("2.1 "))
    section3 = paragraph_starting(doc, "3. Actors, Permissions, and Ownership")
    heading = _clone_paragraph(sub_heading)
    paragraph = _clone_paragraph(intro)
    section3._p.addprevious(heading)
    section3._p.addprevious(paragraph)
    mrd_tools.retitle(Paragraph(heading, section3._parent), f"{number} The School's Day")
    mrd_tools.retitle(Paragraph(paragraph, section3._parent), body)
    edit_value(table_headed(doc, "Section", "Purpose"), "2. Context and Status Model",
               contents_old, contents_new)


def add_module1_dependency(doc, after: str, section: str) -> None:
    """Name Module 1 as the keeper of the time zone the school's day is read in."""
    insert_row_after(row_labelled(doc, after), [
        "Module 1, School & Branch Management",
        "Keeps the school's time zone, display.timezone, from which vs_config.clock gives the "
        f"school's day that every date here reads (Section {section}).",
    ])


# ── M22 Procurement & Requisitions ───────────────────────────────────────────

M22_DIR = "22-procurement-and-requisitions"
M22_STEM = "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document"
M22_SOURCE, M22_TARGET = "1.13", "1.14"

M22_FR004 = [
    ("Requirement",
     "A tenant must have its own approval route rather than share one platform ladder, and "
     "nothing about who approves its spend may be invented for it: who approves is the "
     "school's own answer."),
    ("Current evidence",
     "provision_approval_ladders runs inside the transaction that creates a tenant's books, "
     "registered through finance's entity provisioning, and calls "
     "ensure_tenant_approval_templates with with_default_stages false: one tenant-scoped "
     "route for each of the four approvable document types, published with no steps, and no "
     "approver group at all. The empty route stands in front of the shared platform row, "
     "which ensure_default_approval_templates also publishes with no steps, so a change to "
     "that shared row can never begin governing this tenant's spend. A document submitted "
     "against either is refused with APPROVAL_NOT_CONFIGURED unless its submitter confirms "
     "it (FR-003). A school builds its own steps through Module 7's template publishing, POST "
     "/v1/workflow/templates/publish/ under workflow.template.publish, which needs no "
     "procurement key. seed_procurement_approvals is the one path that still publishes the "
     "two-step default ladder, a manager stage and a senior stage above a threshold, each "
     "naming an approver group it creates empty. It runs for a tenant only when an operator "
     "names it, fills in an empty route rather than skipping it, and leaves a route holding a "
     "live step exactly as it is. Migration 0034 removed the two approver groups earlier "
     "provisioning minted, with the steps naming them, sparing a group with a member and one "
     "whose steps a document has run through."),
    ("Acceptance",
     "ProvisionedBooksCarryNoLadderTests proves that provisioning gives every approvable type "
     "a route of its own with no steps, creates no approver group and keeps every step of a "
     "school that already has a ladder, and that the seeding command still publishes the full "
     "ladder and leaves a school's ladder alone when run twice. SeededApproverGroupCleanupTests "
     "covers the migration: an unused group and its steps go, while a group with a member, a "
     "group whose step has run, a group this app did not seed and another app's groups all "
     "survive."),
    ("Current limit",
     "A tenant whose books predate provisioning's routes holds none of its own and resolves "
     "to the shared platform row, unless somebody has run seed_procurement_approvals for it. "
     "It is refused and confirmable exactly as a newer tenant is, but a change to that shared "
     "row would govern its spend."),
]

M22_FR016_HEADING = "FR-016 Open the Procurement Overview to Everyone in Procurement"
M22_FR016 = [
    ("Requirement",
     "Anybody working in procurement must be able to open the overview, and see only the "
     "figures behind the keys they hold, answered for the branches they work in, so that each "
     "figure matches a list they can open."),
    ("Current evidence",
     "GET /reports/dashboard/ opens to any procurement.* key and refuses a caller holding "
     "none. ?window= takes the finance dashboard's windows, this month, the school's term or "
     "the year to date, always read by date and ending on the school's day (Section 2.3), and "
     "the payload names the window chosen, the "
     "windows on offer, the books, the reader's first name, and narrowed, true for a reader "
     "bound to some branches. Each block is computed only for a reader holding its key, at a "
     "school whose plan reaches that key, and is null otherwise, so a block below the school's "
     "band is absent rather than refused (Module 25). Spend, compared with the same number of days before the window, spend "
     "by category, top vendors and committed_vs_spent need procurement.analytics.view; open "
     "purchase orders need procurement.purchase_order.view; overdue bills, bills_due and the "
     "bill exceptions need procurement.vendor_invoice.view; contracts_ending needs "
     "procurement.contract.view; the vendor count needs procurement.vendor.view. The pipeline "
     "runs from requisition to payment in six stages, requisitions, RFQs, orders, received "
     "not billed, bills and paid, each needing its own list's view key and absent without it; "
     "paid covers the window and every other stage is open now. committed_vs_spent sets "
     "issued orders by order date against posted bills by invoice date over the twelve months "
     "from the fiscal year's start. exceptions carries only the non-empty ones of "
     "match_failed, unpaid bills under-received or over-billed; price_variance, bills whose "
     "unit price is outside tolerance against their order rather than the catalogue; "
     "vendor_on_hold, posted unpaid bills of a vendor on hold; and unbilled_receipts, "
     "receipts unbilled after 30 days, which needs procurement.goods_receipt.view. "
     "contracts_ending lists active contracts ending within 90 days with what was ordered on "
     "each. The reader's own approval queue is always present, and its KPI gives the whole "
     "queue's count and amount, how many have waited more than five days, and across how "
     "many document types."),
    ("Acceptance",
     "Every document figure is narrowed to the reader's branches through the same narrowing "
     "their lists use, so a branch-bound reader's figures reconcile with what they can open; "
     "vendors and contracts are master data every branch shares and stay entity-wide. Two "
     "figures that cannot be narrowed are withheld from such a reader instead: a contract's "
     "ordered value, which sums every branch's orders, is null, and the activity feed, read "
     "from the finance audit log, which carries no branch, goes only to an analytics reader "
     "who sees the whole school. Only display-safe fields leave the dashboard; raw audit and "
     "workflow metadata is never exposed. ProcurementDashboardAccessTests proves a "
     "requisition raiser gets their queue and no spend, a school-wide analyst every analytics "
     "block, a branch analyst no activity feed, and a reader with no procurement key a "
     "refusal; PipelineTests, ExceptionTests, BillsAndContractsTests, SpendTests and "
     "EndpointTests cover the blocks."),
    ("Current limit",
     "The thresholds, five days for a slow approval, 30 days for an unbilled receipt and 90 "
     "days for a contract ending, are fixed rather than set by the school."),
]

M22_FR017_HEADING = "FR-017 Show Spend Against the Plan and How Vendors Perform"
M22_FR017 = [
    ("Requirement",
     "Whoever chooses and manages vendors needs spend set against the school's plan, how "
     "concentrated it is, how well vendors deliver, how much is bought without an order and "
     "how long buying takes, under the same keys and branch reading as the overview."),
    ("Current evidence",
     "GET /reports/dashboard/suppliers/ opens on the overview's terms and takes the same "
     "?window=. spend, under procurement.analytics.view, carries value, prior_value, "
     "delta_pct, vendors_with_spend and vendors_for_80pct, how few vendors take four fifths of "
     "it, and plan: net spend this fiscal year to date against the school's approved budget "
     "on the expense accounts purchasing posts to, those on bill lines, order lines, vendors, "
     "categories and catalogue items, so salary lines in the same budget never count. plan is "
     "null with no approved school budget, and always null for a branch-bound reader, because "
     "a school plan measures the whole school. vendors_paid needs "
     "procurement.vendor_payment.view; deliveries, the on-time share of rated receipts, the "
     "accepted share of received quantity and the rejected lines, needs "
     "procurement.goods_receipt.view; non_po, posted bills without an order and their share "
     "of what was billed beside the school's limit, needs procurement.vendor_invoice.view; "
     "open_rfqs, issued RFQs soonest closing first with invited against quoted, needs "
     "procurement.rfq.view; savings needs procurement.quotation.view; vendor_base, active, on "
     "hold, awaiting KYC, ordered from once this fiscal year and added in the window, needs "
     "procurement.vendor.view. The scorecard of the eight largest vendors, with spend, "
     "on-time and accepted shares, open orders and latest assessment grade, and cycle_times "
     "need procurement.analytics.view; by_branch, posted bills by branch with school-wide "
     "bills as their own row, goes only to such a reader who sees the whole school. Savings "
     "compare each awarded quote with the highest submitted quote on the same RFQ, for RFQs "
     "awarded this fiscal year. cycle_times gives median days for each step that ended in "
     "the window: approval, from request date to the first REQUISITION_APPROVED audit record, "
     "then ordering, delivery and payment. The limit is "
     "ProcurementSettings.non_po_spend_limit_pct, a whole percentage from 0 to 100, default 2 "
     "(migration 0035), read and written through GET and PATCH /settings/ under "
     "procurement.settings.view and procurement.settings.update."),
    ("Acceptance",
     "PlanTests proves spend is read against the approved plan on purchasing accounts, that "
     "no approved plan means no plan figure, and that a branch reader never gets the school "
     "plan. DeliveryAndBillingTests covers the on-time and accepted shares, spend without an "
     "order against the school's limit, and the bill-to-payment step; VendorTests the latest "
     "grade, the saving below the highest quote and blocks following their keys; "
     "SuppliersEndpointTests that the route opens to a procurement reader and refuses one "
     "with no procurement key."),
    ("Current limit",
     "The non-PO limit marks spend over it; nothing refuses a bill because of it. KYC is a "
     "status with no expiry date, so the vendor base counts vendors still awaiting KYC rather "
     "than KYC about to lapse."),
]

M22_SUMMARY = (
    "Minor revision. A school's books arrive holding one approval route per procurement "
    "document type with no steps and no approver group, because who approves a school's "
    "spend is its own answer, built through Module 7. FR-004 is rewritten to say so; the "
    "POST /approvals/default-templates/ route it named is gone, seed_procurement_approvals is "
    "the one path that still publishes the two-step default ladder, and migration 0034 "
    "removed the approver groups earlier provisioning minted. Every submit endpoint takes "
    "confirm_without_approval and reason, so a document refused with APPROVAL_NOT_CONFIGURED "
    "against an empty route can be sent on unreviewed and recorded against the confirmer, "
    "and confirming never bypasses a live step (FR-003, the typed errors, and a third "
    "bullet in the open risk). New FR-016 records the overview: open to any procurement key, "
    "each block behind its own key and narrowed to the reader's branches, read by month, term "
    "or year to date, with a six-stage pipeline, committed against spent, top vendors, "
    "exceptions, bills due and contracts ending; purchase_order_status and "
    "monthly_spend_trend are removed and kpis.total_spend_mtd is renamed kpis.spend. New "
    "FR-017 records the Spend & suppliers tab and the school's non-PO spend limit, and FR-015 "
    "keeps the requisition summary. The scope, the actors, the Module 7, finance abstraction "
    "layer, Module 19 and seeding dependencies, the routes, Further Gaps and traceability to "
    f"MRD v{MRD_VERSION} follow; the layer's submit cannot pass the confirmation, and "
    "procurement.approval.update guards no route. Every date that means today is the "
    "school's own day, a school's day being its own time zone's, Africa/Lagos by default "
    "(new Section 2.3): the budget check's default date, a vendor's quotation date, quotation "
    "validity at listing and award, the drafted order's date, the requisition and RFQ "
    "summaries and the dashboard windows, with a Module 1 dependency for the time zone. A "
    "dashboard block needs the school's plan as well as the reader's key, so one below the "
    "school's band is absent rather than refused (FR-016). The cover's four lines, which "
    "had all printed at title size over the year, each take their own size again. Module 22 "
    "stays at 16 capabilities. " + TEST_EVIDENCE
)


def patch_m22() -> None:
    doc = start(M22_DIR, M22_STEM, M22_SOURCE, M22_TARGET,
                source_mrd=f"XVS Module Requirements Document v{MRD_VERSION}")

    # Scope.
    edit_value(table_headed(doc, "Area", "Responsibility"), "Approval submission",
               "Handing a requisition to Module 7, and reacting to the outcome.",
               "Handing a requisition to Module 7, including the confirmation that sends one "
               "past a route with no steps, and reacting to the outcome.")
    insert_row_after(row_labelled(doc, "The vendor's side"), [
        "Dashboards",
        "The Procurement overview and its Spend & suppliers tab, each block behind the key of "
        "the list it summarises and answered under the reader's branches.",
    ])
    edit_value(table_headed(doc, "Not owned here", "Owner"),
               "Who must approve, and how approvers are resolved",
               "This module publishes its default ladders through the engine's service and "
               "reads the engine's answers; it holds no approver rules.",
               "This module publishes each tenant's empty route through the engine's service "
               "when its books are created, and reads the engine's answers; it holds no "
               "approver rules, and a school builds its steps in Module 7.")

    # Actors.
    requester = row_labelled(doc, "Requester")
    keep_format(requester.cells[1],
                "Create and edit a requisition, check budget availability, submit it, and "
                "where its route has no steps yet, confirm it through unreviewed.")
    keep_format(requester.cells[2],
                "procurement.requisition.create/.update, .submit (SENSITIVE); confirming needs "
                "no further key and is recorded against the confirmer")
    admin = row_labelled(doc, "Approval administrator")
    keep_format(admin.cells[1],
                "Read the coverage report: who can approve at each branch and stage, and where "
                "nobody can. The ladders, and the approver groups their steps name, are built "
                "in Module 7.")
    keep_format(admin.cells[2],
                "procurement.approval.view. Publishing a ladder is Module 7's "
                "workflow.template.publish; procurement.approval.update is seeded but no route "
                "checks it.")

    # Requirements.
    fr003 = fr_table(doc, "FR-003")
    append_to(fr003, "Current evidence",
              " Every submit endpoint, for requisitions here and for orders, bills and payments "
              "in Module 23, takes confirm_without_approval and reason and passes both through "
              "submit_for_approval to the engine. Against a route with no live step the engine "
              "refuses with APPROVAL_NOT_CONFIGURED (409) and nothing is written: the document "
              "stays NOT_SUBMITTED. Asked again with confirm_without_approval true, the "
              "document is approved at once and a POSTED_WITHOUT_APPROVAL audit event records "
              "the confirmer and their reason. The engine consults the flag only once it has "
              "found no live step, so it never bypasses a ladder that has one.")
    append_to(fr003, "Acceptance",
              " UnconfiguredSubmitIsRefusedThenConfirmableTests proves the 409 and its code on "
              "all four submit endpoints, that a refusal leaves the document as it was, that "
              "confirming approves it and records who said so, and that a reason alone is not "
              "a confirmation. ConfirmationNeverBypassesARealLadderTests proves that a "
              "confirmation against an unstaffed ladder parks, one against a staffed ladder "
              "waits for the approver, and a ladder whose every step is retired is "
              "confirmable again.")

    retitle_fr(doc, "FR-004", "FR-004 Give Each Tenant Its Own Approval Route, Empty")
    fr004 = fr_table(doc, "FR-004")
    for label, text in M22_FR004:
        set_value(fr004, label, text)

    edit_value(fr_table(doc, "FR-005"), "Current evidence",
               "Both seeded stages set skip_if_no_approvers to false",
               "Both stages of the default ladder seed_procurement_approvals publishes set "
               "skip_if_no_approvers to false")

    retitle_fr(doc, "FR-015", "FR-015 Summarise Requisitions")
    fr015 = fr_table(doc, "FR-015")
    set_value(fr015, "Current evidence",
              "A requisition summary endpoint counts the pipeline by status over exactly the "
              "rows the caller's own requisition list returns, with the same branch narrowing. "
              "The dashboards that summarise the rest of procurement are FR-016 and FR-017.")
    set_value(fr015, "Acceptance",
              "The summary cannot report spend the list does not show, because it counts the "
              "list's own rows.")
    add_fr(doc, M22_FR016_HEADING, "FR-016 | Implemented", M22_FR016)
    add_fr(doc, M22_FR017_HEADING, "FR-017 | Implemented", M22_FR017)

    # Routes and refusals.
    edit_paragraph(doc, "All routes are mounted under /v1/procurement/",
                   "belong to Modules 21, 23 and 24.",
                   "belong to Modules 21, 23 and 24. The dashboard's Stock & receiving tab and "
                   "the restock draft are Module 24's.")
    edit_cell(row_labelled(doc, "POST /requisitions/{id}/submit/").cells[-1],
              "Submit into approval; the response carries the parked warning.",
              "Submit into approval; the response carries the parked warning. "
              "confirm_without_approval and reason send it past a route with no steps (FR-003).")
    retired = row_labelled(doc, "POST /approvals/default-templates/")
    retired._tr.getparent().remove(retired._tr)
    # Inserted after the same anchor, so the overview row lands above the tab's.
    coverage = row_labelled(doc, "GET /approvals/coverage/")
    insert_row_after(coverage, [
        "GET /reports/dashboard/suppliers/",
        "The Spend & suppliers tab, on the same terms (FR-017).",
    ])
    insert_row_after(coverage, [
        "GET /reports/dashboard/",
        "The overview, open to any procurement key; ?window= picks the span, and each block "
        "needs its own key (FR-016).",
    ])
    errors = table_headed(doc, "Condition", "Answer")
    anchor = row(errors, "Reaching an instance that belongs to another entity")
    insert_row_after(anchor, [
        "Opening a dashboard with no procurement key",
        "Refused.",
    ])
    insert_row_after(anchor, [
        "Submitting against a route with no live step, without confirming",
        "409 APPROVAL_NOT_CONFIGURED; nothing is written. With confirm_without_approval true "
        "the document is approved and the confirmer recorded.",
    ])

    # Dependencies.
    fal = row_labelled(doc, "Finance abstraction layer").cells[-1]
    keep_format(fal, fal.text.rstrip() + (
        " Its submit passes no confirmation, so a document it submits against a route with no "
        "steps is refused with no way through; that is harmless only while its procurement "
        "actions are not exposed over HTTP."))
    edit_cell(row_labelled(doc, "Module 7, Workflow & Approval Engine").cells[-1],
              "publishes its default ladders through the engine's publish service",
              "publishes each tenant's empty route through the engine's publish service, passes "
              "the submitter's confirmation through to the engine's submission")
    finance = row_labelled(doc, "Module 19, Finance & Accounting").cells[-1]
    keep_format(finance, finance.text.rstrip() + (
        " It also supplies the dashboard windows and the approved school budget the Spend & "
        "suppliers tab measures against."))
    keep_format(row_labelled(doc, "Seeding").cells[-1],
                "seed_procurement_permissions registers the keys. seed_procurement_approvals "
                "publishes the two-step default ladder, with its approver groups empty, for each "
                "tenant an operator names, and with --platform the stageless shared row. "
                "seed_procurement_dashboard_demo and seed_procurement_suppliers_demo build worked "
                "dashboard data and refuse to run outside DEBUG.")

    # Needs Attention.
    edit_box(table_with_header(doc, "OPEN RISK").rows[0].cells[0], [
        ("replace", "OPEN RISK",
         "OPEN RISK - THREE WAYS PAST AN APPROVER, ALL DELIBERATE"),
        ("after", "• Separately,",
         "• A school's books arrive with every route empty, so until the school builds its "
         "steps, whoever can submit a document can confirm it through unreviewed. It is refused "
         "first, the confirmation has to be asked for by name, and a POSTED_WITHOUT_APPROVAL "
         "audit event records who confirmed and why; once a route holds a live step, "
         "confirming changes nothing."),
    ])
    edit_box(table_with_header(doc, "FURTHER GAPS").rows[0].cells[0], [
        ("replace", "• Tenants created before the ladders",
         "• A tenant whose books predate provisioning's routes holds none of its own and "
         "resolves to the shared platform row, unless somebody has run "
         "seed_procurement_approvals for it. It is refused and confirmable exactly as a newer "
         "tenant is, but a change to that shared row would govern its spend, and nothing "
         "reports which tenants these are."),
        ("after", "• A tenant whose books predate",
         "• The finance abstraction layer's procurement submit cannot pass the confirmation, so "
         "a school module that submits against an empty route is refused with no way through. "
         "It is harmless only while those actions are not exposed over HTTP."),
        ("after", "• Coverage is reported on demand.",
         "• procurement.approval.update is seeded and grantable, but no route checks it, so "
         "ticking it on a role grants nothing."),
    ])

    # Traceability.
    edit_paragraph(doc, "Module 22 carries 16 capability entries", "MRD v2.91",
                   f"MRD v{MRD_VERSION}")
    paragraph = paragraph_starting(doc, "Module 22 carries 16 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Books arriving with an empty route, and the confirmation that sends a document past "
        "one, strengthen workflow submission and the typed refusals without changing the "
        "count; the two dashboards strengthen requisition summaries."))
    trace = table_headed(doc, "MRD capability", "Requirements")
    set_value(trace, "Requisition summaries", "FR-015, FR-016, FR-017")
    set_value(trace, "Typed refusals for missing rules, parked work, and override",
              "FR-003, FR-004, FR-005, FR-007")

    # The school's day.
    add_school_day(
        doc, "2.3", SCHOOL_DAY_RULE + (
            " Here it decides the date a budget-availability check reads when none is given, "
            "the date a vendor's quotation carries when the vendor opens or saves it on the "
            "public form, whether a quotation has passed its validity date, both where it is "
            "listed and when it is awarded, the order date of the purchase order an award "
            "drafts, the month the requisition summary compares, the RFQs the RFQ summary counts "
            "as closing soon, and the day the overview and Spend & suppliers windows end on."),
        contents_old="branch scope, and requirement states",
        contents_new="branch scope, the school's day, and requirement states")
    append_to(fr_table(doc, "FR-002"), "Current evidence",
              " With no date given, the check reads the plan as at the school's day (Section 2.3).")
    append_to(fr_table(doc, "FR-012"), "Current evidence",
              " A quotation the vendor opens or saves is dated on the school's day.")
    append_to(fr_table(doc, "FR-013"), "Current evidence",
              " A quotation whose validity date is before the school's day is flagged expired on "
              "the list and in the drawer and cannot be awarded, and the drafted order is dated "
              "on the school's day unless the award names a date.")
    append_to(fr_table(doc, "FR-015"), "Current evidence",
              " Its month to date, set against the same number of days of the month before, ends "
              "on the school's day.")
    add_module1_dependency(doc, "Module 23, Purchase Orders, Delivery & AP", "2.3")
    paragraph = paragraph_starting(doc, "Module 22 carries 16 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Dating every procurement day by the school's own time zone strengthens budget "
        "checks, quotation capture, the award decision and requisition summaries without "
        "changing the count."))

    close(doc, M22_DIR, M22_STEM, M22_TARGET, M22_SUMMARY,
          "MRD v2.91", "default-templates", "publishes its default ladders",
          "an endpoint offers the same", "procurement.approval.view or procurement.approval.update")


def flatten_closing_paragraph(doc) -> None:
    """Shrink the empty paragraph a document must end on to a hairline.

    A body ends on a paragraph, and here it follows the change log table. When the
    table's last row reaches the foot of a page, that paragraph at full size spills
    onto a page of its own, which prints blank.
    """
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    closing = doc.paragraphs[-1]
    if closing.text.strip() or closing._p.getprevious().tag != f"{_W}tbl":
        raise ValueError("The document does not end on an empty paragraph after a table")
    fmt = closing.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(1)
    mark = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(f"{_W}val", "2")
    mark.append(size)
    closing._p.get_or_add_pPr().append(mark)


# ── M23 Purchase Orders, Delivery & AP ───────────────────────────────────────

M23_DIR = "23-purchase-orders-delivery-and-ap"
M23_STEM = "XVS_M23_Purchase_Orders_Delivery_and_AP_Functional_Requirements_Document"
M23_SOURCE, M23_TARGET = "1.13", "1.14"

#: Module 23 alone carries the vendor-payment branch changes, so its baseline names them.
M23_CODE_BASELINE = (
    "Backend main at 7860ea86, 28 September 2026, carrying a vendor payment's branch rules on "
    "edit and on post (71c8f89e, merged at 0ea6a97c), the eligible bill's branch_id "
    "(cee3af8d, merged at 864468ba), and the finance pass merged at 32194627: the school's day "
    "in procurement (e987ba7f, 5071e156), the same-branch bank rule (61836114, 1fb5bec5) and "
    "plan-aware dashboard blocks (4af5d331)"
)

M23_FR016_HEADING = "FR-016 Pay a Branch's Bills From That Branch's Money"
M23_FR016 = [
    ("Requirement",
     "A payment settling one branch's bills must come out of that branch's money or the "
     "school's, and nobody may name a bank account outside the branches they work in."),
    ("Current evidence",
     "A vendor payment names its bank account through finance's bank resolver, which reads "
     "only the caller's reach: their own branches' accounts and the school-wide ones. An "
     "account outside it answers 404, No bank account in this entity, exactly as an unknown id "
     "does, so a payment neither draws on another branch's money nor confirms that the account "
     "exists. A new payment takes its branch from the bills it settles, and once they are "
     "resolved its account must be that branch's or a school-wide one; a payment settling "
     "bills of several branches, which only a caller bound to no branch can select, belongs to "
     "the school and may use any account in reach. Another branch's account is refused 400 on "
     "bank_account before anything is written, naming the branch: This vendor payment belongs "
     "to Ikeja Branch. Pay it from an Ikeja Branch account or a school-wide one. The bills "
     "themselves resolve within the caller's branches, read as every procurement document is, "
     "on create, on edit and when an advance is applied, so a bill another branch keeps "
     "answers 400 on allocations exactly as an unknown one does: Every invoice must be posted "
     "and belong to the selected vendor. An edit re-derives the payment's branch from its new "
     "bills as a new payment does, saves it, and checks the bank account against it with the "
     "same 404 and 400. Posting asks the three questions again before any money moves, since a "
     "draft can be posted later, by someone else, after a bill or an account has changed: "
     "every bill must still be in the caller's reach (the 400 of an unknown bill); the bills "
     "must still give the branch the payment carries, which is the branch its journal books "
     "to, or the answer is 400 on allocations, The bills this payment settles now belong to "
     "another branch. Edit the draft to bring it up to date before posting.; and the bank "
     "account must be in reach (404) and of that branch or school-wide (400). The eligible-bill "
     "picker, GET /vendor-payments/eligible-invoices/, offers only bills in reach and returns "
     "each bill's branch_id, so a screen can offer only the accounts the chosen bills allow. "
     "The same reach governs "
     "every ledger account a procurement route names, a receipt or bill line's expense account "
     "included: a ledger account behind a bank account outside the caller's reach is refused "
     "as unknown."),
    ("Acceptance",
     "Mrs Okafor is bursar at Ikeja and at Lekki, so both accounts are in her bank list. "
     "Settling an Ikeja bill from the Lekki account is refused with that message and writes no "
     "payment, and from the Ikeja account or the school-wide one it goes through. A bursar at "
     "Ikeja alone who names the Lekki account gets the 404 of an unknown account. "
     "BankAccountNamedInAPostingTests.test_a_vendor_payment and "
     "DocumentPaidFromItsOwnBranchTests.test_a_vendor_payment cover both. "
     "VendorPaymentBranchTests in tests_vendor_payment_branch.py proves a Lekki bill is "
     "unknown to Ikeja's officer on create and on edit, an edit moves the payment to its "
     "bills' branch and checks the bank there, and posting a draft whose bank moved to another "
     "branch, or that settles another branch's bill, is refused with nothing posted."),
    ("Current limit", "None."),
]

M23_SUMMARY = (
    "Minor revision. Submitting a purchase order, vendor invoice or vendor payment takes "
    "confirm_without_approval and reason, as a requisition does. Against a route with no "
    "steps, which is what a school's books now arrive holding, the submit is refused with "
    "APPROVAL_NOT_CONFIGURED and writes nothing, and a confirmed retry approves the document "
    "and records the confirmer; confirming never bypasses a live step. FR-001, FR-009, the "
    "submit routes, the typed errors and the Module 7 dependency follow, and the finance "
    "abstraction layer's row and Further Gaps record that its submit cannot pass the "
    "confirmation. FR-004's limit and the price-variance gap are restated: the procurement "
    "overview counts unpaid bills whose match flagged a price variance, while the variance "
    "account itself is still reported nowhere. Every date that means today is the school's "
    "own day, a school's day being its own time zone's, Africa/Lagos by default (new Section "
    "2.3): overdue bills and the bill summary, the purchase-order summary, AP ageing, cash "
    "requirements, GR/IR ageing and detail, a vendor's open bills and a payment reversal's "
    "default date, with a Module 1 dependency; a goods receipt through the finance abstraction "
    "layer names its acting user as the caller. New FR-016 records that a vendor payment names "
    "its bank account within the caller's branches, another branch's answering 404, and is "
    "paid from an account of its bills' branch or a school-wide one, another branch's refused "
    "400 naming the branch. Its bills resolve within the caller's reach on create, edit and "
    "advance allocation, an out-of-reach bill answering 400 as an unknown one, where create "
    "had answered 403 \"This document belongs to another branch.\"; an edit re-derives the "
    "payment's branch from its new bills and checks the bank against it; and posting re-checks "
    "the bills' reach, that they still give the stored branch (400 asking for an edit), and the "
    "bank's reach (404) and branch (400) before any money moves (71c8f89e). The eligible-bill "
    "list returns each bill's branch_id (cee3af8d). FR-008, the routes, the typed errors, the "
    "vendor payment record and the Module 19 dependency follow. The cover's four lines, which had all printed at title size over the "
    "year, each take their own size again. "
    f"Traceability is reconciled to MRD v{MRD_VERSION} at 25 capabilities. " + TEST_EVIDENCE + (
    " The bank rule is covered by BankAccountNamedInAPostingTests and "
    "DocumentPaidFromItsOwnBranchTests in vs_finance at 27f1da0f, and the edit and post "
    "rules by VendorPaymentBranchTests in vs_procurement, read at 7860ea86 and not re-run.")
)


def patch_m23() -> None:
    doc = start(M23_DIR, M23_STEM, M23_SOURCE, M23_TARGET,
                source_mrd=f"XVS Module Requirements Document v{MRD_VERSION}")
    set_control(doc, "Code baseline", M23_CODE_BASELINE)

    edit_value(fr_table(doc, "FR-001"), "Current evidence",
               "A draft is editable intent; submission hands the order to the approval engine.",
               "A draft is editable intent; submission hands the order to the approval engine, "
               "and takes confirm_without_approval and reason, which send it past a route with "
               "no steps and are recorded against the confirmer (Module 22 FR-003).")
    set_value(fr_table(doc, "FR-004"), "Current limit",
              "Nothing reports on the purchase price variance account itself. The procurement "
              "overview counts the unpaid bills whose match flagged a price variance (Module 22 "
              "FR-016), but the amount that reached the account is visible only to somebody "
              "reading the ledger.")
    fr009 = fr_table(doc, "FR-009")
    append_to(fr009, "Current evidence",
              " Submitting a bill or a payment takes the same confirm_without_approval and "
              "reason as every procurement submit: against a route with no live step the engine "
              "refuses with APPROVAL_NOT_CONFIGURED and nothing is written, and a confirmed "
              "retry approves the document and records the confirmer. The flag is never "
              "consulted while a live step exists, so it cannot skip a ladder the school has "
              "built.")
    append_to(fr009, "Acceptance",
              " UnconfiguredSubmitIsRefusedThenConfirmableTests runs the refusal and the "
              "confirmation on the purchase-order, vendor-invoice and vendor-payment submit "
              "endpoints as well as the requisition's.")

    edit_cell(row_labelled(doc, "POST /purchase-orders/{id}/submit/").cells[-1],
              "Hand off to approval.",
              "Hand off to approval; confirm_without_approval and reason send it past a route "
              "with no steps.")
    edit_cell(row_labelled(doc, "POST /vendor-invoices/{id}/submit/, /post/").cells[-1],
              "Approval, then AP posting.",
              "Approval, then AP posting. The submit takes the same confirmation.")
    edit_cell(row_labelled(doc, "POST /vendor-payments/{id}/submit/, /post/").cells[-1],
              "Approval, then the posting.",
              "Approval, then the posting. The submit takes the same confirmation.")
    errors = table_headed(doc, "Condition or route", "Answer")
    insert_row_after(row(errors, "Reversing the approval behind a posted bill or payment"), [
        "Submitting an order, bill or payment against a route with no live step, without "
        "confirming",
        "409 APPROVAL_NOT_CONFIGURED; nothing is written. With confirm_without_approval true the "
        "document is approved and the confirmer recorded.",
    ])

    fal = row_labelled(doc, "Finance abstraction layer").cells[-1]
    keep_format(fal, fal.text.rstrip() + (
        " Its submit passes no confirmation, so an order, bill or payment it submits against a "
        "route with no steps is refused with no way through; that is harmless only while its "
        "procurement actions are not exposed over HTTP."))
    edit_cell(row_labelled(doc, "Module 7, Workflow & Approval Engine").cells[-1],
              "through the ladders Module 22 publishes.",
              "through each tenant's own route, which its books arrive holding with no steps "
              "until the school builds them (Module 22 FR-004).")

    edit_box(table_with_header(doc, "FURTHER GAPS").rows[0].cells[0], [
        ("sub", "• A price variance outside tolerance",
         "but nothing reports on that account, so it is visible only to somebody reading the "
         "ledger.",
         "but nothing reports on that account. The procurement overview counts the unpaid bills "
         "whose match flagged a variance; the amount that reached the account is visible only to "
         "somebody reading the ledger."),
        ("append",
         "• The finance abstraction layer's submit cannot pass the confirmation, so an order, "
         "bill or payment a school module submits against an empty route is refused with no way "
         "through. It is harmless only while those actions are not exposed over HTTP."),
    ])

    edit_paragraph(doc, "Module 23 carries 25 capability entries",
                   "in MRD v2.78. Each maps to the requirements below. Cancelling a purchase "
                   "order with a recorded reason is the entry added in that version.",
                   f"in MRD v{MRD_VERSION}. Each maps to the requirements below. Cancelling a "
                   "purchase order with a recorded reason is the entry added in MRD v2.78. The "
                   "confirmation that sends an order, bill or payment past a route with no "
                   "steps strengthens workflow approval integration without changing the count.")

    # The school's day.
    add_school_day(
        doc, "2.3", SCHOOL_DAY_RULE + (
            " Here it decides when a bill is overdue, on its detail, in the overdue tab and in "
            "the bill summary; the month the purchase-order summary compares; the day AP "
            "ageing, cash requirements, GR/IR ageing and its detail, and a vendor's open bills "
            "are read as at when none is given; and the date a payment reversal takes when the "
            "caller supplies none."),
        contents_old="money units, and requirement states",
        contents_new="money units, the school's day, and requirement states")
    edit_value(fr_table(doc, "FR-010"), "Current evidence",
               "on a date the caller supplies.",
               "on a date the caller supplies, or the school's day when they supply none.")
    append_to(fr_table(doc, "FR-012"), "Current evidence",
              " Each is read as at the school's day unless the caller names a date, and a bill "
              "is overdue from the day after its due date in the school's own time zone "
              "(Section 2.3).")
    fal = row_labelled(doc, "Finance abstraction layer").cells[-1]
    keep_format(fal, fal.text.rstrip() + (
        " A goods receipt it records is dated on the school's day unless it names one, and "
        "names its acting user as the caller whose reach the receipt's line accounts are read "
        "under."))
    add_module1_dependency(doc, "Module 6, Configuration & Capability Management", "2.3")

    # A branch's bills are paid from that branch's money.
    add_fr(doc, M23_FR016_HEADING, "FR-016 | Implemented", M23_FR016)
    errors = table_headed(doc, "Condition or route", "Answer")
    anchor = row(errors, "Paying a vendor on hold or failing KYC")
    insert_row_after(anchor, [
        "Posting a vendor payment whose bills no longer give the branch it carries",
        "400 on allocations: The bills this payment settles now belong to another branch. Edit "
        "the draft to bring it up to date before posting. Nothing is posted. A bill or bank "
        "account that has left the caller's reach, or a bank of another branch, answers as on "
        "create.",
    ])
    insert_row_after(anchor, [
        "Naming a bill outside the caller's branches on a vendor payment's create, edit or "
        "advance allocation",
        "400 on allocations, exactly as an unknown bill: Every invoice must be posted and belong "
        "to the selected vendor.",
    ])
    insert_row_after(anchor, [
        "Paying a branch's bills from another branch's bank account",
        "400 on bank_account, naming the branch: pay it from that branch's account or a "
        "school-wide one. Nothing is written.",
    ])
    insert_row_after(anchor, [
        "Naming a bank account outside the caller's branches",
        "404, exactly as an unknown account.",
    ])
    edit_cell(row_labelled(doc, "GET /vendor-payments/eligible-invoices/").cells[-1],
              "What can be settled.",
              "What can be settled: up to 100 posted, unpaid bills within the caller's branches, "
              "oldest due first, each with its branch_id, which decides the bank accounts the "
              "payment may use.")
    append_to(fr_table(doc, "FR-008"), "Current evidence",
              " Each bill is resolved within the caller's branches, so another branch's bill "
              "answers 400 as an unknown one does (FR-016).")
    edit_value(table_headed(doc, "Model", "Holds", "Notes"), "VendorPayment",
               "Approval authorises the plan; posting creates the history.",
               "Approval authorises the plan; posting creates the history. It takes its branch "
               "from the bills it settles and is paid from that branch's account or a "
               "school-wide one; an edit re-derives the branch and posting re-checks it.")
    finance = row_labelled(doc, "Module 19, Finance & Accounting").cells[-1]
    keep_format(finance, finance.text.rstrip() + (
        " Its bank resolver decides which bank accounts a caller may name and which a branch's "
        "payment may be paid from (FR-016)."))
    set_value(table_headed(doc, "MRD capability", "Requirements"), "Vendor payment creation",
              "FR-007, FR-009, FR-016")
    paragraph = paragraph_starting(doc, "Module 23 carries 25 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Paying a branch's bills from that branch's money strengthens vendor payment creation, "
        "and dating by the school's own day the payables reports, without changing the count."))

    flatten_closing_paragraph(doc)
    close(doc, M23_DIR, M23_STEM, M23_TARGET, M23_SUMMARY,
          "the ladders Module 22 publishes", "surfaced on no screen", "MRD v2.81",
          "held to neither", "checks the account against the branch the payment carries")


# ── M24 Inventory & Stock Ledger ─────────────────────────────────────────────

M24_DIR = "24-inventory-and-stock-ledger"
M24_STEM = "XVS_M24_Inventory_and_Stock_Ledger_Functional_Requirements_Document"
M24_SOURCE, M24_TARGET = "1.4", "1.5"

M24_FR012_HEADING = "FR-012 Show the Stores Their Stock and What Is Arriving"
M24_FR012 = [
    ("Requirement",
     "Whoever keeps the stores needs one view of what stock is worth and where, what is "
     "running low and how long it will last, what went out and to whom, and what is due in, "
     "answered for the stores they work in."),
    ("Current evidence",
     "GET /reports/dashboard/stock/ opens to any procurement key and takes the overview's "
     "?window= (Module 22 FR-016), ending on the school's day (Section 2.2). position, running_low, by_store, issued, movements, turns "
     "and adjustments need procurement.stock.view; receipts and unbilled need "
     "procurement.goods_receipt.view; expected needs procurement.purchase_order.view; a block "
     "the reader may not see is null, and so is one whose key the school's plan does not "
     "reach, which is absent rather than refused (Module 25). Stock figures answer for the reader's stores, their "
     "branches' and the school-wide ones, as the stock screens read them (FR-009); receipts "
     "and orders are documents and answer under the reader's branches. position counts what "
     "is below its reorder level and still holds something apart from what is out of stock "
     "and holds nothing, as the stock summary counts them, so no item is in both, and adds "
     "stock unmoved for 90 days. An item is running low at or below its reorder level; days "
     "left is on-hand divided by the average daily quantity issued over the last 90 days, or "
     "since the item first moved when that is shorter, and an item never issued has none; "
     "the suggested order is the reorder quantity or enough to reach the reorder level again, "
     "whichever is more. issued sums the window's issues by the cost centre named on each, "
     "untagged issues as their own row, at the cost they left stock at. turns divides the "
     "cost issued in the window by the average of opening and closing stock value, and names "
     "the fastest store. adjustments counts and values the window's adjustments. receipts "
     "gives this week's deliveries and the share of the window's received lines that came "
     "short or rejected; unbilled the receipts still waiting for a bill, and how many after "
     "30 days; expected the orders not fully received that are late or due within 14 days of the school's day."),
    ("Acceptance",
     "LowStockTests proves an item at or below its reorder level is low with its days left, "
     "that a Lekki storekeeper reads the school-wide store and their own and no other, that "
     "the view counts what is low and what went to whom, and that the stock blocks need the "
     "stock key. StockEndpointTests proves the route opens to a stock reader."),
    ("Current limit",
     "Stock has no transfers and no stock-take record, so the view shows the window's "
     "adjustments rather than a stock-take variance."),
]

M24_FR013_HEADING = "FR-013 Draft a Requisition for What Is Low"
M24_FR013 = [
    ("Requirement",
     "A storekeeper who sees stock running low must be able to start the purchase from there, "
     "without retyping the items or choosing the quantities."),
    ("Current evidence",
     "POST /stock-items/restock-requisition/ drafts one requisition, at the branch a "
     "requisition raised by the caller takes (Module 22, Section 2.2), with a line for every "
     "active item at or below its reorder level in the caller's stores, read by the same "
     "function as the running-low block, so the view and the draft never differ. Each line "
     "takes the suggested quantity and is estimated at the item's average cost on hand, else "
     "its last receipt cost, else its catalogue price; the needed-by date is the procurement "
     "settings' default lead time from the school's day, which is also its request date, "
     "where one is set. item_ids narrows the draft to "
     "some of those items, and an id that is not low in the caller's stores is ignored, so "
     "the quantities are always the server's own. The requisition is a DRAFT and is never "
     "submitted here: it goes through review and approval like any other."),
    ("Acceptance",
     "It needs procurement.requisition.create and procurement.stock.view, and is refused 403 "
     "without the second, as it is where the school's plan does not reach stock; when nothing is low it answers 400 and drafts nothing. "
     "RestockDraftTests covers one line per low item at the suggested quantity, the refusal "
     "without stock access and the empty case."),
    ("Current limit",
     "A person has to ask for the draft; nothing raises one on a schedule or when an item "
     "crosses its level."),
]

M24_SUMMARY = (
    "Minor revision. An issue can name the cost centre the stock went to, an active cost "
    "centre of the same books, kept on the movement (migration 0036) and on the expense "
    "journal line, and returned on the movement ledger (FR-005). New FR-012 records the Stock "
    "& receiving view of the Procurement dashboard, each block behind the stock, receipt or "
    "order key, with stock read for the caller's stores; new FR-013 records the restock "
    "requisition drafted from what is low, which needs the requisition create key and stock "
    "view. The scope, actors, routes, typed errors, the stock movement record, the Module 19 "
    "and Module 22 dependencies and two Further Gaps follow; the stock-take gap now says the "
    "view shows adjustments in its place. Every date that means today is the school's own "
    "day, a school's day being its own time zone's, Africa/Lagos by default (new Section 2.2): "
    "an issue's and an adjustment's default date, the restock draft's request and needed-by "
    "dates and the Stock & receiving view, with a Module 1 dependency. A dashboard block, and "
    "the restock draft's stock check, need the school's plan as well as the reader's key "
    f"(FR-012, FR-013). Traceability is reconciled to MRD v{MRD_VERSION} at "
    "13 capabilities. " + TEST_EVIDENCE
)


def patch_m24() -> None:
    doc = start(M24_DIR, M24_STEM, M24_SOURCE, M24_TARGET,
                source_mrd=f"XVS Module Requirements Document v{MRD_VERSION} | Module 24")

    scope = table_headed(doc, "Area", "Responsibility")
    set_value(scope, "Issues",
              "Valuing an outflow at the current moving average, posting it to expense, and "
              "recording the cost centre it went to.")
    set_value(scope, "Movement ledger and reporting",
              "The immutable movement history, valuation and reorder reports, the Stock & "
              "receiving view, and the restock requisition drafted from what is low.")

    issuer = row_labelled(doc, "Issuer")
    keep_format(issuer.cells[1],
                "Issue stock to an expense account, naming the cost centre it went to.")
    keep_format(issuer.cells[2],
                "procurement.stock.issue (SENSITIVE), which also opens Module 19's cost-centre "
                "list so the issuer can pick one")
    insert_row_after(row_labelled(doc, "Stock counter"), [
        "Restock drafter",
        "Draft one requisition for everything at or below its reorder level in their stores.",
        "procurement.requisition.create, plus procurement.stock.view",
    ])

    fr005 = fr_table(doc, "FR-005")
    append_to(fr005, "Current evidence",
              " An issue may name the cost centre the stock went to, by id or code. It must be "
              "an active cost centre of the same books, and it is kept on the movement and on "
              "the expense line of the journal, so the issue reports under that cost centre "
              "wherever spending is read by one. The movement ledger returns cost_center_id "
              "and cost_center_name.")
    append_to(fr005, "Acceptance",
              " IssueCostCentreTests proves the cost centre is kept on the movement and the "
              "expense line, that one from other books is refused, and that a storekeeper may "
              "read the cost centres to pick one.")
    add_fr(doc, M24_FR012_HEADING, "FR-012 | Implemented", M24_FR012)
    add_fr(doc, M24_FR013_HEADING, "FR-013 | Implemented", M24_FR013)

    edit_cell(row_labelled(doc, "POST /stock-items/{id}/issue/").cells[-1],
              "Issue at the current average.",
              "Issue at the current average, optionally naming the cost centre it went to.")
    insert_row_after(row_labelled(doc, "GET /stock-items/summary/"), [
        "POST /stock-items/restock-requisition/",
        "Draft one requisition for what is low in the caller's stores (FR-013).",
    ])
    insert_row_after(row_labelled(doc, "GET /stock-movements/"), [
        "GET /reports/dashboard/stock/",
        "The Stock & receiving view, each block behind its key (FR-012).",
    ])
    errors = table_headed(doc, "Condition", "Answer")
    anchor = row(errors, "Moving stock without naming a location when the caller can reach "
                         "more than one")
    insert_row_after(anchor, [
        "Drafting a restock requisition without stock access, or when nothing is low",
        "403 without procurement.stock.view or where the school's plan does not reach it; 400 "
        "when nothing in the caller's stores is at or "
        "below its reorder level.",
    ])
    insert_row_after(anchor, [
        "Issuing to a cost centre that is inactive or belongs to other books",
        "Refused: no such active cost centre in this entity.",
    ])

    movement = row_labelled(doc, "Stock movement")
    keep_format(movement.cells[1],
                "Kind, quantity, value, the running balance after, and for an issue the cost "
                "centre it went to.")
    keep_format(movement.cells[2],
                "Immutable; corrections are further movements. The cost centre is optional, and "
                "cannot be deleted while a movement names it.")

    finance = row_labelled(doc, "Module 19, Finance & Accounting").cells[-1]
    keep_format(finance, finance.text.rstrip() + (
        " It owns cost centres too: an issue's cost centre rides on the expense journal line, "
        "and the cost-centre list opens to procurement.stock.issue so a storekeeper can pick "
        "one. The Stock & receiving view reads its dashboard windows."))
    insert_row_after(row_labelled(doc, "Module 23, Purchase Orders, Delivery & AP"), [
        "Module 22, Procurement & Requisitions",
        "Receives the restock requisition this module drafts, as an ordinary draft that is "
        "reviewed and approved like any other. The overview and Spend & suppliers tabs beside "
        "the Stock & receiving view are its FR-016 and FR-017.",
    ])

    edit_box(table_with_header(doc, "FURTHER GAPS").rows[0].cells[0], [
        ("sub", "• There is no cycle-count", "and no variance approval.",
         "and no variance approval. The Stock & receiving view therefore shows the window's "
         "adjustments rather than a stock-take variance."),
        ("replace", "• Reorder indicators are reported on demand",
         "• Reorder indicators are reported on demand, and a person can turn them into a draft "
         "requisition in one step (FR-013), but nothing raises a reorder or warns on a "
         "schedule."),
    ])

    edit_paragraph(doc, "Module 24 carries 13 capability entries", "MRD v2.77",
                   f"MRD v{MRD_VERSION}")
    paragraph = paragraph_starting(doc, "Module 24 carries 13 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " The cost centre on an issue strengthens stock issues, the Stock & receiving view "
        "inventory summaries and reports, and the restock draft reorder indicators, without "
        "changing the count."))
    trace = table_headed(doc, "MRD capability", "Requirements")
    set_value(trace, "Reorder indicators", "FR-008, FR-013")
    set_value(trace, "Inventory summaries and reports", "FR-008, FR-012")

    # The school's day.
    add_school_day(
        doc, "2.2", SCHOOL_DAY_RULE + (
            " Here it decides the date an issue or an adjustment takes when none is given, the "
            "request and needed-by dates of a restock requisition, and the day the Stock & "
            "receiving view is read as at: its window, this week's receipts, the orders late or "
            "due within 14 days, and the stock unmoved for 90 days."),
        contents_old="Valuation without floats, and requirement states",
        contents_new="Valuation without floats, the school's day, and requirement states")
    append_to(fr_table(doc, "FR-005"), "Current evidence",
              " An issue that names no date is dated on the school's day (Section 2.2).")
    append_to(fr_table(doc, "FR-006"), "Current evidence",
              " An adjustment that names no date is dated on the school's day.")
    add_module1_dependency(doc, "Module 6, Configuration & Capability Management", "2.2")
    paragraph = paragraph_starting(doc, "Module 24 carries 13 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Dating issues, adjustments and the restock draft by the school's own day changes no "
        "entry."))

    close(doc, M24_DIR, M24_STEM, M24_TARGET, M24_SUMMARY, "v2.77", "19a50b3")


# ── M21 Vendor Management ────────────────────────────────────────────────────

M21_DIR = "21-vendor-management"
M21_STEM = "XVS_M21_Vendor_Management_Functional_Requirements_Document"
M21_SOURCE, M21_TARGET = "1.5", "1.6"

M21_SUMMARY = (
    "Minor revision. A vendor's nine contact, tax and bank fields are decided per role by "
    "Module 4's Field Access switches rather than by the retired "
    "procurement.vendor.view_sensitive key. A role that cannot read one gets the vendor "
    "without it, one that can read but not change it finds it named in _read_only_fields, and "
    "a submitted field it cannot write is refused 403 field_write_denied with nothing saved. "
    "The conversion set each role's switches from the key it held, so nobody's access "
    "changed. The scope, the payments-facing reader, the field protection box, FR-002, the "
    "Vendor record, the typed errors, the Module 4 and Module 22 dependencies and "
    f"traceability to MRD v{MRD_VERSION} follow. Every date that means today is the "
    "school's own day, a school's day being its own time zone's, Africa/Lagos by default (new "
    "Section 2.2): a milestone's completion date, an assessment's default date, a contract's "
    "lapse, the renewals list and the contract summary, and the vendor year to date, with a "
    "Module 1 dependency (FR-006 to FR-008). Module 21 stays at 13 capabilities. "
    + TEST_EVIDENCE
)


def patch_m21() -> None:
    doc = start(M21_DIR, M21_STEM, M21_SOURCE, M21_TARGET,
                source_mrd=f"XVS Module Requirements Document v{MRD_VERSION} | Module 21")

    edit_value(table_headed(doc, "Area", "Responsibility"), "Vendor profile",
               "with field-level protection over the sensitive parts.",
               "with each contact, tax and bank field read and written only by a role whose "
               "switch allows it.")
    reader = row_labelled(doc, "Payments-facing reader")
    keep_format(reader.cells[1],
                "See, and where allowed change, a vendor's contact, tax and bank detail.")
    keep_format(reader.cells[2],
                "Module 4's per-role Field Access switches: Read to see a field, Write to change "
                "it. Each field is declared sensitive, so a role reaches it only when its switch "
                "is on.")
    edit_box(table_with_header(doc, "FIELD-LEVEL PROTECTION").rows[0].cells[0], [
        ("replace", "• Contact, tax, bank name",
         "• Email, phone, address, contact people, tax ID, bank name, bank code, account number "
         "and account name are the vendor's registered fields, each with a Read and a Write "
         "switch per role (Module 4). A role that cannot read one gets the vendor without it; "
         "one that can read but not change it finds it named in _read_only_fields; a submitted "
         "field it cannot write is refused 403 field_write_denied, naming every such field, "
         "and nothing is saved."),
    ])

    fr002 = fr_table(doc, "FR-002")
    edit_value(fr002, "Current evidence", "Sensitive fields are gated on read and write.",
               "Contact, tax and bank fields are governed per role by Module 4's Field Access "
               "switches: the vendor detail omits one the caller cannot read and names in "
               "_read_only_fields one they can read but not change, and the create and update "
               "routes refuse, before anything is written, a submitted field the caller cannot "
               "write.")
    edit_value(fr002, "Acceptance",
               "A caller without the sensitive key neither sees nor can change bank detail.",
               "A role whose switches are off neither sees nor can change bank detail, and the "
               "contact people nested under a vendor are held to their own switch rather than "
               "escaping below it (VendorDeepPayloadTests).")

    edit_value(table_headed(doc, "Model", "Holds", "Notes"), "Vendor",
               "Sensitive fields gated on read and write;",
               "Contact, tax and bank fields governed per role;")
    errors = table_headed(doc, "Condition", "Answer")
    denied = row(errors, "Reading or writing bank detail without the sensitive key")
    keep_format(denied.cells[0],
                "Submitting a contact, tax or bank field the caller's role cannot write")
    keep_format(denied.cells[-1],
                "Refused 403 field_write_denied, naming each such field; nothing is saved. A "
                "field the role cannot read is absent from the response rather than refused.")

    edit_cell(row_labelled(doc, "Module 4, Roles & Permissions").cells[-1],
              "the field-level policy over contacts, tax and bank detail,",
              "the Field Access switches that decide per role who reads and writes the vendor's "
              "nine contact, tax and bank fields,")
    module22 = row_labelled(doc, "Module 22, Procurement & Requisitions").cells[-1]
    keep_format(module22, module22.text.rstrip() + (
        " Its Spend & suppliers tab scores the largest vendors and counts the vendor base (its "
        "FR-017)."))

    edit_paragraph(doc, "Module 21 carries 13 capability entries", "MRD v2.76",
                   f"MRD v{MRD_VERSION}")
    paragraph = paragraph_starting(doc, "Module 21 carries 13 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Governing the vendor's contact, tax and bank fields by role switches strengthens the "
        "contacts and verified bank details entry without changing the count."))

    # The school's day.
    add_school_day(
        doc, "2.2", SCHOOL_DAY_RULE + (
            " Here it decides the date a milestone is completed on and an assessment is "
            "recorded on when neither names one, when a contract has lapsed, both where its "
            "list and detail flag it expired and in the expiry and missed-milestone runs, which "
            "contracts the renewals list and the contract summary count as ending, and where "
            "the vendor summary and insights start the year to date."),
        contents_old="Snapshots, stable identifiers, and requirement states",
        contents_new="Snapshots, stable identifiers, the school's day, and requirement states")
    append_to(fr_table(doc, "FR-006"), "Current evidence",
              " A milestone completed without a date is completed on the school's day, and a "
              "contract past its end date reads expired from the school's next day (Section "
              "2.2).")
    append_to(fr_table(doc, "FR-007"), "Current evidence",
              " The list, the contract summary and the expiry run measure from the school's "
              "day.")
    append_to(fr_table(doc, "FR-008"), "Current evidence",
              " An assessment recorded without a date is dated on the school's day. The model "
              "field keeps the server's date as its default, because a model default cannot see "
              "a school, and the one route that creates an assessment always passes the "
              "school's day.")
    add_module1_dependency(doc, "Module 6, Configuration & Capability Management", "2.2")
    paragraph = paragraph_starting(doc, "Module 21 carries 13 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Dating contracts, milestones and assessments by the school's own day changes no "
        "entry."))

    close(doc, M21_DIR, M21_STEM, M21_TARGET, M21_SUMMARY,
          "v2.76", "sensitive key", "sensitive-field key", "sensitive policy",
          "field-level RBAC", "13675db")


def main() -> None:
    patch_m22()
    patch_m23()
    patch_m24()
    patch_m21()


if __name__ == "__main__":
    main()
