#!/usr/bin/env python3
"""Cut M11 v2.12: the switches decide a student's fields, and the server holds two rules.

What the backend does, and therefore what the document must say:

* A child's blood group, allergies and conditions are sensitive Field Access
  fields. ``school.students.view_sensitive`` and ``FieldSecurityMixin`` are
  retired (vs_rbac 0025 turned every grant into switches, 0026 deleted the
  key), and a field a role may not write is refused 403 ``field_write_denied``
  on every surface (4767f637, c5549965).
* Enrolment is judged by the same switches as an edit: a typed medical value
  without the Write switch is refused and nothing is created, a blank one is
  dropped, and the fields a new record needs are open on create, the
  enrolment date and a new guardian's phone among them (9557ad6e, cf9cdbad).
  The directory row enforces the enrolment date's Read switch.
* ``school.guardians`` carries fields and no permission keys (196e300a).
* Both spreadsheet imports read a guardian's name in parts, with the one-line
  column an optional fallback that is split and flagged (e420d602).
* Confirming an applicant keeps the school's admission-number rule
  (79d01ebc), and the 2-to-25 age bound is the server's on enrolment and edit
  as well as on the import (0a11452a).
* A promotion that would fill a target class past its capacity is refused
  until acknowledged, as enrolling one child is, and the preview lists those
  classes first (94b4eebe; the school app's confirm step, school-fe 08a5cd5).

Five sentences the module's code had already outgrown are corrected with it: the
standard validation refusals answer 400, not 422; setting the admission-number
policy needs ``school.students.update``; the capacity override needs no key
beyond the route's; section 5.2 lists the students resource's keys as seeded;
and section 8.1 counts them.

The MRD and every other FRD are left alone.

    python tools/patch_student_switches_and_server_rules_docs.py
"""
from __future__ import annotations

from docx import Document

from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    append_rows,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    table_with_header,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    all_text_outside_log,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    insert_after,
    normalise_change_log,
    repair_ooxml,
    row,
    row_starting,
)

import patch_record_history_docs

REVIEW_DATE = "27 September 2026"
patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

M11_DIR = "11-student-management"
M11_STEM = "XVS_M11_Student_Management_Functional_Requirements_Document"
SOURCE, TARGET = "2.11", "2.12"
SOURCE_MRD = "XVS Module Requirements Document v2.93"

CODE_BASELINE = (
    "Backend main at e12e7fbf, where the student module stands as 94b4eebe left it, "
    "27 September 2026, with the school app at 35124d1 and 99ba49b for what it hides and "
    "greys and at 08a5cd5 for the promotion's capacity step"
)
TEST_EVIDENCE = (
    "schools.vs_students ran 334 tests, all passing, at 94b4eebe, as that commit records. "
    "Backend evidence only; nothing here claims deployment."
)


def table_with_row(doc, label: str):
    """The one table holding a row whose first cell reads ``label``."""
    hits = [t for t in doc.tables if any(r.cells[0].text.strip() == label for r in t.rows)]
    if len(hits) != 1:
        raise ValueError(f"{label!r} labels rows in {len(hits)} tables")
    return hits[0]


def set_row(table, label: str, text: str) -> None:
    keep_format(row(table, label).cells[-1], text)


# ── section 2: why this revision exists, and its corrections ─────────────────

CORRECTION_TAIL_OLD = "most of them in places versions 2.7 and 2.8 did not reach."
CORRECTION_TAIL_NEW = (
    CORRECTION_TAIL_OLD
    + " Version 2.11 adds three rows after those, for record history, Field Access over the "
    "whole record and a guardian's name in parts. Version 2.12 adds five, each correcting this "
    "document against the module's own code: the retired medical key, the 403 a refused field "
    "answers, the write rule on enrolment, the guardian name columns on both imports, and the "
    "age bound the server now holds as well as the form."
)

CORRECTIONS = [
    [
        "v2.4 to v2.11 (sections 3.7 and 7.1, FR-001, FR-004 and 12.1): "
        "school.students.view_sensitive gates reading and writing blood_group, allergies and "
        "conditions, enforced by FieldSecurityMixin.",
        "Neither exists. The three medical fields are sensitive Field Access fields of "
        "school.students: closed by default, and read and corrected only by a role whose "
        "switches a school has turned on. vs_rbac migration 0025 turned every grant of the key "
        "into those switches, on school roles, on the prebuilt library and on personal "
        "exceptions, so nobody's access moved; 0026 deleted the key with its grants, and "
        "vs_rbac/fls.py, which held FieldSecurityMixin, is deleted. A switch carries no plan "
        "band, so no plan buys a child's medical record.",
        "apps/schools/vs_students/field_access.py; apps/vs_rbac/field_enforcement.py; "
        "apps/vs_rbac/field_conversion.py, the school.students.view_sensitive entry; "
        "apps/vs_rbac/migrations/0025 and 0026; commits 4767f637 and c5549965.",
    ],
    [
        "v2.4 to v2.11 (FR-004 rule 2 and acceptance (1) and (2), section 12.1): a field the "
        "caller may not write answers a per-field 422.",
        "403 field_write_denied, naming every offending field at once, and nothing is saved. A "
        "hidden field and a read-only one are refused in the same words, so a refusal tells a "
        "caller nothing they did not already know. On an edit, a value equal to the stored one "
        "is dropped rather than refused for a role that may read the field, because a form "
        "echoes the whole record.",
        "apps/vs_rbac/field_enforcement.py, FieldWriteDenied and FieldAccessMixin; "
        "apps/schools/vs_students/tests/test_field_write_paths.py.",
    ],
    [
        "v2.4 to v2.11 (FR-002 and section 7.1): enrolment carries no field write rule, and "
        "enrolment_date is writable only by the student lifecycle keys.",
        "Enrolment is judged by the same switches as an edit. A typed blood group, allergy or "
        "condition from a role whose Write switch does not reach it is refused 403 and nothing "
        "is created; a blank, whitespace-only or absent value is dropped, because the enrol "
        "form posts every input. The enrolment date is open on create: whoever enrols a pupil, "
        "by form or by spreadsheet, sets it, and its Write switch decides only who corrects it "
        "on an existing record. The conversion wrote Write off on that switch for every role "
        "that did not hold the key which used to decide it, so no role gained the correction.",
        "apps/schools/vs_students/serializers.py, EnrolmentWriteSerializer and "
        "StudentWriteSerializer; field_access.py; apps/vs_rbac/field_conversion.py, the "
        "school.students.manage entry; tests/test_field_write_paths.py; commit 9557ad6e.",
    ],
    [
        "v2.11 (FR-012, FR-024 and FR-028 rule 2): both spreadsheet imports carry a guardian's "
        "name on one line, so every imported guardian arrives split and flagged.",
        "Both templates carry Guardian First, Middle and Last Name ahead of Guardian Name, which "
        "is optional. A row naming the parts creates a guardian with no name to check; a row "
        "with only the one line imports with a warning and is split and flagged; once any part "
        "is given, first and last are both required. Every template column is a required "
        "header, so a file built on a template downloaded before the name columns existed is "
        "refused until the school downloads the new one.",
        "apps/schools/vs_students/imports.py, read_guardian_name; guardian_imports.py; "
        "apps/vs_import_data/migrations/0020_guardian_name_part_columns.py; "
        "apps/vs_import_data/services/template_validation.py; commit e420d602.",
    ],
    [
        "v2.9 to v2.11 (FR-012 rule 9): the 2-to-25 age bound is something the enrolment form "
        "refuses, which the import repeats.",
        "The form refused it and the server did not: the enrol and edit endpoints accepted any "
        "date, so a direct request, or a form that skipped its own check, could put a "
        "28-year-old or an unborn child on the roll. The bound, and a birth date not in the "
        "future, live in ages.py, and the import, the enrolment serializer and the edit "
        "serializer all ask it, in one set of words.",
        "apps/schools/vs_students/ages.py; serializers.py, _plausible_birth_date; imports.py; "
        "commit 0a11452a.",
    ],
]

# ── section 3.7 and section 7 ────────────────────────────────────────────────

MEDICAL_KEY_ROW = [
    "No key: blood_group, allergies and conditions",
    "Sensitive field",
    "school_admin, branch_admin, as Read and Write switches",
    "Who reads and corrects a child's medical fields is a Field Access switch on the role "
    "(FR-027), so a school opens them to its nurse without handing over anything else on the "
    "record. The key that decided it, school.students.view_sensitive, is retired: vs_rbac "
    "migration 0025 turned its grants into switches and 0026 deleted it. Note teacher holds "
    "neither switch.",
]

STUDENT_FIELDS = {
    "date_of_birth": (
        "DateField(). Required. Used by the advisory duplicate check in FR-002. A pupil is at "
        "least 2 and at most 25 years old, counted in calendar years, and not born in the "
        "future: the enrol and edit endpoints and both imports refuse any other date in one "
        "set of words, from ages.py (FR-002 rule 12)."
    ),
    "blood_group": (
        'CharField(max_length=4, blank=True, default=""). A sensitive Field Access field '
        "(FR-027): absent unless the role's Read switch is on, and written only under its "
        "Write switch, on enrolment as on an edit."
    ),
    "allergies": (
        'CharField(max_length=200, blank=True, default=""). Sensitive, as blood_group is. '
        "Never present in a list serializer."
    ),
    "conditions": (
        'CharField(max_length=200, blank=True, default=""). Sensitive, as blood_group is. '
        "Never present in a list serializer."
    ),
    "emergency_contact_name": (
        'CharField(max_length=150, blank=True, default=""). A registered field whose switches '
        "start open. An emergency contact only a school administrator can read is useless in "
        "the emergency it exists for, so a school closes it for a role only by choice; and "
        "the field carries an adult's name rather than a child's medical history."
    ),
    "emergency_contact_phone": (
        'CharField(max_length=32, blank=True, default=""). Open by default, for the reason '
        "above. It is still absent from every list serializer and from the export dataset."
    ),
    "enrolment_date": (
        "DateField(default=timezone.localdate). A registered field, open by default and open "
        "on create: whoever enrols a pupil, by form or by spreadsheet, sets the date, and its "
        "Write switch decides who may correct it on an existing record. The conversion that "
        "retired the keys wrote Write off for every role that could not correct it before, so "
        "no role gained the correction. The directory row enforces its Read switch as the "
        "profile does."
    ),
}

GUARDIAN_NAME_LABEL = "Guardian.first_name / middle_name / last_name / full_name"
GUARDIAN_NAME = (
    "first_name, middle_name and last_name are CharField(max_length=100), the name as the "
    "school holds it; full_name is CharField(max_length=150), the one-line form every list, "
    "search and label reads; name_needs_review is a BooleanField set on a name split from one "
    "line until a person confirms it. While the flag is set, full_name keeps the text exactly "
    "as typed; once confirmed, it is composed from the parts on every save. A first and a last "
    "name are required of every route that names a guardian in parts. FR-028."
)
GUARDIAN_PHONE_TAIL = (
    " A Field Access field open on create: every route that adds a guardian requires it, so "
    "whoever adds the guardian sets it, and its Write switch decides only who corrects it "
    "(FR-027)."
)

# ── requirements ─────────────────────────────────────────────────────────────

FR001_EXPOSURE = (
    "The list serializer carries no medical field at all and no guardian contact details. The "
    "detail serializer carries blood_group, allergies and conditions only for a role whose Read "
    "switch on each is on; they are sensitive, so closed until a school opens them, and absent "
    "from every list. emergency_contact_name and emergency_contact_phone are absent from the "
    "list and present on the detail, their switches open by default (section 7.1). Every "
    "registered field follows its Read switch on the list, the detail and the type-ahead "
    "alike, the enrolment date on the directory row included (FR-027). The list row also "
    "carries applied_on, which the applicant board sorts by."
)

FR002_RULES_TAIL = (
    " (12) A date of birth that makes the child under 2 or over 25, counted in calendar years, "
    "or that lies in the future, is refused 400 on date_of_birth, in the sentence the import "
    "uses, from schools.vs_students.ages; the edit route refuses the same dates. The rule is "
    "the server's and not only the form's, so a direct request cannot put a 28-year-old or an "
    "unborn child on the roll. (13) Enrolment is judged by Field Access exactly as an edit is "
    "(FR-027). A typed blood_group, allergies or conditions from a role whose Write switch "
    "does not reach it is refused 403 field_write_denied, naming every such field, and nothing "
    "is created; a blank, whitespace-only or absent value is dropped, because the enrol form "
    "posts every input. The name parts, the date of birth, the gender, the student number, the "
    "enrolment date and a new guardian's first name, last name and phone are open on create: "
    "whoever enrols the pupil sets them, and their Write switches decide only who corrects "
    "them afterwards."
)
FR002_ACCEPTANCE_TAIL = (
    " (19) A date of birth 28 years back, one in the current year and one next year are each "
    "refused 400 and create nothing (tests/test_behaviour.py, "
    "test_a_birth_date_no_pupil_has_is_refused_as_the_import_refuses_it). (20) A role without "
    "the medical Write switches is refused 403 for any typed medical value, naming each, and "
    "creates nothing; the same role enrols with blank or absent medical values, and sets an "
    "enrolment date it could not correct afterwards (tests/test_field_write_paths.py, "
    "EnrolmentMedicalFieldTests and EnrolmentDateTests)."
)

FR003_RULES_TAIL = (
    " (4) The school's admission-number rule holds at confirmation as it does at enrolment "
    "(FR-019). Where the rule requires a number, confirming an applicant whose record carries "
    "none answers 422 ADMISSION_NUMBER_REQUIRED and leaves them an applicant, unless the "
    "confirmation supplies one; a number supplied is held to the pattern and to the "
    "per-school uniqueness exactly as an enrolled child's is. An applicant whose record "
    "already carries a number confirms without one. The rule is the server's, not only the "
    "confirm form's."
)
FR003_ACCEPTANCE_TAIL = (
    " (4) With the rule requiring a number, confirming without one answers 422 "
    "ADMISSION_NUMBER_REQUIRED and the applicant stays an applicant; confirming with one "
    "succeeds; an applicant whose record already carries a number confirms without one "
    "(tests/test_behaviour.py, StatusMachineTests)."
)

FR004_KEY = (
    "school.students.update. Every registered field it carries also needs its Field Access "
    "Write switch (FR-027). The medical fields are sensitive, so closed until a school opens "
    "them for the role; every other field, the enrolment date and the emergency contact "
    "included, is open by default, and the conversion wrote Write off on the enrolment date "
    "for every role that could not correct it before (section 7.1)."
)
FR004_RULE2_OLD = (
    "(2) Field-level write gates are declared on the serializer through "
    "FieldSecurityMixin.write_permissions, so an unauthorised field answers a per-field 422 "
    "naming it, rather than being silently dropped."
)
FR004_RULE2_NEW = (
    "(2) Field write rules are Field Access switches, enforced by the serializer through "
    "FieldAccessMixin: a submitted field the role may not write is refused 403 "
    "field_write_denied, naming every offending field at once, and nothing is saved, rather "
    "than being silently dropped. Every present value is judged, a blank included, because on "
    "an existing record a blank erases what is stored; a value equal to the stored one is "
    "dropped for a role that may read the field, because a form echoes the whole record."
)
FR004_ACCEPTANCE_1_OLD = (
    "(1) A caller without view_sensitive attempting to write allergies, conditions or "
    "blood_group answers 422 naming the field, and the record is unchanged; the same caller "
    "may write emergency_contact_name and emergency_contact_phone."
)
FR004_ACCEPTANCE_1_NEW = (
    "(1) A role without the Write switch on allergies, conditions or blood_group attempting "
    "to write one, a blank included, answers 403 field_write_denied naming the field, and the "
    "record is unchanged; the same role may write emergency_contact_name and "
    "emergency_contact_phone, whose switches start open."
)
FR004_ACCEPTANCE_2_OLD = (
    "(2) A caller without the Field Access write switch for school.students.enrolment_date "
    "attempting to write enrolment_date answers 422."
)
FR004_ACCEPTANCE_2_NEW = (
    "(2) A role without the Write switch for school.students.enrolment_date attempting to "
    "write enrolment_date answers 403 field_write_denied and the stored date is unchanged "
    "(tests/test_field_write_paths.py, EditRouteIsUnchangedTests)."
)

FR012_COLUMNS_OLD = (
    "first_name, last_name, middle_name, date_of_birth, gender, student_number, branch, class, "
    "guardian_full_name, guardian_phone, guardian_email, guardian_relationship, address, "
    "previous_school."
)
FR012_COLUMNS_NEW = (
    "first_name, middle_name, last_name, date_of_birth, gender, student_number, admission_date, "
    "branch, class, guardian_first_name, guardian_middle_name, guardian_last_name, "
    "guardian_full_name, guardian_phone, guardian_email, guardian_relationship, address, "
    "previous_school: eighteen, in the order a school reads them."
)
FR012_RULE9_OLD = (
    "(9) A file is refused everything the enrolment form refuses and more, before any row is "
    "written: a date of birth that makes the child under 2 or over 25,"
)
FR012_RULE9_NEW = (
    "(9) A file is refused everything the enrol endpoint refuses and more, before any row is "
    "written: a date of birth that makes the child under 2 or over 25 or lies in the future, "
    "asked of the same ages.py the enrol and edit endpoints ask (FR-002 rule 12),"
)
FR012_RULES_TAIL = (
    " (11) The guardian's name is read from Guardian First, Middle and Last Name where the row "
    "gives them, which creates a guardian with no name to check. Guardian Name, the one-line "
    "column, is optional and a fallback: a row that fills only it imports with a warning, and "
    "the guardian is split and flagged for a person to confirm (FR-028). Once any part is "
    "given, first and last are both required, and a row naming no guardian at all is refused. "
    "Every template column is a required header, so a file built on a template downloaded "
    "before the name columns existed is refused until the school downloads the new one. "
    "(12) The admission date is open on create, like the enrolment date it fills: an uploader "
    "whose role may not correct a pupil's enrolment date still imports one (FR-027)."
)
FR012_ACCEPTANCE_TAIL = (
    " (10) A guardian named in parts arrives with no name to check; a one-line name imports "
    "and is flagged; a first name without a last is refused, and so is a row naming no "
    "guardian (ImportExecutionTests). (11) The template carries eighteen columns, the "
    "guardian name parts ahead of the optional one-line name "
    "(test_the_template_exists_with_all_eighteen_columns and "
    "test_both_templates_ask_for_the_guardian_name_in_parts). (12) An uploader without the "
    "enrolment date's Write switch validates and executes a file carrying admission dates, "
    "and the dates are saved (tests/test_field_write_paths.py, ImportEnrolmentDateTests)."
)

FR024_COLUMNS_OLD = (
    "guardian_full_name, guardian_phone, guardian_email, student_number, student_first_name, "
    "student_last_name, student_date_of_birth, relationship, is_primary, occupation, address."
)
FR024_COLUMNS_NEW = (
    "guardian_first_name, guardian_middle_name, guardian_last_name, guardian_full_name, "
    "guardian_phone, guardian_email, student_number, student_first_name, student_last_name, "
    "student_date_of_birth, relationship, is_primary, occupation, address: fourteen."
)
FR024_RULES_TAIL = (
    " (9) The guardian's name is read as FR-012 rule 11 reads it: the parts where given, "
    "creating a guardian with no name to check, and the optional one-line Guardian Name only "
    "as a fallback that is split and flagged; first and last are both required once any part "
    "is given."
)
FR024_ACCEPTANCE_OLD = (
    "(6) The template exists with all eleven columns, and a school may import its own "
    "households."
)
FR024_ACCEPTANCE_NEW = (
    "(6) The template exists with all fourteen columns, the guardian name parts ahead of the "
    "optional one-line name, and a school may import its own households. (7) A guardian named "
    "in parts is created with no name to check, and a row naming only a last name is refused."
)

FR027_RULE1_TAIL_OLD = (
    "school.guardians registers the three name parts and photo beside the contact fields."
)
FR027_RULE1_TAIL_NEW = (
    FR027_RULE1_TAIL_OLD
    + " It carries fields and no permission keys: a guardian is read with school.students.view "
    "and corrected with school.students.update, so the resource exists to hang switches on and "
    "mints no second set of keys."
)
FR027_RULE3_OLD = (
    "(3) The name parts, date of birth, gender and student number are open on create for a "
    "student, and a guardian's first and last name are open on create: enrolling is not "
    "correcting."
)
FR027_RULE3_NEW = (
    "(3) The name parts, date of birth, gender, student number and enrolment date are open on "
    "create for a student, and a guardian's first name, last name and phone are open on "
    "create: enrolling is not correcting. Every route that adds a guardian requires a phone, "
    "so without the flag a role that may not correct one could add no guardian at all. A "
    "guardian's email, address and occupation stay governed on create, a blank one dropped. "
    "The flags reach the database through sync_field_registry, which seed_all_permissions "
    "runs on every deploy."
)
FR027_RULES_TAIL = (
    " (6) The switches decide what every surface shows and accepts, in place of the keys that "
    "were once checked inside serializers. A field a role may not read is absent, with no "
    "placeholder; one it may read and not change is present and named in _read_only_fields "
    "on a detail; a submitted field it may not write is refused 403 field_write_denied naming "
    "each, and nothing is saved. The student directory row enforces the enrolment date's Read "
    "switch as the profile does. (7) /me and the sign-in response carry the caller's "
    "field_access map, whose open_on_create names the fields an Add form may still ask for. "
    "The school app hides a field a role may not read, greys one it may not change, and keeps "
    "an Add form's open-on-create fields (school-fe 35124d1 and 99ba49b); nothing here claims "
    "deployment."
)
FR027_ACCEPTANCE_4_OLD = (
    "(4) The deep payload case renders every surface closed and finds no registered name at "
    "any depth."
)
FR027_ACCEPTANCE_4_NEW = (
    "(4) The deep payload case renders every surface closed, the directory row included, and "
    "finds no registered name at any depth (tests/test_field_access_deep_payload.py). (5) A "
    "new guardian's phone is saved by a role that may not correct one, on enrolment and on the "
    "link form; correcting it later is refused 403, and a new guardian's email without its "
    "switch is still refused (tests/test_guardian_phone_on_create.py)."
)

FR028_RULE2_OLD = (
    "(2) A new guardian given only a one-line name, as the spreadsheet imports still do, is "
    "split and flagged the same way; one given its parts is composed and not flagged."
)
FR028_RULE2_NEW = (
    "(2) A new guardian given only a one-line name, by an older client or by a spreadsheet row "
    "that fills only Guardian Name, is split and flagged the same way; one given its parts, on "
    "any route and on both imports (FR-012 rule 11, FR-024 rule 9), is composed and not "
    "flagged."
)
FR028_TRIGGER_TAIL = (
    " The students and guardians spreadsheet imports, with Guardian First, Middle and Last Name."
)

# ── sections 11, 12 and 14 ───────────────────────────────────────────────────

REFUSALS = [
    [
        "A submitted field the caller's role may not write, on an edit or on enrolment",
        "403",
        "field_write_denied, naming every offending field. Raised by FieldAccessMixin; a hidden "
        "field and a read-only one are refused in the same words, and nothing is saved.",
    ],
    [
        "A date of birth that makes the student under 2 or over 25, or lies in the future, on "
        "enrolment or an edit",
        "400",
        "REQUEST_ERROR, detail naming date_of_birth in the import's own sentence, from "
        "schools.vs_students.ages.",
    ],
]

LIST_READ_OLD = (
    "and absent from the detail response for a caller without school.students.view_sensitive. "
    "emergency_contact_name and emergency_contact_phone are absent from every list and present "
    "on the detail for every caller, which is section 7.1's deliberate exception."
)
LIST_READ_NEW = (
    "and absent from the detail response for a role whose Read switch on them is off. "
    "emergency_contact_name and emergency_contact_phone are absent from every list and present "
    "on the detail for every role, because their switches start open, which is section 7.1's "
    "deliberate exception."
)
PATCH_REFUSAL_START = "A caller without school.students.view_sensitive who submits"
PATCH_REFUSAL_NEW = (
    "A role without the Write switch on allergies or conditions that submits either, on a "
    "PATCH or on enrolment, is refused 403 field_write_denied naming each, and nothing is "
    "stored or created."
)
POLICY_ANCHOR = "Two tenants with different admission-number policies each validate"
CONFIRM_BULLET = (
    "Confirming an applicant whose record carries no number, where the school's policy "
    "requires one, is refused 422 ADMISSION_NUMBER_REQUIRED and the applicant stays an "
    "applicant (FR-003)."
)
AGE_BULLET = (
    "A date of birth making the child under 2 or over 25, or in the future, is refused by the "
    "enrol and edit endpoints and by the import, in the same words (FR-002, FR-004, FR-012)."
)

BAND_OLD_ROW = (
    "and view_sensitive carries no band at all, so no plan buys a child's medical record."
)
BAND_NEW_ROW = (
    "and a child's medical fields are Field Access switches, which carry no band at all, so no "
    "plan buys a child's medical record."
)
BAND_OLD_DECISION = (
    "and view_sensitive carries no band at all, so a school cannot buy its way into a child's "
    "medical record."
)
BAND_NEW_DECISION = (
    "and a child's medical fields are Field Access switches, which carry no band at all, so a "
    "school cannot buy its way into a child's medical record."
)

SUPERSEDED_FLS = (
    " Superseded by the v2.12 rows below: the key and the mixin are both retired, and the "
    "medical fields are Field Access switches."
)
SUPERSEDED_MEDICAL = (
    " Superseded in its mechanism by the v2.11 and v2.12 rows below: the first three are "
    "sensitive Field Access fields, and the last two are registered with switches that start "
    "open."
)

# ── corrections found while reading the module against its code ──────────────

#: Section 11 rows that named the handler's standard validation refusal as 422.
VALIDATION_ROWS = {
    "Required field missing or malformed": (
        "400",
        "REQUEST_ERROR, the field errors in detail, for a serializer's own refusal; a Django "
        "validation error answers 400 VALIDATION_ERROR with the same field-keyed detail. "
        "Standard, rendered by the handler. Every such refusal in this module is a "
        "serializer's or a DRF ValidationError, so it answers REQUEST_ERROR.",
    ),
    "A branch omitted where the caller could write": (
        "400",
        "REQUEST_ERROR naming the field. Standard. There is no null branch for a student to "
        "fall back to, so an omitted branch that cannot be inferred is a missing field and not "
        "a shared row.",
    ),
    "A status change or placement with no effective date": (
        "400",
        "REQUEST_ERROR naming the field. Standard. The serializer defaults it to today, so "
        "this is reached only by a caller that sent null on purpose.",
    ),
    "Attaching a document of a type the module does not define": (
        "400",
        "REQUEST_ERROR naming document_type. Standard. The five types in FR-015 are a closed "
        "set.",
    ),
}

FR005_RULE10_OLD = "answers a per-field 422 rather than being coerced to OTHER."
FR005_RULE10_NEW = "answers 400 on that field rather than being coerced to OTHER."
FR015_RULE1_OLD = (
    "(1) The type must be one of the five in section 7.6 and anything else answers a per-field "
    "422."
)
FR015_RULE1_NEW = (
    "(1) The type must be one of the five in section 7.6 and anything else answers 400 on "
    "document_type."
)

LIFECYCLE_KEYS = (
    "school.students.transition, .transfer, .suspend or .reactivate, according to the operation"
)
FR019_KEYS_OLD = f"{LIFECYCLE_KEYS} to set it."
FR019_KEYS_NEW = "school.students.update to set it."
FR019_ACCEPTANCE_OLD = "(1) A caller without manage cannot set the policy and receives 403;"
FR019_ACCEPTANCE_NEW = (
    "(1) A caller without school.students.update cannot set the policy and receives 403;"
)

CAPACITY_BULLET_START = "Assigning into a full class is refused"
CAPACITY_BULLET = (
    "Assigning into a full class is refused 422 CLASS_AT_CAPACITY, and the override, "
    "allow_over_capacity=true, succeeds for any caller the route already admits, with no "
    "further key, and writes an audit event recording that capacity was exceeded (FR-002 "
    "rule 3, FR-006 rule 3)."
)

STUDENT_KEYS_EVIDENCE = (
    "school.students carries view, create, update, transition, transfer, suspend, reactivate, "
    "promote, import and export. None of them is assign or enrol: transfer records a child "
    "leaving for another school (FR-017), not a move between classes, and the assign verb was "
    "not attached to that resource although it was available. Placement was never modelled as "
    "an operation on a student."
)
STUDENT_KEYS_WHERE = "apps/core/management/commands/seed_school_permissions.py, the school.students rows"

KEY_COUNT_OLD = (
    "The school.students resource already exists with five keys, listed in section 3.7, and "
    "all five are reused unchanged. Three new keys are needed and no new resource is."
)
KEY_COUNT_NEW = (
    "The school.students resource carries ten keys. Section 3.7 lists the ones every read and "
    "every lifecycle move uses, reused unchanged; the three below, import, export and "
    "promote, are this module's additions, and no new resource was needed."
)

CORRECTIONS_SUMMARY = (
    " Five corrections against the code go with it. Section 11's standard validation refusals "
    "answer 400 REQUEST_ERROR (a Django validation error 400 VALIDATION_ERROR), not 422, and "
    "FR-005 and FR-015 say the same of a bad relationship or document type. Setting the "
    "admission-number policy needs school.students.update, in FR-019 and section 10. The "
    "capacity override in section 12.2 needs no key beyond the route's, as FR-002 and FR-006 "
    "say. Section 5.2 lists the students resource's keys as they are, without manage or "
    "view_sensitive. Section 8.1 counts ten keys on that resource."
)

# ── FR-010: a promotion keeps the capacity rule ─────────────────────────────

FR010_RULES_TAIL = (
    " (14) Capacity holds for a promotion as it does for enrolling one child (FR-002 rule 3, "
    "FR-006 rule 3). The preview returns over_capacity: one entry per target class the run "
    "would fill past its capacity, carrying class, class_name, capacity, used (the seats "
    "already taken in the target year), adding (the pupils this run places there) and "
    "over_by, sorted by class name. A pupil already placed in the target year is counted "
    "once, as a seat taken, and never again as an arrival; only PROMOTE and REPEAT place a "
    "pupil; a class with no capacity is unlimited and never listed. The run refuses 422 "
    "PROMOTION_OVER_CAPACITY, naming each class with its count against its capacity and "
    "carrying the same list in extra.classes, before a batch is created or any pupil moves, "
    "until it is sent allow_over_capacity=true. No further key is asked, as for one child. "
    "The school app's confirm step lists those classes and keeps Run promotion disabled until "
    "the registrar ticks \"Go ahead and put these classes over capacity\" (school-fe "
    "08a5cd5); nothing here claims deployment."
)
FR010_PREVIEW_TAIL = (
    " It also returns over_capacity (rule 14), worked out in two queries whatever the size of "
    "the cohort."
)
FR010_ACCEPTANCE_TAIL = (
    " (17) The preview names a class the run would overfill with its capacity, seats taken, "
    "arrivals and overflow, and a class with room, or with no capacity, is not listed. (18) A "
    "run that would overfill a class answers 422 PROMOTION_OVER_CAPACITY and writes no "
    "enrolment in the target year; the same run with allow_over_capacity=true promotes "
    "(tests/test_behaviour.py, PromotionTests: "
    "test_the_preview_names_a_class_the_run_would_overfill, "
    "test_a_run_that_overfills_a_class_waits_for_an_acknowledgement and "
    "test_a_class_with_room_or_no_capacity_is_not_listed)."
)
PROMOTION_REFUSAL = [
    "A promotion run that would fill a target class past its capacity, unacknowledged",
    "422",
    "PROMOTION_OVER_CAPACITY, extra.classes lists each class with its capacity, used, adding "
    "and over_by. Cleared by allow_over_capacity=true, as CLASS_AT_CAPACITY is for one child, "
    "and refused before a batch is created.",
]
PROMOTION_BULLET = (
    "A promotion that would fill a class past its capacity is refused 422 "
    "PROMOTION_OVER_CAPACITY and writes nothing until it is acknowledged, and the preview "
    "lists that class first (FR-010)."
)
PROMOTION_SUMMARY = (
    " A promotion keeps the capacity rule enrolling one child keeps: the preview lists every "
    "target class the run would overfill as over_capacity, and the run refuses 422 "
    "PROMOTION_OVER_CAPACITY until it is sent allow_over_capacity=true; FR-010's rules, "
    "preview and acceptance, sections 10, 11 and 12.2 follow."
)

SUMMARY = (
    "Minor revision. The medical key is retired and the switches decide: blood_group, "
    "allergies and conditions are sensitive Field Access fields, read and corrected only where "
    "a school turns a role's switches on, and a field a role may not write is refused 403 "
    "field_write_denied rather than 422, on enrolment as on an edit. Enrolment refuses a typed "
    "medical value without the switch and drops a blank one; the enrolment date, a new "
    "pupil's name, birth date, gender and number, and a new guardian's name and phone are "
    "open on create. The directory row enforces the enrolment date's Read switch, and "
    "school.guardians carries fields and no keys. Both imports read a guardian's name in "
    "parts, with the one-line column an optional fallback that is flagged for a check. "
    "Confirming an applicant keeps the admission-number rule, and the 2-to-25 age bound is the "
    "server's on enrolment and edit as well as on the import. Sections 3.7, 7.1, 7.2, 11, 12 "
    "and 14 and FR-001 to FR-004, FR-012, FR-024, FR-027 and FR-028 follow; five rows join "
    "section 2.2, and section 9 counts twenty-eight requirements." + PROMOTION_SUMMARY
    + CORRECTIONS_SUMMARY + " "
    + TEST_EVIDENCE
)

STALE = (
    "FieldSecurityMixin.write_permissions",
    "Gated by FieldSecurityMixin",
    "Deliberately NOT gated",
    "medical_notes is absent",
    "a per-field 422 naming it",
    "without school.students.view_sensitive",
    "as the spreadsheet imports still do",
    "all eleven columns",
    "a fifteenth column",
    "Twenty-five requirements",
    "view_sensitive carries no band",
    "already exists with five keys",
    "without manage cannot set",
    "anything else answers a per-field 422",
    "answers a per-field 422 rather than being coerced",
    "update, manage and view_sensitive",
)


def patch() -> None:
    require_newest(str(ROOT / "functional-requirements" / M11_DIR / f"{M11_STEM}_v*.docx"),
                   SOURCE)
    doc = Document(str(frd_path(M11_DIR, M11_STEM, SOURCE)))
    set_control(doc, "Version", TARGET)
    set_control(doc, "Date", REVIEW_DATE)
    set_control(doc, "Supersedes", f"v{SOURCE}")
    set_control(doc, "Verified against", CODE_BASELINE)
    set_control(doc, "Source MRD", SOURCE_MRD)

    # Section 2.
    edit_paragraph(doc, "The table below lists every claim", CORRECTION_TAIL_OLD,
                   CORRECTION_TAIL_NEW)
    corrections = table_with_header(doc, "What an earlier version said")
    append_to_row = row_starting(corrections, "Restrict medical_notes exposure")
    keep_format(append_to_row.cells[1], append_to_row.cells[1].text.rstrip() + SUPERSEDED_FLS)
    medical = row_starting(corrections, "v2.3.1 (section 7.1): medical_notes")
    keep_format(medical.cells[1], medical.cells[1].text.rstrip() + SUPERSEDED_MEDICAL)
    band = row_starting(corrections, "v2.8 (row R31 above")
    edit_cell(band.cells[1], BAND_OLD_ROW, BAND_NEW_ROW)
    append_rows(corrections, CORRECTIONS)

    # Section 3.7.
    keys = table_headed(doc, "Key", "Sensitivity", "Default holders", "Use in M11")
    medical_key = row(keys, "school.students.view_sensitive")
    for cell, text in zip(medical_key.cells, MEDICAL_KEY_ROW):
        keep_format(cell, text)

    # Section 7.1 and 7.2.
    student = table_with_row(doc, "blood_group")
    for label, text in STUDENT_FIELDS.items():
        set_row(student, label, text)
    guardian = table_with_row(doc, "Guardian.full_name")
    name_row = row(guardian, "Guardian.full_name")
    keep_format(name_row.cells[0], GUARDIAN_NAME_LABEL)
    keep_format(name_row.cells[-1], GUARDIAN_NAME)
    append_to(guardian, "Guardian.phone", GUARDIAN_PHONE_TAIL)

    # Section 9.
    edit_paragraph(doc, "Twenty-five requirements.", "Twenty-five requirements.",
                   "Twenty-eight requirements.")

    fr001 = fr_table(doc, "FR-001")
    set_row(fr001, "Field exposure", FR001_EXPOSURE)
    edit_value(fr001, "Acceptance",
               "(6) medical_notes is absent from the response for a caller without "
               "view_sensitive, and present for one with it.",
               "(6) blood_group, allergies and conditions are absent from the detail for a role "
               "whose Read switch is off, and present for one whose switch is on.")

    fr002 = fr_table(doc, "FR-002")
    append_to(fr002, "Business rules", FR002_RULES_TAIL)
    append_to(fr002, "Acceptance", FR002_ACCEPTANCE_TAIL)

    fr003 = fr_table(doc, "FR-003")
    edit_value(fr003, "Trigger", "POST /v1/students/<id>/confirm/.",
               "POST /v1/students/<id>/confirm/, with an optional student_number, reason and "
               "effective_date.")
    append_to(fr003, "Business rules", FR003_RULES_TAIL)
    append_to(fr003, "Acceptance", FR003_ACCEPTANCE_TAIL)

    fr004 = fr_table(doc, "FR-004")
    set_row(fr004, "Key", FR004_KEY)
    edit_value(fr004, "Business rules", FR004_RULE2_OLD, FR004_RULE2_NEW)
    append_to(fr004, "Business rules",
              " (6) A date of birth is held to the bound in FR-002 rule 12.")
    edit_value(fr004, "Acceptance", FR004_ACCEPTANCE_1_OLD, FR004_ACCEPTANCE_1_NEW)
    edit_value(fr004, "Acceptance", FR004_ACCEPTANCE_2_OLD, FR004_ACCEPTANCE_2_NEW)
    append_to(fr004, "Acceptance",
              " (6) A date of birth thirty years back answers 400 and the stored date is "
              "unchanged (test_editing_a_birth_date_keeps_the_same_bounds).")

    fr012 = fr_table(doc, "FR-012")
    edit_value(fr012, "Columns", FR012_COLUMNS_OLD, FR012_COLUMNS_NEW)
    edit_value(fr012, "Columns", "admission_date is added as a fifteenth column, optional, "
               "because", "admission_date is optional, because")
    append_to(fr012, "Columns",
              " The guardian's name is three columns and a fallback: Guardian First, Middle and "
              "Last Name come first, and Guardian Name, the name on one line, is optional "
              "(rule 11).")
    edit_value(fr012, "Business rules", FR012_RULE9_OLD, FR012_RULE9_NEW)
    append_to(fr012, "Business rules", FR012_RULES_TAIL)
    append_to(fr012, "Acceptance", FR012_ACCEPTANCE_TAIL)

    fr024 = fr_table(doc, "FR-024")
    edit_value(fr024, "Columns", FR024_COLUMNS_OLD, FR024_COLUMNS_NEW)
    append_to(fr024, "Business rules", FR024_RULES_TAIL)
    edit_value(fr024, "Acceptance", FR024_ACCEPTANCE_OLD, FR024_ACCEPTANCE_NEW)

    fr027 = fr_table(doc, "FR-027")
    edit_value(fr027, "Business rules", FR027_RULE1_TAIL_OLD, FR027_RULE1_TAIL_NEW)
    edit_value(fr027, "Business rules", FR027_RULE3_OLD, FR027_RULE3_NEW)
    append_to(fr027, "Business rules", FR027_RULES_TAIL)
    edit_value(fr027, "Acceptance", FR027_ACCEPTANCE_4_OLD, FR027_ACCEPTANCE_4_NEW)

    fr028 = fr_table(doc, "FR-028")
    edit_value(fr028, "Business rules", FR028_RULE2_OLD, FR028_RULE2_NEW)
    append_to(fr028, "Trigger", FR028_TRIGGER_TAIL)

    # Section 5.2.
    evidence = row_starting(table_headed(doc, "The evidence", "What it says", "Where"),
                            "The students resource has no key")
    keep_format(evidence.cells[1], STUDENT_KEYS_EVIDENCE)
    keep_format(evidence.cells[2], STUDENT_KEYS_WHERE)

    # Section 8.1.
    edit_paragraph(doc, "The school.students resource already exists", KEY_COUNT_OLD,
                   KEY_COUNT_NEW)

    # FR-005, FR-015 and FR-019.
    edit_value(fr_table(doc, "FR-005"), "Business rules", FR005_RULE10_OLD, FR005_RULE10_NEW)
    edit_value(fr_table(doc, "FR-015"), "Business rules", FR015_RULE1_OLD, FR015_RULE1_NEW)
    fr019 = fr_table(doc, "FR-019")
    edit_value(fr019, "Keys", FR019_KEYS_OLD, FR019_KEYS_NEW)
    edit_value(fr019, "Acceptance", FR019_ACCEPTANCE_OLD, FR019_ACCEPTANCE_NEW)

    # FR-010.
    fr010 = fr_table(doc, "FR-010")
    append_to(fr010, "Business rules", FR010_RULES_TAIL)
    append_to(fr010, "Preview", FR010_PREVIEW_TAIL)
    append_to(fr010, "Acceptance", FR010_ACCEPTANCE_TAIL)

    # Section 10.
    routes = table_headed(doc, "Method", "Endpoint", "Key", "Requirement")
    for endpoint, old, new in (
        ("/v1/students/promotions/preview/", "FR-010. Classification, writes nothing.",
         "FR-010. Classification, writes nothing. Returns over_capacity."),
        ("/v1/students/promotions/", "FR-010. Run.",
         "FR-010. Run. Takes allow_over_capacity; without it, a class the preview listed "
         "answers 422 PROMOTION_OVER_CAPACITY."),
    ):
        hits = [r for r in routes.rows if r.cells[0].text.strip() == "POST"
                and r.cells[1].text.strip() == endpoint]
        if len(hits) != 1 or hits[0].cells[3].text.strip() != old:
            raise ValueError(f"The {endpoint} row is not what it was expected to be")
        keep_format(hits[0].cells[3], new)
    put = [r for r in routes.rows if r.cells[0].text.strip() == "PUT"
           and r.cells[1].text.strip() == "/v1/students/admission-number-policy/"]
    if len(put) != 1 or put[0].cells[2].text.strip() != LIFECYCLE_KEYS:
        raise ValueError("The policy PUT row is not where, or what, it was expected to be")
    keep_format(put[0].cells[2], "school.students.update")

    # Section 12.2.
    edit_paragraph(doc, CAPACITY_BULLET_START, paragraph_text(doc, CAPACITY_BULLET_START),
                   CAPACITY_BULLET)
    insert_after(doc, CAPACITY_BULLET_START, PROMOTION_BULLET)

    # Section 11.
    refusals = table_headed(doc, "Condition", "Status", "Code")
    for start, (status, code) in VALIDATION_ROWS.items():
        target = row_starting(refusals, start)
        keep_format(target.cells[1], status)
        keep_format(target.cells[2], code)
    edit_cell(row_starting(refusals, "A student number omitted where").cells[0],
              "A student number omitted where the school's policy requires one",
              "A student number omitted where the school's policy requires one, at enrolment "
              "or when confirming an applicant whose record carries none")
    append_rows(refusals, REFUSALS + [PROMOTION_REFUSAL])

    # Section 12.
    edit_paragraph(doc, "blood_group, allergies and conditions are absent from every list",
                   LIST_READ_OLD, LIST_READ_NEW)
    edit_paragraph(doc, PATCH_REFUSAL_START,
                   paragraph_text(doc, PATCH_REFUSAL_START), PATCH_REFUSAL_NEW)
    insert_after(doc, POLICY_ANCHOR, CONFIRM_BULLET)
    insert_after(doc, CONFIRM_BULLET[:60], AGE_BULLET)

    # Section 14, decision 10.
    decision = row_starting(table_headed(doc, "Question", "Why it cannot be answered here",
                                         "Cost of deferring"), "10. Is Student Management")
    edit_cell(decision.cells[1], BAND_OLD_DECISION, BAND_NEW_DECISION)

    log_change(doc, TARGET, SUMMARY)
    assert_absent_outside_log(doc, *STALE)
    report_remaining(doc, "view_sensitive")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M11_DIR, M11_STEM, TARGET),
           f"{M11_STEM.replace('_', ' ')} v{TARGET}", TARGET)


def paragraph_text(doc, start: str) -> str:
    hits = [p.text for p in doc.paragraphs if p.text.strip().startswith(start)]
    if len(hits) != 1:
        raise ValueError(f"{start!r} starts {len(hits)} paragraphs")
    return hits[0]


def report_remaining(doc, needle: str) -> None:
    """Print each surviving mention outside the log, so a reviewer sees they are historical."""
    for line in all_text_outside_log(doc).split("\n"):
        if needle in line:
            print(f"  kept ({needle}): {line[:110]}...")


if __name__ == "__main__":
    patch()
