#!/usr/bin/env python3
"""Cut MRD v2.89, M04 FRD v1.28, M07 FRD v1.16 and M12 FRD v2.11: a restricted role grant waits for approval.

What changed in the backend, and therefore in the documents:

* Giving somebody a tenant role whose restricted permissions the granter does
  not hold was refused with 403 and the full list of missing keys. In a school
  with one administrator that could never be satisfied: nobody holds the
  finance keys until a Finance Admin exists, so nobody could seat the first
  one, and only the Vision super admin could.
* ``vs_rbac.services.grant_role`` is the one place a school role is granted. A
  granter holding every restricted key the role carries grants directly;
  anybody else raises a ``TenantRoleGrantRequest`` (vs_rbac 0029) submitted to
  the new ``rbac.role_grant`` ladder (vs_workflow 0018), shaped like the
  ``rbac.role_change`` ladder, with continue-without-approval refused. The
  requester may decide it only where nobody else is eligible
  (``self_approval_only_when_alone``), so a school with one administrator
  closes its own request and a school with two gets the other's decision; the
  role-change ladder keeps its wider exemption. The assignment create and replace
  endpoints answer 202 with the waiting request, and the replaced grant stays
  in force until approval.
* The staff bulk grant applied no restricted rule at all, so an administrator
  could bulk-grant themselves Finance Admin. It now calls ``grant_role`` for
  each person and names the waiting people under ``pending_approval``.
* On approval the grant is written naming the requester as granter, and its
  ROLE_ASSIGNED audit row carries the request, requester, approver and a
  self_approved flag. An archived role fails to apply; a grant already held is
  not written twice; a reversal is refused once the grant is written.
* The staff roles read lists waiting requests under ``pending``.
* Creating a person with such a role, and editing the role on an existing
  grant to one, are still refused, now in one sentence naming the role. That
  is recorded under Needs Attention in M04 and in M12's refusal list, because
  M12 carries no Needs Attention section.

Two stale labels are corrected on the way: M04's contents named FR-001 to
FR-026 while FR-027 existed, and M07's traceability paragraph named 27 entries
in MRD v2.80 while the table and the MRD carry 28.

    python tools/patch_restricted_grant_ladder_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_record_history_docs import (
    ROOT,
    add_fr,
    append_rows,
    change_log_table,
    finish,
    frd_path,
    keep_format,
    log_change,
    replace_cell,
    set_control,
    set_cover_version,
    update_reconciliation,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    insert_after,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
    row_with_cell_starting,
    set_value,
    table_with_header,
)

import patch_record_history_docs

REVIEW_DATE, SHORT_DATE = "26 September 2026", "26 Sep 2026"
MRD_SOURCE, MRD_TARGET = "2.88", "2.89"
CODE_BASELINE = (
    "Backend main at ac42a93d with the restricted role grant ladder in its working tree, "
    "26 September 2026. The school and console integrations are implemented and pending "
    "release"
)
TEST_EVIDENCE = (
    "Verified by the full vs_rbac (890 tests), vs_workflow (420), schools.vs_staff (317) and "
    "vs_user (421) suites, slow tests included, each app run on its own and all passing. The "
    "school and console integrations are implemented and pending release; nothing here claims "
    "deployment."
)

# log_change writes the module's own review date, read from this module.
patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


# ── editing helpers ──────────────────────────────────────────────────────────


def row_labelled(doc, label: str):
    """The one row in the document whose first cell reads ``label``."""
    hits = [r for t in doc.tables for r in t.rows if r.cells[0].text.strip() == label]
    if len(hits) != 1:
        raise ValueError(f"{label!r} labels {len(hits)} rows")
    return hits[0]


def table_headed(doc, *headers: str):
    """The one table whose header row reads exactly ``headers``."""
    hits = [t for t in doc.tables
            if tuple(c.text.strip() for c in t.rows[0].cells) == headers]
    if len(hits) != 1:
        raise ValueError(f"{headers!r} heads {len(hits)} tables")
    return hits[0]


def insert_row_after(anchor, values: list[str]):
    """Clone ``anchor`` directly below itself and write ``values`` into the copy."""
    clone = copy.deepcopy(anchor._tr)
    anchor._tr.addnext(clone)
    table = anchor._parent._parent
    new = next(r for r in table.rows if r._tr is clone)
    cells = list({id(c._tc): c for c in new.cells}.values())
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        keep_format(cell, value)
    return new


def edit_box(cell, edits) -> None:
    """Edit a heading-and-bullets box line by line, whichever shape it has.

    Edits are ("sub", start, old, new) inside the one line starting with
    ``start``, ("after", start, new) to add a line after it, ("replace", start,
    new) and ("append", new). A one-run box keeps its lines as breaks in that
    run; a paragraph-per-line box reuses its paragraphs positionally.
    """
    lines = [line.strip() for p in cell.paragraphs for line in p.text.split("\n") if line.strip()]

    def at(start):
        hits = [i for i, line in enumerate(lines) if line.startswith(start)]
        if len(hits) != 1:
            raise ValueError(f"{start!r} starts {len(hits)} box lines")
        return hits[0]

    for edit in edits:
        kind = edit[0]
        if kind == "sub":
            _, start, old, new = edit
            i = at(start)
            if lines[i].count(old) != 1:
                raise ValueError(f"{old[:60]!r} occurs {lines[i].count(old)} times in the line")
            lines[i] = lines[i].replace(old, new)
        elif kind == "after":
            lines.insert(at(edit[1]) + 1, edit[2])
        elif kind == "replace":
            lines[at(edit[1])] = edit[2]
        elif kind == "append":
            lines.append(edit[1])
        else:
            raise ValueError(f"Unknown box edit {kind!r}")
    if len(cell.paragraphs) == 1 and len(cell.paragraphs[0].runs) == 1:
        cell.paragraphs[0].runs[0].text = "\n".join(lines)
    else:
        mrd_tools.set_lines(cell, lines)


def append_to(table, label: str, tail: str) -> None:
    cell = row(table, label).cells[-1]
    keep_format(cell, cell.text.rstrip() + tail)


def retitle_paragraph(doc, start: str, text: str) -> None:
    mrd_tools.retitle(paragraph_starting(doc, start), text)


def all_text_outside_log(doc) -> str:
    log = change_log_table(doc)._tbl
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        if table._tbl is log:
            continue
        for r in table.rows:
            texts.extend(c.text for c in r.cells)
    return "\n".join(texts)


def assert_absent_outside_log(doc, *needles: str) -> None:
    """Refuse a stale claim anywhere but the change log, which records history."""
    blob = all_text_outside_log(doc)
    for needle in needles:
        hit = blob.find(needle)
        if hit >= 0:
            context = blob[max(0, hit - 160):hit + 80].replace("\n", " / ")
            raise ValueError(f"{needle!r} is still in the document: ...{context}...")


# ── M04 Roles & Permissions (RBAC) ───────────────────────────────────────────

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.27", "1.28"

M04_FR028_HEADING = "FR-028  Send a Restricted Role Grant to Approval Rather Than Refusing It"
M04_FR028 = [
    ("Requirement",
     "Giving somebody a role must never be refused only because nobody in the school yet "
     "holds its restricted permissions. A granter who holds every restricted key the role "
     "carries grants it at once; anybody else's grant, whoever it is for and the granter "
     "included, waits for an approval ladder that records who decided it, and confers nothing "
     "until then. Every entry point that grants a school role keeps the same rule, so a bulk "
     "grant is not a way round it."),
    ("Current evidence",
     "vs_rbac.services.grant_role is the one place a tenant role is granted: the assignment "
     "create view, the replace endpoint and the school staff app's bulk grant (grant_to_many) "
     "all call it after refusing what no approval can make right, a role or branch from "
     "another tenant, a duplicate grant, or a reach outside a branch-bound caller's branches. "
     "grant_needs_approval asks missing_restricted_grant_authority; where nothing is missing "
     "the grant is written at once with assigned_by the granter, and otherwise "
     "raise_role_grant_request creates a TenantRoleGrantRequest and submits it to the "
     "rbac.role_grant ladder in one transaction, so a request the engine cannot route leaves "
     "no row behind. A second request for the same person, role and reach while one waits is "
     "refused under role. RoleGrantWorkflowHandler resolves the tenant-less role-grant "
     "template for a school, one stage, Role administrator approval, naming the school_admin "
     "system role with advance rule ANY and no skipping, or role-grant-platform, owned by the "
     "platform tenant and naming xvs_platform_admin; vs_workflow migration 0018 seeds both. "
     "Like the role-change handler it declares allows_requester_self_approval and refuses "
     "continue-without-approval; unlike it, it also declares self_approval_only_when_alone, so "
     "the engine keeps the requester on the stage's approver list only where nobody else is "
     "eligible. A school with one administrator therefore closes its own request, and a "
     "school with a second administrator gets that person's decision while the requester's "
     "own vote is refused. Its decision layout names the person, the role, the reach, "
     "the grant it replaces and the reason, and lists only the role's restricted permissions. "
     "on_approved runs apply_role_grant_request, the only place a request's grant is written: "
     "it locks the request, writes the grant with assigned_by the requester, withdraws a "
     "replaced grant in the same transaction, links the grant to the request and records the "
     "reviewer. A role archived while the request waited marks it APPLY_FAILED with the "
     "reason, and a person who already holds the role at that reach by then leaves it "
     "APPROVED with nothing written twice. Rejection, withdrawal and cancellation close it as "
     "DENIED. A reversal is refused once the grant is written, naming withdrawing the role "
     "from the person's profile as the remedy."),
    ("Acceptance",
     "A school's only administrator can raise Finance Admin for herself, approve it and hold "
     "it, and the grant's audit row says she approved her own. Where the school has a second "
     "administrator, that person is the only one on the approver list and the requester's own "
     "vote is refused, and the second administrator's approval is not self-approval. Granting yourself a restricted role raises a ladder and "
     "grants nothing until it completes; a role with no restricted key, or a granter holding "
     "every one, is granted on the spot; rejecting grants nothing; a role archived while "
     "waiting is not granted; a grant already held by approval time is not written twice; an "
     "applied grant cannot be reversed through the engine; replacing a grant keeps the old one "
     "until approval; and a bulk grant raises one request per person and reports somebody "
     "already waiting rather than asking twice. RoleGrantLadderTests pins each case in "
     "fourteen tests, and TenantUserRoleAssignmentViewTests pins the 202 answers on assign "
     "and replace and the second request refused under role."),
    ("Limit",
     "The requester may approve their own grant where nobody else is eligible, which is "
     "every school with one administrator; the approval is recorded as self-approved, visible "
     "afterwards rather than reviewed beforehand. Creating a person with such a role, through account creation, "
     "draft submission or the school staff create, and editing the role on an existing grant "
     "to one, have no ladder behind them and are still refused, in one sentence naming the "
     "role and saying to grant it from the person's profile (Needs Attention). The approver "
     "list is frozen when the stage activates, so an administrator appointed while a request "
     "waits does not take it from a requester who was alone. A grant request is decided in the workflow "
     "engine's approval screens; unlike a role change it has no decide route of its own "
     "under /rbac/."),
]

M04_SELF_APPROVAL = (
    "A restricted addition to any role needs only the approval of whoever raised it. The "
    "role-change ladder keeps the requester on its one stage and completes on one approval in "
    "every school, so a head teacher can raise a restricted addition to School Admin and "
    "approve it themselves even where a second administrator sits on the same stage. The "
    "approval is recorded as self_approved_change_request, which makes it visible afterwards, "
    "not reviewed beforehand. A restricted role grant is narrower: its requester is eligible "
    "only where nobody else is (FR-028)."
)
M04_SELF_APPROVAL_FIX = (
    "Exclude the requester whenever the stage resolves another eligible approver, keeping the "
    "exemption for the school with one administrator, or record here that one-person approval "
    "is the accepted rule. The role-grant handler already does this through "
    "self_approval_only_when_alone; declaring it on the role-change handler would close this "
    "item."
)
M04_CREATION_GAP = (
    "Creating a person with a restricted role is refused rather than routed for approval. "
    "Account creation, draft submission and the school staff create answer 403 where the "
    "role carries restricted permissions the creator does not hold, and editing the role on "
    "an existing grant to such a role is refused the same way, because none of them has a "
    "ladder behind it. The sentence names the role and the way through: add the person "
    "without it, then grant it from their profile, where it waits for approval. A school "
    "adding its bursar makes two moves where one would do."
)
M04_CREATION_FIX = (
    "Route a restricted role chosen at creation, or on a grant edit, through grant_role once "
    "the account exists, so creation writes the person and raises the request in one "
    "transaction; or record here that two steps is the accepted rule."
)

M04_RECONCILIATION = [
    ("replace", "• MRD v2.88 lists Module 4",
     f"• MRD v{MRD_TARGET} lists Module 4 as Roles & Permissions (RBAC), Phase V1, Backend "
     "Complete, In use Complete, code vs_rbac, with twenty-three capability entries. The "
     "module number, name, phase, states and ownership agree with this revision."),
    ("append", "• A restricted role grant decided by an approval ladder extends the existing "
     "assignment and role-change request entries rather than adding one, so the count stays "
     "at twenty-three."),
]

M04_SUMMARY = (
    "Minor revision. A role carrying restricted permissions its granter does not hold is no "
    "longer refused: grant_role, now the one place a school role is granted, writes it at "
    "once for a granter holding every restricted key and otherwise raises a "
    "TenantRoleGrantRequest decided by the new rbac.role_grant ladder, one School Admin stage "
    "that the requester may decide only where nobody else is eligible, so a school with a "
    "second administrator gets that person's decision. Assign and replace answer 202 with the waiting request, a "
    "replaced grant stays in force until approval, and a second request for the same grant is "
    "refused under role. The staff bulk grant, which applied no restricted rule at all and let "
    "an administrator bulk-grant themselves Finance Admin, goes through the same function. "
    "The grant is written on approval naming the requester, and its ROLE_ASSIGNED audit row "
    "records the request, the approver and whether it was self-approved. Account creation, "
    "draft submission and an edit of a grant's role still refuse such a role, in one sentence "
    "naming it. FR-028 is added; FR-003, FR-017, FR-019 and FR-024, the scope, actors, "
    "withdrawal routes, data model, contracts, dependencies, verification, Needs Attention "
    "(one row added, and the self-approval row confined to role changes) and traceability "
    "follow. The contents line now names "
    f"FR-001 to FR-028. MRD v{MRD_TARGET}. " + TEST_EVIDENCE
)


def patch_m04() -> None:
    doc = Document(str(frd_path(M04_DIR, M04_STEM, M04_SOURCE)))
    set_cover_version(doc, M04_SOURCE, M04_TARGET)
    set_control(doc, "Version", M04_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Source scope", CODE_BASELINE)
    set_control(doc, "Code inspected", CODE_BASELINE)
    set_control(doc, "MRD baseline", f"XVS Module Requirements Document v{MRD_TARGET}")

    keep_format(row_labelled(doc, "4. Functional Requirements").cells[1],
                "FR-001 to FR-028, each with the code evidence behind it")
    keep_format(row_labelled(doc, "10. MRD Traceability").cells[1],
                f"Agreement with MRD v{MRD_TARGET}'s twenty-three capability entries")

    edit_cell(row_labelled(doc, "Assignments").cells[-1],
              "Assign, replace, account creation, and draft submission all refuse a restricted "
              "role unless the actor already holds every restricted key in it, whoever the role "
              "is being given to, the actor included.",
              "Every new grant, single, replacing or bulk, goes through grant_role: a granter "
              "holding every restricted key the role carries grants it at once, and anybody "
              "else's grant waits for the role-grant ladder, whoever it is for, the granter "
              "included. Account creation, draft submission and an edit of an existing grant's "
              "role still refuse a restricted role the actor does not hold.")
    insert_row_after(row_labelled(doc, "Role-change approval"), [
        "Role-grant approval",
        "TenantRoleGrantRequest, raised by grant_role for a role carrying restricted keys the "
        "granter does not hold and submitted to the rbac.role_grant workflow ladder in the same "
        "transaction. The handler this module registers writes the grant when the ladder "
        "approves, naming the requester as granter and recording who approved it. The "
        "requester may decide it only where nobody else is eligible.",
    ])

    edit_paragraph(doc, "Five boundaries hold this module up",
                   "and cannot be handed to anybody by an assigner who does not already hold "
                   "it.",
                   "and cannot be handed to anybody by an assigner who does not already hold it "
                   "without a role grant decided by the role-grant workflow ladder.")
    edit_cell(row_labelled(doc, "School administrator").cells[1],
              "assign and revoke roles within their own grant ceiling, and decide role changes",
              "assign and revoke roles, where a role carrying restricted permissions they do "
              "not hold goes for approval rather than being granted, and decide role changes "
              "and restricted role grants")
    edit_cell(row_labelled(doc, "School administrator").cells[1],
              "the ladder names, their own included.",
              "the ladder names, their own included, though a restricted role grant of their "
              "own only where nobody else is eligible to decide it.")

    fr003 = fr_table(doc, "FR-003")
    edit_value(fr003, "Current evidence",
               "and assignment, replace, account creation, and draft submission refuse a role "
               "whose restricted keys the assigner does not already hold "
               "(missing_restricted_grant_authority), whoever receives it.",
               "and every new grant, single, replacing or bulk, goes through grant_role, which "
               "writes it directly only where the granter already holds every restricted key "
               "the role carries (missing_restricted_grant_authority) and otherwise raises a "
               "TenantRoleGrantRequest for the rbac.role_grant ladder, whoever receives it "
               "(FR-028). Account creation, draft submission and an edit of the role on an "
               "existing grant have no ladder behind them and still refuse such a role, in one "
               "sentence naming the role rather than a list of keys.")
    edit_value(fr003, "Acceptance",
               "Giving yourself a role that already carries a restricted key is closed by the "
               "assignment ceiling.",
               "Giving yourself, or anybody, a role that already carries a restricted key you "
               "do not hold waits for the role-grant ladder rather than taking effect.")
    edit_value(fr003, "Acceptance",
               "test_giving_yourself_that_role_is_still_refused_by_the_grant_ceiling and "
               "test_giving_it_to_somebody_else_is_refused_by_the_same_ceiling",
               "test_giving_yourself_that_role_waits_for_approval and "
               "test_giving_it_to_somebody_else_waits_the_same_way")
    append_to(fr003, "Limit", " A restricted role grant likewise takes effect only once "
              "approved, and its requester may approve it only where nobody else is eligible "
              "(FR-028).")

    fr017 = fr_table(doc, "FR-017")
    edit_value(fr017, "Acceptance", "The exemption applies to this document type alone",
               "The exemption applies to this document type and, narrowed to a requester "
               "nobody else can relieve, to restricted role grants (FR-028), and to no other "
               "type")

    append_to(fr_table(doc, "FR-019"), "Current evidence",
              " A grant written by an approved role grant request carries grant_request_id, "
              "requested_by_id, approved_by_id and self_approved in its ROLE_ASSIGNED audit "
              "row, so who approved their own grant is a filter rather than a reconstruction.")
    edit_value(fr_table(doc, "FR-024"), "Current evidence",
               "so neither can drift from the schema or from the other.",
               "so neither can drift from the schema or from the other. Approving a role grant "
               "request asks it again at the moment of writing, so a grant that became a "
               "duplicate while it waited is approved without a second row.")

    add_fr(doc, M04_FR028_HEADING, "FR-028 | Implemented with limits", M04_FR028)

    edit_cell(row_labelled(doc, "Replace an assignment").cells[1],
              "The old role, atomically, as the new one is issued.",
              "The old role, atomically, as the new one is issued. Where the new role waits for "
              "approval, the old one stays in force until the approval writes the new one.")

    insert_row_after(row_labelled(doc, "TenantRoleChangeRequest / TenantRoleChangeDeltaItem"), [
        "TenantRoleGrantRequest",
        "A role grant waiting on approval: the person, the role, its reach, the grant it "
        "replaces, the reason, and who asked.",
        "PENDING, APPROVED, DENIED and APPLY_FAILED. Every reference sits inside the request's "
        "tenant. Routed as workflow document type rbac.role_grant. assignment links, one to "
        "one, the grant the approval wrote, and is null while pending and when the request "
        "closed without one. Indexed on tenant, status and submission time, and on tenant, "
        "person and status. Migration 0029.",
    ])

    edit_cell(row_labelled(doc, "GET, POST /rbac/tenants/{slug}/role-assignments/").cells[-1],
              "A restricted role is refused unless the actor already holds all its restricted "
              "keys; a second grant at another branch remains valid.",
              "A role carrying restricted keys the actor does not hold is not granted: the "
              "answer is 202 with the waiting request (id, status, person, role, reach, the "
              "grant it replaces, reason, requester and submission time) and a message that "
              "the role takes effect once approved; otherwise 201 with the grant. A second "
              "request for the same person, role and reach while one waits is 400 under role. "
              "A second grant at another branch remains valid.")
    edit_cell(row_labelled(doc, "POST /rbac/tenants/{slug}/role-assignments/{id}/replace/")
              .cells[-1],
              "Swap one role for another atomically. The assign keys.",
              "Swap one role for another atomically. The assign keys. A target role carrying "
              "restricted keys the caller does not hold answers 202 with the waiting request, "
              "and the grant being replaced stays in force until the approval withdraws it in "
              "the same transaction that writes the new one.")

    restricted = row_labelled(
        doc, "A restricted role assigned by somebody who lacks one of its restricted keys")
    keep_format(restricted.cells[0],
                "A restricted role granted by somebody who lacks one of its restricted keys")
    keep_format(restricted.cells[-1],
                "On assign, replace and bulk grant: not refused. 202 with the waiting request "
                "(a bulk grant names the people under pending_approval), whoever the recipient "
                "is, the granter included, and nothing is granted until the role-grant ladder "
                "approves. On account creation, draft submission and an edit of an existing "
                "grant's role, which have no ladder: 403 with one sentence naming the role and "
                "saying to grant it from the person's profile instead.")
    added = insert_row_after(restricted, [
        "A second grant request for the same person, role and reach while one waits",
        "400 under role, saying the role is already waiting for approval for this person.",
    ])
    added = insert_row_after(added, [
        "A requester deciding their own role grant request",
        "Accepted only where nobody else is eligible on the stage, and the grant's "
        "ROLE_ASSIGNED audit row records self_approved. Where a second administrator is "
        "eligible, the requester is not on the approver list and the engine refuses the vote "
        "under its own code.",
    ])
    insert_row_after(added, [
        "Reversing the approval of a role grant already written",
        "Refused by the engine with the handler's sentence: withdraw the role from the "
        "person's profile instead.",
    ])
    insert_after(doc, "Role-change approval detail contract.",
                 "Role-grant approval detail contract. The summary names the role and the "
                 "person, with who raised it and the reach. The frozen details carry the "
                 "person, role, reach, the grant it replaces and the reason where present, and "
                 "a change list of the role's restricted permissions only, each marked "
                 "restricted.")

    append_to(table_with_header(doc, "Dependency"), "Module 3, Identity and Team",
              " They refuse a role beyond it rather than routing it for approval, naming the "
              "role in one sentence.")
    module7 = row_labelled(doc, "Module 7, Workflow and Approval")
    edit_cell(module7.cells[-1],
              "The engine's requester self-approval exemption is declared for this document "
              "type alone.",
              "A grant of a role carrying restricted keys its granter does not hold is routed "
              "as document type rbac.role_grant through the role-grant and role-grant-platform "
              "templates (vs_workflow migration 0018), and that handler writes the grant on "
              "approval. The engine's requester self-approval exemption is declared for these "
              "two document types alone, and the role-grant handler narrows it with "
              "self_approval_only_when_alone to a requester nobody else can relieve.")
    insert_row_after(module7, [
        "Module 12, Staff Management",
        "Its bulk grant calls grant_role for each person, so a restricted role waits for "
        "approval one request per person rather than being written, and the staff roles read "
        "lists the waiting requests under pending, apart from the roles held.",
    ])

    retitle_paragraph(doc, "•  Both escalations end to end", (
        "•  Both escalations end to end: scope refuses a platform-only key on role and override "
        "paths; restriction refuses payments.payout.create in a group and as an ALLOW override, "
        "holds a restricted addition back for approval on every role save, whether the role is "
        "the editor's own, somebody else's or new, and giving a role that carries it, by "
        "anybody who does not hold it, waits for the role-grant ladder rather than taking "
        "effect."))
    retitle_paragraph(doc, "•  Bootstrap and ceiling", (
        "•  Bootstrap and ceiling: a platform actor may still grant a platform-scoped key "
        "inside CodeX; the Vision super administrator, and any granter holding every "
        "restricted key a role carries, grant it directly; anybody else's grant, whoever "
        "receives it, raises a role grant request, so a school's only administrator can seat "
        "its first Finance Admin through the ladder."))
    edit_paragraph(doc, "•  Self-service on the two controls",
                   "a role change may be decided by its requester and is then recorded under "
                   "self_approved_change_request",
                   "a role change may be decided by its requester, and a restricted role grant "
                   "only where nobody else is eligible, and each is then recorded as "
                   "self-approved")
    insert_after(doc, "•  The role-change ladder:",
                 "•  The role-grant ladder: nothing granted before approval; the grant written "
                 "on approval naming the requester and recording the approver and "
                 "self-approval; the requester deciding only when she is the only eligible "
                 "approver, and refused when a second administrator is; rejection granting nothing; an archived role failing to apply; "
                 "a grant already held not written twice; reversal refused once written; a "
                 "replaced grant kept until approval; a second request for the same grant "
                 "refused; and a bulk grant raising one request per person.")

    attention = table_with_header(doc, "Pri.")
    self_row = row_with_cell_starting(attention, 1, "A restricted addition to any role needs")
    keep_format(self_row.cells[1], M04_SELF_APPROVAL)
    keep_format(self_row.cells[2], M04_SELF_APPROVAL_FIX)
    keep_format(self_row.cells[3], "FR-017")
    insert_row_after(self_row, ["P2", M04_CREATION_GAP, M04_CREATION_FIX, "FR-003, FR-028"])
    edit_box(table_with_header(doc, "REMOVAL RULE").rows[0].cells[0], [
        ("replace", "• Against v1.19",
         "• Against v1.27 no item leaves. One is added: creating a person with a restricted "
         "role is refused rather than routed for approval. The self-approval item stays "
         "confined to role changes and now says so: a restricted role grant lets its requester "
         "decide only when nobody else can."),
    ])

    trace = next(t for t in doc.tables
                 if t.rows[0].cells[0].text.strip().startswith("MRD Module"))
    assign = row(trace, "Assign, revoke, and replace user roles")
    keep_format(assign.cells[1], "FR-009, FR-016, FR-028")
    keep_format(assign.cells[2], (
        "Implemented. A grant of a role carrying restricted keys its granter does not hold, "
        "single, replacing or bulk, waits for the role-grant ladder whoever receives it, and "
        "is written on approval naming the requester; a replaced grant stays until then. "
        "Account creation and an edit of a grant's role still refuse such a role (P2). A null "
        "assignment inherits all selected role branches through one grant."))
    append_to(trace, "Role-change request workflow",
              " A restricted role grant is decided by a ladder of the same shape (FR-028).")
    edit_paragraph(doc, "MRD v2.88 records Module 4", "MRD v2.88", f"MRD v{MRD_TARGET}")
    update_reconciliation(doc, M04_RECONCILIATION)

    log_change(doc, M04_TARGET, M04_SUMMARY)
    assert_absent_outside_log(
        doc, "MRD v2.88", "FR-001 to FR-026", "within their own grant ceiling",
        "may assign the first restricted holder", "assignment ceiling still refuses",
        "refused by the same ceiling", "is_still_refused_by_the_grant_ceiling",
        "Against v1.19", "declared for this document type alone",
        "A restricted role is refused unless the actor",
    )
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M04_DIR, M04_STEM, M04_TARGET),
           f"{M04_STEM.replace('_', ' ')} v{M04_TARGET}", M04_TARGET)


# ── M07 Workflow & Approval Engine ───────────────────────────────────────────

M07_DIR = "07-workflow-and-approval-engine"
M07_STEM = "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document"
M07_SOURCE, M07_TARGET = "1.15", "1.16"

M07_SUMMARY = (
    "Minor revision. Registers a fifteenth document type, rbac.role_grant: a grant of a role "
    "carrying restricted permissions its granter does not hold is a document decided by a "
    "ladder shaped like the role-change one, a tenant-less role-grant template naming "
    "school_admin for every school and role-grant-platform owned by the platform tenant, both "
    "seeded by vs_workflow migration 0018, one stage, ANY, never skipped. Its handler lets the "
    "requester decide only where nobody else is eligible, through the new handler attribute "
    "self_approval_only_when_alone (False on the base handler, not declared by role changes), "
    "refuses continue-without-approval, writes the grant on approval, closes "
    "the request on rejection, withdrawal or cancellation, and refuses a reversal once the "
    "grant is written. Self-approval is now declared by two document types rather than one, "
    "which the actors, ownership rules, FR-009, FR-012, FR-019, FR-021, FR-033, the "
    "resolution order, the Module 4 dependency and the further gaps now say. The "
    "traceability paragraph named 27 entries in MRD v2.80; it names the 28 the table and the "
    f"MRD carry, in MRD v{MRD_TARGET}, and the FR-033 heading is kept with its table. "
    + TEST_EVIDENCE
)


def patch_m07() -> None:
    doc = Document(str(frd_path(M07_DIR, M07_STEM, M07_SOURCE)))
    set_cover_version(doc, M07_SOURCE, M07_TARGET)
    set_control(doc, "Version", M07_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")
    set_control(doc, "Supporting apps", (
        "vs_rbac, vs_user, vs_tenants, vs_notifications, vs_audit; handlers registered by "
        "vs_finance, vs_procurement, vs_payments, vs_rbac (role changes and restricted role "
        "grants), vs_user (user creation) and the school staff app (leave)"))

    edit_cell(row_labelled(doc, "Requester").cells[-1],
              "except a role change, whose handler permits it (FR-012).",
              "except a role change, or a restricted role grant where nobody else is "
              "eligible, whose handlers permit it (FR-012).")
    edit_cell(row_labelled(doc, "The requester is never an approver, with one declared "
                                "exception").cells[-1],
              "role changes are the one type that does.",
              "role changes and restricted role grants are the two types that do. A handler "
              "may also narrow its exemption to a requester nobody else can relieve "
              "(self_approval_only_when_alone), which role grants do.")
    keep_format(row_labelled(doc, "The requester is never an approver, with one declared "
                                  "exception").cells[0],
                "The requester is never an approver, with two declared exceptions")

    append_to(fr_table(doc, "FR-009"), "Acceptance",
              " Restricted role grants run on ladders of the same shape: the tenant-less "
              "role-grant template names school_admin for every school, and "
              "role-grant-platform, owned by the platform tenant, names xvs_platform_admin. A "
              "grant's requester is on its stage only where nobody else holds that role "
              "(FR-012).")

    fr012 = fr_table(doc, "FR-012")
    edit_value(fr012, "Current evidence",
               "True for rbac.role_change alone, because in most schools the one person who "
               "administers roles is also the only one who can approve a change,",
               "True for rbac.role_change and rbac.role_grant alone, because in most schools "
               "the one person who administers roles is also the only one who can approve a "
               "change or seat the first holder of a restricted role,")
    append_to(fr012, "Current evidence",
              " A handler may also declare self_approval_only_when_alone, False on the base "
              "handler. Where it and allows_requester_self_approval are both True, "
              "resolve_approvers keeps the requester only where nobody else resolves, so the "
              "exemption covers the school with one administrator and nothing wider. Of the two "
              "exempt types only rbac.role_grant declares it.")
    append_to(fr012, "Acceptance",
              " A restricted role grant keeps its requester only where she is alone: with a "
              "second administrator the frozen list holds that person alone and the requester's "
              "vote is refused with NOT_ELIGIBLE_APPROVER, and the second administrator's "
              "approval is recorded as not self-approved (RoleGrantLadderTests).")
    append_to(fr012, "Current limit",
              " A role grant is not affected: its requester decides only where nobody else "
              "can. Module 4 marks each self-approval: self_approved_change_request on a role "
              "change, and self_approved on a role grant's ROLE_ASSIGNED audit row.")

    fr019 = fr_table(doc, "FR-019")
    edit_value(fr019, "Current evidence", "All fourteen registered types declare an answer.",
               "All fifteen registered types declare an answer.")
    edit_value(fr019, "Current evidence", "role changes refuse once the request is decided.",
               "role changes refuse once the request is decided, and role grants once the grant "
               "is written, naming withdrawing the role from the person's profile as the "
               "remedy.")
    edit_value(fr019, "Acceptance", "so a fifteenth type cannot arrive silently",
               "so a sixteenth type cannot arrive silently")
    edit_value(fr019, "Acceptance", "UserCreationReversalTests and the leave",
               "UserCreationReversalTests, RoleGrantLadderTests and the leave")

    edit_value(fr_table(doc, "FR-021"), "Current limit",
               "as blocking role changes, although that person may decide their own.",
               "as blocking role changes and restricted role grants, although that person may "
               "decide their own of both, a sole holder being exactly the case in which a role "
               "grant's requester stays eligible.")

    fr033 = fr_table(doc, "FR-033")
    edit_value(fr033, "Current evidence", "role change, platform user creation,",
               "role change, restricted role grant, platform user creation,")
    edit_value(fr033, "Acceptance",
               "focused domain tests exercise all fourteen registered document types and the "
               "payout masking rule.",
               "focused domain tests exercise the fourteen document types registered before "
               "role grants and the payout masking rule, and RoleGrantLadderTests exercises the "
               "role grant's layout.")

    # The FR-033 heading was the one requirement heading not kept with its table.
    paragraph_starting(doc, "FR-033 Present").paragraph_format.keep_with_next = True

    step3 = row(table_headed(doc, "Step", "Rule"), "3")
    edit_cell(step3.cells[-1], "which only role changes do.",
              "which only role changes and restricted role grants do; a role grant's handler "
              "also declares self_approval_only_when_alone, so its requester stays only where "
              "nobody else resolves.")

    edit_cell(row_labelled(doc, "Module 4, Roles & Permissions").cells[-1],
              "A role permission change is itself a document routed here: its rbac.role_change "
              "handler applies the change on approval, and it is the one type that lets the "
              "requester decide their own.",
              "A role permission change is itself a document routed here, and so is a grant of "
              "a role carrying restricted permissions its granter does not hold: the "
              "rbac.role_change handler applies the change on approval and the rbac.role_grant "
              "handler writes the grant, and they are the two types that let the requester "
              "decide their own, the role grant only where nobody else can.")

    edit_paragraph(doc, "These are current risks and gaps, not history. The controlled-spend",
                   "the one document type whose requester may decide it",
                   "the two document types whose requester may decide them")
    edit_box(table_with_header(doc, "FURTHER GAPS").rows[0].cells[0], [
        ("sub", "• Continue-without-approval", "Payouts and role changes forbid it",
         "Payouts, role changes and role grants forbid it"),
        ("sub", "• A role change may be decided",
         "Module 4's audit marks the change as self-approved.",
         "Module 4's audit marks the change as self-approved. A restricted role grant is "
         "narrower: its handler declares self_approval_only_when_alone, so where a second "
         "administrator exists that person decides it and the requester cannot."),
        ("sub", "• Reversal is answered", "All fourteen registered types answer",
         "All fifteen registered types answer"),
    ])

    retitle_paragraph(doc, "Module 7 carries 27 capability entries", (
        f"Module 7 carries 28 capability entries in MRD v{MRD_TARGET}. Each maps to the "
        "requirements below. Requiring a document type to declare what its approval releases "
        "before it may register strengthens action reversal and the domain handler entries "
        "without changing the count, and the restricted role grant ladder is one more document "
        "type on the role-based resolution, central-template and reversal entries, adding "
        "none."))

    log_change(doc, M07_TARGET, M07_SUMMARY)
    assert_absent_outside_log(
        doc, "fourteen registered", "True for rbac.role_change alone", "one document type whose",
        "role changes are the one type", "which only role changes do", "MRD v2.80",
        "MRD v2.81", "27 capability entries",
    )
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M07_DIR, M07_STEM, M07_TARGET),
           f"{M07_STEM.replace('_', ' ')} v{M07_TARGET}", M07_TARGET)


# ── M12 Staff Management ─────────────────────────────────────────────────────

M12_DIR = "12-staff-management"
M12_STEM = "XVS_M12_Staff_Management_Functional_Requirements_Document"
M12_SOURCE, M12_TARGET = "2.10", "2.11"

M12_SUMMARY = (
    "Minor revision. FR-019's bulk grant applied no restricted rule at all, so a school "
    "administrator could bulk-grant themselves Finance Admin; it now calls Module 4's "
    "grant_role for each person, so a role carrying restricted permissions the granter does "
    "not hold raises one approval request per person, named under pending_approval, and "
    "somebody already waiting is reported there rather than asked for twice. FR-005's roles "
    "read lists waiting requests under pending, apart from the roles, and answers it empty on "
    "an as-at read. FR-001 records that creating a person with such a role is refused, in one "
    "sentence saying to add them without it and grant it from their profile; this document has "
    "no Needs Attention section, so the limitation joins section 3.4 and Module 4 carries it as "
    "a gap. The school permission table said the bulk grant called the RBAC endpoint; it calls "
    "the service. The roles read joins the endpoint table. "
    f"MRD v{MRD_TARGET}. " + TEST_EVIDENCE
)


def patch_m12() -> None:
    doc = Document(str(frd_path(M12_DIR, M12_STEM, M12_SOURCE)))
    set_control(doc, "Version", M12_TARGET)
    set_control(doc, "Date", REVIEW_DATE)
    set_control(doc, "Supersedes", f"v{M12_SOURCE}")
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")
    set_control(doc, "Verified against", CODE_BASELINE)

    append_rows(table_headed(doc, "What is asked for", "Why it cannot be done",
                             "What a school gets instead"), [[
        "Adding a person with a restricted role in one step.",
        "Creation has no approval ladder behind it, and a role carrying restricted permissions "
        "the creator does not hold cannot be written without one, so the create refuses it "
        "rather than granting it or leaving the person half-made. Module 4 records it under "
        "Needs Attention.",
        "Add the person with a role that carries no restricted permission, then grant the "
        "restricted role from their profile, where it waits for approval and is listed as "
        "pending until decided.",
    ]])

    edit_cell(row_labelled(doc, "school.roles.view / .assign").cells[-1],
              "FR-019's bulk grant calls the RBAC endpoint once per person under "
              "school.roles.assign rather than inventing a bulk key.",
              "FR-019's bulk grant calls RBAC's grant_role once per person under "
              "school.roles.assign rather than inventing a bulk key, so a restricted role waits "
              "for approval exactly as a single grant does.")

    append_to(fr_table(doc, "FR-001"), "Refusals",
              " A role carrying restricted permissions the creator does not hold is refused "
              "with 403 and one sentence naming the role and saying to add the person without "
              "it, then grant it from their profile, where it goes for approval, because "
              "creation has no approval ladder behind it (section 3.4).")
    append_to(fr_table(doc, "FR-005"), "Business rules",
              " (6) A grant waiting on the approval ladder is listed under pending, apart from "
              "the roles: a role carrying restricted permissions its granter does not hold is "
              "requested rather than written (M04 FR-028), and a request confers nothing until "
              "it is approved. Each carries its role, reach, the grant it replaces, the reason, "
              "who asked and when. An as-at read answers pending empty, because it asks what "
              "was held.")

    fr019 = fr_table(doc, "FR-019")
    edit_value(fr019, "Permission",
               "it calls the existing role-assignment endpoint once per person, inside one "
               "transaction, so a bulk grant obeys exactly the rules a single grant does.",
               "it calls RBAC's grant_role once per person, inside one transaction, so a bulk "
               "grant obeys exactly the rules a single grant does, the restricted rule "
               "included.")
    append_to(fr019, "Business rules",
              " (7) A role carrying restricted permissions the granter does not hold is not "
              "written: grant_role raises one approval request per person, the response names "
              "them under pending_approval and the message counts them. Somebody already "
              "waiting for that role at that reach is reported there too, and not asked for "
              "twice.")

    endpoints = table_headed(doc, "Method and path", "Permission", "Requirement")
    keep_format(row(endpoints, "POST /v1/i/me/staff/roles/bulk/").cells[-1],
                "FR-019. data names the people under granted, already_held and "
                "pending_approval.")
    insert_row_after(row(endpoints, "GET /v1/i/me/staff/<id>/history/"), [
        "GET /v1/i/me/staff/<id>/roles/", "school.teachers.view",
        "FR-005. The grants, reach, history and exceptions, and pending: grant requests "
        "waiting for approval, empty on an ?as_at= read.",
    ])

    refusals = table_headed(doc, "Condition", "Status", "Code")
    added = insert_row_after(row(refusals, "A bulk role grant naming somebody this caller "
                                           "cannot see"), [
        "A role carrying restricted permissions the creator does not hold, on the staff create",
        "403",
        "Not a code. One sentence naming the role: add the person without it, then grant it "
        "from their profile, where it goes for approval.",
    ])
    insert_row_after(added, [
        "The same role in a bulk grant", "200",
        "Not a refusal. Each person gets an approval request, named under pending_approval; "
        "nothing is granted until it is approved.",
    ])

    paragraph = paragraph_starting(doc, "A bulk role grant over five people")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " A bulk grant of a role carrying restricted permissions the granter does not hold "
        "writes no grant: it raises one approval request per person, names them under "
        "pending_approval, and names somebody already waiting there rather than asking twice "
        "(RoleGrantLadderTests)."))
    insert_after(doc, "Every list endpoint returns [] for an empty result",
                 "A person's roles read answers pending as [] when nothing waits, and lists a "
                 "grant waiting for approval under pending and never under roles "
                 "(EmptyShapeTests).")

    edit_paragraph(doc, "MRD v2.87 records Module 12", "MRD v2.87", f"MRD v{MRD_TARGET}")

    log_change(doc, M12_TARGET, M12_SUMMARY)
    assert_absent_outside_log(doc, "MRD v2.87", "calls the existing role-assignment endpoint",
                              "calls the RBAC endpoint once per person")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M12_DIR, M12_STEM, M12_TARGET),
           f"{M12_STEM.replace('_', ' ')} v{M12_TARGET}", M12_TARGET)


# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_CONTENTS_NOTE = "A restricted role grant waits for approval"
MRD_INTRO = (
    "This revision records that a role carrying restricted permissions its granter does not "
    "hold is granted through an approval ladder rather than refused, that a bulk grant keeps "
    "the same rule, and that creating a person with such a role is still refused."
)
MRD_CHANGE_SUMMARY = (
    "Giving somebody a role whose restricted permissions the granter does not hold is no "
    "longer refused: it raises a request that a one-stage School Admin ladder decides, the "
    "requester only where nobody else can, and the grant is written on approval naming the requester. A school "
    "with one administrator can therefore seat its first Finance Admin itself; only CodeX "
    "could. A single grant, a change of a grant's role and the staff bulk grant all obey the "
    "rule; the bulk grant had applied no restricted check, so an administrator could "
    "bulk-grant themselves Finance Admin. Creating a person with such a role is still refused. "
    "Capability entries unchanged at 510 across 31 modules; statuses do not move. M04 v1.28, "
    "M07 v1.16 and M12 v2.11. Backend evidence; the school and console integrations are "
    "implemented and pending release; no deployment claim."
)
MRD_DELTA_ROWS = [
    ["Restricted role grants", "Wait for approval, not refused",
     "Giving somebody a role whose restricted permissions the granter does not hold raises a "
     "request that a one-stage School Admin ladder decides, the requester only where nobody "
     "else can; the grant is written on approval naming the requester."],
    ["The first Finance Admin", "Seatable by the school",
     "A school with one administrator could not grant Finance Admin to anybody, because nobody "
     "held its permissions, so only CodeX could. The administrator now raises and approves it, "
     "recorded as self-approved; where a second administrator exists, that person decides."],
    ["Bulk role grant", "Obeys the restricted rule",
     "The staff bulk grant applied no restricted check, so an administrator could bulk-grant "
     "themselves Finance Admin. It raises one request per person and names them under "
     "pending_approval."],
    ["Changing a grant's role", "Old role kept until approval",
     "Replacing a grant with a restricted role waits for approval, and the grant it replaces "
     "stays in force until the approval withdraws it in the same transaction."],
    ["Adding a person with one", "Still refused",
     "Account and staff creation have no ladder, so a restricted role the creator does not hold "
     "is refused in one sentence: add them without it, then grant it from their profile."],
    ["Module FRDs", "Three revised", "M04 v1.28, M07 v1.16, M12 v2.11."],
]
MRD_M04_DECISION = (
    "• A restricted role is granted through a ladder of the same kind, not refused. Giving "
    "somebody a role whose restricted permissions the granter did not hold was refused, and in "
    "a school with one administrator nothing could ever satisfy it: nobody holds the finance "
    "permissions until a Finance Admin exists, so nobody could seat the first one, and only "
    "CodeX could. The grant now raises a request that one School Admin stage decides, and "
    "unlike a role change its requester may decide it only where nobody else can, so a school "
    "with a second administrator gets that person's decision. It is written on approval naming the requester as granter. A "
    "single grant, a change of a grant's role and a bulk grant keep the same rule; the bulk "
    "grant applied none and let an administrator bulk-grant themselves Finance Admin. Creating "
    "a person with such a role is still refused: add them without it and grant it from their "
    "profile."
)


def module_status_cell(doc, module: int, label: str):
    status = doc.tables[mrd_tools.CAPABILITIES[module] - 1]
    for cell in status.rows[0].cells:
        if cell.paragraphs[0].text.strip().startswith(f"{label}:"):
            return cell
    raise ValueError(f"Module {module} has no {label} status cell")


def box_of(doc, module: int, heading: str):
    grid = doc.tables[mrd_tools.CAPABILITIES[module]]
    cells = list({id(c._tc): c for r in grid.rows for c in r.cells
                  if c.text.startswith(heading)}.values())
    if len(cells) != 1:
        raise ValueError(f"Module {module} has {len(cells)} {heading} boxes")
    return cells[0]


def edit_blurb(doc, module: int, old: str, new: str) -> None:
    blurb = mrd_tools.module_blurbs(doc)[module]
    if blurb.text.count(old) != 1:
        raise ValueError(f"Module {module}'s description holds {old!r} "
                         f"{blurb.text.count(old)} times")
    mrd_tools.retitle(blurb, blurb.text.replace(old, new))


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    doc = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta, log = tables[76], tables[78]
    assert delta.rows[0].cells[0].text.strip().endswith("capability delta")
    assert index.rows[0].cells[5].text.strip() == "Entries"

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for r in control.rows:
        label = r.cells[0].text.strip()
        if label == "Version":
            replace_cell(r.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(r.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(r.cells[1], CODE_BASELINE, size=9)
    for r in contents.rows:
        if r.cells[0].text.strip().startswith("5."):
            keep_format(r.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(r.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (text.startswith("5. ") and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    edit_box(box_of(doc, 4, "CURRENT DECISION"), [
        ("after", "• A role change is decided by a ladder", MRD_M04_DECISION),
    ])
    edit_blurb(doc, 4, "Documented by M04 FRD v1.27.",
               "A role carrying restricted permissions its granter does not hold is not "
               "refused either: the grant waits for an approval ladder, whether given singly, "
               "as a replacement or in bulk. Documented by M04 FRD v1.28.")

    code = module_status_cell(doc, 7, "Code")
    runs = code.paragraphs[0].runs
    if runs[1].text.count("role-change, user-creation") != 1:
        raise ValueError("Module 7's code cell no longer lists the role-change handler")
    runs[1].text = runs[1].text.replace("role-change, user-creation",
                                        "role-change, role-grant, user-creation")
    blurb7 = mrd_tools.module_blurbs(doc)[7]
    if "Documented by M07" in blurb7.text:
        raise ValueError("Module 7's description already names an FRD version")
    mrd_tools.retitle(blurb7, blurb7.text.rstrip() + (
        " A grant of a role carrying restricted permissions its granter does not hold is a "
        "document type of its own, decided by a one-stage ladder shaped like the role-change "
        "one, whose requester may decide it only where nobody else is eligible. Documented by "
        "M07 FRD v1.16."))
    edit_box(box_of(doc, 7, "NEEDS ATTENTION"), [
        ("sub", "• Continue-without-approval", "Payouts and role changes forbid it",
         "Payouts, role changes and role grants forbid it"),
        ("replace", "• A role change may be decided",
         "• A role change may be decided by the person who raised it, and its stage completes "
         "on any one approval, so an administrator can give their own role a restricted "
         "permission with nobody else looking, even where the school has a second "
         "administrator; the ladder and Module 4's audit record that they did. A restricted "
         "role grant also exempts its requester, but only where nobody else is eligible: a "
         "second administrator, where there is one, decides it."),
        ("sub", "• Reversal is answered", "All fourteen registered types answer",
         "All fifteen registered types answer"),
    ])

    edit_blurb(doc, 12, "Documented by M12 FRD v2.10.",
               "A bulk role grant keeps the restricted rule a single grant keeps: a role "
               "carrying restricted permissions the granter does not hold waits for approval, "
               "one request per person, and the profile lists what waits apart from the roles "
               "held. Documented by M12 FRD v2.11.")
    edit_box(box_of(doc, 12, "NEEDS ATTENTION"), [
        ("append", "• Adding a person with a restricted role is refused rather than routed for "
         "approval. Creation has no ladder behind it, so a role carrying restricted permissions "
         "the creator does not hold is refused with a sentence saying to add the person without "
         "it and grant it from their profile, where it waits for approval. A school adding its "
         "bursar makes two moves where one would do. Module 4 carries the item."),
    ])

    total = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total != 510:
        raise ValueError(f"Capability total is {total}, expected 510")

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)
    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if len(intro) != 1:
        raise ValueError(f"{len(intro)} delta introductions found")
    mrd_tools.retitle(intro[0], MRD_INTRO)

    # Copied from the latest version row, so the new row reads at the same size.
    log.rows[1]._tr.addprevious(copy.deepcopy(log.rows[1]._tr))
    for cell, text in zip(log.rows[1].cells, (MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)):
        keep_format(cell, text)
    assert_absent_outside_log(doc, "All fourteen registered", "the one document type with that",
                              "Payouts and role changes forbid")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m04()
    patch_m07()
    patch_m12()
    patch_mrd()


if __name__ == "__main__":
    main()
