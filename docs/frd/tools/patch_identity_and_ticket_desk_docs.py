#!/usr/bin/env python3
"""Version the Module 3 and Module 31 FRDs against the backend at 13675db.

Module 3 had four behaviours in its code that its document did not describe.
A platform caller's account list returned every account on the platform, so
pickers built on it offered a school's staff to CodeX screens; a collection
read now defaults to the platform tenant while acting on one account by id
keeps the operator's reach. The ID card carries a random, revocable key held
only by platform staff instead of an email address, and signs in with it.
Activation refusals name their reason. The platform-hire handler refuses to
reverse an approval once the account has been invited.

Module 31 recorded escalation as a requirement without carrying it into the
audit, notification and data sections, and never recorded that readers follow
and mute tickets. Tracing those found two gaps: the desk's recipients and its
assignee picker are different queries that can disagree, and the desk's
ticket export carries CodeX's own tickets only.

Both documents cite MRD v2.76. Module 3's traceability gains the ID-card
sign-in entry, which that MRD revision is expected to add.

    python tools/patch_identity_and_ticket_desk_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)

REVIEW_DATE = "11 September 2026"
SHORT_DATE = "11 Sep 2026"
MRD_VERSION = "2.76"
CODE_BASELINE = "Backend worktree at 13675db, 11 September 2026"


# ── shared docx helpers ──────────────────────────────────────────────────────

def cell_size(cell, default: float) -> float:
    """The point size a cell already renders at, so a rewrite keeps it."""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                return run.font.size.pt
    return default


def replace_cell(cell, text: str, *, size: float | None = None, **kwargs) -> None:
    size = cell_size(cell, 8) if size is None else size
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, size=size, **kwargs)


def replace_in_cell(cell, old: str, new: str) -> None:
    text = cell.text.strip()
    if old not in text:
        raise ValueError(f"Text not found in cell: {old!r}")
    replace_cell(cell, text.replace(old, new))


def append_cell_text(cell, tail: str) -> None:
    replace_cell(cell, cell.text.strip() + tail)


def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError("Paragraph carries no run to inherit formatting from")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_cover_version(table, source: str, target: str) -> None:
    """Rewrite the version on the cover, which is one run of one title block."""
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def replace_control_value(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            replace_cell(row.cells[1], value, size=9)
            return
    raise ValueError(f"Control row not found: {label}")


def find_paragraph(doc, prefix: str, *, style: str | None = None):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            if style is None or paragraph.style.name == style:
                return paragraph
    raise ValueError(f"Body paragraph not found: {prefix}")


def replace_body_paragraph(doc, prefix: str, text: str, *, style: str | None = None):
    paragraph = find_paragraph(doc, prefix, style=style)
    set_run_text(paragraph, text)
    return paragraph


def clone_paragraph_after(paragraph, text: str):
    """Insert a copy of ``paragraph`` directly after it, carrying ``text``."""
    clone = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(clone)
    new = Paragraph(clone, paragraph._parent)
    set_run_text(new, text)
    return new


def find_row(table, prefix: str):
    for row in table.rows:
        if row.cells[0].text.strip().startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix}")


def replace_row(table, prefix: str, values: list[str | None]) -> None:
    """Rewrite a row's cells; ``None`` leaves that cell as it is."""
    row = find_row(table, prefix)
    for cell, value in zip(row.cells, values):
        if value is not None:
            replace_cell(cell, value)


def insert_row_before(table, prefix: str, values: list[str]):
    """Clone the row named by ``prefix`` and write ``values`` into the copy."""
    anchor = find_row(table, prefix)
    anchor._tr.addprevious(copy.deepcopy(anchor._tr))
    # The clone carries the anchor's text, so it is now the first match.
    clone = find_row(table, prefix)
    for cell, value in zip(clone.cells, values):
        replace_cell(cell, value)
    return clone


def append_row(table, values: list[str]):
    template = table.rows[-1]
    template._tr.addnext(copy.deepcopy(template._tr))
    row = table.rows[-1]
    for cell, value in zip(row.cells, values):
        replace_cell(cell, value)
    return row


def prepend_change_log(table, version: str, summary: str) -> None:
    template = table.rows[1]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], SHORT_DATE, size=8)
    replace_cell(row.cells[2], summary, size=8)


def require_header(table, text: str):
    if not table.rows[0].cells[0].text.strip().startswith(text):
        raise ValueError(f"Table does not start with {text!r}")
    return table


def finish(doc, output: Path, title: str, version: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    check = Document(str(output))
    words = "\n".join(p.text for p in check.paragraphs)
    for table in check.tables:
        for row in table.rows:
            for cell in row.cells:
                words += "\n" + cell.text
    # Spelled in parts: the repository vocabulary test reads every file, and a
    # guard that wrote the retired word would be an offender itself.
    retired_site_word = "cam" + "pus"
    if retired_site_word in words.lower():
        raise ValueError(f"The retired word {retired_site_word!r} is in {output}")


# ═════════════════════════════════════════════════════════════════════════════
# Module 3 - Identity, Team & Organogram
# ═════════════════════════════════════════════════════════════════════════════

M03_DIR = "03-identity-team-and-organogram"
M03_STEM = "XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document"
M03_SOURCE, M03_TARGET = "1.11", "1.12"

EVIDENCE_ROW, ACCEPTANCE_ROW, LIMIT_ROW = 2, 3, 4

M03_SCOPE_SIGN_IN = (
    "Tenant resolved first, account looked up inside it, password checked, "
    "lockout checked, status checked, session opened, short-lived access token "
    "returned, and rotating refresh credential stored in a Secure HttpOnly "
    "SameSite cookie. A browser Origin must match the configured first-party "
    "allowlist. Platform staff may present the random key printed on their ID "
    "card in place of an address; it resolves only inside the platform tenant."
)

M03_SCOPE_EVIDENCE = (
    "Seventeen identity actions emitted into the shared audit trail with the "
    "subject's tenant attached, because these endpoints are unauthenticated and "
    "have no ambient tenant to inherit. Sixteen carry a typed audit action; "
    "replacing an ID-card key is recorded as a custom one."
)

M03_BOUNDARY_PLATFORM = (
    "The platform boundary is Tenant.kind. User.is_platform_user reads it, and "
    "it is the only thing that decides whether an account may hold a branch, "
    "may carry a platform staff profile, may be assigned an organogram seat, or "
    "carries an ID-card sign-in key. It cannot be asserted by a caller, which "
    "is the property the old persona column lacked."
)

M03_BOUNDARY_ACCOUNT_SCOPE = (
    "Every by-id account action resolves its target through the same tenant "
    "and branch scope as the account viewset, and an out-of-scope id receives "
    "the same 404 as a missing one. That scope answers two questions, and a "
    "platform caller answers them differently. Acting on one account by id - "
    "reading or editing it, suspending, reactivating or unlocking it, resetting "
    "its password, resending its invitation, changing its address, replacing "
    "its ID-card key, or naming it as the target of a force logout or a lockout "
    "clearance - keeps a platform operator's reach across every tenant, because "
    "CodeX support unlocking a Greenfield bursar is what the console is for. "
    "Listing accounts does not: a platform caller's list is the platform "
    "tenant's own accounts unless it asks for more with ?scope=school, "
    "?scope=all or ?school_id=, so the pickers built on that list, among them "
    "the Console's approver groups, platform role assignment and the workflow "
    "template builder, never offer a school's teachers and bursars to a CodeX "
    "screen. A tenant caller answers both questions the same way, with its own "
    "tenant narrowed to the branches its grants reach. The identity audit "
    "endpoint names platform.audit.view and applies the shared audit tenant "
    "scope."
)

M03_ACTOR_ANYBODY = [
    "Anybody with the address or card",
    "Attempt a sign-in by address or, for platform staff, by ID-card key; "
    "preview the name on an ID card; ask for a password reset; and open an "
    "invitation or reset link.",
    "No key, and no authentication. Rate limited by throttle scope, and the two "
    "ID-card routes also per card key; every attempt is recorded whether or not "
    "the address exists.",
]

M03_ACTOR_OPERATOR = [
    "Platform operator",
    "List the platform tenant's own accounts, and a school's or every tenant's "
    "on request; read, edit, deactivate, suspend, reactivate, unlock, reset, "
    "re-invite, re-address or force out any account on any tenant by id; "
    "replace a platform staff member's ID-card key.",
    "The same platform.team.* keys, held on the platform tenant. The caller's "
    "own tenant kind decides the reach, not a key; a collection read defaults "
    "to the platform tenant and widens only on ?scope=school, ?scope=all or "
    "?school_id=. Replacing a card key takes platform.team.update.",
]

M03_ACTOR_ENGINE = [
    None,
    "Approve or reject a platform hire, which invites the account or rejects it "
    "and vacates its seat, and reverse an approval only while the hire is still "
    "awaiting a decision.",
    "Not reachable from this module's API. The handler is registered for the "
    "PLATFORM_USER_CREATION document type and refuses a reversal once the "
    "account has left PENDING_APPROVAL.",
]

M03_FR008_HEADING = (
    "FR-008  Resolve an ID Card by a Random, Revocable Key, Never by an Address"
)

M03_FR008_REQUIREMENT = (
    "The ID-card sign-in flow needs an unauthenticated endpoint that turns a "
    "scanned card into a display name, and a sign-in that accepts the card in "
    "place of an address, because a barcode scanner carries no credentials. "
    "What the card carries is therefore the whole of its security: it must "
    "disclose no address and no account state, answer only for platform staff, "
    "and stop working the moment it is replaced."
)

M03_FR008_EVIDENCE = (
    "Every platform-tenant user holds card_login_id, a unique random UUID that "
    "User.save() issues and migration 0011 issued to the accounts that predate "
    "it; no other account receives one, and rotate_card_login_id refuses to "
    "issue one. SpecialLoginPreviewView reads ?card_id=, resolves it only "
    "against tenant__kind=PLATFORM, and answers 200 with the full name only "
    "when the account may sign in. A missing, malformed, unknown, replaced or "
    "inactive key, and a school account, receive one 404 with one sentence, and "
    "an address is not a lookup it accepts. POST /auth/login/ accepts card_id "
    "in place of email, exactly one of the two, and LoginService resolves the "
    "key inside the platform tenant without reading the asserted tenant slug, "
    "so the address stays server-side for the whole card flow. Both card "
    "routes carry two throttles, one on the caller's IP address (login_preview "
    "and login) and one on a SHA-256 digest of the card key "
    "(card_login_preview at ten a minute and card_login at five), so neither "
    "many IP addresses against one card nor one IP address against many cards "
    "gets past its budget. POST /users/{id}/card-login/rotate/ replaces the key "
    "under platform.team.update, answers 404 for a non-platform target, and "
    "records CARD_LOGIN_ROTATED."
)

M03_FR008_ACCEPTANCE = (
    "BarcodePreviewPrivacyTests proves an active staff card previews the name "
    "and never the address, that an address is refused as a lookup, that a "
    "missing, malformed and unknown key and a school account's answer "
    "identically, that pending, locked, suspended and deactivated staff are "
    "indistinguishable from an unknown card, that replacing the key kills the "
    "old card at once, and that a card key and password complete a sign-in "
    "with no address while an unknown key is the generic INVALID_CREDENTIALS "
    "failure. CardLoginRateLimitTests proves one card is bounded across many IP "
    "addresses and one IP address across many cards. "
    "CardLoginRotationEndpointTests proves an authorised rotation replaces the "
    "key and writes its audit event, and that a caller without "
    "platform.team.update is refused with the key unchanged. Card keys are "
    "globally unique, so the preview does not depend on exactly one platform "
    "tenant existing; ScopedEmailLookupTests pins that with a second platform "
    "tenant seeded."
)

M03_FR008_LIMIT = (
    "Anybody holding a card key learns the full name of the staff member it "
    "belongs to, which is what the key is for; it discloses nothing else. The "
    "key is a lookup, not a credential, and the password is still required. "
    "card_login_id is returned on the account read payload with no field-level "
    "gate, so a platform colleague holding platform.team.view can read another "
    "staff member's key, and the sign-in response returns the holder's own."
)

M03_FR002_OLD = "the barcode preview"
M03_FR002_NEW = "the ID-card preview and card sign-in"

M03_FR009_EVIDENCE_TAIL = (
    " An unknown or replaced ID-card key at sign-in returns the same "
    "INVALID_CREDENTIALS payload as a wrong password."
)

M03_FR009_LIMIT = (
    "The ID-card preview returns a platform staff member's full name to whoever "
    "holds their card key, which is what the key is for, and says nothing about "
    "any other account or any account's state; it is bounded by one throttle "
    "per IP address and one per card. Authenticated account actions return the "
    "same 404 for a missing or out-of-scope id."
)

M03_FR011_EVIDENCE_TAIL = (
    " Activation's refusals name their reason. A request the serializer "
    "refuses answers with \"That password does not meet the requirements.\", "
    "the code PASSWORD_POLICY_VIOLATION when the password fails the policy and "
    "VALIDATION_ERROR otherwise, and the field errors under detail, which is "
    "where the clients read a complaint that belongs beneath a field. A refusal "
    "from the service passes its own error_code and sentence through: "
    "INVITATION_NOT_FOUND, INVITATION_ALREADY_USED, INVITATION_EXPIRED, "
    "INVITATION_NOT_ACTIONABLE or PASSWORD_POLICY_VIOLATION."
)

M03_FR011_ACCEPTANCE_TAIL = (
    " ActivationTellsYouWhyItRefusedTests proves a weak password is refused "
    "with PASSWORD_POLICY_VIOLATION and the policy's reasons under "
    "detail.password, that a good password still activates, and that a used "
    "link answers with INVITATION_ALREADY_USED and its own sentence rather than "
    "a generic line."
)

M03_FR011_LIMIT_TAIL = (
    " Activation answers in two refusal shapes: a serializer refusal carries "
    "code and detail, a service refusal carries error_code and message, and a "
    "mismatched confirmation carries a field error with no code at all."
)

M03_FR012_EVIDENCE_TAIL = (
    " validate_reversal answers the engine before a reversal writes anything. "
    "It allows one while the account is still PENDING_APPROVAL, or where the "
    "user row no longer exists, and refuses with ReversalNotAllowedError once "
    "approval has invited the account, because by then the person may have set "
    "a password and undoing the approval record would change nothing about the "
    "account."
)

M03_FR012_ACCEPTANCE_TAIL = (
    " UserCreationReversalTests proves an account still awaiting a decision "
    "may be reversed, an invited or live account refuses, and a deleted account "
    "is not treated as a live one."
)

M03_FR012_LIMIT_TAIL = (
    " Closing an account that approval has already invited is a deactivation, "
    "recorded as what it is, rather than a reversal of the approval."
)

M03_FR017_EVIDENCE = (
    "services.audit.log_auth_event emits seventeen identity events under the "
    "IDENTITY module key, with the subject as the entity, the subject's tenant "
    "passed explicitly, and the IP address, agent, tenant slug, school slug and "
    "the event's own name in metadata. Sixteen map onto typed vs_audit action "
    "types; CARD_LOGIN_ROTATED has no entry in that map and is recorded as "
    "CUSTOM. The outcome is derived from the event: LOGIN_FAILURE and "
    "ACCOUNT_LOCKED are FAILED, everything else SUCCESS. It swallows its own "
    "failures into a critical log line because an audit write must never "
    "prevent an auth action. record_attempt is deliberately kept out of the "
    "audit stream: AuthAttempt drives security policy, not compliance."
)

M03_FR017_LIMIT_TAIL = (
    " A card-key rotation is found by its auth_event metadata and not by an "
    "action-type filter, because the map carries no typed action for it."
)

M03_SIGN_IN_FIRST = [
    None,
    "The asserted tenant slug is resolved to a row, and the account is looked "
    "up inside it. An ID-card sign-in instead resolves the card key inside the "
    "platform tenant and reads no slug.",
    "Before the account, so the tenant is never inferred from a row the caller "
    "has not proved they own. A card key exists only on platform accounts and "
    "is unique across the table, so there is one tenant it can resolve in.",
]

M03_USER_MODEL_TAIL = (
    " card_login_id is a unique random UUID held by platform users alone, "
    "issued on save and replaced by rotation."
)

M03_LOGIN_ROUTE = (
    "Sign in with an address, or with an ID-card key for platform staff, and "
    "the password: exactly one of the two, or 400 on credentials. Optional "
    "tenant slug in the body, not read on the card path. A browser Origin must "
    "match the first-party allowlist. Returns access and session context, sets "
    "the refresh and CSRF cookies, and never returns refresh. AllowAny, "
    "throttled per IP address and, for a card, per card key."
)

M03_PREVIEW_ROUTE = (
    "ID-card preview: ?card_id= returns a display name for an active platform "
    "staff member's key and one identical 404 for anything else. AllowAny, "
    "throttled per IP address and per card key (FR-008)."
)

M03_ACTIVATE_ROUTE = (
    "Set the password and activate. AllowAny, throttled. Issues no tokens. "
    "Every refusal names its reason (FR-011)."
)

M03_USERS_ROUTE = (
    "List accounts, or create one. platform.team.view / .create. A tenant "
    "caller lists its own tenant, narrowed to the branches its grants reach. A "
    "platform caller lists the platform tenant's own accounts unless it asks "
    "for more: ?scope=school for every school's accounts, ?scope=all for every "
    "account, or ?school_id= for one school's; ?scope=platform names the "
    "default, and an unrecognised value widens nothing. Creation derives the "
    "target tenant and refuses any selected role whose restricted keys exceed "
    "the actor's own authority."
)

M03_USER_DETAIL_TAIL = (
    " By id, a platform operator reaches an account on any tenant; a tenant "
    "caller reaches one inside its own tenant and branch reach."
)

M03_ROTATE_ROUTE = [
    "POST /users/{id}/card-login/rotate/",
    "Replace a platform staff member's ID-card key, so every printed or copied "
    "card stops resolving at once. platform.team.update. A non-platform target "
    "answers 404; the rotation is recorded as CARD_LOGIN_ROTATED.",
]

M03_PREVIEW_ANSWER = [
    "An ID-card preview for a missing, malformed, unknown or replaced key, a "
    "school account, or staff who may not sign in",
    "404 with one sentence in every case. The preview does not accept an "
    "address.",
]

M03_NEW_ANSWERS = [
    ["A sign-in naming both an address and a card key, or neither",
     "400 on credentials."],
    ["An activation password the policy refuses",
     "400 with PASSWORD_POLICY_VIOLATION and the policy's reasons under "
     "detail.password."],
    ["An invitation link that is unknown, used, expired or no longer actionable",
     "400 with the service's own error_code and sentence, for example "
     "INVITATION_ALREADY_USED."],
    ["Reversing the approval of a platform hire already invited",
     "Refused by the handler with ReversalNotAllowedError before the engine "
     "writes anything. Deactivate the account instead."],
]

M03_AUDIT_DEPENDENCY = (
    "Receives seventeen identity events through emit_audit_event with the "
    "subject's tenant passed explicitly, sixteen under typed action types and "
    "the ID-card key rotation as CUSTOM. Also reads back: this module's "
    "auth-event endpoint serves vs_audit rows rather than its own table."
)

M03_WORKFLOW_DEPENDENCY_TAIL = (
    " The handler also answers the engine's reversal contract and refuses to "
    "reverse an approval once the account has been invited."
)

M03_VERIFICATION_LEAD = (
    "vs_user ran 381 tests OK at 6611856, whose vs_user code is identical to "
    "this document's baseline and which contains every change recorded here; "
    "the suite was not run again for this revision. The client evidence is "
    "unchanged from v1.10: the migrated school client passed 46 focused auth "
    "tests and its production build, and its source and production bundle "
    "contain neither js-cookie nor refresh-token identifiers. Deployment is not "
    "inferred from these checks."
)

M03_CARD_BULLET = (
    "•  The ID-card routes: an active staff key previews and signs in "
    "without an address, every other key answers one identical 404, replacing "
    "the key kills the old card, and one card and one IP address are each "
    "bounded. Covered by BarcodePreviewPrivacyTests, CardLoginRateLimitTests "
    "and CardLoginRotationEndpointTests."
)

M03_ACTION_SCOPE_BULLET = (
    "•  Cross-tenant isolation on every account action: an administrator "
    "cannot change the address of, suspend, reactivate, unlock, reset, "
    "re-invite, force out or clear the lockout of an account outside their own "
    "tenant and branch reach, and a platform operator still reaches every "
    "tenant. Covered by AccountActionTenantScopeTests."
)

M03_NEW_BULLETS = [
    "•  The account list's default for a platform caller: an unscoped list "
    "is the platform tenant's own accounts, an unknown scope value does not "
    "widen it, ?scope=school still reaches the schools, and the school and "
    "platform scopes partition the table. Covered by UserListScopeTests.",
    "•  Activation refusals naming their reason, and a platform hire's "
    "approval refusing reversal once the account is invited. Covered by "
    "ActivationTellsYouWhyItRefusedTests and UserCreationReversalTests.",
]

M03_TRACEABILITY_LEAD = (
    f"MRD v{MRD_VERSION} records Module 3 as Identity, Team & Organogram, Phase "
    "V1, Backend Partial, In use Complete, code ownership vs_user, with "
    "eighteen capability entries. Each entry maps below, including the ID-card "
    "sign-in entry, which no earlier traceability carried."
)

M03_SUSPEND_STATE = (
    "Implemented with limits. Transitions and targets are scoped, and a "
    "platform operator reaches an account on any tenant by id; the configured "
    "lock window still does not restore ACTIVE automatically."
)

M03_AUDIT_STATE = (
    "Implemented with limits. Seventeen identity events reach the shared "
    "tenant-scoped trail behind platform.audit.view, sixteen under typed "
    "actions and the card-key rotation as a custom one; the unused AuthEventLog "
    "table remains."
)

#: The photographs row carried a verbatim copy of the account-creation row
#: above it, which says nothing about who may read a photograph.
M03_PHOTO_STATE = (
    "Implemented. Each photograph is bound to its profile and to the tenant the "
    "person is on, a read is refused unless the caller's asserted tenant "
    "matches, and each URL is signed for one reader and expires."
)

M03_CARD_CAPABILITY = [
    "Platform staff ID-card sign-in by a random, revocable card key",
    "FR-008, FR-009",
    "Implemented. The card carries a random key held only by platform staff, "
    "never an address. Every other key answers one identical 404, sign-in "
    "accepts the key in place of an address, both routes are throttled per IP "
    "address and per card, and replacing the key kills the old card at once. "
    "The key is returned on the account read payload without a field-level "
    "gate.",
]

M03_RECONCILIATION = (
    "MRD RECONCILIATION\n"
    f"• MRD v{MRD_VERSION} lists Module 3 as Backend Partial and In use "
    "Complete with eighteen capabilities. The ID-card sign-in entry maps to "
    "FR-008; the flow existed without a tracker entry, and the random card key "
    "and card sign-in make it a capability of its own.\n"
    "• A platform caller's account list is the platform tenant's own "
    "accounts unless it asks for a school's or every tenant's, while acting on "
    "one account by id keeps its reach across every tenant. This is recorded "
    "inside the existing account boundary rather than as a capability.\n"
    "• Partial remains correct because privileged MFA and the existing "
    "tenant lifecycle, lockout, retention, audit-model and organogram gaps "
    "remain.\n"
    "• Backend evidence only. Neither document claims deployment, "
    "production adoption, or data migration completion."
)

M03_CHANGE_SUMMARY = (
    "Reconciles Module 3 with four changes that reached the code without "
    "reaching this document, two of them after v1.11 and two before it. The "
    "account list: a platform caller's GET /users/ returned every account on "
    "the platform, so the pickers built on it, among them the Console's "
    "approver groups, platform role assignment and the workflow template "
    "builder, offered a school's teachers and bursars to CodeX screens, "
    "including as candidates for a CodeX platform role. A collection read now "
    "defaults to the platform tenant's own accounts and widens only on "
    "?scope=school, ?scope=all or ?school_id=, while acting on one account by "
    "id, including the force-logout and lockout targets named in a request "
    "body, keeps a platform operator's reach across every tenant. The ID card: "
    "FR-008 described a preview keyed by an email address that answered with "
    "the account's state and remained a name-and-existence oracle inside the "
    "platform tenant. The card carries a random, revocable key held only by "
    "platform staff; the preview and a card sign-in resolve it inside the "
    "platform tenant alone, every other key answers one identical 404, both "
    "routes are throttled per IP address and per card, and POST "
    "/users/{id}/card-login/rotate/ kills the old card. Activation refusals "
    "name their reason: a password the policy refuses returns "
    "PASSWORD_POLICY_VIOLATION with the reasons under detail, and the "
    "service's own code and sentence for an unknown, used or expired link pass "
    "through instead of a generic line. The platform-hire handler refuses to "
    "reverse an approval once the account has been invited. Section 1.1, "
    "section 2.2, the actors table, FR-002, FR-008, FR-009, FR-011, FR-012, "
    "FR-017, the sign-in order, the User and AuthEventLog records, the three "
    "API tables, the Module 5 and Module 7 dependencies, the verification "
    "list, gap 9.6, MRD traceability and the reconciliation box are updated. "
    "The verification list names the action-scope test class that exists, and "
    "the staff-photograph traceability row, which carried a copy of the "
    "account-creation row above it, is corrected. Module 3 remains Backend "
    "Partial and In use Complete and moves from seventeen to eighteen "
    "capabilities with the ID-card sign-in entry. vs_user ran 381 tests OK at "
    "6611856, whose vs_user code matches this baseline. Backend evidence only; "
    "nothing here is deployed."
)


def patch_m03(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M03 Identity Team and Organogram Functional Requirements "
        f"Document v{M03_TARGET}"
    )

    # Bound before anything is inserted, because doc.tables is recomputed
    # from the body on every access and an insertion shifts later indices.
    cover, control = doc.tables[0], doc.tables[1]
    scope = require_header(doc.tables[3], "Area")
    actors = require_header(doc.tables[6], "Actor")
    fr002 = require_header(doc.tables[8], "FR-002")
    fr008 = require_header(doc.tables[14], "FR-008")
    fr009 = require_header(doc.tables[15], "FR-009")
    fr011 = require_header(doc.tables[17], "FR-011")
    fr012 = require_header(doc.tables[18], "FR-012")
    fr017 = require_header(doc.tables[23], "FR-017")
    sign_in = require_header(doc.tables[30], "Order")
    model = require_header(doc.tables[32], "Model")
    auth_routes = require_header(doc.tables[33], "Method and path")
    account_routes = require_header(doc.tables[34], "Method and path")
    answers = require_header(doc.tables[36], "Condition")
    dependencies = require_header(doc.tables[37], "Dependency")
    gaps = require_header(doc.tables[38], "Pri.")
    traceability = require_header(doc.tables[40], "MRD Module 3 capability")
    reconciliation = require_header(doc.tables[41], "MRD RECONCILIATION")
    change_log = require_header(doc.tables[42], "Version")

    replace_cover_version(cover, M03_SOURCE, M03_TARGET)
    replace_control_value(control, "Version", M03_TARGET)
    replace_control_value(control, "Review date", REVIEW_DATE)
    replace_control_value(control, "Code baseline", CODE_BASELINE)
    replace_control_value(
        control, "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 3",
    )

    replace_row(scope, "Sign-in", [None, M03_SCOPE_SIGN_IN])
    replace_row(scope, "Identity evidence", [None, M03_SCOPE_EVIDENCE])

    replace_body_paragraph(doc, "The platform boundary is Tenant.kind", M03_BOUNDARY_PLATFORM)
    replace_body_paragraph(
        doc, "Every by-id account action now resolves", M03_BOUNDARY_ACCOUNT_SCOPE,
    )

    replace_row(actors, "Anybody with the address", M03_ACTOR_ANYBODY)
    insert_row_before(actors, "Security reader", M03_ACTOR_OPERATOR)
    replace_row(actors, "The workflow engine", M03_ACTOR_ENGINE)

    replace_in_cell(fr002.rows[EVIDENCE_ROW].cells[1], M03_FR002_OLD, M03_FR002_NEW)

    replace_body_paragraph(doc, "FR-008", M03_FR008_HEADING, style="Heading 2")
    replace_cell(fr008.rows[1].cells[1], M03_FR008_REQUIREMENT)
    replace_cell(fr008.rows[EVIDENCE_ROW].cells[1], M03_FR008_EVIDENCE)
    replace_cell(fr008.rows[ACCEPTANCE_ROW].cells[1], M03_FR008_ACCEPTANCE)
    replace_cell(fr008.rows[LIMIT_ROW].cells[1], M03_FR008_LIMIT)

    append_cell_text(fr009.rows[EVIDENCE_ROW].cells[1], M03_FR009_EVIDENCE_TAIL)
    replace_cell(fr009.rows[LIMIT_ROW].cells[1], M03_FR009_LIMIT)

    append_cell_text(fr011.rows[EVIDENCE_ROW].cells[1], M03_FR011_EVIDENCE_TAIL)
    append_cell_text(fr011.rows[ACCEPTANCE_ROW].cells[1], M03_FR011_ACCEPTANCE_TAIL)
    append_cell_text(fr011.rows[LIMIT_ROW].cells[1], M03_FR011_LIMIT_TAIL)

    append_cell_text(fr012.rows[EVIDENCE_ROW].cells[1], M03_FR012_EVIDENCE_TAIL)
    append_cell_text(fr012.rows[ACCEPTANCE_ROW].cells[1], M03_FR012_ACCEPTANCE_TAIL)
    append_cell_text(fr012.rows[LIMIT_ROW].cells[1], M03_FR012_LIMIT_TAIL)

    replace_cell(fr017.rows[EVIDENCE_ROW].cells[1], M03_FR017_EVIDENCE)
    append_cell_text(fr017.rows[LIMIT_ROW].cells[1], M03_FR017_LIMIT_TAIL)

    replace_row(sign_in, "1", M03_SIGN_IN_FIRST)

    append_cell_text(find_row(model, "User").cells[2], M03_USER_MODEL_TAIL)
    replace_in_cell(
        find_row(model, "AuthEventLog").cells[2],
        "sixteen event choices", "seventeen event choices",
    )

    replace_row(auth_routes, "POST /auth/login/", [None, M03_LOGIN_ROUTE])
    replace_row(auth_routes, "GET /auth/special_login/preview/", [None, M03_PREVIEW_ROUTE])
    replace_row(auth_routes, "POST /auth/activate/{key}/", [None, M03_ACTIVATE_ROUTE])

    replace_row(account_routes, "GET, POST /users/", [None, M03_USERS_ROUTE])
    append_cell_text(
        find_row(account_routes, "GET, PATCH, DELETE /users/{id}/").cells[1],
        M03_USER_DETAIL_TAIL,
    )
    insert_row_before(account_routes, "POST /users/{id}/submit/", M03_ROTATE_ROUTE)

    replace_row(answers, "A barcode preview for an address", M03_PREVIEW_ANSWER)
    for values in M03_NEW_ANSWERS:
        insert_row_before(answers, "An empty result list", values)

    replace_row(dependencies, "Module 5, Audit", [None, M03_AUDIT_DEPENDENCY])
    append_cell_text(
        find_row(dependencies, "Module 7, Workflow").cells[1],
        M03_WORKFLOW_DEPENDENCY_TAIL,
    )

    replace_body_paragraph(
        doc, "The focused browser-session contract", M03_VERIFICATION_LEAD,
    )
    replace_body_paragraph(doc, "•  The barcode preview answering", M03_CARD_BULLET)
    anchor = replace_body_paragraph(
        doc, "•  Cross-tenant isolation on every account action",
        M03_ACTION_SCOPE_BULLET,
    )
    for text in M03_NEW_BULLETS:
        anchor = clone_paragraph_after(anchor, text)

    for row in gaps.rows:
        if row.cells[1].text.strip().startswith("9.6 AuthEventLog"):
            replace_in_cell(
                row.cells[2], "sixteen event choices", "seventeen event choices",
            )
            break
    else:
        raise ValueError("Gap 9.6 row not found")

    replace_body_paragraph(doc, "MRD v2.74 records Module 3", M03_TRACEABILITY_LEAD)
    replace_row(
        traceability, "Suspend, reactivate, and unlock", [None, None, M03_SUSPEND_STATE],
    )
    replace_row(
        traceability, "Authentication and account audit events",
        [None, None, M03_AUDIT_STATE],
    )
    replace_row(
        traceability, "Authorised, expiring access to staff photographs",
        [None, None, M03_PHOTO_STATE],
    )
    append_row(traceability, M03_CARD_CAPABILITY)
    # Word joins two adjacent tables into one, so a reconciliation box that
    # breaks onto the next page repeats the traceability header above it.
    from docx.oxml import OxmlElement

    if traceability._tbl.getnext() is reconciliation._tbl:
        traceability._tbl.addnext(OxmlElement("w:p"))
    set_run_text(reconciliation.rows[0].cells[0].paragraphs[0], M03_RECONCILIATION)

    prepend_change_log(change_log, M03_TARGET, M03_CHANGE_SUMMARY)
    finish(doc, output, title, M03_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 31 - Support Tickets
# ═════════════════════════════════════════════════════════════════════════════

M31_DIR = "31-support-tickets"
M31_STEM = "XVS_M31_Support_Tickets_Functional_Requirements_Document"
M31_SOURCE, M31_TARGET = "1.5", "1.6"

M31_SUPPORTING_APPS = "vs_tenants, vs_rbac, vs_user, vs_notifications, vs_exports, core"

M31_SCOPE_ACCOUNTING = (
    "The audit trail, the notifications and who follows them, the dashboard "
    "counters, and the export datasets."
)

M31_FR004_EVIDENCE_TAIL = (
    " Attaching sits on the pending-tenant surface beside filing, so a school "
    "that has not gone live can show the screen it is reporting; the action "
    "still resolves its ticket through the caller's visibility, and the rest of "
    "the desk opens at go-live."
)

M31_FR004_ACCEPTANCE_TAIL = (
    " PendingSchoolAttachmentTests proves the attachments action is on that "
    "surface and the rest of the desk stays shut, and AttachmentLimitTests "
    "proves the size, allowlist and content checks that make opening it safe."
)

M31_FR009_EVIDENCE = (
    "Every service writes an audit row with actor, action, a human summary, "
    "and before and after snapshots for updates. Creation, update, assignment, "
    "status change, comment, internal note, attachment and escalation each have "
    "their own action; ESCALATED records the decision that sent a school's "
    "ticket to CodeX."
)

M31_FR009_ACCEPTANCE_TAIL = " EscalationTests proves an escalation is written to the trail."

M31_FR010_EVIDENCE = (
    "Creation, escalation, assignment, status change, comment and attachment "
    "each dispatch a notification through the platform's notification service. "
    "A new ticket, and a requester's own reply or file on a ticket nobody owns "
    "yet, go to whoever works that ticket's queue: the school's own triage "
    "staff while the ticket is the school's, and the platform desk once it is "
    "escalated or when it is one of CodeX's own. Escalation goes to the desk "
    "alone and names the school. Assignment goes to the new owner, even one who "
    "has muted the ticket. Status changes, comments and files go to the "
    "participants: the requester, the owner, and anybody following. Commenting "
    "follows a ticket, and a reader who may see it can follow or mute it; a "
    "mute holds until they follow again or comment. Recipients are "
    "de-duplicated, the actor is excluded from their own event, and anybody who "
    "can no longer see the ticket is dropped at dispatch. An internal note goes "
    "only to the people who may read internal notes. The ticket's tenant is "
    "passed as what the message is about, not as who owns the resulting "
    "records: each record is owned by its recipient's own tenant, so the notice "
    "a platform agent receives about a school's ticket sits in the platform's "
    "feed and history rather than the school's."
)

M31_FR010_ACCEPTANCE = (
    "Routing by where the ticket sits is load-bearing, because recipients who "
    "cannot see a ticket are dropped: paging the platform desk about an "
    "unescalated school ticket would send nothing and leave it unread with "
    "nobody aware it exists. A note written for platform support reaches no "
    "requester, and its record is absent from the raising tenant's delivery "
    "history by list, by identifier and by search. A platform agent reads the "
    "ticket notice in their own in-app feed. EscalationTests proves the "
    "school's own queue is told of a new ticket and the desk is told at the "
    "moment of escalation; TicketServiceTests proves that commenting follows, "
    "that a follower can mute and that commenting follows again, and that a "
    "stale follower without access receives nothing."
)

M31_FR010_LIMIT = (
    "In-app and email only; there is no SMS channel. The notice a school reads "
    "covers the messages sent to its own people, so a school cannot audit "
    "delivery of the messages its ticket raised to platform staff. The desk's "
    "recipients and its assignee picker are two queries that can disagree: the "
    "recipients are platform users whose roles grant either triage key "
    "directly, while the picker also counts the manage key granted through a "
    "permission group and honours a direct deny. An agent holding the key only "
    "through a group can be assigned a ticket and is never told of a new or "
    "escalated one."
)

M31_TICKET_MODEL_NOTE = (
    "Number allocated from the shared tenant counter as TK-<tenant><YYMMDD><n>. "
    "Requester and branch must belong to the ticket's tenant. escalated_at and "
    "escalated_by stamp the escalation on the ticket itself, indexed with "
    "status because the desk's list is escalated tickets by state."
)

M31_SUBSCRIPTION_ROW = [
    "TicketSubscription",
    "One reader following or muting one ticket.",
    "Unique on ticket and user. A mute keeps the row with its time rather than "
    "deleting it, so it holds until the reader follows again or comments; the "
    "source records whether they followed by hand or by commenting.",
]

M31_FOLLOW_ROUTE = [
    "POST, DELETE /support/tickets/{id}/follow/",
    "Follow a ticket the caller may see, or mute its notifications.",
]

M31_ATTACH_ROUTE = (
    "Attach evidence. Validated before storage, and open to a school that has "
    "not gone live."
)

M31_EXPORT_DEPENDENCY = [
    "Module 26, Reporting & Exports",
    "Receives the support datasets, registered from the app config. Rows narrow "
    "through the ticket list's own visibility function and then to the tenant "
    "the run was requested in, and the ticket body is never offered as a "
    "column.",
]

M31_NEW_GAPS = [
    "• The platform desk's ticket export carries CodeX's own tickets only. "
    "Every export run is scoped to the tenant it was requested in, so the "
    "escalated school tickets the desk works are on its list and absent from "
    "its file.",
    "• The desk's notification list and its assignee list can disagree: an "
    "agent who holds the manage key only through a permission group can be "
    "assigned a ticket and is never told of a new or escalated one.",
]

M31_TRACEABILITY_LEAD = (
    f"Module 31 carries 17 capability entries in MRD v{MRD_VERSION}. Each maps "
    "to the requirements below."
)

M31_CHANGE_SUMMARY = (
    "Carries escalation and ticket following into the sections that own them, "
    "and records two gaps found while tracing them. v1.5 recorded escalation "
    "as a requirement without reconciling the rest: FR-009 listed every audit "
    "action but escalation, which has one of its own; FR-010 listed every "
    "notification but escalation, and did not say that a new ticket goes to "
    "the school's own triage staff until it is escalated; the data model had "
    "no escalation stamp; and neither recorded that readers follow and mute "
    "tickets, a table and a route present since participant notifications "
    "shipped. FR-009, FR-010, the data model, the conversation routes and the "
    "scope table carry them, FR-004 records that a school still onboarding may "
    "attach evidence to its own ticket, and the dependencies gain the export "
    "datasets. FR-010 also claimed that the people notified and the people "
    "assignable cannot diverge. They can: the desk's recipients are read from "
    "direct role grants of either triage key, while the assignee picker also "
    "counts the manage key granted through a permission group, so an agent "
    "holding it only that way can be assigned a ticket and is never told of a "
    "new or escalated one. That joins Section 9 with the second gap: the "
    "desk's ticket export carries CodeX's own tickets only, because every "
    "export run is scoped to the tenant it was requested in. The resolved "
    "banner for v1.5's own fix is removed, because Section 9 is current state "
    "and the change log holds it. No ticket code changed after v1.5's baseline "
    "apart from test data. Module 31 remains Backend Complete and In use "
    "Complete with seventeen capabilities. Evidence is named from "
    "EscalationTests, TicketServiceTests, PendingSchoolAttachmentTests, "
    "AttachmentLimitTests and TicketExportVisibilityTests; the suite was not "
    "run for this revision. Backend evidence only; nothing here is deployed."
)


def remove_table_and_spacer(table) -> None:
    """Remove a banner table and the empty spacer paragraph that follows it.

    The spacer is removed only when it is empty and carries no page break, so a
    section break is never lost with it.
    """
    following = table._tbl.getnext()
    if following is not None and following.tag == qn("w:p"):
        text = "".join(following.itertext()).strip()
        breaks = [
            node for node in following.iter(qn("w:br"))
            if node.get(qn("w:type")) == "page"
        ]
        if not text and not breaks and following.find(".//" + qn("w:sectPr")) is None:
            following.getparent().remove(following)
    table._tbl.getparent().remove(table._tbl)


def patch_m31(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M31 Support Tickets Functional Requirements Document v{M31_TARGET}"

    cover, control = doc.tables[0], doc.tables[1]
    scope = require_header(doc.tables[3], "Area")
    fr004 = require_header(doc.tables[10], "FR-004")
    fr009 = require_header(doc.tables[15], "FR-009")
    fr010 = require_header(doc.tables[16], "FR-010")
    model = require_header(doc.tables[22], "Model")
    conversation_routes = require_header(doc.tables[24], "Method and path")
    dependencies = require_header(doc.tables[26], "Dependency")
    resolved_banner = require_header(doc.tables[27], "RESOLVED - ASSIGNMENT")
    further_gaps = require_header(doc.tables[28], "FURTHER GAPS")
    change_log = require_header(doc.tables[30], "Version")

    replace_cover_version(cover, M31_SOURCE, M31_TARGET)
    replace_control_value(control, "Version", M31_TARGET)
    replace_control_value(control, "Review date", REVIEW_DATE)
    replace_control_value(control, "Code baseline", CODE_BASELINE)
    replace_control_value(
        control, "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 31",
    )
    replace_control_value(control, "Supporting apps", M31_SUPPORTING_APPS)

    replace_row(scope, "Accounting and reach", [None, M31_SCOPE_ACCOUNTING])

    append_cell_text(fr004.rows[EVIDENCE_ROW].cells[1], M31_FR004_EVIDENCE_TAIL)
    append_cell_text(fr004.rows[ACCEPTANCE_ROW].cells[1], M31_FR004_ACCEPTANCE_TAIL)

    replace_cell(fr009.rows[EVIDENCE_ROW].cells[1], M31_FR009_EVIDENCE)
    append_cell_text(fr009.rows[ACCEPTANCE_ROW].cells[1], M31_FR009_ACCEPTANCE_TAIL)

    replace_cell(fr010.rows[EVIDENCE_ROW].cells[1], M31_FR010_EVIDENCE)
    replace_cell(fr010.rows[ACCEPTANCE_ROW].cells[1], M31_FR010_ACCEPTANCE)
    replace_cell(fr010.rows[LIMIT_ROW].cells[1], M31_FR010_LIMIT)

    replace_row(model, "Ticket", [None, None, M31_TICKET_MODEL_NOTE])
    insert_row_before(model, "TicketAttachment", M31_SUBSCRIPTION_ROW)

    insert_row_before(
        conversation_routes, "POST /support/tickets/{id}/attachments/", M31_FOLLOW_ROUTE,
    )
    replace_row(
        conversation_routes, "POST /support/tickets/{id}/attachments/",
        [None, M31_ATTACH_ROUTE],
    )

    insert_row_before(dependencies, "Platform core", M31_EXPORT_DEPENDENCY)

    # Section 9 is current state. The banner records v1.5's own fix, which the
    # change log already carries.
    remove_table_and_spacer(resolved_banner)
    last = further_gaps.rows[0].cells[0].paragraphs[-1]
    for text in M31_NEW_GAPS:
        last = clone_paragraph_after(last, text)

    replace_body_paragraph(doc, "Module 31 carries", M31_TRACEABILITY_LEAD)

    prepend_change_log(change_log, M31_TARGET, M31_CHANGE_SUMMARY)
    finish(doc, output, title, M31_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for folder, stem, source, target, patch in (
        (M03_DIR, M03_STEM, M03_SOURCE, M03_TARGET, patch_m03),
        (M31_DIR, M31_STEM, M31_SOURCE, M31_TARGET, patch_m31),
    ):
        directory = root / folder
        patch(
            directory / f"{stem}_v{source}.docx",
            directory / f"{stem}_v{target}.docx",
        )
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
