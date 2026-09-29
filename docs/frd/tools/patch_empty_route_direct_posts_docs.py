#!/usr/bin/env python3
"""Cut M07 v1.22: journals, expense claims and direct entries honour the empty-route rule.

What changed in the backend, and therefore in the document:

* The /post/ routes for journals and expense claims call finance's
  guard_direct_post, as refunds, write-offs, credit and debit notes and
  concessions already did (d7a700a8). A route with live steps answers 400 and
  sends the caller to submit; a route with none answers 409
  APPROVAL_NOT_CONFIGURED and writes nothing until confirm_without_approval is
  sent, and the post is then recorded as POSTED_WITHOUT_APPROVAL; no route at
  all posts as before.
* POST /finance/direct-entries/ is a journal, so the school's journal route
  governs it (88873c1c). A route with steps does not refuse it: the entry is
  created and submitted into the route. An empty route answers 409 until
  confirmed, and no route posts directly.

FR-031, the Modules 17 to 24 dependency and traceability carry it. The MRD was
checked and left at v2.93: the fail-closed and finance-handler entries already
name the behaviour, and no capability is added.

    python tools/patch_empty_route_direct_posts_docs.py
"""
from __future__ import annotations

from docx import Document

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_procurement_dashboards_and_routes_docs import spread_cover_lines
from patch_record_history_docs import (
    ROOT,
    finish,
    frd_path,
    log_change,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    row_labelled,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_value,
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    set_value,
)

import patch_record_history_docs

REVIEW_DATE = "28 September 2026"
CODE_BASELINE = (
    "Backend main at 27f1da0f, 28 September 2026, holding the finance pass merged at "
    "32194627: journal and expense-claim posts confirm an empty route (d7a700a8), and a "
    "direct entry goes through the school's journal route (88873c1c)"
)
TEST_EVIDENCE = (
    "Traced from the code at 27f1da0f; JournalDirectPostTests, ExpenseClaimDirectPostTests "
    "and DirectEntryApprovalRouteTests in vs_finance cover the three cases, and were not "
    "re-run for this revision. Nothing here claims deployment."
)

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

M07_DIR = "07-workflow-and-approval-engine"
M07_STEM = "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document"
M07_SOURCE, M07_TARGET = "1.21", "1.22"

FR031_EVIDENCE = (
    "submit_for_approval raises ApprovalNotConfiguredError (409, APPROVAL_NOT_CONFIGURED) when "
    "the resolved template has no live stages, after rolling back the pending flip so a "
    "refused submit writes nothing. A retired stage is history rather than configuration, so "
    "a template holding only retired stages is refused the same way instead of being routed "
    "past each of them to an APPROVED outcome. A caller may retry with an explicit "
    "confirmation, which terminates the instance as approved and records who confirmed and "
    "why through the POSTED_WITHOUT_APPROVAL audit action. A document posted without ever "
    "being submitted meets the same live-stage rule at the finance boundary, through "
    "approval_unconfigured: guard_direct_post stands in front of the direct post of every "
    "finance document type with an approval route, journals, expense claims, refunds, "
    "write-offs, credit and debit notes and concessions. A route with a live step answers 400 "
    "and sends the caller to submit; a route with none answers 409 APPROVAL_NOT_CONFIGURED "
    "and writes nothing until confirm_without_approval is sent, when the post goes through "
    "and is recorded as POSTED_WITHOUT_APPROVAL against the confirmer and their reason; no "
    "route at all posts directly. A direct entry, POST /finance/direct-entries/, is a journal "
    "and meets the journal route in the same three ways, except that a route with steps "
    "takes it rather than refusing it: the entry is created and submitted into the route, "
    "answering 201 with Direct entry <no> is waiting for approval. It reaches the books once "
    "it is approved., and the same approval block the journal submit returns."
)

M07_SUMMARY = (
    "Minor revision. Journals, expense claims and direct entries honour the empty-route "
    "rule. The journal and expense-claim /post/ routes go through finance's direct-post "
    "guard, as refunds, write-offs, credit and debit notes and concessions already did: a "
    "route with steps answers 400 and sends the caller to submit, an empty route answers 409 "
    "APPROVAL_NOT_CONFIGURED and writes nothing until confirm_without_approval is sent and "
    "then records the post as made without approval, and no route posts directly. Every "
    "school's books publish an empty expense-claim route, so each claim posted directly had "
    "skipped both the question and the record. A direct entry is a journal and meets the "
    "journal route the same way, except that a route with steps takes it: the entry is "
    "created and submitted into the route rather than posted. FR-031, the Modules 17 to 24 "
    "dependency and traceability follow; the MRD was checked and stays at v2.93. The cover's "
    "four lines, which had all printed at title size over the year, each take their own size "
    "again, and every table row is kept whole on its page. "
    + TEST_EVIDENCE
)


def patch_m07() -> None:
    require_newest(str(ROOT / "functional-requirements" / M07_DIR / f"{M07_STEM}_v*.docx"),
                   M07_SOURCE)
    doc = Document(str(frd_path(M07_DIR, M07_STEM, M07_SOURCE)))
    set_cover_version(doc, M07_SOURCE, M07_TARGET)
    spread_cover_lines(doc)
    if f"Version: {M07_TARGET}" not in doc.tables[0].rows[0].cells[0].text:
        raise ValueError("The cover does not carry the new version")
    set_control(doc, "Version", M07_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)

    # 4. Requirements
    fr031 = fr_table(doc, "FR-031")
    set_value(fr031, "Current evidence", FR031_EVIDENCE)
    append_to(fr031, "Acceptance",
              " At the finance boundary, JournalDirectPostTests and ExpenseClaimDirectPostTests "
              "prove an empty route refuses until confirmed and writes nothing, a confirmed "
              "post goes through and is recorded against the confirmer, and a route with steps "
              "still sends the caller to submit; DirectEntryApprovalRouteTests proves a direct "
              "entry waits for approval under a route with steps, is refused and then "
              "confirmable under an empty one, and posts when there is no route.")

    # 8. Dependencies
    handlers = row_labelled(doc, "Modules 17 to 24, finance and procurement").cells[-1]
    edit_cell(handlers, "Procurement's central ladder is published by this engine's publish "
                        "service.",
              "Procurement's central ladder is published by this engine's publish service. "
              "Finance holds the empty-route rule of FR-031 at its direct posts, journals, "
              "expense claims, refunds, write-offs, credit and debit notes and concessions, "
              "and submits a direct entry into the journal route where that route has steps.")

    # 10. Traceability
    paragraph = paragraph_starting(doc, "Module 7 carries 28 capability entries in MRD")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Journals, expense claims and direct entries honouring the empty-route rule "
        "strengthen the fail-closed and finance workflow handler entries without changing the "
        "count."))
    trace = table_headed(doc, "MRD capability", "Requirements")
    edit_value(trace, "Finance workflow handlers", "FR-016, FR-019", "FR-016, FR-019, FR-031")

    log_change(doc, M07_TARGET, M07_SUMMARY)
    retired_site_word = "cam" + "pus"
    assert_absent_outside_log(doc, retired_site_word, retired_site_word.capitalize(),
                              "The finance direct-post gate applies the same live-stage rule")
    # Every row below the cover stays whole on its page.
    for table in doc.tables[1:]:
        mrd_tools.keep_rows_whole(table)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M07_DIR, M07_STEM, M07_TARGET),
           f"{M07_STEM.replace('_', ' ')} v{M07_TARGET}", M07_TARGET)


def main() -> None:
    patch_m07()


if __name__ == "__main__":
    main()
