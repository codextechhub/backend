#!/usr/bin/env python3
"""Version the MRD and two FRDs for selling depth rather than whole modules.

A school used to be sold a list of modules and given a plan that decided
nothing. This records the change that joins the two: every school holds every
module, the plan it pays for decides how far into each one it reaches, and the
subscription date that was already on the row now ends the grants it pays for.

    python tools/patch_capability_depth_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from generate_requirements_documents import (
    BLACK,
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)

REVIEW_DATE = "6 September 2026"
CHANGE_DATE = "6 Sep 2026"

MRD_SOURCE_VERSION = "2.65"
MRD_TARGET_VERSION = "2.66"
M01_SOURCE_VERSION = "1.18"
M01_TARGET_VERSION = "1.19"
M04_SOURCE_VERSION = "1.13"
M04_TARGET_VERSION = "1.14"

MRD_SOURCE_SCOPE = (
    "Backend change selling depth rather than whole modules. Every school is "
    "granted every module at the depth its plan reaches, a plan sets no size "
    "ceiling, and a subscription expiry now ends the grants it pays for "
    "(6 September 2026)"
)

M01_SOURCE_SCOPE = (
    "Backend change replacing the picked-module package step with a plan that "
    "sets one depth for every module, removing every plan size ceiling, and "
    "writing the subscription expiry onto each grant (6 September 2026)"
)

M04_SOURCE_SCOPE = (
    "Backend change recording on each permission the capability that governs "
    "it, and refusing a request the role allows but the plan does not reach "
    "(6 September 2026)"
)

MRD_CHANGE_SUMMARY = (
    "Joined the plan a school pays for to the product it gets. A school was "
    "sold a tier on one step and hand-picked modules on another, and nothing "
    "reconciled them, so a school on the cheapest tier could be ticked into "
    "everything and a school never saw a module nobody thought to sell it. "
    "Every school is now granted every module, and the plan sets how far into "
    "each one it reaches. Depth is a field on the capability row with a "
    "module above it, so the catalogue is two levels deep and moving a "
    "feature between bands is one edit that every school sees on its next "
    "request. A grant naming no depth still reaches every band, which is what "
    "every grant written before depth existed means. An uplift given for a "
    "deal is a row of its own with its own dates, so it expires back to the "
    "tier rather than taking the module away. Three defects closed with it: "
    "entitlements gated nothing at all, because the evaluation was called in "
    "one place to decorate a screen and no route ever refused a request; the "
    "subscription expiry was stored and never written onto a grant, so a "
    "school that stopped paying kept everything switched on; and the plan "
    "size ceilings are removed, three of which were read by nothing while the "
    "branch ceiling refused a proprietor the site she had planned for months. "
    "Size is priced rather than fenced. Nothing yet prices it: there is no "
    "rate, no floor and no headcount snapshot, which is recorded as a new "
    "priority gap. Verified by 94 configuration tests, 466 RBAC tests, 301 "
    "school tests and 135 core tests. Backend evidence only; nothing here is "
    "deployed."
)

M01_CHANGE_SUMMARY = (
    "Replaces the picked-module package step with a plan that decides depth. "
    "School creation granted whatever the wizard ticked, which made the tier "
    "and the product two unrelated facts about the same school. The plan now "
    "grants every module at its own depth, through one service both creation "
    "and a later plan change pass through, and the module list a frontend "
    "still sends is accepted and narrows nothing. Every grant carries the "
    "subscription expiry, exclusive and rounded to the end of the paid day, "
    "so a school paid to the 31st still works on the 31st and a school that "
    "stopped paying does not keep the product. Every plan size ceiling is "
    "removed: the student, teacher and administrator caps were stored and "
    "read by nothing, and the branch ceiling was enforced on both creation "
    "paths, so the cheapest tier refused a proprietor a fourth site. A "
    "management command applies a school's existing plan without disturbing "
    "an uplift given for a deal. Verified by 301 school tests including 31 "
    "focused package tests. Backend evidence only; nothing here is deployed."
)

M04_CHANGE_SUMMARY = (
    "Records on each permission the capability that governs it, and adds the "
    "second gate that reads it. A hardcoded map joined the two vocabularies "
    "at module and resource level only, so it could say Finance and never the "
    "bulk generation inside Finance, which is the granularity depth is sold "
    "at. The permission row now names its capability, the map remains the "
    "fallback for keys nobody has classified, and a null means the key is "
    "core and every school holds it. The gate runs where every route's "
    "permission key is already checked, after the role question and never "
    "before it, and refuses with a code of its own naming the module, the "
    "depth needed and the depth held, so a refusal only a proprietor can act "
    "on does not read like one a school administrator can fix. It is off "
    "until the permission-to-band map is complete, settable at platform scope "
    "only so a school cannot disable its own gate, and a school with no "
    "package grants at all is treated as unprovisioned and never refused. "
    "Verified by 466 RBAC tests including 20 focused gate tests. Backend "
    "evidence only; nothing here is deployed."
)

# --- new capability entries -------------------------------------------------

M06_NEW_CAPABILITIES = [
    ("▸  Module and band catalogue with ordered depth",
     "▸  Depth-limited capability entitlements"),
    ("▸  Time-boxed tenant depth grants",
     "▸  Depth-aware effective-capability evaluation"),
]

M04_NEW_CAPABILITIES = [
    ("▸  Permission mapped to the capability that governs it",
     "▸  Plan depth gate with its own refusal"),
]

M01_PACKAGE_CAPABILITY = "▸  Package plan setting one depth for every module"
M01_EXTRA_CAPABILITY = "▸  Subscription expiry written onto every capability grant"

M06_DECISION = (
    "\n• A module is sold whole and used in slices. The catalogue is two "
    "levels deep: a module is what a school holds, and a band is that module "
    "cut at Core, Plus or Advanced. A band is never granted on its own; it "
    "answers when the depth on its module's grant reaches it, so moving a "
    "feature between bands is one field on one row and every school sees it "
    "on the next request, with nothing rewritten per tenant. A grant that "
    "names no depth reaches every band, which is what every grant written "
    "before depth existed means and why a missing depth cannot be read as "
    "the shallowest one."
    "\n• An uplift given for a deal is a row of its own, with its own dates. "
    "Ending it on the entitlement would end the grant rather than the uplift, "
    "taking the whole module away instead of returning the school to the tier "
    "it actually pays for. It only ever adds, so a school that upgrades past "
    "an old uplift is not demoted by a row nobody deleted."
)

M04_DECISION = (
    "\n• A permission now names the capability that governs it, and the plan "
    "is asked after the role, never before it. The hardcoded map could answer "
    "only at module and resource level, so it could say Finance and never the "
    "bulk generation inside it; it remains the fallback for keys nobody has "
    "classified, and a null means the key is core. The gate refuses with a "
    "code of its own naming the module, the depth needed and the depth held, "
    "because a refusal only a proprietor can act on must not read like one a "
    "school administrator can fix. It is off until the permission-to-band map "
    "is complete, settable at platform scope only so a school cannot disable "
    "its own gate, and a school with no package grants at all is unprovisioned "
    "rather than unentitled and is never refused."
)

M01_ATTENTION = (
    "\n• A plan no longer caps a school's size. The student, teacher and "
    "administrator ceilings were stored and read by nothing, and the branch "
    "ceiling was enforced on both creation paths, so the cheapest tier "
    "refused a proprietor the fourth site she had planned for months. All "
    "four are removed and size is priced rather than fenced, which leaves the "
    "depth wall as the only thing that moves a school up a tier."
    "\n• Nothing yet prices it. The plan carries no rate, no floor and no "
    "billing unit, and no term headcount is taken from the student register, "
    "so what a school is charged is still decided outside the platform."
)

M01_DESCRIPTION = (
    "Shared-row multi-tenant school and branch administration. The current "
    "backend supports creation, listing, editing, statistics, package setup, "
    "primary administrators, and Branch lifecycle operations. A required "
    "administrator is provisioned inside the same transaction as the School "
    "or standalone Branch, and the invitation email is queued only once that "
    "transaction commits. School status mirroring now calls one transactional "
    "Tenant transition service, which permanently ends active impersonation "
    "sessions whenever the Tenant leaves ACTIVE and cannot race a new session "
    "into existence. Package setup grants every module at the depth the plan "
    "reaches and writes the subscription expiry onto each grant; no plan caps "
    "how many students, staff, administrators or branches a school may have. "
    "Commercial lifecycle controls, subscription pricing, normal "
    "login-session revocation, safe deletion, and closed-Branch access rules "
    "remain open."
)

M06_DESCRIPTION = (
    "Typed platform, school, and branch configuration plus capability "
    "entitlements. The capability catalogue is two levels deep: a module is "
    "what a school is sold, and a band is that module cut at an ordered "
    "depth. A grant carries the depth it reaches, a time-boxed uplift can "
    "carry one school deeper for a while, and evaluation answers for a band "
    "by asking its module. Current uncommitted platform, security, "
    "integration, mail, and onboarding-default work is counted as complete."
)

M04_DESCRIPTION = (
    "Two-layer platform and school authorization with reusable permission "
    "definitions, roles, templates, assignments, and controlled overrides. "
    "Platform-only scope is enforced independently from restricted "
    "sensitivity. The registry requires that scope when a key is created, and "
    "its module, resource and action identity cannot be renamed through a "
    "supported update. Each permission records the product capability that "
    "governs it, and a second gate refuses a request the caller's role allows "
    "but the school's plan does not reach, with a refusal of its own. "
    "Restricted keys cannot enter direct role writes, groups, or ALLOW "
    "overrides; they enter a custom role only through a change request "
    "decided by a different reviewer who already holds every restricted key "
    "being granted. Branch reach remains grant-derived, and one person may "
    "hold the same role at more than one branch. Every supported role "
    "permission or group change records its actor, reason, source, approval "
    "reference where applicable, and complete before-and-after access "
    "configuration in the same transaction as the mutation."
)

DELTA_ROWS = [
    ["Depth in the catalogue", "Built",
     "Capability gains a parent and an ordered depth, with a check constraint "
     "keeping the catalogue two levels deep. Eleven modules and three bands "
     "each are seeded, and rebanding a feature is one field on one row."],
    ["What a tier grants", "Changed",
     "A plan sets one depth for every module, with exceptions held as rows. "
     "Onboarding no longer picks modules: every school holds every module and "
     "the plan decides how far in."],
    ["A deal that ends by itself", "Built",
     "A time-boxed depth uplift is a row of its own, so it expires back to "
     "the tier the school pays for instead of taking the module away."],
    ["Subscription expiry", "Enforced",
     "The grant carries the paid-to date, exclusive and rounded to the end of "
     "that day. It was stored and never written onto a grant, so a school "
     "that stopped paying kept everything switched on."],
    ["Plan size ceilings", "Removed",
     "Every plan capacity column is gone, with the enforced branch refusal. "
     "Three were read by nothing; the branch ceiling refused a planned site. "
     "Size is priced rather than fenced."],
    ["The plan gate", "Built, switched off",
     "A refusal of its own naming the module, the depth needed and the depth "
     "held, at the one place every route's permission key is already checked. "
     "Off until the permission-to-band map is complete."],
    ["What a school pays", "Not built",
     "No rate, no floor, no billing unit and no term headcount. Depth decides "
     "what a school reaches; nothing yet decides its invoice."],
]

PRIORITY_GAP_ROW = [
    "P1",
    "Subscription pricing and metering",
    "Put a rate, a floor and a billing unit on the plan, snapshot each term's "
    "headcount from the student register, and invoice the school through the "
    "platform ledger entity. Depth now decides what a school reaches and "
    "nothing decides what it pays, and with the size ceilings removed the "
    "headcount is the only remaining way a larger school pays more.",
    "M1, M6, M17",
]


def retitle(paragraph, text: str) -> None:
    """Rewrite a paragraph's words while leaving its formatting alone.

    Rebuilding the paragraph would drop the style it inherits, which is how a
    cloned heading ends up a different colour and size from the siblings it
    was copied from.
    """
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def update_control_table(table, *, version, source_scope, entries=None,
                         mrd_baseline=None):
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], source_scope, size=9)
        elif label == "Capability entries" and entries:
            replace_cell(row.cells[1], entries, size=9)
        elif label == "MRD baseline" and mrd_baseline:
            replace_cell(row.cells[1], mrd_baseline, size=9)


def prepend_change_log(table, version, date, summary) -> None:
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], date, size=8)
    replace_cell(row.cells[2], summary, size=8)


def append_capability_rows(table, pairs, *, size=8.5) -> None:
    """Add capability rows by cloning the shape an existing row already has."""
    template = table.rows[0]
    anchor = table.rows[-1]
    for left, right in pairs:
        new_tr = copy.deepcopy(template._tr)
        anchor._tr.addprevious(new_tr)
        row = table.rows[len(table.rows) - 2]
        replace_cell(row.cells[0], left, size=size)
        replace_cell(row.cells[1], right, size=size)


def append_to_cell(cell, addition, *, size) -> None:
    replace_cell(cell, cell.text.rstrip() + addition, size=size)


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = (
        f"XVS Module Requirements Document v{MRD_TARGET_VERSION}"
    )
    doc.core_properties.version = MRD_TARGET_VERSION

    replace_cover_version(doc.tables[0], MRD_SOURCE_VERSION, MRD_TARGET_VERSION)
    update_control_table(
        doc.tables[1],
        version=MRD_TARGET_VERSION,
        source_scope=MRD_SOURCE_SCOPE,
        entries="486",
    )

    # Contents: the delta section is named after the revision it belongs to.
    for row in doc.tables[2].rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(
                row.cells[0], f"5. v{MRD_TARGET_VERSION} Capability Delta",
                size=9, bold=True, color=BLUE,
            )
            replace_cell(
                row.cells[1],
                "Depth sold instead of whole modules, and the plan that sets it",
                size=9,
            )

    # Module index counts.
    index = doc.tables[5]
    for row in index.rows:
        number = row.cells[0].text.strip()
        if number == "1":
            replace_cell(row.cells[5], "19", size=8.5)
        elif number == "4":
            replace_cell(row.cells[5], "20", size=8.5)
        elif number == "6":
            replace_cell(row.cells[5], "24", size=8.5)

    # Module descriptions and the delta heading.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Shared-row multi-tenant school and branch"):
            write_paragraph(paragraph, M01_DESCRIPTION, size=9, space_after=5)
        elif text.startswith("Two-layer platform and school authorization"):
            write_paragraph(paragraph, M04_DESCRIPTION, size=9, space_after=5)
        elif text.startswith("Typed platform, school, and branch configuration"):
            write_paragraph(paragraph, M06_DESCRIPTION, size=9, space_after=5)
        elif text == f"5. v{MRD_SOURCE_VERSION} Capability Delta":
            write_paragraph(
                paragraph, f"5. v{MRD_TARGET_VERSION} Capability Delta",
                size=17, bold=True, space_before=15, space_after=8,
            )
        elif text.startswith("This revision records one module arriving"):
            write_paragraph(
                paragraph,
                "This revision records the plan a school pays for deciding "
                "what that school gets. Depth arrives in the capability "
                "catalogue, the tier sets it, the subscription date ends it, "
                "and the size ceilings that fenced a school rather than "
                "pricing it are removed. What a school is charged is not here "
                "and is recorded as a gap.",
                size=9, space_after=5,
            )

    # Module 1: the package entry, one new entry, and the attention note.
    m01 = doc.tables[8]
    replace_cell(m01.rows[2].cells[0], M01_PACKAGE_CAPABILITY, size=8.5)
    append_to_cell(
        m01.rows[7].cells[1], "\n" + M01_EXTRA_CAPABILITY, size=8.5,
    )
    append_to_cell(m01.rows[8].cells[0], M01_ATTENTION, size=8.2)

    # Module 4: one new row, then the decision note.
    m04 = doc.tables[14]
    append_capability_rows(m04, M04_NEW_CAPABILITIES)
    append_to_cell(m04.rows[-1].cells[0], M04_DECISION, size=8.2)

    # Module 6: two new rows, then the decision note.
    m06 = doc.tables[18]
    append_capability_rows(m06, M06_NEW_CAPABILITIES)
    append_to_cell(m06.rows[-1].cells[0], M06_DECISION, size=8.2)

    # Priority gaps: pricing is a new material gap, not an optional idea.
    gaps = doc.tables[75]
    template = gaps.rows[-1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addnext(new_tr)
    row = gaps.rows[-1]
    # The priority and the gap name carry the emphasis every sibling row has;
    # writing them plain leaves one row looking like an afterthought.
    for idx, value in enumerate(PRIORITY_GAP_ROW):
        emphasised = idx < 2
        replace_cell(
            row.cells[idx], value, size=8.2,
            bold=emphasised, color=BLUE if emphasised else BLACK,
        )

    rebuild_table(
        doc.tables[76],
        [f"v{MRD_TARGET_VERSION} capability delta", "Decision", "Evidence"],
        DELTA_ROWS,
        [1.75, 0.95, 4.57],
        font_size=8.2,
    )

    prepend_change_log(
        doc.tables[78], MRD_TARGET_VERSION, CHANGE_DATE, MRD_CHANGE_SUMMARY,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(
        output, f"XVS Module Requirements Document v{MRD_TARGET_VERSION}"
    )
    shrink_inherited_media(output)
    assert_no_em_dash(output)


M01_DECISION = (
    "\n• The plan decides the product, and no plan decides the size. School "
    "creation used to grant whatever list the wizard sent, so the tier a "
    "school paid for and the modules it held were two unrelated facts about "
    "the same school. One service now grants every module at the depth the "
    "plan reaches into it, and both the creation path and a later plan change "
    "call it, so the two cannot answer differently. Every grant carries the "
    "subscription expiry, which was recorded on the setup row and never "
    "written anywhere the evaluation could read it."
    "\n• Every plan capacity column is removed with the branch refusal that "
    "read one of them. Three were stored and read by nothing; the branch "
    "ceiling was enforced, so the cheapest tier refused a proprietor the "
    "fourth site she had planned for months. Size is priced rather than "
    "fenced, and nothing prices it yet."
)

FR007_ROWS = {
    "Requirement": (
        "School creation selects an active PackagePlan, records the "
        "subscription expiry, and grants every module at the depth that plan "
        "reaches. No plan caps how many students, staff, administrators or "
        "Branches a School may have."
    ),
    "Current evidence": (
        "apply_plan_entitlements grants every active, entitlement-gated module "
        "through vs_config's set_entitlement, carrying the depth "
        "PackagePlan.depth_for() returns for that module: the plan's "
        "default_depth, or the PackagePlanModuleDepth exception where one is "
        "named. School creation and the apply_plans command both call that one "
        "service, so a School created with a plan and a School moved onto "
        "another cannot end up with different grants. Each grant carries the "
        "subscription expiry as an exclusive ends_at rounded to the end of the "
        "paid day, which vs_config's evaluation already honoured and was never "
        "given. The enabled_modules list a client still sends is accepted and "
        "narrows nothing, so a frontend that has not caught up keeps working. "
        "PackagePlan carries no capacity columns, and branch_ceiling_refusal "
        "is removed from both paths by which a School gains a site."
    ),
    "Acceptance": (
        "Creating a School with a plan grants every sellable module, including "
        "one the client did not name, each at the plan's depth for that "
        "module. A module the plan names an exception for takes that depth and "
        "its siblings do not. A School paid to a date still works on that date "
        "and is refused the day after. A grant written before depth existed, "
        "carrying no depth, still reaches every band. No School is refused a "
        "student, a member of staff, an administrator or a Branch on account "
        "of its plan."
    ),
    "Current limit": (
        "The plan decides what a School reaches and nothing decides what it "
        "pays. PackagePlan carries no rate, no floor and no billing unit, no "
        "term headcount is taken from the student register, and no "
        "subscription invoice is raised, so the commercial half of the plan is "
        "still settled outside the platform."
    ),
}

M01_GAP_ROW = [
    "P1",
    "Nothing prices the plan",
    "Put a rate, a floor and a billing unit on PackagePlan, snapshot each "
    "term's headcount from the student register, and raise the subscription "
    "invoice through the platform ledger entity. With the capacity ceilings "
    "removed, the headcount is the only remaining way a larger School pays "
    "more.",
    "FR-007",
]

M01_TRACE_PACKAGE = [
    "Package plan setting one depth for every module",
    "FR-007",
    "Implemented; the plan grants every module at the depth it reaches into "
    "that module, exceptions included, and no plan caps a School's size.",
]

M01_TRACE_EXPIRY = [
    "Subscription expiry written onto every capability grant",
    "FR-007",
    "Implemented; the grant carries the paid-to date as an exclusive bound "
    "rounded to the end of that day, so a School paid to the 31st works on "
    "the 31st and not on the 1st.",
]

M04_FR025_TITLE = "FR-025  Refuse What the Role Allows and the Plan Does Not Reach"

#: The banner row is one merged cell carrying both halves, so it is written
#: once. Writing the two column indexes in turn replaces the same cell twice
#: and loses the requirement number.
M04_FR025_BANNER = "FR-025 | Implemented, disabled by default"

M04_FR025_ROWS = [
    ("Requirement",
     "A request whose permission key is governed by a product capability the "
     "school's plan does not reach is refused, and the refusal names the plan "
     "rather than the role."),
    ("Current evidence",
     "Permission.capability records the capability governing each key; null "
     "means the key is core and every school holds it, and "
     "vs_rbac.capability_map remains the fallback for keys nobody has "
     "classified. plan_gate.plan_refusal runs inside HasRBACPermission after "
     "the role check has passed and never before it, evaluates the capability "
     "through vs_config's effective_capability, and raises PlanUpgradeRequired "
     "naming the module, the depth needed and the depth held. The gate is "
     "governed by platform.entitlements.enforce, which is settable at platform "
     "scope only so a school holding config.value.update cannot switch off its "
     "own gate, and a tenant holding no PACKAGE entitlement at all is treated "
     "as unprovisioned rather than unentitled and is never refused."),
    ("Acceptance",
     "A caller whose role carries the key and whose school reaches the depth "
     "is served. A caller whose role carries the key and whose school does not "
     "reach the depth receives 403 PLAN_UPGRADE_REQUIRED naming the module and "
     "both depths. A caller whose role does not carry the key receives the "
     "ordinary role refusal under its own code, so the two are distinguishable "
     "by a client and by the person reading them. A school with no package "
     "grants is served either way."),
    ("Limit",
     "Switched off until every permission key carries a band, because a "
     "half-mapped catalogue refuses real work for no commercial reason. Only "
     "rbac_permission is gated: a view gated solely by rbac_group_permission "
     "or by HasAnyModuleAccess passes untouched. Enforcement is one platform "
     "switch rather than one per school."),
]

M04_REFUSAL_ROW = [
    "A key the caller's role allows but the school's plan does not reach",
    "403 PLAN_UPGRADE_REQUIRED, naming the module, the depth needed and the "
    "depth held. Deliberately not the role refusal: the person who can act on "
    "it is the proprietor, not the school administrator.",
]

M04_TRACE_ROWS = [
    ["Permission mapped to the capability that governs it", "FR-025",
     "Implemented. The permission row names its capability, a null means the "
     "key is core, and the hardcoded map is the fallback for keys nobody has "
     "classified yet."],
    ["Plan depth gate with its own refusal", "FR-025",
     "Implemented and switched off. The gate runs where every route's key is "
     "already checked and refuses under a code of its own; it waits on the "
     "permission-to-band map before it is enabled."],
]

M04_RECONCILIATION = (
    "MRD RECONCILIATION\n"
    "• MRD v2.66 lists Module 4 as Roles & Permissions (RBAC), Phase V1, "
    "Backend Complete, In use Complete, code vs_rbac, with twenty capability "
    "entries. The module number, name, phase, states and ownership agree with "
    "this revision.\n"
    "• The two entries added at this revision are the capability recorded on "
    "each permission and the plan depth gate that reads it. Both are traced to "
    "FR-025 above. The gate is built and disabled, which is why the module "
    "stays Complete: nothing it promises is missing, and what is switched off "
    "is a commercial rollout rather than a capability."
)


def insert_row_after(table, index, values, *, size=8.2, styled=()):
    """Clone a row's shape rather than adding a bare one, so borders survive."""
    template = table.rows[index]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addnext(new_tr)
    row = table.rows[index + 1]
    for idx, value in enumerate(values):
        emphasised = idx in styled
        replace_cell(
            row.cells[idx], value, size=size,
            bold=emphasised, color=BLUE if emphasised else BLACK,
        )
    return row


def delete_row(table, index) -> None:
    row = table.rows[index]
    row._tr.getparent().remove(row._tr)


def patch_m01(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = (
        "XVS M01 School and Branch Management Functional Requirements "
        f"Document v{M01_TARGET_VERSION}"
    )
    doc.core_properties.version = M01_TARGET_VERSION

    replace_cover_version(doc.tables[0], M01_SOURCE_VERSION, M01_TARGET_VERSION)
    update_control_table(
        doc.tables[1],
        version=M01_TARGET_VERSION,
        source_scope=M01_SOURCE_SCOPE,
        mrd_baseline=(
            f"XVS Module Requirements Document v{MRD_TARGET_VERSION}, Module 1, "
            "nineteen capability entries"
        ),
    )

    decision = doc.tables[6].rows[0].cells[0]
    replace_cell(
        decision,
        decision.text.rstrip().replace(
            "Module 1 remains Backend Partial and In use Complete in MRD v2.50.",
            f"Module 1 remains Backend Partial and In use Complete in MRD "
            f"v{MRD_TARGET_VERSION}, with nineteen capability entries.",
        ) + M01_DECISION,
        size=8.2,
    )

    # Prose that still describes the capacities and the ceiling. Left as found
    # it would read as a contract the code no longer keeps.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("FR-007  Assign Package Capacity"):
            retitle(paragraph, "FR-007  Grant Every Module at the Plan's Depth")
        elif text.startswith("• Package capacity, module entitlements"):
            retitle(
                paragraph,
                "• Package depth, module entitlements, dependency resolution, "
                "and scoped overrides used during onboarding.",
            )
        elif text.startswith("• A plan's Branch ceiling refused on both"):
            retitle(
                paragraph,
                "• A plan granting every module at its own depth, an exception "
                "overriding that depth for one module alone, the subscription "
                "expiry reaching every grant, and a School on the cheapest plan "
                "opening a fourth Branch and being refused none of it.",
            )

    workflow = doc.tables[30]
    replace_cell(
        workflow.rows[5].cells[1],
        "Create SchoolPackageSetup, then grant every sellable module at the "
        "depth the plan reaches into it, each grant carrying the subscription "
        "expiry.",
        size=8.2,
    )
    replace_cell(
        workflow.rows[5].cells[2],
        "Reject an invalid plan or expiry. Grant every module, at one depth "
        "per module, with no capacity to validate.",
        size=8.2,
    )

    fr007 = doc.tables[15]
    for row in fr007.rows:
        label = row.cells[0].text.strip()
        if label in FR007_ROWS:
            replace_cell(row.cells[1], FR007_ROWS[label], size=8.2)

    models = doc.tables[33]
    replace_cell(
        models.rows[5].cells[2],
        "Unique code/name; one default depth with per-module exceptions; no "
        "capacity columns; active catalogue flag",
        size=8.2,
    )
    replace_cell(models.rows[6].cells[1], "Selected plan and subscription expiry", size=8.2)
    insert_row_after(models, 6, [
        "PackagePlanModuleDepth",
        "One module a plan reaches differently from the rest",
        "Unique per (plan, capability key); holds only the exceptions, so a "
        "plan selling its own depth everywhere needs no rows at all",
    ])

    # The two package capacity refusals describe rules that no longer exist.
    refusals = doc.tables[38]
    for index in sorted(
        [
            i for i, row in enumerate(refusals.rows)
            if row.cells[0].text.strip() in ("Package capacity", "Package Branch capacity")
        ],
        reverse=True,
    ):
        delete_row(refusals, index)

    gaps = doc.tables[40]
    insert_row_after(gaps, len(gaps.rows) - 1, M01_GAP_ROW, styled=(0, 1))

    trace = doc.tables[42]
    for index, row in enumerate(trace.rows):
        if row.cells[0].text.strip().startswith("Package plan"):
            for idx, value in enumerate(M01_TRACE_PACKAGE):
                replace_cell(row.cells[idx], value, size=8.2)
            insert_row_after(trace, index, M01_TRACE_EXPIRY)
            break

    prepend_change_log(
        doc.tables[43], M01_TARGET_VERSION, CHANGE_DATE, M01_CHANGE_SUMMARY,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(
        output,
        "XVS M01 School and Branch Management Functional Requirements "
        f"Document v{M01_TARGET_VERSION}",
    )
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_m04(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = (
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{M04_TARGET_VERSION}"
    )
    doc.core_properties.version = M04_TARGET_VERSION

    replace_cover_version(doc.tables[0], M04_SOURCE_VERSION, M04_TARGET_VERSION)
    update_control_table(
        doc.tables[1],
        version=M04_TARGET_VERSION,
        source_scope=M04_SOURCE_SCOPE,
        mrd_baseline=(
            f"XVS Module Requirements Document v{MRD_TARGET_VERSION}, Module 4, "
            "twenty capability entries"
        ),
    )

    # Bound before anything is inserted. Adding the FR-025 table shifts every
    # later index by one, so a lookup made afterwards reads its neighbour.
    models = doc.tables[36]
    refusals = doc.tables[40]
    dependencies = doc.tables[41]
    trace = doc.tables[44]
    reconciliation = doc.tables[45]
    change_log = doc.tables[46]

    # FR-025 is cloned from FR-024, heading and table together, so it inherits
    # the shape every other requirement on the page already has.
    heading = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("FR-024"):
            heading = paragraph
            break
    fr024 = doc.tables[32]
    new_heading = copy.deepcopy(heading._p)
    fr024._tbl.addnext(new_heading)
    new_table = copy.deepcopy(fr024._tbl)
    new_heading.addnext(new_table)

    retitle(Paragraph(new_heading, heading._parent), M04_FR025_TITLE)
    fr025 = Table(new_table, fr024._parent)
    replace_cell(fr025.rows[0].cells[0], M04_FR025_BANNER, size=8.5, bold=True, color=BLUE)
    for index, (label, value) in enumerate(M04_FR025_ROWS, start=1):
        replace_cell(fr025.rows[index].cells[0], label, size=8.2, bold=True, color=BLUE)
        replace_cell(fr025.rows[index].cells[1], value, size=8.2)

    append_to_cell(
        models.rows[2].cells[2],
        " Also carries the product capability that governs the key, null when "
        "the key is core and every school holds it.",
        size=8.2,
    )

    insert_row_after(refusals, len(refusals.rows) - 1, M04_REFUSAL_ROW)

    for row in dependencies.rows:
        if row.cells[0].text.strip() == "vs_config":
            append_to_cell(
                row.cells[1],
                " It also supplies the capability catalogue and the depth a "
                "tenant reaches into each module, which the plan gate reads "
                "through effective_capability, and the enforcement switch that "
                "decides whether that gate answers at all.",
                size=8.2,
            )
            break

    anchor = len(trace.rows) - 1
    for values in M04_TRACE_ROWS:
        insert_row_after(trace, anchor, values)
        anchor += 1

    replace_cell(reconciliation.rows[0].cells[0], M04_RECONCILIATION, size=8.2)

    prepend_change_log(
        change_log, M04_TARGET_VERSION, CHANGE_DATE, M04_CHANGE_SUMMARY,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(
        output,
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{M04_TARGET_VERSION}",
    )
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root)

    mrd_dir = root / "module-requirements"
    patch_mrd(
        mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_SOURCE_VERSION}.docx",
        mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_TARGET_VERSION}.docx",
    )
    print(f"Wrote MRD v{MRD_TARGET_VERSION}")

    frd = root / "functional-requirements"
    m01 = frd / "01-school-and-branch-management"
    patch_m01(
        m01 / ("XVS_M01_School_and_Branch_Management_Functional_Requirements_"
               f"Document_v{M01_SOURCE_VERSION}.docx"),
        m01 / ("XVS_M01_School_and_Branch_Management_Functional_Requirements_"
               f"Document_v{M01_TARGET_VERSION}.docx"),
    )
    print(f"Wrote M01 FRD v{M01_TARGET_VERSION}")

    m04 = frd / "04-roles-and-permissions-rbac"
    patch_m04(
        m04 / ("XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_"
               f"Document_v{M04_SOURCE_VERSION}.docx"),
        m04 / ("XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_"
               f"Document_v{M04_TARGET_VERSION}.docx"),
    )
    print(f"Wrote M04 FRD v{M04_TARGET_VERSION}")


if __name__ == "__main__":
    main()
