#!/usr/bin/env python3
"""Cut MRD v2.88 and M04 FRD v1.27: a restricted permission waits for approval on every role.

What changed in the backend, and therefore in the documents:

* A role save never grants a restricted permission the role does not already
  hold directly, whoever holds the role and whether the role is new or old.
  The role serializer grants everything else and raises one role change
  request for the restricted keys in the same transaction; a key already
  waiting on the role is not asked for twice, and a ladder that cannot run
  refuses the whole save. The role detail lists what waits as
  ``pending_additions``.
* That replaces the rule that only an addition to the editor's own role needed
  approval. Creating a role with a restricted permission had been refused
  outright, with no way to raise the request from creation.
* ``RestrictedNeedsApprovalError`` and its 409 ``RESTRICTED_NEEDS_APPROVAL``
  are gone; nothing raises them.
* ``set_role_access`` keeps a role's explicit denies when its caller names no
  deny set, lifting only a deny on a key the call grants. The role editor's
  save and the approval of a change request both name none, so each deleted
  every deny on the role it touched.

M04 also cited MRD v2.85 with twenty-two capability entries in its traceability
paragraph, while its reconciliation box said v2.86 with twenty-three. Both now
name v2.88 and twenty-three. The MRD's Module 4 decision still described a
reviewer other than the requester, holding the key, as the one who decides a
restricted grant; the ladder replaced that and the bullet now says so.

    python tools/patch_restricted_approval_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_record_history_docs import (
    ROOT,
    change_log_table,
    finish,
    frd_path,
    keep_format,
    log_change,
    replace_cell,
    set_control,
    set_cover_version,
    update_reconciliation,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
    row_with_cell_starting,
    set_value,
    table_with_header,
)

import patch_record_history_docs

REVIEW_DATE, SHORT_DATE = "26 September 2026", "26 Sep 2026"
MRD_SOURCE, MRD_TARGET = "2.87", "2.88"
CODE_BASELINE = (
    "Backend main at 38b062ef, with school-fe main at 67288aa and console-fe main at "
    "b537fb9, 26 September 2026"
)
TEST_EVIDENCE = (
    "Verified by the vs_rbac suite (875 tests) and the core suite (160 tests), each run on "
    "its own and all passing; and driven against the running school app, where a role "
    "created with Delete calendar kept View calendar at once and listed Delete calendar as "
    "waiting for approval."
)

# log_change writes the module's own review date, read from this module.
patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


# ── editing helpers ──────────────────────────────────────────────────────────


def set_paragraph(doc, start: str, text: str) -> None:
    mrd_tools.retitle(paragraph_starting(doc, start), text)


def row_labelled(doc, label: str):
    """The one row in the document whose first cell reads ``label``."""
    hits = [r for t in doc.tables for r in t.rows if r.cells[0].text.strip() == label]
    if len(hits) != 1:
        raise ValueError(f"{label!r} labels {len(hits)} rows")
    return hits[0]


def edit_run_text(cell, old: str, new: str) -> None:
    """Swap a passage inside a one-run cell whose lines are breaks, keeping the breaks."""
    runs = [r for p in cell.paragraphs for r in p.runs if old in r.text]
    if len(runs) != 1 or runs[0].text.count(old) != 1:
        raise ValueError(f"{old[:60]!r} is not in exactly one run once")
    runs[0].text = runs[0].text.replace(old, new)


def assert_absent_outside_log(doc, *needles: str) -> None:
    """Refuse a stale claim anywhere but the change log, which records history."""
    log = change_log_table(doc)._tbl
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        if table._tbl is log:
            continue
        for r in table.rows:
            texts.extend(c.text for c in r.cells)
    blob = "\n".join(texts)
    for needle in needles:
        at = blob.find(needle)
        if at >= 0:
            context = blob[max(0, at - 160):at + 80].replace("\n", " / ")
            raise ValueError(f"{needle!r} is still in the document: ...{context}...")


# ── M04 Roles & Permissions (RBAC) ───────────────────────────────────────────

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.26", "1.27"

M04_FR003_EVIDENCE = (
    "Scope remains enforced on every grant model, including bulk writes. Restricted authority "
    "has one rule on the role path, whoever holds the role: a save never grants a restricted "
    "key the role does not already hold directly. TenantRoleTemplateDetailSerializer splits "
    "the submitted keys in _hold_back_restricted, passes the rest to set_role_access, and hands "
    "the restricted remainder to raise_role_change_request through _raise_for_approval in the "
    "same transaction, on creation and update alike. set_role_access keeps its own guard for "
    "callers that have not settled the question: it refuses a restricted key the role does not "
    "already hold directly unless the caller passes allow_restricted, which only an approved "
    "change request, prebuilt provisioning, role suggestions, the reach settlement and the "
    "seeders do. GroupPermission and TenantRoleGroup reject restricted membership; "
    "UserPermissionOverride rejects restricted ALLOW; and assignment, replace, account "
    "creation, and draft submission refuse a role whose restricted keys the assigner does not "
    "already hold (missing_restricted_grant_authority), whoever receives it."
)
M04_FR003_ACCEPTANCE = (
    "The role path, group path, prebuilt path, assignment path, override path, account "
    "creation path, and draft-submission path all enforce scope. A restricted key reaches a "
    "custom role by one route: a role change decided through the workflow ladder, which a role "
    "save raises on its own. Giving yourself a role that already carries a restricted key is "
    "closed by the assignment ceiling. test_create_role_holds_a_restricted_permission_for_"
    "approval, test_a_restricted_addition_to_a_role_the_actor_holds_waits_for_approval, "
    "test_a_restricted_addition_to_somebody_elses_role_waits_too, "
    "test_a_restricted_key_the_role_already_holds_survives_a_save, "
    "test_giving_yourself_that_role_is_still_refused_by_the_grant_ceiling and "
    "test_giving_it_to_somebody_else_is_refused_by_the_same_ceiling pin that, and "
    "test_field_access_restriction repeats the rule in a school with one branch and a school "
    "with two."
)

M04_FR017_REQUIREMENT = (
    "Editing what a role may do changes every holder. Ordinary additions and removals may be "
    "written directly by a role administrator. A restricted addition to any role, a new one "
    "included and whoever holds it, must not take effect on save: the save grants everything "
    "else and raises a request carrying the administrator's reason as its justification, and "
    "an approval ladder that records who decided it grants the restricted keys."
)
M04_FR017_EVIDENCE_TAIL = (
    " A role save raises the request itself: the role serializer holds back every restricted "
    "key the role does not already hold directly and passes them to raise_role_change_request "
    "in the save's own transaction, skipping a key already waiting on that role in a PENDING "
    "request, so saving again does not stack a second request. The role detail's "
    "pending_additions lists each restricted key waiting on the role with the request that "
    "carries it."
)
M04_FR017_ACCEPTANCE_TAIL = (
    " The role save's side is pinned by TenantRoleTemplateViewTests: "
    "test_create_role_holds_a_restricted_permission_for_approval, "
    "test_a_restricted_addition_to_somebody_elses_role_waits_too, "
    "test_saving_again_while_a_request_waits_does_not_raise_a_second and "
    "test_a_ladder_that_cannot_run_refuses_the_whole_create."
)
M04_FR017_LIMIT_TAIL = (
    " A ladder that cannot run refuses the whole role save, so a tenant whose role-change "
    "template cannot be resolved cannot save a role that adds a restricted key, even where its "
    "ordinary permissions would have saved alone. A key left out of a later save stays in its "
    "pending request: withdrawing the request, not unticking the box, is what stops it."
)

M04_FR019_DENIES_OLD = (
    "Supplying grants to set_role_access without a deny set clears the role's explicit denies, "
    "which is the service's stated backward-compatible rule; the direct update path and an "
    "approved role change supply grants alone, so each deletes any explicit deny on the role "
    "it touches, and the audit's before-and-after shows it going. The reach settlement is the "
    "exception: it names the role's existing denies in the same write, so a key a school "
    "denied on purpose is still denied after a revocation rather than becoming merely "
    "ungranted and returning through a group. The API exposes no deny field, so the denies at "
    "risk are those a migration or seeder wrote. "
)

M04_ROLES_ROUTE = "GET, POST /rbac/tenants/{slug}/roles/"
M04_ROLE_ROUTE = "GET, PATCH, DELETE /rbac/tenants/{slug}/roles/{key}/"

M04_RECONCILIATION = [
    ("replace", "• MRD v2.86 lists Module 4",
     f"• MRD v{MRD_TARGET} lists Module 4 as Roles & Permissions (RBAC), Phase V1, Backend "
     "Complete, In use Complete, code vs_rbac, with twenty-three capability entries. The "
     "module number, name, phase, states and ownership agree with this revision."),
    ("append", "• A role save holding restricted keys back for approval restates the existing "
     "create-and-update and role-change request entries rather than adding one."),
]

M04_SUMMARY = (
    "Minor revision. A restricted permission waits for approval on every role save, whoever "
    "holds the role: the save grants the rest and raises one role change request for the "
    "restricted keys in the same transaction, a key already waiting is not asked for twice, a "
    "ladder that cannot run refuses the whole save, and the role detail lists what waits in "
    "pending_additions. This replaces the rule that only an addition to the editor's own role "
    "needed approval, and creating a role with a restricted permission, which was refused "
    "outright with no way to raise the request, now creates the role and raises it. 409 "
    "RESTRICTED_NEEDS_APPROVAL is retired. set_role_access keeps a role's explicit denies when "
    "its caller names none, so saving a role or approving a request no longer deletes them. "
    "FR-003, FR-017 and FR-019, the actor table, the role contracts, verification, one Needs "
    "Attention row and traceability are updated, and the traceability paragraph now names the "
    f"same MRD as its reconciliation box. MRD v{MRD_TARGET}. " + TEST_EVIDENCE
)


def patch_m04() -> None:
    doc = Document(str(frd_path(M04_DIR, M04_STEM, M04_SOURCE)))
    set_cover_version(doc, M04_SOURCE, M04_TARGET)
    set_control(doc, "Version", M04_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Source scope", CODE_BASELINE)
    set_control(doc, "Code inspected", CODE_BASELINE)
    set_control(doc, "MRD baseline", f"XVS Module Requirements Document v{MRD_TARGET}")

    edit_paragraph(doc, "Five boundaries hold this module up",
                   "cannot land on its editor's own role without a role change",
                   "cannot land on any role without a role change")
    edit_cell(row_labelled(doc, "Role-change approval").cells[-1],
              "It is the route a restricted key takes onto a role its editor holds.",
              "It is the only route a restricted key takes onto a tenant role after "
              "provisioning: a role save raises the request itself for every restricted key "
              "the role does not already hold.")
    edit_cell(row_labelled(doc, "School administrator").cells[1],
              "Build the tenant's role catalogue, including restricted additions to roles they "
              "do not hold, assign and revoke roles within their own grant ceiling, raise a "
              "role change for a restricted addition to a role they hold, and decide",
              "Build the tenant's role catalogue, whose saves send any new restricted "
              "permission for approval rather than granting it, assign and revoke roles within "
              "their own grant ceiling, and decide")

    fr003 = fr_table(doc, "FR-003")
    set_value(fr003, "Current evidence", M04_FR003_EVIDENCE)
    set_value(fr003, "Acceptance", M04_FR003_ACCEPTANCE)
    edit_value(fr003, "Limit",
               "An addition to a role the editor does not hold reaches that role's current "
               "holders at once with no second person involved; the assignment ceiling governs "
               "who can be given the role, not what its existing holders receive.",
               "A restricted addition reaches a role's current holders only once its request is "
               "approved, and the person who raised it may be the one who approves it (Needs "
               "Attention).")

    fr017 = fr_table(doc, "FR-017")
    set_value(fr017, "Requirement", M04_FR017_REQUIREMENT)
    for label, tail in (("Current evidence", M04_FR017_EVIDENCE_TAIL),
                        ("Acceptance", M04_FR017_ACCEPTANCE_TAIL),
                        ("Limit", M04_FR017_LIMIT_TAIL)):
        set_value(fr017, label, row(fr017, label).cells[-1].text.strip() + tail)

    fr019 = fr_table(doc, "FR-019")
    edit_value(fr019, "Current evidence", "A group-only edit preserves existing explicit denies.",
               "An edit or an approved change that names no deny set keeps the role's existing "
               "explicit denies, lifting only a deny on a key it now grants, and the reach "
               "settlement names them outright.")
    edit_value(fr019, "Limit", M04_FR019_DENIES_OLD, "")

    edit_cell(row_labelled(doc, M04_ROLES_ROUTE).cells[-1],
              "A restricted permission is refused at creation; the creator, who does not hold "
              "the new role, may add it by editing the role afterwards.",
              "A restricted permission is not granted at creation: the role is created with the "
              "rest, one role change request carries the restricted keys, and the response's "
              "pending_additions names them.")
    role_route = row_labelled(doc, M04_ROLE_ROUTE).cells[-1]
    edit_cell(role_route,
              "One role, including permission_keys, group_ids and held_by_me, which says "
              "whether the reader holds an active assignment to it, so the screen can show "
              "before saving whether a restricted addition will save or need approval.",
              "One role, including permission_keys, group_ids, held_by_me, which says whether "
              "the reader holds an active assignment to it, and pending_additions, each "
              "restricted key waiting on the role with its request_id.")
    edit_cell(role_route,
              "A restricted addition to a role the reader holds is refused with 409 "
              "RESTRICTED_NEEDS_APPROVAL listing restricted_additions; to anybody else's role "
              "it saves.",
              "A restricted addition, to any role, saves the rest and raises one role change "
              "request for the restricted keys, skipping any already waiting; a ladder that "
              "cannot run refuses the whole save.")
    keep_format(
        row_labelled(doc, "A direct role addition, group member, or ALLOW override using a "
                          "restricted key").cells[-1],
        "A group member or ALLOW override: 400 naming the restricted key. A role addition, at "
        "creation or on update: the role saves without it, 201 or 200, the key appears in "
        "pending_additions and a role change request is raised. A role-change ladder that "
        "cannot be resolved refuses the whole save under the engine's own code.")

    set_paragraph(doc, "•  Both escalations end to end", (
        "•  Both escalations end to end: scope refuses a platform-only key on role and override "
        "paths; restriction refuses payments.payout.create in a group and as an ALLOW override, "
        "holds a restricted addition back for approval on every role save, whether the role is "
        "the editor's own, somebody else's or new, and the assignment ceiling still refuses "
        "giving a role that carries it to anybody whose assigner does not hold it."))
    edit_paragraph(doc, "•  Direct permission and group edits through the shared service",
                   "and a newly created direct deny stored as granted=False and proved to "
                   "override the same key supplied by an attached group.",
                   "a newly created direct deny stored as granted=False and proved to override "
                   "the same key supplied by an attached group, and a save or an approval that "
                   "names no deny set keeping every deny it does not grant.")
    edit_paragraph(doc, "•  The role-change ladder:",
                   "and the requester exemption confined to rbac.role_change.",
                   "the requester exemption confined to rbac.role_change, and a role save "
                   "raising its own request, not raising a second for a key already waiting, and "
                   "refusing whole when the ladder cannot run.")

    attention = table_with_header(doc, "Pri.")
    edit_cell(row_with_cell_starting(attention, 1, "A restricted addition to your own role").cells[1],
              "A restricted addition to your own role needs only your own approval.",
              "A restricted addition to any role needs only the approval of whoever raised it.")

    trace = next(t for t in doc.tables
                 if t.rows[0].cells[0].text.strip().startswith("MRD Module"))
    keep_format(row(trace, "Create and update roles").cells[2], (
        "Implemented. Ordinary permission and group changes require a reason and cross one "
        "locked, transactional mutation service. A restricted addition to any role, new or "
        "existing, is held back and raised as one role change request in the same save, and "
        "the role detail lists what waits through pending_additions. Branch reach can select "
        "several equal tenant branches and requires a reason to change. A holder cannot widen "
        "their own role's reach."))
    edit_cell(row(trace, "Role-change request workflow").cells[2],
              "The reviewer grant ceiling is gone.",
              "Every role save raises one for the restricted keys it adds. The reviewer grant "
              "ceiling is gone.")
    keep_format(row_labelled(doc, "10. MRD Traceability").cells[1],
                f"Agreement with MRD v{MRD_TARGET}'s twenty-three capability entries")
    edit_paragraph(doc, "MRD v2.85 records Module 4", "MRD v2.85", f"MRD v{MRD_TARGET}")
    edit_paragraph(doc, f"MRD v{MRD_TARGET} records Module 4",
                   "with twenty-two capability entries", "with twenty-three capability entries")
    update_reconciliation(doc, M04_RECONCILIATION)

    log_change(doc, M04_TARGET, M04_SUMMARY)
    assert_absent_outside_log(
        doc, "RESTRICTED_NEEDS_APPROVAL", "restricted_additions", "_actor_holds_this_role",
        "role the editor holds", "role they hold", "roles they do not hold",
        "clears the role's explicit denies", "editor's own role without",
        "refused at creation", "twenty-two capability", "MRD v2.83", "MRD v2.85", "MRD v2.86",
        "test_create_role_rejects_restricted_permission",
        "test_the_same_addition_to_somebody_elses_role_goes_through",
    )
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M04_DIR, M04_STEM, M04_TARGET),
           f"{M04_STEM.replace('_', ' ')} v{M04_TARGET}", M04_TARGET)


# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_CONTENTS_NOTE = "Restricted permissions wait for approval on every role"
MRD_INTRO = (
    "This revision records that a restricted permission waits for approval on every role save, "
    "whoever holds the role, that creating a role with one no longer fails, and that a role's "
    "explicit denies survive a save and an approval."
)
MRD_CHANGE_SUMMARY = (
    "A role save never grants a restricted permission the role does not already hold: on "
    "creation and update, and whoever holds the role, it grants the rest and raises one role "
    "change request for the restricted keys, which the role lists as pending until the ladder "
    "decides. This replaces approval for the editor's own role only, and creating a role with "
    "a restricted permission, refused outright before, now works. A role's explicit denies "
    "survive a save and an approval; both used to delete them. The console's role creation "
    "sends the reason the server requires. Capability entries unchanged at 510 across 31 "
    "modules; statuses do not move. M04 v1.27. Backend evidence and local school-app "
    "verification; no deployment claim."
)
MRD_DELTA_ROWS = [
    ["Restricted permissions", "Wait for approval on every role",
     "A role save, on creation or update and whoever holds the role, grants everything else "
     "and raises one role change request for the restricted keys; the role lists them as "
     "pending until the ladder decides."],
    ["Role creation", "Takes a restricted permission",
     "Creating a role with a restricted permission was refused with no way to raise the "
     "request; it now creates the role and raises it."],
    ["Explicit denies", "Kept through saves and approvals",
     "An edit or an approval that names no deny set keeps the role's denies, lifting only a "
     "deny on a key it grants; both used to delete every deny on the role."],
    ["Console role creation", "Sends its reason",
     "The create screen never sent the reason the server requires with permissions, and hid "
     "the refusal; it asks for one and shows any refusal."],
    ["Module FRDs", "One revised", "M04 v1.27."],
]
MRD_M04_DECISION_EDITS = [
    ("A restricted permission added to a role the editor holds is raised as a request in the "
     "same transaction,",
     "A restricted permission added to any role, a new one included and whoever holds it, is "
     "held back from the save and raised as a request in the same transaction,"),
    ("Separately, a restricted key is refused on direct role writes, in permission groups, and "
     "as an ALLOW override; it may enter a custom role only through a change request decided "
     "by somebody other than the requester, and that reviewer may grant only restricted keys "
     "they already hold.",
     "Separately, a restricted key is never granted by a role save and is refused in "
     "permission groups and as an ALLOW override; it enters a custom role only through an "
     "approved change request, which the role save raises itself."),
]
MRD_M04_BLURB_TAIL = (
    " A restricted permission never takes effect on a role save: the save grants the rest and "
    "raises an approval request for it. Documented by M04 FRD v1.27."
)


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    doc = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta, log = tables[76], tables[78]
    assert delta.rows[0].cells[0].text.strip().endswith("capability delta")
    assert index.rows[0].cells[5].text.strip() == "Entries"

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for r in control.rows:
        label = r.cells[0].text.strip()
        if label == "Version":
            replace_cell(r.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(r.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(r.cells[1], CODE_BASELINE, size=9)
    for r in contents.rows:
        if r.cells[0].text.strip().startswith("5."):
            keep_format(r.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(r.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (text.startswith("5. ") and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    grid = tables[mrd_tools.CAPABILITIES[4]]
    cells = list({id(c._tc): c for r in grid.rows for c in r.cells
                  if c.text.startswith("CURRENT DECISION")}.values())
    if len(cells) != 1:
        raise ValueError(f"Module 4 has {len(cells)} CURRENT DECISION cells")
    for old, new in MRD_M04_DECISION_EDITS:
        edit_run_text(cells[0], old, new)

    blurb = mrd_tools.module_blurbs(doc)[4]
    if "Documented by M04" in blurb.text:
        raise ValueError("Module 4's description already names an FRD version")
    mrd_tools.retitle(blurb, blurb.text.rstrip() + MRD_M04_BLURB_TAIL)

    total = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total != 510:
        raise ValueError(f"Capability total is {total}, expected 510")

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)
    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if len(intro) != 1:
        raise ValueError(f"{len(intro)} delta introductions found")
    mrd_tools.retitle(intro[0], MRD_INTRO)

    # Copied from the latest version row, so the new row reads at the same size.
    log.rows[1]._tr.addprevious(copy.deepcopy(log.rows[1]._tr))
    for cell, text in zip(log.rows[1].cells, (MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)):
        keep_format(cell, text)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m04()
    patch_mrd()


if __name__ == "__main__":
    main()
