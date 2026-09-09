#!/usr/bin/env python3
"""Version the MRD and Module 3 FRD for browser refresh-cookie sessions."""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)


REVIEW_DATE = "9 September 2026"
SHORT_DATE = "9 Sep 2026"
MRD_SOURCE_VERSION = "2.69"
MRD_TARGET_VERSION = "2.70"
FRD_SOURCE_VERSION = "1.9"
FRD_TARGET_VERSION = "1.10"


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_control_value(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            replace_cell(row.cells[1], value, size=9)
            return
    raise ValueError(f"Control row not found: {label}")


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def prepend_change_log(table, version: str, summary: str) -> None:
    template = table.rows[1]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], SHORT_DATE, size=8)
    replace_cell(row.cells[2], summary, size=8)


def insert_row_before(table, row_index: int, values: list[str]) -> None:
    template = table.rows[row_index]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[row_index]
    for cell, value in zip(row.cells, values):
        replace_cell(cell, value, size=8.2)


def append_row(table, values: list[str]) -> None:
    template = table.rows[-1]
    template._tr.addnext(copy.deepcopy(template._tr))
    row = table.rows[-1]
    for cell, value in zip(row.cells, values):
        replace_cell(cell, value, size=8.2)


def finish(doc, output: Path, title: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS Module Requirements Document v{MRD_TARGET_VERSION}"
    doc.core_properties.title = title
    doc.core_properties.version = MRD_TARGET_VERSION
    replace_cover_version(doc.tables[0], MRD_SOURCE_VERSION, MRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Version", MRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Review date", REVIEW_DATE)
    replace_control_value(
        doc.tables[1],
        "Source scope",
        "Backend and first-party client change moving browser refresh credentials into a Secure HttpOnly SameSite cookie, keeping access tokens in memory, requiring CSRF on refresh and logout, and removing the response-body refresh-token contract (9 September 2026)",
    )

    replace_cell(
        doc.tables[2].rows[5].cells[0],
        f"5. v{MRD_TARGET_VERSION} Capability Delta",
        size=9,
        bold=True,
        color=BLUE,
    )
    replace_cell(
        doc.tables[2].rows[5].cells[1],
        "Browser refresh credentials leave JavaScript and the legacy body contract is removed",
        size=9,
    )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Authentication, invitations, account security, sessions"):
            write_paragraph(
                paragraph,
                "Authentication, invitations, account security, sessions, platform staff, and organization-position structures. Browser login returns only the short-lived access token and stores the rotating refresh credential in a host-only Secure HttpOnly SameSite cookie. Refresh and logout accept that cookie alone and require CSRF, while first-party clients keep the access token in memory and rebuild context through refresh and /auth/me/ after a reload. Invitation and password-reset links remain request-bound, and account creation applies the same restricted-role grant ceiling as the dedicated assignment endpoints.",
                size=9,
                space_after=5,
            )
        elif text == f"5. v{MRD_SOURCE_VERSION} Capability Delta":
            write_paragraph(
                paragraph,
                f"5. v{MRD_TARGET_VERSION} Capability Delta",
                size=17,
                bold=True,
                space_before=15,
                space_after=8,
            )
        elif text.startswith("This revision records the product catching up with its own gate"):
            write_paragraph(
                paragraph,
                "This revision closes the browser refresh-token exposure at the shared contract. The Console and school client restore sessions through the HttpOnly cookie, keep access tokens in memory, and persist no credential-bearing auth state. Login and refresh never return a refresh credential, and refresh and logout no longer accept one from the request body.",
                size=9,
                space_after=5,
            )

    replace_cell(
        doc.tables[12].rows[0].cells[1],
        "▸  JWT access with HttpOnly refresh-token rotation",
        size=8.2,
    )
    replace_cell(
        doc.tables[12].rows[1].cells[0],
        "▸  Cookie-scoped refresh-token blacklisting and logout",
        size=8.2,
    )

    rebuild_table(
        doc.tables[76],
        [f"v{MRD_TARGET_VERSION} capability delta", "Decision", "Evidence"],
        [
            [
                "Browser refresh credential",
                "Removed from JavaScript",
                "Login stores the one-day rotating credential in a host-only Secure HttpOnly SameSite cookie. Neither first-party client can read it.",
            ],
            [
                "Access credential",
                "Memory only",
                "Login and refresh return the fifteen-minute access token. The clients keep it in module memory rather than cookies, Redux, localStorage, or sessionStorage.",
            ],
            [
                "Reload recovery",
                "Restored from the server",
                "A reload uses the refresh cookie to obtain a new access token, then /auth/me/ rebuilds the user, tenant, permissions, session, and active proxy context.",
            ],
            [
                "Refresh and logout",
                "Cookie plus CSRF",
                "Both routes read the refresh cookie only. A missing or invalid double-submit CSRF token is refused before rotation or revocation.",
            ],
            [
                "Browser origin",
                "First-party allowlist",
                "Login accepts the Console and one school subdomain per configured base domain. An untrusted Origin is refused before credentials are processed.",
            ],
            [
                "Legacy token body",
                "Removed",
                "Login and refresh responses never contain refresh. A refresh value submitted in a request body is ignored and cannot authenticate the request.",
            ],
        ],
        [1.7, 1.15, 4.42],
        font_size=8.2,
    )
    prepend_change_log(
        doc.tables[78],
        MRD_TARGET_VERSION,
        "Closed the browser refresh-token exposure at the shared authentication contract. Login returns only the short-lived access token and stores the rotating refresh credential in a host-only Secure HttpOnly SameSite cookie. Refresh and logout accept that cookie alone and require CSRF; a refresh value in the request body is ignored. Login and refresh responses never contain refresh. The Console and school client keep access tokens in module memory, persist no credential-bearing auth state, and restore context through cookie refresh and /auth/me/. Browser login is limited to configured first-party origins, including one school subdomain per base domain. Module 3 remains Backend Partial and In use Complete with seventeen capabilities; MFA and the existing lifecycle gaps are unchanged. Verified by 11 focused backend session tests, 46 school-client auth tests, the school production build, and a production-bundle token audit. Backend and client evidence only; nothing here is deployed.",
    )
    finish(doc, output, title)


def patch_frd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M03 Identity Team and Organogram Functional Requirements "
        f"Document v{FRD_TARGET_VERSION}"
    )
    doc.core_properties.title = title
    doc.core_properties.version = FRD_TARGET_VERSION
    replace_cover_version(doc.tables[0], FRD_SOURCE_VERSION, FRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Version", FRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Review date", REVIEW_DATE)
    replace_control_value(
        doc.tables[1],
        "Code baseline",
        "Backend worktree and migrated first-party browser clients at 9 September 2026",
    )
    replace_control_value(
        doc.tables[1],
        "Source MRD",
        f"XVS Module Requirements Document v{MRD_TARGET_VERSION} | Module 3",
    )

    replace_cell(
        doc.tables[3].rows[2].cells[1],
        "Tenant resolved first, account looked up inside it, password checked, lockout checked, status checked, session opened, short-lived access token returned, and rotating refresh credential stored in a Secure HttpOnly SameSite cookie. A browser Origin must match the configured first-party allowlist.",
        size=8.5,
    )
    replace_cell(
        doc.tables[3].rows[6].cells[1],
        "One LoginSession per sign-in, keyed to the refresh token's JTI. The raw refresh credential stays in a host-only HttpOnly cookie and reaches only the CSRF-protected refresh and logout routes, so one device, every other device, or every device at once can still be ended and blacklisted.",
        size=8.5,
    )

    fr015 = doc.tables[21]
    replace_cell(
        fr015.rows[1].cells[1],
        "A browser refresh credential must stay outside JavaScript, and a session row that says \"ended\" must correspond to a token that no longer works. Login and rotation must not return the refresh credential in a response body or accept it from a request body.",
        size=8.5,
    )
    replace_cell(
        fr015.rows[2].cells[1],
        "LoginSession stores the refresh token's JTI, the address, the agent, a derived device label and the last-seen time. Login returns the short-lived access token and places the rotating refresh credential in a host-only Secure HttpOnly SameSite cookie. Refresh and logout read only that cookie and enforce Django's double-submit CSRF check. Rotation registers the new token, moves the session JTI forward, sets the replacement cookie, and returns no refresh value. Logout blacklists only the cookie's session, while the self-service and administrative routes retain their existing one-session and all-session behavior.",
        size=8.5,
    )
    replace_cell(
        fr015.rows[3].cells[1],
        "BrowserSessionContractTests proves login and refresh responses never contain refresh, the cookie is HttpOnly, Secure, SameSite Strict and path-wide, refresh rotation replaces it, refresh and logout require CSRF, a submitted body token cannot authenticate refresh, a school origin is accepted, and an attacker origin is refused. SessionScopedLogoutTests proves logout still ends only the session named by the cookie.",
        size=8.5,
    )
    replace_cell(
        fr015.rows[4].cells[1],
        "Access tokens already issued remain usable until their fifteen-minute expiry and remain readable to the active page while held in memory. expire_stale_login_sessions still runs only when somebody opens a session list, so an expired refresh token can leave its row reading active until then. Administrative force-logout targets remain scoped through administrable_user.",
        size=8.5,
    )

    replace_cell(
        doc.tables[29].rows[8].cells[2],
        "A session row, a fifteen-minute access token in the response, a rotating refresh credential in a Secure HttpOnly SameSite cookie, the effective permission list, and the tenant and school context blocks. A reload rotates the cookie and rebuilds context through /auth/me/.",
        size=8.2,
    )
    replace_cell(
        doc.tables[32].rows[3].cells[2],
        "Tenant-aware manager with an unscoped escape hatch. Carries the refresh JTI, the address, the agent, a derived device label, and the end reason. The raw refresh credential is not stored on this row and is delivered to browsers only through the HttpOnly cookie. Two indexes on (user, is_active) and (is_active, last_seen_at).",
        size=8.2,
    )

    auth_routes = doc.tables[33]
    insert_row_before(
        auth_routes,
        1,
        [
            "GET /auth/csrf/",
            "Set the readable CSRF cookie and return the same non-secret token for first-party hosts that cannot read the API host's cookie. AllowAny.",
        ],
    )
    replace_cell(
        auth_routes.rows[2].cells[1],
        "Sign in. Optional tenant slug in the body. A browser Origin must match the first-party allowlist. Returns access and session context, sets the refresh and CSRF cookies, and never returns refresh. AllowAny, throttled.",
        size=8.2,
    )
    replace_cell(
        auth_routes.rows[3].cells[1],
        "Blacklist the refresh cookie and end only its session. The request body is ignored. CSRF required; idempotent after a valid cookie reaches revocation.",
        size=8.2,
    )
    replace_cell(
        auth_routes.rows[4].cells[1],
        "Rotate the refresh cookie, register the new token, move the session JTI forward, and return access plus session id only. The request body is ignored. AllowAny with cookie validity and CSRF as the gates.",
        size=8.2,
    )

    conditions = doc.tables[36]
    append_row(
        conditions,
        [
            "A browser login from an Origin outside the configured Console and school-app origins",
            "403 before account lookup or password processing, and no refresh cookie is set.",
        ],
    )
    append_row(
        conditions,
        [
            "Refresh or logout with no valid CSRF token, or refresh with a body token but no cookie",
            "403 for the CSRF failure; 401 when CSRF is valid but no refresh cookie exists. The body token is never considered.",
        ],
    )

    replace_cell(
        doc.tables[37].rows[8].cells[1],
        "Supplies the short-lived access token, rotating refresh token, outstanding-token table and blacklist. CodeXRefreshToken adds tenant id, tenant slug, branch id, account status and full name as claims. Browser delivery is narrower than the library serializer: access is returned to the client, while refresh is stored only in the HttpOnly cookie and is absent from login and refresh response bodies.",
        size=8.2,
    )

    traceability = doc.tables[40]
    replace_cell(
        traceability.rows[2].cells[0],
        "JWT access with HttpOnly refresh-token rotation",
        size=8.2,
    )
    replace_cell(
        traceability.rows[2].cells[2],
        "Implemented. Rotation registers the new token and moves the session JTI forward. Browsers receive the rotating credential only through the Secure HttpOnly SameSite cookie, and the response carries access plus session id only.",
        size=8.2,
    )
    replace_cell(
        traceability.rows[3].cells[0],
        "Cookie-scoped refresh-token blacklisting and logout",
        size=8.2,
    )
    replace_cell(
        traceability.rows[3].cells[2],
        "Implemented. Refresh and logout read only the CSRF-protected refresh cookie; a body token is ignored. Logout remains scoped to that cookie's session, and all-device revocation stays in the administrative and status paths.",
        size=8.2,
    )
    replace_cell(
        doc.tables[41].rows[0].cells[0],
        "MRD RECONCILIATION\n"
        f"• MRD v{MRD_TARGET_VERSION} lists Module 3 as Backend Partial and In use Complete with seventeen capabilities. This revision renames the two refresh capabilities to state the cookie boundary without adding a product surface or changing the count.\n"
        "• Login and refresh responses carry no refresh credential. Browser refresh and logout use the host-only HttpOnly cookie and require CSRF; the request-body compatibility path is absent.\n"
        "• Partial remains correct because privileged MFA and the existing tenant lifecycle, lockout, retention, audit-model and organogram gaps remain.\n"
        "• Backend and client evidence only. Neither document claims deployment, production adoption, or data migration completion.",
        size=8.5,
    )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "FR-015  Track a Session Against Its Refresh Token, and End It Cryptographically":
            write_paragraph(
                paragraph,
                "FR-015  Keep Refresh Credentials Outside JavaScript and End Sessions Cryptographically",
                size=13.5,
                bold=True,
                space_before=11,
                space_after=6,
            )
        elif text.startswith("Routes are mounted under /v1/user/"):
            write_paragraph(
                paragraph,
                "Routes are mounted under /v1/user/. Browser login, refresh and logout use credentialed cross-origin requests. Login and refresh return no refresh value; refresh and logout read the host-only HttpOnly cookie and require CSRF. List routes paginate and return the shared success envelope, which returns an empty list as an empty list and coerces only a missing payload to an empty object. Authenticated routes require the ?tenant= assertion unless the view declares tenant_param_required = False.",
                size=9,
                space_after=5,
            )
        elif text.startswith("The current vs_user suite contains"):
            write_paragraph(
                paragraph,
                "The focused browser-session contract contains nine tests and passed in full, alongside the two existing session-scoped logout tests. The migrated school client passed 46 focused auth tests and its production build; its source and production bundle contain neither js-cookie nor refresh-token identifiers. The Console's cookie flow is already present, and its focused auth test passed during this review. Deployment is not inferred from these checks.",
                size=9,
                space_after=5,
            )
        elif text.startswith("MRD v2.61 records Module 3"):
            write_paragraph(
                paragraph,
                f"MRD v{MRD_TARGET_VERSION} records Module 3 as Identity, Team & Organogram, Phase V1, Backend Partial, In use Complete, code ownership vs_user, with seventeen capability entries. Each entry maps below.",
                size=9,
                space_after=5,
            )

    prepend_change_log(
        doc.tables[42],
        FRD_TARGET_VERSION,
        "Closes the browser refresh-token exposure and removes the shared response-body compatibility path. Login returns only the fifteen-minute access token and stores the rotating one-day refresh credential in a host-only Secure HttpOnly SameSite cookie. Refresh and logout accept that cookie alone, require Django CSRF, and ignore request-body refresh values. Rotation sets the replacement cookie and returns access plus session id, never refresh. Browser login is limited to configured first-party origins. The Console and school client keep access tokens in memory, persist no credential-bearing auth state, and restore sessions through refresh followed by /auth/me/. FR-015, the sign-in and session ownership map, first-login workflow, LoginSession record, API and refusal contracts, SimpleJWT dependency, verification evidence, and MRD traceability are updated. Module 3 remains Backend Partial and In use Complete with seventeen capabilities. Verified by nine browser contract tests, two session-scoped logout tests, 46 school-client auth tests, the school production build, and a production-bundle token audit. Backend and client evidence only; nothing here is deployed.",
    )
    finish(doc, output, title)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[3])
    args = parser.parse_args()
    root = args.root.resolve()
    mrd_dir = root / "docs/frd/module-requirements"
    frd_dir = root / "docs/frd/functional-requirements/03-identity-team-and-organogram"
    patch_mrd(
        mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_SOURCE_VERSION}.docx",
        mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_TARGET_VERSION}.docx",
    )
    patch_frd(
        frd_dir / f"XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document_v{FRD_SOURCE_VERSION}.docx",
        frd_dir / f"XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document_v{FRD_TARGET_VERSION}.docx",
    )


if __name__ == "__main__":
    main()
