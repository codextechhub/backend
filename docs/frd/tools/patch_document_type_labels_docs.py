#!/usr/bin/env python3
"""Cut MRD v2.91, M07 FRD v1.17 and M22 FRD v1.13: every document type has one human name.

What changed in the backend, and therefore in the documents:

* Every workflow document type is named by its handler's ``noun``, and
  ``register_handler`` refuses a handler whose noun is empty, beside the
  reversal-contract refusal, so a type cannot reach a screen as its code.
* ``vs_workflow.conditions.fields.document_type_label`` is the one place a
  type becomes words: the noun, or the code's last segment in words for a type
  with no registered handler.
* ``document_type_label`` is returned on workflow instance list and detail
  rows (and so on the pending and submitted dashboards), templates, approval
  delegations, team-load rows and a Dynamic Role's ``used_by`` entries.
  Instance rows also carry ``document_title``, read from the stored summary.
* Procurement's two private label maps are gone. Its approval queue and
  coverage report read the same function, so a requisition reads "Purchase
  requisition" and the queue's labels are in sentence case.

The approval inbox had shown approvers names such as "Rbac Role Grant",
because the shared frontend title-cased the raw code and no queue payload
carried a label.

M23 was checked and left: it quotes no document type name and no queue shape.
M19, M18, M12, M04 and M03 were checked for the same and state nothing this
changes.

Each document starts from the newest version in its folder and is written at
the next free number; the script refuses to run if a newer version has
appeared, and ``finish`` refuses to overwrite one.

    python tools/patch_document_type_labels_docs.py
"""
from __future__ import annotations

import copy
import glob

from docx import Document

import patch_mrd_v2_79_docs as mrd_tools
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    replace_cell,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    row_labelled,
)
from patch_staff_id_and_auth_events_docs import (
    edit_paragraph,
    edit_value,
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
)

import patch_record_history_docs

REVIEW_DATE, SHORT_DATE = "26 September 2026", "26 Sep 2026"
MRD_SOURCE, MRD_TARGET = "2.90", "2.91"
CODE_BASELINE = (
    "Backend main at 9fa0bede with the document type label change in its working tree, "
    "26 September 2026. The FinPro approval screens' integration is implemented and pending "
    "release"
)
TEST_EVIDENCE = (
    "Verified by the full vs_workflow (424 tests), vs_rbac (891), vs_procurement (614), "
    "vs_finance (872), vs_payments (212), vs_user (421) and schools.vs_staff (317) suites, "
    "slow tests included, each app run on its own and all passing. The FinPro approval "
    "screens' integration is implemented and pending release; nothing here claims deployment."
)

# log_change writes the module's own review date, read from this module.
patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


def _version_key(path: str):
    return [int(x) for x in path.rsplit("_v", 1)[1][:-5].split(".")]


def require_newest(pattern: str, source: str) -> None:
    """Refuse to build from ``source`` when a newer version has been written since."""
    paths = [p for p in glob.glob(pattern) if "~$" not in p]
    newest = max(paths, key=_version_key)
    found = newest.rsplit("_v", 1)[1][:-5]
    if found != source:
        raise SystemExit(f"Newest is v{found}, not v{source}: rebase this script on it first")


# ── M07 Workflow & Approval Engine ───────────────────────────────────────────

M07_DIR = "07-workflow-and-approval-engine"
M07_STEM = "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document"
M07_SOURCE, M07_TARGET = "1.16", "1.17"

NOUNS = (
    "Journal entry, Customer refund, Bad-debt write-off, Concession, Credit or debit note, "
    "Expense claim, Payout batch, Platform user account, Role permission change, Restricted "
    "role grant, Leave request, Purchase requisition, Purchase order, Vendor invoice and "
    "Vendor payment"
)

M07_FR034_HEADING = "FR-034  Give Every Document Type One Human Name"
M07_FR034 = [
    ("Requirement",
     "Every approvable document type has one name people use for it, and every surface that "
     "lists approval work shows that name rather than the type's code. A type that has not "
     "said what it is called does not register."),
    ("Current evidence",
     "BaseWorkflowHandler.noun names one document of the type. register_handler refuses a "
     "handler whose noun is empty with a TypeError while the apps load, beside the "
     "reversal-contract refusal (FR-019). vs_workflow.conditions.fields.document_type_label "
     "is the one place a type becomes words: the handler's noun, or, for a type with no "
     "registered handler such as one retired while its instances remain, the last segment of "
     f"its code in words. The fifteen registered nouns are {NOUNS}. document_type_label is "
     "returned on instance list and detail rows, which the pending and submitted dashboards "
     "use, on templates, on approval delegations (blank for one covering every type), on "
     "team-load rows and on a Dynamic Role's used_by entries. Instance rows also carry "
     "document_title, the title the handler gave the document at submission, read from the "
     "stored document_summary so a page of rows costs no extra query, and blank where the "
     "handler gave none. Procurement's approval queue and coverage report read the same "
     "function (Module 22), so no module keeps a label map of its own."),
    ("Acceptance",
     "An approver's inbox names a restricted role grant Restricted role grant and a role "
     "change Role permission change, never rbac.role_grant or a title-cased code. A handler "
     "with no noun cannot register, and every registered type's label is its own noun "
     "(test_registry: test_a_handler_with_no_noun_cannot_register and "
     "test_every_registered_type_has_a_label_of_its_own). A queue row names the document in "
     "words (RoleGrantLadderTests.test_the_queue_row_names_the_document_in_words)."),
    ("Current limit",
     "document_title is the summary frozen at submission, so a document renamed afterwards "
     "keeps the title it was submitted under, and an instance whose summary carries no title "
     "has a blank one. The shared approval screens read the label and the title and fall "
     "back to title-casing the code only where they are absent; that integration is "
     "implemented and pending release, and this document does not evidence it deployed."),
]

M07_SUMMARY = (
    "Minor revision. Adds FR-034: every document type has one human name, its handler's "
    "noun, and register_handler refuses a handler without one, beside the reversal-contract "
    "refusal. document_type_label is the one place a type becomes words, and it is returned "
    "on instance rows and so on the pending and submitted dashboards, on templates, "
    "delegations, team-load rows and a Dynamic Role's used_by entries; instance rows also "
    "carry document_title from the stored summary. The approval inbox had shown approvers "
    "names such as Rbac Role Grant, because the shared frontend title-cased the raw code and "
    "no queue payload carried a label. FR-023, the template, Dynamic Role, instance, "
    "dashboard and delegation contracts, the domain-app dependency and traceability follow. "
    f"MRD v{MRD_TARGET}. " + TEST_EVIDENCE
)


def patch_m07() -> None:
    require_newest(str(ROOT / "functional-requirements" / M07_DIR / f"{M07_STEM}_v*.docx"),
                   M07_SOURCE)
    doc = Document(str(frd_path(M07_DIR, M07_STEM, M07_SOURCE)))
    set_cover_version(doc, M07_SOURCE, M07_TARGET)
    set_control(doc, "Version", M07_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")

    append_to(fr_table(doc, "FR-023"), "Current evidence",
              " Each queue row carries document_type_label and document_title, and each "
              "team-load row document_type_label, so a screen never names a type by its code "
              "(FR-034).")
    add_fr(doc, M07_FR034_HEADING, "FR-034 | Implemented", M07_FR034)

    for label, tail in (
        ("GET /templates/", " Each carries document_type_label."),
        ("GET /templates/{id}/", " It carries document_type_label."),
        ("GET, POST /dynamic-roles/",
         " Each used_by entry names its template's type with document_type_label."),
        ("GET /instances/", " Each row carries document_type_label and document_title."),
        ("GET /instances/{id}/", " It carries document_type_label and document_title too."),
        ("GET /dashboard/pending/",
         " Rows as GET /instances/, with document_type_label and document_title."),
        ("GET /dashboard/submitted/",
         " Rows as GET /instances/, with document_type_label and document_title."),
        ("GET /dashboard/team-load/", " Each row carries document_type_label."),
        ("GET, POST /delegations/",
         " Each carries document_type_label, blank for a delegation covering every type."),
    ):
        cell = row_labelled(doc, label).cells[-1]
        keep_format(cell, cell.text.rstrip() + tail)

    cell = row_labelled(doc, "Domain apps, including the school staff app").cells[-1]
    keep_format(cell, cell.text.rstrip() + (
        " Each handler also names its type with a noun, which every approval surface shows "
        "and without which the type does not register."))

    retitle = paragraph_starting(doc, "Module 7 carries 28 capability entries")
    mrd_tools.retitle(retitle, (
        f"Module 7 carries 28 capability entries in MRD v{MRD_TARGET}. Each maps to the "
        "requirements below. Requiring a document type to declare what its approval releases "
        "before it may register strengthens action reversal and the domain handler entries "
        "without changing the count, and the restricted role grant ladder is one more document "
        "type on the role-based resolution, central-template and reversal entries, adding "
        "none. Giving every type one human name strengthens the dashboards and the approval "
        "detail entries without adding one."))
    trace = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "MRD capability")
    keep_format(row(trace, "Pending, submitted, and team-load dashboards").cells[-1],
                "FR-023, FR-034")

    log_change(doc, M07_TARGET, M07_SUMMARY)
    assert_absent_outside_log(doc, "MRD v2.89", "MRD v2.80")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M07_DIR, M07_STEM, M07_TARGET),
           f"{M07_STEM.replace('_', ' ')} v{M07_TARGET}", M07_TARGET)


# ── M22 Procurement & Requisitions ───────────────────────────────────────────

M22_DIR = "22-procurement-and-requisitions"
M22_STEM = "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document"
M22_SOURCE, M22_TARGET = "1.12", "1.13"

M22_SUMMARY = (
    "Minor revision. The approval queue and the coverage report name the four procurement "
    "document types through the engine's document_type_label (Module 7 FR-034) rather than "
    "two label maps of this module's own, which disagreed with each other and with the "
    "engine. Visible effect: a requisition reads Purchase requisition on both, where it read "
    "Requisition, and the queue's Purchase Order, Vendor Invoice and Vendor Payment are "
    "sentence case, as the coverage report already had them. FR-008, FR-009, the queue and "
    "coverage routes, the Module 7 dependency and the traceability paragraph are updated. "
    f"MRD v{MRD_TARGET}. " + TEST_EVIDENCE
)


def patch_m22() -> None:
    require_newest(str(ROOT / "functional-requirements" / M22_DIR / f"{M22_STEM}_v*.docx"),
                   M22_SOURCE)
    doc = Document(str(frd_path(M22_DIR, M22_STEM, M22_SOURCE)))
    set_cover_version(doc, M22_SOURCE, M22_TARGET)
    set_control(doc, "Version", M22_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")

    append_to(fr_table(doc, "FR-008"), "Current evidence",
              " Each document type is named through the engine's document_type_label, so the "
              "report says Purchase requisition, Purchase order, Vendor invoice and Vendor "
              "payment, the words the engine's own inbox uses.")
    append_to(fr_table(doc, "FR-009"), "Current evidence",
              " Each queue row names its document type through the same function (Module 7 "
              "FR-034), in sentence case; this module keeps no label map of its own.")

    for label, tail in (
        ("GET /approvals/", " Each row carries document_type_label, the engine's name for "
                            "the type."),
        ("GET /approvals/coverage/", " Each document type is named by document_type_label."),
    ):
        cell = row_labelled(doc, label).cells[-1]
        keep_format(cell, cell.text.rstrip() + tail)

    cell = row_labelled(doc, "Module 7, Workflow & Approval Engine").cells[-1]
    keep_format(cell, cell.text.rstrip() + (
        " Its document_type_label names the four types on this module's queue and coverage "
        "report from each handler's noun: Purchase requisition, Purchase order, Vendor invoice "
        "and Vendor payment."))

    edit_paragraph(doc, "Module 22 carries 16 capability entries", "MRD v2.78",
                   f"MRD v{MRD_TARGET}")
    paragraph = paragraph_starting(doc, "Module 22 carries 16 capability entries")
    mrd_tools.retitle(paragraph, paragraph.text.rstrip() + (
        " Naming the queue's document types through the engine strengthens the approval queue "
        "entry without changing the count."))

    log_change(doc, M22_TARGET, M22_SUMMARY)
    assert_absent_outside_log(doc, "MRD v2.78", "MRD v2.81")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M22_DIR, M22_STEM, M22_TARGET),
           f"{M22_STEM.replace('_', ' ')} v{M22_TARGET}", M22_TARGET)


# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_CONTENTS_NOTE = "Every document type has one human name"
MRD_INTRO = (
    "This revision records that every approval document type has one human name, that "
    "registration refuses a type without one, and that every approval surface, procurement's "
    "own queue included, shows that name rather than the type's code."
)
MRD_CHANGE_SUMMARY = (
    "Every workflow document type is named by its handler's noun, and a handler without one "
    "does not register. One engine function turns a type into words, and its label is on "
    "every approval surface: instance rows and so the pending and submitted inboxes, "
    "templates, delegations, team load and Dynamic Roles, with the document's own title on "
    "queue rows. The inbox had shown approvers names such as Rbac Role Grant. Procurement's "
    "queue and coverage report use the same label, so a requisition reads Purchase "
    "requisition. Capability entries unchanged; statuses do not move. M07 v1.17 and M22 "
    "v1.13. Backend evidence; the FinPro approval screens' integration is implemented and "
    "pending release; no deployment claim."
)
MRD_DELTA_ROWS = [
    ["Document type names", "One per type, from its handler",
     "Every handler declares a noun, such as Customer refund or Restricted role grant, and "
     "registration refuses a handler without one while the apps load."],
    ["Approval inbox", "Names documents in words",
     "Instance rows, and so the pending and submitted inboxes, carry document_type_label and "
     "the document's own title; the inbox had shown names such as Rbac Role Grant."],
    ["Other approval surfaces", "Carry the same label",
     "Templates, delegations, team-load rows and Dynamic Role usage carry document_type_label "
     "from the one engine function."],
    ["Procurement queue", "Uses the engine's names",
     "The queue and coverage report drop their own label maps: Requisition becomes Purchase "
     "requisition, and the queue's labels move to sentence case."],
    ["Module FRDs", "Two revised",
     "M07 v1.17, M22 v1.13. M23, M19, M18, M12, M04 and M03 checked and left."],
]


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    require_newest(str(folder / "XVS_Module_Requirements_Document_v*.docx"), MRD_SOURCE)
    doc = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta, log = tables[76], tables[78]
    assert delta.rows[0].cells[0].text.strip().endswith("capability delta")
    assert index.rows[0].cells[5].text.strip() == "Entries"
    total_before = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for r in control.rows:
        label = r.cells[0].text.strip()
        if label == "Version":
            replace_cell(r.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(r.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(r.cells[1], CODE_BASELINE, size=9)
    for r in contents.rows:
        if r.cells[0].text.strip().startswith("5."):
            keep_format(r.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(r.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (text.startswith("5. ") and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    blurbs = mrd_tools.module_blurbs(doc)
    blurb7 = blurbs[7]
    if blurb7.text.count("Documented by M07 FRD v1.16.") != 1:
        raise ValueError("Module 7's description does not name M07 FRD v1.16 once")
    mrd_tools.retitle(blurb7, blurb7.text.replace(
        "Documented by M07 FRD v1.16.",
        "Every document type has one human name, its handler's noun, without which it does not "
        "register, and every approval surface shows that name, with the document's own title "
        "on queue rows. Documented by M07 FRD v1.17."))
    blurb22 = blurbs[22]
    if "Documented by M22" in blurb22.text:
        raise ValueError("Module 22's description already names an FRD version")
    mrd_tools.retitle(blurb22, blurb22.text.rstrip() + (
        " The approval queue and coverage report name document types through the engine, so a "
        "requisition reads Purchase requisition on both. Documented by M22 FRD v1.13."))

    total_after = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total_after != total_before:
        raise ValueError(f"Capability total moved from {total_before} to {total_after}")

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)
    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if len(intro) != 1:
        raise ValueError(f"{len(intro)} delta introductions found")
    mrd_tools.retitle(intro[0], MRD_INTRO)

    # Copied from the latest version row, so the new row reads at the same size.
    log.rows[1]._tr.addprevious(copy.deepcopy(log.rows[1]._tr))
    for cell, text in zip(log.rows[1].cells, (MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)):
        keep_format(cell, text)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m07()
    patch_m22()
    patch_mrd()


if __name__ == "__main__":
    main()
