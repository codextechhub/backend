#!/usr/bin/env python3
"""Cut M17 v1.7: a school bills fees through the cohort route, at the price list's branch.

Queue entry D39 (324401bc), and therefore the document:

* The school's cohort billing route, POST
  /v1/school-finance/fee-structures/<pk>/generate-invoices/, is how a school
  raises fee invoices. It bills a named cohort only; there is no bill-everyone
  form.
* A named child outside the caller's branch reach answers 404 and nothing is
  billed.
* A fee structure that carries a branch is that branch's price list: a run
  naming a child at another branch is refused 409 WRONG_BRANCH, even for a
  caller who sees every branch. A school-wide structure bills every branch.
* The link-term route answers a read under finance.feestructure.view.
* The dry run returns the due date the run will write, resolved by the
  school's fee due-date rule, and the fee structure read carries branch_id.

The school app no longer calls the finance engine's own generate route: it
fills FinPro v0.7.24's generation slot with a cohort drawer (school-fe
9686919), committed and pending release. The console still uses that route.

Queue entry D48 (8cd97cae) bounds that route the same way. all_active bills
only the active customers the caller reaches, and a named customer outside
that reach answers 404. A structure with a branch is that branch's price list:
a named customer filed anywhere else, school-wide included, is refused 409
WRONG_BRANCH, and all_active from it selects only that branch's customers. A
dunning run started by a branch-bound caller stays in their reach; the
scheduled run is school-wide. That closes the all_active branch gap.

Two entries queued for Module 19 contradict sentences in this document, so
their Module 17 side is carried here: D16 (dcc692e2) narrows AR aging to the
reader's branches, and D17 (4e3f3456) builds a branch-bound reader's AR
reconciliation and statements from their own journals, which retires FR-013's
limit that the statements aggregate across branches.

Queue entry D62 (0c59e23d, merged 864468ba) holds a receipt's deposit account
to the receipt's branch: a customer receipt and an invoice payment work out the
branch they inherit before the account is named, and a ledger account behind
another branch's bank account is refused 400, "Deposit it into ...".

M19 and the MRD are not revised here.

    python tools/patch_cohort_billing_route_docs.py
"""
from __future__ import annotations

from docx import Document

from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    edit_box,
    insert_row_after,
    row_labelled,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    normalise_change_log,
    repair_ooxml,
    row,
    set_value,
)

from patch_backend_owned_permission_registry_docs import keep_rows_whole
from patch_record_history_docs import change_log_table

import patch_record_history_docs

REVIEW_DATE = "28 September 2026"
CODE_BASELINE = (
    "Backend main at 7860ea86, 28 September 2026, carrying 0c59e23d (a receipt deposits into "
    "its own branch's bank ledger or a school-wide one, merged at 864468ba) and 8cd97cae, where "
    "both fee runs, the school's cohort route (324401bc) and the engine's own generate route, "
    "bill only within the caller's branches and at the price list's branch. The school app bills through the cohort route "
    "(FinPro v0.7.24, school-fe 9686919), committed and pending release"
)
SOURCE_MRD = "XVS Module Requirements Document v2.93 | Module 17"
TEST_EVIDENCE = (
    "Verified by the schools.core.fal suite (231 tests, all passing, twelve of them the cohort "
    "route's), by vs_finance's branch-scope and branch-ledger tests (43 tests, all passing) "
    "and by its branch-reference tests (29 tests, all passing, twelve of them the "
    "engine fee run's and the dunning run's). The receipt's deposit rule rests on "
    "test_a_customer_receipt_deposited_into_it, test_an_invoice_payment_deposited_into_it and "
    "test_a_school_wide_customers_receipt_may_land_in_any_bank_she_reaches in vs_finance's "
    "tests_ledger_reach.py, read at 7860ea86 and not rerun. Nothing here claims deployment."
)

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

M17_DIR = "17-billing-and-invoicing"
M17_STEM = "XVS_M17_Billing_and_Invoicing_Functional_Requirements_Document"
M17_SOURCE, M17_TARGET = "1.6", "1.7"

M17_SUMMARY = (
    "Minor revision. A school raises its fee invoices through the cohort route in the finance "
    "abstraction layer, and the school app no longer bills every active customer. A named "
    "child outside the caller's branches answers 404 and nothing is billed, where the route "
    "had checked only that the child was the school's, so a bursar at one branch could bill "
    "a child at another by id. A fee structure that carries a branch is that branch's price "
    "list: a run naming a child at another branch is refused with 409 WRONG_BRANCH, even for "
    "a caller who sees every branch, and a school-wide structure bills every branch. The "
    "link-term route answers a read under finance.feestructure.view, the dry run returns the "
    "due date the run writes, and the fee structure read carries branch_id. The school app "
    "fills the finance console's generation slot with a cohort drawer (FinPro v0.7.24, "
    "school-fe 9686919), committed and pending release. The engine's own generate route, "
    "which the console still uses, is bounded the same way: all_active bills only the active "
    "customers the caller reaches, where it had billed every active customer in the books, "
    "and a structure with a branch refuses a named customer filed anywhere else, school-wide "
    "included, with 409 WRONG_BRANCH, while all_active from it selects only that branch's "
    "customers. A dunning run started by a branch-bound caller stays in their reach, and the "
    "scheduled run stays school-wide (FR-008). FR-002, FR-008, FR-011, FR-013, FR-014 and "
    "sections 6 and 7 follow. Needs Attention restates the two engine-route gaps around that "
    "route's remaining reach. FR-009 records that aging and the AR "
    "reconciliation answer a branch-bound reader for their branches, and FR-013 drops its "
    "limit that the financial statements aggregate across branches, which Module 19 no "
    "longer does. A receipt, recorded against a customer or against one invoice, works out "
    "the branch it inherits before its deposit account is named, and a deposit account behind "
    "another branch's bank account is refused 400 naming the branch, \"Deposit it into\" "
    "that branch's account or a school-wide one, while a school-wide customer's receipt may "
    "land in any bank the caller reaches (0c59e23d; FR-004 and section 7). Section 7 also "
    "records the customer receipt route as the POST that records a receipt, where it had "
    "listed a GET of a receipt document. Traceability is checked against MRD v2.93, where Module 17 still carries "
    "20 entries. " + TEST_EVIDENCE
)


def gaps_box(doc):
    """The one-cell FURTHER GAPS box in Needs Attention."""
    hits = [t for t in doc.tables
            if t.rows[0].cells[0].text.strip().startswith("FURTHER GAPS")]
    if len(hits) != 1:
        raise ValueError(f"FURTHER GAPS heads {len(hits)} tables")
    return hits[0].rows[0].cells[0]


def patch_control(doc) -> None:
    set_cover_version(doc, M17_SOURCE, M17_TARGET)
    set_control(doc, "Version", M17_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", SOURCE_MRD)


def patch_engine_generation(doc) -> None:
    """FR-002: the engine route's two bounds, and its reach now that a school bills elsewhere."""
    fr002 = fr_table(doc, "FR-002")
    append_to(
        fr002, "Current evidence",
        " The route takes a named customer list or all_active, and who it bills is bounded "
        "twice for both. The caller's branch reach: all_active selects the active customers "
        "the caller reaches, their branches' and the school-wide ones, and a named customer "
        "outside that reach answers 404 exactly as an unknown code does. The structure's "
        "branch: a structure with a branch is that branch's price list, the rule the cohort "
        "route holds for children (FR-011), so a named customer filed anywhere else, "
        "school-wide included, is refused with 409 WRONG_BRANCH, and all_active from it "
        "selects only that branch's customers. A school-wide structure prices every branch "
        "and narrows nothing further.",
    )
    append_to(
        fr002, "Acceptance",
        " FeeRunBillsOnlyWhatTheCallerReachesTests in tests_branch_reference.py proves a "
        "pinned bursar's all_active run bills only her reach, another branch's family named "
        "answers as an unknown code does, a whole-school caller's run is unchanged, a branch "
        "price list refuses a family filed at another branch and a school-wide family, bills "
        "its own branch's family and, with all_active, only that branch, and a single-branch "
        "school bills every family.",
    )
    set_value(
        fr002, "Current limit",
        "Selection of customers is supplied by the caller; the module has no rule about who "
        "should be billed, because that rule belongs to a domain it does not know. The route "
        "serves the console's own books; a school bills through the cohort route (FR-011), "
        "which names children by class and has no bill-everyone form. The route applies no "
        "school's due-date rule (FR-014), and all_active from a school-wide structure bills "
        "every active customer the caller reaches, so a school caller who uses it directly "
        "over the API bills every family in reach, whatever class the structure was for, on "
        "a date its own rule did not choose (Section 9).",
    )


def patch_dunning(doc) -> None:
    """FR-008: a caller's dunning run stays in their reach; the scheduled run does not narrow."""
    fr008 = fr_table(doc, "FR-008")
    append_to(
        fr008, "Current evidence",
        " A run started by a branch-bound caller raises notices only on the invoices in their "
        "reach, their branches' and the school-wide ones; the scheduled daily run covers the "
        "whole entity.",
    )
    append_to(
        fr008, "Acceptance",
        " A pinned bursar's run chases only her reach and a whole-school caller's run is "
        "unchanged (BulkRunsNarrowToTheCallersReachTests in tests_branch_reference.py).",
    )


def patch_receivable_reports(doc) -> None:
    """FR-009: aging and the AR reconciliation answer a branch-bound reader for their branches."""
    fr009 = fr_table(doc, "FR-009")
    append_to(
        fr009, "Current evidence",
        " A reader bound to branches is answered for those branches and the school-wide "
        "documents beside them: aging, with its CSV, XLSX and PDF exports, narrows as the "
        "lists it is built from do, and the AR reconciliation compares their receivables "
        "with a control account built from their own and school-wide journals (Module 19). "
        "Each answer carries narrowed, and a narrowed export's subtitle says it covers the "
        "reader's branches only.",
    )
    append_to(
        fr009, "Acceptance",
        " A whole-school reader's figures are unchanged, and a branch reader's reconciliation "
        "still reconciles, because every journal balances on its own "
        "(test_the_branch_ar_reconciliation_compares_like_with_like in tests_branch_ledger.py).",
    )


def patch_cohort_billing(doc) -> None:
    """FR-011: the cohort route is the school's billing route."""
    fr011 = fr_table(doc, "FR-011")
    edit_value(
        fr011, "Current evidence",
        "Both steps are served at /v1/school-finance/: POST fee-structures/{id}/link-term/ "
        "under finance.feestructure.edit,",
        "Both steps are served at /v1/school-finance/: fee-structures/{id}/link-term/, which "
        "reads the link with GET under finance.feestructure.view and sets it with POST under "
        "finance.feestructure.edit, and answers linked false for a structure with no term "
        "rather than 404, so a billing screen can name the term and load its classes before "
        "anybody is chosen;",
    )
    edit_value(
        fr011, "Current evidence",
        "The structure is resolved inside the caller's tenant and branches, so another "
        "tenant's or another branch's answers 404.",
        "The structure is resolved inside the caller's tenant and branches, so another "
        "tenant's or another branch's answers 404. Each named child is then checked against "
        "the caller's branch reach before anything is priced: a child the caller cannot see "
        "answers 404 and nothing is billed. A structure that carries a branch is that "
        "branch's price list, so a run naming a child at another branch is refused with 409 "
        "WRONG_BRANCH even for a caller who sees every branch, while a school-wide structure, "
        "with no branch, bills children at every branch. The engine's own generate route holds "
        "the same price-list rule for the customers it bills (FR-002).",
    )
    edit_value(
        fr011, "Current evidence",
        "Each bill carries the due date the school's fee due policy resolves (FR-014),",
        "Each bill carries the due date the school's fee due policy resolves (FR-014), and "
        "the answer returns it as due_date, a dry run included, resolved against the "
        "structure's own term by the code that posts, so the date a bursar is shown before "
        "billing is the date written;",
    )
    edit_value(
        fr011, "Current evidence",
        "so a CodeX operator acting inside a school's session raises it as that school's act.",
        "so a CodeX operator acting inside a school's session raises it as that school's act. "
        "This is the route a school bills through, and the school app no longer bills every "
        "active customer. It fills the finance console's generation slot (FinPro v0.7.24) "
        "with a cohort drawer (school-fe 9686919) that names the term the structure bills, "
        "linking one where there is none, lists that year's classes with their pupil counts, "
        "limited to the structure's branch where it has one, previews the pupils to bill, "
        "the total with tax and the due date, and then bills. That frontend work is "
        "committed and pending release; nothing here claims it is deployed.",
    )
    edit_value(
        fr011, "Acceptance",
        "A repeated run bills nobody twice",
        "There is no bill-everyone default: an empty cohort is refused rather than read as "
        "the whole roll. A repeated run bills nobody twice",
    )
    edit_value(
        fr011, "Acceptance",
        "The refusals are 409 TERM_NOT_LINKED for an unlinked structure,",
        "The refusals are 404 for a named child outside the caller's branches, 409 "
        "WRONG_BRANCH for a child at a branch other than the structure's, 409 "
        "TERM_NOT_LINKED for an unlinked structure,",
    )
    edit_value(
        fr011, "Acceptance",
        "an empty cohort and a child named twice;",
        "an empty cohort and a child named twice, a branch bursar refused another branch's "
        "child and billing her own, a school-wide bursar billing every branch, a branch "
        "price list refusing another branch's child while a school-wide one bills every "
        "branch, the link read under the view key alone, the unlinked and linked reads, "
        "another school's and another branch's read answering 404, and a preview carrying "
        "the due date the run writes;",
    )


def patch_branch_lookup(doc) -> None:
    """FR-013: the cohort run narrows the children it names; the statements gap is gone."""
    fr013 = fr_table(doc, "FR-013")
    edit_value(
        fr013, "Current evidence",
        "The finance abstraction layer's fee structure lookup applies the identical rule.",
        "The finance abstraction layer's fee structure lookup applies the identical rule, and "
        "its cohort run holds each child it is given to the caller's branches as well "
        "(FR-011). The bulk runs keep to the same reach: the engine's fee run (FR-002) and a "
        "dunning run a branch-bound caller starts (FR-008).",
    )
    set_value(
        fr013, "Current limit",
        "The rule governs reference lookups, lists and the receivable reports (FR-009). The "
        "financial statements a branch-bound reader sees are built from their own and "
        "school-wide journals, and belong to Module 19.",
    )


def patch_due_policy(doc) -> None:
    """FR-014: the rule reaches the school's own billing route, and the preview shows it."""
    fr014 = fr_table(doc, "FR-014")
    edit_value(
        fr014, "Current evidence",
        "Cohort billing (FR-011) hands fees.generate_invoices the resolved date as a plain date.",
        "Cohort billing (FR-011) hands fees.generate_invoices the resolved date as a plain "
        "date, and returns it on a dry run, so the date previewed is the date written.",
    )
    append_to(
        fr014, "Acceptance",
        " test_the_preview_carries_the_due_date_the_run_will_write in test_route.py proves the "
        "cohort preview returns the date the run writes.",
    )
    set_value(
        fr014, "Current limit",
        "The policy reaches bills raised through cohort billing, which is the route the school "
        "app bills through. This module's own generation route (FR-002) still takes a "
        "caller-supplied date or the entity's default payment terms: the console bills its "
        "own books through it, and a school caller who reaches it directly over the API gets "
        "a date its own rule did not choose. Changing the rule dates future bills and leaves "
        "bills already raised as they were.",
    )


def patch_data_and_api(doc) -> None:
    """Sections 6 and 7: branch_id, the link read, the preview's date, and two refusals."""
    fee_row = row_labelled(doc, "FeeStructure, FeeItem")
    edit_cell(
        fee_row.cells[-1],
        "carried on every invoice it raises as FEE:<code>.",
        "carried on every invoice it raises as FEE:<code>. The read carries branch_id, the "
        "branch whose price list it is, or null for a structure the whole school shares.",
    )

    engine = row_labelled(doc, "POST /fee-structures/{id}/generate/")
    keep_format(
        engine.cells[-1],
        "Bulk-raise posted invoices for a named customer list or all_active, within the "
        "caller's branches and, from a branch's structure, that branch's customers only. The "
        "console's route; a school bills through generate-invoices below.",
    )
    link = row_labelled(doc, "POST /v1/school-finance/fee-structures/{id}/link-term/")
    keep_format(link.cells[0], "GET, POST /v1/school-finance/fee-structures/{id}/link-term/")
    keep_format(
        link.cells[-1],
        "Read the academic session and term a fee structure bills (finance.feestructure.view), "
        "or set it (finance.feestructure.edit).",
    )
    cohort = row_labelled(doc, "POST /v1/school-finance/fee-structures/{id}/generate-invoices/")
    keep_format(
        cohort.cells[-1],
        "Bill a named cohort, or preview the run with dry_run. The answer carries the due "
        "date the run writes.",
    )

    errors = table_headed(doc, "Condition or route", "Answer")
    anchor = next(r for r in errors.rows
                  if r.cells[0].text.strip().startswith("Billing a cohort from a structure"))
    added = insert_row_after(anchor, [
        "Billing from a branch's structure a child or customer filed anywhere else",
        "409 WRONG_BRANCH, even for a caller who sees every branch, on the cohort route and "
        "the engine's generate route alike; the engine route counts a school-wide customer "
        "as filed elsewhere.",
    ])
    insert_row_after(added, [
        "Naming a child or customer outside the caller's branches in a fee run",
        "404, as for an unknown reference, and nothing is billed.",
    ])
    dunning = row_labelled(doc, "POST /dunning/generate/")
    keep_format(
        dunning.cells[-1],
        "Raise notices for everything overdue, within the caller's branches; the scheduled "
        "run covers the whole school.",
    )


def patch_receipt_deposit(doc) -> None:
    """FR-004 and section 7: a receipt deposits into its own branch's bank (D62)."""
    fr004 = fr_table(doc, "FR-004")
    append_to(
        fr004, "Current evidence",
        " A receipt is recorded against a customer (POST /customers/{id}/receipt/, allocated to "
        "their open invoices oldest first) or against one invoice (POST /invoices/{id}/pay/). "
        "Either takes its branch from the chain it continues, the customer's or the invoice's, "
        "and works it out before the deposit account is resolved. A deposit account behind a "
        "bank account must then be that branch's or a school-wide one: another branch's bank "
        "ledger is refused 400 on deposit_account, naming the branch, \"This receipt belongs to "
        "Ikeja Branch. Deposit it into an Ikeja Branch account or a school-wide one.\", and "
        "nothing is written (Module 19 FR-020). A school-wide customer's receipt may land in any "
        "bank the caller reaches, and a ledger account behind no bank account is unaffected.",
    )
    append_to(
        fr004, "Acceptance",
        " Mrs Okafor, who covers Ikeja and Lekki, cannot deposit an Ikeja family's receipt, or "
        "a payment against an Ikeja invoice, into Lekki's bank ledger, and can deposit a "
        "school-wide customer's receipt into either (test_a_customer_receipt_deposited_into_it, "
        "test_an_invoice_payment_deposited_into_it and "
        "test_a_school_wide_customers_receipt_may_land_in_any_bank_she_reaches in vs_finance's "
        "tests_ledger_reach.py).",
    )

    receipt = row_labelled(doc, "GET /customers/{id}/receipt/")
    keep_format(receipt.cells[0], "POST /customers/{id}/receipt/")
    keep_format(receipt.cells[-1],
                "Record a receipt for the customer, allocated to their open invoices oldest "
                "first. The deposit account follows the customer's branch.")

    errors = table_headed(doc, "Condition or route", "Answer")
    insert_row_after(row(errors, "Allocating more than the receipt holds"), [
        "Depositing a branch's receipt into the ledger account behind another branch's bank "
        "account",
        "400 on deposit_account, naming the branch: \"This receipt belongs to Ikeja Branch. "
        "Deposit it into an Ikeja Branch account or a school-wide one.\" Nothing is written.",
    ])


def patch_needs_attention(doc) -> None:
    """Section 9: the engine route's reach, restated as it stands."""
    edit_box(gaps_box(doc), [
        ("replace", "• The module holds no rule about who should be billed",
         "• The module holds no rule about who should be billed, by design, so a run bills "
         "exactly the selection it is given. The school app names that selection by class "
         "through the cohort route, which has no bill-everyone form. The engine's own "
         "generation route still accepts all_active, bounded by the caller's branches and "
         "by a branch structure's branch: the console uses it for its own books, but a "
         "school caller holding finance.feestructure.generate can still reach it over the "
         "API, and all_active from a school-wide structure bills every active family in "
         "their reach, whatever class the structure was for."),
        ("replace", "• The school fee due policy is applied only by cohort billing",
         "• The school fee due policy is applied by cohort billing, which is how the school "
         "app bills. A structure billed through this module's own generation route still "
         "takes the caller's date or the entity's default payment terms, so a school caller "
         "who uses that route directly gets a date its own rule did not choose."),
    ])


def patch_traceability(doc) -> None:
    """Section 10: checked against the latest MRD; the count is unchanged."""
    edit_paragraph(doc, "Module 17 carries 20 capability entries", "MRD v2.76", "MRD v2.93")
    lookup = row_labelled(doc, "Branch-scoped customer, invoice and fee-structure lookup")
    keep_format(lookup.cells[-1], "FR-013, FR-011")


def keep_requirement_rows_whole(doc) -> None:
    """No row of a requirement, the change log or any other table breaks across a page.

    A row split at a page foot leaves a line or two of a limit, or of a version
    summary, stranded above the page break, where it reads as a separate row.
    """
    for table in doc.tables:
        if table.rows[0].cells[0].text.strip().startswith("FR-"):
            keep_rows_whole(table)
    keep_rows_whole(change_log_table(doc))
    # Route, error and dependency rows stay whole too, so none prints half on each page.
    for table in doc.tables[1:]:
        if len(table.rows[0].cells) > 1:
            keep_rows_whole(table)
    # The change log opens its own page, so its first, longest row never has to follow a
    # table's tail down a page.
    heading = [p for p in doc.paragraphs if p.text.strip() == "11. Change Log"]
    if len(heading) != 1:
        raise ValueError("Section 11's heading is not unique")
    heading[0].paragraph_format.page_break_before = True


def patch_m17() -> None:
    require_newest(str(ROOT / "functional-requirements" / M17_DIR / f"{M17_STEM}_v*.docx"),
                   M17_SOURCE)
    doc = Document(str(frd_path(M17_DIR, M17_STEM, M17_SOURCE)))
    patch_control(doc)
    patch_engine_generation(doc)
    patch_dunning(doc)
    patch_receivable_reports(doc)
    patch_cohort_billing(doc)
    patch_branch_lookup(doc)
    patch_due_policy(doc)
    patch_data_and_api(doc)
    patch_receipt_deposit(doc)
    patch_needs_attention(doc)
    patch_traceability(doc)

    log_change(doc, M17_TARGET, M17_SUMMARY)
    assert_absent_outside_log(
        doc,
        "Version: 1.6",
        "MRD v2.76",
        "a bulk run bills exactly the selection it is given",
        "applied only by cohort billing",
        "POST fee-structures/{id}/link-term/ under finance.feestructure.edit",
        "still aggregate across branches",
        "without the caller's branch narrowing",
        "bills every other branch's families",
        "GET /customers/{id}/receipt/",
    )
    keep_requirement_rows_whole(doc)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M17_DIR, M17_STEM, M17_TARGET),
           f"{M17_STEM.replace('_', ' ')} v{M17_TARGET}", M17_TARGET)


def main() -> None:
    patch_m17()


if __name__ == "__main__":
    main()
