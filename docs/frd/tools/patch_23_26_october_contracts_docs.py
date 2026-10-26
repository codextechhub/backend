#!/usr/bin/env python3
"""Version the MRD and affected FRDs for the 23 to 26 October 2026 queue."""

from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

import patch_record_history_docs as common
from patch_record_history_docs import add_fr, append_rows, finish, keep_format, log_change, set_control, set_cover_version
from patch_staff_id_and_auth_events_docs import normalise_change_log, repair_ooxml
from patch_22_october_contracts_docs import restore_clean_change_log_rows, sort_and_deduplicate_change_log
from generate_requirements_documents import assert_no_em_dash, shrink_inherited_media, update_extended_title


ROOT = Path(__file__).resolve().parents[1]
REVIEW_DATE = "26 October 2026"
MRD_SOURCE = "2.114"
MRD_TARGET = "2.115"
CODE_BASELINE = (
    "Backend main at 2fb1f1fe (26 October 2026), covering the committed behavior "
    "through that date. Frontend evidence is limited to the commits named in the "
    "requirements queue. Production deployment and adoption are not inferred."
)
common.REVIEW_DATE = REVIEW_DATE


DOCS = {
    "M01": ("01-school-and-branch-management", "XVS_M01_School_and_Branch_Management_Functional_Requirements_Document", "1.35", "1.35.1"),
    "M03": ("03-identity-team-and-organogram", "XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document", "1.24.1", "1.25"),
    "M04": ("04-roles-and-permissions-rbac", "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document", "1.42", "1.43"),
    "M05": ("05-audit-and-activity-logging", "XVS_M05_Audit_and_Activity_Logging_Functional_Requirements_Document", "1.13", "1.13.1"),
    "M06": ("06-configuration-and-capability", "XVS_M06_Configuration_and_Capability_Management_Functional_Requirements_Document", "1.11", "1.12"),
    "M07": ("07-workflow-and-approval-engine", "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document", "1.27", "1.28"),
    "M08": ("08-notifications-and-delivery", "XVS_M08_Notifications_and_Delivery_Functional_Requirements_Document", "1.19", "1.20"),
    "M11": ("11-student-management", "XVS_M11_Student_Management_Functional_Requirements_Document", "2.17", "2.18"),
    "M12": ("12-staff-management", "XVS_M12_Staff_Management_Functional_Requirements_Document", "2.21", "2.22"),
    "M15": ("15-attendance-management", "XVS_M15_Attendance_Management_Functional_Requirements_Document", "1.0.1", "1.1"),
    "M17": ("17-billing-and-invoicing", "XVS_M17_Billing_and_Invoicing_Functional_Requirements_Document", "1.14", "1.15"),
    "M18": ("18-payments-and-collections", "XVS_M18_Payments_and_Collections_Functional_Requirements_Document", "1.19", "1.20"),
    "M19": ("19-finance-and-accounting", "XVS_M19_Finance_and_Accounting_Functional_Requirements_Document", "1.23", "1.24"),
    "M20": ("20-adjustments-and-concessions", "XVS_M20_Adjustments_and_Concessions_Functional_Requirements_Document", "1.12", "1.13"),
    "M27": ("27-student-portal", "XVS_M27_Student_Portal_Functional_Requirements_Document", "1.1", "1.2"),
    "M28": ("28-parent-portal", "XVS_M28_Parent_Portal_Functional_Requirements_Document", "1.4", "1.5"),
    "M30": ("30-system-health-and-monitoring", "XVS_M30_System_Health_and_Monitoring_Functional_Requirements_Document", "1.6", "1.6.1"),
}


REQUIREMENTS = {
    "M01": ("FR-025  Provision School Services and Brand School Delivery", [
        ("Description", "Provision the school-owned approval and delivery seams needed by leave, email, push, native sign-in, and account closure without moving their domain rules into school setup."),
        ("Provisioning", "A new school receives its one-stage line-manager leave route with the rest of its school setup. Existing schools receive a missing leave route from the deploy command, and an existing route of any shape is never replaced."),
        ("Branding", "Delivery resolves the recipient's school at send time. School email uses the school's name as the From display name and its uploaded logo alone in the header; the address remains the platform address, a missing logo falls back to XVS, and the footer remains Powered by XVS. Platform mail keeps CodeX branding."),
        ("Public and signed-in seams", "The school's public name and logo routes let a store app resolve one school by its code without listing schools. The school exposes the shared native-session, phone-push, app-version, and account-closure routes owned by Modules 3, 8, 11, 27, 28, and 30."),
        ("Acceptance", "A newly created school can file leave through its own route; an email to Lagoon View Academy shows that school name and logo; an unknown school code reveals no list; and the shared endpoints retain their owning modules' authentication, tenant, branch, and Field Access rules."),
        ("Current limit", "The school setup proves the backend seams only. Store-app use, production phone delivery, and account-closure screens remain separate evidence."),
    ]),
    "M03": ("FR-029  Keep Native and Browser Sessions Recoverable and Close Eligible Accounts", [
        ("Description", "Run browser and store-app sessions through one rotating session model, allow a lost refresh answer to recover safely, and let the product decide which people may close their own account."),
        ("Native origin", "An allowed NATIVE_APP_ORIGINS request must ask for session=device. Login and device refresh return refresh in the body with no cookie; refresh and sign-out accept it in the body without CSRF. Other origins keep the Secure HttpOnly cookie, require CSRF, and ignore a body refresh. Staff cannot use a store-app device session."),
        ("Parent sessions", "A guardian may hold a parent-app device session for parents.portal.session_days, default 30 and valid from 1 to 90, sliding on use. A parent-only browser sign-in names /parent/ as device_app; a staff member who is also a parent keeps the staff app as the browser destination and may separately use the parent app."),
        ("Refresh rotation", "LoginSession remembers the current and immediately previous refresh jti and the rotation time. Within SESSION_REFRESH_GRACE_SECONDS, default 120, the previous credential returns the existing current credential and a fresh access token. After that window, or for an older device credential, reuse ends the session as CREDENTIAL_REUSED, records TOKEN_REVOKED, and ends that person's impersonations. TOKEN_STALE is retired."),
        ("Account closure", "The product policy in vs_user.services.account_closure decides immediate closure, tenant review, or refusal. Closure deactivates the account, makes the password unusable, ends every browser and device session as ACCOUNT_CLOSED, forgets browser and phone push devices, and records ACCOUNT_DEACTIVATED with origin SELF_SERVICE or TENANT_REVIEW. Reactivation requires forgot password before sign-in."),
        ("Routes and refusals", "GET and POST /v1/user/me/account-closure/ serve the signed-in flow. The public request and confirmation routes use an indistinguishable request answer and a single-use 24-hour link. Review list, close, and decline routes are tenant and branch scoped. Impersonation is refused. TOKEN_EXPIRED, TOKEN_REVOKED, and TOKEN_INVALID remain the native-client terminal codes."),
        ("Acceptance", "Two racing refreshes converge on one current credential; a delayed prior credential inside the window succeeds; reuse after it ends the traceable session; a deleted user answers TOKEN_INVALID instead of 500; and an eligible account closure ends every way that account can authenticate."),
    ]),
    "M04": ("FR-035  Expose Pending Grants and Enforce the New Field Contracts", [
        ("Description", "Expose the narrow state a role screen needs for pending grants and register every new leave, account-closure, finance-account, refund, and virtual-account field at its owning boundary."),
        ("Pending grants", "GET /rbac/tenants/{slug}/role-grant-requests/ lists only PENDING rbac.role_grant requests, filters by role key or id and numeric user id, paginates, requires the role-view keys, and applies the role-holder branch reach. Each row carries id, status, user_id, role_key, branch_id, and submitted_at only. A pending request blocks Give only at the same role reach."),
        ("Leave fields", "school.leave registers leave_type and dates as Read and Write open on create, days and note as Read and Write, and status and balance as read-only. The owner always reads and files their own leave. The switches reach list, echo, as-at, preview, filing warnings, section and audit summaries, approval cards, registers, and on_leave_until."),
        ("Account closure fields", "school.account_closures registers read-only person, kind, channel, requested_at, status, decided_at, and decided_by, enforced on the review list, close and decline actions, and reviewer notices."),
        ("Finance fields", "finance.account.managed_by is read-only in Ledger; finance.account.ifrs_line is the statutory report line and is accepted on create. payments.virtual_account.account_number and account_name govern the staff invoice view. The refund bank_account body key is a Write alias only and no longer a Read alias."),
        ("Acceptance", "A branch-bound role reader sees only pending requests for people in reach; a hidden leave or closure field disappears from every alternate path; a reader without the virtual-account number sees neither the pupil account nor the school's bank block; and a hidden managed_by value cannot reveal that a zero-balance row is system managed."),
    ]),
    "M05": ("FR-019  Record Leave and Account Closure Evidence Without Leaking Protected Fields", [
        ("Description", "Keep leave and self-service account closure actions readable in the append-only trail while applying the owning module's field decisions to summaries and notices."),
        ("Leave evidence", "Leave filing, change, and workflow evidence names protected leave words through the school.leave resource. The approval card marks the leave type and uses Leave request as its subtitle. Rebuilding approval snapshots upgrades in-flight cards without rewriting their decisions."),
        ("Closure evidence", "ACCOUNT_DEACTIVATED records metadata.origin as SELF_SERVICE or TENANT_REVIEW. The trail retains the subject and decision even though the user account is inactive, while reviewer-facing payloads apply school.account_closures Field Access."),
        ("Acceptance", "A role that cannot read leave dates finds none in an audit summary or approval card; an account closure remains attributable after deactivation; and an audit write failure still cannot roll back the business change."),
    ]),
    "M06": ("FR-024  Configure Parent Sessions and Family Card Payments", [
        ("Description", "Declare the school-wide settings that control parent device-session length and whether families may start card payments."),
        ("Parent session setting", "parents.portal.session_days is INTEGER, default 30, valid from 1 to 90. A missing or out-of-range value resolves to 30, and each successful parent-device use starts the period again."),
        ("Family card setting", "fees.family_card_payments_enabled is BOOLEAN, default false. The school finance settings route exposes family_card_payments and permits a change only to school.fees.update with whole-school reach; a branch-bound bursar is refused."),
        ("Business rules", "The card switch changes family and public invoice checkout only. It does not stop virtual-account issuance, transfer booking, staff collections, terminals, refunds, or payroll. The minimum part-payment applies to card only."),
        ("Acceptance", "A parent session lasts the configured number of sliding days; an invalid setting falls back to 30; card payment is off for a new school; and a branch-bound settings writer cannot change the whole-school choice."),
        ("Current limit", "The generic configuration store carries Parent App Sign-In Length, but no school screen edits it yet."),
    ]),
    "M07": ("FR-036  Route Leave to the Subject's Manager and Provision Hand Journals", [
        ("Description", "Resolve leave from the person the request concerns and give every school an explicit, editable route for hand journals."),
        ("Leave routing", "A document type that declares request_for starts an organogram climb from that person, not from the administrator who filed it. The standard school leave route has one DIRECT_MANAGER organogram stage, school scope, never skips, and never self-approves. The parked sentence names the person whose manager is missing, and that person reads their own approval card."),
        ("Journal routing", "Each school owns a published finance.journal route named Journal entry approval, code standard, with no steps or approver group. It is created with the books and by provision_journal_routes on deploy, never replacing an existing route. Schools add their steps in Workflow Templates."),
        ("No-step behavior", "A hand journal with no configured steps answers 409 APPROVAL_NOT_CONFIGURED until confirm_without_approval is supplied, then records that choice. Once steps exist, it submits and posts only after approval. System journals raised by invoices, receipts, or payroll bypass this route and follow their source document."),
        ("Acceptance", "Leave filed by an administrator reaches the absent person's manager; missing management parks with the exact remedy; a new school sees its Journal entry approval template; and no-step confirmation is explicit rather than silently treating the journal as approved."),
    ]),
    "M08": ("FR-018  Brand School Email and Deliver to Browser and Phone Devices", [
        ("Description", "Resolve school email branding at delivery, deliver one push message across browser and phone transports, and send the new account and transfer notices without exposing protected fields."),
        ("Email branding", "A recipient in a school tenant receives the school's name as the From display name at the platform sending address. The header shows only the uploaded school logo in a white fitted square, or the XVS logo from the school's subdomain, and the footer says Powered by XVS. Platform mail retains CodeX System and CodeX Secure notification. Staff-authored HTML is unchanged."),
        ("Phone push", "The push channel builds one message and sends it to browser subscriptions and native devices through Firebase Cloud Messaging, including Apple delivery. Registration is available only when FCM_SERVICE_ACCOUNT_FILE is configured, otherwise 503 PUSH_NOT_CONFIGURED. Registering the same token refreshes it; deletion forgets only the caller's row and works while phone push is off."),
        ("Device lifecycle", "One row belongs to one account on one phone. An account may retain four phones and four browsers separately; a fifth drops the least recently seen of that kind. A 404 or UNREGISTERED provider answer deletes the token; another failure keeps it and never stops the remaining devices. Signing one account out removes only its devices."),
        ("New events", "user.deletion_requested and user.account_deleted notify authorized closure reviewers by in-app and email, using Someone where person is hidden; user.account_closure_link sends the confirmation link. A booked pupil transfer tells that pupil, while pupil payments are on, and every guardian with a parent-app account by in-app and push, with no amount on the lock screen."),
        ("Acceptance", "A school email carries the school's own identity; two children sharing one phone each receive their own push; a forgotten token is removed without blocking other delivery; closure notices obey Field Access; and a transfer notification opens the receipt."),
        ("Current limit", "No store app registers a phone yet, and phone delivery has not been tried against real Firebase."),
    ]),
    "M11": ("FR-029  Decide School Account Closure for Guardians and Pupils", [
        ("Description", "Supply the XVS policy that decides which school accounts close immediately, which wait for school review, and which are refused."),
        ("Policy", "A guardian-only account closes immediately. A pupil request enters tenant review. Staff and platform accounts, including a teacher who is also a parent, are refused 403 ACCOUNT_CLOSURE_NOT_OFFERED with an administrator pointer. Impersonation is refused."),
        ("Record retention", "Closing the account deactivates the User and ends access, but the Guardian, Student, relationships, enrollment, finance records, notices, and audit evidence the school must keep are untouched. An administrator may reactivate the account, after which the person uses forgot password."),
        ("Review scope", "school.account_closures.view and update govern the pending list and decisions. A branch-bound reviewer sees a pupil in their branch, or a guardian whose linked children are all in one branch they can read. A request with no branch is visible only to a whole-school reviewer. Cross-tenant ids answer 404."),
        ("Acceptance", "A guardian closes at once; a pupil receives one pending request when asking twice; a teacher-parent is refused; a reviewer cannot close another school's request; and a closure never deletes the school records attached to the person."),
        ("Current limit", "The school application has no account-deletion page or review list, and the store apps have no deletion screens."),
    ]),
    "M12": ("FR-028  File Leave Through the Live Ladder and Count a Draft Payslip in Its Year", [
        ("Description", "Give staff a live leave preview and approval path, apply Field Access to every leave surface, and keep a draft payslip's year-to-date figures internally consistent."),
        ("Leave preview", "GET /v1/i/me/staff/{id}/leave/preview/ takes start_date, end_date, optional leave_type and exclude, uses working days less closures at the person's branch, and returns days, resumption_date, allowance, taken, pending, remaining, remaining_after, over_allowance_by, and session. It shares own and other-person leave permissions with filing, returns zero days as data, and refuses bad order, over 366 days, or an unknown type."),
        ("Leave filing", "Every school has a one-stage direct-manager route. Filing remains successful when nobody can approve and returns NO_APPROVER with the named organogram remedy. A person's own leave is always readable and writable by that person; other readers and writers follow school.leave Field Access across lists, history, preview, approval, registers, and directory state."),
        ("Payslip year to date", "A payslip's YTD includes its own payroll line exactly once whatever the run status. Other runs count only once posted or paid, matching the PAYE working. The yearly tax summary remains limited to posted and paid runs."),
        ("Acceptance", "The form can show allowance, taken, pending, remaining, requested days, remaining after, and return date before filing; an administrator's filing routes to the subject's manager; a hidden leave field stays hidden on every alternate read; and a draft payslip shows the YTD its paid form will show."),
    ]),
    "M15": ("FR-011  Apply Leave Field Decisions to Attendance Surfaces", [
        ("Description", "Show leave in pupil, staff, and class attendance only to the extent the reader may read the corresponding leave fields."),
        ("Business rules", "Register leave blocks and on-leave labels use school.leave. Dates require dates; type-specific balances and days taken require leave_type; totals and day counts require days; status requires status. The owner rule remains with the leave module and does not widen a colleague's attendance view."),
        ("Acceptance", "A register reader without leave dates cannot infer an end date; a reader without leave type sees no type label or per-type balance; a person still reads their own leave in full; and the attendance state remains correct when a protected label is omitted."),
    ]),
    "M17": ("FR-016  Show Owing and Credit Separately and Instruct Family Transfers", [
        ("Description", "Present what a family owes without netting away held credit, and render the payment method the school currently offers."),
        ("Balances", "Portal fee payloads carry owing from open invoices and debit notes, credit from unapplied credit, and balance as owing minus credit. Status is OWING whenever owing is positive, IN_CREDIT only when nothing is owed and credit is held, otherwise CLEAR. A family total sums each child's owing and credit separately, so one child's credit never hides another's debt."),
        ("Payment methods", "fees.family_card_payments_enabled defaults off. While off, family and public pay routes start no checkout and answer CARD_PAYMENTS_OFF after the pupil-payment gate. A pupil's dedicated transfer account leads the fee screen. Parent pay-all is replaced by one account per child; debit notes are paid by transfer, not card."),
        ("Invoice instructions", "When card is off and the pupil account exists, invoice email, document, PDF, and family screen show its bank, account number, account name, and automatic-matching instruction, and omit the school's bank block. Without an account they show contact the office and the school's bank. With card on they keep online payment. Empty email button destinations are dropped."),
        ("Receipt wording", "Account balance after this payment reads an amount owing, an amount in credit, or zero, with no minus sign or Dr or Cr. The per-invoice balance remains unchanged."),
        ("Acceptance", "A child owing 100 with 20 credit still shows 100 owing and 20 held; a sibling's credit does not reduce another child's amount; a transfer-only invoice carries the child's account; and hidden account-number Field Access removes both the pupil account and school bank fallback from the staff invoice view."),
    ]),
    "M18": ("FR-028  Issue and Operate Family Transfer Accounts While Card Is Optional", [
        ("Description", "Issue each enrolled pupil a dedicated Paystack transfer account, keep family card payment optional, and notify the family when a transfer books."),
        ("Issuance", "Issue the account on every move into enrolled, including readmission, and at the first bill. The daily fal-issue-missing-transfer-accounts sweep runs at 01:30, makes at most 300 provider asks, asks a refusing branch once, and skips a school after three failures for the run. fees.student_payment_enabled governs pupil payment screens, not issuance."),
        ("Transfer booking", "The account belongs to the pupil's finance customer and branch. A booked transfer follows existing webhook verification and ledger booking, then raises the pupil and guardian receipt notice. The apps may poll deposits every five seconds while visible and move from On its way to Received with the receipt link."),
        ("Card switch", "Family card routes honor fees.family_card_payments_enabled, default false. Staff-side collections, terminals, refunds, payroll, and virtual-account transfer booking are unchanged. The minimum part-payment applies to card only."),
        ("Acceptance", "An enrolled pupil receives an account even when pupil payment is off; a missing account produces the setup message rather than a card fallback; a transfer books once and notifies each linked parent account; and a family card attempt is refused while staff collection remains available."),
        ("Current limit", "Paystack rate limits for bulk enrollment are untested, and a real end-to-end transfer has not been run on staging."),
    ]),
    "M19": ("FR-027  Guide the Chart, Gate Hand Journals, and Settle Receivables Clearly", [
        ("Description", "Explain the chart of accounts, separate school-managed and system-managed rows, allocate safe codes, route hand journals explicitly, and keep payroll, receivable, and held-credit totals truthful."),
        ("Chart guide", "Every default account has a one-line purpose back-filled only where description is blank. The chart shows code, name, and tags; detail exposes readable description from an info icon. Every account read carries managed_by or null. Zero-balance managed accounts fold under Managed by the system; a readable nonzero balance moves the row into the main chart. A hidden balance is treated as zero, and hidden managed_by produces the single chart."),
        ("Account creation", "POST /finance/accounts/ may take account_type without code and allocates ten above the highest free code for that type, skipping every default and mapping-reserved code. The next-code preview and suggestions routes require finance.account.create. Generic suggestions come from finance and school suggestions from the FAL. Choosing a reserved code names the kept system account and the next free code, including on bank and petty-cash splits."),
        ("Journal approval", "Each school owns a finance.journal route with no steps until configured. A direct or human-posted draft journal answers APPROVAL_NOT_CONFIGURED until explicitly confirmed without approval; once steps exist it posts only after workflow approval. Journals raised by invoices, receipts, and payroll remain governed by their source."),
        ("Payroll and balances", "A draft payslip includes its own line once in YTD; other runs count only when posted or paid. Budget headlines count the plan's one direction. Owing and credit remain separate. Receipt account balance uses owing, in credit, or zero wording without a sign convention."),
        ("Held credit", "At the same branch and under auto_apply_customer_credit, receipt credit settles a newly posted debit note and eligible held credit settles only bills reopened by voiding a receipt, credit note, concession, refund, or credit transfer. Credit-note credit never settles a debit note. A closed period skips the pairing without blocking the void. Receivable-move voids are excluded."),
        ("Field Access", "finance.account.managed_by and ifrs_line apply end to end; claimant and pay figures remain as documented; virtual-account values on staff invoices use the payments resource. The refund bank_account request name is Write-only under its existing switch."),
        ("Current limits", "Credit held at another branch deliberately remains beside the bill. Invoice posting still fails if held credit cannot be applied in a closed period; voiding does not. The real family-transfer flow remains unproved on staging."),
    ]),
    "M20": ("FR-013  Apply Held Credit When an Adjustment Reopens or Creates a Bill", [
        ("Description", "Keep debit-note debt visible and use eligible held credit when an adjustment creates or reopens a receivable."),
        ("Business rules", "Posting a debit note may use receipt credit at the same branch, never credit-note credit. Voiding a receipt, credit note, concession, refund, or credit transfer applies eligible held credit only to the bills that void reopened and only when auto_apply_customer_credit is on. A closed period skips the pairing and does not block the void. Receivable-move voids are excluded."),
        ("Presentation", "Debit notes contribute to owing without being netted by credit. Parent rows list each child's debit notes and direct payment to that child's transfer account. Account balance after a payment uses owing, in credit, or zero wording."),
        ("Acceptance", "A debit note remains visible even where credit covers it; receipt credit may settle it at the same branch; a cross-branch credit does not; a void in a closed period completes without auto-pairing; and a receivable-move void never runs the rule."),
    ]),
    "M27": ("FR-012  Run the Pupil Portal as a Store App With Transfer-First Fees", [
        ("Description", "Extend the pupil portal from an installable web app to a store-app client with native sessions, phone push, version checks, account closure, transfer-first fees, and recoverable refresh rotation."),
        ("Store app", "The app resolves one school by code, signs in with session=device from an allowed native origin, keeps refresh in the phone keychain, registers its own Firebase token, and asks the public app-versions route whether it is current, may update, or must block. It drops stored refresh on TOKEN_EXPIRED, TOKEN_REVOKED, or TOKEN_INVALID."),
        ("Fees", "Fees lead with owing and show held credit separately. Card availability comes from card_payments. The pupil's transfer account leads the page and deposits are polled while visible; a received transfer opens its receipt. A missing account shows the setup message."),
        ("Account closure", "A pupil may request closure in-app or by the public email link. The request goes to school review, preserves the pupil record, and becomes inaccessible only when the account is closed. Duplicate requests return the same pending request."),
        ("Refresh", "A lost refresh answer or two racing app requests may replay the immediately previous credential inside the grace window and both receive the existing current credential. Reuse after the window ends the session."),
        ("Acceptance", "A below-minimum app version cannot continue; a phone sign-out forgets only that pupil's devices; owing is never reduced by displayed credit; a transfer receipt opens from the notice; and a closure request remains tenant and branch scoped."),
        ("Current limits", "No store app calls the version or phone-registration routes yet, no versions are configured, no deletion screen exists, and real Firebase and transfer delivery remain unproved."),
    ]),
    "M28": ("FR-015  Keep Parents Signed In and Lead Each Child With Their Transfer Account", [
        ("Description", "Give the parent app a sliding device session, native keychain refresh, phone push, version checks, account closure, and a per-child transfer-first fee view."),
        ("Session and destination", "A guardian gets a device session for parents.portal.session_days, default 30, sliding on use. Web and store clients use the device routes. A parent-only school sign-in names /parent/ as device_app and may hand the session across; a teacher-parent keeps staff browser routing and uses a separate parent device session."),
        ("Family fees", "Each child shows owing, held credit, debit notes, and their own transfer account. Family outstanding sums owing and family credit sums credit without netting children. Pay-all stays invoices-only when card is enabled and carries charges_total; while card is off it is replaced by per-child account cards and Copy."),
        ("Native client", "An allowed native origin carries refresh in the body and keychain, registers the account's phone token, checks the public version policy, and recovers the immediately previous refresh credential inside the grace window. A push opens the app attached to the device session, so a teacher-parent's staff push still opens staff."),
        ("Account closure", "A guardian-only account closes immediately, ends every session and device, and preserves the Guardian and child links the school keeps. A teacher-parent is refused self-closure. The public confirmation link follows the same product policy."),
        ("Acceptance", "A parent remains signed in for the configured sliding period; one child's credit cannot hide another's debt; two accounts on one phone receive separate pushes; a below-minimum app blocks; and immediate guardian closure ends all access without deleting school records."),
        ("Current limits", "No native parent app uses phone registration or version checks yet, no versions are configured, the deletion screen is absent, and real Firebase and transfer delivery remain unproved."),
    ]),
    "M30": ("FR-017  Validate the Store-App Version Policy", [
        ("Description", "Report malformed global store-app version configuration before release while leaving the public version route available."),
        ("Rule", "NATIVE_APP_VERSIONS defines minimum, latest, and store_url per known app and android or ios. Versions contain one to three numeric dotted parts and compare numerically, so 1.10 is newer than 1.9. A malformed entry produces system check core.W002 and never breaks the request route."),
        ("Public contract", "GET /v1/app-versions/ is unauthenticated, throttled, cacheable for five minutes, tenant-independent, and returns update_required, update_available, current, or null status when no installed version is supplied. Unknown app or platform is 404 and bad input is 400."),
        ("Acceptance", "A bad deployment setting is visible to system checks; a below-minimum version is update_required; a version between minimum and latest is update_available; and one school's request cannot alter another's answer."),
        ("Current limit", "No store-app versions are configured and no released store app calls the route yet."),
    ]),
}


SUMMARIES = {
    "M01": "Reconciles school provisioning, school-branded email, and the public and signed-in seams used by native clients and account closure. Aligns with MRD v2.115.",
    "M03": "Adds native-origin sessions, parent device sessions, refresh replay grace, and product-governed account closure. Aligns with MRD v2.115.",
    "M04": "Adds the pending role-grant list and the leave, account-closure, finance-account, refund, and virtual-account Field Access contracts. Aligns with MRD v2.115.",
    "M05": "Records protected leave evidence and durable self-service account-closure audit evidence. Aligns with MRD v2.115.",
    "M06": "Adds the parent device-session length and family card-payment settings. Aligns with MRD v2.115.",
    "M07": "Routes leave from the person it concerns and provisions each school's editable hand-journal route. Aligns with MRD v2.115.",
    "M08": "Adds school email identity, native Firebase delivery, closure notices, and pupil-transfer notices. Aligns with MRD v2.115.",
    "M11": "Adds the guardian and pupil account-closure policy and branch-scoped school review. Aligns with MRD v2.115.",
    "M12": "Adds live leave preview and approval behavior, leave Field Access, and draft-payslip YTD consistency. Aligns with MRD v2.115.",
    "M15": "Applies leave Field Access to attendance and register surfaces. Aligns with MRD v2.115.",
    "M17": "Separates owing from held credit and documents transfer-first family payment instructions. Aligns with MRD v2.115.",
    "M18": "Adds enrolled-pupil transfer-account coverage, optional family card payment, and booked-transfer notices. Aligns with MRD v2.115.",
    "M19": "Reconciles chart guidance, managed accounts, code allocation, hand-journal approval, YTD, held credit, and receipt balances. Aligns with MRD v2.115.",
    "M20": "Applies eligible held credit when a debit note posts or an adjustment void reopens a bill. Aligns with MRD v2.115.",
    "M27": "Adds the store-app session, push, version, closure, refresh, and transfer-first fee contracts for pupils. Aligns with MRD v2.115.",
    "M28": "Adds sliding parent sessions, native delivery, version and closure behavior, and per-child transfer-first fees. Aligns with MRD v2.115.",
    "M30": "Adds the global store-app version policy and its core.W002 deployment check. Aligns with MRD v2.115.",
}


TRACE = {
    "M01": ("School service provisioning and branding seams", "FR-025"),
    "M03": ("Native sessions, refresh recovery, and account closure", "FR-029"),
    "M04": ("Pending grant state and new Field Access resources", "FR-035"),
    "M05": ("Leave and account-closure audit evidence", "FR-019"),
    "M06": ("Parent-session and family-card settings", "FR-024"),
    "M07": ("Subject-based leave and hand-journal routes", "FR-036"),
    "M08": ("School email branding and native push", "FR-018"),
    "M11": ("Guardian and pupil account closure", "FR-029"),
    "M12": ("Leave preview, approval, Field Access, and draft YTD", "FR-028"),
    "M15": ("Leave-aware attendance surfaces", "FR-011"),
    "M17": ("Owing, credit, and family transfer instructions", "FR-016"),
    "M18": ("Pupil transfer accounts and optional family cards", "FR-028"),
    "M19": ("Chart guidance, hand journals, and held credit", "FR-027"),
    "M20": ("Held credit after debit notes and voids", "FR-013"),
    "M27": ("Pupil store app and transfer-first fees", "FR-012"),
    "M28": ("Parent device sessions and per-child transfers", "FR-015"),
    "M30": ("Store-app version policy check", "FR-017"),
}


MRD_SUMMARIES = {
    1: "School setup now supplies the standard leave route and school delivery resolves the school's own email name and logo.",
    3: "Native and parent device sessions keep refresh credentials in the phone keychain, recover one lost rotation inside a short grace window, and eligible people may close their own account under product policy.",
    4: "Pending role grants have a narrow branch-scoped read, and Field Access covers leave, account closure, managed accounts, statutory lines, virtual-account values, and refund bank writes.",
    5: "Leave and account-closure evidence stays durable without exposing fields their reader cannot read.",
    6: "Schools configure parent device-session length and whether families may start card payments.",
    7: "Leave routes from the person it concerns, and every school owns an editable no-step hand-journal route that fails closed until configured or explicitly bypassed.",
    8: "School email carries the school's own identity, phone push uses Firebase beside Web Push, and account closure and booked transfers raise their dedicated notices.",
    11: "Guardian-only accounts close immediately, pupil requests wait for branch-scoped school review, and retained school records are never deleted with the login account.",
    12: "Leave preview, line-manager approval, leave Field Access, and draft-payslip year-to-date figures share the live business rules.",
    15: "Attendance surfaces apply the leave resource's type, date, day, status, and balance decisions.",
    17: "Family fees show owing and held credit separately and render the pupil's transfer account when card payment is off.",
    18: "Every enrolled pupil receives a dedicated transfer account, family card collection is optional, and a booked transfer notifies the pupil and linked guardians.",
    19: "The chart explains and folds system-managed accounts, allocates safe codes, hand journals use the school's route, and held credit settles only eligible same-branch receivables.",
    20: "A debit note and a qualifying adjustment void may apply eligible same-branch held credit without blocking a closed-period void.",
    27: "The pupil portal now has the backend contract for store-app sessions, phone push, version enforcement, account closure, refresh recovery, and transfer-first fees.",
    28: "The parent portal uses a sliding device session and per-child transfer accounts, with the same store-app delivery, version, closure, and refresh contracts.",
    30: "The public store-app version policy is validated by core.W002 and remains tenant-independent.",
}


def replace_runs(paragraph: Paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    return True


def set_optional_control(doc, labels: tuple[str, ...], value: str) -> None:
    for table in doc.tables[:3]:
        for row in table.rows:
            if row.cells[0].text.strip() in labels:
                keep_format(row.cells[-1], value)
                return


def update_mrd_reference(doc) -> None:
    set_optional_control(doc, ("Source MRD", "MRD baseline"), MRD_TARGET)
    change_log = common.change_log_table(doc)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Module ") and "MRD v" in paragraph.text:
            text = re.sub(r"MRD v[0-9.]+", f"MRD v{MRD_TARGET}", paragraph.text)
            replace_runs(paragraph, paragraph.text, text)
    for table in doc.tables:
        if table._tbl is change_log._tbl:
            continue
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip().startswith("MRD RECONCILIATION"):
                    for paragraph in cell.paragraphs:
                        text = re.sub(r"MRD v[0-9.]+", f"MRD v{MRD_TARGET}", paragraph.text)
                        replace_runs(paragraph, paragraph.text, text)


def append_traceability(doc, capability: str, requirement: str) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(cell.text for cell in table.rows[0].cells).lower()
        if "mrd" not in header or "capabil" not in header:
            continue
        columns = len(table.rows[0].cells)
        values = [capability, requirement]
        if columns >= 3:
            values.append("Implemented with recorded limits")
        append_rows(table, [values])
        return


def append_gap_line(doc, starts: tuple[str, ...], text: str) -> bool:
    for table in doc.tables:
        if not table.rows:
            continue
        cell = table.rows[0].cells[0]
        if any(cell.text.strip().startswith(start) for start in starts):
            template = cell.paragraphs[-1]
            clone = copy.deepcopy(template._p)
            template._p.addnext(clone)
            paragraph = Paragraph(clone, template._parent)
            if paragraph.runs:
                paragraph.runs[0].text = f"• {text}"
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(f"• {text}")
            return True
    return False


def insert_gap_paragraph(doc, heading: str, text: str) -> bool:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() != heading:
            continue
        sibling = paragraph._p.getnext()
        if sibling is None or not sibling.tag.endswith("}p"):
            return False
        clone = copy.deepcopy(sibling)
        sibling.addnext(clone)
        created = Paragraph(clone, paragraph._parent)
        if created.runs:
            created.runs[0].text = text
            for run in created.runs[1:]:
                run.text = ""
        else:
            created.add_run(text)
        return True
    return False


def remove_branch_synonym(doc) -> None:
    def clean(text: str) -> str:
        return re.sub(r"\bcampuses\b", "branches", re.sub(r"\bcampus\b", "branch", text, flags=re.I), flags=re.I)
    for paragraph in doc.paragraphs:
        if re.search(r"\bcampus(?:es)?\b", paragraph.text, re.I):
            replace_runs(paragraph, paragraph.text, clean(paragraph.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if re.search(r"\bcampus(?:es)?\b", paragraph.text, re.I):
                        replace_runs(paragraph, paragraph.text, clean(paragraph.text))


def ensure_ordinary_change_log_row(doc, folder: str, stem: str) -> None:
    """Repair inherited logs whose version rows all carry the header style."""
    table = common.change_log_table(doc)
    if any(not common.is_header_styled(table, row) for row in table.rows[1:]):
        return
    candidates = sorted(
        (ROOT / "functional-requirements" / folder).glob(f"{stem}_v*.docx"),
        key=lambda path: tuple(int(part) for part in path.stem.rsplit("_v", 1)[1].split(".")),
    )
    for path in candidates:
        candidate = Document(path)
        try:
            candidate_table = common.change_log_table(candidate)
        except ValueError:
            continue
        if candidate_table.rows[0].cells[0].text.strip() != "Version":
            continue
        if any(not common.is_header_styled(candidate_table, row) for row in candidate_table.rows[1:]):
            restore_clean_change_log_rows(doc, path)
            return
    raise ValueError(f"No clean change-log row exists for {stem}")


def clean_change_log_source(folder: str, stem: str) -> Document:
    """Return the newest earlier document with a real three-cell log header."""
    candidates = sorted(
        (ROOT / "functional-requirements" / folder).glob(f"{stem}_v*.docx"),
        key=lambda path: tuple(int(part) for part in path.stem.rsplit("_v", 1)[1].split(".")),
        reverse=True,
    )
    for path in candidates:
        candidate = Document(path)
        try:
            header = [cell.text.strip() for cell in common.change_log_table(candidate).rows[0].cells]
        except ValueError:
            continue
        if header in (["Version", "Date", "Summary"], ["Version", "Date", "Change"]):
            return candidate
    raise ValueError(f"No clean change-log header exists for {stem}")


def ensure_clean_change_log_header(doc, folder: str, stem: str) -> None:
    """Replace a header whose Date and Summary cells were overwritten by a version."""
    table = common.change_log_table(doc)
    header = [cell.text.strip() for cell in table.rows[0].cells]
    if header in (["Version", "Date", "Summary"], ["Version", "Date", "Change"]):
        return
    clean = common.change_log_table(clean_change_log_source(folder, stem)).rows[0]
    table.rows[0]._tr.addprevious(copy.deepcopy(clean._tr))
    table._tbl.remove(table.rows[1]._tr)


def trim_trailing_empty_paragraphs(doc) -> None:
    """Remove empty body paragraphs after the final table that create a blank page."""
    body = doc.element.body
    while len(body) >= 2:
        element = body[-2]
        if not element.tag.endswith("}p"):
            break
        paragraph = Paragraph(element, doc)
        if paragraph.text.strip() or element.xpath('.//w:drawing | .//w:pict | .//w:br'):
            break
        body.remove(element)


def apply_module_gaps(key: str, doc) -> None:
    if key == "M08":
        append_gap_line(doc, ("CURRENT GAPS",), "No store app registers a phone yet, and native delivery has not been proved against real Firebase.")
    elif key == "M11":
        table = doc.tables[58]
        append_rows(table, [["Account-closure screens", "The backend policy and branch-scoped review are complete, but the staff review list and parent and pupil deletion screens do not exist.", "People cannot exercise or review the feature in the product UI."]])
    elif key == "M18":
        append_gap_line(doc, ("FURTHER GAPS",), "Paystack rate limits for bulk enrollment are untested, and no real end-to-end family transfer has been run on staging.")
    elif key == "M19":
        append_gap_line(doc, ("NEEDS ATTENTION",), "Credit held at another branch remains beside a bill by design. Invoice posting still refuses when a closed period prevents required credit application, while a void skips the pairing and completes.")
    elif key == "M27":
        insert_gap_paragraph(doc, "10. Needs Attention", "No store app calls phone registration or version policy yet, no versions are configured, the deletion screen is absent, and real Firebase and transfer delivery remain unproved.")
    elif key == "M28":
        append_gap_line(doc, ("FURTHER GAPS", "OPEN"), "No native parent app calls phone registration or version policy yet, no versions are configured, the deletion screen is absent, and real Firebase and transfer delivery remain unproved.")
    elif key == "M30":
        for table in doc.tables:
            if table.rows and table.rows[0].cells[0].text.strip() == "Priority":
                append_rows(table, [["P2", "No store-app versions are configured", "Set NATIVE_APP_VERSIONS for each released parent and pupil app and have the clients call the public policy route before sign-in.", "FR-017"]])
                break


def patch_frd(key: str) -> Path:
    folder, stem, source, target = DOCS[key]
    source_path = ROOT / "functional-requirements" / folder / f"{stem}_v{source}.docx"
    output = ROOT / "functional-requirements" / folder / f"{stem}_v{target}.docx"
    if output.exists():
        raise SystemExit(f"{output.name} already exists")
    doc = Document(source_path)
    set_cover_version(doc, source, target)
    set_control(doc, "Version", target)
    set_optional_control(doc, ("Review date", "Date"), REVIEW_DATE)
    set_optional_control(doc, ("Code baseline", "Source scope", "Verified against"), CODE_BASELINE)
    update_mrd_reference(doc)
    heading, rows = REQUIREMENTS[key]
    add_fr(doc, heading, heading, rows)
    append_traceability(doc, *TRACE[key])
    apply_module_gaps(key, doc)
    ensure_ordinary_change_log_row(doc, folder, stem)
    ensure_clean_change_log_header(doc, folder, stem)
    log_change(doc, target, SUMMARIES[key], header_source=clean_change_log_source(folder, stem))
    repair_ooxml(doc)
    normalise_change_log(doc)
    sort_and_deduplicate_change_log(doc)
    remove_branch_synonym(doc)
    trim_trailing_empty_paragraphs(doc)
    title = f"{stem.replace('_', ' ')} v{target}"
    finish(doc, output, title, target)
    return output


def append_module_summary(doc, module: int, sentence: str, version: str) -> None:
    heading = f"Module {module}:"
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith(heading):
            description = doc.paragraphs[index + 1]
            text = re.sub(rf"Documented by M{module:02d} FRD v[0-9.]+\.?", "", description.text).rstrip()
            if text and not text.endswith("."):
                text += "."
            replace_runs(description, description.text, f"{text} {sentence} Documented by M{module:02d} FRD v{version}.")
            return
    raise ValueError(f"MRD module heading not found: {module}")


def mrd_header_source() -> Document:
    paths = sorted(
        ROOT.joinpath("module-requirements").glob("XVS_Module_Requirements_Document_v*.docx"),
        key=lambda path: tuple(int(part) for part in path.stem.rsplit("_v", 1)[1].split(".")),
        reverse=True,
    )
    for path in paths:
        doc = Document(path)
        try:
            table = common.change_log_table(doc)
        except ValueError:
            continue
        if table.rows[0].cells[0].text.strip() == "Version":
            return doc
    raise ValueError("No MRD change log with a header")


def patch_mrd() -> Path:
    source = ROOT / "module-requirements" / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"
    output = ROOT / "module-requirements" / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx"
    if output.exists():
        raise SystemExit(f"{output.name} already exists")
    doc = Document(source)
    set_control(doc, "Version", MRD_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Source scope", CODE_BASELINE)
    doc.core_properties.version = MRD_TARGET

    for module, sentence in MRD_SUMMARIES.items():
        key = f"M{module:02d}"
        append_module_summary(doc, module, sentence, DOCS[key][3])

    change_heading = next(p for p in doc.paragraphs if p.text.strip() == "7. Change Log")
    template_heading = next(p for p in doc.paragraphs if p.text.strip().startswith("6.5 22 October 2026"))
    heading_xml = copy.deepcopy(template_heading._p)
    change_heading._p.addprevious(heading_xml)
    heading = Paragraph(heading_xml, template_heading._parent)
    replace_runs(heading, heading.text, "6.6 23 to 26 October 2026 Capability Contracts")
    template_body = template_heading._p.getnext()
    if template_body is None or not template_body.tag.endswith("}p"):
        raise ValueError("The 6.5 heading has no body paragraph to clone")
    paragraphs = [
        "Leave now has one complete school contract: a direct-manager ladder provisioned without replacing a school's choices, a live working-day preview, subject-based organogram routing, Field Access on every read and approval surface, and attendance views that reveal only allowed leave facts.",
        "Identity now supports store-app origins and parent device sessions, keeps native refresh credentials in the phone keychain, recovers the immediately previous rotated credential inside a two-minute default grace window, and applies one product policy to self-service account closure and school review.",
        "Notifications resolve each school's email name and logo at delivery and send the same push message to browser subscriptions and Firebase phone devices. Account-closure notices and booked family-transfer notices use the same tenant, branch, preference, and Field Access boundaries as their records.",
        "Family fees are transfer-first by default. Owing and held credit are shown separately, each enrolled pupil receives a dedicated transfer account, family card payment is a school-wide option off by default, and invoices, receipts, and portals state which balance and payment method they show.",
        "Finance now explains and folds system-managed accounts, allocates safe account codes, reserves the default chart, routes hand journals through each school's editable route, counts a draft payslip in its own year to date, and applies held credit only to eligible same-branch receivables.",
        "The pupil and parent portals now have the backend contracts for native sign-in, phone push, release-version policy, refresh recovery, account closure, and transfer-first fees. The parent app also has a sliding school-configured device-session length. Store-app release and production evidence remain open.",
        "No module phase, ownership, priority order, or capability count changes. These behaviors deepen existing capability entries, so the tracker remains at 32 modules and 527 capability entries. The seventeen affected FRDs advance independently and reconcile to this MRD.",
    ]
    for text in paragraphs:
        clone = copy.deepcopy(template_body)
        change_heading._p.addprevious(clone)
        paragraph = Paragraph(clone, change_heading._parent)
        replace_runs(paragraph, paragraph.text, text)

    summary = (
        "Reconciles the 23 to 26 October behavior batch across leave, role-grant state, native and parent sessions, refresh recovery, account closure, school-branded email, Firebase phone push, app-version policy, transfer-first family fees, chart guidance, hand-journal approval, held credit, receipt balances, and draft-payslip year to date. No module phase, ownership, priority order, or capability count changes; the tracker remains at 32 modules and 527 entries. Reconciled to M01 v1.35.1, M03 v1.25, M04 v1.43, M05 v1.13.1, M06 v1.12, M07 v1.28, M08 v1.20, M11 v2.18, M12 v2.22, M15 v1.1, M17 v1.15, M18 v1.20, M19 v1.24, M20 v1.13, M27 v1.2, M28 v1.5, and M30 v1.6.1."
    )
    log_change(doc, MRD_TARGET, summary, header_source=mrd_header_source())
    repair_ooxml(doc)
    normalise_change_log(doc)
    sort_and_deduplicate_change_log(doc)
    remove_branch_synonym(doc)
    doc.core_properties.title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.save(output)
    update_extended_title(output, doc.core_properties.title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    print(f"Wrote {output.name}")
    return output


def main() -> None:
    outputs = [patch_frd(key) for key in DOCS]
    outputs.append(patch_mrd())
    print(f"Created {len(outputs)} documents")


if __name__ == "__main__":
    main()
