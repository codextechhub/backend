#!/usr/bin/env python3
"""Version the MRD and Module 8 and 9 FRDs for onboarding email copy."""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)


REVIEW_DATE = "4 September 2026"
SHORT_DATE = "4 Sep 2026"
MRD_SOURCE_VERSION = "2.61"
MRD_TARGET_VERSION = "2.62"


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_paragraph(paragraph, text: str, **kwargs) -> None:
    write_paragraph(paragraph, text, **kwargs)


def replace_cover_versions(table, versions: tuple[str, ...], target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            for source in versions:
                if source in run.text:
                    run.text = run.text.replace(source, target)


def replace_control_value(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            replace_cell(row.cells[1], value, size=9)
            return
    raise ValueError(f"Control row not found: {label}")


def prepend_change_log(table, version: str, summary: str) -> None:
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], SHORT_DATE, size=8)
    replace_cell(row.cells[2], summary, size=8)


def append_styled_row(table, values: list[str], *, size: float = 8.2) -> None:
    template = table.rows[-1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addnext(new_tr)
    row = table.rows[-1]
    for cell, value in zip(row.cells, values):
        replace_cell(cell, value, size=size)


def finish(doc, output: Path, title: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS Module Requirements Document v{MRD_TARGET_VERSION}"
    doc.core_properties.title = title
    doc.core_properties.version = MRD_TARGET_VERSION
    replace_cover_versions(doc.tables[0], (MRD_SOURCE_VERSION,), MRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Version", MRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Review date", REVIEW_DATE)
    replace_control_value(
        doc.tables[1],
        "Source scope",
        "Backend change refining all six onboarding lifecycle email templates across Modules 8 and 9. Progress, readiness, review, activation, expiry, and operator follow-up mail now separates the event details from the next action, uses human-readable local dates, and states the correct school or platform responsibility. Migration 0016 refreshes only platform-maintained templates. Verified by 139 notification tests and 171 onboarding tests on 4 September 2026. Backend evidence only; nothing here is deployed.",
    )

    replace_cell(
        doc.tables[2].rows[5].cells[0],
        f"5. v{MRD_TARGET_VERSION} Capability Delta",
        size=9,
        bold=True,
        color=BLUE,
    )
    replace_cell(
        doc.tables[2].rows[5].cells[1],
        "Clear onboarding lifecycle mail for schools and platform operators",
        size=9,
    )

    paragraphs = doc.paragraphs
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Event-driven in-app and email delivery"):
            replace_paragraph(
                paragraph,
                "Event-driven in-app and email delivery. A record is owned by the tenant of the recipient who reads it, and the tenant an event is about is kept separately as its origin. Transactional operational alarms, including Module 30 health alerts, bypass delivery preferences and use both channels. Standard email templates resolve issuer, entity, or tenant branding and its available logo at render time and use one polished, email-client-safe layout. Invitation, password-reset, and all six onboarding lifecycle messages use purpose-specific copy and tell recipients what happened, what it means, and what they can do next. Onboarding review, activation, and expiry mail uses local display dates while older machine-readable context keys remain available to staff-authored templates. Migrations refresh platform-maintained markup while preserving every staff-authored override. One-time invitation and reset credentials are replaced only in the delivery worker, so notification records retain a marker rather than the raw secret. This is a notification engine, not a person-to-person chat product.",
                size=9,
                space_after=5,
            )
        elif text.startswith("The onboarding tracker and the go-live gate"):
            replace_paragraph(
                paragraph,
                "The onboarding tracker and the go-live gate, built as a module of its own. A School that is not yet live signs in, works a checklist, submits a go-live request, and is activated only after platform review. The six lifecycle messages now match that ownership: the ready message tells the School to submit for review rather than telling an administrator to activate it, reviewed and activated mail shows readable local times, expiry mail states the exact deadline and sign-in consequence, and the operator report separates ageing Schools from those already suspended. A School that abandons the process is warned, expires, and can be put back by hand. Approving, rejecting and reinstating remain reserved to the platform Tenant.",
                size=9,
                space_after=5,
            )
        elif text == f"5. v{MRD_SOURCE_VERSION} Capability Delta":
            replace_paragraph(
                paragraph,
                f"5. v{MRD_TARGET_VERSION} Capability Delta",
                size=17,
                bold=True,
                space_before=15,
                space_after=8,
            )
        elif text.startswith("This revision improves existing password and security communication"):
            replace_paragraph(
                paragraph,
                "This revision improves the existing onboarding notification capability shared by Modules 8 and 9. It does not add a route, permission, event, or capability entry. The six messages now distinguish school actions from platform decisions, use readable lifecycle dates, and preserve staff-authored template overrides.",
                size=9,
                space_after=5,
            )

    rebuild_table(
        doc.tables[76],
        [f"v{MRD_TARGET_VERSION} capability delta", "Decision", "Evidence"],
        [
            [
                "Six onboarding messages",
                "Purpose-specific",
                "Step completion, readiness, review, activation, expiry warning, and the stale operator report now present structured details and one accurate next action.",
            ],
            [
                "Go-live hand-off",
                "Corrected",
                "The ready email tells the School to submit a request for platform review. It no longer tells a School administrator to make the School live, an action reserved to platform staff.",
            ],
            [
                "Lifecycle dates and outcomes",
                "Human-readable",
                "Review and activation use local display timestamps; expiry uses a display deadline and correct singular or plural day wording. The operator report separates ageing Schools from recently suspended Schools.",
            ],
            [
                "Existing customized templates",
                "Preserved",
                "Migration 0016 changes only html_is_custom=false rows. Older machine-readable timestamp and date context keys remain available beside the new display values.",
            ],
        ],
        [1.75, 0.95, 4.57],
        font_size=8.2,
    )
    prepend_change_log(
        doc.tables[78],
        MRD_TARGET_VERSION,
        "Refined all six onboarding lifecycle email templates across Modules 8 and 9. Step completion now names the completed item and checklist position. Readiness tells the School to submit a go-live request for platform review instead of telling an administrator to activate the School. Review mail separates approved and rejected outcomes and keeps the rejection reason. Activation shows when the workspace became live and accurately limits access to enabled areas permitted by each person's roles. Expiry mail uses a readable deadline, correct day grammar, and the exact suspension and sign-in consequence. The platform report separates ageing Schools from those recently suspended. Review and activation retain their existing machine-readable context while adding local display timestamps; expiry retains its ISO date beside the display date. Migration 0016 refreshes only platform-maintained templates and preserves staff-authored content. Module states, ownership, capability counts, priority gaps, and build order do not change. Verified by 139 notification tests and 171 onboarding tests. Backend evidence only; nothing here is deployed.",
    )
    finish(doc, output, title)


def patch_m08(source: Path, output: Path) -> None:
    source_version, target_version = "1.7", "1.8"
    doc = Document(str(source))
    title = f"XVS M08 Notifications and Delivery Functional Requirements Document v{target_version}"
    doc.core_properties.title = title
    doc.core_properties.version = target_version
    replace_cover_versions(doc.tables[0], (source_version,), target_version)
    replace_control_value(doc.tables[1], "Version", target_version)
    replace_control_value(doc.tables[1], "Review date", REVIEW_DATE)
    replace_control_value(
        doc.tables[1],
        "Code baseline",
        "Commit 1d68109 plus the onboarding lifecycle email refinements, migration 0016, and passing Module 8 and Module 9 suites under review (4 September 2026)",
    )
    replace_control_value(
        doc.tables[1],
        "Source MRD",
        f"XVS Module Requirements Document v{MRD_TARGET_VERSION} | Module 8",
    )
    replace_control_value(
        doc.tables[1],
        "Supporting apps",
        "vs_tenants, vs_rbac, vs_user, vs_config, vs_health, schools/vs_onboarding, core",
    )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Module 8 is the platform's one way"):
            replace_paragraph(
                paragraph,
                "Module 8 is the platform's one way of telling somebody something happened. A domain module raises a named event with a context; this module decides which channels may carry it for that tenant, renders the message for each recipient, writes the in-app record, queues the email, and keeps the outcome. Standard email renders use one polished, email-client-safe layout and resolve the displayed sender name and available logo from the issuer, entity, or tenant at delivery time. Invitation, password-reset, and all six onboarding lifecycle messages now use purpose-specific copy. The onboarding family covers progress, readiness, review, activation, expiry, and platform follow-up, with structured details, accurate ownership, and readable local dates. Transactional operational alarms such as a Module 30 health incident bypass preferences and use both email and in-app. Templates and ordinary rendered content remain visible to administrators. A one-time invitation or reset credential is the deliberate exception: history stores an inert marker and only the delivery worker sees the raw value.",
                size=9,
                space_after=5,
            )
        elif text.startswith("Module 8 carries 22 capability entries"):
            replace_paragraph(
                paragraph,
                f"Module 8 carries 22 capability entries in MRD v{MRD_TARGET_VERSION}. Each maps to the requirements below. Refining the six onboarding templates and their display context strengthens the existing template, preview, rendering, and delivery capabilities without changing the count.",
                size=9,
                space_after=5,
            )

    fr006 = doc.tables[12]
    replace_cell(
        fr006.rows[2].cells[1],
        "A template carries a flag saying who maintains its markup. While it is false the markup is regenerated from the shared layout on every save; the first hand edit sets it, after which the markup is kept verbatim. Clearing it restores the standard design. Standard markup carries a domain-neutral email_brand placeholder and an optional brand-logo placeholder resolved from issuer, entity, or tenant context at delivery, with CodeX Vision as fallback. Migrations 0012 through 0015 refreshed standard design and purpose-specific account mail. Migration 0016 refines all six standard onboarding email rows. Every migration filters html_is_custom=false.",
        size=8.5,
    )
    replace_cell(
        fr006.rows[3].cells[1],
        "Changing the message of a standard template changes its markup. Changing the message of a hand-edited one does not. Saving regenerated standard markup back is not treated as a hand edit. A tenant-branded message renders its name and available logo while missing branding renders the CodeX Vision fallback. Each design migration changes only html_is_custom=false. A customized invitation, password reset, account-lock, or onboarding template keeps its subject, body, action, and HTML verbatim.",
        size=8.5,
    )

    fr008 = doc.tables[14]
    replace_cell(
        fr008.rows[2].cells[1],
        "Preview renders through the same path dispatch uses and needs no payload: sample values are generated from the placeholders the template itself uses. Exact onboarding samples cover approved or rejected decisions, readable review and activation times, an expiry deadline, singular or plural remaining days, counts, and both operator lists. Preview also accepts unsaved editor content, so the console re-renders as an author types, and returns the markup alongside the rendered copy.",
        size=8.5,
    )
    replace_cell(
        fr008.rows[3].cells[1],
        "A preview of any active template returns a complete HTML document. Preview writes nothing: no record is created and no mail is sent. Supplying context overrides the generated samples for those keys only. Every standard onboarding template renders without unresolved variables, including both branches of its conditional copy.",
        size=8.5,
    )

    append_styled_row(
        doc.tables[28],
        [
            "Module 9, School Onboarding",
            "Raises six onboarding events. It supplies the School name and available logo, lifecycle details, readable local review and activation timestamps, the readable expiry deadline, and separate ageing and suspended lists. Existing machine-readable date and timestamp keys remain available for customized templates.",
        ],
    )

    prepend_change_log(
        doc.tables[33],
        target_version,
        "Refines all six standard onboarding lifecycle email templates. Progress and readiness mail names the completed work and the correct next step. The ready message now tells a School to submit a go-live request for platform review instead of telling an administrator to activate it. Review mail separates approved and rejected outcomes and includes the reason when rejected. Activation states when the workspace became live and limits access claims to enabled areas permitted by roles. Expiry warning mail shows a readable deadline, correct singular or plural day wording, and the suspension and sign-in consequence. The platform report separates ageing Schools from those recently suspended. Module 9 retains the existing ISO context values and adds local display timestamps and a display date; preview supplies exact samples for every new variable. Migration 0016 updates only platform-maintained rows and preserves staff-authored content. Updates scope, FR-006, FR-008, the Module 9 dependency, and MRD traceability. Module 8 remains Backend Complete and In use Complete with 22 capabilities. The complete 139-test Module 8 suite and 171-test Module 9 suite passed. Backend evidence only; nothing here is deployed.",
    )
    finish(doc, output, title)


def patch_m09(source: Path, output: Path) -> None:
    source_version, target_version = "2.7", "2.8"
    doc = Document(str(source))
    title = f"XVS M09 School Onboarding Functional Requirements Document v{target_version}"
    doc.core_properties.title = title
    doc.core_properties.version = target_version
    replace_cover_versions(doc.tables[0], (source_version,), target_version)
    replace_control_value(doc.tables[1], "Version", target_version)
    replace_control_value(doc.tables[1], "Review date", REVIEW_DATE)
    replace_control_value(
        doc.tables[1],
        "Code baseline",
        "Commit 1d68109 plus onboarding lifecycle email refinements and migration 0016 under review (4 September 2026)",
    )
    replace_control_value(
        doc.tables[1],
        "Source MRD",
        f"XVS Module Requirements Document v{MRD_TARGET_VERSION} | Module 9",
    )
    replace_control_value(
        doc.tables[1],
        "Supersedes",
        "v2.7 and all earlier versions, retained unchanged",
    )

    replace_cell(
        doc.tables[4].rows[0].cells[0],
        "CURRENT CREATION, READINESS, AND COMMUNICATION BOUNDARY\n"
        "• Module 1 returns 503 and rolls the whole new School transaction back when its required administrator cannot be provisioned. Module 9 receives no half-created School to repair.\n"
        "• FIRST_ADMIN still checks later invitation activation and account or assignment liveness. ROLE_BASELINE still checks that the School administrator role grants authority before go-live.\n"
        "• A School completes the seven-step catalog and submits a request. Platform staff alone approve, reject, activate, and reinstate; the ready email now states that hand-off accurately.\n"
        "• Six email and in-app events cover progress, readiness, review, activation, expiry warning, and platform follow-up. Human-readable local dates sit beside the older machine-readable context for customized templates.\n"
        "• Onboarding expiry and reinstatement still use the shared locked Tenant transition service. Suspension closes active impersonation sessions in the same transaction, and reinstatement never revives them.",
        size=8.5,
    )

    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if text.startswith("Version 2.7 makes onboarding suspension"):
            replace_paragraph(
                paragraph,
                "Version 2.8 makes the six onboarding messages follow the lifecycle they describe. Readiness now tells the School to submit for platform review, review and activation use readable local timestamps, expiry uses a readable deadline and precise sign-in consequence, and the operator report separates ageing Schools from those already suspended. Existing machine-readable context keys remain available to customized templates.",
                size=9,
                space_after=5,
            )
        elif text.startswith("Module 9 carries 20 capability entries"):
            replace_paragraph(
                paragraph,
                f"Module 9 carries 20 capability entries in MRD v{MRD_TARGET_VERSION}. Each maps to the requirements below. Clearer lifecycle mail strengthens the existing onboarding notification, expiry warning, and operator report capabilities without adding an event or changing the count.",
                size=9,
                space_after=5,
            )
        elif text == "FR-008  Approve or Reject a Go-Live Request":
            previous = paragraphs[index - 1]
            for page_break in previous._p.xpath(".//w:br"):
                page_break.getparent().remove(page_break)

    fr011 = doc.tables[20]
    replace_cell(
        fr011.rows[2].cells[1],
        "Eleven action types are registered under the ONBOARDING module key. Six notification events exist with in-app and email templates: step completed, ready, reviewed, activated, the expiry warning to the School, and the stale report to operators. The emails now present structured details and an accurate next action. The ready message tells the School to submit for platform review; reviewed mail branches between approval and rejection; activated mail states the access boundary; expiry warning names its deadline and sign-in consequence; and the stale report separates ageing from recently suspended Schools. Review and activation add local display timestamps while retaining reviewed_at and go_live_at; expiry adds expires_on_display while retaining expires_on. Audit and notification remain queued in one after-commit callback, so the trail is written before anybody is told.",
        size=8.5,
    )
    replace_cell(
        fr011.rows[3].cells[1],
        "Every transition writes exactly one registered audit event. School-facing onboarding events resolve whole-tenant onboarding.progress.view holders with branch None, so the audience and the control-room gate cannot drift apart. The stale report resolves platform onboarding.go_live.approve holders because those operators own cross-School follow-up. Approved and rejected review messages render their distinct paths, human-readable dates contain no ISO timestamp, one remaining day is grammatically singular, and the six templates render with no unresolved variable.",
        size=8.5,
    )
    replace_cell(
        fr011.rows[4].cells[1],
        "Notification dispatch is wrapped and never fails a request. Where a template is missing the channel is skipped and a warning is logged, so seeding is a deploy precondition rather than a runtime guarantee. No call-to-action URL is stored for this family because the backend has no reliable tenant-specific frontend route builder; each email names the next action in text.",
        size=8.5,
    )

    replace_cell(
        doc.tables[34].rows[4].cells[1],
        "Carries six onboarding events on in-app and email. Standard templates use the shared tenant-aware layout and available School logo, preserve customized rows, and render the display values this module supplies. Dispatch is wrapped: a template problem never fails a request.",
        size=8.2,
    )

    prepend_change_log(
        doc.tables[41],
        target_version,
        "Refines all six onboarding lifecycle messages without changing the workflow. Step-completion mail names the completed item and checklist position. Readiness now tells the School to submit a go-live request for platform review, correcting the former instruction that told an administrator to make the School live. Review mail separates approved and rejected outcomes and includes the rejection reason and readable review time. Activation shows its readable local time and promises only the areas enabled for the School and permitted by each person's roles. Expiry warning shows a readable deadline, correct day grammar, and the precise suspension and sign-in consequence. The platform report separates ageing Schools from those recently suspended. Existing reviewed_at, go_live_at, and expires_on context remains for customized templates; reviewed_at_display, go_live_at_display, and expires_on_display are added for the standard copy. Migration 0016 updates only platform-maintained rows. FR-011, the current boundary, Module 8 dependency, and MRD traceability are reconciled. Module 9 stays Backend Complete and In use Partial with twenty capability entries. Verified by 171 onboarding tests and 139 notification tests. Backend evidence only; nothing here is deployed.",
    )
    finish(doc, output, title)


def main() -> None:
    parser = argparse.ArgumentParser()
    for name in ("mrd", "m08", "m09"):
        parser.add_argument(f"--{name}-source", type=Path, required=True)
        parser.add_argument(f"--{name}-output", type=Path, required=True)
    args = parser.parse_args()
    patch_mrd(args.mrd_source, args.mrd_output)
    patch_m08(args.m08_source, args.m08_output)
    patch_m09(args.m09_source, args.m09_output)
    for output in (args.mrd_output, args.m08_output, args.m09_output):
        print(output)


if __name__ == "__main__":
    main()
