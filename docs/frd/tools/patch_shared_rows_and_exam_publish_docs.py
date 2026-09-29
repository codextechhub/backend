#!/usr/bin/env python3
"""Cut M13 v2.12 and M14 v3.5: shared rows are read-only to a branch administrator,
and an exam timetable publishes over a shared hall.

What changed in the backend, and therefore in the documents:

* 33473d77. A branch-bound caller reads the school-wide rows beside their own
  branch's, and changes a row only when every branch it belongs to is theirs.
  One predicate, ``vs_rbac.scoping.caller_may_change``, carries the rule for
  staff, academic and calendar rows, and a refusal is 403
  SHARED_RECORD_READ_ONLY. A row is judged by its own branch, else its
  parent's: a session by the branches it covers, a term by its session, an exam
  by its exam period, a lesson and an exam paper by their class. It is applied
  at each module's ``get_object`` and at every hand-written write, and every row
  carries ``can_manage``. Subject offerings keep the offerings a branch caller
  may not change, and a replace touches only the year being edited. Detail
  routes that resolved a calendar row by tenant alone, and the exam, paper and
  session lists, are narrowed. A branch head can no longer create an exam under
  a school-wide exam period, nor name the class teacher or give teaching duties
  on a school-wide class.
* 9526abd2. ``publish_exam`` no longer refuses over a room holding several
  classes' papers or one invigilator between two rooms; both stay warnings on
  write. A class sitting two papers at once is still refused at the write, and
  a class timetable still blocks on its clashes.

M14 also stops saying that "today" is the UTC day and that no school carries a
time zone. Each school keeps its own, display.timezone, Africa/Lagos by default,
and the calendar's "today" is the calendar day in it (480a4c87, 39c256a6). Only
those statements are corrected: FR-005's time-zone rule, FR-006's trigger,
section 12's time-zone row and section 13's decision 9, which closes.

Both documents' control pages, left reading as the version before by the
revision that cut v2.11 and v3.4, are brought up to date, and those two
change-log rows, written at the top of logs that run oldest first, are moved to
their places.

The MRD and the Module 12 and Module 4 documents are not revised here.

    python tools/patch_shared_rows_and_exam_publish_docs.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from patch_document_type_labels_docs import require_newest
from patch_students_academics_calendar_docs import (
    add_bullets_before,
    add_row_after,
    append_cell,
    append_paragraph,
    append_row,
    body_paragraph,
    clone_fr_block,
    edit_cell,
    find_row,
    finish,
    fr_heading,
    replace_text,
    restore_row_format,
    row_by_label,
    set_cell,
    set_cover,
    set_run_text,
)

ROOT = Path(__file__).resolve().parents[1] / "functional-requirements"

REVIEW_DATE = "27 September 2026"
SHORT_DATE = "27 Sep 2026"
MRD_VERSION = "2.93"


def table_headed(doc, first: str):
    """The one table whose first cell reads exactly ``first``."""
    hits = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == first]
    if len(hits) != 1:
        raise ValueError(f"Table headed {first!r}: {len(hits)} matches")
    return hits[0]


def table_with_row(doc, label: str):
    """The one table holding a row whose first cell reads exactly ``label``."""
    hits = [t for t in doc.tables if any(r.cells[0].text.strip() == label for r in t.rows)]
    if len(hits) != 1:
        raise ValueError(f"Table with row {label!r}: {len(hits)} matches")
    return hits[0]


def fr_table(doc, fr: str):
    hits = [t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith(fr + " ")]
    if len(hits) != 1:
        raise ValueError(f"{fr}: {len(hits)} tables")
    return hits[0]


def reorder_log(log, version: str) -> None:
    """Move a row written at the top of an oldest-first log to its end."""
    stray = log.rows[1]
    if stray.cells[0].text.strip() != version:
        raise ValueError(f"Expected the v{version} row at the top of the change log")
    template = log.rows[-1]
    tr = stray._tr
    tr.getparent().remove(tr)
    template._tr.addnext(tr)
    restore_row_format(log.rows[-1], log.rows[-2])


def keep_rows_whole(table) -> None:
    """Stop a control-page row breaking across two pages."""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def source_path(folder: str, stem: str, version: str) -> Path:
    require_newest(str(ROOT / folder / f"{stem}_v*.docx"), version)
    return ROOT / folder / f"{stem}_v{version}.docx"


SHARED_TESTS = "schools/vs_calendar/tests/test_shared_rows_read_only.py"


# ═════════════════════════════════════════════════════════════════════════════
# Module 13 - Academic Structure
# ═════════════════════════════════════════════════════════════════════════════

M13_DIR = "13-academic-structure"
M13_STEM = "XVS_M13_Academic_Structure_Functional_Requirements_Document"
M13_SOURCE, M13_TARGET = "2.11", "2.12"

M13_STATUS = (
    "Built. Version 2.12 records that a shared row is read-only to a branch administrator, "
    "refused 403 SHARED_RECORD_READ_ONLY, and that every row carries can_manage (FR-018). "
    "Subject offerings are replaced one year at a time."
)

M13_VERIFIED = (
    "Backend main at 9526abd2, 27 September 2026, holding 33473d77. Every claim this "
    "version adds was read from the committed code and test files there; nothing "
    "uncommitted was read."
)

M13_RULE_51_TAIL = (
    " Nor may they change one: a shared row is read-only to a branch-bound caller and a "
    "write to it is refused 403 SHARED_RECORD_READ_ONLY (FR-018)."
)

M13_FR011_WRITES_TAIL = (
    " Changing an existing row is narrower than creating one or reading it: a branch-bound "
    "caller changes a row only when every branch it belongs to is theirs (FR-018)."
)

M13_FR006_RULE_TAIL = (
    " Module 12's route refuses a branch-bound caller the class teacher of a school-wide "
    "class, 403 SHARED_RECORD_READ_ONLY, and so does its teaching duty route (FR-018)."
)

M13_FR007_OLD = (
    "(2) The offerings endpoint takes the complete set of level ids and replaces the "
    "existing set in one transaction, so a client never has to diff."
)
M13_FR007_NEW = (
    "(2) The offerings endpoint takes the complete set of level ids and replaces the "
    "existing set for the year being edited in one transaction, so a client never has to "
    "diff. Another year's offerings are that year's record and are left as they are; a "
    "replace used to delete the subject's offerings in every year. On a shared subject, a "
    "branch-bound caller changes only the offerings at their own branch's levels and every "
    "other offering is carried over, including those at levels they cannot see (FR-018)."
)

M13_FR013_OLD = "None on activation."
M13_FR013_NEW = (
    "None on activation for a caller not narrowed to any branch. A branch-bound caller "
    "names only their own branches in the set, and is refused 403 SHARED_RECORD_READ_ONLY "
    "when activation would narrow or end an ACTIVE year covering more than their branches, "
    "because the narrowing below changes that year for every branch it covers (FR-018)."
)

M13_FR016_RULE = (
    " (8) Whole-tenant only. A branch-bound caller is refused 403 SHARED_RECORD_READ_ONLY "
    "before anything is read or written, because copying a year copies its shared "
    "structure, which is not a branch's to create (FR-018)."
)
M13_FR016_ACCEPTANCE = (
    " (9) A branch-bound caller is refused 403 SHARED_RECORD_READ_ONLY and nothing is "
    "written (test_copying_a_year_forward_is_a_school_wide_act)."
)

M13_FR018 = [
    ("Description",
     "A branch-bound caller reads the shared rows beside their own branch's (FR-011) and "
     "changes a row only when every branch it belongs to is theirs. A shared row, or one "
     "that also belongs to a branch they do not cover, is read-only to them, and a write to "
     "it is refused 403 SHARED_RECORD_READ_ONLY. A caller not narrowed to any branch is "
     "unaffected."),
    ("Why",
     "Reading is inclusive because a school-wide subject, class or year belongs to every "
     "branch. That is also why a branch may not change one: a correction made at Ikeja to "
     "the school-wide JSS1 reaches Lekki's screens with nobody at Lekki knowing. Reads were "
     "narrowed inclusively and writes reused the same narrowing, so anything a branch "
     "administrator could see, they could change. The two behaviour changes a school will "
     "notice, the class teacher below and M14 FR-020's exam, were confirmed by the product "
     "owner on 26 September 2026."),
    ("Actors",
     "Every write in FR-001 to FR-007, FR-013 and FR-016. This requirement adds no "
     "permission key, for the reason FR-011 gives: which rows a caller may change is a "
     "question about where they work, not about what they hold."),
    ("The rule",
     "One predicate for every module, vs_rbac.scoping.caller_may_change(user, tenant, "
     "branch_ids): a whole-tenant caller may change any row they can see; a branch-bound "
     "caller may change a row only when its branch set is non-empty and every branch in it "
     "is one of theirs. assert_caller_may_change raises vs_rbac.exceptions."
     "SharedRecordReadOnly, 403 SHARED_RECORD_READ_ONLY, whose message names who can make "
     "the change, because forbidden alone sends a branch administrator looking for a "
     "permission they will never be given. Module 12's staff records and M14's calendar "
     "rows read through the same predicate."),
    ("Whose row it is",
     "A row is judged by its own branch, else its parent's, in schools.vs_academics."
     "services.scoping.row_branch_ids: a department, programme, level, class or subject by "
     "its branch column; a session by the branches it covers, a school-wide session being "
     "shared; a term by its session. A null branch is shared."),
    ("Where it applies",
     "At the module's get_object, so at every detail view at once: another branch's row "
     "answers 404, as it always did, and a shared row answers 403 to a write. At every "
     "hand-written write as well: archive and restore of a department, programme, level, "
     "class and subject; a year's create and the branch set on its edit; archiving and "
     "activating a year; adding a term; and subject offerings. Module 12's class teacher "
     "and teaching duty routes call the same guard on the class, so a branch administrator "
     "no longer names the class teacher of a school-wide class or gives teaching duties on "
     "one."),
    ("Years",
     "A branch-bound caller sets up a year for their own branches only: an empty branch set "
     "is a school-wide year, and neither that nor a branch they do not cover may be named on "
     "a create or on an edit of the branch set. Activating a year is refused when it would "
     "narrow or end an ACTIVE year covering more than their branches (FR-013). Archiving a "
     "school-wide year is refused. Copying a year forward is whole-tenant only (FR-016). The "
     "session list, which showed a branch-bound caller every branch's years, is narrowed "
     "inclusively: school-wide years and years at their branches."),
    ("Offerings",
     "An offering belongs to the narrower of its subject and its level. On a subject the "
     "caller may change, the set they send is written. On a shared subject each offering is "
     "judged by its level: a branch-bound caller adds or removes offerings at their own "
     "branch's levels only, every other offering the subject has in that year is carried "
     "over untouched, including those at levels they cannot see, and asking to add or "
     "remove one at a shared level is refused 403 SHARED_RECORD_READ_ONLY rather than "
     "ignored. The replace touches only the year being edited (FR-007)."),
    ("can_manage",
     "Every row this module returns to a request carries can_manage, true when the viewer "
     "may change the row rather than only read it, so a screen hides the controls the "
     "server would refuse. The caller's branches are resolved once and memoised, so a page "
     "of rows costs no query per row. It is absent where a serializer runs without a "
     "request, and a screen reads a missing flag as manageable. The school app hides write "
     "controls where it is false and offers the year copy only to a whole-school reader "
     "(school-fe 2b53c61); for a reader who works in one branch it also stops drawing where "
     "a row applies, while the API still returns the branch fields."),
    ("Known gap",
     "The offerings endpoint lets a branch head add their own levels to a shared subject, "
     "but the school app hides the subject's edit control wherever can_manage is false, so "
     "it offers them no way to do it."),
    ("Acceptance",
     "(1) A shared class and a shared subject are read-only to a branch head, their own "
     "class is editable, archiving a shared class is refused, and a school-wide "
     "administrator still edits shared rows. (2) The class list says which rows the viewer "
     "may change. (3) A branch head adds their own level to a shared subject's offerings "
     "and keeps everybody else's; removing a shared level's offering is refused; rewriting "
     "one year's offerings leaves another year's alone. (4) A branch head cannot set up a "
     "school-wide year or archive one, and copying a year forward is refused to them. "
     + SHARED_TESTS + ", AcademicRowsTests, OfferingsTests and SessionTests."),
]

M13_API_TAIL = (
    " Every row these endpoints return carries can_manage, and every write refuses a "
    "branch-bound caller a row that is not wholly their own branch's (FR-018)."
)

M13_REFUSAL = [
    "A branch-bound caller changing a shared row, or one that also belongs to a branch "
    "they do not cover",
    "403",
    "SHARED_RECORD_READ_ONLY, with a message naming who can make the change (FR-018). "
    "This includes a school-wide year, activating over a wider year, copying a year "
    "forward, and an offering at a shared level",
]

M13_BULLETS = {
    "11.2 Behaviour": [
        "A branch-bound caller is refused 403 SHARED_RECORD_READ_ONLY on a write to a "
        "shared class, subject or year, and on archiving one, while a school-wide "
        "administrator is not, and every row says which the viewer may change through "
        "can_manage (FR-018).",
    ],
    "11.3 Shape and scale": [
        "A branch head adding their own level to a shared subject's offerings keeps every "
        "other offering the subject has, and rewriting one year's offerings leaves another "
        "year's alone (FR-007, FR-018).",
    ],
}

M13_CHANGE_SUMMARY = (
    "Minor revision. A shared row is read-only to a branch administrator, and FR-018 is "
    "new to say so. A branch-bound caller still reads the school-wide rows beside their own "
    "branch's (FR-011), but changes a row only when every branch it belongs to is theirs; "
    "otherwise the write is refused 403 SHARED_RECORD_READ_ONLY, from one predicate in "
    "vs_rbac.scoping that Module 12's staff records and M14's calendar rows share. A row is "
    "judged by its own branch, a session by the branches it covers and a term by its "
    "session. The rule reaches every detail view and every hand-written write: archive and "
    "restore, a year's branch set, archiving and activating a year, adding a term, and "
    "subject offerings. A branch-bound caller sets up a year only for their own branches, "
    "may not activate one that would narrow or end a wider year (FR-013), and may not copy "
    "a year forward (FR-016). On a shared subject they change only the offerings at their "
    "own levels and every other offering is kept, and a replace now touches only the year "
    "being edited, where it had deleted the subject's offerings in every year (FR-007). "
    "Module 12 no longer lets a branch administrator name the class teacher of a "
    "school-wide class (FR-006). The session list, which showed every branch's years, is "
    "narrowed inclusively. Every row carries can_manage, and the school app hides write "
    "controls where it is false. Section 10 gains the refusal and section 11 its "
    "acceptance. FR-018 records one gap: the school app offers a branch head no way to add "
    "their own levels to a shared subject's offerings, though the endpoint allows it. The "
    "version 2.11 row is moved to its place in this log, and the control page, which "
    "version 2.11 left reading as version 2.10, is brought up to date. Verified against "
    "backend 33473d77; backend evidence only, and nothing here claims deployment."
)


def patch_m13() -> None:
    doc = Document(str(source_path(M13_DIR, M13_STEM, M13_SOURCE)))
    cover = doc.tables[0]
    rule51 = table_headed(doc, "Branch scope of academic structure")
    fr006, fr007 = fr_table(doc, "FR-006"), fr_table(doc, "FR-007")
    fr011, fr013 = fr_table(doc, "FR-011"), fr_table(doc, "FR-013")
    fr016, fr017 = fr_table(doc, "FR-016"), fr_table(doc, "FR-017")
    refusals = table_headed(doc, "Condition")
    log = table_headed(doc, "Version")

    # Control page.
    status_cell = row_by_label(cover, "Status").cells[1]
    supersedes = row_by_label(cover, "Supersedes").cells[1].text.strip()
    set_cover(cover, {
        "Version": M13_TARGET,
        "Date": REVIEW_DATE,
        "Supersedes": f"v{M13_SOURCE} (September 2026), v2.10 (September 2026), {supersedes}",
        "Status": M13_STATUS,
        "Verified against": M13_VERIFIED,
        "Source MRD": f"XVS Module Requirements Document v{MRD_VERSION} | Module 13",
    }, status_cell)

    # Section 5.1.
    append_cell(row_by_label(rule51, "Who may create a shared item").cells[1], M13_RULE_51_TAIL)

    # Section 8.
    append_cell(row_by_label(fr011, "The rule, on writes").cells[1], M13_FR011_WRITES_TAIL)
    append_cell(row_by_label(fr006, "Business rules").cells[1], M13_FR006_RULE_TAIL)
    edit_cell(row_by_label(fr007, "Business rules").cells[1], M13_FR007_OLD, M13_FR007_NEW)
    edit_cell(row_by_label(fr013, "Refusals").cells[1], M13_FR013_OLD, M13_FR013_NEW)
    append_cell(row_by_label(fr016, "Business rules").cells[1], M13_FR016_RULE)
    append_cell(row_by_label(fr016, "Acceptance").cells[1], M13_FR016_ACCEPTANCE)
    clone_fr_block(
        fr017._tbl, fr_heading(fr017), fr017,
        "FR-018  Shared rows are read-only to a branch administrator",
        "FR-018  Shared Rows Are Read-Only to a Branch Administrator", M13_FR018,
    )

    # Section 9.
    append_paragraph(body_paragraph(doc, "All endpoints are under /v1/academics/"),
                     M13_API_TAIL)

    # Section 10.
    anchor = find_row(refusals, "A branch-bound caller creating a school-wide row")
    add_row_after(anchor, anchor, M13_REFUSAL)

    # Section 11.
    for heading_text, bullets in M13_BULLETS.items():
        add_bullets_before(doc, heading_text, bullets)

    # Section 14: oldest first.
    reorder_log(log, M13_SOURCE)
    append_row(log, log.rows[-2], [M13_TARGET, SHORT_DATE, M13_CHANGE_SUMMARY])

    title = f"XVS M13 Academic Structure Functional Requirements Document v{M13_TARGET}"
    finish(doc, ROOT / M13_DIR / f"{M13_STEM}_v{M13_TARGET}.docx", title, M13_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 14 - Academic Calendar and Timetables
# ═════════════════════════════════════════════════════════════════════════════

M14_DIR = "14-timetable-and-calendar"
M14_STEM = "XVS_M14_Academic_Calendar_and_Timetables_Functional_Requirements_Document"
M14_SOURCE, M14_TARGET = "3.4", "3.5"

M14_STATUS = (
    "Built and verified. Version 3.5 records that an exam timetable publishes over a shared "
    "hall or an invigilator between two rooms (FR-017), and that a shared row is read-only "
    "to a branch administrator (FR-020). The school year and the term lifecycle remain "
    "owned upstream by M13."
)

M14_VERIFIED = (
    "Backend main at 9526abd2, 27 September 2026, holding 33473d77, and at 232c45a9 for the "
    "school's own day (480a4c87, 39c256a6). Every claim this version adds was read from the "
    "committed code and test files there; nothing uncommitted was read. schools.vs_calendar "
    "ran 251 tests, all passing, at 33473d77 and at 9526abd2. An inherited line "
    "reference points to a file rather than to a line."
)

#: The school's own day, which answers FR-005's "today" and section 13, decision 9.
M14_TIMEZONE_RULE = (
    "The date defaults to the school's own today: the calendar day in the school's own time "
    "zone, display.timezone, Africa/Lagos by default, read through vs_config.clock."
    "tenant_today. The server still runs TIME_ZONE = \"UTC\" with USE_TZ = True, and the day "
    "is the school's rather than the server's, so a resolution made just after midnight in "
    "Lagos answers with the new day. Callers who want another date pass ?on= explicitly. "
    "Section 13, decision 9."
)
M14_FR006_TRIGGER_OLD = "omitting on means today, with the timezone caveat in FR-005."
M14_FR006_TRIGGER_NEW = "omitting on means the school's own today, as in FR-005."
M14_TIMEZONE_DEPENDENCY = (
    "Answered. Each school keeps its own time zone in display.timezone, a configuration value "
    "at platform and school scope, Africa/Lagos by default, and FR-005's today is the "
    "calendar day in it (480a4c87, 39c256a6). Section 13, decision 9 closes."
)
M14_DECISION9_QUESTION = "9. Whose \"today\" does FR-005 use? CLOSED in version 3.5."
M14_DECISION9_ANSWER = (
    "Answered by the platform. Each school keeps its own time zone, display.timezone, "
    "Africa/Lagos by default, and this module's today is the calendar day in that zone rather "
    "than the server's UTC day (480a4c87, 39c256a6). The server still runs TIME_ZONE = "
    "\"UTC\"; a stored instant is unchanged, and only the day is the school's."
)
M14_OPEN_DECISIONS_OLD = (
    "all but decision 11, which the platform's plan gate answers at version 3.3, are still open"
)
M14_OPEN_DECISIONS_NEW = (
    "all but decisions 9 and 11, which the school's own time zone answers at version 3.5 and "
    "the platform's plan gate at version 3.3, are still open"
)

M14_FR002_WRITES_TAIL = (
    " Changing an existing event is narrower than creating one: a branch-bound caller "
    "changes only an event at one of their own branches, and a school-wide event is "
    "read-only to them (FR-020)."
)
M14_FR002_ACCEPTANCE = (
    " (9) A branch-bound caller's edit of a school-wide event is refused 403 "
    "SHARED_RECORD_READ_ONLY (" + SHARED_TESTS + ", CalendarRowsTests)."
)

M14_FR013_RULE = (
    " A shared class's grid is read-only to a branch-bound caller: every write to it, and "
    "the lesson preview, is refused 403 SHARED_RECORD_READ_ONLY (FR-020)."
)

M14_FR016_RULE7_OLD = "A school-wide period stays every branch's to build on."
M14_FR016_RULE7_NEW = (
    "An exam hung off a school-wide period is school-wide too, so it is created by a "
    "school-wide administrator: a branch-bound caller is refused 403 "
    "SHARED_RECORD_READ_ONLY, and adds their own classes' papers once it exists (FR-020)."
)
M14_FR016_RULE9 = (
    " (9) A paper belongs to its class: a branch-bound caller schedules only their own "
    "branch's classes, inside a school-wide exam too, and a shared class is refused 403 "
    "SHARED_RECORD_READ_ONLY on a paper's create or update (FR-020)."
)
M14_FR016_ACC_OLD = (
    "and a school-wide period still can be (tests/test_exams.py, ExamSecurityTests)."
)
M14_FR016_ACC_NEW = (
    "and an exam under a school-wide period is refused to a branch administrator, 403 "
    "SHARED_RECORD_READ_ONLY (tests/test_exams.py, ExamSecurityTests, "
    "test_a_school_wide_exam_is_set_up_by_a_school_wide_administrator)."
)

M14_FR017_DESCRIPTION = (
    "Move a class timetable or an exam timetable from draft to published, and refuse a "
    "class timetable while its grid still contradicts itself. An exam timetable is not "
    "refused on a clash: its one impossible clash is refused when the paper is written. "
    "Version 1.0 gives publication to exams alone; it applies to both, and the class "
    "timetable is the one a school actually publishes every term."
)
M14_FR017_GUARD = (
    "Publishing a class timetable recomputes FR-014's three rules over its grid and "
    "refuses if any clash remains, 409 TIMETABLE_HAS_CLASHES, with the offending slot ids "
    "in error.detail. It recomputes rather than reading a stored flag, deliberately: a "
    "clash is a relationship between two rows, and editing either of them can create or "
    "resolve one in a slot nobody touched, including one at another branch, so a cached "
    "flag is a cache with no invalidation. Section 6.5. An exam timetable is not gated. A "
    "class sitting two papers in one sitting is refused when the paper is written, 409 "
    "CLASS_ALREADY_SITTING (FR-016). A room holding several classes' papers and one "
    "invigilator between two rooms are things a school does on purpose, since the Main "
    "Hall seats JSS1 to JSS3 together, and the editor warns about each as the paper is "
    "written. Refusing them again here refused the same guess with no way for the school "
    "to say it meant it, so a school with one hall could never publish. Both publish."
)
M14_FR017_ACC_EDITS = [
    ("(3) The same for a room clash.", "(3) The same for a room clash on a class timetable."),
    ("(5) An exam holding only an invigilator clash publishes, because that rule warns "
     "rather than refusing, asserted so the asymmetry is deliberate.",
     "(5) An exam holding a shared hall publishes, and so does one holding an invigilator "
     "between two rooms, because both warn rather than refuse (tests/test_exams.py, "
     "test_a_shared_hall_publishes and test_an_invigilator_between_two_rooms_publishes)."),
]
M14_FR017_ACC_TAIL = (
    " (8) A branch-bound caller may not publish a school-wide exam or a shared class's "
    "timetable, 403 SHARED_RECORD_READ_ONLY (FR-020)."
)

M14_EXAM_HOLDERS = (
    "school_admin, branch_admin (delete: school_admin only; view also teacher). .delete and "
    ".publish are SENSITIVE"
)

M14_FR020 = [
    ("Description",
     "A branch-bound caller reads the school-wide rows beside their own branch's (FR-002, "
     "FR-019) and changes a row only when it is wholly their own branch's. A shared row is "
     "read-only to them, and a write to it is refused 403 SHARED_RECORD_READ_ONLY. It is the "
     "rule M13 FR-018 carries for the structure and Module 12 for staff records, from the "
     "same predicate, vs_rbac.scoping.caller_may_change."),
    ("Why",
     "The school's holidays, its everyday bell schedule and a school-wide exam week belong "
     "to every branch, which is why a branch reads them and why it may not change them: a "
     "holiday moved at Ikeja moves at Lekki too, with nobody at Lekki told. The behaviour change a "
     "school will notice, confirmed by the product owner on 26 September 2026: a branch "
     "head no longer creates an exam under a school-wide exam period, and adds their own "
     "classes' papers to the exam a school-wide administrator creates."),
    ("Actors",
     "Every write in FR-001, FR-011 to FR-013, FR-016 and FR-017. This requirement adds no "
     "permission key, for the reason FR-002 gives."),
    ("Whose row it is",
     "An event, a room and a period by their own branch. An exam by the exam period it "
     "hangs off. A lesson, an exam paper and a class's grid by the class. So Ikeja "
     "schedules its own classes' papers inside a school-wide exam, but cannot create, "
     "rename or publish that exam. The function is M13's, schools.vs_academics.services."
     "scoping.row_branch_ids, so the two modules cannot judge one row differently."),
    ("Where it applies",
     "At the module's get_object, so at every detail view at once: another branch's event, "
     "room, period, exam, exam paper or lesson answers 404, and a shared one answers 403 to "
     "a write. At every hand-written write as well: a class timetable's put, duplicate, "
     "clear and publish; a lesson's create and the lesson preview; an exam's create, judged "
     "by its period; a paper's create and update, judged by its class; and an exam's "
     "publish."),
    ("Reads it closed",
     "The event, period, room, exam, exam paper and lesson detail routes resolved a row by "
     "tenant alone, so another branch's row was readable and writable by id, and the exam "
     "list and each exam's paper list were not narrowed by branch at all. Both lists now "
     "read inclusively: another branch's exam period is not listed, and another branch's "
     "classes' papers inside a shared exam are not listed either."),
    ("can_manage",
     "Every event, room, period, exam, exam paper and lesson this module returns to a "
     "request carries can_manage, and so does each class in the class timetable picker, for "
     "its grid, so a screen hides the controls the server would refuse. The school app does "
     "(school-fe 2b53c61)."),
    ("Acceptance",
     "(1) Another branch's event is not found by id, a shared event is read-only, and the "
     "school-wide bell schedule is read-only. (2) A branch head schedules their own class "
     "in a shared exam, cannot schedule a shared class, does not see another branch's "
     "papers in it, and cannot publish it; another branch's exam and its papers are not "
     "found. (3) A branch head fills their own class's grid, a shared class's grid is "
     "read-only, another branch's lesson is not found by id, and the class picker says "
     "which grids the viewer may change. " + SHARED_TESTS + ", CalendarRowsTests, "
     "ExamTests and TimetableTests; the exam create in tests/test_exams.py, "
     "ExamSecurityTests."),
]

M14_BULLETS = {
    "11.2 Behaviour": [
        "A branch-bound caller is refused 403 SHARED_RECORD_READ_ONLY on a school-wide "
        "event, the school-wide bell schedule, a shared class's grid and a school-wide "
        "exam, and reads another branch's room, period, exam, exam paper or lesson by id "
        "as 404 (FR-020).",
    ],
}

M14_CHANGE_SUMMARY = (
    "Minor revision. Two changes. First, an exam timetable publishes over a shared hall and "
    "over an invigilator between two rooms. Both were refused at publication although the "
    "editor had already warned about each as it was written, with no way for a school to "
    "say it meant it, so a school with one hall could never publish. Both stay warnings on "
    "write and neither gates; a class sitting two papers in one sitting is still refused at "
    "the write, 409 CLASS_ALREADY_SITTING, and a class timetable still may not publish over "
    "a clash, 409 TIMETABLE_HAS_CLASHES. FR-017, FR-014, section 6.10, section 10 and "
    "section 11 follow, and FR-017's acceptance, which said a room clash blocked an exam "
    "and an invigilator clash did not, is corrected on both counts. Second, FR-020 is new: "
    "a shared row is read-only to a branch administrator, the rule M13 FR-018 carries for "
    "the structure, refused 403 SHARED_RECORD_READ_ONLY. An event, a room and a period are "
    "judged by their own branch, an exam by its exam period, and a lesson or an exam paper "
    "by its class, so a branch head schedules their own classes' papers inside a "
    "school-wide exam but no longer creates an exam under a school-wide exam period, which "
    "FR-016 had said stayed every branch's to build on. Detail routes that resolved another "
    "branch's event, period, room, exam, exam paper or lesson by id now answer 404, the "
    "exam list and the paper lists are narrowed inclusively, and every row carries "
    "can_manage. One correction: the exam keys' holders in sections 4.7 and 7.1 named "
    "academics.exam.manage, a key that no longer exists; they now read as seed_school_permissions grants them, "
    ".delete to school_admin alone and .delete and .publish SENSITIVE, which is what FR-016 "
    "and FR-017 already said. A second correction: FR-005's today, and FR-006's through it, is "
    "the calendar day in the school's own time zone, Africa/Lagos by default, not the server's "
    "UTC day, as this document had said (480a4c87, 39c256a6), so section 12's time-zone row "
    "and section 13, decision 9 close. The version 3.4 row is moved to its place in this log, and the control "
    "page, which version 3.4 left reading as version 3.3, is brought up to date. Verified "
    "against backend 33473d77 and 9526abd2, and the school's day at 232c45a9; backend evidence "
    "only, and nothing here claims "
    "deployment."
)


def patch_m14() -> None:
    doc = Document(str(source_path(M14_DIR, M14_STEM, M14_SOURCE)))
    cover = doc.tables[0]
    exam_slot = table_headed(doc, "ExamSlot")
    branch_611 = table_with_row(doc, "May one class's timetable span two branches?")
    fr002, fr013 = fr_table(doc, "FR-002"), fr_table(doc, "FR-013")
    fr014, fr016 = fr_table(doc, "FR-014"), fr_table(doc, "FR-016")
    fr017, fr019 = fr_table(doc, "FR-017"), fr_table(doc, "FR-019")
    refusals = table_headed(doc, "Condition")
    log = table_headed(doc, "Version")

    # Control page.
    status_cell = row_by_label(cover, "Status").cells[1]
    set_cover(cover, {
        "Version": M14_TARGET,
        "Date": REVIEW_DATE,
        "Supersedes": f"v{M14_SOURCE}",
        "Status": M14_STATUS,
        "Verified against": M14_VERIFIED,
        "Source MRD": f"XVS Module Requirements Document v{MRD_VERSION} | Module 14",
    }, status_cell)
    upstream = row_by_label(cover, "Upstream document").cells[1]
    keep_rows_whole(cover)
    edit_cell(upstream, "M13 Academic Structure FRD v2.10,", f"M13 Academic Structure FRD v{M13_TARGET},")
    edit_cell(upstream,
              "whose FR-002, FR-003 and FR-011 are the authority on the session lifecycle, the "
              "term lifecycle and the branch scope of everything this module schedules.",
              "whose FR-002, FR-003, FR-011 and FR-018 are the authority on the session "
              "lifecycle, the term lifecycle, the branch scope of everything this module "
              "schedules and who may change a shared row.")

    # Sections 4.7 and 7.1: the exam keys' holders, as seed_school_permissions grants them.
    exam_rows = [row for table in doc.tables for row in table.rows
                 if row.cells[0].text.strip().startswith("academics.exam.view / .create")]
    if len(exam_rows) != 2:
        raise ValueError(f"Expected the exam key row in sections 4.7 and 7.1, found {len(exam_rows)}")
    for row in exam_rows:
        edit_cell(row.cells[1], "school_admin, branch_admin (manage: school_admin; view also teacher)",
                  M14_EXAM_HOLDERS)

    # Section 4.3.
    append_paragraph(body_paragraph(doc, "Holding a key is not the same as being scoped to a row."),
                     " Changing a row is narrower than reading it: FR-020.")

    # Section 6.10.
    edit_cell(row_by_label(exam_slot, "invigilator").cells[1],
              "as a warning in FR-014 and a refusal on publish in FR-017.",
              "as a warning in FR-014. It does not block publication (FR-017).")
    replace_text([body_paragraph(doc, "ℹ Note which of the exam clashes is a constraint")],
                 "FR-014 states all three and FR-017 says which of them block publication.",
                 "FR-014 states all three. Neither warning blocks publication: FR-017 "
                 "publishes an exam timetable over both, and the constraint is the only "
                 "refusal.")

    # Section 6.11.
    edit_cell(row_by_label(branch_611, "May one class's timetable span two branches?").cells[2],
              "It is not a theoretical case: a school-wide class is visible to both branches' "
              "admins under the inclusive read, so two branch admins can each start building "
              "a grid for JSS1 A in their own rooms without either knowing about the other. "
              "This rule is what stops them, and a single-branch school can never reach it.",
              "A school-wide class's grid is read-only to a branch administrator (FR-020), "
              "so two branch admins can no longer each start one in their own rooms; a "
              "school-wide administrator building JSS1 A's week can still reach for a room at "
              "the other branch by mistake, and this rule is what stops them. A "
              "single-branch school can never reach it.")

    # Section 8.
    intro = body_paragraph(doc, "Nineteen requirements, every one of them this module's own.")
    replace_text([intro], "Nineteen requirements, every one of them",
                 "Twenty requirements, every one of them")
    append_paragraph(intro, " Version 3.5 adds FR-020, the rule that a shared row is "
                            "read-only to a branch administrator.")

    append_cell(row_by_label(fr002, "The rule, on writes").cells[1], M14_FR002_WRITES_TAIL)
    append_cell(row_by_label(fr002, "Acceptance").cells[1], M14_FR002_ACCEPTANCE)

    append_cell(row_by_label(fr013, "Business rules").cells[1], M14_FR013_RULE)

    edit_cell(row_by_label(fr014, "Actors").cells[1],
              "and on the publish guard in FR-017", "and on the class timetable's publish "
              "guard in FR-017")
    edit_cell(row_by_label(fr014, "Warn, do not refuse").cells[1],
              "The refusal happens once, at publication, and is FR-017.",
              "The refusal happens once, when a class timetable is published, and is FR-017; "
              "an exam timetable has no such gate, because its one impossible clash is "
              "refused at the write.")

    fr016_rules = row_by_label(fr016, "Business rules").cells[1]
    edit_cell(fr016_rules, M14_FR016_RULE7_OLD, M14_FR016_RULE7_NEW)
    append_cell(fr016_rules, M14_FR016_RULE9)
    edit_cell(row_by_label(fr016, "Acceptance").cells[1], M14_FR016_ACC_OLD, M14_FR016_ACC_NEW)

    set_cell(row_by_label(fr017, "Description").cells[1], M14_FR017_DESCRIPTION)
    set_cell(row_by_label(fr017, "The guard, which is the point of the requirement").cells[1],
             M14_FR017_GUARD)
    edit_cell(row_by_label(fr017, "Where the refusal sits, and why not earlier").cells[1],
              "This is the module's only hard refusal on a clash, and",
              "This is the module's only hard refusal on a clash, for a class timetable, and")
    fr017_acceptance = row_by_label(fr017, "Acceptance").cells[1]
    for old, new in M14_FR017_ACC_EDITS:
        edit_cell(fr017_acceptance, old, new)
    append_cell(fr017_acceptance, M14_FR017_ACC_TAIL)

    clone_fr_block(
        fr019._tbl, fr_heading(fr019), fr019,
        "FR-020  Shared rows are read-only to a branch administrator",
        "FR-020  Shared Rows Are Read-Only to a Branch Administrator", M14_FR020,
    )

    # Section 10.
    clash = find_row(refusals, "Publishing a timetable that still holds a teacher, room or class clash")
    set_cell(clash.cells[0], "Publishing a class timetable that still holds a teacher, room or "
                             "class clash")
    append_cell(clash.cells[2], ". An exam timetable is never refused on a clash; its one "
                                "impossible clash is CLASS_ALREADY_SITTING, at the write")
    other = find_row(refusals, "An event of another branch, outside the caller's visible set")
    set_cell(other.cells[0], "An event, room, period, exam, exam paper or lesson of another "
                             "branch, outside the caller's visible set")
    edit_cell(other.cells[2], "(FR-002)", "(FR-002, FR-020)")
    add_row_after(other, other, [
        "A branch-bound caller changing a shared row",
        "403",
        "SHARED_RECORD_READ_ONLY, with a message naming who can make the change (FR-020). A "
        "school-wide event, bell period or exam, including an exam created under a "
        "school-wide exam period, and a shared class's grid or papers",
    ])

    # Section 11.
    add_bullets_before(doc, "11.2 Behaviour", M14_BULLETS["11.2 Behaviour"])
    set_run_text(
        body_paragraph(doc, "An exam holding only an invigilator clash publishes"),
        "An exam holding a shared hall publishes, and so does one holding an invigilator "
        "between two rooms, while a class timetable holding a room clash does not: the "
        "pair proves that an exam timetable's clashes warn and never gate (FR-017).",
    )

    # Closing notes.
    replace_text([body_paragraph(doc, "ℹ Two places where this document and the revised design brief")],
                 "and the only refusal is on publication.",
                 "and the only refusal is on publishing a class timetable; an exam timetable "
                 "publishes over its warnings.")

    # The school's own day: FR-005, FR-006, section 12 and section 13, decision 9.
    set_cell(row_by_label(fr_table(doc, "FR-005"), "Timezone").cells[1], M14_TIMEZONE_RULE)
    edit_cell(row_by_label(fr_table(doc, "FR-006"), "Trigger").cells[1],
              M14_FR006_TRIGGER_OLD, M14_FR006_TRIGGER_NEW)
    zone = row_by_label(table_with_row(doc, "A per-tenant timezone"), "A per-tenant timezone")
    set_cell(zone.cells[2], M14_TIMEZONE_DEPENDENCY)
    decision9 = row_by_label(table_with_row(doc, '9. Whose "today" does FR-005 use?'),
                             '9. Whose "today" does FR-005 use?')
    set_cell(decision9.cells[0], M14_DECISION9_QUESTION)
    set_cell(decision9.cells[1], M14_DECISION9_ANSWER)
    replace_text([body_paragraph(doc, "Each of the following is a product question")],
                 M14_OPEN_DECISIONS_OLD, M14_OPEN_DECISIONS_NEW)

    # Section 14: oldest first.
    reorder_log(log, M14_SOURCE)
    append_row(log, log.rows[-2], [M14_TARGET, SHORT_DATE, M14_CHANGE_SUMMARY])

    title = (
        "XVS M14 Academic Calendar and Timetables Functional Requirements Document "
        f"v{M14_TARGET}"
    )
    finish(doc, ROOT / M14_DIR / f"{M14_STEM}_v{M14_TARGET}.docx", title, M14_TARGET)


def main() -> None:
    patch_m13()
    print(f"Wrote {M13_DIR} v{M13_TARGET}")
    patch_m14()
    print(f"Wrote {M14_DIR} v{M14_TARGET}")


if __name__ == "__main__":
    main()
