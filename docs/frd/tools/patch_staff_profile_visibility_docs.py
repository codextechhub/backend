#!/usr/bin/env python3
"""Cut M12 v2.15, M04 v1.30 and MRD v2.95: a school decides what each
relationship reads of a colleague's staff profile.

What changed in the backend (ac0b4c3e, de392b79), and therefore in the documents:

* Every member of staff opens any colleague from the organogram. The school's
  staff profile policy, the vs_config value staff.profile_visibility set at
  /v1/i/me/settings/staff-profiles/, decides which sections each relationship
  reads: the person themselves, anybody above them in the reporting line
  through main, acting and dotted posts, and any other colleague. A reader whose
  keys reach the person reads as their role allows. The contact card always
  shows. Field Access stays the ceiling for fields.
* The Teacher role no longer holds school.teachers.view, in the library and in
  every school's native Teacher role (vs_rbac 0030).
* GET /v1/i/me/staff/mine/ answers the caller's own staff id.
* The organogram summary counts a branch-bound reader's branches; the chart
  stays whole. A post a live approval step or group names is not deleted.
* A staff document's file asks school.staff_records.view, the key its tab asks;
  it had opened on school.teachers.view. The roles read leaves the overrides key
  out for a reader without the override key.
* The staff dates that mean today are the school's own day.

The owner confirmed on 28 September 2026 that the organogram sits in Core.

    python tools/patch_staff_profile_visibility_docs.py
"""
from __future__ import annotations

import copy

from docx import Document

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    replace_cell,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    box_of,
    edit_blurb,
    edit_box,
    insert_row_after,
    row_labelled,
    table_headed,
)
from patch_school_organogram_docs import row_starting_cell, table_headed_prefix
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    fr_table,
    insert_after,
    normalise_change_log,
    repair_ooxml,
    row,
    row_starting,
)

import patch_record_history_docs

REVIEW_DATE = "29 September 2026"
SHORT_DATE = "29 Sep 2026"
CODE_BASELINE = (
    "Backend main at 281d3e66, 29 September 2026, read for the staff profile policy, the "
    "Teacher role's directory key and the own-record route (ac0b4c3e), and the document "
    "file key and the roles read (de392b79). The school app's screens are committed "
    "(school-fe afb08c1, 43fb127), pending release"
)
TEST_EVIDENCE = (
    "Verified by schools.vs_staff (422 tests, 32 of them the profile policy's and 3 the "
    "own-record route's), schools.vs_schools (430), vs_config (127) and vs_rbac (895), each "
    "run on its own and all passing, and by opening a colleague, a report and one's own "
    "record as a Holy Cross teacher on the running app. Nothing here claims deployment."
)
MRD_SOURCE, MRD_TARGET = "2.94", "2.95"

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


# ── M12 Staff Management ─────────────────────────────────────────────────────

M12_DIR = "12-staff-management"
M12_STEM = "XVS_M12_Staff_Management_Functional_Requirements_Document"
M12_SOURCE, M12_TARGET = "2.14", "2.15"

M12_SUMMARY = (
    "Minor revision. Adds FR-026: a school decides how much of a colleague's profile each "
    "relationship reads. Every member of staff opens any colleague from the chart, and the "
    "policy staff.profile_visibility, set at /v1/i/me/settings/staff-profiles/ under "
    "school.settings.view and school.field_access.update, grants sections to the person "
    "themselves, to anybody above them through main, acting and dotted posts, and to any other "
    "colleague; a reader whose keys reach the person reads as their role allows, the contact "
    "card always shows, the most generous standing wins, and Field Access stays the ceiling. "
    "The record carries profile_view and visible_sections, and a restricted read carries only "
    "the sections shown. The Teacher role no longer holds school.teachers.view (vs_rbac 0030). "
    "GET /v1/i/me/staff/mine/ answers the caller's own staff id. The organogram summary "
    "counts a branch-bound reader's branches while the chart stays whole, and a post an "
    "approval step or group names is not deleted. FR-006 and FR-007 named school.teachers.view "
    "where the records keys have decided since v2.2, and a document's file had opened on the "
    "directory key; both now say, and do, school.staff_records. The roles read leaves the "
    "overrides key out for a reader without its key, the staff dates meaning today are the "
    "school's own day, and the organogram's Core band is confirmed. " + TEST_EVIDENCE
)

FR026_ROWS = [
    ("Description",
     "Let every member of staff open any colleague from the organogram, and let the school decide "
     "how much of the profile each relationship reads. The chart itself is not narrowed by this "
     "requirement (FR-024)."),
    ("Permission",
     "No key of its own. A reader whose keys reach the person, the key each section is read under "
     "held and the person inside the reader's branch scope, reads as their role allows. Anybody "
     "else inside the school reads what the policy grants their relationship. The policy is read "
     "under school.settings.view and written under school.field_access.update at GET and PUT "
     "/v1/i/me/settings/staff-profiles/, by a caller whose reach is the whole school."),
    ("Scoping",
     "The tenant first: another school's person is 404 on every route, whoever reads. Inside the "
     "school a relationship crosses branches, so a Lekki teacher opening an Ikeja colleague from "
     "the chart gets the contact card where the key path alone would answer 404. A reader counts "
     "as a colleague while their own staff record is not invited, resigned or terminated, or, "
     "with no staff record, while they hold school.organogram.view; a leaver is not open to "
     "colleagues."),
    ("Business rules",
     "(1) The profile is eight sections: contact, employment, personal, records, leave, teaching, "
     "history and roles. (2) Four audiences: ADMIN, whose keys reach the person and who is not "
     "configurable; SELF, the person themselves; LINE, anybody above the person in the reporting "
     "line at any depth, through a main, acting or dotted post; and COLLEAGUE, anybody else at the "
     "school. (3) The school sets, per configurable audience, which sections it shows. A school "
     "that has saved nothing shows the person their whole record, a line manager contact, "
     "employment, leave and teaching, and a colleague the contact card. (4) contact is shown to "
     "every audience and is added to a save that leaves it out; it carries the name, sign-in "
     "email, phone, photograph, post and branch. (5) The most generous standing wins: a reader "
     "reads the union of what their keys reach and what each audience they belong to is shown. "
     "(6) The policy grants by relationship and never beyond it: a teacher who holds no leave key "
     "reads the leave of the people under them because the school ticked leave for line managers, "
     "and nobody else's. (7) Field Access stays the ceiling: a field the reader's role has switched "
     "off is removed before a section is considered, except that a person reads their own record "
     "through the owner rule. (8) The record carries profile_view, full for a reader whose keys "
     "reach the person or who is the person, restricted otherwise, and visible_sections; a "
     "restricted read carries only the fields of the sections shown, and every tab outside them "
     "answers 403. A document file is readable where records are shown (FR-007) and a photograph "
     "wherever the contact card is. (9) Reading a past day (FR-022) stays with a full reader. (10) "
     "The relationship costs three queries at most, whatever the size of the chart, and is cached "
     "on the request. (11) A save is audited through the configuration engine with the policy "
     "before and after."),
    ("Acceptance",
     "(1) Another school's person is 404 on every route whatever the reader is. (2) A colleague at "
     "another branch reads the contact card only, with the branch named, and every other tab and "
     "a past day are refused (ColleagueTests). (3) A line manager two levels up across branches "
     "reads employment and leave; acting posts and dotted lines put a post under its manager; a "
     "peer is only a colleague; and a document opens for a line manager once records are shown to "
     "them (LineManagerTests). (4) A teacher reads leave for their report and not for a peer, and "
     "the school can take it away (TeacherLineManagerLeaveTests). (5) A switched-off field stays "
     "hidden whatever the grid says, while a person reads and corrects their own phone with the "
     "switch off (FieldAccessCeilingTests). (6) The default applies until a school saves one, one "
     "school's policy does not reach another, and a malformed policy is refused on every write "
     "path (PolicyTests). (7) The line check does not grow with the chart (QueryBudgetTests, "
     "schools.vs_staff.tests.test_profile_visibility, 32 tests)."),
]


def patch_m12() -> None:
    require_newest(str(ROOT / "functional-requirements" / M12_DIR / f"{M12_STEM}_v*.docx"),
                   M12_SOURCE)
    doc = Document(str(frd_path(M12_DIR, M12_STEM, M12_SOURCE)))
    set_control(doc, "Version", M12_TARGET)
    set_control(doc, "Date", REVIEW_DATE)
    set_control(doc, "Supersedes", f"v{M12_SOURCE}")
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")
    set_control(doc, "Verified against", CODE_BASELINE)

    # 8.1 Permission keys
    keys = table_headed(doc, "Key", "State", "Sensitivity", "Default holders", "Use")
    directory = row(keys, "school.teachers.view")
    keep_format(directory.cells[3], "school_admin, branch_admin")
    keep_format(directory.cells[-1], directory.cells[-1].text.rstrip() +
                " Not granted to teacher: vs_rbac 0030 takes it off every school's native Teacher "
                "role, and a teacher reads a colleague as far as the school's staff profile policy "
                "shows them (FR-026).")
    records = row(keys, "school.staff_records.view / .update")
    keep_format(records.cells[-1], records.cells[-1].text.rstrip() +
                " The file behind a document asks the same read key (FR-007).")

    # FR-004
    fr004 = fr_table(doc, "FR-004")
    append_to(fr004, "Scoping",
              " A reader outside that narrowing, or holding no staff key, opens a colleague "
              "inside the school as far as FR-026 shows them.")
    append_to(fr004, "Business rules",
              " (5) GET /v1/i/me/staff/mine/ answers the caller's own staff id, so a person "
              "holding no staff key reaches their own record, and 404 where they have no staff "
              "record at this school.")
    append_to(fr004, "Response",
              " profile_view reads full or restricted and visible_sections lists the sections "
              "the reader is shown (FR-026); a restricted read carries only those sections' "
              "fields. Tenure is counted to the school's own day.")

    # FR-005
    edit_cell(row(fr_table(doc, "FR-005"), "Notes").cells[-1],
              "shows nothing at all, not an empty block, to anyone else.",
              "leaves the overrides key out of the payload for anyone else, rather than "
              "sending it empty or null.")

    # FR-006 and FR-007: the records keys
    fr006 = fr_table(doc, "FR-006")
    keep_format(row(fr006, "Permission").cells[-1],
                "GET school.staff_records.view; POST, PATCH and DELETE school.staff_records.update. "
                "A reader the school's profile policy shows records to reads them without the key "
                "(FR-026).")
    edit_cell(row(fr006, "Business rules").cells[-1],
              "(2) The year, where given, is validated as plausible and no further.",
              "(2) The year, where given, is validated as plausible, no later than the school's "
              "own year, and no further.")
    fr007 = fr_table(doc, "FR-007")
    keep_format(row(fr007, "Permission").cells[-1],
                "GET school.staff_records.view; POST and DELETE school.staff_records.update. A "
                "reader the school's profile policy shows records to reads them without the key "
                "(FR-026).")
    edit_cell(row(fr007, "Business rules").cells[-1],
              "may never reach anybody else's without school.teachers.view.",
              "may never reach anybody else's without school.staff_records.view inside their "
              "branches, or a relationship the school's profile policy shows records to "
              "(FR-026). The file behind a document asks the same, so a role that lists staff "
              "but may not read their records cannot open a CV by its link.")

    # FR-024
    fr024 = fr_table(doc, "FR-024")
    append_to(fr024, "Scoping",
              " The summary counts a branch-bound reader's own branches and the school-wide "
              "staff; the chart and its lists are not narrowed.")
    edit_cell(row(fr024, "Business rules").cells[-1],
              "(8) A unit or post in use is not deleted: 409 with the blocking counts.",
              "(8) A unit or post in use is not deleted: 409 with the blocking counts, which for "
              "a post include the live approval steps and approver groups naming it "
              "(ORGANOGRAM_IN_USE, counted by the workflow engine's position_references).")
    append_to(fr024, "Business rules",
              " (10) An appointment starts, and ends by default, on the school's own day. (11) "
              "The chart is registered with the workflow engine, which climbs it for a school "
              "requester and names posts on it by code, asking this module to find a post, "
              "describe it and list its holders able to act (M07 FR-008).")
    edit_cell(row(fr024, "Acceptance").cells[-1],
              "(schools.vs_staff.tests.test_organogram, 41 tests)",
              "(10) A branch-bound reader's summary counts their branches and the chart stays "
              "whole (test_the_summary_follows_the_readers_branches_and_the_chart_does_not). (11) A "
              "post an approval step or group names is not deleted (test_a_post_a_step_or_a_group_"
              "names_is_not_deleted, schools.vs_staff.tests.test_organogram, 60 tests)")

    add_fr(doc, "FR-026  A school decides what each relationship reads of a profile",
           "FR-026  A School Decides What Each Relationship Reads of a Profile", FR026_ROWS)

    # 10. API endpoints
    insert_row_after(row_starting(table_headed(doc, "Method and path", "Permission", "Requirement"),
                                  "GET /v1/i/me/staff/search/"), [
        "GET /v1/i/me/staff/mine/",
        "Any active member of the school",
        "FR-004. The caller's own staff id, 404 where they have none; open before go-live.",
    ])
    insert_row_after(row_starting(table_headed(doc, "Method and path", "Permission", "Requirement"),
                                  "GET /v1/i/me/staff/organogram/summary/"), [
        "GET, PUT /v1/i/me/settings/staff-profiles/",
        "school.settings.view / school.field_access.update",
        "FR-026. The policy with the sections and audiences the grid draws from. Mounted "
        "beside the school's other settings; the policy is this module's.",
    ])

    # 11. Refusals
    anchor = row_starting(table_headed(doc, "Condition", "Status", "Code"),
                          "A roster asked for a branch the caller does not work in")
    anchor = insert_row_after(anchor, [
        "A staff profile policy naming an unknown audience or section, setting ADMIN, or "
        "leaving an audience out",
        "422",
        "INVALID_CONFIGURATION_VALUE, one sentence per audience in error.detail. contact is added "
        "where a save leaves it out, and is not a refusal (FR-026).",
    ])
    anchor = insert_row_after(anchor, [
        "A branch-bound caller saving the staff profile policy",
        "403",
        "SHARED_RECORD_READ_ONLY: only a school-wide administrator changes who reads staff "
        "profiles. Nothing is written.",
    ])
    insert_row_after(anchor, [
        "A tab outside the sections a restricted reader is shown",
        "403",
        "Not a code. The record read still answers, restricted (FR-026).",
    ])

    # 12. Acceptance
    insert_after(doc, "A person with no staff permission at all reaches their own record",
                 "A teacher holding no staff key opens a colleague at another branch from the "
                 "chart and reads the contact card, a line manager reads what the school shows "
                 "line managers, and nobody reads another school's person (FR-026). A role "
                 "holding school.teachers.view without school.staff_records.view cannot open a "
                 "staff document's file (DocumentFileTests).")

    # 14. Decisions
    edit_cell(row_starting_cell(doc, "11. Is staff management entitlement-gated?").cells[1],
              "school.staff_records and school.leave at Plus,",
              "school.staff_records and school.leave at Plus, and school.organogram at Core, which "
              "the owner confirmed on 28 September 2026, so every school draws its chart,")

    # 15. Traceability
    trace = table_headed(doc, "MRD capability", "Requirements", "State at e12e7fbf")
    keep_format(trace.rows[0].cells[2], "State at 281d3e66")
    audit = row(trace, "Audit history and field-level access")
    keep_format(audit.cells[1], "FR-012, FR-022, FR-023, FR-026")
    keep_format(audit.cells[2], audit.cells[2].text.rstrip() +
                " Beyond the keys, the school's staff profile policy decides what each "
                "relationship reads of a colleague's profile, never past Field Access (FR-026).")
    edit_paragraph(doc, "MRD v2.94 records Module 12", "MRD v2.94", f"MRD v{MRD_TARGET}")
    edit_paragraph(doc, "📌  MRD v2.94 keeps Module 12", "MRD v2.94", f"MRD v{MRD_TARGET}")

    log_change(doc, M12_TARGET, M12_SUMMARY)
    assert_absent_outside_log(doc, "State at e12e7fbf", "MRD v2.94",
                              "without school.teachers.view. This is stated")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M12_DIR, M12_STEM, M12_TARGET),
           f"{M12_STEM.replace('_', ' ')} v{M12_TARGET}", M12_TARGET)


# ── M04 Roles & Permissions ──────────────────────────────────────────────────

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.29", "1.30"

M04_SUMMARY = (
    "Minor revision. The Teacher role no longer holds school.teachers.view. The seed leaves it "
    "off the library template, and vs_rbac 0030 takes a granted row off the template and every "
    "school's native Teacher role, keyed teacher or teacher-<branch>, leaving an explicit deny, a "
    "custom role and every other role alone; a school that gave Teacher the key on purpose loses "
    "it and may grant it again, and the migration reverses. A teacher now reads a colleague "
    "through Module 12's staff profile policy, which shows sections by relationship without a "
    "key, never beyond the relationship, under Field Access as the ceiling, and is written under "
    "school.field_access.update. FR-021, FR-022's owner rule, the Module 12 dependency and "
    "traceability follow. " + TEST_EVIDENCE
)


def patch_m04() -> None:
    require_newest(str(ROOT / "functional-requirements" / M04_DIR / f"{M04_STEM}_v*.docx"),
                   M04_SOURCE)
    doc = Document(str(frd_path(M04_DIR, M04_STEM, M04_SOURCE)))
    set_cover_version(doc, M04_SOURCE, M04_TARGET)
    set_control(doc, "Version", M04_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Source scope", CODE_BASELINE)
    set_control(doc, "Code inspected", CODE_BASELINE)
    set_control(doc, "MRD baseline", f"XVS Module Requirements Document v{MRD_TARGET}")

    fr021 = fr_table(doc, "FR-021")
    append_to(fr021, "Current evidence",
              " Teacher does not carry school.teachers.view: the seed leaves it off the library "
              "template, and vs_rbac 0030 takes a granted row off the template and every "
              "school's native Teacher role, keyed teacher or teacher-<branch>, leaving an "
              "explicit deny, a custom role and every other role alone. A teacher reads "
              "colleagues through Module 12's staff profile policy instead.")
    append_to(fr021, "Acceptance",
              " Every school's Teacher role loses the directory key, custom roles, other roles "
              "and denies are left alone, and the reverse grants it again "
              "(TeacherDirectoryKeyWithdrawalTests).")
    append_to(fr021, "Limit",
              " A school that gave its Teacher role the directory key on purpose loses it with "
              "the default, because a migration cannot tell the two grants apart; it may grant "
              "it again from its roles screen.")

    edit_cell(row(fr_table(doc, "FR-022"), "Limit").cells[-1],
              "such as a member of staff reading and changing their own payroll bank details.",
              "such as a member of staff reading and changing their own payroll bank details, "
              "and their own middle name, date of birth, photograph and phone on the school "
              "staff record.")

    append_to(doc.tables[[i for i, t in enumerate(doc.tables)
                          if any(r.cells[0].text.strip() == "Module 12, Staff Management"
                                 for r in t.rows)][0]],
              "Module 12, Staff Management",
              " A key is not the only way into a staff profile: Module 12 shows a colleague's "
              "sections by the reader's relationship to them under the school's staff profile "
              "policy (staff.profile_visibility), never beyond the relationship and never past "
              "Field Access, and the policy is written under school.field_access.update by a "
              "caller whose reach is the whole school. Teacher holds no school.teachers.view.")

    keep_format(row_starting(doc.tables[2], "10. MRD Traceability").cells[1],
                f"Agreement with MRD v{MRD_TARGET}'s twenty-three capability entries")
    edit_paragraph(doc, "MRD v2.94 records Module 4", "MRD v2.94", f"MRD v{MRD_TARGET}")
    edit_box(table_headed_prefix(doc, "MRD RECONCILIATION").rows[0].cells[0], [
        ("sub", "• MRD v2.94 lists Module 4", "MRD v2.94", f"MRD v{MRD_TARGET}"),
        ("sub", "• The field registry and its tree", "MRD v2.94 words the entry", f"MRD v{MRD_TARGET} words the entry"),
    ])
    templates = row_labelled(doc, "Role templates")
    keep_format(templates.cells[-1], templates.cells[-1].text.rstrip() +
                " Teacher holds no school.teachers.view.")

    log_change(doc, M04_TARGET, M04_SUMMARY)
    assert_absent_outside_log(doc, "MRD v2.94")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M04_DIR, M04_STEM, M04_TARGET),
           f"{M04_STEM.replace('_', ' ')} v{M04_TARGET}", M04_TARGET)


# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_CONTENTS_NOTE = "Two module FRDs reconciled; no capability added"
MRD_INTRO = (
    "This revision reconciles the tracker to M12 v2.15 and M04 v1.30: a school decides what each "
    "relationship reads of a colleague's staff profile, the Teacher role gives up the staff "
    "directory key, and the organogram's Core band is confirmed. No capability is added and no "
    "status moves."
)
MRD_DELTA_ROWS = [
    ["Staff profile visibility", "Module 12 entry strengthened",
     "Every member of staff opens any colleague from the chart. The school's policy decides what "
     "the person, the line above them through main, acting and dotted posts, and any other "
     "colleague read; a reader whose keys reach the person reads as their role allows; the "
     "contact card always shows; Field Access stays the ceiling. The count stays at 11."],
    ["Teacher and the directory key", "Module 4 default changed",
     "Teacher no longer holds school.teachers.view, in the library and every school's native "
     "Teacher role (vs_rbac 0030); the profile policy decides what a teacher reads."],
    ["Own staff record", "Module 12",
     "GET /v1/i/me/staff/mine/ finds the caller's record, so a person holding no staff key "
     "reaches it, applies for leave and corrects their photograph, middle name, date of birth "
     "and phone."],
    ["Staff documents", "Module 12 correction",
     "A document's file asks school.staff_records.view, the key its tab asks; it had opened on "
     "the directory key."],
    ["Organogram", "Core, confirmed",
     "school.organogram sits in Core in the plan map, confirmed by the owner on 28 September "
     "2026. A branch-bound reader's summary counts their branches while the chart stays whole."],
    ["Module FRDs", "Two reconciled", "M04 v1.30, M12 v2.15."],
]
MRD_CHANGE_SUMMARY = (
    "A school decides what each relationship reads of a colleague's staff profile: the person, "
    "the line above them through main, acting and dotted posts, and any other colleague, with "
    "the contact card always shown, a reader whose keys reach the person reading as their role "
    "allows, and Field Access as the ceiling. The Teacher role gives up school.teachers.view "
    "(vs_rbac 0030). A person reaches their own record through GET /v1/i/me/staff/mine/. A staff "
    "document's file asks the records key its tab asks. The organogram's Core band is confirmed. "
    "Capability entries unchanged at 512; statuses do not move. M04 v1.30 and M12 v2.15. Backend "
    "main at 281d3e66 with ac0b4c3e and de392b79. Backend evidence; the school app's screens are "
    "committed and pending release; no deployment claim."
)


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    require_newest(str(folder / "XVS_Module_Requirements_Document_v*.docx"), MRD_SOURCE)
    doc = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta = table_headed_prefix(doc, f"v{MRD_SOURCE} capability delta")
    log = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Version"
               and len(t.rows[0].cells) == 3)
    assert index.rows[0].cells[5].text.strip() == "Entries"
    total_before = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for r in control.rows:
        label = r.cells[0].text.strip()
        if label == "Version":
            replace_cell(r.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(r.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(r.cells[1], CODE_BASELINE, size=9)
    for r in contents.rows:
        if r.cells[0].text.strip().startswith("5."):
            keep_format(r.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(r.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (text.startswith("5. ") and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    edit_blurb(doc, 4, "Documented by M04 FRD v1.29.",
               "The Teacher role holds no staff directory key: a school's staff profile policy "
               "decides what a teacher reads of a colleague. Documented by M04 FRD v1.30.")
    edit_blurb(doc, 12, "Documented by M12 FRD v2.14.",
               "Every member of staff opens any colleague from the chart, and the school decides "
               "what the person, the line above them and any other colleague read of the "
               "profile, the contact card always shown and Field Access the ceiling; a person "
               "reaches their own record whatever keys they hold. A staff document's file asks "
               "the records key its tab asks. Documented by M12 FRD v2.15.")

    total_after = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total_after != total_before:
        raise ValueError(f"Capability total moved from {total_before} to {total_after}")

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)
    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if len(intro) != 1:
        raise ValueError(f"{len(intro)} delta introductions found")
    mrd_tools.retitle(intro[0], MRD_INTRO)

    log.rows[1]._tr.addprevious(copy.deepcopy(log.rows[1]._tr))
    for cell, text in zip(log.rows[1].cells, (MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)):
        keep_format(cell, text)
    assert_absent_outside_log(doc, "Documented by M12 FRD v2.14", "Documented by M04 FRD v1.29",
                              "v2.94 capability delta")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m12()
    patch_m04()
    patch_mrd()


if __name__ == "__main__":
    main()
