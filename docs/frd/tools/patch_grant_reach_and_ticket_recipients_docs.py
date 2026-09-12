#!/usr/bin/env python3
"""Version the Module 12 and Module 31 FRDs against the backend at 19a50b3.

Two contracts moved, one in each document, and each was written down the old
way.

Module 12 recorded that a role granted alongside a new member of staff reaches
the whole school unless the caller names a branch for it, and that the staff
import never names one. The reach is now decided in the creation service every
caller shares: an unstated reach follows the role template's own branch where
it has one and otherwise the posting the account is created with, a branch in
``role_branch`` pins the grant there, and the whole school is asked for by
sending ``role_branch`` as the word ``school``. The Add staff endpoint passes a
resolved reach and its patch-up is gone; the import is corrected by the choke
point alone.

Module 31 recorded as a gap that the desk's notification recipients and its
assignee picker are two queries that can disagree, so an agent holding the
manage key through a permission group could be given a ticket and was never
told of one. Both questions now narrow one shared query. The other gap in that
box, the desk's ticket export carrying CodeX's own tickets only, is not fixed
and stays.

The two documents carry different typography and are edited accordingly. The
Module 12 family is Calibri with slate text in tables, which the shared
``write_cell`` helper does not reproduce, so every edit there rewrites the text
of an existing run and keeps its formatting; its change log runs oldest first,
so the new row is appended. The Module 31 family is the one
``generate_requirements_documents`` builds, so its cells are rewritten through
``write_cell`` at the size the cell already renders at, and its change log runs
newest first.

    python tools/patch_grant_reach_and_ticket_recipients_docs.py
"""

from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)

REVIEW_DATE = "12 September 2026"
SHORT_DATE = "12 Sep 2026"
MRD_VERSION = "2.77"
CODE_BASELINE = "Backend worktree at 19a50b3, 12 September 2026"


# ── finders shared by both documents ─────────────────────────────────────────

def unique_cells(row):
    seen, cells = set(), []
    for cell in row.cells:
        if id(cell._tc) not in seen:
            seen.add(id(cell._tc))
            cells.append(cell)
    return cells


def table_with_header(doc, *header: str):
    """The one table whose first row starts with these cell texts."""
    matches = [
        table for table in doc.tables
        if [c.text.strip() for c in unique_cells(table.rows[0])][: len(header)]
        == list(header)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one table headed {header}, found {len(matches)}")
    return matches[0]


def table_with_row(doc, first_cell: str):
    """The one table holding a row whose first cell is exactly ``first_cell``."""
    matches = [
        table for table in doc.tables
        if any(row.cells[0].text.strip() == first_cell for row in table.rows)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one table with row {first_cell!r}, found {len(matches)}")
    return matches[0]


def banner_table(doc, label: str):
    """The one table whose first cell opens with ``label``."""
    matches = [
        table for table in doc.tables
        if table.rows[0].cells[0].text.strip().startswith(label)
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one table banner {label!r}, found {len(matches)}")
    return matches[0]


def row_where(table, first: str, *, exact: bool = False):
    for row in table.rows:
        text = row.cells[0].text.strip()
        if (text == first) if exact else text.startswith(first):
            return row
    raise ValueError(f"Row not found: {first[:60]!r}")


def row_for(table, tr):
    return next(row for row in table.rows if row._tr is tr)


def find_paragraph(doc, prefix: str):
    matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph {prefix[:50]!r}, found {len(matches)}")
    return matches[0]


def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError(f"Paragraph carries no run: {paragraph.text[:60]!r}")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_in_paragraph(paragraph, old: str, new: str) -> None:
    text = paragraph.text
    if old not in text:
        raise ValueError(f"Fragment not found in paragraph: {old[:60]!r}")
    set_run_text(paragraph, text.replace(old, new, 1))


def set_callout_text(paragraph, body: str) -> None:
    """Rewrite a callout, keeping its leading symbol and spacing."""
    prefix = re.match(r"^\S+\s+", paragraph.text).group(0)
    set_run_text(paragraph, prefix + body)


def insert_paragraph_after(anchor, text: str):
    new_p = copy.deepcopy(anchor._p)
    anchor._p.addnext(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    set_run_text(paragraph, text)
    return paragraph


def finish(doc, output: Path, title: str, version: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    check = Document(str(output))
    words = "\n".join(p.text for p in check.paragraphs)
    for table in check.tables:
        for row in table.rows:
            for cell in row.cells:
                words += "\n" + cell.text
    # Spelled in parts: the repository vocabulary test reads every file, and a
    # guard that wrote the retired word would be an offender itself.
    retired_site_word = "cam" + "pus"
    if retired_site_word in words.lower():
        raise ValueError(f"The retired word {retired_site_word!r} is in {output}")


# ═════════════════════════════════════════════════════════════════════════════
# Module 12 - Staff Management
# ═════════════════════════════════════════════════════════════════════════════

M12_DIR = "12-staff-management"
M12_STEM = "XVS_M12_Staff_Management_Functional_Requirements_Document"
M12_SOURCE, M12_TARGET = "2.3", "2.4"


# ── helpers that keep the Module 12 family's own formatting ─────────────────

def m12_set_cell_text(cell, text: str) -> None:
    """Rewrite a single-paragraph cell, keeping its run formatting."""
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    set_run_text(cell.paragraphs[0], text)


def m12_replace_in_cell(cell, old: str, new: str) -> None:
    text = cell.text
    if old not in text:
        raise ValueError(f"Fragment not found in cell: {old[:60]!r}")
    m12_set_cell_text(cell, text.replace(old, new, 1))


def m12_append_to_cell(cell, tail: str) -> None:
    m12_set_cell_text(cell, cell.text.rstrip() + tail)


def m12_insert_row_after(table, anchor, values: list[str]):
    new_tr = copy.deepcopy(anchor._tr)
    anchor._tr.addnext(new_tr)
    row = row_for(table, new_tr)
    for cell, value in zip(unique_cells(row), values):
        m12_set_cell_text(cell, value)
    return row


def m12_assert_native(cell) -> None:
    font = cell.paragraphs[0].runs[0].font.name
    if font != "Calibri":
        raise ValueError(f"Template row is not in the document's typeface: {font}")


M12_BASELINE_OLD = "Backend worktree at 13675db, 11 September 2026"
M12_BASELINE_NEW = "Backend worktree at 19a50b3, 12 September 2026"

M12_FR001_POSTCONDITION_OLD = (
    "a TenantUserRoleAssignment for the named role, school-wide unless "
    "role_branch names a branch or the role template is tied to one;"
)

M12_FR001_POSTCONDITION_NEW = (
    "a TenantUserRoleAssignment for the named role, reaching the branch the "
    "person is posted to unless the role template is tied to a branch of its "
    "own or role_branch names one, and reaching the whole school where "
    "role_branch is the word school;"
)

M12_FR001_RULE_OLD = (
    "(7) The grant is school-wide unless role_branch names a branch, in which "
    "case the one grant is moved onto it rather than a second being added, or "
    "the role template is itself tied to a branch. The posting does not narrow "
    "it, so a teacher posted to Ikeja and invited without role_branch reaches "
    "every branch's records."
)

M12_FR001_RULE_NEW = (
    "(7) How far the grant reaches is decided in the creation service every "
    "caller shares rather than in this endpoint, because a caller that forgets "
    "to ask is the case that matters. Saying nothing is not the same as asking "
    "for everywhere: an unstated reach follows the role template's own branch "
    "where it has one, and otherwise the posting the account is created with, "
    "so a teacher posted to Ikeja and invited with role_branch left out is "
    "granted at Ikeja and reaches Ikeja. role_branch names a branch to pin the "
    "grant there instead, which is what a deputy based at Lekki and made "
    "Branch Admin of Ikeja needs, and carries the word school to ask for the "
    "whole school deliberately. The grant is written once with the reach it is "
    "meant to have, rather than written school-wide and rewritten afterwards. "
    "A school with one branch is unaffected: there is one branch for a grant "
    "to follow, and the field is not rendered."
)

M12_FR001_REFUSAL_TAIL = (
    " A role_branch that is unknown, malformed or another school's collapses "
    "into the one answer every branch reference collapses into, so the field "
    "cannot be used to probe for sites."
)

M12_FR010_RULE_TAIL = (
    " (8) That the grant written beside a new account follows the posting "
    "(FR-001) does not make a later posting change move a grant. The reach is "
    "chosen once, when the grant is written; moving somebody's base afterwards "
    "leaves every grant where it is, and a school that wants their reach to "
    "move changes the grant."
)

M12_FR016_RULE_OLD = (
    "(3) The branch column is resolved through resolve_branch_reference; blank "
    "means across the whole school and is a valid answer rather than a missing "
    "one."
)

M12_FR016_RULE_NEW = (
    "(3) The branch column is resolved through resolve_branch_reference and is "
    "the posting; blank means across the whole school and is a valid answer "
    "rather than a missing one. The file carries no second column for the "
    "reach and needs none: the grant follows the posting at the choke point "
    "every caller shares, so a row naming Ikeja is granted at Ikeja and a row "
    "leaving the column blank is granted across the school "
    "(AnImportedGrantFollowsItsRowTests)."
)

M12_ACCEPTANCE = [
    "A teacher added at Ikeja with role_branch left out is granted at Ikeja "
    "and not across the school; one whose role_branch is the word school is "
    "granted across the school; one whose role_branch names Lekki is granted "
    "at Lekki whatever their posting; and somebody posted school-wide is "
    "granted across the school. Asserted at a school with two branches and at "
    "a school with one (AGrantFollowsThePostingTests).",
    "An imported row naming a branch is granted at that branch and a row "
    "leaving the branch column blank is granted across the school, through the "
    "same service a single add uses and with no import-only rule of its own "
    "(AnImportedGrantFollowsItsRowTests).",
]

M12_CREATE_ROUTE = (
    "FR-001. Moved from school.administrators.create when this module was "
    "built, and writes the staff record in the same transaction. role_branch "
    "is optional and decides how far the role grant reaches: left out it "
    "follows the posting, the word school asks for the whole school, and a "
    "branch reference pins the grant to that branch."
)

M12_BRANCH_REFUSAL_TAIL = (
    " The same answer covers a posting, a role_branch and the directory's "
    "branch filter alike."
)

M12_REACH_REFUSAL = [
    "A staff create that names no role_branch",
    "201",
    "Not a refusal. The grant is written with the reach the posting implies: "
    "the role template's own branch where it has one, otherwise the branch the "
    "person is posted to, and the whole school where that posting is "
    "school-wide.",
]

M12_IMPORT_DEPENDENCY_OLD = (
    "An imported person's role is granted school-wide unless the role template "
    "is tied to a branch: the file carries a branch for the posting and "
    "nothing for the reach."
)

M12_IMPORT_DEPENDENCY_NEW = (
    "An imported person's role reaches the branch the row posts them to, "
    "because the reach follows the posting at the choke point every caller "
    "shares; the file needs no second column for it."
)

M12_TRACE_STATE = "State at 19a50b3"

M12_TRACE_BRANCH_TAIL = (
    " A role grant written beside a new account reaches the branch the person "
    "is posted to unless the caller pins it to another or asks for the whole "
    "school by name, and the staff import takes the same rule from the same "
    "choke point."
)

M12_CHANGE_SUMMARY = (
    "Records how far a role grant written beside a new member of staff "
    "reaches. The reach was read off the role template rather than off the "
    "posting, so any school-wide template granted across the school by "
    "omission: a teacher hired at Ikeja and given Teacher without anybody "
    "naming a branch could read Lekki's records, and the staff import named no "
    "branch at all, so a file of forty teachers uploaded by one branch office "
    "produced forty people posted there and granted everywhere. The reach is "
    "now decided in the creation service every caller shares. An unstated "
    "reach follows the role template's own branch where it has one and "
    "otherwise the posting the account is created with; role_branch pins the "
    "grant to a named branch; and the whole school has to be asked for, by "
    "sending role_branch as the word school. The Add staff endpoint passes a "
    "resolved reach, and its patch-up, which wrote the grant school-wide and "
    "rewrote the row afterwards, is gone; the import is corrected by the choke "
    "point alone and needs no column of its own. A school with one branch is "
    "unaffected. FR-001's postconditions, its seventh business rule and its "
    "refusals, FR-010's rules, FR-016's rules, the create route, the branch "
    "refusal and one new row beside it, section 12.3's acceptance, section "
    "13's import row and section 15's traceability record it, and the control "
    "page moves to 19a50b3. No model, route, permission key or audit action "
    "changed. Covered by AGrantFollowsThePostingTests and "
    "AnImportedGrantFollowsItsRowTests; schools.vs_staff ran 225 tests OK. "
    "Backend evidence only; nothing here is deployed."
)

M12_MINOR_NOTE = (
    "Why version 2.4 is minor. The functional baseline is intact and so are "
    "its contracts in every other place: the same six models, the same two "
    "status vocabularies, the same posting-and-reach separation, the same "
    "routes, keys, refusal codes and audit actions. What changes is one rule, "
    "how far the grant written beside a new account reaches, and the one word "
    "that asks for the whole school deliberately."
)


def patch_m12(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M12 Staff Management Functional Requirements Document v{M12_TARGET}"

    control = table_with_row(doc, "Supersedes")
    fr001 = banner_table(doc, "FR-001  ")
    fr010 = banner_table(doc, "FR-010  ")
    fr016 = banner_table(doc, "FR-016  ")
    routes = table_with_header(doc, "Method and path", "Permission", "Requirement")
    refusals = table_with_row(doc, "A branch that is unknown, malformed or another tenant's")
    dependencies = table_with_header(doc, "Dependency", "Needed by", "Current state")
    trace = table_with_header(doc, "MRD capability", "Requirements")
    change_log = table_with_header(doc, "Version", "Date", "Summary")

    # ── control page ──
    m12_set_cell_text(row_where(control, "Version", exact=True).cells[1], M12_TARGET)
    m12_set_cell_text(row_where(control, "Date", exact=True).cells[1], REVIEW_DATE)
    m12_set_cell_text(
        row_where(control, "Supersedes", exact=True).cells[1], f"v{M12_SOURCE}",
    )
    m12_set_cell_text(
        row_where(control, "Source MRD", exact=True).cells[1],
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 12",
    )
    verified = row_where(control, "Verified against", exact=True).cells[1]
    m12_replace_in_cell(verified, M12_BASELINE_OLD, M12_BASELINE_NEW)
    m12_replace_in_cell(verified, "What version 2.3 changes", "What version 2.4 changes")

    # ── requirements ──
    m12_replace_in_cell(
        row_where(fr001, "Postconditions", exact=True).cells[1],
        M12_FR001_POSTCONDITION_OLD, M12_FR001_POSTCONDITION_NEW,
    )
    m12_replace_in_cell(
        row_where(fr001, "Business rules", exact=True).cells[1],
        M12_FR001_RULE_OLD, M12_FR001_RULE_NEW,
    )
    m12_append_to_cell(
        row_where(fr001, "Refusals", exact=True).cells[1], M12_FR001_REFUSAL_TAIL,
    )
    m12_append_to_cell(
        row_where(fr010, "Business rules", exact=True).cells[1], M12_FR010_RULE_TAIL,
    )
    m12_replace_in_cell(
        row_where(fr016, "Business rules", exact=True).cells[1],
        M12_FR016_RULE_OLD, M12_FR016_RULE_NEW,
    )

    # ── routes and refusals ──
    m12_set_cell_text(
        row_where(routes, "POST /v1/i/me/staff/", exact=True).cells[2], M12_CREATE_ROUTE,
    )
    branch_refusal = row_where(
        refusals, "A branch that is unknown, malformed or another tenant's", exact=True,
    )
    m12_append_to_cell(branch_refusal.cells[2], M12_BRANCH_REFUSAL_TAIL)
    m12_insert_row_after(refusals, branch_refusal, M12_REACH_REFUSAL)

    # ── acceptance criteria, section 12.3 ──
    anchor = find_paragraph(doc, "Changing a posting leaves every role grant untouched")
    for text in M12_ACCEPTANCE:
        anchor = insert_paragraph_after(anchor, text)

    # ── dependencies and traceability ──
    m12_replace_in_cell(
        row_where(dependencies, "The staff import dataset", exact=True).cells[2],
        M12_IMPORT_DEPENDENCY_OLD, M12_IMPORT_DEPENDENCY_NEW,
    )
    replace_in_paragraph(
        find_paragraph(doc, "MRD v2.76 records Module 12"), "MRD v2.76", "MRD v2.77",
    )
    m12_set_cell_text(unique_cells(trace.rows[0])[2], M12_TRACE_STATE)
    m12_append_to_cell(
        row_where(trace, "School and branch assignments", exact=True).cells[2],
        M12_TRACE_BRANCH_TAIL,
    )

    # ── change log, oldest first ──
    last = change_log.rows[len(change_log.rows) - 1]
    m12_assert_native(last.cells[0])
    m12_insert_row_after(change_log, last, [M12_TARGET, SHORT_DATE, M12_CHANGE_SUMMARY])
    minor_notes = [p for p in doc.paragraphs if "Why version 2.3 is minor" in p.text]
    if len(minor_notes) != 1:
        raise ValueError(f"Expected one minor-version note, found {len(minor_notes)}")
    set_callout_text(minor_notes[0], M12_MINOR_NOTE)

    finish(doc, output, title, M12_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 31 - Support Tickets
# ═════════════════════════════════════════════════════════════════════════════

M31_DIR = "31-support-tickets"
M31_STEM = "XVS_M31_Support_Tickets_Functional_Requirements_Document"
M31_SOURCE, M31_TARGET = "1.6", "1.7"

EVIDENCE_ROW, ACCEPTANCE_ROW, LIMIT_ROW = 2, 3, 4


def m31_cell_size(cell, default: float) -> float:
    """The point size a cell already renders at, so a rewrite keeps it."""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                return run.font.size.pt
    return default


def m31_replace_cell(cell, text: str, *, size: float | None = None, **kwargs) -> None:
    size = m31_cell_size(cell, 8) if size is None else size
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, size=size, **kwargs)


def m31_append_cell_text(cell, tail: str) -> None:
    m31_replace_cell(cell, cell.text.strip() + tail)


def m31_replace_cover_version(table, source: str, target: str) -> None:
    """Rewrite the version on the cover, which is one run of one title block."""
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def m31_replace_control_value(table, label: str, value: str) -> None:
    m31_replace_cell(row_where(table, label, exact=True).cells[1], value, size=9)


def m31_prepend_change_log(table, version: str, summary: str) -> None:
    template = table.rows[1]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[1]
    m31_replace_cell(row.cells[0], version, size=8)
    m31_replace_cell(row.cells[1], SHORT_DATE, size=8)
    m31_replace_cell(row.cells[2], summary, size=8)


def m31_drop_cell_paragraph(cell, fragment: str) -> None:
    """Remove the one paragraph of ``cell`` carrying ``fragment``."""
    matches = [p for p in cell.paragraphs if fragment in p.text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph {fragment!r}, found {len(matches)}")
    matches[0]._p.getparent().remove(matches[0]._p)


M31_FR005_EVIDENCE_TAIL = (
    " The query that resolves them is named and shared rather than written "
    "here: users_holding_ticket_keys_qs answers both who a ticket may be given "
    "to and who its notifications reach."
)

M31_FR005_ACCEPTANCE_TAIL = (
    " The ticket's notifications narrow that same query, so everybody the "
    "picker offers is somebody the desk's notices reach "
    "(TicketRecipientsHoldingAKeyThroughAGroupTests)."
)

M31_FR010_EVIDENCE_TAIL = (
    " Who works a queue is resolved through the query the assignee picker "
    "runs, so a triage key held through a permission group counts exactly as "
    "one written on the role, an explicit direct deny beats either and is "
    "judged key by key, and branch reach is matched the way the permission "
    "gate matches it."
)

M31_FR010_ACCEPTANCE_TAIL = (
    " The people a ticket can be given to and the people it can reach are one "
    "query rather than two that have to agree by inspection. "
    "TicketRecipientsHoldingAKeyThroughAGroupTests proves an agent whose "
    "manage key arrives through a permission group is told of a new ticket and "
    "of an escalation, that a key written on the role still reaches the desk, "
    "that a platform account holding no ticket key is told nothing, that "
    "everybody the picker offers can be reached, and that a school's own "
    "triage queue answers the same way."
)

M31_FR010_LIMIT = (
    "In-app and email only; there is no SMS channel. The notice a school reads "
    "covers the messages sent to its own people, so a school cannot audit "
    "delivery of the messages its ticket raised to platform staff. Recipients "
    "are resolved at the moment of the event, so somebody granted a triage key "
    "afterwards is not told about a ticket raised before they held it."
)

M31_RBAC_DEPENDENCY = (
    "Supplies the manage, assign, comment, internal-note, attach and audit "
    "keys, and the effective-grant evaluation behind the single query that "
    "answers both who may be assigned a ticket and who its notifications "
    "reach: a key held through a permission group counts as one written on the "
    "role, and an explicit direct deny beats either."
)

M31_GAP_REMOVED = "The desk's notification list and its assignee list can disagree"

M31_TRACEABILITY_LEAD = (
    f"Module 31 carries 17 capability entries in MRD v{MRD_VERSION}. Each maps "
    "to the requirements below."
)

M31_CHANGE_SUMMARY = (
    "Closes the gap v1.6 recorded between who may be given a ticket and who is "
    "told about one. The recipient query read only the keys written directly "
    "on a role while the assignee picker also counted a key held through a "
    "permission group, so a desk agent whose tickets.ticket.manage arrived "
    "through a group stood in the picker, could be handed a ticket, and was "
    "never told of a new or escalated one. Both questions now narrow one "
    "shared query and cannot disagree: a key held through a permission group "
    "counts exactly as one written on the role, an explicit direct deny beats "
    "either and is judged key by key, and branch reach is matched the way the "
    "permission gate matches it. The school-side triage queue had the same "
    "fault and is corrected by the same change, so a school's own triage staff "
    "holding the key through a group are told of a ticket raised at their "
    "school. FR-005, FR-010 and the Module 4 dependency record it, and the gap "
    "leaves Section 9. The other gap in that box stays as it was: the desk's "
    "ticket export carries CodeX's own tickets only, because every export run "
    "is scoped to the tenant it was requested in. No route, permission key, "
    "refusal code, audit action, model or traceability mapping changed, and "
    "Module 31 remains Backend Complete and In use Complete with seventeen "
    "capabilities. Covered by TicketRecipientsHoldingAKeyThroughAGroupTests; "
    "vs_tickets ran 113 tests OK. Backend evidence only; nothing here is "
    "deployed."
)


def patch_m31(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M31 Support Tickets Functional Requirements Document v{M31_TARGET}"

    cover = doc.tables[0]
    control = table_with_row(doc, "Code baseline")
    fr005 = banner_table(doc, "FR-005")
    fr010 = banner_table(doc, "FR-010")
    dependencies = table_with_row(doc, "Module 4, Roles & Permissions")
    gaps = banner_table(doc, "FURTHER GAPS")
    change_log = table_with_header(doc, "Version", "Date", "Summary")

    m31_replace_cover_version(cover, M31_SOURCE, M31_TARGET)
    m31_replace_control_value(control, "Version", M31_TARGET)
    m31_replace_control_value(control, "Review date", REVIEW_DATE)
    m31_replace_control_value(control, "Code baseline", CODE_BASELINE)
    m31_replace_control_value(
        control, "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 31",
    )

    m31_append_cell_text(fr005.rows[EVIDENCE_ROW].cells[1], M31_FR005_EVIDENCE_TAIL)
    m31_append_cell_text(fr005.rows[ACCEPTANCE_ROW].cells[1], M31_FR005_ACCEPTANCE_TAIL)

    m31_append_cell_text(fr010.rows[EVIDENCE_ROW].cells[1], M31_FR010_EVIDENCE_TAIL)
    m31_append_cell_text(fr010.rows[ACCEPTANCE_ROW].cells[1], M31_FR010_ACCEPTANCE_TAIL)
    m31_replace_cell(fr010.rows[LIMIT_ROW].cells[1], M31_FR010_LIMIT)

    m31_replace_cell(
        row_where(dependencies, "Module 4, Roles & Permissions", exact=True).cells[1],
        M31_RBAC_DEPENDENCY,
    )

    # Section 9 is current state, so a resolved gap leaves it rather than
    # being retained as history; the change log carries what was fixed.
    m31_drop_cell_paragraph(gaps.rows[0].cells[0], M31_GAP_REMOVED)

    set_run_text(find_paragraph(doc, "Module 31 carries"), M31_TRACEABILITY_LEAD)
    m31_prepend_change_log(change_log, M31_TARGET, M31_CHANGE_SUMMARY)

    finish(doc, output, title, M31_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for folder, stem, source, target, patch in (
        (M12_DIR, M12_STEM, M12_SOURCE, M12_TARGET, patch_m12),
        (M31_DIR, M31_STEM, M31_SOURCE, M31_TARGET, patch_m31),
    ):
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists():
            raise SystemExit(f"{output.name} already exists; refusing to overwrite it")
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
