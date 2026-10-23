#!/usr/bin/env python3
"""Version the MRD and affected FRDs for the 22 October 2026 contract batch."""

from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

import patch_record_history_docs as common
from patch_record_history_docs import (
    add_fr,
    append_rows,
    finish,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
)
from patch_staff_id_and_auth_events_docs import normalise_change_log, repair_ooxml
from generate_requirements_documents import assert_no_em_dash, shrink_inherited_media, update_extended_title


ROOT = Path(__file__).resolve().parents[1]
REVIEW_DATE = "23 October 2026"
CODE_BASELINE = (
    "Backend main at 4030229d (23 October 2026), covering the committed behaviour changes "
    "through 22 October 2026. No frontend deployment or production adoption is inferred."
)
MRD_SOURCE = "2.110"
MRD_TARGET = "2.111"
common.REVIEW_DATE = REVIEW_DATE


DOCS = {
    "M01": ("01-school-and-branch-management", "XVS_M01_School_and_Branch_Management_Functional_Requirements_Document", "1.34", "1.35"),
    "M03": ("03-identity-team-and-organogram", "XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document", "1.24", "1.24.1"),
    "M04": ("04-roles-and-permissions-rbac", "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document", "1.37", "1.38"),
    "M06": ("06-configuration-and-capability", "XVS_M06_Configuration_and_Capability_Management_Functional_Requirements_Document", "1.10", "1.11"),
    "M07": ("07-workflow-and-approval-engine", "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document", "1.26", "1.27"),
    "M10": ("10-bulk-data-import", "XVS_M10_Bulk_Data_Import_Functional_Requirements_Document", "1.11", "1.12"),
    "M12": ("12-staff-management", "XVS_M12_Staff_Management_Functional_Requirements_Document", "2.20", "2.21"),
    "M19": ("19-finance-and-accounting", "XVS_M19_Finance_and_Accounting_Functional_Requirements_Document", "1.19", "1.20"),
    "M22": ("22-procurement-and-requisitions", "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document", "1.20", "1.21"),
    "M32": ("32-messaging-and-conversations", "XVS_M32_Messaging_and_Conversations_Functional_Requirements_Document", "1.1", "1.2"),
}


REQUIREMENTS = {
    "M01": (
        "FR-024  Give Every Member of Staff the Branch Reference List",
        "FR-024  Give Every Member of Staff the Branch Reference List",
        [
            ("Description", "Serve the school's branch reference list to every current member of staff, while keeping branch administration behind its own permission."),
            ("Permission", "GET /v1/i/me/branches/ and its detail use IsTenantStaff. school.branches.view still opens the Branches administration screen and is not required for this reference list."),
            ("Scoping", "The caller must be active staff of the asserted tenant. Another school's staff, a guardian, a pupil, and a person whose last staff grant was revoked receive no list. Platform staff pass only at the platform tenant."),
            ("Business rules", "Staff means an active role or an in-force personal ALLOW grant at the tenant, evaluated by vs_rbac.evaluator.is_tenant_staff. The rule is shared with tenant branch options, configuration capability reads, workflow, and procurement approvals. A school always has at least one branch; a null branch elsewhere still means school-wide."),
            ("Acceptance", "A teacher without school.branches.view reads the branches needed by a picker. A parent using the same address is refused. Removing the person's final active staff grant closes the list on the next request. Tests: schools.vs_schools.tests_my_branches and vs_tenants.tests_branch_options."),
            ("Current limit", "None. This is a reference-list contract, not authority to create, edit, transition, or delete a branch."),
        ],
    ),
    "M03": (
        "FR-028  Name People by Their Names",
        "FR-028  Name People by Their Names",
        [
            ("Description", "Return a person's first and last names from the shared user model wherever a module asks for a human name."),
            ("Permission", "No new key and no new Field Access field. Each consuming surface keeps its own authorization and field rules."),
            ("Scoping", "The identity model supplies only the label for a person already within the caller's authorized tenant and resource scope."),
            ("Business rules", "get_full_name() joins the stored first and last names; get_short_name() uses the same human-name basis. An account with no name falls back to its email so an operational row is never blank. Finance audit keeps actor email beside actor_name so two people with one name remain distinguishable."),
            ("Acceptance", "A petty-cash keeper, claimant, workflow actor, and finance-audit option show Ada Okafor when those names are stored. An unnamed service account shows its email. Tests: vs_finance.tests_people_named."),
            ("Current limit", "The email fallback is deliberate for an account with no name; it is not used when either stored name produces a label."),
        ],
    ),
    "M04": (
        "FR-034  Apply Staff Membership and New Field Access Contracts",
        "FR-034  Apply Staff Membership and New Field Access Contracts",
        [
            ("Description", "Use one tenant-staff rule for staff-only reference and approval surfaces, and register the claimant and message-receipt fields that the new finance and messaging payloads expose."),
            ("Permission", "IsTenantStaff accepts an active tenant role or an in-force personal ALLOW grant. The finance.expenseclaim.claimant field has Read and Write switches over claimant_id, claimant_name, and claimant. school.messages.read_receipts is read-only over receipt and receipts."),
            ("Scoping", "Staff membership is evaluated inside the asserted tenant. Parent and pupil accounts have neither a role nor an ALLOW grant and are refused. Revoked roles, expired grants, and another tenant's staff do not qualify. Platform staff qualify at the platform tenant."),
            ("Business rules", "Staff-only gates protect branch and capability reference lists, workflow delegations, pending and instance decisions, procurement approvals, and every write that names an approver. Unknown, foreign, and non-staff ids receive the same Choose a member of staff response. Expense claimant values are open by default but not open on create. Read receipts are open by default and never writable. The school.messages resource now has nine fields."),
            ("Acceptance", "A guardian cannot reach an approval queue by knowing its URL. A revoked staff account cannot be seated in a stage. A hidden claimant disappears from lists, exports, approval cards, and dashboard summaries, and a forbidden write is refused. A hidden receipt removes receipt and receipts without hiding the message. Tests: the staff-only suites named in the queue, vs_finance.tests_field_access_deep_payload, schools.vs_messages.tests.test_receipts, and vs_rbac.tests.test_field_registry."),
            ("Current limit", "The names-only staff picker returns an id with its transport label and does not expose a staff profile field, so school.teachers switches do not apply to that narrow reference list."),
        ],
    ),
    "M06": (
        "FR-023  Restrict Capability Reads to Staff and Configure Read Receipts",
        "FR-023  Restrict Capability Reads to Staff and Configure Read Receipts",
        [
            ("Description", "Keep the caller's capability reference list staff-only and give each school one switch for message read receipts."),
            ("Permission", "GET /v1/config/my-capabilities/ requires IsTenantStaff. messages.read_receipts is a whole-school BOOLEAN setting, default true, read through the existing configuration resolver."),
            ("Scoping", "The capability list resolves only for staff of the asserted tenant. The receipt setting is school-wide and has no branch override."),
            ("Business rules", "Turning messages.read_receipts off hides receipts from every participant without changing stored read watermarks. Turning it back on makes the current state visible again. Inspectors and guardians watching a pupil conversation never contribute a read. The setting does not change unread counts, alerts, exports, or audit."),
            ("Acceptance", "A staff member can populate capability-aware screens; a parent and pupil receive no capability catalogue. A school can disable every receipt at once, and no branch can diverge. Tests: vs_config.tests_my_capabilities, schools.vs_messages.tests.test_receipts, and seed_config_catalogue coverage."),
            ("Current limit", "The setting is all or nothing for the school; there is no per-conversation, per-person, or branch choice."),
        ],
    ),
    "M07": (
        "FR-035  Keep Approval Seats Staff Only and Name the Trail",
        "FR-035  Keep Approval Seats Staff Only and Name the Trail",
        [
            ("Description", "Limit approval participation and administration to current tenant staff, and give every person on the readable trail a names-only label."),
            ("Permission", "Existing workflow keys still decide which staff member may configure, read, or decide. IsTenantStaff is an additional membership boundary on delegations, queues, instance reads, and votes."),
            ("Scoping", "Every named approver, replacement, Dynamic Role result, delegation delegate, and frozen seat must belong to the instance tenant's staff. The resolver rechecks staff when a stage activates; a stage with nobody parks or skips only where its own policy permits."),
            ("Business rules", "Writes naming an unknown, foreign, guardian, pupil, or former staff account answer 400 Choose a member of staff. Seats already frozen on a waiting stage are not re-resolved. Payloads carry requested_by_name; actor_name, on_behalf_of_name, and reversed_by_name on votes; user_name and on_behalf_of_name on eligible approvers; actor_name on audit rows; and on_behalf_of_name on the pending queue. acted_label uses the same names-only label. An account with no name reads Unknown user, never an email."),
            ("Acceptance", "A parent cannot receive a delegation or vote. A row left by bad historical configuration cannot seat a non-staff account when the stage activates. A reader who lacks the staff-list permission still reads the people named on an approval they may open. Tests: schools.vs_staff.tests.test_approval_seats, vs_workflow.tests.test_reassignment, and vs_workflow.tests.test_trail_names."),
            ("Current limit", "A waiting stage's already frozen seats remain fixed; later staff membership changes do not rewrite its recorded decision boundary."),
        ],
    ),
    "M10": (
        "FR-021  Import a School Owned Bank Statement from Its Account",
        "FR-021  Import a School Owned Bank Statement from Its Account",
        [
            ("Description", "Treat bank_statements as a tenant dataset entered only through the selected bank account's import wizard."),
            ("Permission", "The account wizard uses finance.bankaccount.import, the same key as manual statement entry. The generic batch upload refuses bank_statements for every caller."),
            ("Scoping", "The account is resolved inside the caller's tenant, finance entity, and branch reach before a batch exists. Publishing may write only that account's statement and lines."),
            ("Business rules", "The school template picker does not offer a bank-statement dataset. The account route supplies the module entry and account context that a generic batch cannot safely infer. bank_statements is tenant-owned rather than platform-only, but it remains unavailable at the generic upload endpoint."),
            ("Acceptance", "An Ikeja bursar imports an Ikeja account's file and cannot name Lekki's account. Posting the same dataset to the generic batch route is refused for CodeX and school callers alike. Tests: vs_finance.test_statement_import_wizard and vs_import_data.tests_dataset_ownership."),
            ("Current limit", "Only the owning account's wizard starts this dataset; there is intentionally no school-wide bank-statement upload."),
        ],
    ),
    "M12": (
        "FR-027  Supply Narrow Staff References and Payroll Documents",
        "FR-027  Supply Narrow Staff References and Payroll Documents",
        [
            ("Description", "Provide one names-only staff reference list for staff pickers and present payslips and annual tax summaries in the current payroll document design."),
            ("Permission", "GET /v1/i/me/staff/names/?q= is allowed when the caller holds any supported staff-picker key for expense claims, salary rows, workflow templates, approver groups, or approver assignment. Payroll document routes retain their existing pay-figure checks."),
            ("Scoping", "The picker returns the tenant's active staff across branches, capped at 500, searched on name parts only, as id and name. Employed people with suspended accounts are excluded. A payslip shows its branch line only when the school has several branches."),
            ("Business rules", "The reference list exposes no email, phone, posting, role, or staff-profile field. Payslips render from the shared HTML design with WeasyPrint and fall back to ReportLab where Pango is unavailable. They show net pay, gross allocation, earnings, deductions, employer contributions, year to date, and prior-employer figures. Later pages name the person, document number, and page count. The tax summary puts annual PAYE in its summary band and a full year's months in one page. Payslip JSON adds show_branch, employer_total, and gross_shares from already authorized figures."),
            ("Acceptance", "A finance officer finds Amaka by name without receiving her email. A suspended former worker is absent. A one-branch school's payslip omits a repeated branch line; a multi-branch school's includes it. A twelve-month tax summary fits one page. Tests: schools.vs_staff.tests.test_staff_names and vs_finance.tests_payslip_document."),
            ("Current limit", "The picker is intentionally not narrowed by branch because the named workflows can legitimately choose staff across the tenant; the write that consumes the id applies its own scope."),
        ],
    ),
    "M19": (
        "FR-026  Reconcile Payroll, Budgets, Bank Imports, Claims, and Human Labels",
        "FR-026  Reconcile Payroll, Budgets, Bank Imports, Claims, and Human Labels",
        [
            ("Description", "Keep the current finance contracts for payroll documents, directional budget totals, bank-statement entry, expense claimants, and human-readable person labels in one reconciled requirement."),
            ("Permission", "Existing payslip, report, bank-account import, expense-claim, audit, and dashboard keys remain. finance.expenseclaim.claimant adds Read and Write Field Access over claimant_id, claimant_name, and claimant; it is open by default and not open on create."),
            ("Scoping", "A bank import starts from an account resolved inside tenant, entity, and branch reach. A claimant id resolves through the payroll person resolver inside the entity's tenant. Budget headlines count the direction the plan budgets, while rows of the other direction remain visible without changing the headline."),
            ("Business rules", "The redesigned payslip and tax summary use WeasyPrint with the ReportLab fallback described by Module 12. A spending plan totals spending; an income-only plan totals income. Expense claims accept claimant as an account id and copy claimant_name; when no person is supplied a typed claimant_name is required. Unknown, malformed, foreign, and non-staff ids share one 400 refusal. Migration 0076 links no old claim, so its typed name remains. Claimant Field Access reaches searches, both exports, approval cards, and the spending dashboard; older frozen approval cards keep their saved name. Human labels use first and last names and fall back to email only for an unnamed account. Finance audit carries actor_name beside actor email and orders filter options by name."),
            ("Acceptance", "An unbudgeted fee-income row remains in a spending report but does not inflate actual spending or the consumed share. A claim can select a current employee and rejects another school's account without revealing it. Hiding claimant removes it from every alternate read and blocks writes. A finance audit distinguishes two Ada Okafors by retaining their emails. Tests: vs_finance.tests_payslip_document, vs_finance.tests, vs_finance.test_statement_import_wizard, vs_finance.tests_expense_claim_claimant, vs_finance.tests_field_access_deep_payload, and vs_finance.tests_people_named."),
            ("Current limit", "Older expense claims remain typed names because the migration deliberately guesses no identity. ReportLab remains a fallback where the HTML renderer's native libraries are unavailable."),
        ],
    ),
    "M22": (
        "FR-019  Keep Procurement Approval Surfaces Staff Only",
        "FR-019  Keep Procurement Approval Surfaces Staff Only",
        [
            ("Description", "Require current tenant staff membership on the procurement approval queue, its detail, and every decision action."),
            ("Permission", "The existing procurement approval and workflow keys still decide what a staff member may see or do. IsTenantStaff is an additional boundary and grants no procurement authority by itself."),
            ("Scoping", "The caller must be staff of the asserted tenant, and the document must remain inside their tenant, entity, and branch reach. Guardians, pupils, former staff, and another school's staff receive no queue access."),
            ("Business rules", "The procurement queue follows the same staff definition as workflow: an active role or in-force personal ALLOW grant. A decision still passes through the workflow instance, frozen seat, branch reach, and conflict checks. Removing staff membership closes the route even if an old token remains valid."),
            ("Acceptance", "A parent who knows an approval id cannot open or decide it. A branch approver sees only documents in reach. Existing staff with the right decision key continue unchanged. Tests: vs_procurement.tests, vs_procurement.tests_approval_resubmit, and vs_workflow.tests.test_tenant_scoping."),
            ("Current limit", "This rule does not change the deliberate override and submitter-release risks already recorded in Needs Attention."),
        ],
    ),
    "M32": (
        "FR-027  Show Optional Read Receipts and Poll the Fixed Thread",
        "FR-027  Show Optional Read Receipts and Poll the Fixed Thread",
        [
            ("Description", "Show one tick when the sender's message is stored and two when every current recipient has read it, while keeping an open conversation current without moving its header or composer."),
            ("Permission", "The caller must already read the conversation. messages.read_receipts must be true, and school.messages.read_receipts must be readable. Receipt values are server-owned and never writable."),
            ("Scoping", "Only the caller's own messages carry receipt. Direct, family, and pupil conversations count their current other participant. Groups and broadcasts count only recipients who could see the message: a later joiner, removed member, inactive account, former staff member, or pupil off the roll is excluded. An inspector and a guardian watching a pupil conversation never count."),
            ("Business rules", "receipt is {state, read_count, recipient_count}; detail carries receipts true or false. GET messages/?after=<id> on staff, parent, and pupil routes returns newer messages plus receipts for the caller's recent messages and consumes no send limit. Receipt counts come from read watermarks in one query per page, with no migration. Receipts never enter exports, alerts, audit, or another person's message row. The apps keep the person header and composer fixed, scroll only messages, hide the app header and tab bar on a phone, poll every ten seconds while open and visible, and mark read only when the reader has the message in view."),
            ("Acceptance", "A direct message moves from one tick to two after its recipient sees it. A group sender sees Read by n of m and receives two ticks only at n equals m. Turning the school setting off removes every receipt without changing unread state. A parent observing a pupil thread and a welfare inspector change no count. Tests: schools.vs_messages.tests.test_receipts and the three frontend conversation suites."),
            ("Current limit", "Polling is ten-second foreground refresh, not a live socket. Receipts are school-wide, with no per-conversation choice."),
        ],
    ),
}


SUMMARIES = {
    "M01": "Makes the branch reference list available to every current member of staff through the shared tenant-staff boundary, without granting branch administration. Aligns with MRD v2.111.",
    "M03": "Records the shared human-name contract: first and last names lead, and email is only the fallback for an unnamed account. Aligns with MRD v2.111.",
    "M04": "Records the shared tenant-staff boundary, the expense-claim claimant field, the read-only message receipt field, and nine message Field Access switches. Aligns with MRD v2.111.",
    "M06": "Makes my-capabilities staff-only and adds the school-wide messages.read_receipts switch, default on. Aligns with MRD v2.111.",
    "M07": "Limits approval participation and administration to current staff and adds names-only labels to every readable approval person. Aligns with MRD v2.111.",
    "M10": "Makes bank_statements a tenant dataset entered only from the selected bank account's import wizard. Aligns with MRD v2.111.",
    "M12": "Adds the names-only staff picker contract and reconciles the redesigned payslip and annual tax summary. Aligns with MRD v2.111.",
    "M19": "Reconciles payroll documents, directional budget totals, bank-statement entry, expense claimants, claimant Field Access, and human labels. Aligns with MRD v2.111.",
    "M22": "Adds the current-staff boundary to procurement approval lists, detail, and decisions. Aligns with MRD v2.111.",
    "M32": "Adds optional read receipts, after polling, receipt-count rules, and the fixed conversation layout across staff, parent, and pupil apps. Aligns with MRD v2.111.",
}


TRACE = {
    "M01": ("Create, list, retrieve, and update Branches", "FR-005, FR-010, FR-024"),
    "M03": ("Tenant-aware email and password login", "FR-001, FR-028"),
    "M04": ("Field Access", "FR-023, FR-034"),
    "M06": ("Typed configuration", "FR-001, FR-023"),
    "M07": ("Role-based approver resolution", "FR-005, FR-035"),
    "M10": ("CSV and XLSX upload", "FR-003, FR-021"),
    "M12": ("Staff profile linked to identity", "FR-001, FR-027"),
    "M19": ("Budget", "FR-013, FR-026"),
    "M22": ("Approval queue and decision actions", "FR-009, FR-019"),
    "M32": ("Inbox, search, unread counts and read state", "FR-014, FR-027"),
}


def replace_runs(paragraph: Paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    return True


def set_optional_control(doc, labels: tuple[str, ...], value: str) -> None:
    for table in doc.tables[:3]:
        for row in table.rows:
            if row.cells[0].text.strip() in labels:
                keep_format(row.cells[-1], value)
                return


def table_starting(doc, start: str) -> Table:
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip().startswith(start):
            return table
    raise ValueError(f"Table not found: {start}")


def cell_starting(doc, start: str, contains: str | None = None):
    """Return the first table cell matching its opening and optional contents."""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip().startswith(start) and (contains is None or contains in cell.text):
                    return cell
    raise ValueError(f"Cell not found: {start}")


def replace_table_text(doc, old: str, new: str) -> int:
    """Replace current-state table wording without rewriting historical rows."""
    change_log = common.change_log_table(doc)
    count = 0
    for table in doc.tables:
        if table._tbl is change_log._tbl:
            continue
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if replace_runs(paragraph, old, new):
                        count += 1
    return count


def reconcile_m32_current_state(doc) -> None:
    """Bring inherited Module 32 counts and evidence through the receipt change."""
    replacements = {
        "The school.messages resource, its eight fields and its exemptions":
            "The school.messages resource, its nine fields and its exemptions",
        "field_access.register publishes eight fields on the school.messages resource":
            "field_access.register publishes nine fields on the school.messages resource",
        "The three school settings, read with the student rules resolver.":
            "The four school settings, read with the student rules resolver.",
    }
    for old, new in replacements.items():
        if replace_table_text(doc, old, new) != 1:
            raise ValueError(f"Module 32 current-state text not found once: {old}")

    evidence = table_starting(doc, "INSPECTED EVIDENCE")
    keep_format(
        evidence.rows[0].cells[0],
        "INSPECTED EVIDENCE The module's receipt behavior is covered by "
        "schools.vs_messages.tests.test_receipts. The behavior commit records that "
        "focused module test as passing; this documentation-only revision did not rerun "
        "application tests. The earlier 21 October 2026 module run covered 386 tests and "
        "passed (OK). The seed and registry checks in core.test_seed_school_permissions, "
        "vs_rbac.tests.test_field_registry and vs_notifications.tests_in_app_subjects, and "
        "the full repository suite, were not run for this revision. Backend test evidence "
        "does not prove a live broker, worker, push service or inbox. Frontend main and "
        "staging evidence is recorded separately and does not prove production deployment.",
    )


def reconcile_mrd_current_state(doc) -> None:
    """Replace inherited planning language with the committed Module 32 state."""
    paragraph_replacements = {
        "Current uncommitted platform, security, integration, mail, and onboarding-default work is counted as complete.":
            "Current platform, security, integration, mail, and onboarding-default work is counted as complete.",
        "The FRDs for Modules 8, 27, 28, 4 and 5 are queued for revision.":
            "The supporting FRDs were reconciled in MRD v2.110.",
    }
    for old, new in paragraph_replacements.items():
        matches = sum(replace_runs(paragraph, old, new) for paragraph in doc.paragraphs)
        if matches != 1:
            raise ValueError(f"MRD current-state paragraph not found once: {old}")

    table_replacements = {
        "Inbox, search, unread counts and monotonic read state":
            "Inbox, search, unread counts, read state and optional read receipts",
        "Field Access over messages, names, recipients and times":
            "Field Access over messages, names, recipients, times and read receipts",
        "Ten permission keys, default roles and three school settings":
            "Ten permission keys, default roles and four school settings",
    }
    for old, new in table_replacements.items():
        if replace_table_text(doc, old, new) != 1:
            raise ValueError(f"MRD Module 32 capability text not found once: {old}")

    decision = cell_starting(doc, "CURRENT DECISION", "Module 32")
    keep_format(
        decision,
        "CURRENT DECISION\n"
        "• Module 32 is Backend Complete and In use Partial with fourteen capability entries, reconciled to M32 FRD v1.2.\n"
        "• Conversations remain separate from Module 8. Module 8 delivers four new-message events and their templates through the CONVERSATION record family; no message text appears in any channel or stored notification.\n"
        "• The owner decisions are confirmed. Staff direct conversations are visible to their two participants only; every pupil conversation is readable by the primary guardian and inspectable read-only and audited under the welfare key; groups and broadcasts are staff only; messages are text only.\n"
        "• The backend and the staff, parent, and student screens are committed on main and present on staging. Production deployment and adoption are not inferred.\n"
        "• In use is Partial because the pupil and parent home/ responses carry no unread-messages block; each portal serves its own unread-count route instead.\n"
        "• Optional read receipts are school-wide, default on, and read-only under Field Access. An open conversation polls every ten seconds while visible and keeps its person header and composer fixed around the scrolling messages.\n"
        "• The receipt behavior commit records schools.vs_messages.tests.test_receipts as passing. This documentation-only revision did not rerun application tests, and the full repository suite was not run.\n"
        "• Field Access: school.messages registers nine fields, open by default; portal payloads remain a fixed narrow shape that no role switch reaches. Audit has the MESSAGES module key and six action types; RBAC has ten school keys and the verbs inspect and office; configuration has four school settings.\n"
        "• Needs Attention remains: production deployment is not evidenced, portal home integration is incomplete, staff reach follows the home branch only, a leaver's family conversations are not handed over, and a lost alert is not redispatched.",
    )


def remove_inherited_branch_synonym(doc) -> None:
    """Keep the current document free of the retired branch synonym."""
    old = "Document-only patch. Retired the word campus, which had drifted in as a synonym for the site primitive this module defines."
    new = "Document-only patch. Retired an incorrect synonym for the site primitive this module defines."
    count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if replace_runs(paragraph, old, new):
                        count += 1
    if count != 1:
        raise ValueError(f"Inherited branch synonym summary not found once: {count}")


def update_traceability(doc, needle: str, requirements: str) -> None:
    matches = []
    for table in doc.tables:
        if not table.rows or "MRD" not in table.rows[0].cells[0].text:
            continue
        for row in table.rows[1:]:
            if needle.lower() in row.cells[0].text.lower():
                matches.append(row)
    if len(matches) != 1:
        raise ValueError(f"Traceability row {needle!r} matched {len(matches)} rows")
    keep_format(matches[0].cells[1], requirements)


def update_current_mrd_reference(doc) -> None:
    set_optional_control(doc, ("Source MRD", "MRD baseline"), MRD_TARGET)
    for table in doc.tables:
        if not table.rows:
            continue
        text = table.rows[0].cells[0].text.strip()
        if text.startswith("MRD RECONCILIATION"):
            cell = table.rows[0].cells[0]
            for p in cell.paragraphs:
                replace_runs(p, f"MRD v{MRD_SOURCE}", f"MRD v{MRD_TARGET}")
    for p in doc.paragraphs:
        if p.text.startswith("Module ") and f"MRD v{MRD_SOURCE}" in p.text:
            replace_runs(p, f"MRD v{MRD_SOURCE}", f"MRD v{MRD_TARGET}")


def restore_clean_change_log_rows(doc, clean_source: Path) -> None:
    """Rebuild a drifted change log from a known clean header and data row."""
    clean = Document(clean_source)
    target_table = common.change_log_table(doc)
    clean_table = common.change_log_table(clean)
    values = [[cell.text for cell in row.cells] for row in target_table.rows]
    for row in list(target_table.rows):
        target_table._tbl.remove(row._tr)
    target_table._tbl.append(copy.deepcopy(clean_table.rows[0]._tr))
    for row_values in values[1:]:
        target_table._tbl.append(copy.deepcopy(clean_table.rows[1]._tr))
        row = target_table.rows[-1]
        for cell, value in zip(row.cells, row_values):
            keep_format(cell, value)


def sort_and_deduplicate_change_log(doc) -> None:
    """Keep one row per semantic version and order newest first."""
    table = common.change_log_table(doc)
    rows = list(table.rows)
    header = rows[0]
    version_rows = {}
    other_rows = []
    for row in rows[1:]:
        value = row.cells[0].text.strip()
        if re.fullmatch(r"[0-9]+(?:[.][0-9]+)+", value):
            version_rows.setdefault(tuple(int(part) for part in value.split(".")), row)
        else:
            other_rows.append(row)
    for row in rows[1:]:
        table._tbl.remove(row._tr)
    for version in sorted(version_rows, reverse=True):
        table._tbl.append(copy.deepcopy(version_rows[version]._tr))
    for row in other_rows:
        table._tbl.append(copy.deepcopy(row._tr))


def patch_frd(key: str) -> Path:
    folder, stem, source, target = DOCS[key]
    source_path = ROOT / "functional-requirements" / folder / f"{stem}_v{source}.docx"
    output = ROOT / "functional-requirements" / folder / f"{stem}_v{target}.docx"
    if output.exists():
        raise SystemExit(f"{output.name} already exists")
    doc = Document(source_path)
    set_cover_version(doc, source, target)
    set_control(doc, "Version", target)
    set_optional_control(doc, ("Review date", "Date"), REVIEW_DATE)
    set_optional_control(doc, ("Code baseline", "Source scope", "Verified against"), CODE_BASELINE)
    update_current_mrd_reference(doc)
    heading, header, rows = REQUIREMENTS[key]
    add_fr(doc, heading, header, rows)
    update_traceability(doc, *TRACE[key])

    if key == "M32":
        fields = table_starting(doc, "Field")
        append_rows(fields, [["read_receipts", "Read receipts", "Read only", "Not writable", "receipt, receipts"]])
        settings = table_starting(doc, "Setting")
        append_rows(settings, [["messages.read_receipts", "On", "Show stored and read receipt state to authorized participants; school-wide, with no branch override"]])
        contracts = table_starting(doc, "Item")
        for row in contracts.rows:
            if row.cells[0].text.strip() == "Message":
                keep_format(row.cells[1], "id, body, sender (user_id, name), sent_at, mine and, only on the caller's own messages when enabled and readable, receipt {state, read_count, recipient_count}")
        for table in doc.tables:
            for row in table.rows:
                if row.cells and row.cells[0].text.strip().startswith("GET conversations/{id}/messages/"):
                    keep_format(row.cells[-1], row.cells[-1].text.rstrip() + "; after returns newer messages and recent receipts without using the send limit")
                    break
        reconcile_m32_current_state(doc)
    log_change(doc, target, SUMMARIES[key], header_source=common.oldest_with_header(folder, stem, source))
    if key == "M01":
        remove_inherited_branch_synonym(doc)
    repair_ooxml(doc)
    normalise_change_log(doc)
    if key == "M06":
        restore_clean_change_log_rows(
            doc,
            ROOT / "functional-requirements" / folder / f"{stem}_v1.0.docx",
        )
    sort_and_deduplicate_change_log(doc)
    title = f"{stem.replace('_', ' ')} v{target}"
    finish(doc, output, title, target)
    return output


def paragraph_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    clone = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(clone)
    created = Paragraph(clone, paragraph._parent)
    if style:
        created.style = style
    if created.runs:
        created.runs[0].text = text
        for run in created.runs[1:]:
            run.text = ""
    else:
        created.add_run(text)
    return created


def mrd_header_source() -> Document:
    folder = ROOT / "module-requirements"
    paths = sorted(folder.glob("XVS_Module_Requirements_Document_v*.docx"), key=lambda p: [int(x) for x in p.stem.rsplit("_v", 1)[1].split(".")], reverse=True)
    for path in paths:
        doc = Document(path)
        try:
            table = common.change_log_table(doc)
        except ValueError:
            continue
        if table.rows[0].cells[0].text.strip() == "Version":
            return doc
    raise ValueError("No MRD change log with a header")


def append_module_summary(doc, module: int, sentence: str, frd_version: str) -> None:
    heading = f"Module {module}:"
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(heading):
            desc = doc.paragraphs[i + 1]
            text = desc.text.rstrip()
            text = re.sub(rf"Documented by M{module:02d} FRD v[0-9.]+\.?", "", text).rstrip()
            if text and not text.endswith("."):
                text += "."
            replace_runs(desc, desc.text, f"{text} {sentence} Documented by M{module:02d} FRD v{frd_version}.")
            return
    raise ValueError(f"MRD module heading not found: {module}")


def patch_mrd() -> Path:
    source = ROOT / "module-requirements" / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"
    output = ROOT / "module-requirements" / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx"
    if output.exists():
        raise SystemExit(f"{output.name} already exists")
    doc = Document(source)
    set_control(doc, "Version", MRD_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Source scope", CODE_BASELINE)
    doc.core_properties.version = MRD_TARGET

    summaries = {
        1: ("Every current member of staff reads the branch reference list, while branch administration remains separately permissioned.", DOCS["M01"][3]),
        3: ("Shared person labels use first and last names, with email only for an unnamed account.", DOCS["M03"][3]),
        4: ("One tenant-staff rule protects staff-only references and approvals; Field Access adds expense claimant and read-only message receipts.", DOCS["M04"][3]),
        6: ("The caller capability list is staff-only, and messages.read_receipts is a school-wide default-on setting.", DOCS["M06"][3]),
        7: ("Approval seats, delegations, queues, and decisions are staff-only, while every readable trail person carries a names-only label.", DOCS["M07"][3]),
        10: ("A school imports a bank statement only from the selected bank account's wizard; the generic batch route refuses that dataset.", DOCS["M10"][3]),
        12: ("A narrow active-staff names list feeds shared pickers, and payroll documents use the current payslip and tax-summary design.", DOCS["M12"][3]),
        19: ("Finance reconciles payroll documents, directional budget totals, account-bound statement imports, staff-linked claimants, claimant Field Access, and human labels.", DOCS["M19"][3]),
        22: ("The procurement approval queue, detail, and decisions require current tenant staff membership in addition to their existing keys and reach.", DOCS["M22"][3]),
        32: ("Authorized participants see optional stored and read receipts, including group and broadcast counts, through foreground ten-second polling and a fixed thread layout.", DOCS["M32"][3]),
    }
    for module, (sentence, version) in summaries.items():
        append_module_summary(doc, module, sentence, version)

    reconcile_mrd_current_state(doc)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "Inbox, search, unread counts and read state" in cell.text:
                    for p in cell.paragraphs:
                        replace_runs(p, "Inbox, search, unread counts and read state", "Inbox, search, unread counts, read state and optional read receipts")

    change_heading = next(p for p in doc.paragraphs if p.text.strip() == "7. Change Log")
    template_heading = next(p for p in doc.paragraphs if p.text.strip().startswith("6.4 21 October 2026"))
    new_heading = copy.deepcopy(template_heading._p)
    change_heading._p.addprevious(new_heading)
    hp = Paragraph(new_heading, template_heading._parent)
    replace_runs(hp, hp.text, "6.5 22 October 2026 Capability Contracts")
    paragraphs = [
        "Messaging adds optional read receipts. A stored message shows one tick, a fully read message two, and a group or broadcast reports Read by n of m from current eligible recipients. Inspectors and a guardian watching a pupil conversation do not count. The three apps keep an open thread fixed around its scrolling messages and poll every ten seconds while visible.",
        "The shared tenant-staff boundary now protects branch and capability references, workflow administration and decisions, procurement approvals, and every write that names an approver. Active tenant roles and in-force personal ALLOW grants qualify; guardians, pupils, revoked holders, and another tenant's staff do not.",
        "Approval payloads name requesters, voters, delegates, eligible approvers, reversers, and audit actors without exposing their profiles. The names-only staff reference list supplies shared claimant, payroll, and approver pickers, while each consuming write keeps its own tenant and business scope.",
        "Finance adds the current payroll document design, one-direction budget headlines, tenant-owned bank-statement imports entered from their account, staff-linked expense claimants with end-to-end Field Access, and human-name labels with an email fallback only for an unnamed account.",
        "No module status, ownership, priority gap, recommended build order, or capability count changes. The platform remains at 32 modules and 527 capability entries. The ten affected FRDs advance independently and reconcile to this MRD.",
    ]
    template_contract_paragraph = template_heading._p.getnext()
    if template_contract_paragraph is None or not template_contract_paragraph.tag.endswith("}p"):
        raise ValueError("The 6.4 heading has no body paragraph to clone")
    for text in paragraphs:
        p = copy.deepcopy(template_contract_paragraph)
        change_heading._p.addprevious(p)
        created = Paragraph(p, change_heading._parent)
        replace_runs(created, created.text, text)

    summary = (
        "Reconciles the 22 October contract batch across messaging receipts, staff-only reference and approval surfaces, names-only workflow trails and staff pickers, payroll documents, directional budget totals, account-bound bank-statement imports, expense claimants, claimant Field Access, and shared human labels. No module status, priority, ownership, build order, or capability count changes; the tracker remains at 32 modules and 527 entries. Reconciled to M01 v1.35, M03 v1.24.1, M04 v1.38, M06 v1.11, M07 v1.27, M10 v1.12, M12 v2.21, M19 v1.20, M22 v1.21, and M32 v1.2."
    )
    log_change(doc, MRD_TARGET, summary, header_source=mrd_header_source())
    repair_ooxml(doc)
    normalise_change_log(doc)
    sort_and_deduplicate_change_log(doc)
    doc.core_properties.title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.save(output)
    update_extended_title(output, doc.core_properties.title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    print(f"Wrote {output.name}")
    return output


def main() -> None:
    outputs = [patch_frd(key) for key in DOCS]
    outputs.append(patch_mrd())
    print(f"Created {len(outputs)} documents")


if __name__ == "__main__":
    main()
