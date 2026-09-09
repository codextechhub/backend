#!/usr/bin/env python3
"""Create the Module 26 and Module 25 FRDs.

Two modules documented in one change because they are two halves of the same
question. Module 26 is an engine with an app of its own; Module 25 is not an
app at all, and its FRD exists largely to say so and to name where each
dashboard actually lives. Written from the code rather than from the MRD's
summary of it.

    python tools/create_analytics_frds.py
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_requirements_documents import (
    GREY,
    add_body,
    add_callout as add_reference_callout,
    add_cover,
    add_heading,
    add_metadata_table,
    add_page_break,
    add_requirement,
    add_status_key,
    add_table,
    assert_no_em_dash,
    remove_body_content,
    set_headers,
    shrink_inherited_media,
    update_extended_title,
    write_paragraph,
)

FRD_VERSION = "1.0"
REVIEW_DATE = "9 September 2026"
MRD_VERSION = "2.72"

M26_BASELINE = (
    "apps/vs_exports on main with data export reachable at Core on every "
    "plan, reviewed on 9 September 2026"
)
M25_BASELINE = (
    "The dashboard and summary endpoints across vs_admin_console, vs_audit, "
    "vs_finance, vs_procurement, vs_payments, vs_workflow, vs_todo, "
    "vs_tickets, vs_health and vs_exports on main, reviewed on "
    "9 September 2026"
)


# ── Module 26 ────────────────────────────────────────────────────────────────

M26_REQUIREMENTS = [
    {
        "id": "FR-001",
        "title": "Publish What May Be Exported, and What Each Dataset Can Do",
        "status": "Implemented",
        "requirement": (
            "A reader must be able to see which datasets they may export and, for "
            "one dataset, which columns, filters, formats and value modes it "
            "supports, without guessing from a failed run."
        ),
        "evidence": (
            "The catalogue endpoints list datasets and describe one, and the "
            "capabilities endpoint reports what the caller may do. The list is "
            "narrowed to the caller's own tenant and reach, so a dataset nobody may "
            "export is never offered."
        ),
        "acceptance": (
            "A dataset absent from the catalogue cannot be run. A described "
            "dataset's columns and formats are the ones a run will accept."
        ),
        "limit": (
            "The catalogue describes datasets the engine ships with. A domain "
            "adding one registers it in code rather than through the API."
        ),
    },
    {
        "id": "FR-002",
        "title": "Preview Before Committing to a Run",
        "status": "Implemented",
        "requirement": (
            "A reader must be able to see the shape and a sample of what an export "
            "would produce before producing it."
        ),
        "evidence": (
            "The preview endpoint applies the same dataset, columns and filters a "
            "run would, and returns a bounded sample rather than the whole result."
        ),
        "acceptance": (
            "A preview and the run built from it describe the same columns. A "
            "preview never writes a run, a file or an analytics event of the kind a "
            "run writes."
        ),
        "limit": (
            "A preview is bounded, so it cannot prove a large export will complete."
        ),
    },
    {
        "id": "FR-003",
        "title": "Export the Screen in Front of You",
        "status": "Implemented",
        "requirement": (
            "Taking the filtered list a reader is already looking at must not "
            "require them to build a definition first."
        ),
        "evidence": (
            "The quick and from-screen endpoints accept a dataset and the filters "
            "in force and produce a run directly, through the same engine a saved "
            "definition uses."
        ),
        "acceptance": (
            "A quick export and a saved definition with the same filters produce "
            "the same rows. Neither path can reach a dataset the caller may not "
            "export."
        ),
        "limit": (
            "A quick export is not kept. Returning to it means saving a definition."
        ),
    },
    {
        "id": "FR-004",
        "title": "Save an Export and Return to It",
        "status": "Implemented",
        "requirement": (
            "A reader must be able to name and keep an export they run again, and "
            "to build a new one from an existing one."
        ),
        "evidence": (
            "ExportDefinition holds dataset, columns, filters, format and values "
            "mode. The duplicate endpoint copies one into a new definition owned by "
            "the caller rather than sharing the original."
        ),
        "acceptance": (
            "Editing a definition does not change a run already produced from it, "
            "because the run froze its own configuration (FR-006)."
        ),
        "limit": (
            "A definition belongs to one owner. There is no team-owned definition, "
            "only a shared one."
        ),
    },
    {
        "id": "FR-005",
        "title": "Share Sight of a Definition Without Sharing Data Access",
        "status": "Implemented",
        "requirement": (
            "Sharing an export must let a colleague see it and its files without "
            "silently lending them the sharer's reach over the data."
        ),
        "evidence": (
            "ExportDefinitionShare grants sight of the definition and its files. "
            "The run always executes as the owner, and every download is "
            "re-authorised against the person downloading rather than the person "
            "who shared."
        ),
        "acceptance": (
            "A recipient sees the definition and is refused a file whose contents "
            "they may not read. Revoking a share removes sight without touching the "
            "runs already produced."
        ),
        "limit": (
            "A share is per person. There is no share to a role or to a branch."
        ),
    },
    {
        "id": "FR-006",
        "title": "Make a Run Describe Itself, Not Its Definition",
        "status": "Implemented",
        "requirement": (
            "A run must record what it actually did, so that reading it months "
            "later describes the file it produced rather than whatever the "
            "definition has since become."
        ),
        "evidence": (
            "ExportRun carries frozen_config: dataset, entity, columns, filters, "
            "format and values mode exactly as they were when the run began, and a "
            "sequenced human reference."
        ),
        "acceptance": (
            "Editing the definition after a run leaves that run's detail unchanged. "
            "Two runs of an edited definition describe different configurations."
        ),
        "limit": (
            "The frozen configuration is a record, not a replay: re-running it uses "
            "the definition as it now stands."
        ),
    },
    {
        "id": "FR-007",
        "title": "Keep a Run's Status a Closed Set and Never Rewrite It",
        "status": "Implemented",
        "requirement": (
            "A run has one of a fixed set of statuses, terminal states are final, "
            "and nothing that happens to a file afterwards may change what the run "
            "says happened."
        ),
        "evidence": (
            "RunStatus is the closed set queued, running, completed, completed with "
            "omissions, failed and cancelled, and the terminal set is named "
            "explicitly. Expired is deliberately not a run status: expiry is a "
            "property of the file, derived at read time, so history is never "
            "overwritten to represent it."
        ),
        "acceptance": (
            "A file expiring leaves its run reading completed, because the run "
            "genuinely did succeed. A terminal run is not mutated again."
        ),
        "limit": (
            "A reader wanting 'succeeded but the file is gone' has to read the run "
            "and the file together. That is the intended shape."
        ),
    },
    {
        "id": "FR-008",
        "title": "Explain a Failure With a Code, Never a Traceback",
        "status": "Implemented",
        "requirement": (
            "A failed run must carry a machine-readable reason, a message safe to "
            "show a user, and a reference support can trace, and must never expose "
            "an internal error."
        ),
        "evidence": (
            "FailureCode is a closed vocabulary carried by every failure. The API "
            "returns the code, a user-safe message and a support reference, and the "
            "client maps the code to a recommended action."
        ),
        "acceptance": (
            "A failure renders as an action a reader can take. No traceback, query "
            "or internal path reaches the response."
        ),
        "limit": (
            "A code the client does not recognise renders as a generic failure "
            "rather than a specific action."
        ),
    },
    {
        "id": "FR-009",
        "title": "Explain a Partial Result Structurally",
        "status": "Implemented",
        "requirement": (
            "A run that produced a file while leaving something out must say what "
            "was left out, in a form a screen can render rather than infer."
        ),
        "evidence": (
            "Completed with omissions is its own status, and OmissionCode gives the "
            "reasons as a structured list rather than prose, because the client "
            "renders them."
        ),
        "acceptance": (
            "A reader sees which columns or rows were omitted and why, without "
            "reading a sentence and guessing."
        ),
        "limit": (
            "Omission reasons are a closed vocabulary. A new reason is a code "
            "change on both sides."
        ),
    },
    {
        "id": "FR-010",
        "title": "Cancel and Retry Without Losing the Original",
        "status": "Implemented",
        "requirement": (
            "A run still going must be stoppable, and a failed one repeatable, "
            "without either action erasing what happened the first time."
        ),
        "evidence": (
            "The cancel and retry endpoints act on a run; cancellation moves it to "
            "a terminal status and a retry produces a new run rather than reopening "
            "the old one."
        ),
        "acceptance": (
            "Cancelling a queued or running export produces a cancelled run that "
            "stays cancelled. A retry leaves the failed run readable beside its "
            "successor."
        ),
        "limit": (
            "Cancellation covers the statuses before execution finishes. A run "
            "already writing its file finishes it."
        ),
    },
    {
        "id": "FR-011",
        "title": "Finish Runs Whose Worker Died",
        "status": "Implemented",
        "requirement": (
            "A run whose worker vanished must not sit queued or running forever, "
            "looking to a reader exactly like one that is about to finish."
        ),
        "evidence": (
            "sweep_abandoned_runs runs every half hour and finishes runs that never "
            "reported back, moving them to a terminal status with a failure code."
        ),
        "acceptance": (
            "An abandoned run reaches a terminal status without human intervention "
            "and carries a reason a reader can act on."
        ),
        "limit": (
            "Detection is on the sweep's cadence, so an abandoned run can look live "
            "for up to half an hour."
        ),
    },
    {
        "id": "FR-012",
        "title": "Hold a File for a Stated Period, Then Destroy It",
        "status": "Implemented",
        "requirement": (
            "A produced file must be available for a stated period and then removed "
            "from storage, and its removal must not rewrite the run that made it."
        ),
        "evidence": (
            "ExportFile computes availability from available_until on read. A "
            "nightly job purges storage past the thirty-day window and stamps "
            "purged_at. Nothing about the run changes."
        ),
        "acceptance": (
            "A file past its window is reported unavailable rather than missing, "
            "and its bytes are gone from storage. The run still reads completed."
        ),
        "limit": (
            "The window is a platform constant rather than a per-definition or "
            "per-tenant setting."
        ),
    },
    {
        "id": "FR-013",
        "title": "Re-authorise Every Download, and Log the Refusals",
        "status": "Implemented",
        "requirement": (
            "Holding a link to a file must not be authority to read it, and a "
            "refused attempt must leave a trace."
        ),
        "evidence": (
            "Every download is authorised against the downloader at the moment of "
            "download, not against whoever produced or shared the file. "
            "ExportDownload records every attempt, allowed and refused alike, "
            "because who tried and was told no is exactly the question a compliance "
            "review asks and a refusal that leaves no trace cannot be answered."
        ),
        "acceptance": (
            "A shared file is refused to a recipient who may not read its contents, "
            "and the refusal appears in the download log."
        ),
        "limit": (
            "The log records the attempt, not the reason the caller thought they "
            "were entitled."
        ),
    },
    {
        "id": "FR-014",
        "title": "Require a Second Key Before Sensitive Columns Leave",
        "status": "Implemented",
        "requirement": (
            "Including a restricted column in an export must need more than the "
            "right to export that dataset."
        ),
        "evidence": (
            "exports.sensitive_field.export is required in addition to the "
            "dataset's own key. It is seeded to school_admin alone, so holding the "
            "run keys is not holding the right to take restricted columns out."
        ),
        "acceptance": (
            "A reader with the dataset key and without the sensitive key produces a "
            "file without those columns, reported as an omission rather than "
            "silently."
        ),
        "limit": (
            "The key is all-or-nothing across every dataset. It cannot be granted "
            "for one dataset's sensitive columns and withheld for another's."
        ),
    },
    {
        "id": "FR-015",
        "title": "Schedule an Export, and Survive a Missed Window",
        "status": "Implemented",
        "requirement": (
            "An export a school needs regularly must run on a timetable, in the "
            "school's own timezone, and a window missed because nothing was running "
            "must be handled deliberately rather than skipped in silence."
        ),
        "evidence": (
            "ExportSchedule carries the cadence and the scheduling module computes "
            "the next occurrence in a named zone, clamping a day that does not exist "
            "in a shorter month. should_run_missed decides whether a window passed "
            "while nothing was running is still worth running. Schedules can be "
            "paused and resumed."
        ),
        "acceptance": (
            "A monthly schedule on the 31st fires in February. A paused schedule "
            "produces no runs and resumes without back-filling every window it "
            "missed."
        ),
        "limit": (
            "A schedule produces a run like any other, so a missed window competes "
            "for the same workers as everything else."
        ),
    },
    {
        "id": "FR-016",
        "title": "Scope Every Read to the Caller's Tenant and Branch Reach",
        "status": "Implemented",
        "requirement": (
            "Definitions, runs, files and activity belong to a tenant, and a "
            "branch-pinned reader must not see beyond their own branches."
        ),
        "evidence": (
            "Reads are scoped to the asserted tenant and narrowed by the caller's "
            "branch reach, with dedicated coverage for branch scoping and for what "
            "a school's granted roles reach."
        ),
        "acceptance": (
            "One school cannot list another's runs. A branch-pinned reader's list "
            "excludes other branches' exports rather than refusing the whole list."
        ),
        "limit": (
            "Scope is applied on read. A definition written before a reader was "
            "pinned still exists; it simply stops being listed for them."
        ),
    },
    {
        "id": "FR-017",
        "title": "Measure Behaviour Without Touching Data",
        "status": "Implemented",
        "requirement": (
            "Product telemetry about how the Export Centre is used must not become "
            "a second copy of the data being exported, and must not pretend to be "
            "the audit trail."
        ),
        "evidence": (
            "ExportAnalyticsEvent is deliberately unlike an audit event: prunable "
            "rather than immutable, kept for a retention window rather than forever, "
            "and carrying counts and bucket labels only. Its properties are filtered "
            "through a schema on the way in, so no row data or field value can reach "
            "it, and it carries no entity column because it measures behaviour "
            "rather than data. A nightly job prunes past the window."
        ),
        "acceptance": (
            "A property outside the schema is dropped rather than stored. The "
            "boundary that applies is the tenant."
        ),
        "limit": (
            "Because it is prunable, analytics cannot answer a question older than "
            "the retention window. The audit trail is where durable answers live."
        ),
    },
    {
        "id": "FR-018",
        "title": "Let an Administrator Read Other People's Activity, Audibly",
        "status": "Implemented",
        "requirement": (
            "Someone must be able to review what the school has been exporting, and "
            "that review must itself be visible."
        ),
        "evidence": (
            "exports.activity.view is administrator-only and reading the activity "
            "history is itself audited, so looking at colleagues' exports is not a "
            "silent act."
        ),
        "acceptance": (
            "A reader without the key sees only their own activity. A reader with "
            "it leaves a record of having looked."
        ),
        "limit": (
            "The key is per tenant. It cannot be narrowed to one branch's activity."
        ),
    },
    {
        "id": "FR-019",
        "title": "Reach This Engine on Every Plan",
        "status": "Implemented",
        "requirement": (
            "Taking a school's own records out must not be priced above the "
            "cheapest plan."
        ),
        "evidence": (
            "Every export key answers to the data_export band, which sits at Core, "
            "so the catalogue, saved definitions, runs, files, schedules and the "
            "activity log are reachable on every plan. The gate that decides whether "
            "restricted columns may leave carries no band at all and is core to "
            "every school already."
        ),
        "acceptance": (
            "A school on the shallowest plan reaches the whole engine. What still "
            "governs an export is the reader's role and FR-014."
        ),
        "limit": (
            "A future decision to price part of this engine would move one band and "
            "close all of it, because the keys answer to one band."
        ),
    },
]

M26_TRACEABILITY = [
    ["Report catalogue and capability detail", "FR-001", "Implemented"],
    ["Report preview", "FR-002", "Implemented"],
    ["Quick export", "FR-003", "Implemented"],
    ["Saved export definitions", "FR-004", "Implemented"],
    ["Duplicate and share definitions", "FR-004, FR-005", "Implemented"],
    ["Run export definitions", "FR-006", "Implemented"],
    ["Sequenced export-run references, list, and detail", "FR-006", "Implemented"],
    ["Cancel and retry runs", "FR-010, FR-011", "Implemented"],
    ["Generated-file listing", "FR-012", "Implemented"],
    ["Controlled file download", "FR-013, FR-014", "Implemented"],
    ["Download logs", "FR-013", "Implemented"],
    ["Export analytics", "FR-017", "Implemented"],
    ["Delivery-access revocation", "FR-005, FR-013", "Implemented"],
    ["Export activity history", "FR-018", "Implemented"],
    ["File-expiry background tasks", "FR-011, FR-012, FR-017", "Implemented"],
    ["Domain-specific finance, procurement, audit, and import exports",
     "FR-001, FR-016", "Implemented"],
]

M26_NEEDS_ATTENTION = [
    ["1", "One band closes the whole engine",
     "Every export key answers to data_export. That is what made moving it to "
     "Core open the catalogue, the runs and the files together, and it equally "
     "means any future decision to price part of the engine would close all of "
     "it. There is no way to sell scheduled exports separately from manual "
     "ones.",
     "No action while the engine is Core. Recorded so a future repricing is "
     "understood to be all-or-nothing."],
    ["2", "The sensitive-field key is all-or-nothing",
     "exports.sensitive_field.export governs restricted columns across every "
     "dataset at once. A school wanting a bursar to export sensitive finance "
     "columns but not sensitive student ones cannot express that.",
     "Per-dataset sensitive grants, if a school ever asks for the "
     "distinction."],
    ["3", "An abandoned run looks live for up to half an hour",
     "The sweep that finishes runs whose worker died is on a thirty-minute "
     "cadence, so a reader watching a stalled export sees the same thing they "
     "would see one minute in.",
     "A staleness hint on the run detail, so a long-running export can be told "
     "from a lost one before the sweep catches it."],
    ["4", "File availability is a platform constant",
     "The thirty-day window is fixed for every school and every dataset. A "
     "school with a shorter retention obligation cannot hold files for less "
     "time, and one needing longer cannot hold them for more.",
     "A configuration definition for the window, per school, if retention "
     "policy ever requires it."],
    ["5", "Analytics cannot answer beyond its retention window",
     "Product telemetry is prunable by design, so a question about behaviour "
     "older than the window has no answer. This is correct and is recorded so "
     "the analytics summary is not mistaken for a durable record.",
     "No action. The audit trail is where durable answers live."],
]


# ── Module 25 ────────────────────────────────────────────────────────────────

M25_REQUIREMENTS = [
    {
        "id": "FR-001",
        "title": "Answer Each Dashboard From the Module That Owns Its Data",
        "status": "Implemented",
        "requirement": (
            "A dashboard must be served by the module whose records it summarises, "
            "so that a number on a screen and the list behind it cannot disagree."
        ),
        "evidence": (
            "There is no vs_dashboards application and no aggregation layer. The "
            "finance overview is served by vs_finance, procurement summaries by "
            "vs_procurement, payment movement by vs_payments, approval queues by "
            "vs_workflow, task views by vs_todo, the ticket dashboard by "
            "vs_tickets, the audit summary by vs_audit, the health overview by "
            "vs_health, export analytics by vs_exports, and the platform overview "
            "by vs_admin_console."
        ),
        "acceptance": (
            "Every dashboard figure is computed from the same queryset its own "
            "module's list endpoint uses, with the same tenant and branch scoping "
            "applied."
        ),
        "limit": (
            "A cross-domain figure has no owner, so there is no single endpoint "
            "answering 'how is the school doing' across modules."
        ),
    },
    {
        "id": "FR-002",
        "title": "Gate a Dashboard on Its Domain's Key, Not a Dashboard Key",
        "status": "Implemented",
        "requirement": (
            "Seeing a summary of something must require the same right as seeing "
            "the thing itself. There must be no key whose only effect is to reveal "
            "figures."
        ),
        "evidence": (
            "The finance overview requires finance.report.view, the procurement "
            "dashboard procurement.analytics.view, the audit summary "
            "platform.audit.view, and the team approval-load view the workflow "
            "instance key. Ticket and task dashboards reuse their own modules' "
            "permission sets."
        ),
        "acceptance": (
            "A reader refused a domain's records is refused its dashboard. Granting "
            "a dashboard does not require inventing a key that exists only for it."
        ),
        "limit": (
            "A reader holding a domain key sees the whole of that domain's "
            "dashboard. There is no per-card permission."
        ),
    },
    {
        "id": "FR-003",
        "title": "Serve a Reader Their Own Work Without a Key",
        "status": "Implemented",
        "requirement": (
            "A worklist made only of the caller's own items needs no permission, "
            "because holding a permission is not what makes it theirs."
        ),
        "evidence": (
            "Pending approvals returns instances the caller is eligible to act on "
            "and requires authentication alone; the personal task view and the "
            "platform console overview do the same. A reader with nothing waiting "
            "gets an empty list rather than a refusal."
        ),
        "acceptance": (
            "A teacher's approval queue is empty rather than forbidden. A "
            "self-scoped view cannot be pointed at somebody else."
        ),
        "limit": (
            "The distinction is per endpoint. A view mixing personal and team data "
            "has to gate on the team half."
        ),
    },
    {
        "id": "FR-004",
        "title": "Compose a Domain Overview as One Read",
        "status": "Implemented",
        "requirement": (
            "A dashboard is a screen, not a list of endpoints. Each domain "
            "overview must arrive in a single call carrying every block it draws."
        ),
        "evidence": (
            "The finance dashboard returns every overview block in one payload. "
            "Procurement, tickets, tasks and health follow the same shape, with "
            "narrower summaries such as accounts-receivable ageing and payment "
            "movement served alongside as their own reads."
        ),
        "acceptance": (
            "Opening a dashboard costs one request per domain rather than one per "
            "card. The blocks in one payload are internally consistent."
        ),
        "limit": (
            "A screen drawing blocks from two domains still makes two requests. "
            "That is the boundary in FR-001 rather than an oversight."
        ),
    },
    {
        "id": "FR-005",
        "title": "Scope Every Figure to the Caller's Tenant and Branch Reach",
        "status": "Implemented",
        "requirement": (
            "A summary must count only what the reader could have listed, so a "
            "figure never discloses the existence of records they may not open."
        ),
        "evidence": (
            "Each dashboard is computed from its own module's scoped queryset, so "
            "the tenant boundary and the caller's branch reach are applied by the "
            "same code that governs the list. No dashboard assembles its own "
            "queryset from raw models."
        ),
        "acceptance": (
            "A branch-pinned reader's totals cover their branches only. One "
            "school's figures never include another's rows."
        ),
        "limit": (
            "A count the reader may see, of rows they may not open individually, "
            "is possible where a module's list and detail scope differently."
        ),
    },
    {
        "id": "FR-006",
        "title": "Report Operational Health as Evidence, Never as a Default",
        "status": "Implemented",
        "requirement": (
            "An operational figure with no measurement behind it must say so "
            "rather than report a comfortable number."
        ),
        "evidence": (
            "The health overview is Module 30's and is evidence-based: seeding "
            "creates services, checks and rules but never fabricates measurements. "
            "Module 30's own FRD carries the one known exception, where a "
            "per-service uptime window returns a full result with no daily rollup "
            "behind it."
        ),
        "acceptance": (
            "A dashboard with no data behind it renders as unknown rather than as "
            "zero or as healthy."
        ),
        "limit": (
            "The exception above is tracked against Module 30, not here. This "
            "module consumes that surface rather than owning it."
        ),
    },
    {
        "id": "FR-007",
        "title": "Keep Product Telemetry Out of the Operational Picture",
        "status": "Implemented",
        "requirement": (
            "How people use a feature and what the business is doing are different "
            "questions, and the answer to one must not be presented as the other."
        ),
        "evidence": (
            "Export analytics is a separate pipeline from the audit trail: "
            "prunable, retention-bounded, carrying counts and bucket labels only, "
            "with properties filtered through a schema so no row data reaches it."
        ),
        "acceptance": (
            "A behavioural summary carries no record content. A durable question is "
            "answered from the audit trail rather than from analytics."
        ),
        "limit": (
            "Analytics answers nothing older than its retention window."
        ),
    },
    {
        "id": "FR-008",
        "title": "Leave the Unbuilt Domains Visibly Absent",
        "status": "Partial",
        "requirement": (
            "A dashboard must not invent a figure for a domain that does not exist "
            "yet, and the absence must be legible rather than look like a zero."
        ),
        "evidence": (
            "Academic, attendance, student and guardian indicators have no backend "
            "behind them because the modules that would produce them are planned "
            "rather than built. Nothing fabricates them, and no endpoint returns a "
            "placeholder figure."
        ),
        "acceptance": (
            "No dashboard reports an academic or attendance figure. Their absence "
            "is the honest state and is recorded in Needs Attention."
        ),
        "limit": (
            "This is the module's principal gap and is why it is not Complete. It "
            "closes when Modules 15 and 16 and the school-domain modules land."
        ),
    },
    {
        "id": "FR-009",
        "title": "Give the Platform Operator a Cross-School View",
        "status": "Implemented",
        "requirement": (
            "CodeX must be able to see the estate rather than one school at a time."
        ),
        "evidence": (
            "The console overview and the audit dashboard summary answer for the "
            "platform. The audit summary is gated on platform.audit.view; the "
            "console overview requires authentication as a platform actor and is "
            "self-scoped to what that actor may already reach."
        ),
        "acceptance": (
            "A school user cannot reach a cross-school figure. A platform operator "
            "sees the estate without asserting each tenant in turn."
        ),
        "limit": (
            "The cross-school view is operational. There is no commercial estate "
            "view, because nothing prices a school yet."
        ),
    },
    {
        "id": "FR-010",
        "title": "Own No Storage, and Therefore No Second Truth",
        "status": "Implemented",
        "requirement": (
            "This module must persist nothing of its own, so that a dashboard "
            "cannot drift from the records it describes."
        ),
        "evidence": (
            "There is no vs_dashboards application, no model, no migration and no "
            "scheduled rollup owned here. Every figure is computed at read time "
            "from its own module's rows."
        ),
        "acceptance": (
            "A record changed in a domain changes that domain's dashboard on the "
            "next read, with no rebuild, no cache to invalidate and no job to run."
        ),
        "limit": (
            "Computing at read time means an expensive summary is expensive every "
            "time. No dashboard is currently materialised."
        ),
    },
]

M25_TRACEABILITY = [
    ["Platform administration overview", "FR-009, FR-003", "Implemented"],
    ["Audit dashboard summary", "FR-002, FR-009", "Implemented"],
    ["Finance dashboard and statements", "FR-002, FR-004", "Implemented"],
    ["Accounts-receivable summaries", "FR-004, FR-005", "Implemented"],
    ["Procurement summaries and insights", "FR-002, FR-004", "Implemented"],
    ["Workflow pending and submitted views", "FR-003", "Implemented"],
    ["Team approval-load view", "FR-002", "Implemented"],
    ["Task personal, team, and organization views", "FR-003, FR-005", "Implemented"],
    ["Support-ticket dashboard", "FR-002, FR-004", "Implemented"],
    ["System-health overview", "FR-006", "Implemented"],
    ["Payment movement summary", "FR-004, FR-005", "Implemented"],
    ["Export analytics", "FR-007", "Implemented"],
]

M25_NEEDS_ATTENTION = [
    ["1", "The academic half of the school has no dashboard",
     "Academic, attendance, student and guardian indicators are blocked by the "
     "modules that would produce them. Nothing fabricates a figure, which is "
     "right, and it means a head teacher's dashboard today is an operations and "
     "money picture rather than a school one.",
     "Attendance and Gradebook, and the student-domain indicators that depend "
     "on them. This is why the module is Partial."],
    ["2", "No configurable layout",
     "Every dashboard draws the blocks its module ships. A school cannot "
     "reorder them, hide one, or choose which figures a role sees on opening "
     "the product.",
     "User-configurable widgets and per-role layouts, once the set of figures "
     "is stable enough to be worth arranging."],
    ["3", "No cross-domain figure has an owner",
     "Because each dashboard is served by the module owning its data, a number "
     "combining two domains has nowhere to live. There is no endpoint answering "
     "how the school as a whole is doing.",
     "A deliberate decision about where a cross-domain figure would be "
     "computed, before one is built somewhere by accident."],
    ["4", "Every summary is computed on read",
     "Nothing is materialised, which is what keeps a dashboard from drifting "
     "from its records and equally means an expensive summary is expensive on "
     "every open.",
     "Materialise a summary only where a measured read is too slow, and record "
     "the staleness it introduces when doing so."],
]


# ── Building ─────────────────────────────────────────────────────────────────


def add_real_bullets(doc, items, *, size=9):
    for item in items:
        paragraph = doc.add_paragraph()
        set_bullet_numbering(paragraph)
        write_paragraph(paragraph, item, size=size, space_after=2)


def set_bullet_numbering(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(ilvl)
    num_pr.append(num_id)


def add_callout(doc, title, lines, *, kind="info"):
    table = add_reference_callout(doc, title, lines, kind=kind)
    cell = table.cell(0, 0)
    for paragraph in cell.paragraphs[1:]:
        text = paragraph.text.removeprefix("• ")
        set_bullet_numbering(paragraph)
        write_paragraph(paragraph, text, size=8.6, space_after=2)
    return table


def hold_requirement_together(doc):
    """Keep a requirement's heading and banner with the rows beneath them."""
    table = doc.tables[-1]
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
    for row in table.rows[:2]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    for paragraph in reversed(doc.paragraphs):
        if paragraph.text.strip().startswith("FR-"):
            paragraph.paragraph_format.keep_with_next = True
            break


def contents_table(doc, rows):
    add_heading(doc, "Table of Contents", level=1)
    add_table(doc, ["Section", "Purpose"], rows, [2.35, 4.92], font_size=8.5)
    add_body(
        doc,
        "Use the Word Navigation pane to jump between headings. This contents "
        "page is intentionally static for reliable headless rendering.",
        size=8.5,
        color=GREY,
    )
    add_page_break(doc)


def requirements_section(doc, requirements):
    add_heading(doc, "4. Functional Requirements", level=1)
    add_body(
        doc,
        "Each requirement records the required behaviour, the inspected "
        "evidence, the acceptance boundary and the current limit. Status is "
        "current state, not revision history.",
    )
    for requirement in requirements:
        add_requirement(doc, requirement)
        hold_requirement_together(doc)
    add_page_break(doc)


def tail_sections(doc, *, needs, traceability, capability_count, module_number,
                  change_text):
    add_heading(doc, "9. Needs Attention", level=1)
    add_body(
        doc,
        "Current state, not history. An item leaves this section when "
        "implementation and verification resolve it, and is rewritten when the "
        "risk changes shape.",
        size=9,
        color=GREY,
    )
    add_table(
        doc,
        ["Pri.", "Current gap", "Detail", "Required completion"],
        needs,
        [0.45, 1.55, 3.3, 1.97],
        font_size=8.2,
    )
    add_page_break(doc)

    add_heading(doc, "10. MRD Traceability", level=1)
    add_body(
        doc,
        f"Module {module_number} of XVS Module Requirements Document "
        f"v{MRD_VERSION} lists {capability_count} capabilities. Each maps to "
        "the requirements above.",
        size=9,
    )
    add_table(
        doc,
        ["MRD capability", "FRD requirement", "State"],
        traceability,
        [3.3, 1.97, 2.0],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "11. Change Log", level=1)
    add_table(
        doc,
        ["Version", "Date", "Change"],
        [[FRD_VERSION, REVIEW_DATE, change_text]],
        [0.85, 1.25, 5.17],
        font_size=8.2,
    )


def start_document(reference, *, title, header, cover_title, control_rows,
                   boundary_lines):
    doc = Document(str(reference))
    remove_body_content(doc)
    doc.core_properties.title = title
    doc.core_properties.author = "CodeX Team"
    doc.core_properties.version = FRD_VERSION
    set_headers(doc, header)
    add_cover(
        doc,
        family="Functional Requirements Document",
        title=cover_title,
        subtitle="XVision Systems | Code-aligned functional baseline",
        version=FRD_VERSION,
    )
    add_heading(doc, "Document Control", level=1)
    add_metadata_table(doc, control_rows)
    add_callout(doc, "Evidence boundary", boundary_lines, kind="info")
    add_page_break(doc)
    return doc


def build_m26(reference: Path, output: Path) -> None:
    title = (
        "XVS M26 Reporting and Exports Functional Requirements Document "
        f"v{FRD_VERSION}"
    )
    doc = start_document(
        reference,
        title=title,
        header="CodeX | Reporting & Exports | Functional Requirements Document (FRD)",
        cover_title="Module 26: Reporting & Exports",
        control_rows=[
            ("Document", "Functional Requirements Document (FRD)"),
            ("Module", "M26 | Reporting & Exports"),
            ("Version", FRD_VERSION),
            ("Review date", REVIEW_DATE),
            ("Code baseline", M26_BASELINE),
            ("Source MRD",
             f"XVS Module Requirements Document v{MRD_VERSION} | Module 26, "
             "sixteen capability entries"),
            ("Primary app", "vs_exports"),
            ("Supporting apps",
             "core, vs_tenants, vs_rbac, vs_config, vs_audit, and every domain "
             "that registers a dataset"),
            ("Status", "Code-aligned baseline for Product and Engineering review"),
            ("Owner", "CodeX Team"),
        ],
        boundary_lines=[
            "Implemented means the stated backend path is present in the "
            "inspected code. It does not prove frontend completion, "
            "deployment, production adoption or data migration.",
            "A produced file is real data leaving the platform, so every claim "
            "about who may read one is a claim about disclosure. Where a rule "
            "is enforced elsewhere, that place is named.",
            "The MRD and this FRD carry independent versions.",
        ],
    )

    contents_table(doc, [
        ["Document Control", "Version, ownership, baseline and evidence rules"],
        ["1. Purpose and Scope", "What the Export Centre owns and what it refuses"],
        ["2. Context and Status Model", "Runs, files and the vocabularies that describe them"],
        ["3. Actors, Permissions, and Ownership", "Who may define, run, download and review"],
        ["4. Functional Requirements", "Testable behaviour, evidence, acceptance and limits"],
        ["5. Workflows and Lifecycle Rules", "From a definition to a file, and to its destruction"],
        ["6. Data Model and Relationships", "What is persisted, and the safety each row carries"],
        ["7. API and Validation Contracts", "Routes, keys, refusals and response rules"],
        ["8. Dependencies and Operational Evidence", "Workers, storage, RBAC and the plan"],
        ["9. Needs Attention", "Current gaps and required completion"],
        ["10. MRD Traceability", "All sixteen Module 26 capability mappings"],
        ["11. Change Log", "Independent FRD revision history"],
    ])

    add_heading(doc, "1. Purpose and Scope", level=1)
    add_body(
        doc,
        "Module 26 is how a school's own records leave the platform: a "
        "catalogue of what may be exported, definitions a reader saves and "
        "returns to, runs that produce files, and controlled access to those "
        "files for as long as they exist. It is domain-neutral. Each domain "
        "registers the datasets it is willing to expose, and this engine knows "
        "nothing about what a student, an invoice or a purchase order is.",
    )
    add_heading(doc, "1.1 In Scope", level=2)
    add_real_bullets(doc, [
        "The dataset catalogue, per-dataset capability detail, and preview.",
        "Quick exports from a filtered screen, and saved definitions that are "
        "duplicated, shared and run.",
        "Runs, their frozen configuration, their closed status vocabulary, "
        "failure codes and omission reasons.",
        "Files, their availability window, controlled download and the log of "
        "every attempt.",
        "Schedules in the school's own timezone, including a window missed "
        "while nothing was running.",
        "Product analytics about how the Export Centre is used, kept "
        "deliberately apart from the audit trail.",
    ])
    add_heading(doc, "1.2 Out of Scope", level=2)
    add_real_bullets(doc, [
        "What a dataset contains. A domain registers its own dataset and owns "
        "the query behind it.",
        "Whether a field is sensitive. Field-level rules belong to the module "
        "that owns the field; this engine enforces the extra key.",
        "Whether the school may export at all. That is the plan question, "
        "answered by Module 6 and enforced by Module 4.",
        "Dashboards. Reading a summary on screen is Module 25; taking it away "
        "as a file is here.",
    ])
    add_page_break(doc)

    add_heading(doc, "2. Context and Status Model", level=1)
    add_body(
        doc,
        "Two rows describe one export: the run, which is an attempt and is "
        "immutable once terminal, and the file, whose availability is "
        "computed and whose bytes are eventually destroyed. Keeping them "
        "separate is what lets a file expire without the history of a "
        "successful run being rewritten.",
    )
    add_heading(doc, "2.1 Requirement Status", level=2)
    add_status_key(doc)
    add_heading(doc, "2.2 Run and File Vocabulary", level=2)
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Queued, Running", "The run has not finished. Cancellable."],
            ["Completed", "The file was produced in full."],
            ["Completed with omissions",
             "A file exists and something was left out, explained by a "
             "structured reason list rather than by prose."],
            ["Failed", "No file. Carries a machine reason code, a user-safe "
                       "message and a support reference."],
            ["Cancelled", "Stopped deliberately. Terminal."],
            ["Expired", "Not a run status. A property of the file, derived at "
                        "read time, so a successful run stays successful."],
            ["Purged", "The bytes are gone from storage. Stamped on the file, "
                       "never on the run."],
        ],
        [1.75, 5.52],
        font_size=8.5,
    )
    add_page_break(doc)

    add_heading(doc, "3. Actors, Permissions, and Ownership", level=1)
    add_heading(doc, "3.1 Permission Matrix", level=2)
    add_table(
        doc,
        ["Actor", "May do", "Governed by"],
        [
            ["Any holder of the run keys",
             "See the catalogue, preview, run a quick export, and download "
             "what they produced.",
             "exports.catalogue.view, .run.view, .run.create, .file.download."],
            ["A definition's owner",
             "Save, edit, duplicate, delete, share and schedule it.",
             "exports.definition.* and exports.schedule.*."],
            ["A person a definition is shared with",
             "See the definition and its files. Data access is not shared: the "
             "run executes as the owner and every download is re-authorised "
             "against the downloader.",
             "The share row, plus the downloader's own reach."],
            ["A school administrator",
             "Include restricted columns, and read other people's export "
             "activity. Reading it is itself audited.",
             "exports.sensitive_field.export and exports.activity.view."],
        ],
        [1.55, 3.6, 2.12],
        font_size=8.3,
    )
    add_heading(doc, "3.2 Ownership Boundaries", level=2)
    add_real_bullets(doc, [
        "A domain owns its datasets and the query behind each. This engine "
        "owns the run, the file and who may take it.",
        "Module 6 owns whether the school reaches this engine at all. It sits "
        "at Core, so every school does.",
        "Module 4 enforces the keys. The extra key for restricted columns is "
        "checked here, in addition to the dataset's own.",
        "vs_audit receives the durable record. Product analytics is a separate "
        "and prunable pipeline that must not be mistaken for it.",
    ])
    add_page_break(doc)

    requirements_section(doc, M26_REQUIREMENTS)

    add_heading(doc, "5. Workflows and Lifecycle Rules", level=1)
    add_heading(doc, "5.1 From a Definition to a File", level=2)
    add_table(
        doc,
        ["Step", "What happens", "Effect"],
        [
            ["1", "A run is created and freezes its configuration.",
             "The run describes this file, not the definition's future."],
            ["2", "A worker produces the bytes.",
             "Sensitive columns are included only with the extra key."],
            ["3", "The run reaches a terminal status.",
             "Completed, completed with omissions, failed or cancelled. Never "
             "changed again."],
            ["4", "The file carries an availability window.",
             "Availability is computed on read rather than stored as a state."],
            ["5", "A nightly job purges expired bytes and stamps the file.",
             "The run still reads completed, because it did."],
        ],
        [0.6, 3.4, 3.27],
        font_size=8.4,
    )
    add_heading(doc, "5.2 Every Download", level=2)
    add_table(
        doc,
        ["Step", "What happens", "Effect"],
        [
            ["1", "The downloader is authorised, not the producer or sharer.",
             "A link is not authority."],
            ["2", "The attempt is recorded, allowed or refused.",
             "'Who tried and was told no' has an answer."],
            ["3", "An expired file is reported unavailable.",
             "Distinct from a file that never existed."],
        ],
        [0.6, 3.4, 3.27],
        font_size=8.4,
    )
    add_heading(doc, "5.3 Scheduled and Missed Windows", level=2)
    add_body(
        doc,
        "A schedule's next occurrence is computed in the school's own "
        "timezone, clamping a day that a shorter month does not have, so a "
        "monthly export on the 31st still fires in February. A window that "
        "passed while nothing was running is judged deliberately rather than "
        "skipped in silence, and a paused schedule resumes without "
        "back-filling everything it missed.",
        size=9,
    )
    add_page_break(doc)

    add_heading(doc, "6. Data Model and Relationships", level=1)
    add_table(
        doc,
        ["Model", "Holds", "Safety contract"],
        [
            ["ExportDefinition", "A saved export: dataset, columns, filters, "
                                 "format and values mode.",
             "Owned by one person. Editing it never alters a run already "
             "produced."],
            ["ExportDefinitionShare", "One person a definition is shared with.",
             "Grants sight, never the sharer's data access."],
            ["ExportSchedule", "The cadence and timezone a definition runs on.",
             "Pausable. A missed window is judged, not assumed."],
            ["ExportRun", "One attempt, with its frozen configuration and a "
                          "sequenced reference.",
             "Terminal statuses are final and are never rewritten."],
            ["ExportFile", "The produced bytes and how long they stay.",
             "Availability computed on read; purging stamps the file and "
             "leaves the run alone."],
            ["ExportDownload", "Every download attempt.",
             "Refusals are logged as well as successes."],
            ["ExportAnalyticsEvent", "Product telemetry: counts and bucket "
                                     "labels.",
             "Schema-filtered so no row data enters, prunable, and carrying no "
             "entity column because it measures behaviour."],
        ],
        [1.7, 2.75, 2.82],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "7. API and Validation Contracts", level=1)
    add_table(
        doc,
        ["Method and path", "Purpose and permission"],
        [
            ["GET /exports/catalogue/ and /{key}/",
             "What may be exported, and one dataset in detail. "
             "exports.catalogue.view."],
            ["GET /exports/capabilities/", "What this caller may do."],
            ["POST /exports/preview/", "A bounded sample before committing."],
            ["POST /exports/quick/ and /from-screen/",
             "Run without saving a definition. exports.run.create."],
            ["GET, POST /exports/definitions/",
             "Saved definitions. exports.definition.view and .create."],
            ["GET, PATCH, DELETE /exports/definitions/{pk}/",
             "One definition. .update and .delete."],
            ["POST /exports/definitions/{pk}/duplicate/",
             "Copy into a new definition owned by the caller."],
            ["POST /exports/definitions/{pk}/share/",
             "Share sight. exports.definition.share."],
            ["POST /exports/definitions/{pk}/run/", "Produce a run."],
            ["GET /exports/runs/ and /{pk}/", "Runs and one run. .run.view."],
            ["POST /exports/runs/{pk}/cancel/ and /retry/",
             "Stop one, or produce a successor. .run.cancel and .run.create."],
            ["GET /exports/files/", "Produced files and their availability."],
            ["GET /exports/files/{pk}/download/",
             "Take the bytes. exports.file.download, re-authorised per "
             "download."],
            ["GET /exports/files/{pk}/downloads/",
             "Who tried, and whether they were allowed."],
            ["GET, POST /exports/schedules/ and /{pk}/",
             "Timetables. exports.schedule.view, .create and .manage."],
            ["POST /exports/schedules/{pk}/pause/ and /resume/",
             "Stop and restart a timetable. .schedule.manage."],
            ["POST /exports/analytics/ and GET /analytics/summary/",
             "Product telemetry in, and the behavioural summary out."],
            ["GET /exports/activity/",
             "Other people's export activity. exports.activity.view, and "
             "reading it is audited."],
        ],
        [2.7, 4.57],
        font_size=8.3,
    )
    add_page_break(doc)
    add_heading(doc, "7.1 Refusals", level=2)
    add_table(
        doc,
        ["Condition", "Answer"],
        [
            ["A dataset absent from the caller's catalogue",
             "Refused. It was never offered."],
            ["A restricted column without the sensitive key",
             "The file is produced without it, reported as a structured "
             "omission rather than silently."],
            ["A download by somebody who may not read the contents",
             "Refused, and the refusal is logged."],
            ["A download of an expired file",
             "Reported unavailable, distinctly from a file that never "
             "existed."],
            ["Cancelling a run that has already finished",
             "Refused. Terminal statuses are final."],
            ["A run whose worker died",
             "Finished by the sweep within half an hour, with a failure code."],
            ["An analytics property outside the schema",
             "Dropped on the way in. No row data can enter the pipeline."],
        ],
        [3.0, 4.27],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "8. Dependencies and Operational Evidence", level=1)
    add_table(
        doc,
        ["Dependency", "Contract"],
        [
            ["Celery and a worker",
             "Produces every file, sweeps abandoned runs every half hour, and "
             "runs the nightly expiry and analytics pruning. Without a worker "
             "nothing is produced."],
            ["File storage",
             "Holds the bytes for the availability window and releases them on "
             "purge."],
            ["Module 6, Configuration & Capability",
             "Decides whether the school reaches this engine. data_export sits "
             "at Core, so every school does."],
            ["Module 4, Roles & Permissions",
             "Enforces the fifteen export keys, including the extra one for "
             "restricted columns."],
            ["Every domain registering a dataset",
             "Owns the query behind its dataset and the field-level rules on "
             "it. Finance, procurement, audit and import each register their "
             "own."],
            ["vs_audit",
             "Receives the durable record. Product analytics is separate and "
             "prunable."],
        ],
        [2.1, 5.17],
        font_size=8.3,
    )
    add_heading(doc, "8.1 Verification Evidence", level=2)
    add_body(
        doc,
        "The module's own suite covers the run lifecycle, file availability "
        "and download authorisation, with dedicated coverage for branch "
        "scoping and for what a school's granted roles reach. Backend "
        "evidence only; nothing here is deployed.",
        size=9,
    )
    add_page_break(doc)

    tail_sections(
        doc,
        needs=M26_NEEDS_ATTENTION,
        traceability=M26_TRACEABILITY,
        capability_count="sixteen",
        module_number=26,
        change_text=(
            "First code-aligned baseline for Module 26. Records the catalogue, "
            "preview and quick export; saved definitions, duplication and a "
            "share that grants sight without lending data access; the run as "
            "an attempt that freezes its own configuration and whose terminal "
            "status is never rewritten; failure codes and structured omission "
            "reasons instead of tracebacks and prose; the separation of run "
            "from file that lets bytes expire without a successful run "
            "becoming unsuccessful; download re-authorised against the "
            "downloader with refusals logged; the extra key restricted columns "
            "need; schedules in the school's own timezone including a missed "
            "window; and a product-analytics pipeline kept deliberately apart "
            "from the audit trail. Five current gaps are recorded, the first "
            "being that every export key answers to one band, so any future "
            "repricing of part of the engine would close all of it. Backend "
            "evidence only; nothing here is deployed."
        ),
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def build_m25(reference: Path, output: Path) -> None:
    title = (
        "XVS M25 Dashboards and Operational Analytics Functional Requirements "
        f"Document v{FRD_VERSION}"
    )
    doc = start_document(
        reference,
        title=title,
        header=(
            "CodeX | Dashboards & Operational Analytics | Functional "
            "Requirements Document (FRD)"
        ),
        cover_title="Module 25: Dashboards & Operational Analytics",
        control_rows=[
            ("Document", "Functional Requirements Document (FRD)"),
            ("Module", "M25 | Dashboards & Operational Analytics"),
            ("Version", FRD_VERSION),
            ("Review date", REVIEW_DATE),
            ("Code baseline", M25_BASELINE),
            ("Source MRD",
             f"XVS Module Requirements Document v{MRD_VERSION} | Module 25, "
             "twelve capability entries"),
            ("Primary app",
             "None. This module owns no application, and that is the decision "
             "it records."),
            ("Supporting apps",
             "vs_admin_console, vs_audit, vs_finance, vs_procurement, "
             "vs_payments, vs_workflow, vs_todo, vs_tickets, vs_health, "
             "vs_exports"),
            ("Status", "Code-aligned baseline for Product and Engineering review"),
            ("Owner", "CodeX Team"),
        ],
        boundary_lines=[
            "Implemented means the stated backend path is present in the "
            "inspected code. It does not prove frontend completion, "
            "deployment or production adoption, and this module in particular "
            "is consumed almost entirely by screens.",
            "This module is a set of contracts rather than an application. "
            "Every requirement below is enforced by the module that owns the "
            "data, and that module is named.",
            "The MRD and this FRD carry independent versions.",
        ],
    )

    contents_table(doc, [
        ["Document Control", "Version, ownership, baseline and evidence rules"],
        ["1. Purpose and Scope", "Why this module owns no application"],
        ["2. Context and Status Model", "Where each dashboard is served from"],
        ["3. Actors, Permissions, and Ownership", "Who sees which figures, and why"],
        ["4. Functional Requirements", "Testable behaviour, evidence, acceptance and limits"],
        ["5. Workflows and Lifecycle Rules", "How a figure is produced and scoped"],
        ["6. Data Model and Relationships", "What is persisted here, and why it is nothing"],
        ["7. API and Validation Contracts", "Every dashboard route and its key"],
        ["8. Dependencies and Operational Evidence", "The ten modules this one reads"],
        ["9. Needs Attention", "Current gaps and required completion"],
        ["10. MRD Traceability", "All twelve Module 25 capability mappings"],
        ["11. Change Log", "Independent FRD revision history"],
    ])

    add_heading(doc, "1. Purpose and Scope", level=1)
    add_body(
        doc,
        "Module 25 is the platform's answer to 'how are things going', and it "
        "is deliberately not an application. There is no vs_dashboards, no "
        "model, no migration and no rollup job owned here. Every dashboard is "
        "served by the module whose records it summarises, and what this "
        "module owns is the set of rules those modules follow so their "
        "dashboards behave like one product.",
    )
    add_heading(doc, "1.1 In Scope", level=2)
    add_real_bullets(doc, [
        "The contract that a dashboard is served by the module owning its "
        "data, computed from that module's own scoped queryset.",
        "The rule that a summary is gated on its domain's key rather than a "
        "key invented for the dashboard.",
        "The distinction between a self-scoped worklist, which needs no key, "
        "and a team or domain view, which does.",
        "The cross-school operator views, and the boundary that stops a school "
        "reaching them.",
        "The separation of operational evidence from product telemetry.",
    ])
    add_heading(doc, "1.2 Out of Scope", level=2)
    add_real_bullets(doc, [
        "The figures themselves. Each belongs to its own module and is "
        "specified in that module's FRD.",
        "Taking a summary away as a file. That is Module 26.",
        "The health signals behind the operational overview. Module 30 "
        "collects and owns them.",
        "Screen layout, ordering and which cards a role sees. No backend "
        "decides those today, which is recorded as a gap rather than "
        "described as a feature.",
    ])
    add_page_break(doc)

    add_heading(doc, "2. Context and Status Model", level=1)
    add_body(
        doc,
        "The module is Partial, and the reason is worth stating plainly: the "
        "operational and financial half of the school has dashboards and the "
        "academic half does not, because the modules that would produce "
        "academic figures are planned rather than built. Nothing fabricates "
        "the missing half, which is the correct behaviour and also the "
        "principal gap.",
    )
    add_heading(doc, "2.1 Requirement Status", level=2)
    add_status_key(doc)
    add_heading(doc, "2.2 Where Each Dashboard Lives", level=2)
    add_table(
        doc,
        ["Surface", "Served by", "Gated on"],
        [
            ["Platform administration overview", "vs_admin_console",
             "Authentication as a platform actor; self-scoped."],
            ["Audit dashboard summary", "vs_audit", "platform.audit.view"],
            ["Finance overview and statements", "vs_finance",
             "finance.report.view"],
            ["Accounts-receivable ageing", "vs_finance", "finance.report.view"],
            ["Procurement dashboard and insights", "vs_procurement",
             "procurement.analytics.view"],
            ["Payment movement and collection summaries", "vs_payments",
             "The payments module's own keys."],
            ["Pending approvals and my submissions", "vs_workflow",
             "Authentication; the list is the caller's own."],
            ["Team approval load", "vs_workflow", "The workflow instance key."],
            ["My tasks, team and organisation views", "vs_todo",
             "The task module's own permission set."],
            ["Support-ticket dashboard", "vs_tickets",
             "The ticket module's own permission set."],
            ["System-health overview", "vs_health",
             "Module 30's platform health keys."],
            ["Export analytics summary", "vs_exports",
             "The export module's own keys."],
        ],
        [2.35, 1.9, 3.02],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "3. Actors, Permissions, and Ownership", level=1)
    add_heading(doc, "3.1 Permission Matrix", level=2)
    add_table(
        doc,
        ["Actor", "May do", "Governed by"],
        [
            ["Any authenticated user",
             "See their own worklists: what is waiting on them to approve, and "
             "their own tasks.",
             "No key. The list is made of the caller's own items."],
            ["A domain's reader",
             "See that domain's dashboard, scoped to their tenant and branch "
             "reach.",
             "The same key that lets them list the domain's records."],
            ["CodeX platform staff",
             "See the estate: the console overview and the cross-school audit "
             "summary.",
             "Platform keys, and the platform tenant boundary."],
            ["Nobody",
             "See a figure for a domain that does not exist yet.",
             "There is no endpoint. Absence is the honest state (FR-008)."],
        ],
        [1.55, 3.6, 2.12],
        font_size=8.3,
    )
    add_heading(doc, "3.2 Ownership Boundaries", level=2)
    add_real_bullets(doc, [
        "A dashboard belongs to the module that owns its records. This module "
        "owns the rules, not the figures.",
        "A dashboard key is never invented. Seeing a summary requires the "
        "right to see the thing summarised.",
        "Module 30 owns the health signals; this module consumes the overview "
        "and does not fabricate a reading.",
        "Module 26 owns taking a summary away as a file. Reading it on screen "
        "is here.",
    ])
    add_page_break(doc)

    requirements_section(doc, M25_REQUIREMENTS)

    add_heading(doc, "5. Workflows and Lifecycle Rules", level=1)
    add_heading(doc, "5.1 Producing One Figure", level=2)
    add_table(
        doc,
        ["Step", "What happens", "Effect"],
        [
            ["1", "The caller reaches a domain's dashboard endpoint.",
             "The request is bound to one asserted tenant."],
            ["2", "The domain's own permission is checked.",
             "No dashboard-only key exists to be granted by mistake."],
            ["3", "The figure is computed from that module's scoped queryset.",
             "The reader's branch reach applies exactly as it does to the "
             "list."],
            ["4", "Every block for that domain returns in one payload.",
             "One request per domain rather than one per card."],
            ["5", "Nothing is stored.",
             "The next read reflects the records as they now are."],
        ],
        [0.6, 3.4, 3.27],
        font_size=8.4,
    )
    add_heading(doc, "5.2 Why There Is No Rollup", level=2)
    add_body(
        doc,
        "A materialised dashboard is a second copy of the truth, and a second "
        "copy drifts. Computing at read time means a figure and the list "
        "behind it cannot disagree, that there is no cache to invalidate when "
        "a record changes, and that no job has to run for a dashboard to be "
        "correct. The cost is that an expensive summary is expensive on every "
        "open, which is recorded as a gap to be closed one measured case at a "
        "time rather than by materialising everything.",
        size=9,
    )
    add_page_break(doc)

    add_heading(doc, "6. Data Model and Relationships", level=1)
    add_body(
        doc,
        "This module persists nothing. There is no model, no migration and no "
        "table owned by Module 25, and that is the decision rather than an "
        "omission: see FR-010. The rows every dashboard reads belong to the "
        "modules listed in section 2.2, and each is specified in that module's "
        "own document.",
        size=9,
    )
    add_heading(doc, "6.1 The One Thing Stored Nearby", level=2)
    add_body(
        doc,
        "Export analytics events are persisted, by Module 26, and are the "
        "closest thing to a dashboard-owned table. They are deliberately not "
        "one: prunable rather than immutable, carrying counts and bucket "
        "labels only, schema-filtered so no row data can enter, and measuring "
        "how a feature is used rather than what the business did.",
        size=9,
    )
    add_page_break(doc)

    add_heading(doc, "7. API and Validation Contracts", level=1)
    add_table(
        doc,
        ["Method and path", "Purpose and permission"],
        [
            ["GET /admin-console/dashboard/overview/",
             "The platform estate. Platform actor; self-scoped."],
            ["GET /audit/dashboard-summary/",
             "Cross-school audit activity. platform.audit.view."],
            ["GET /finance/reports/dashboard/",
             "Every finance overview block in one payload. "
             "finance.report.view."],
            ["GET /finance/reports/ar-aging/",
             "Receivables by age. finance.report.view."],
            ["GET /procurement/reports/dashboard/",
             "Procurement overview. procurement.analytics.view."],
            ["GET /procurement/categories/insights/ and vendor and item "
             "insights",
             "Narrower procurement views on the same key."],
            ["GET /payments/movements/summary/ and the collection, payout and "
             "batch summaries",
             "Money moving, on the payments module's own keys."],
            ["GET /workflow/dashboard/pending/ and /submitted/",
             "The caller's own approval work. Authentication only."],
            ["GET /workflow/dashboard/team-load/",
             "Approval load across a team. The workflow instance key."],
            ["GET /todo/dashboard/mine/, /team/ and /org/",
             "Personal, team and organisation task views."],
            ["GET /tickets/dashboard/",
             "Support workload counters."],
            ["GET /health/overview/",
             "Operational health. Module 30's keys, evidence-based."],
            ["GET /exports/analytics/summary/",
             "How the Export Centre is being used."],
        ],
        [2.9, 4.37],
        font_size=8.3,
    )
    add_page_break(doc)
    add_heading(doc, "7.1 Refusals", level=2)
    add_table(
        doc,
        ["Condition", "Answer"],
        [
            ["A reader without the domain's key",
             "403 from that domain, exactly as its list would answer."],
            ["A reader with nothing waiting on them",
             "An empty list, not a refusal. A self-scoped worklist is theirs "
             "whether or not it has anything in it."],
            ["A school user asking for a cross-school figure",
             "Refused at the tenant boundary before the figure is computed."],
            ["A branch-pinned reader",
             "Served, with totals covering their branches only."],
            ["A request for an academic or attendance figure",
             "There is no endpoint. Nothing returns a placeholder."],
            ["An operational reading with no measurement behind it",
             "Reported as unknown rather than as a comfortable default. One "
             "exception is tracked against Module 30."],
        ],
        [3.0, 4.27],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "8. Dependencies and Operational Evidence", level=1)
    add_table(
        doc,
        ["Dependency", "Contract"],
        [
            ["vs_finance, vs_procurement, vs_payments",
             "Own the money and buying figures, and the keys that gate them."],
            ["vs_workflow, vs_todo, vs_tickets",
             "Own the work queues, and the distinction between a personal "
             "worklist and a team one."],
            ["vs_audit, vs_admin_console",
             "Own the cross-school operator views and the platform tenant "
             "boundary that protects them."],
            ["vs_health",
             "Owns the operational overview, and the rule that a figure "
             "without evidence reports unknown."],
            ["vs_exports",
             "Owns the behavioural summary, kept apart from the audit trail."],
            ["vs_rbac, vs_tenants",
             "Enforce the domain keys and the tenant and branch scoping every "
             "figure inherits."],
            ["Modules 15, 16 and the school domains",
             "Would produce the academic half. Planned rather than built, "
             "which is why this module is Partial."],
        ],
        [2.1, 5.17],
        font_size=8.3,
    )
    add_heading(doc, "8.1 Verification Evidence", level=2)
    add_body(
        doc,
        "This module has no suite of its own, because it has no code of its "
        "own. Each dashboard is covered by the tests of the module that serves "
        "it, and the scoping every figure inherits is covered by that module's "
        "tenant and branch tests. Backend evidence only; nothing here is "
        "deployed.",
        size=9,
    )
    add_page_break(doc)

    tail_sections(
        doc,
        needs=M25_NEEDS_ATTENTION,
        traceability=M25_TRACEABILITY,
        capability_count="twelve",
        module_number=25,
        change_text=(
            "First code-aligned baseline for Module 25. Records the decision "
            "that this module owns no application: every dashboard is served "
            "by the module whose records it summarises, computed from that "
            "module's own scoped queryset, so a figure and the list behind it "
            "cannot disagree and nothing is materialised to drift. Records the "
            "rule that a summary is gated on its domain's key rather than a "
            "key invented for a dashboard, the distinction between a "
            "self-scoped worklist needing no key and a team view that does, "
            "the cross-school operator views and the boundary protecting them, "
            "and the separation of operational evidence from product "
            "telemetry. Names where each of the twelve surfaces is served from "
            "and what gates it. Four current gaps are recorded, the first "
            "being that the academic half of the school has no dashboard "
            "because the modules that would produce one are planned rather "
            "than built, which is why the module is Partial. Backend evidence "
            "only; nothing here is deployed."
        ),
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root)
    reference = (
        root / "functional-requirements" / "10-bulk-data-import"
        / "XVS_M10_Bulk_Data_Import_Functional_Requirements_Document_v1.2.docx"
    )

    m26 = root / "functional-requirements" / "26-reporting-and-exports"
    build_m26(reference, m26 / (
        "XVS_M26_Reporting_and_Exports_Functional_Requirements_Document_"
        f"v{FRD_VERSION}.docx"
    ))
    print(f"Wrote M26 FRD v{FRD_VERSION}")

    m25 = root / "functional-requirements" / "25-dashboards-and-analytics"
    build_m25(reference, m25 / (
        "XVS_M25_Dashboards_and_Operational_Analytics_Functional_Requirements_"
        f"Document_v{FRD_VERSION}.docx"
    ))
    print(f"Wrote M25 FRD v{FRD_VERSION}")


if __name__ == "__main__":
    main()
