#!/usr/bin/env python3
"""Reconcile the MRD's Module 25 and Module 26 entries to their new FRDs.

The same job Module 6 had: both modules had been carrying their detail in the
tracker because neither had a document of its own. Module 26's box held two
bullets narrating one past revision; Module 25's held two gaps out of the four
its FRD now records. Both are rewritten as current state pointing at the FRD.

    python tools/patch_analytics_reconciliation_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)

REVIEW_DATE = "9 September 2026"
CHANGE_DATE = "9 Sep 2026"

MRD_SOURCE, MRD_TARGET = "2.72", "2.73"
FRD_VERSION = "1.0"

SOURCE_SCOPE = (
    "Documentation change reconciling Modules 25 and 26 to their first "
    "Functional Requirements Documents, and moving their accumulated detail "
    "out of the tracker and into them (9 September 2026)"
)

INTRO = (
    "This revision adds no capability and changes no status. Two more modules "
    "stop keeping their documentation in the tracker: Reporting & Exports and "
    "Dashboards & Operational Analytics each have an FRD now, and their "
    "entries here say where they stand rather than how they behave."
)

DELTA_ROWS = [
    ["Module 26's decision box", "Rewritten as current state",
     "Two bullets narrating one past revision's reference format replaced by "
     "what is true now. M26 FRD v1.0 states the engine as nineteen testable "
     "requirements, including the sequenced reference those bullets "
     "described."],
    ["Module 25's gap list", f"Reconciled to M25 FRD v{FRD_VERSION}",
     "The tracker carried two of the four gaps the FRD now records, and none "
     "of the shape that produces them: this module owns no application, and "
     "every dashboard is served by the module whose records it summarises."],
    ["Both modules' traceability", "Reconciled",
     "All sixteen Module 26 capabilities and all twelve Module 25 "
     "capabilities map to a requirement. Counts, module statuses and "
     "integration statuses are unchanged."],
    ["Module 25 remains Partial", "Unchanged, deliberately",
     "The academic half of the school has no dashboard because the modules "
     "that would produce one are planned rather than built. In use is left as "
     "recorded: adoption is not something backend evidence can establish."],
]

M26_DECISION = [
    "Module 26 remains Backend Complete and In use Complete with sixteen "
    f"capability entries, reconciled to M26 FRD v{FRD_VERSION}.",
    "The detail now lives in the FRD: nineteen requirements covering the "
    "catalogue and preview, saved definitions and shares, the run and its "
    "frozen configuration, file availability and controlled download, "
    "schedules, and the analytics pipeline. What stays here is what other "
    "modules have to know.",
    "A run and its file are separate rows, and that is what lets bytes expire "
    "without a successful run becoming unsuccessful. Expired is a property of "
    "the file derived at read time, never a run status, so history is not "
    "rewritten to represent it.",
    "A share grants sight and never data access. The run executes as the "
    "owner, every download is re-authorised against the person downloading, "
    "and refusals are logged as well as successes, because who tried and was "
    "told no is what a compliance review asks.",
    "Restricted columns need a second key beyond the dataset's own, held by "
    "school_admin alone, and a file produced without it reports the omission "
    "rather than dropping the columns silently.",
    "Every school reaches this module in full: its keys answer to the "
    "data_export band, which sits at Core.",
    "Five current gaps are recorded in the FRD rather than here. The first is "
    "that one band closes the whole engine, so a future decision to price part "
    "of it would close all of it.",
]

M25_ATTENTION = [
    f"Reconciled to M25 FRD v{FRD_VERSION}, which records the shape these "
    "gaps follow from: this module owns no application, and every dashboard "
    "is served by the module whose records it summarises, gated on that "
    "domain's own key and computed at read time.",
    "The academic half of the school has no dashboard. Academic, attendance, "
    "student and guardian indicators are blocked by the modules that would "
    "produce them, nothing fabricates a figure in their place, and that is "
    "why this module is Partial.",
    "Dashboard layout is not configurable. Every dashboard draws the blocks "
    "its own module ships; a school cannot reorder them, hide one, or choose "
    "which figures a role sees on opening the product.",
    "No cross-domain figure has an owner. Because each dashboard belongs to "
    "the module owning its data, a number combining two domains has nowhere "
    "to live, and no endpoint answers how the school as a whole is doing.",
    "Every summary is computed on read. Nothing is materialised, which is "
    "what keeps a dashboard from drifting from its records and equally means "
    "an expensive summary is expensive on every open.",
]

CHANGE_SUMMARY = (
    "Reconciled Modules 25 and 26 to their first FRDs and moved their "
    "accumulated detail into them. Module 26's decision box carried two "
    "bullets describing one past revision's run-reference format; M26 FRD "
    "v1.0 states the engine as nineteen testable requirements and the box now "
    "carries what crosses module boundaries: that a run and its file are "
    "separate rows so bytes can expire without a successful run becoming "
    "unsuccessful, that a share grants sight and never data access with every "
    "download re-authorised against the downloader and refusals logged, that "
    "restricted columns need a second key, and that the engine sits at Core "
    "so every school reaches it. Module 25's box carried two of the four gaps "
    "its FRD now records and none of the shape producing them: that module "
    "owns no application, every dashboard is served by the module whose "
    "records it summarises, gated on that domain's own key rather than one "
    "invented for a dashboard, and computed at read time so nothing is "
    "materialised to drift. All sixteen and all twelve capability entries map "
    "to a requirement. No capability was added and no status changed: Module "
    "25 stays Partial because the academic half of the school has no "
    "dashboard, and its In use state is left as recorded because adoption is "
    "not something backend evidence can establish. The count stays at 492 "
    "across 31 modules. Documentation only; no code was touched."
)


def replace_cell(cell, text, **kwargs):
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def write_box(cell, title, lines, *, size=8.2):
    """Rewrite a decision or attention box, keeping its banner a banner.

    Only the paragraphs are replaced, so the cell's own shading survives and a
    Needs Attention box stays the colour that makes it one.
    """
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_paragraph(
        cell.paragraphs[0], title.upper(), size=size, bold=True, color=BLUE,
        space_after=3,
    )
    for line in lines:
        paragraph = cell.add_paragraph()
        write_paragraph(paragraph, f"• {line}", size=size, space_after=2)


def retitle(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_cover_version(table, source, target):
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def update_control_table(table, *, version, source_scope):
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], source_scope, size=9)


def prepend_change_log(table, version, date, summary):
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], date, size=8)
    replace_cell(row.cells[2], summary, size=8)


def keep_rows_whole(table):
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.core_properties.version = MRD_TARGET

    replace_cover_version(doc.tables[0], MRD_SOURCE, MRD_TARGET)
    update_control_table(
        doc.tables[1], version=MRD_TARGET, source_scope=SOURCE_SCOPE,
    )
    for row in doc.tables[2].rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(
                row.cells[0], f"5. v{MRD_TARGET} Capability Delta",
                size=9, bold=True, color=BLUE,
            )
            replace_cell(
                row.cells[1],
                "Modules 25 and 26 reconciled to their first FRDs",
                size=9,
            )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == f"5. v{MRD_SOURCE} Capability Delta":
            retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")
        elif text.startswith("This revision") and len(text) > 80:
            retitle(paragraph, INTRO)

    write_box(doc.tables[61].rows[-1].cells[0], "Needs attention", M25_ATTENTION)
    write_box(doc.tables[63].rows[-1].cells[0], "Current decision", M26_DECISION)

    rebuild_table(
        doc.tables[76],
        [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
        DELTA_ROWS,
        [1.85, 1.35, 3.85],
    )
    keep_rows_whole(doc.tables[76])

    prepend_change_log(doc.tables[78], MRD_TARGET, CHANGE_DATE, CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, f"XVS Module Requirements Document v{MRD_TARGET}")
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    mrd = Path(args.root) / "module-requirements"
    patch_mrd(
        mrd / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx",
        mrd / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
    )
    print(f"Wrote MRD v{MRD_TARGET}")


if __name__ == "__main__":
    main()
