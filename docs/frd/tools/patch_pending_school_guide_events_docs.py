#!/usr/bin/env python3
"""Cut M31 v1.10: a school being set up can record guide events, and its searches keep their words.

What changed in the backend (bf518f2c, b4a3d00a, 6cf38070), and therefore in the document:

* GuideAnalyticsEventView declares pending_tenant_surface, so a user whose
  school has not gone live records guide events as anyone else does. Before,
  every guide such a school opened was refused with TENANT_NOT_LIVE, and the
  school app read that refusal as a closed screen. An event carries no tenant
  or user, so nothing is withheld by closing it. The platform-wide summary
  declares nothing and stays closed.
* The ticket viewset's pending surface names every action a school uses (list,
  retrieve, create, update, transition, escalate, follow, comments,
  attachments and their download, audit) and the dashboard counts are open,
  so a school that has not gone live works its own desk (6cf38070). assign
  and eligible_assignees stay closed and destroy stays off; scoping and
  visibility are unchanged.
* SAFE_SEARCH_TERMS gains the school app's task nouns and verbs with the
  plurals people type, so a school's no-result search keeps its task words
  where it used to be stored as [redacted]. Names still redact.

What the school app holds, committed on main and pending release (school-fe
dfb6090, 3f629f8, d2e976f, cb933d8, ca372ce): a guide system of its own, with
97 guides its tests hold to cover every screen it mounts, 40 walkthroughs,
guides in the header search and the help panel, guides shown by role and plan,
events posted as silent requests with school.-prefixed guide ids, and the
screen, product area and guide on every ticket it raises.

Found while checking it, and recorded as current state:

* FR-008 described four context fields and omitted the keys other modules
  register, which the MRD capability it traces to is about.
* The console's editorial queue ranks only its own registry's guides.

The MRD was not revised here; it still carries seventeen Module 31 capability
entries under the same names, so only the Source MRD reference advances.

    python tools/patch_pending_school_guide_events_docs.py
"""
from __future__ import annotations

from docx import Document

from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    finish,
    frd_path,
    log_change,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    insert_row_after,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
    row_starting,
)

import patch_record_history_docs

REVIEW_DATE = "28 September 2026"
CODE_BASELINE = (
    "Backend commits bf518f2c, b4a3d00a and 6cf38070, 27 and 28 September 2026. The school app's guides and ticket "
    "context are committed on main (school-fe dfb6090, 3f629f8, d2e976f, cb933d8, "
    "ca372ce) and pending release"
)
MRD_SOURCE, MRD_TARGET = "v2.80", "v2.93"

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

M31_DIR = "31-support-tickets"
M31_STEM = "XVS_M31_Support_Tickets_Functional_Requirements_Document"
M31_SOURCE, M31_TARGET = "1.9", "1.10"

M31_SUMMARY = (
    "Minor revision. A school that has not gone live can record guide events: the intake "
    "declares the pending-tenant surface (bf518f2c), because a school being set up reads "
    "the guides more than anyone and an event carries no tenant or user. The summary stays "
    "closed. Before, such a school was refused TENANT_NOT_LIVE on every guide it opened, and "
    "the school app moved the reader to the not-live screen. The school app's own guide "
    "system is recorded as committed on main and pending release: 97 guides covering every "
    "screen it mounts, 40 walkthroughs, guides in the header search and the help panel, "
    "guide ids beginning \"school.\", and the screen, product area and guide on each ticket it "
    "raises. The no-result search vocabulary gains the school app's task words with the "
    "plurals people type (b4a3d00a), so a school's missed search keeps them while a name is "
    "still redacted. A school that has not gone live now works its own support desk "
    "(6cf38070): it lists and opens its tickets, reads and posts replies, attaches and "
    "downloads files, follows and edits them, and its desk staff move status, escalate and "
    "read the trail, with the dashboard counts, exactly as after go-live. Assigning an "
    "owner and the eligible-assignee list stay the platform desk's and closed, and scoping "
    "and visibility are unchanged. Before, such a school could file and attach but not read "
    "the answer, while the school app offered it the desk. FR-002 and FR-004 record it. "
    "FR-008 also records the module-registered context keys it had left out. FR-011, "
    "FR-012, the actor table, the routes, the typed errors and the dependencies follow. No capability or mapping changed, and the traceability is "
    "reconciled to MRD v2.93. Covered by test_a_school_being_set_up_can_record_guide_events, "
    "test_a_school_search_keeps_its_task_words_and_drops_the_child, "
    "PendingSchoolSupportDeskTests and PendingSchoolSurfaceTests; the suite was not "
    "run for this revision. Nothing here claims deployment."
)


def patch_m31() -> None:
    require_newest(str(ROOT / "functional-requirements" / M31_DIR / f"{M31_STEM}_v*.docx"),
                   M31_SOURCE)
    doc = Document(str(frd_path(M31_DIR, M31_STEM, M31_SOURCE)))

    # Cover and document control.
    set_cover_version(doc, M31_SOURCE, M31_TARGET)
    if f"Version: {M31_TARGET}" not in doc.tables[0].rows[0].cells[0].text:
        raise ValueError("The cover version was not rewritten")
    set_control(doc, "Version", M31_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    control = doc.tables[1]
    edit_cell(row(control, "Source MRD").cells[-1], MRD_SOURCE, MRD_TARGET)

    # 1.1 In Scope.
    scope = table_headed(doc, "Area", "Responsibility")
    edit_cell(row(scope, "Guide editorial signals").cells[-1],
              "Closed how-to event intake, ",
              "Closed how-to event intake from the console and the school app, open to a "
              "school that has not gone live, ")

    # 3. Actors.
    actors = table_headed(doc, "Actor", "May do", "Governed by")
    active = row(actors, "Active user")
    edit_cell(active.cells[1],
              "settled no-result-search event.",
              "settled no-result-search event, from the console or the school app, whether "
              "or not their school has gone live.")
    edit_cell(active.cells[2],
              "Active authentication, closed fields, and the guide_analytics throttle scope",
              "Active authentication, closed fields, the guide_analytics throttle scope, and "
              "the pending-tenant surface")

    # FR-008: the registered keys, and what the school app sends.
    fr008 = fr_table(doc, "FR-008")
    edit_cell(row(fr008, "Current evidence").cells[-1],
              "Creation accepts only four context fields - guide id, normalised route pattern, "
              "product area, and app version - each pattern-checked, with unknown fields "
              "rejected rather than dropped.",
              "Creation accepts four context fields of this module's own - guide id, "
              "normalised route pattern, product area from a closed list of twenty, and app "
              "version - each checked against its pattern or list, and any closed-vocabulary "
              "key another module registers from its own app: Module 9 registers "
              "onboarding_task_key and onboarding_readiness_state, so this module imports "
              "nothing of a school's. Unknown fields are rejected rather than dropped.")
    append_to(fr008, "Current evidence",
              " The school app (committed on main, pending release) sends the route pattern "
              "and product area of the screen a ticket is raised from, and the guide id of the "
              "guide written for that screen where there is one; its tests walk the mounted "
              "routes so that no screen resolves to no product area. While the school is "
              "being set up it adds the two onboarding keys.")
    edit_cell(row(fr008, "Current limit").cells[-1], "None.",
              "The product-area list is platform-wide and has no finer entry for a school's "
              "own screens, so the school app files tickets from Students, Staff, Academic "
              "Structure, the calendar, timetables and Branches under School management; the "
              "route pattern is what tells them apart.")

    # FR-011: the pending surface, the second client, and the vocabulary gap.
    fr011 = fr_table(doc, "FR-011")
    append_to(fr011, "Current evidence",
              " The intake declares the pending-tenant surface, so a user whose school has not "
              "gone live records events as anyone else does; the summary declares none and "
              "stays closed to that school. Opening the intake withholds nothing, because an "
              "event carries no tenant or user. Two applications post to it: the console, and "
              "the school app (committed on main, pending release), whose guide ids start with "
              "\"school.\" and walkthrough ids with \"walkthrough.school.\", and which sends every "
              "event as a silent request, so a refused or failed event never moves the reader "
              "off the guide they are reading.")
    append_to(fr011, "Acceptance",
              " test_a_school_being_set_up_can_record_guide_events proves the intake is on the "
              "pending-tenant surface and the summary is not. The vocabulary carries the "
              "school app's task words beside the console's - student, enrol, class, subject, "
              "teacher, guardian, term, timetable, exam and promote among them, with the "
              "plurals people type - so a school's missed search for \"enrol student Chiamaka "
              "into JSS1 class\" is stored as \"enrol student [redacted] class\": the task "
              "survives and the child's name does not "
              "(test_a_school_search_keeps_its_task_words_and_drops_the_child).")

    # FR-012: whose guides the aggregate carries, and who ranks them.
    fr012 = fr_table(doc, "FR-012")
    append_to(fr012, "Current limit",
              " The summary counts the school app's guides under their \"school.\" ids beside the "
              "console's, and both applications' no-result phrases and walkthrough exits "
              "share the one top-25 and top-50 list. The console's editorial queue ranks only "
              "the guides in its own registry, so a school guide's counts are returned but no "
              "screen yet prioritises them.")

    # FR-002 and FR-004: the desk before go-live.
    fr002 = fr_table(doc, "FR-002")
    append_to(fr002, "Current evidence",
              " A school that has not gone live works its own desk as it will after go-live: "
              "every action it uses is named on the viewset's pending-tenant surface, and the "
              "visibility rules apply unchanged. Setting an owner and the eligible-assignee "
              "list stay off it, as the platform desk's rota.")
    append_to(fr002, "Acceptance",
              " PendingSchoolSupportDeskTests proves such a school lists and opens only its "
              "own tickets, answers a reply, attaches and downloads a file, triages, reads its "
              "dashboard counts, and is still refused assignment.")
    fr004 = fr_table(doc, "FR-004")
    edit_cell(row(fr004, "Current evidence").cells[-1],
              "Attaching sits on the pending-tenant surface beside filing, so a school that has "
              "not gone live can show the screen it is reporting; the action still resolves its "
              "ticket through the caller's visibility, and the rest of the desk opens at "
              "go-live.",
              "Attaching and downloading sit on the pending-tenant surface with the rest of a "
              "school's desk (FR-002), so a school that has not gone live can show the screen "
              "it is reporting; both still resolve the ticket through the caller's "
              "visibility.")
    edit_cell(row(fr004, "Acceptance").cells[-1],
              "PendingSchoolAttachmentTests proves the attachments action is on that surface "
              "and the rest of the desk stays shut,",
              "PendingSchoolSurfaceTests proves the school's desk actions are on that surface "
              "and the platform desk's assignment actions are not,")

    # 7. Routes and refusals.
    edit_paragraph(doc, "Routes are mounted under /v1/support/.",
                   "rather than by an asserted entity.",
                   "rather than by an asserted entity. A school that has not gone live reaches "
                   "every route here except setting an owner, the eligible-assignee list and "
                   "the guide summary.")
    routes = [t for t in doc.tables
              if t.rows[0].cells[0].text.strip() == "Method and path"
              and any(r.cells[0].text.startswith("POST /support/guides/") for r in t.rows)]
    if len(routes) != 1:
        raise ValueError("The guide routes are not in exactly one route table")
    attach = row_starting(routes[0], "POST /support/tickets/{id}/attachments/")
    edit_cell(attach.cells[-1], "Validated before storage, and open to a school that has not "
              "gone live.", "Validated before storage.")
    events = row_starting(routes[0], "POST /support/guides/analytics/events/")
    edit_cell(events.cells[-1], "from an active user,",
              "from an active user, including one whose school has not gone live,")
    summary = row_starting(routes[0], "GET /support/guides/analytics/summary/")
    edit_cell(summary.cells[-1], "Needs platform.health.view.",
              "Needs platform.health.view, and stays closed to a school that has not gone "
              "live.")

    errors = table_headed(doc, "Condition or route", "Answer")
    insert_row_after(row_starting(errors, "Reading a ticket the caller may not see"), [
        "A school that has not gone live, setting or clearing an owner, listing eligible "
        "assignees, or reading the guide summary",
        "403 TENANT_NOT_LIVE. Every other route here is open to it, scoped as after go-live.",
    ])

    # 8. Dependencies.
    deps = table_headed(doc, "Dependency", "Contract")
    insert_row_after(row_starting(deps, "Module 3, Identity & Team"), [
        "Module 9, School Onboarding",
        "Owns the rule that closes a school's platform until go-live and registers the "
        "onboarding context keys. This module opens its whole desk and the guide event intake "
        "to such a school; setting an owner, the eligible-assignee list and the guide summary "
        "stay closed.",
    ])
    insert_row_after(row_starting(deps, "Platform core"), [
        "Client applications",
        "The console and the school app file tickets and post guide events. The school app's "
        "side is committed on main and pending release: a guides home and articles under "
        "/support/guides, 97 guides that its tests hold to cover every screen it mounts, 40 "
        "walkthroughs, guides in the header search and in the help panel for the screen "
        "underneath, guides shown by the reader's role and the school's plan, and the screen "
        "and guide on a ticket. Nothing here claims deployment.",
    ])

    # Keep the heading and its one-line intro on the page with the box.
    heading = paragraph_starting(doc, "9. Needs Attention")
    intro = paragraph_starting(doc, "These are current risks and gaps")
    heading.paragraph_format.keep_with_next = True
    intro.paragraph_format.keep_with_next = True

    # 10. Traceability.
    edit_paragraph(doc, "Module 31 carries 17 capability entries",
                   f"MRD {MRD_SOURCE}.", f"MRD {MRD_TARGET}.")

    log_change(doc, M31_TARGET, M31_SUMMARY)
    assert_absent_outside_log(doc, f"MRD {MRD_SOURCE}", "Creation accepts only four",
                              "lacks most of a school",
                              "rest of the desk opens at go-live",
                              "PendingSchoolAttachmentTests")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M31_DIR, M31_STEM, M31_TARGET),
           f"{M31_STEM.replace('_', ' ')} v{M31_TARGET}", M31_TARGET)


def main() -> None:
    patch_m31()


if __name__ == "__main__":
    main()
