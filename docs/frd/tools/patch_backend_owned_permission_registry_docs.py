#!/usr/bin/env python3
"""Version the MRD and RBAC FRD for a backend-owned access registry."""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)


REVIEW_DATE = "21 September 2026"
CHANGE_DATE = "21 Sep 2026"
MRD_SOURCE = "2.81"
MRD_TARGET = "2.82"
FRD_SOURCE = "1.21"
FRD_TARGET = "1.22"

MRD_SOURCE_SCOPE = (
    "Backend and console worktrees making permission definitions and shipped "
    "defaults backend-owned, with read-only catalogue screens and readable "
    "allocation labels, 21 September 2026"
)
FRD_SOURCE_SCOPE = (
    "Backend and console worktrees removing permission-definition writes while "
    "preserving custom roles and personal exceptions, 21 September 2026"
)

MRD_CHANGE_SUMMARY = (
    "Made the existing Module 4 permission registry a backend-owned definition "
    "surface. The global module, resource, action, permission, dependency, "
    "permission-group and prebuilt-role definitions remain relational database "
    "records because authorization, grants and dependency validation reference "
    "them, but their supported APIs now accept reads only. POST, PUT, PATCH and "
    "DELETE return 405. Definitions and shipped defaults are reconciled through "
    "backend seeders, and show_access_registry presents the complete hierarchy "
    "by module, resource and action with groups and prebuilt roles. Registry and "
    "personal-exception responses carry backend-owned readable labels, and the "
    "console removes definition forms, mutation hooks, dotted-key display and "
    "the four obsolete registry-write permission constants. The platform seeder "
    "creates only platform.permissions.view and marks legacy create, update, "
    "manage and delete keys inactive without cascading historical relationships. "
    "Custom tenant role creation, role permission allocation and personal ALLOW "
    "or DENY exceptions remain supported. Verified by the complete 798-test RBAC "
    "run, focused seeder and override tests, 15 frontend exception tests, guide "
    "checks and production builds. Module 4 remains Backend Complete and In use "
    "Complete with twenty-two capability entries. Nothing here claims deployment."
)

FRD_CHANGE_SUMMARY = (
    "Makes permission definitions and shipped defaults backend-owned while "
    "keeping their stored relational model. Global registry endpoints are now "
    "GET-only and require platform.permissions.view; all definition mutations "
    "return 405. Seeders remain the supported write boundary, and the new "
    "show_access_registry command gives a module, resource and action overview "
    "with default groups and prebuilt roles. Permission, dependency and personal "
    "exception responses carry readable backend labels. The console removes "
    "definition create, edit and delete routes and mutation hooks, while custom "
    "roles and personal ALLOW or DENY exceptions remain supported. The platform "
    "seeder registers only the registry view key and reversibly deactivates four "
    "legacy write keys. The complete 798-test RBAC run, focused seeder and "
    "override checks, 15 frontend exception tests, guide checks and production "
    "builds pass. The capability count remains twenty-two. Nothing here claims "
    "deployment."
)

MRD_DELTA_ROWS = [
    [
        "Permission-definition ownership",
        "Backend only",
        "Seeders own modules, resources, actions, permissions, dependencies, "
        "permission groups and prebuilt roles. Stored models remain because "
        "authorization and grants reference them.",
    ],
    [
        "Global registry API",
        "Read only",
        "Every definition collection and detail accepts GET. POST, PUT, PATCH "
        "and DELETE return 405 under one platform.permissions.view gate.",
    ],
    [
        "Registry overview",
        "Added",
        "show_access_registry prints definitions in module, resource and action "
        "order and includes default groups and prebuilt roles; JSON and module "
        "filters support review and automation.",
    ],
    [
        "Readable allocation",
        "Backend labelled",
        "Permission, dependency and personal-exception responses carry readable "
        "labels. The console no longer asks an operator to interpret dotted keys.",
    ],
    [
        "Custom tenant roles",
        "Preserved",
        "A tenant administrator may still create a custom role and allocate "
        "backend-defined permissions and groups through the existing controls.",
    ],
    [
        "Personal exceptions",
        "Preserved",
        "Personal grants and denies remain available with their existing scope, "
        "restricted-key, reason, expiry and self-edit controls.",
    ],
    [
        "Registry write permissions",
        "Retired reversibly",
        "Fresh databases receive only platform.permissions.view. Existing create, "
        "update, manage and delete rows are marked inactive so historical "
        "relationships are not cascade-deleted.",
    ],
]


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def update_control_table(table, *, version: str, source_scope: str, mrd_baseline=None):
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label in {"Source scope", "Code inspected"}:
            replace_cell(row.cells[1], source_scope, size=9)
        elif label == "MRD baseline" and mrd_baseline:
            replace_cell(row.cells[1], mrd_baseline, size=9)


def prepend_change_log(table, version: str, date: str, summary: str) -> None:
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], date, size=8)
    replace_cell(row.cells[2], summary, size=8)


def keep_rows_whole(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.core_properties.title = title
    doc.core_properties.version = MRD_TARGET

    replace_cover_version(doc.tables[0], MRD_SOURCE, MRD_TARGET)
    update_control_table(
        doc.tables[1], version=MRD_TARGET, source_scope=MRD_SOURCE_SCOPE,
    )

    for row in doc.tables[2].rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(
                row.cells[0], f"5. v{MRD_TARGET} Capability Delta",
                size=9, bold=True, color=BLUE,
            )
            replace_cell(
                row.cells[1],
                "Permission definitions are backend-owned and readable",
                size=9,
            )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == f"5. v{MRD_SOURCE} Capability Delta":
            write_paragraph(
                paragraph, f"5. v{MRD_TARGET} Capability Delta",
                size=17, bold=True, space_before=15, space_after=8,
            )
        elif text.startswith("Two-layer platform and school authorization"):
            write_paragraph(
                paragraph,
                "Two-layer platform and tenant authorization with reusable, "
                "backend-owned permission definitions, roles, templates, "
                "assignments, and controlled overrides. Modules, resources, "
                "actions, permissions, dependencies, shipped groups and prebuilt "
                "roles are reconciled in backend code and exposed read-only to the "
                "console with readable labels. Custom tenant roles and personal "
                "grants or denies remain supported allocation controls. Platform-only "
                "scope is enforced independently from restricted sensitivity. Each "
                "permission records the product capability that governs it, and a "
                "second gate refuses a request the caller's role allows but the "
                "tenant's plan does not reach. Branch reach remains grant-derived, "
                "and every supported role permission or group change records its "
                "actor, reason, source, approval reference where applicable, and "
                "complete before-and-after access configuration in the same transaction.",
                size=9, space_after=5,
            )

    replace_cell(
        doc.tables[14].rows[0].cells[0],
        "▸  Backend-owned permission registry by module, resource, and action",
        size=8.5,
    )
    decision = doc.tables[14].rows[-1].cells[0].text.rstrip()
    decision += (
        "\n• Global permission definitions and shipped defaults have one write "
        "boundary: backend seeders. The console reads the ordered catalogue, "
        "builds custom roles from it and applies personal exceptions, but it "
        "cannot create, rename or delete the definitions themselves. Legacy "
        "registry-write keys are deactivated without deleting their history."
    )
    replace_cell(doc.tables[14].rows[-1].cells[0], decision, size=8.2)

    rebuild_table(
        doc.tables[76],
        [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
        MRD_DELTA_ROWS,
        [1.85, 1.15, 4.05],
    )
    keep_rows_whole(doc.tables[76])
    prepend_change_log(doc.tables[78], MRD_TARGET, CHANGE_DATE, MRD_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_frd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{FRD_TARGET}"
    )
    doc.core_properties.title = title
    doc.core_properties.version = FRD_TARGET

    replace_cover_version(doc.tables[0], FRD_SOURCE, FRD_TARGET)
    update_control_table(
        doc.tables[1], version=FRD_TARGET, source_scope=FRD_SOURCE_SCOPE,
        mrd_baseline=f"XVS Module Requirements Document v{MRD_TARGET}",
    )

    for row in doc.tables[2].rows:
        if row.cells[0].text.strip() == "10. MRD Traceability":
            replace_cell(
                row.cells[1],
                f"Agreement with MRD v{MRD_TARGET}'s twenty-two capability entries",
                size=9,
            )

    replace_cell(
        doc.tables[3].rows[1].cells[1],
        "PermissionModule, PermissionResource and PermissionAction, from which a "
        "Permission key is composed as module.resource.action and stored as the "
        "primary key. Backend seeders own their creation and maintenance; supported "
        "APIs expose them read-only.",
        size=8.5,
    )
    replace_cell(
        doc.tables[3].rows[2].cells[1],
        "Permission rows carrying a readable label, description, sensitivity, "
        "restricted flag, active flag, scope and product capability. Backend code "
        "owns definitions; the console reads them for review and allocation.",
        size=8.5,
    )
    replace_cell(
        doc.tables[3].rows[4].cells[1],
        "Reusable backend-defined bundles attachable to a custom role. A group "
        "carries its own scope, membership is checked against that declaration, "
        "and restricted permissions are excluded because attachment has no per-key "
        "approval record.",
        size=8.5,
    )
    replace_cell(
        doc.tables[3].rows[6].cells[1],
        "PrebuiltRoleTemplate and its defaults, a backend-owned library of five "
        "templates provisioned into a tenant's own roles rather than shared with it. "
        "Tenant role creation remains separate and mutable.",
        size=8.5,
    )

    replace_cell(
        doc.tables[7].rows[3].cells[1],
        "Read the global permission catalogue and default groups, administer "
        "platform and tenant roles, and manage permitted personal exceptions. "
        "Permission definitions and shipped defaults change through backend code "
        "and seeding, not through the console.",
        size=8.2,
    )
    replace_cell(
        doc.tables[7].rows[3].cells[2],
        "platform.permissions.view for the registry; platform.roles.* and the "
        "relevant override keys for allocation. Definition write keys are retired.",
        size=8.2,
    )

    fr001 = doc.tables[9]
    replace_cell(
        fr001.rows[1].cells[1],
        "A permission key is composed from a module, resource and action that each "
        "exist as a row. Those definitions are owned in backend code and exposed "
        "read-only, so every environment receives the same reviewed vocabulary.",
        size=8.5,
    )
    replace_cell(
        fr001.rows[2].cells[1],
        "Permission.save composes key as f\"{module_id}.{resource.name}.{action_id}\". "
        "The master and module seeders reconcile the vocabulary, dependencies, "
        "default groups and prebuilt roles. Registry collection and detail views "
        "use ListAPIView and RetrieveAPIView under platform.permissions.view, so "
        "POST, PUT, PATCH and DELETE return 405. show_access_registry prints the "
        "hierarchy and shipped defaults in one place.",
        size=8.5,
    )
    replace_cell(
        fr001.rows[3].cells[1],
        "finance.invoice.view appears under Finance, Invoice and View with a "
        "backend-provided readable label. An authenticated reader can inspect it; "
        "no API caller can mint a different definition or rewrite its identity.",
        size=8.5,
    )
    replace_cell(
        fr001.rows[4].cells[1],
        "A semantic rename or retirement requires a deliberate backend transition "
        "that accounts for role grants, groups, dependencies, overrides, defaults "
        "and pending changes. The database models remain because those relationships "
        "are runtime authorization data.",
        size=8.5,
    )

    fr023 = doc.tables[31]
    replace_cell(
        fr023.rows[2].cells[1],
        "seed_all_permissions runs the definition sources in dependency order. "
        "Module seeders classify keys, group seeders remove or skip restricted "
        "members, and prebuilt templates provide trusted bootstrap roles. The "
        "platform seeder registers only platform.permissions.view and marks legacy "
        "registry create, update, manage and delete keys inactive. Registry writes "
        "are unavailable through the API.",
        size=8.5,
    )
    replace_cell(
        fr023.rows[3].cells[1],
        "A fresh environment receives the same backend-defined hierarchy and "
        "defaults. Re-running the seeders is idempotent, preserves custom tenant "
        "roles and deactivates obsolete registry-write keys without deleting their "
        "historical relationships.",
        size=8.5,
    )
    replace_cell(
        fr023.rows[4].cells[1],
        "Seeders are the supported authors, but no aggregate check yet proves that "
        "every active seeded permission carries a scope. A missing scope still "
        "fails closed when a non-platform tenant tries to hold the key.",
        size=8.5,
    )

    replace_cell(
        doc.tables[38].rows[1].cells[2],
        "Backend-owned vocabulary. Modules and actions are slug primary keys; a "
        "resource is unique within its module. Supported APIs are read-only.",
        size=8.2,
    )
    replace_cell(
        doc.tables[38].rows[2].cells[2],
        "Primary key is the dotted key itself. Carries the backend-readable label "
        "source, scope, sensitivity, restricted and active states, and product "
        "capability. Supported APIs are read-only.",
        size=8.2,
    )

    rebuild_table(
        doc.tables[40],
        ["Method and path", "Purpose and permission"],
        [
            ["GET /rbac/vision/permission-modules/", "List backend-defined modules. platform.permissions.view."],
            ["GET /rbac/vision/permission-modules/{name}/", "Read one backend-defined module. platform.permissions.view."],
            ["GET /rbac/vision/permission-resources/", "List resources within modules. platform.permissions.view."],
            ["GET /rbac/vision/permission-resources/{id}/", "Read one resource. platform.permissions.view."],
            ["GET /rbac/vision/permission-actions/", "List reusable actions. platform.permissions.view."],
            ["GET /rbac/vision/permission-actions/{name}/", "Read one action. platform.permissions.view."],
            ["GET /rbac/vision/permissions/", "List permissions with backend-readable module, resource and permission labels. platform.permissions.view."],
            ["GET /rbac/vision/permissions/{key}/", "Read one permission and its relationships. platform.permissions.view."],
            ["GET /rbac/vision/permission-dependencies/", "List prerequisite edges with readable labels. platform.permissions.view."],
            ["GET /rbac/vision/permission-dependencies/{id}/", "Read one prerequisite edge. platform.permissions.view."],
            ["GET /rbac/vision/permission-groups/", "List backend-defined reusable bundles. platform.permissions.view."],
            ["GET /rbac/vision/permission-groups/{id}/", "Read one bundle and its members. platform.permissions.view."],
        ],
        [2.55, 4.5],
        font_size=8.2,
    )
    keep_rows_whole(doc.tables[40])

    for row in doc.tables[42].rows:
        if row.cells[0].text.strip() == "Permission creation without scope":
            replace_cell(
                row.cells[0],
                "A definition mutation on a global registry route",
                size=8.2,
            )
            replace_cell(
                row.cells[1],
                "405. The same route remains readable through GET for a holder of platform.permissions.view.",
                size=8.2,
            )

    for row in doc.tables[44].rows:
        if row.cells[0].text.strip() == "P2" and "platform registry API requires scope" in row.cells[1].text:
            replace_cell(
                row.cells[1],
                "Backend seeders are now the only supported authors of active "
                "permissions, but no aggregate check proves that every active "
                "seeded key is classified. A forgotten scope remains safe at "
                "grant time, where it is refused, but the error appears after "
                "seeding rather than during it.",
                size=8.2,
            )

    replace_cell(
        doc.tables[46].rows[1].cells[0],
        "Backend-owned permission registry by module, resource, and action",
        size=8.2,
    )
    replace_cell(
        doc.tables[46].rows[1].cells[2],
        "Implemented. Seeders own definitions and shipped defaults, the API is "
        "read-only, and backend labels let the console present permissions without "
        "exposing dotted keys. Custom roles remain mutable allocation records.",
        size=8.2,
    )
    replace_cell(
        doc.tables[46].rows[14].cells[2],
        "Implemented with limits. Seeders provide one reviewed source, obsolete "
        "registry-write keys are inactive, system roles provide bootstrap, groups "
        "exclude restricted members and every active key must gate a view; aggregate "
        "scope classification remains open.",
        size=8.2,
    )

    replace_cell(
        doc.tables[47].rows[0].cells[0],
        f"MRD RECONCILIATION\n"
        f"• MRD v{MRD_TARGET} lists Module 4 as Roles & Permissions (RBAC), Phase V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two capability entries. The module number, name, phase, states and ownership agree with this revision.\n"
        "• The count does not change. Backend ownership, read-only definition APIs, readable allocation labels and the registry overview command restate the existing permission-registry and seeding capabilities.\n"
        "• Custom tenant role creation and personal exceptions remain supported. Definition ownership changes without removing the allocation surfaces tenants use.\n"
        "• Backend and console verification are recorded separately from deployment and production adoption.",
        size=8.5,
    )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Module 4 is the platform's authorization boundary"):
            write_paragraph(
                paragraph,
                "Module 4 is the platform's authorization boundary. Every other "
                "module asks it whether a person may act and whose rows they may "
                "see. Permission definitions and shipped defaults are owned in "
                "backend code and read through one global catalogue. Tenants still "
                "compose custom roles from that catalogue and apply controlled "
                "personal grants or denies. Scope, restricted state, plan depth, "
                "branch reach and audit continue to govern those allocations.",
                size=9, space_after=5,
            )
        elif text.startswith("•  Registry integrity:"):
            write_paragraph(
                paragraph,
                "•  Backend ownership: every global registry collection refuses "
                "POST and every detail refuses PUT, PATCH and DELETE; GET returns "
                "the ordered definitions and readable labels; custom role creation "
                "and personal exceptions still succeed through their tenant routes.",
                size=9, space_after=3,
            )
        elif text.startswith("MRD v2.79 records Module 4"):
            write_paragraph(
                paragraph,
                f"MRD v{MRD_TARGET} records Module 4 as Roles & Permissions (RBAC), Phase V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two capability entries. The module number, name, phase, states and ownership agree with this revision. Each entry maps below.",
                size=9, space_after=5,
            )

    prepend_change_log(doc.tables[48], FRD_TARGET, CHANGE_DATE, FRD_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    mrd_dir = root / "module-requirements"
    frd_dir = root / "functional-requirements" / "04-roles-and-permissions-rbac"

    mrd_source = mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"
    mrd_output = mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx"
    frd_stem = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
    frd_source = frd_dir / f"{frd_stem}_v{FRD_SOURCE}.docx"
    frd_output = frd_dir / f"{frd_stem}_v{FRD_TARGET}.docx"

    patch_mrd(mrd_source, mrd_output)
    patch_frd(frd_source, frd_output)
    print(mrd_output)
    print(frd_output)


if __name__ == "__main__":
    main()
