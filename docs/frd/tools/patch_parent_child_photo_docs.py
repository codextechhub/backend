#!/usr/bin/env python3
"""Patch M11 and M28 for guardian access to a child's photograph."""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
import patch_record_history_docs as common
from patch_record_history_docs import (
    append_rows,
    finish,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
)
from patch_staff_id_and_auth_events_docs import repair_ooxml


ROOT = Path(__file__).resolve().parents[1]
REVIEW_DATE = "28 October 2026"
CODE_BASELINE = (
    "Backend main at 3d0b44c4 (28 October 2026), covering the committed "
    "parent child-photograph behavior at d5e2c5fa. Production deployment "
    "and adoption are not inferred."
)
common.REVIEW_DATE = REVIEW_DATE


M11_DIR = ROOT / "functional-requirements" / "11-student-management"
M11_STEM = "XVS_M11_Student_Management_Functional_Requirements_Document"
M11_SOURCE = M11_DIR / f"{M11_STEM}_v2.18.docx"
M11_TARGET = M11_DIR / f"{M11_STEM}_v2.18.1.docx"

M28_DIR = ROOT / "functional-requirements" / "28-parent-portal"
M28_STEM = "XVS_M28_Parent_Portal_Functional_Requirements_Document"
M28_SOURCE = M28_DIR / f"{M28_STEM}_v1.5.docx"
M28_TARGET = M28_DIR / f"{M28_STEM}_v1.5.1.docx"


def table_starting(doc, text: str):
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text.strip().startswith(text):
            return table
    raise ValueError(f"No table starts with {text!r}")


def set_row(table, label: str, value: str) -> None:
    for row in table.rows[1:]:
        if row.cells[0].text.strip() == label:
            keep_format(row.cells[-1], value)
            return
    raise ValueError(f"No row labelled {label!r} in {table.rows[0].cells[0].text!r}")


def insert_contract(doc, before: str, heading: str, body: list[str]) -> None:
    target_index, target = next(
        (i, p) for i, p in enumerate(doc.paragraphs) if p.text.strip() == before
    )
    prior_heading = next(
        p for p in reversed(doc.paragraphs[:target_index])
        if p.text.strip().startswith("Current Contracts at ")
    )
    prior_heading_index = next(
        i for i, p in enumerate(doc.paragraphs) if p._p is prior_heading._p
    )
    prior_body = next(
        p for p in doc.paragraphs[prior_heading_index + 1 :]
        if p.text.strip()
    )
    for source, text in [(prior_heading, heading), *((prior_body, item) for item in body)]:
        clone = copy.deepcopy(source._p)
        target._p.addprevious(clone)
        paragraph = Paragraph(clone, target._parent)
        mrd_tools.retitle(paragraph, text)


def trim_trailing_empty_paragraphs(doc) -> None:
    """Remove empty body paragraphs after the final table."""
    body = doc.element.body
    while len(body) >= 2:
        element = body[-2]
        if not element.tag.endswith("}p"):
            break
        paragraph = Paragraph(element, doc)
        if paragraph.text.strip() or element.xpath('.//w:drawing | .//w:pict | .//w:br'):
            break
        body.remove(element)


def patch_m11() -> None:
    doc = Document(M11_SOURCE)
    set_cover_version(doc, "2.18", "2.18.1")
    set_control(doc, "Version", "2.18.1")
    set_control(doc, "Date", REVIEW_DATE)
    set_control(doc, "Supersedes", "v2.18")
    set_control(doc, "Verified against", CODE_BASELINE)

    corrections = doc.tables[1]
    append_rows(corrections, [[
        "v2.18 (FR-015) lists only the staff document keys and does not record guardian photograph access.",
        "A pupil's stored files have three independent read paths. Staff holding school.students.view read every file within their branch scope. The pupil reads their own passport photograph only. A guardian whose active account is linked to an on-roll pupil reads that pupil's passport photograph only. Roles and branches do not replace the guardian link, so a teacher-parent may read their own child's photograph at another branch while staff scope still controls every other pupil.",
        "media_policies.py; guardians.py, is_portal_guardian_of; the student and parent-portal child-photo tests; commit d5e2c5fa.",
    ]])

    requirement = table_starting(doc, "FR-015  Student Documents")
    set_row(
        requirement,
        "Keys",
        "school.students.view to list and to read any stored file as staff, within the caller's branch scope; school.students.update to attach or remove. A pupil and a guardian need no staff key for the narrow passport-photograph alternatives below.",
    )
    set_row(
        requirement,
        "Business rules",
        "(1) The type must be one of the five in section 7.6 and anything else answers 400 on document_type. (2) Attaching a type the student already holds replaces it: the old StudentDocument row is deleted first, which is what core.binding's post_delete hook needs in order to RETIRE the superseded file. Without the delete the previous birth certificate keeps its own live URL for ever, and the school believes it has replaced it. (3) The file is written through core.StoredFile and served only through the authenticated media view; no public URL is ever returned. (4) A missing required document never blocks an enrolment, an edit or a status change. The requirement is a prompt, not a gate, and the design says so under the list. Only the birth certificate carries it. (5) Removal is audited separately from attachment, because removal is the direction that loses evidence. (6) Nothing may exceed 5MB, and a PASSPORT_PHOTO must be an image: JPEG, PNG, WebP or HEIC. A PDF remains a valid birth certificate. (7) PASSPORT_PHOTO is what photo_url resolves to. Replacing it from the checklist changes the face on every list, and replacing it from the avatar changes what the checklist shows, because there is one document. (8) Three independent callers may read a passport photograph: staff holding school.students.view inside their branch scope; the active pupil account linked to that pupil; and the active account of a guardian in the same school linked to that pupil while the pupil has a parent-portal status. A guardian's authority is the relationship, not a role or staff branch. (9) Birth certificates, medical letters and every other document remain staff-only. A teacher-parent therefore reads their own child's photograph through the guardian link even when the child is at another branch, while their staff scope still decides every other pupil and every non-photograph file.",
    )
    set_row(
        requirement,
        "Field exposure",
        "The staff document list carries the type, whether it is attached, when, by whom, and a url. It never carries the bytes. The parent child payload already carries photo_url; this change makes that existing signed link readable by the linked guardian and adds no field or Field Access switch. The staff school.students.photo Read switch is unchanged. Every url is signed, user-bound, expiring and absolute, and the media view still requires the JWT.",
    )
    set_row(
        requirement,
        "Acceptance",
        "(1) A caller without school.students.update cannot attach or remove, and receives 403. (2) A document of another tenant's student answers 404 by id. (3) Attaching the same type twice leaves one row. (4) Enrolling with no documents succeeds and the profile shows all five as not on file, with the birth certificate alone marked required. (5) The media path returned refuses an unauthenticated fetch. (6) Removal writes STUDENT_DOCUMENT_REMOVED and attachment writes STUDENT_DOCUMENT_ATTACHED. (7) Attaching a PASSPORT_PHOTO fills photo_url on the profile and directory row, removing it empties both, and the url is absolute. (8) A page of students costs a fixed number of queries because the photograph is prefetched. (9) An active guardian linked to an on-roll pupil may load that pupil's passport photograph. (10) An unrelated guardian, a guardian from another school, and the guardian of a pupil outside the portal statuses are refused. (11) A pupil and guardian cannot load the pupil's birth certificate, medical letter or any other non-photograph document. (12) A teacher-parent may load their own child's photograph at another branch, but cannot use that relationship to read another pupil or a non-photograph file outside staff scope. (13) The guardian rule works at a school with one branch and at a school with several branches.",
    )

    insert_contract(
        doc,
        "15. Change Log",
        "Current Contracts at 28 October 2026",
        [
            "A pupil's stored files have three independent readers. Staff holding school.students.view read every file within their branch scope. The pupil reads their own passport photograph only. An active guardian account linked to an on-roll pupil reads that pupil's passport photograph only.",
            "Birth certificates, medical letters and every other document remain staff-only. A teacher-parent may read their own child's photograph through the guardian link at another branch, while their staff scope still controls every other pupil. The existing photo_url and staff photograph Field Access switch are unchanged.",
        ],
    )
    inherited_note = next(
        p for p in doc.paragraphs if p.text.startswith("ℹ  Version 2.0 was a major version")
    )
    inherited_note._p.getparent().remove(inherited_note._p)
    log_change(
        doc,
        "2.18.1",
        "Clarifies the three independent readers of a pupil photograph, keeps every other student document staff-only, and records the teacher-parent branch boundary. MRD v2.115 is unchanged.",
    )
    trim_trailing_empty_paragraphs(doc)
    repair_ooxml(doc)
    finish(
        doc,
        M11_TARGET,
        "XVS M11 Student Management Functional Requirements Document v2.18.1",
        "2.18.1",
    )


def patch_m28() -> None:
    doc = Document(M28_SOURCE)
    set_cover_version(doc, "1.5", "1.5.1")
    set_control(doc, "Version", "1.5.1")
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(
        doc,
        "Status",
        "The parent backend is mounted and its core child, fee, notice, profile and messaging reads are implemented. Consent, general document actions, notification preferences, Assignments and the unread-messages home block remain incomplete.",
    )

    child_list = table_starting(doc, "FR-002 | Implemented")
    set_row(child_list, "Current evidence", "The parent me response lists only pupils linked through the signed-in guardian account. Each child row includes the existing photo_url, whose media policy now admits that linked guardian while the pupil remains in the parent-portal statuses.")
    set_row(child_list, "Current limit", "The child list and its photograph are implemented. The relationship has no historical version, so the product still cannot answer what a parent could see on an earlier date.")

    child_record = table_starting(doc, "FR-004 | Implemented")
    set_row(child_record, "Current evidence", "The mounted parent child response serves identity, class, enrolment history and photo_url. The signed photograph route admits the linked active guardian without granting staff access or widening branch scope.")
    set_row(child_record, "Acceptance", "A correction made by the school is visible without a second write. A linked guardian can load the on-roll child's photograph, while an unrelated, cross-school or off-roll relationship is refused. A teacher-parent may load their own child's photograph at another branch, and their staff grants still decide every other pupil.")
    set_row(child_record, "Current limit", "The child photograph is implemented. General student documents remain outside this requirement and staff-only under Module 11.")

    documents = table_starting(doc, "FR-010 | Specified")
    set_row(documents, "Current evidence", "The existing child payload and signed media route now serve the pupil's passport photograph to the linked active guardian. No consent action or general parent document list, read or response route is built.")
    set_row(documents, "Acceptance", "A linked guardian may load the on-roll child's passport photograph and no other stored file. Birth certificates, medical letters and every other document remain staff-only. A consent answer, when built, is auditable with its actor and time, and a general document link must remain authorised and expiring.")
    set_row(documents, "Current limit", "Passport-photograph read is implemented. Consent and every general document action remain specified, including any future rule by which the school chooses a document to share.")

    insert_contract(
        doc,
        "9. Needs Attention",
        "Current Contracts at 28 October 2026",
        [
            "The existing child payload carries photo_url, and the signed media route now admits an active guardian linked to that on-roll pupil. An unrelated guardian, a guardian from another school, and the guardian of a pupil outside the portal statuses are refused.",
            "The photograph is the only pupil file a guardian may read. Birth certificates, medical letters and every other document remain staff-only. Consent and general document actions remain unbuilt.",
            "Guardian and staff authority are independent. A teacher-parent reads their own child's photograph through the guardian link even when the child is at another branch, while staff grants and branch scope still decide every other pupil. No field or Field Access switch changed.",
        ],
    )

    mrd_tools.set_lines(doc.tables[29].rows[0].cells[0], [
        "OPEN - THE PARENT BACKEND IS LIVE; PRODUCT SURFACES REMAIN PARTIAL",
        "• The mounted parent backend serves child reads, fees and payments, notices, profile edits, messages and the child's passport photograph.",
        "• Consent actions, general document actions, notification preferences, Assignments and the unread-messages home block remain incomplete.",
        "• No native parent app calls phone registration or version policy yet, no versions are configured, the deletion screen is absent, and real Firebase and transfer delivery remain unproved.",
    ])
    mrd_tools.set_lines(doc.tables[30].rows[0].cells[0], [
        "FURTHER GAPS",
        "• Guardian access follows the current StudentGuardian link and has no history, so the platform cannot reconstruct what a parent could see on an earlier date.",
        "• Consent and general document sharing still need their own product rules and routes. The photograph exception does not make other stored files parent-readable.",
        "• Assignments still have no owning module, and the unread-messages home block and notification preferences remain incomplete.",
        "• The document has no evidence of a completed parent frontend or production adoption.",
    ])

    trace = next(p for p in doc.paragraphs if p.text.startswith("Module 28 carries 10 capability entries"))
    mrd_tools.retitle(
        trace,
        "Module 28 carries 10 capability entries in MRD v2.115. Core parent backend reads, payments, notices, profile, messaging, sessions and the child photograph are implemented. Consent, general document actions, notification preferences, Assignments and the unread-messages home block remain incomplete. The capability count and MRD roadmap meaning do not change in this patch.",
    )
    log_change(
        doc,
        "1.5.1",
        "Records linked-guardian access to the child's passport photograph, keeps every other document staff-only, and closes the teacher-parent cross-role test gap. MRD v2.115 is unchanged.",
    )
    trim_trailing_empty_paragraphs(doc)
    repair_ooxml(doc)
    finish(
        doc,
        M28_TARGET,
        "XVS M28 Parent Portal Functional Requirements Document v1.5.1",
        "1.5.1",
    )


def main() -> None:
    patch_m11()
    patch_m28()


if __name__ == "__main__":
    main()
