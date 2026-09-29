#!/usr/bin/env python3
"""Cut M20 v1.6: refunds and write-offs open on either key, and the empty route.

What the committed backend does, and therefore what the document must say:

* The refunds-and-write-offs list (``GET /finance/ar-adjustments/``) opens on
  ``finance.refund.view`` or ``finance.writeoff.view``. Each kind of row, and each
  figure drawn from it, goes only to a holder of that kind's key; ``kinds`` names
  what was sent, and refundable credit counts the reader's branches only
  (dcc692e2). It becomes FR-012.
* A tenant's books arrive holding one approval route per adjustment type with no
  steps in it and no approver group (3aab9000). The default ladder, with its
  threshold, is published only for a tenant that asks for it.
* The set-of-books picker answers any finance key, so a reader holding only an
  adjustment key reaches their screen (415953a7).
* The finance dashboard reads this module's records behind their own keys
  (fb3ccf16, 2371f364), and a posted credit or debit note emails the customer
  under the note's branch's notification settings (98e98798).
* No school-layer write port exists for concessions (the todo's Undone item 1).

Corrections folded in, each found by reading the code: a refund debits the
customer-credit liability, not the receivable; an unapplied credit note is held
as customer credit; there is no DebitNote model; the credit-note and concession
submit routes were missing from the API tables; the final approval posts the
adjustment itself, and a void leaves it REVERSED; a write-off has no void; the
finance submit routes take no confirmation.

A refund or write-off batch confirms like a single post against an empty route,
records one "posted without approval" per document, confirms every document
before posting any, and resolves its lines within the caller's branch reach
(8cd97cae).

The finance pass (merged at 32194627): a refund names its bank account within
the caller's branches and pays from its own branch's account or a school-wide
one, each batch line checked on its own, and refund-availability rows carry
branch_id (D57: 61836114, 1fb5bec5, 07baab0a); the list counts pending from
the dashboard's own definition, and the approvals block counts credit and debit
notes (D59: 63b3ce27, d6260eec). Section 9 drops the two gaps D59 closes.

The MRD is not revised here.

    python tools/patch_adjustment_readers_and_empty_routes_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    insert_row_after,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    normalise_change_log,
    repair_ooxml,
    row,
    row_starting,
    set_value,
)

import patch_record_history_docs

REVIEW_DATE = "28 September 2026"
LOG_DATE = "28 Sep 2026"
CODE_BASELINE = (
    "Backend main at 6d3251c6, 28 September 2026, carrying the batch fix (8cd97cae) and the "
    "finance pass (merge 32194627: 61836114, 1fb5bec5, 07baab0a, 63b3ce27, d6260eec), "
    "committed code only"
)
SOURCE_MRD = "XVS Module Requirements Document v2.93"

M20_DIR = "20-adjustments-and-concessions"
M20_STEM = "XVS_M20_Adjustments_and_Concessions_Functional_Requirements_Document"
M20_SOURCE, M20_TARGET = "1.5", "1.6"

#: The fill and text colour a status FRD gives an "Implemented with limits" header.
LIMITS_FILL, LIMITS_TEXT = "FDF3D6", "94670A"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

M20_SUMMARY = (
    "Minor revision. Adds FR-012: the refunds-and-write-offs list opens on either "
    "finance.refund.view or finance.writeoff.view, each kind of row and each figure drawn "
    "from it goes only to a holder of that kind's key, kinds names what was sent, and "
    "refundable credit counts the reader's branches only (dcc692e2). Books now arrive with "
    "an empty approval route per adjustment type and no approver group, so every post is "
    "refused until the school adds steps or somebody confirms; the default ladder and its "
    "50,000.00 threshold are published only on request, and a school can set each type's "
    "own conditions, so the one-threshold gap is withdrawn (3aab9000). Section 7 records "
    "that the set-of-books picker answers any finance key (415953a7); Section 8 records the "
    "dashboard blocks drawn from this module (fb3ccf16, 2371f364) and the note emails sent "
    "under the note's branch settings (98e98798); Section 9 records the missing FAL write "
    "port for concessions. CORRECTIONS: a refund debits the customer-credit liability and "
    "does not restore a receivable; an unapplied credit note is held as customer credit; a "
    "debit note is a credit note of kind DEBIT, not a model of its own; the credit-note and "
    "concession submit routes were missing from Section 7; the final approval posts the "
    "adjustment itself, and a void leaves it REVERSED; a write-off has no void; the finance "
    "submit routes take no confirmation. FR-007 becomes Implemented with limits. A refund or "
    "write-off batch confirms like a single post against an empty route, records one "
    "posted-without-approval entry per document, confirms every document before posting "
    "any, and resolves its lines within the caller's branch reach, answering 404 for "
    "another branch's customer or invoice (8cd97cae); FR-006 and FR-008 say so. A refund "
    "names its bank account within the caller's branches, another branch's answering 404 "
    "as an unknown one does, and is paid only from its own branch's account or a school-wide "
    "one, a 400 naming the branch, checked per line in a batch; refund-availability rows "
    "carry the customer's branch_id (61836114, 1fb5bec5, 07baab0a; FR-004, FR-006, Section "
    "7). The list counts pending as the receivables dashboard does, drafts and documents "
    "awaiting approval from one definition, so a voided or cancelled refund is not pending, "
    "and the dashboard's approvals waiting on the reader count credit and debit notes "
    "(63b3ce27, d6260eec; FR-012, Section 8); Section 9 drops both gaps. Verified by reading "
    "the committed code at 6d3251c6; no suite was run for this revision."
)


def set_limits(table) -> None:
    """Mark a requirement Implemented with limits, recolouring its header as the family does."""
    fr = table.rows[0].cells[0].text.split("|")[0].strip()
    header = table.rows[0]
    for shd in header._tr.iter(W + "shd"):
        shd.set(W + "fill", LIMITS_FILL)
    for colour in header._tr.iter(W + "color"):
        colour.set(W + "val", LIMITS_TEXT)
    for cell in header.cells:
        keep_format(cell, f"{fr} | Implemented with limits")


def set_box(cell, lines: list[str]) -> None:
    mrd_tools.set_lines(cell, lines)


def box_cell(doc, heading: str):
    hits = [r.cells[0] for t in doc.tables for r in t.rows
            if r.cells[0].text.startswith(heading)]
    if len(hits) != 1:
        raise ValueError(f"{heading!r} heads {len(hits)} boxes")
    return hits[0]


def route_table(doc, route: str):
    """The one API table, of the two headed 'Method and path', holding a row for ``route``."""
    hits = [t for t in doc.tables
            if t.rows[0].cells[0].text.strip() == "Method and path"
            and any(r.cells[0].text.strip().startswith(route) for r in t.rows)]
    if len(hits) != 1:
        raise ValueError(f"{route!r} is in {len(hits)} route tables")
    return hits[0]


def keep_fr_rows_whole(doc) -> None:
    """Stop a requirement row breaking across a page, or a requirement ending one.

    FR-008's evidence and acceptance rows are long enough that a page break lands
    inside one; a row that cannot split moves to the next page whole instead. A
    requirement's status bar and Requirement row keep with the evidence below
    them, so a page never ends on a requirement's heading and first line alone, and
    a short requirement moves whole rather than leaving its last rows to the next page.
    """
    for table in doc.tables:
        if not table.rows[0].cells[0].text.strip().startswith("FR-"):
            continue
        for r in table.rows:
            props = r._tr.get_or_add_trPr()
            if props.find(W + "cantSplit") is None:
                props.append(props.makeelement(W + "cantSplit", {}))
        short = sum(len(c.text) for r in table.rows for c in r.cells) < 1500
        for r in (table.rows[:-1] if short else table.rows[:2]):
            for cell in r.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def keep_refund_routes_opening_together(doc) -> None:
    """Section 7.2's header row and first two routes stay on one page.

    Without it the sub-heading, the header row and the first route can end a
    page on their own while the rest of the table starts the next.
    """
    table = route_table(doc, "GET /refunds/availability/")
    for r in table.rows[:2]:
        for cell in r.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def keep_boxes_with_their_heading(doc) -> None:
    """Keep Sections 7, 9 and 10's lead lines on the page their first box or table starts.

    The control box is tall enough that it moves to the next page whole, and without
    this the heading and its one-line lead are left alone at the foot of the page.
    """
    lead = next(p for p in doc.paragraphs
                if p.text.strip() == "These are current risks and gaps, not history.")
    lead.paragraph_format.keep_with_next = True
    # Section 7's and Section 10's lead paragraphs travel with the table below them.
    for start in ("All routes are mounted under /v1/finance/",
                  "Module 20 carries 14 capability entries"):
        lead = next(p for p in doc.paragraphs if p.text.strip().startswith(start))
        lead.paragraph_format.keep_with_next = True
    # The typed errors never leave their last row alone under a repeated header.
    for r in table_headed(doc, "Condition", "Answer").rows[-3:-1]:
        for cell in r.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def start_change_log_on_its_own_page(doc) -> None:
    """Begin the change log on a fresh page and never split one of its rows.

    A version summary is a long single row; left to break where it falls it
    splits across two pages, and made unsplittable it would leave the heading and
    the header row stranded at the foot of the previous page.
    """
    from patch_record_history_docs import change_log_table

    heading = next(p for p in doc.paragraphs if p.text.strip() == "11. Change Log")
    heading.paragraph_format.page_break_before = True
    for r in change_log_table(doc).rows:
        props = r._tr.get_or_add_trPr()
        if props.find(W + "cantSplit") is None:
            props.append(props.makeelement(W + "cantSplit", {}))


def space_before_last_fr(doc) -> None:
    """Put the blank paragraph that separates requirements above the one just added."""
    frs = [t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith("FR-")]
    spacer = frs[0]._tbl.getnext()
    if spacer.tag != W + "p" or Paragraph(spacer, doc).text.strip():
        raise ValueError("The first requirement is not followed by a blank paragraph")
    heading = frs[-1]._tbl.getprevious()
    heading.addprevious(copy.deepcopy(spacer))


def patch_m20() -> None:
    require_newest(str(ROOT / "functional-requirements" / M20_DIR / f"{M20_STEM}_v*.docx"),
                   M20_SOURCE)
    doc = Document(str(frd_path(M20_DIR, M20_STEM, M20_SOURCE)))
    set_cover_version(doc, M20_SOURCE, M20_TARGET)
    set_control(doc, "Version", M20_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", SOURCE_MRD)
    set_control(doc, "Supporting apps",
                "vs_workflow, vs_rbac, vs_audit, vs_exports, vs_notifications, vs_tenants")

    # ── 2.1 The five postings ────────────────────────────────────────────────
    postings = table_headed(doc, "Adjustment", "Journal", "Effect on the invoice")
    credit = row(postings, "Credit note")
    keep_format(credit.cells[1],
                "Dr revenue or returns, Dr output tax; Cr receivable control for what is "
                "applied, Cr customer credit (a liability) for the rest.")
    keep_format(credit.cells[2],
                "Applied portion raises the credited amount; the rest is held as customer "
                "credit, so the receivable never carries a credit balance.")
    refund = row(postings, "Refund")
    keep_format(refund.cells[1], "Dr customer credit (a liability); Cr bank.")
    keep_format(refund.cells[2],
                "None. It pays out credit an over-payment or an unapplied credit note left "
                "with the customer; no invoice is reopened.")

    # ── 3. Actors ────────────────────────────────────────────────────────────
    actors = table_headed(doc, "Actor", "May do", "Governed by")
    officer = row(actors, "Refund officer")
    keep_format(officer.cells[1],
                "Check availability, raise a refund, and post it directly or submit it for "
                "approval.")
    keep_format(officer.cells[2],
                "finance.refund.* (checking availability needs finance.refund.create). "
                "Approving is whoever the school's own stages name.")
    writeoff = row(actors, "Write-off authoriser")
    keep_format(writeoff.cells[1],
                "Raise a write-off, and post it directly or submit it for approval; the final "
                "approval posts it.")
    keep_format(writeoff.cells[2],
                "finance.writeoff.*, and finance.invoice.writeoff for the direct invoice "
                "route. Approving is whoever the school's own stages name.")
    insert_row_after(writeoff, [
        "Adjustments reader",
        "Open the refunds-and-write-offs list, seeing only the kinds their keys allow "
        "(FR-012).",
        "finance.refund.view or finance.writeoff.view",
    ])
    set_box(box_cell(doc, "WHERE THE THRESHOLD SITS"), [
        "WHERE APPROVAL COMES FROM, AND WHY THE DEFAULT THRESHOLD IS LOW",
        "• A school's books arrive with an approval route for each of the four adjustment "
        "types, no steps in it, and no approver group. Until the school adds steps, every "
        "direct post of a refund, write-off, concession or credit note, at any amount, is "
        "refused unless somebody confirms it, and the confirmation is recorded against them. "
        "Who approves is the school's own answer, read from the organogram it builds "
        "(Module 7).",
        "• The default ladder is published only for a tenant that asks for it "
        "(seed_finance_approvals). It gates refunds and write-offs at any size: one moves "
        "cash out, the other concedes income, and neither has a size at which a second pair "
        "of eyes stops being worth it.",
        "• It gates concessions and credit notes at or above a threshold, defaulting to "
        "50,000.00. That is far below procurement's 500,000.00 senior bar on purpose: a "
        "purchase at 400,000.00 still buys the entity something, a waiver at 400,000.00 is "
        "income given away, and a term's fees can sit under procurement's bar. Set it to zero "
        "to approve every one. A school that builds its own stages sets its own conditions "
        "for each type.",
    ])

    # ── 4. Functional requirements ───────────────────────────────────────────
    fr001 = fr_table(doc, "FR-001")
    set_value(fr001, "Current evidence",
              "A credit note debits revenue or returns and the output tax. What is applied "
              "to invoices at posting credits the receivable control, a non-cash settlement "
              "that raises the credited amount; the rest is held as customer credit, a "
              "liability, so the receivable never carries a credit balance, and it can be "
              "applied later or refunded. Posting a credit or a debit note emails the "
              "customer at their billing address, best effort, so a failed delivery never "
              "undoes the posting; the note's branch decides whose notification settings "
              "apply, then the school's (Module 8).")

    fr002 = fr_table(doc, "FR-002")
    append_to(fr002, "Current evidence",
              " A debit note is a credit note of kind DEBIT, numbered DRN, raised and posted "
              "through the credit-note routes; receipts settle it as they settle an invoice.")

    fr003 = fr_table(doc, "FR-003")
    edit_value(fr003, "Current evidence",
               "A school's scholarship is simply a concession with a kind of SCHOLARSHIP.",
               "The kinds are DISCOUNT, WAIVER and SCHOLARSHIP; a school's bursary or "
               "scholarship is a concession of kind SCHOLARSHIP. A concession cannot exceed "
               "the invoice's outstanding balance.")

    fr004 = fr_table(doc, "FR-004")
    set_value(fr004, "Current evidence",
              "An availability endpoint answers what may be refunded. Posting a refund "
              "debits the customer-credit liability and credits bank, drawing down the "
              "credit an over-payment or an unapplied credit note left with the customer, "
              "and records which of those it drew on; no invoice is reopened. The refund "
              "carries a submit step and routes through the approval engine where a "
              "template exists.")

    fr005 = fr_table(doc, "FR-005")
    set_value(fr005, "Current limit",
              "A posted write-off cannot be voided (FR-007).")

    fr006 = fr_table(doc, "FR-006")
    set_value(fr006, "Current evidence",
              "A batch endpoint raises up to 100 refunds or write-offs on one date, as "
              "drafts, posted, or submitted for approval, in one transaction. It needs the "
              "kind's create key and the key for the chosen action; a customer or an invoice "
              "may appear once, and refund availability is measured on the batch's date. Each "
              "line resolves within the caller's branch reach, so another branch's customer or "
              "invoice answers 404 as an unknown one does, and takes its branch from its own "
              "customer or invoice. The combined list of refunds and write-offs is FR-012.")
    set_value(fr006, "Acceptance",
              "Each line posts through the same posting service as a single adjustment, so "
              "the ledger guards (period lock, chronology, availability as at the batch date) "
              "hold for every line, and the POST action refuses the whole batch when any "
              "line's ladder would apply a stage. On a route with no steps it asks for the "
              "confirmation a single post asks for: without confirm_without_approval the whole "
              "batch is refused with APPROVAL_NOT_CONFIGURED and nothing is left behind; with "
              "it, every document is confirmed before any is posted and each is recorded as "
              "posted without approval against the caller, one record per document. "
              "AdjustmentBatchConfirmationTests prove both, and that the single route refuses "
              "and records the same way; BulkRunsNarrowToTheCallersReachTests prove a line "
              "naming another branch's invoice is not found.")
    set_value(fr006, "Current limit",
              "The batch is "
              "applied to the selection it is given; the module holds no rule about who should "
              "have been included.")

    fr007 = fr_table(doc, "FR-007")
    set_limits(fr007)
    set_value(fr007, "Current evidence",
              "Credit notes, concessions, refunds and payments each carry a void action that "
              "reverses the posting rather than deleting the row, and leaves the document "
              "REVERSED. A journal a document owns cannot be reversed on its own; the journal "
              "screen sends the reader to the document's void.")
    set_value(fr007, "Current limit",
              "A write-off has no void. A posted write-off stands, and because its journal "
              "belongs to the write-off request the raw journal reversal refuses it too.")

    fr008 = fr_table(doc, "FR-008")
    edit_value(fr008, "Current evidence",
               "; the submit route refuses and records the same way.",
               ". The four finance submit routes refuse the same route with "
               "APPROVAL_NOT_CONFIGURED and take no confirmation, so an adjustment whose "
               "route has no steps reaches the ledger only through a confirmed post. The "
               "final approval posts the adjustment inside the same transaction as the "
               "decision, so a posting failure, a closed period for example, rolls the "
               "decision back and leaves the stage open.")
    edit_value(fr008, "Current evidence",
               "Refund and write-off handlers now freeze",
               "Refund and write-off handlers freeze")
    edit_value(fr008, "Current evidence",
               "and the direct invoice write-off, answer 409",
               "the direct invoice write-off, and a refund or write-off batch posted directly, "
               "answer 409")
    edit_value(fr008, "Current evidence",
               "against the person who confirmed.",
               "against the person who confirmed, one record per document in a batch.")
    edit_value(fr008, "Acceptance",
               "The gate is opt-in by template, matching the rest of the platform, but the "
               "templates are now published when a tenant's books are created rather than by "
               "a command somebody has to remember.",
               "The gate is opt-in by template, matching the rest of the platform. A tenant's "
               "books arrive with one route per adjustment type and no steps in it, and no "
               "approver group. The default ladder, with its threshold and two groups created "
               "empty, is published only for a tenant that asks for it through "
               "seed_finance_approvals; it fills in an empty route and never replaces a route "
               "holding a live step. Migration 0028 removed the groups earlier provisioning "
               "had created, with the steps naming them, sparing a group with members and one "
               "whose steps a document has run through. ProvisionedBooksCarryNoLadderTests and "
               "SeededApproverGroupCleanupTests prove both.")
    set_value(fr008, "Current limit",
              "In a school that has not built its stages, every adjustment is refused at "
              "post until somebody confirms, which is the designed state; the batch POST action "
              "asks for the same confirmation (FR-006). Under the default ladder, "
              "concessions and credit notes are gated only at or above the threshold "
              "(FR-010).")

    fr010 = fr_table(doc, "FR-010")
    edit_value(fr010, "Current evidence",
               "The seeded ladder puts the threshold on both of its stages, so no stage "
               "applies below it, and each stage names an approver group created empty rather "
               "than a role provisioning invented.",
               "The default ladder, where a tenant has asked for it, puts the threshold on "
               "both of its stages, so no stage applies below it, and each stage names an "
               "approver group created empty. A tenant's books do not carry it: they arrive "
               "with an empty route, and a school that builds its own stages sets its own "
               "conditions for each type.")
    edit_value(fr010, "Current evidence",
               "Concession, credit-note and debit-note handlers now freeze",
               "Concession, credit-note and debit-note handlers freeze")
    edit_value(fr010, "Acceptance",
               "Below the threshold the direct post proceeds",
               "Under the default ladder, below the threshold the direct post proceeds")
    set_value(fr010, "Current limit",
              "The seeding command writes one threshold for both types; a school wanting "
              "different sizes sets each type's own stages (Module 7). Because the answer "
              "depends on the document's own amount, it changes when the amount does and "
              "must be re-read after an edit rather than cached against a document id.")

    add_fr(doc, "FR-012 Read Refunds and Write-offs Together, Each Kind to Its Readers",
           "FR-012 | Implemented", [
               ("Requirement",
                "A person who handles only one kind of adjustment must reach the combined "
                "list, and must see only that kind."),
               ("Current evidence",
                "GET /ar-adjustments/ opens on finance.refund.view or finance.writeoff.view. "
                "Refund rows and the refundable-credit figure go only to a holder of the "
                "refund key; write-off rows, posted and not yet posted, and the written-off "
                "this year figure go only to a holder of the write-off key. A figure the "
                "reader may not see is null, pending counts the kinds they see, and only drafts "
                "and documents awaiting approval, the one definition of pending "
                "(PENDING_STATUSES in vs_finance/constants.py) the receivables dashboard also "
                "counts from, so a voided or cancelled refund is finished business on both "
                "screens; kinds "
                "names the row kinds sent (REFUND, WRITEOFF). Every row and figure is "
                "narrowed to the reader's branches plus school-wide rows, and refundable "
                "credit counts only active customers the reader can reach. Rows carry the "
                "approval requirement (FR-011) and filter on ?type=refund or writeoff and "
                "?search=."),
               ("Acceptance",
                "A write-off holder at Ikeja Branch opens the list and sees Ikeja's and the "
                "school-wide write-offs, no refunds, and a null refundable credit. "
                "AdjustmentListAccessTests prove that a reader with neither key is refused, a "
                "write-off holder sees write-offs only, a refund holder sees refunds only, "
                "and a branch write-off holder never sees another branch's write-offs. "
                "AdjustmentPendingCountTests prove voided and cancelled documents are not "
                "pending and that the list and the dashboard count the same documents."),
               ("Current limit",
                "The list is merged in memory from the newest 1,000 refunds, 1,000 posted "
                "write-offs and 1,000 write-off requests not yet posted, so the oldest rows of "
                "a long history fall off the end."),
           ])
    fr004 = fr_table(doc, "FR-004")
    append_to(fr004, "Current evidence",
              " The bank account a refund is paid from resolves among the caller's branches' "
              "accounts and the school-wide ones, so another branch's account answers 404 as "
              "an unknown one does, and is resolved before the customer's credit is measured, "
              "so a refused account is refused first. A refund belongs to its customer's "
              "branch and is paid only from that branch's account or a school-wide one: "
              "another branch's account is a 400 naming the branch, \"This refund belongs to "
              "Ikeja Branch. Pay it from an Ikeja Branch account or a school-wide one.\" Each "
              "availability row carries its customer's branch_id, which the refund inherits, "
              "so a screen offers only the accounts the refund may be paid from (Module 19, "
              "FR-020).")
    append_to(fr004, "Acceptance",
              " In Module 19's tests_bank_account_reach.py, "
              "test_a_refund_names_only_a_reachable_account and "
              "test_a_refund_by_name_is_refused_the_same_way prove another branch's account is "
              "unknown by id and by name, DocumentPaidFromItsOwnBranchTests that a caller "
              "covering two branches cannot pay one's refund from the other's account, and "
              "test_a_school_wide_refund_may_use_any_account_she_reaches the school-wide case.")
    fr006 = fr_table(doc, "FR-006")
    append_to(fr006, "Current evidence",
              " A refund batch names one bank account, resolved within the caller's reach, "
              "and checks it against each line's own branch, since a batch may span branches: "
              "a line whose customer belongs to another branch than the account is refused "
              "400 on that line, and nothing posts.")
    space_before_last_fr(doc)
    keep_fr_rows_whole(doc)
    keep_refund_routes_opening_together(doc)

    # ── 5. Workflows ─────────────────────────────────────────────────────────
    life = table_headed(doc, "State", "Entered when", "Leaves to")
    draft = row(life, "DRAFT")
    keep_format(draft.cells[1],
                "Raised, not yet posted; also where a rejected or returned adjustment comes "
                "back to.")
    keep_format(draft.cells[2],
                "PENDING_APPROVAL on submission, or POSTED directly where no stage applies, "
                "after a confirmation where the route has no steps.")
    pending = row(life, "PENDING")
    keep_format(pending.cells[0], "PENDING_APPROVAL")
    keep_format(pending.cells[2],
                "POSTED by the final approval itself, in the same transaction as the "
                "decision. A rejection or a return sends it back to DRAFT, and a reversed "
                "decision brings it back to PENDING_APPROVAL.")
    posted = row(life, "POSTED")
    keep_format(posted.cells[2],
                "REVERSED by a void, except a write-off, which has none. The approval behind "
                "it can no longer be reversed.")
    void = row(life, "VOID")
    keep_format(void.cells[0], "REVERSED")
    keep_format(void.cells[1],
                "Voided: a reversing journal on its own date, the original kept.")

    gated = table_headed(doc, "Adjustment", "Approval", "Why")
    for label, text in (
        ("Refund", "Default ladder: every one."),
        ("Write-off", "Default ladder: every one."),
        ("Credit note", "Default ladder: at or above the threshold."),
        ("Debit note", "Default ladder: at or above the threshold, as a credit note."),
        ("Concession", "Default ladder: at or above the threshold."),
    ):
        keep_format(row(gated, label).cells[1], text)

    # ── 6. Data model ────────────────────────────────────────────────────────
    model = table_headed(doc, "Model", "Holds", "Notes")
    note = row(model, "CreditNote")
    keep_format(note.cells[1],
                "A credit or a debit note, by kind (CREDIT or DEBIT), numbered CRN or DRN, "
                "with its lines.")
    keep_format(note.cells[2],
                "Applying a credit note raises the invoice's credited amount; a debit note is "
                "never applied against another invoice.")
    debit = row(model, "DebitNote")
    keep_format(debit.cells[0], "DebitNoteAllocation")
    keep_format(debit.cells[1], "A receipt's settlement of a debit note.")
    keep_format(debit.cells[2],
                "A debit note is a CreditNote of kind DEBIT, not a model of its own.")
    keep_format(row(model, "Concession").cells[1],
                "A discount, waiver or scholarship, with its kind; a bursary is a "
                "SCHOLARSHIP.")
    keep_format(row(model, "Refund").cells[2],
                "Draws down customer credit; RefundAllocation records which receipts or "
                "credit notes it drew on.")
    wo = row(model, "WriteOff")
    keep_format(wo.cells[0], "WriteOffRequest")
    keep_format(wo.cells[1], "A conceded debt against one invoice, and its approval state.")
    keep_format(wo.cells[2],
                "Posts to bad-debt expense. The list reads posted write-offs from the finance "
                "audit trail.")

    # ── 7. API ───────────────────────────────────────────────────────────────
    edit_paragraph(doc, "All routes are mounted under /v1/finance/",
                   "require an asserted entity.",
                   "require an asserted entity. The set-of-books picker every finance screen "
                   "opens with answers any finance key (Module 19), so a reader holding only an "
                   "adjustment key, such as finance.concession.view, reaches their screen.")
    notes = route_table(doc, "GET, POST /credit-notes/")
    first = row_starting(notes, "GET, POST /credit-notes/")
    keep_format(first.cells[1], "Raise and read credit and debit notes, by kind.")
    cn = row_starting(notes, "POST /credit-notes/{id}/post/")
    keep_format(cn.cells[0], "POST /credit-notes/{id}/submit/, /post/, /allocate/, /void/")
    keep_format(cn.cells[1], "Approval hand-off, post, apply to invoices, reverse.")
    cc = row_starting(notes, "POST /concessions/{id}/post/")
    keep_format(cc.cells[0], "POST /concessions/{id}/submit/, /post/, /void/")
    keep_format(cc.cells[1], "Approval hand-off, post and reverse.")

    routes = route_table(doc, "GET /refunds/availability/")
    keep_format(row(routes, "GET /refunds/availability/").cells[1],
                "What may be refunded, and why. Needs finance.refund.create.")
    keep_format(row(routes, "POST /write-offs/{id}/submit/, /post/").cells[1],
                "Approval hand-off and posting. There is no void.")
    keep_format(row(routes, "POST /invoices/{id}/write-off/").cells[1],
                "Concede one invoice directly (finance.invoice.writeoff): raises a write-off "
                "request, and submits it where a stage applies or posts it.")
    keep_format(row(routes, "POST /ar-adjustments/batch/").cells[1],
                "Raise up to 100 refunds or write-offs on one date, as drafts, posted or "
                "submitted (FR-006).")
    keep_format(row(routes, "GET /ar-adjustments/").cells[1],
                "Refunds and write-offs in one list, on either view key, each kind to its own "
                "readers (FR-012).")

    errors = table_headed(doc, "Condition", "Answer")
    unconfigured = row_starting(errors, "Posting an adjustment whose resolved ladder has no steps")
    submit_row = insert_row_after(unconfigured, [
        "Submitting an adjustment whose route has no steps",
        "409 APPROVAL_NOT_CONFIGURED. The finance submit routes take no confirmation; the "
        "confirmed post is the way through.",
    ])
    list_row = insert_row_after(submit_row, [
        "Reading the refunds-and-write-offs list with neither view key",
        "403.",
    ])
    bank_row = insert_row_after(list_row, [
        "A refund naming a bank account that is unknown or outside the caller's branches",
        "404, identical for both.",
    ])
    insert_row_after(bank_row, [
        "A refund, or a refund batch line, paid from another branch's bank account",
        "400 naming the branch, on the batch line where it is one: \"This refund belongs to "
        "Ikeja Branch. Pay it from an Ikeja Branch account or a school-wide one.\"",
    ])

    # ── 8. Dependencies ──────────────────────────────────────────────────────
    deps = table_headed(doc, "Dependency", "Contract")
    append_to(deps, "Module 19, Finance & Accounting",
              " Its finance dashboard reads this module's records behind their own keys and "
              "the reader's branches: posted concessions by kind (finance.concession.view), "
              "and credit and debit notes, refunds and write-offs, each on its own view key; "
              "its cash movement counts refund payouts as a step of their own. Its approvals "
              "waiting on the reader count refunds, write-offs, concessions and credit and "
              "debit notes, the notes under one label, and its pending adjustments are counted "
              "from the same definition as this module's list (FR-012).")
    keep_format(row(deps, "Module 7, Workflow & Approval Engine").cells[1],
                "Holds each tenant's approval route per adjustment type, published empty with "
                "the books, and runs whatever stages the school puts in it. It refuses a route "
                "with no steps at submission, asks this module before reversing a decision, "
                "and lets a tenant's named Dynamic Role route a refund on its method.")
    append_to(deps, "Module 4, Roles & Permissions",
              " finance.writeoff.view opens the refunds-and-write-offs list on its own.")
    insert_row_after(row(deps, "Module 5, Audit"), [
        "Module 8, Notifications & Delivery",
        "Delivers the credit-note and debit-note emails under the settings of the note's "
        "branch, then the school's.",
    ])

    # ── 9. Needs Attention ───────────────────────────────────────────────────
    set_box(box_cell(doc, "CURRENT ADJUSTMENT CONTROL"), [
        "CURRENT ADJUSTMENT CONTROL",
        "• A school's books arrive with an empty route per adjustment type, so every refund, "
        "write-off, concession and credit note is refused at post until the school adds steps "
        "or somebody confirms. Under the default ladder a tenant can ask for, refunds and "
        "write-offs are gated at any size and a concession or credit note at or above the "
        "threshold, below which it posts on one person's permission.",
        "• A ladder with no steps is approval undecided, not approval-free. Every direct post "
        "route, the batch included, refuses it with APPROVAL_NOT_CONFIGURED, and a confirmation lets the "
        "adjustment through only with the person's name recorded against it; the submit "
        "routes refuse it and take no confirmation.",
        "• An approval decision cannot be withdrawn once the adjustment has posted; the "
        "correction is a void, which leaves both postings visible. A write-off has no void, so "
        "a mistaken write-off cannot be undone through the platform.",
    ])
    set_box(box_cell(doc, "FURTHER GAPS"), [
        "FURTHER GAPS",
        "• Approval is opt-in by template even for refunds and write-offs: where no template "
        "resolves at all, neither a tenant route nor a shared row, they post directly with no "
        "confirmation. Books created since routes are published with them always resolve one.",
        "• Ladder seeding never replaces a route holding a live step, so a correction to the "
        "default shape reaches existing tenants only through a migration.",
        "• Refund availability is computed from the receivable state, so it cannot tell that an "
        "over-payment has not actually settled at the bank.",
        "• A batch adjustment is applied to the selection it is given; the module holds no rule "
        "about who should have been in it.",
        "• Concession and refund eligibility is not connected to the student and guardian "
        "records, so who qualifies is decided outside the platform and recorded here as a "
        "decision. Tracked against Modules 11 and 20.",
        "• There is no school-layer (FAL) write port for concessions: a school screen granting a "
        "waiver calls /v1/finance/concessions/ directly and has to know the finance customer id, "
        "which the layer exists to hide. It waits, on purpose, for the fees design to say what a "
        "waiver screen needs.",
        "• The approval requirement a read reports does not signal a ladder with no steps. Such "
        "a row reads as needing no approval, and its post is still refused until somebody "
        "confirms, so a client choosing between Post and Submit from that field alone meets a "
        "refusal it was not warned of.",
    ])

    # ── 10. Traceability ─────────────────────────────────────────────────────
    edit_paragraph(doc, "Module 20 carries 14 capability entries",
                   "in MRD v2.76", "in MRD v2.93")
    trace = table_headed(doc, "MRD capability", "Requirements")
    keep_format(row(trace, "Adjustment summaries and export").cells[1],
                "FR-006, FR-009, FR-012")

    patch_record_history_docs.REVIEW_DATE = LOG_DATE
    log_change(doc, M20_TARGET, M20_SUMMARY)
    patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

    assert_absent_outside_log(
        doc,
        "Dr receivable control; Cr bank",
        "restoring the receivable",
        "the submit route refuses and records the same way",
        "The seeded ladder puts",
        "cannot yet be gated at different sizes",
        "post it once approved",
        "plus membership of the approver group",
        "templates are now published",
        "so a batch cannot bypass",
        "A batch skips the confirmation",
        "does not ask for the confirmation",
        "MRD v2.76",
        "a voided one included",
        "but not credit or debit notes",
    )
    repair_ooxml(doc)
    normalise_change_log(doc)
    keep_boxes_with_their_heading(doc)
    start_change_log_on_its_own_page(doc)
    finish(doc, frd_path(M20_DIR, M20_STEM, M20_TARGET),
           f"{M20_STEM.replace('_', ' ')} v{M20_TARGET}", M20_TARGET)


def main() -> None:
    patch_m20()


if __name__ == "__main__":
    main()
