#!/usr/bin/env python3
"""Cut M04 FRD v1.29: every change queued against Module 4 since 14 September.

The Documents owed queue in todo.md held fifteen entries naming this module
that no revision had carried. Each was checked against the code at 7f3756e2
before being written, and the two that followed (2e446e27, a8e0e4d0) against
a8e0e4d0, and the document now says what the code does:

* Field Access as built. The registry of restrictable fields, filed under the
  same module and resource tree as permissions (FR-029); each role's Read and
  Write switch on each field, with one-person exceptions read as at a day
  (FR-030); and enforcement on every surface, which replaced the serializer
  mixin and the view_sensitive keys FR-022 still described (FR-022).
* Shared reference reads that open on module membership, and what the keys
  that used to open them still decide (FR-031).
* The branch-bound write rule, for role grants and for shared rows, and the
  narrowed holder list (FR-032), with the bank reach the finance pass added
  under it: bank accounts and bank ledgers named on a write stay within the
  caller's branches, and a branch's document is paid from its own branch
  (D57: 61836114, 966e5770, 1fb5bec5).
* Dataset import keys, school.staff.import among them, standing in for the
  wizard's engine keys, and past the wizard only for what a dataset declares:
  the bank-statement key's rollback of its own statements (FR-033, a8e0e4d0).
* The sweep that takes a reclassified key back from every tenant surface, the
  provisioning copy that drops such a key, and the scope audit's leftover-grant
  question (FR-002, FR-021; the P1 school-creation item leaves).
* The read refusal beside the write one: FieldAccessMixin.can_read_field and
  403 field_read_denied, for a route whose whole response is one field the
  caller may not read, such as an import file download (FR-022, 2e446e27).
* The RBAC audit refusing an action type the central trail does not register
  (FR-019), and permission exceptions read as at a day (FR-008).
* A dashboard block asks the school's plan as well as the reader's role
  (4af5d331, part of the finance pass), so no gap is recorded for it.

Two entries are stale in detail and are written as the code stands: the field
access update keys are ``.update``, not ``.manage``, since the manage action
was retired, and the conversion ended twelve write abilities, the number the
entry's own list adds up to, not ten.

The traceability follows MRD v2.94, which widens the Field Access entry to
every registered field, by role switch and one-person exception, so the entry
and what it maps agree; the reconciliation box says so. The permission-group and
role-assignment rows take the MRD's names for those entries too.

    python tools/patch_field_access_and_branch_reach_docs.py
"""
from __future__ import annotations

from docx import Document

import patch_record_history_docs
from patch_document_type_labels_docs import require_newest
from patch_backend_owned_permission_registry_docs import keep_rows_whole
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
    table_with_header,
    update_reconciliation,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    edit_box,
    insert_row_after,
    row_labelled,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    insert_after,
    normalise_change_log,
    repair_ooxml,
    row,
    row_with_cell_starting,
    set_status,
    set_value,
)

REVIEW_DATE = "28 September 2026"
MRD_VERSION = "2.94"

#: The Field Access entry, as MRD v2.93 and v2.94 word it.
FIELD_ACCESS_OLD = "Field Access over whole person records, with composite names"
FIELD_ACCESS_NEW = (
    "Field Access over every registered field, by role switch and one-person exception, with "
    "composite names"
)
CODE_BASELINE = (
    "Backend main at 6d3251c6, 28 September 2026, read for every change queued against this "
    "module since 14 September, the field read refusal (2e446e27), the bank-statement key's "
    "rollback (a8e0e4d0) and the finance pass (merge 32194627: bank reach 61836114, 966e5770 and "
    "1fb5bec5; dashboard plan check 4af5d331) among them. Uncommitted work in other apps' "
    "working trees is not part "
    "of it"
)
TEST_EVIDENCE = (
    "Test evidence is what each change recorded when it was committed, and every test named "
    "here was read at the baseline; this revision re-ran no suite. Nothing here claims release "
    "or deployment."
)

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.28", "1.29"


# ── FR-022, rewritten: enforcement ───────────────────────────────────────────

FR022_REQUIREMENT = (
    "Some values are sensitive within a record that the caller is otherwise entitled to read. "
    "Hiding them in the frontend is not protection. What a person may read and change of each "
    "registered field is decided by their roles' switches (FR-030), and every surface that "
    "carries the field keeps that decision, a create path, a nested block and a response built "
    "by hand included."
)
FR022_EVIDENCE = (
    "vs_rbac.field_enforcement is the one place the answer is acted on, through five doors: "
    "FieldAccessMixin for serializers, visible for a response a view builds by hand, "
    "assert_writable for a body applied without a serializer, can_read for a block included "
    "or omitted whole, and FieldAccessMixin.can_read_field for a route that serves one field "
    "of a record on its own. A field the caller cannot read is absent, with no placeholder and no list "
    "of what was withheld. One they can read but not change is present, and a detail response "
    "names it in _read_only_fields. A submitted field they cannot write is refused with 403 "
    "field_write_denied, naming every offending field at once and in the same words for a hidden "
    "field and a read-only one, and nothing is saved. Three submissions are dropped rather than "
    "refused, because none changes anything: on update, a value equal to the one stored, and "
    "only from a caller who may read the field; on create, an empty or null value, which a form "
    "posts for every input its user never touched; and on create, any value of a field declared "
    "open_on_create, which is how a pupil's enrolment date, a guardian's phone and a staff "
    "member's email are still set by whoever creates the record. The caller is the request's "
    "user in the request's tenant with every role they hold counting, evaluated once per "
    "request. A render with no request passes everything; a render inside a request that acts "
    "for nobody must be declared in SYSTEM_SURFACES with its reason, and the one entry is an "
    "invoice's pay-to block, which is addressed to the payer. The sign-in response and /me carry "
    "field_access: per resource, the names hidden from the user, the names read-only to them, "
    "and the open_on_create names an Add form may still offer, each of which also sits in one of "
    "the other two lists; a field the tenant may not hold is left out. The surfaces are vendors, "
    "bank accounts, payroll runs, salaries, virtual accounts, payouts, pupils, guardians, school "
    "staff, user accounts and CX staff profiles, imports, exports, the money movements feed and "
    "approval documents, with their nested copies such as a staff record's account block and a "
    "row of the student directory. Migration 0025 converted every permission key that had "
    "guarded a field into switches on tenant roles, the prebuilt library and personal "
    "exceptions, and verify_field_access_conversion compares every user's access before and "
    "after on a copy of a production database, allowing only the differences "
    "vs_rbac.field_conversion lists. Migration 0026 then deleted the eight keys that guarded "
    "nothing but fields, school.students.view_sensitive, finance.bankaccount.view_sensitive, "
    "finance.payrollrun.view_sensitive, procurement.vendor.view_sensitive, the two payments "
    "view_sensitive keys and platform.staff_payroll.view and .manage, and the serializer mixin "
    "that read them is gone. A route whose whole response is one field, such as the download of "
    "an import batch's uploaded file, has no payload to filter, so it asks "
    "FieldAccessMixin.can_read_field, the rule the record's own serializer applies, owner rule "
    "included, and refuses with 403 field_read_denied (FieldReadDenied) in the envelope "
    "field_write_denied uses, naming the field; it asks before the file is looked up, so the "
    "refusal says nothing about whether the record holds one."
)
FR022_ACCEPTANCE = (
    "A role with Read off does not receive the field on any surface at any depth; one with Read "
    "on and Write off receives it named in _read_only_fields; a submission of either is refused "
    "403 field_write_denied and saves nothing; an unchanged echo on update, a blank on create and "
    "an open-on-create value on create are accepted (test_field_enforcement). Registry tests fail "
    "the build when a serializer emits or writes a registered name without enforcing it or "
    "without being a declared surface (EveryReadPathIsKnownTests, EveryWritePathIsKnownTests), "
    "and when a declared surface rendered for a caller with every field closed carries a "
    "registered name at any depth (EveryDeclaredSurfaceIsRenderedDeepTests and each module's "
    "deep payload test); every allowlist entry carries its reason and must still match a real "
    "hit. The write-path test was proven by negative control when it was written: removing the "
    "enrolment guard failed it on each medical field. The conversion is pinned in a school with "
    "one branch and one with two, on the prebuilt library and on personal exceptions "
    "(test_field_conversion), and the retired keys are neither seeded nor left behind "
    "(test_retired_field_keys). The uploaded file downloads while its Read switch is on and is "
    "refused field_read_denied while it is off, the detail and the download agree, and the "
    "refusal is the same whether or not a file is held (BatchRoutesFollowTheDetailsSwitchesTests)."
)
FR022_LIMIT = (
    "Write implies Read, so twelve write abilities ended at the conversion, by owner decision "
    "on 21 September 2026, each where a key gated only reading and the value was written behind "
    "the endpoint's own key: a finance bank account number, four payroll-line figures, three "
    "salary figures, an import batch file, and a CX staff member's bank name, account name and "
    "account number. A role that needs one is given sight of the field. The map on /me is built "
    "from roles alone; on an existing record _read_only_fields wins, because it carries owner "
    "rules the map cannot know, such as a member of staff reading and changing their own payroll "
    "bank details. A field nobody has registered is open, so the approved fields not yet "
    "registered are not protected (Needs Attention). The school app hides and greys fields from "
    "the map and _read_only_fields on its student and staff screens; coverage elsewhere is not "
    "evidenced here."
)


# ── FR-029: the registry and the tree ────────────────────────────────────────

FR029_HEADING = "FR-029  Name Every Restrictable Field Under Its Resource"
FR029 = [
    ("Requirement",
     "An administrator can restrict only a field the code honours, so which fields exist is the "
     "developer's declaration, never an administrator's. The fields sit in one tree with the "
     "permissions that open the record: module, then resource, then that resource's keys and "
     "fields, each under a readable name."),
    ("Current evidence",
     "Each app declares its fields in its own field_access.py and registers them from "
     "AppConfig.ready() through vs_rbac.field_registry.register_fields, naming the serializer "
     "classes that carry them as surfaces; a duplicate name or an unclassified scope stops the "
     "process at start-up. manage.py sync_field_registry writes the declarations to "
     "FieldDefinition (migration 0022), deactivates rather than deletes one that has gone, "
     "audits FIELD_REGISTRY_SYNCED, offers --check, and runs inside seed_all_permissions, which "
     "the deploy runs. A field carries its label, group, description, the api_names it travels "
     "under, sensitive (closed by default when true, open otherwise), writable (false for a value "
     "no API path can set, which then offers no write switch), open_on_create (migration 0024) "
     "and a scope with the meaning and guard of a permission's. Fields are registered on finance "
     "bank accounts, payroll runs and salaries, payment virtual accounts and payouts, procurement "
     "vendors, import templates, batches and jobs, user accounts and CX staff profiles, and "
     "pupils, guardians and school staff. A resource may carry fields and no keys: "
     "school.guardians, whose records are read and corrected under the school.students keys, is "
     "registered through FIELD_ONLY_RESOURCES, so a school is not handed a second set of keys "
     "governing nothing. PermissionModule and PermissionResource carry a readable label, read as "
     "the slug in words where blank; RESOURCE_LABELS seeds the ones that read wrongly, so "
     "school.teachers reads Staff, because the register covers the bursar and the registrar "
     "too, and the finance, procurement and payments seeds fill a blank label and never "
     "overwrite one. GET /rbac/tenants/{slug}/access-catalogue/ returns the tree, filterable by "
     "module, resource and search, and its permission entries are built by the function "
     "permission-catalogue/ uses, so flattening it gives that catalogue. A tenant that is not "
     "the platform never sees a PLATFORM permission or field in it, whoever reads it, and a "
     "platform operator holding a platform role-view key may assert a school to read that "
     "school's tree. GET /rbac/vision/fields/ lists the whole registry, active or not, "
     "read-only."),
    ("Acceptance",
     "Every field a serializer guards is declared, every declared name reaches a serializer, and "
     "a field's scope equals its guarding key's (RegistryMatchesTheSerializersTests, "
     "RegistryScopeMatchesTheGuardingKeyTests). A bad declaration is refused on registration "
     "(RegistrationValidationTests), and the sync writes, deactivates and checks "
     "(SyncFieldRegistryTests). The tree matches the flat catalogue, hides platform entries from "
     "a school, answers a platform reader asserting a school, and costs the same whatever the "
     "number of resources (AccessCatalogueTests, AccessCatalogueQueryCostTests)."),
    ("Limit",
     "A field exists only once an app declares it, and the approved fields still to be declared "
     "are in Needs Attention. A declaration reaches a database only when sync_field_registry "
     "runs there, which is how a new open_on_create flag arrives. Both applications carry a "
     "Field Access screen on roles in their source; release and deployment are not evidenced "
     "here."),
]


# ── FR-030: per-role switches and one-person exceptions ──────────────────────

FR030_HEADING = "FR-030  Decide Each Field Per Role, With One-Person Exceptions"
FR030 = [
    ("Requirement",
     "A school decides, role by role, who reads and who changes each registered field, and can "
     "make a one-person exception without minting a role, under the rules a permission "
     "exception keeps. Nobody may change a value they cannot see."),
    ("Current evidence",
     "RoleFieldAccess holds one role's Read and Write on one field; a missing row is the field's "
     "default, and a reset deletes the row, so a stored row is always a decision. "
     "PrebuiltRoleFieldAccess holds the library's defaults, copied at provisioning (FR-021). "
     "UserFieldAccessOverride is an ALLOW or DENY on READ or WRITE for one person, with a "
     "required reason and lazy expiry, never on yourself, and one row per person, field and "
     "access, so a new exception replaces the old (migration 0023). A database check constraint "
     "on both switch tables stores Write only with Read, and the guards on save, clean and "
     "bulk_create refuse a field the tenant may not hold and Write on a field that is not "
     "writable. No field is restricted: a switch takes effect without approval, which is why the "
     "key that changes switches is restricted instead. vs_rbac.field_evaluator.get_field_access "
     "answers per field in this order: the Vision super admin reads and writes everything; a "
     "field the tenant may not hold is closed; an unregistered field is open; otherwise the most "
     "generous of the person's active roles wins, those being the roles the permission evaluator "
     "counts, so a role at a branch out of service stops counting here too; a personal ALLOW "
     "READ or ALLOW WRITE adds; a personal DENY beats every role and ALLOW, and DENY READ also "
     "removes Write; and Write is then limited to what is readable and writable. A cold "
     "evaluation costs three queries whatever the number of roles and fields, memoised on the "
     "user for the request and checked against the registry revision. GET and PATCH "
     "/rbac/tenants/{slug}/roles/{key}/field-access/ list and change a role's switches: 1 to 200 "
     "changes per PATCH, each field once, applied atomically with the role and its rows locked; "
     "a change that moves nothing writes nothing; each switch that moves writes one "
     "FIELD_ACCESS_CHANGED row and each reset one FIELD_ACCESS_RESET row, and role.version is "
     "bumped once. The field-access-overrides routes list, create and lift exceptions under the "
     "user_overrides and team_overrides keys that govern permission exceptions, each row showing "
     "role_state, what the person's roles alone say; creating, replacing or lifting one writes "
     "FIELD_OVERRIDE_CREATED or FIELD_OVERRIDE_LIFTED at WARNING with the person's state before "
     "and after. Every one of these audit rows records actor_holds_read, whether the "
     "administrator could read the field themselves, because changing a field you cannot read "
     "is allowed, and a failed audit rolls the change back. UserFieldAccessOverride is tracked "
     "by vs_history, so ?as_at= lists the exceptions as they stood at the end of a day, with "
     "role_state null. school.field_access.view and .update and platform.field_access.view and "
     ".update are CRITICAL, and the school pair sits at Core. The view keys are unrestricted, "
     "because seeing a switch opens nothing; the update keys are restricted, because a switch "
     "opens a field the moment it is saved, so adding update to a role goes through the "
     "role-change ladder and granting a role that carries it, by somebody who does not hold it, "
     "through the role-grant ladder. The school pair defaults to School Admin and the platform "
     "pair to the two platform administrator roles. Both routes are open before go-live, as role "
     "detail is."),
    ("Acceptance",
     "The most generous role wins, a personal DENY beats a role and an ALLOW, the super admin "
     "and unregistered fields are open, and a cold evaluation is three queries "
     "(test_field_evaluator). A switch change is refused across tenants and without the update "
     "key; a platform field is refused exactly as an unknown one; a field exception on yourself "
     "is refused under either identity of a proxy session; two administrators on one role "
     "serialise; and every change is audited and rolls back with a failed audit "
     "(CrossTenantIsolationTests, MissingKeysTests, PlatformFieldTests, SelfExceptionTests, "
     "TwoAdministratorsOnOneRoleTests, FieldAccessAuditTests). A past list keeps the view key and "
     "another school's administrator cannot read it (FieldExceptionsAsAtTests), and the "
     "restricted update key holds in a school with one branch and one with two "
     "(test_field_access_restriction)."),
    ("Limit",
     "A role's switches keep no history, so exceptions read as at a day carry no role comparison "
     "(Needs Attention). Switch changes take effect without approval by design; what needs "
     "approval is holding the key that makes them. The role switch route is not reachable across "
     "tenants: a platform operator records an exception for a school user, and does not change a "
     "school's role switches."),
]


# ── FR-031: shared reads on module membership ───────────────────────────────

FR031_HEADING = "FR-031  Open a Shared Reference Read to Anyone Working in the Module"
FR031 = [
    ("Requirement",
     "A list every screen of a module needs, such as which books the school keeps, its periods, "
     "cost centres and dimensions, or the module's dashboard, carries no privilege of its own. "
     "It opens to anybody holding any key in that module, what it returns stays scoped, and each "
     "block or row inside it still answers to its own key. A key that used to open such a list "
     "keeps whatever it still decides."),
    ("Current evidence",
     "HasAnyModuleAccess admits a caller holding any live key in the modules a view names in "
     "rbac_modules, runs the tenant surface check itself because it can be a view's only gate, "
     "and refuses a view that names none. Finance opens this way the ledger entity list (the "
     "set-of-books picker every finance screen starts from), fiscal years and periods, cost "
     "centres, which also stay open to the two requisition keys, dimensions, the posting window, "
     "the chart of accounts without balances, tax codes, currencies, and its dashboards, the "
     "overview on any finance or payments key; procurement opens its dashboards on any "
     "procurement key. Rows stay the caller's own tenant and entity, and another tenant's code "
     "answers 404. Inside a dashboard each block asks its own key and is null for a reader "
     "without it, and the branch rule decides what a narrowed reader's figures cover (Module 19 "
     "and the procurement FRD). What the older keys decide now: finance.entity.view gates no "
     "backend data and still decides who is shown the Entities setup screen, its palette action "
     "and the procurement settings link to it in both frontends, and platform roles keep it; "
     "finance.period.view, finance.costcenter.view and finance.dimension.view still gate their "
     "setup screens, and finance.period.view also the period close checklist and the "
     "dashboard's period-close block; finance.report.view no longer opens the finance dashboard "
     "and decides its ledger blocks; procurement.analytics.view no longer opens the procurement "
     "dashboard and decides its spend blocks. Creating a ledger entity, cost centre or dimension "
     "keeps its create key. An either-key list is the other shape: a view naming several keys "
     "admits a holder of any one, and GET /finance/ar-adjustments/ opens on finance.refund.view "
     "or finance.writeoff.view, sending refund rows and the refundable-credit figure only to "
     "refund-key holders and write-off rows and the written-off figure only to write-off-key "
     "holders, so finance.writeoff.view now opens a screen on its own."),
    ("Acceptance",
     "A bursar holding only finance.report.view reads the period, cost-centre and dimension "
     "options on the reports she may run, a branch admin holding only finance.concession.view "
     "reaches the set-of-books picker, and neither can create a cost centre or dimension "
     "(test_reference_list_access). Each dashboard opens to any key in its module and gives a "
     "reader without a block's key null for that block (tests_dashboard_access in finance and in "
     "procurement). A write-off holder without the refund key opens the adjustments list and "
     "sees write-off rows only (tests_adjustment_and_aging_access)."),
    ("Limit",
     "HasAnyModuleAccess is weaker than a key check by design and is used only for reads whose "
     "payload is not sensitive on its own; tenant, entity and branch scoping is the view's job, "
     "not the gate's. The plan gate does not read it (FR-025). A dashboard block asks its key of "
     "the school's plan as well as of the role, through vs_rbac.plan_gate.keys_within_plan, so a "
     "block whose key sits above the school's band is absent, as the screen behind it would "
     "refuse it (4af5d331)."),
]


# ── FR-032: branch-bound writes ──────────────────────────────────────────────

FR032_HEADING = "FR-032  Let a Branch-Bound Caller Change Only What Is Wholly Theirs"
FR032 = [
    ("Requirement",
     "A branch administrator reads what their branches share with the rest of the school, and "
     "changes only what sits wholly inside their branches. Access itself follows the same rule: "
     "they may grant, change or take away a role only where the grant's reach and the person "
     "holding it both sit inside their branches, so nobody hands out reach they do not have."),
    ("Current evidence",
     "vs_rbac.scoping.caller_may_change states the rule once: a whole-tenant caller may change "
     "any row they can see, and a branch-bound caller a row only when its branch set is "
     "non-empty and every branch in it is theirs, so a row with no branch is read-only to them. "
     "assert_caller_may_change refuses with 403 SHARED_RECORD_READ_ONLY "
     "(vs_rbac.exceptions.SharedRecordReadOnly), whose sentence names who can make the change. "
     "Staff, the academic structure, the calendar and finance budgets judge their writes through "
     "it, and their rows carry can_manage so a screen hides what the reader may not change; "
     "which rows, and which parent decides a row's branches, are each module's own. "
     "vs_rbac.grant_reach.assert_caller_may_grant applies it to role grants, for a caller inside "
     "the tenant whose scope is a set of branches: the grant's reach, its own branch or else the "
     "role's selected branches, must be non-empty and inside theirs, or it is refused with 400 "
     "under branch; and the holder, unless the holder is the caller, must be posted only to "
     "their branches, or it is refused with 400 under user. A grant naming neither reaches the "
     "whole school, so only a whole-tenant caller may write one. It runs on assignment create, "
     "update, revoke and replace and on the staff bulk role grant; a platform operator acting "
     "on a school is governed by the platform keys instead. The assignment list and detail show "
     "a branch-bound caller only grants held by people posted at their branches or school-wide, "
     "the same inclusive reading the staff directory applies. Money follows the same reach "
     "(Module 19). A bank account a posting names resolves among the caller's own branches' "
     "accounts and the school-wide ones, the bank list's rule, and another branch's account "
     "answers 404 as an unknown one does. A ledger account that backs a bank account answers "
     "to the same list wherever a write names it (vs_finance.accounts."
     "accounts_a_caller_may_name): another branch's bank ledger is refused as an unknown "
     "account, 400 on the field or 404 on the account's own edit route, while a ledger account "
     "behind no bank and a caller bound to no branch are untouched, and reads are not "
     "narrowed by it. A document that belongs to a branch is paid only from that branch's "
     "account or a school-wide one, so a caller who covers two branches still cannot pay one "
     "branch's refund from the other's account (400, naming the branch)."),
    ("Acceptance",
     "A branch admin grants at their own branch and is refused a school-wide grant, a grant at "
     "another branch, a grant to the school-wide registrar and revoking the registrar's "
     "school-wide role; their holder list shows their own people and the school-wide ones; a "
     "whole-school administrator is not narrowed (GrantReachTests). The shared-row refusal is "
     "pinned where each module applies it (test_branch_write_scope for staff, "
     "test_shared_rows_read_only for the calendar). The money rule is pinned in Module 19's "
     "tests_bank_account_reach (BankAccountNamedInAPostingTests, "
     "DocumentPaidFromItsOwnBranchTests) and tests_ledger_reach (ResolverTests, "
     "LedgerNamedOnAWriteTests)."),
    ("Limit",
     "The role list is not narrowed: a branch-bound administrator reads every role in the "
     "school (Needs Attention). The holder rule reads postings, so a person with no posting is "
     "school-wide and read-only to a branch administrator."),
]


# ── FR-033: dataset import keys ──────────────────────────────────────────────

FR033_HEADING = ("FR-033  Let a Dataset's Import Key Take Its Own File Through the Wizard, "
                 "and Past It Only Where Declared")
FR033 = [
    ("Requirement",
     "A school administrator loads their own records from a spreadsheet with that dataset's key "
     "and without the import engine's broad keys. The dataset key covers taking one file through "
     "the wizard, and nothing that unwinds or administers data already written, unless the "
     "owning module declares an engine key its dataset cannot be corrected without."),
    ("Current evidence",
     "A domain app registers its dataset's own key with the engine from AppConfig.ready() "
     "(register_dataset_import_key): school.students.import for students and guardians, "
     "school.staff.import for staff, and academics.structure.import for the academic structure "
     "and subjects; finance.bankaccount.import for bank statements is held by the engine itself, "
     "so that it stands even where finance's ready() is not reached. HasImportBatchRBACPermission, "
     "a HasRBACPermission, first asks for the engine key the view names. Failing that, the "
     "dataset key stands in where the view asks for a wizard key, import.batches.view, "
     ".create, .run or .import, import.validations.view or import.jobs.view, and only for a batch "
     "of that dataset in the asserted tenant and inside the caller's branches, so one module's "
     "key never opens another's file, and a bank statement batch must also resolve to the "
     "tenant's own books. A registration may also declare extra engine keys its dataset's key "
     "covers past the wizard (_DATASET_EXTRA_ENGINE_KEYS). Every school dataset declares none. "
     "Bank statements declare import.rollbacks.run and nothing else, because finance refuses to "
     "edit a bulk-imported statement and corrects it by rolling the import back, and a declared "
     "key counts only against a batch of its own dataset: the bank-statement key rolls back a "
     "statement batch and never a students batch, whatever other keys its holder has. The "
     "rollback history, delete, edit, issue resolution, the import audit log and import "
     "notifications need the engine's own key for every dataset, and for a school dataset so "
     "does rollback; the refusal is the same 403 whether or not the batch exists. "
     "school.staff.import is SENSITIVE and so "
     "restricted, reaching a role only through the role-change ladder; it is held by School "
     "Admin and Branch Admin by default, where school.students.import stops at School Admin, "
     "because a staff file adds people to the branch the uploader already staffs. It sits on no "
     "plan band, so every school reaches it, where the student import answers to the "
     "bulk_import capability, itself at Core."),
    ("Acceptance",
     "A holder of a school dataset's key validates, starts, reads and abandons her own batch; "
     "rollback, delete, rewriting the batch, the administrative reads and resolving an issue in "
     "place are refused, the same for a batch that does not exist; the engine's rollback key "
     "rolls back; and neither key reaches another school's batch (tests_dataset_key_actions). A "
     "staff key holder takes a staff file through the wizard without any import key "
     "(tests_branch_scope). A holder of finance.bankaccount.import alone rolls back her own "
     "statement import, is refused delete and the administrative reads, cannot roll back a "
     "students batch even holding the students key as well, and cannot reach another school's "
     "statement (tests_bank_statement_key_rollback); RegisterDatasetImportKeyTests pins that "
     "the school datasets declare nothing past the wizard and bank statements the rollback "
     "alone."),
    ("Limit",
     "Rolling a school dataset's import back is a support action on the engine's key: a school "
     "corrects its roll, staff or structure by uploading a fixed file. A dataset whose app "
     "registers no key of its own, such as calendar events, needs the engine's keys throughout."),
]


# ── Needs Attention ──────────────────────────────────────────────────────────

NA_ROLE_LIST = (
    "The role list still returns every role to a branch-bound administrator. The platform-wide "
    "narrowing has landed: procurement, finance, its statements read from the reader's own "
    "journals included, workflow, tickets, accounts, configuration and import all resolve their "
    "querysets from the grants, and this module's assignment list and detail show a "
    "branch-bound caller only the holders at their branches or school-wide (FR-032). The role "
    "catalogue is the one left: the Ikeja administrator reads every role the school has, "
    "Lekki's branch roles included."
)
NA_ROLE_LIST_FIX = (
    "Narrow the role list to roles whose reach meets the caller's branches, with the "
    "school-wide ones, or state in code why an access administrator legitimately reads the "
    "whole catalogue, so the absence is a decision rather than an omission."
)
NA_ROLE_HISTORY = (
    "A past view cannot say what a person's roles allowed that day. Field exceptions and role "
    "grants keep a record history, and a list of exceptions reads as at a day, but a role's "
    "field switches and its permissions keep none, so a past exception carries no role "
    "comparison."
)
NA_ROLE_HISTORY_FIX = (
    "Track RoleFieldAccess and TenantRolePermission in vs_history, or record here that a past "
    "view shows exceptions without the role state they departed from."
)
D13_SENTENCE = (
    " At a live school the staff create grants the school's Teacher role and takes no other "
    "(Module 12), so there it meets this refusal only where the school has given Teacher a "
    "restricted permission the person adding staff does not hold."
)


M04_SUMMARY = (
    "Minor revision, carrying every change queued against this module since 14 September. "
    "Field Access is documented as built: FR-029 adds the field registry, filed under the same "
    "module and resource tree as permissions with readable labels and fields-only resources; "
    "FR-030 adds each role's Read and Write switch on each field and one-person field "
    "exceptions, with the evaluation order, the routes, the four field access keys, their audit "
    "and exceptions read as at a day; and FR-022 is rewritten from the retired serializer mixin "
    "to the enforcement contract, in which a field the caller may not read is absent, one they "
    "may not change is named read-only, a write of either is refused 403 field_write_denied, a "
    "route whose whole response is one unreadable field, such as an import file download, is "
    "refused 403 field_read_denied through FieldAccessMixin.can_read_field (2e446e27), the "
    "old keys were converted into switches by migration 0025 and deleted by 0026, and twelve "
    "write abilities ended because write implies read. FR-031 adds shared reference reads "
    "opened on module membership, with what finance.entity.view, the three finance setup view "
    "keys, finance.report.view, procurement.analytics.view and finance.writeoff.view now decide. "
    "FR-032 adds the branch-bound write rule, for role grants and shared rows, and the narrowed "
    "holder list, and records that money follows the same reach: a bank account or bank "
    "ledger another branch keeps is refused as unknown wherever a write names it, and a "
    "branch's document is paid only from its own branch's account or a school-wide one "
    "(61836114, 966e5770, 1fb5bec5). FR-033 adds dataset import keys, school.staff.import among them, standing in "
    "for the wizard's engine keys, and past the wizard only for an engine key the dataset "
    "declares, which bank statements alone do, for the rollback of a statement import "
    "(a8e0e4d0). FR-002 and FR-021 record migration 0021's sweep of "
    "reclassified keys from every tenant surface, organization tenants included, the "
    "provisioning copy that drops such a key, and the scope audit's leftover-grant question, "
    "and FR-002 no longer names the deleted platform.staff_payroll keys. FR-019 records that the "
    "RBAC audit refuses an action type the central trail does not register; FR-008 that "
    "permission exceptions read as at a day; FR-010 the session's branch_reach. FR-025, FR-027 "
    "and FR-028 follow, with the scope, actors, withdrawal routes, data model, contracts, "
    "dependencies, verification and traceability. Needs Attention: the P1 school-creation item "
    "leaves; the unnarrowed-lists item narrows to the role list; the field-exception history "
    "item is rewritten to the role history still missing. FR-031 records that a dashboard block "
    "asks its key of the school's plan as well as of the role (4af5d331). FR-010 records that "
    "the reader's reach also decides which bank accounts and bank ledgers a write may name. "
    "The contents name FR-001 to FR-033. "
    f"MRD v{MRD_VERSION}, whose Field Access entry is widened to every registered field, by "
    "role switch and one-person exception, the wording the traceability takes, as do the "
    "permission-group and role-assignment entries. " + TEST_EVIDENCE
)


def keep_table_rows_whole(doc) -> None:
    """No row of a multi-column table breaks across a page; a boxed note may."""
    for table in doc.tables:
        if len(table.rows[0].cells) > 1:
            keep_rows_whole(table)


def keep_dependency_table_opening_together(doc) -> None:
    """Section 8's header row and first two dependencies stay on one page.

    Without it the heading, the header row and the first dependency can end a
    page on their own while the rest of the table starts the next.
    """
    hits = [t for t in doc.tables
            if [c.text.strip() for c in t.rows[0].cells] == ["Dependency", "Contract"]]
    if len(hits) != 1:
        raise ValueError(f"The dependency table was found {len(hits)} times")
    for table_row in hits[0].rows[:2]:
        for cell in table_row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def keep_requirement_tails(doc) -> None:
    """A requirement's last row never starts a page alone, and Section 4 opens with FR-001.

    Each requirement's second-last row keeps with its last, so a Limit travels with
    its Acceptance; Section 4's lead line keeps with the first requirement, so the
    heading never ends a page on its own.
    """
    for table in doc.tables:
        if table.rows[0].cells[0].text.strip().startswith("FR-"):
            for cell in table.rows[-2].cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    lead = [p for p in doc.paragraphs
            if p.text.strip().startswith("Each requirement records the behaviour required")]
    if len(lead) != 1:
        raise ValueError(f"Section 4's lead line was found {len(lead)} times")
    lead[0].paragraph_format.keep_with_next = True


def patch_m04() -> None:
    require_newest(str(ROOT / "functional-requirements" / M04_DIR / f"{M04_STEM}_v*.docx"),
                   M04_SOURCE)
    doc = Document(str(frd_path(M04_DIR, M04_STEM, M04_SOURCE)))

    # Control page and contents.
    set_cover_version(doc, M04_SOURCE, M04_TARGET)
    set_control(doc, "Version", M04_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Source scope", CODE_BASELINE)
    set_control(doc, "Code inspected", CODE_BASELINE)
    set_control(doc, "MRD baseline", f"XVS Module Requirements Document v{MRD_VERSION}")
    keep_format(row_labelled(doc, "4. Functional Requirements").cells[1],
                "FR-001 to FR-033, each with the code evidence behind it")
    keep_format(row_labelled(doc, "10. MRD Traceability").cells[1],
                f"Agreement with MRD v{MRD_VERSION}'s twenty-three capability entries")

    # 1.1 In scope.
    append_to(table_with_header(doc, "Area"), "The permission vocabulary",
              " Each module and resource carries a readable label, seeded where the slug reads "
              "wrongly.")
    field_row = row_labelled(doc, "Field-level security")
    keep_format(field_row.cells[0], "Field Access")
    keep_format(field_row.cells[1], (
        "FieldDefinition, the code-owned registry of fields an administrator may restrict, "
        "filed under the same resources as permissions; each role's Read and Write switch on "
        "each field, with one-person exceptions; and vs_rbac.field_enforcement, which makes "
        "every surface leave out what the caller may not read and refuse what they may not "
        "write."))
    insert_row_after(row_labelled(doc, "Branch scope"), [
        "Branch-bound writes",
        "caller_may_change, the one rule by which a branch-bound caller changes a row only when "
        "every branch it belongs to is theirs, and grant_reach, which keeps their role grants "
        "inside their branches and to people posted only there.",
    ])

    # 3 Actors.
    admin = row_labelled(doc, "School administrator")
    keep_format(admin.cells[1], admin.cells[1].text.rstrip() + (
        " They also decide which registered fields each role reads and writes, and record "
        "one-person field exceptions."))
    append_to(table_with_header(doc, "Actor"), "School administrator",
              " school.field_access.view to read a role's field switches and "
              "school.field_access.update, restricted, to change them; the override keys also "
              "govern field exceptions.")
    insert_row_after(admin, [
        "A branch-bound administrator",
        "Read shared rows and change only what sits wholly inside their branches, and grant, "
        "change or revoke a role only within their branches, for somebody posted only there.",
        "The same keys as any administrator, narrowed by their own grants' reach through "
        "visible_branch_ids (FR-032).",
    ])
    append_to(table_with_header(doc, "Actor"), "CodeX platform staff",
              " platform.field_access.view and .update, the second restricted, for platform "
              "roles' field switches.")

    # FR-002: the sweep that finished the reclassifications.
    fr002 = fr_table(doc, "FR-002")
    edit_value(fr002, "Current evidence",
               "0018 did not do the same for finance.entity.create (see Needs Attention).",
               "0018 did not do the same for finance.entity.create, and no sweep reached tenant "
               "permission groups or looked past schools at organization tenants. Migration 0021 "
               "finished the job from Permission.scope rather than from a list: "
               "vs_rbac.scope_withdrawal.withdraw_from_tenants deletes every tenant-side grant of "
               "a key a tenant may no longer hold, on a tenant's own roles, the prebuilt library, "
               "tenant-scoped groups, ALLOW overrides and the ADD lines of pending role changes, "
               "in every tenant that is not the platform, school or organization alike, and "
               "spares the platform tenant's own rows, platform-scoped groups, DENY overrides and "
               "decided requests. It found the library's finance.entity.create and "
               "import.templates.create inside two tenant groups left by migration 0008. A later "
               "reclassification calls the same function in its own migration. "
               "audit_permission_scope asks, beside its questions about routes, whether any "
               "tenant still holds a key that is now platform-only, and --strict fails on a "
               "leftover grant as it does on a global-table write.")
    edit_value(fr002, "Acceptance",
               "platform.team_overrides.* or platform.staff_payroll.* is refused",
               "platform.team_overrides.* or platform.field_access.update is refused")
    append_to(fr002, "Acceptance",
              " WithdrawalReachesEverySurfaceTests, WithdrawalSparesWhatItShouldTests and "
              "WithdrawalReachesAnyTenantKindTests pin the sweep, an organization tenant "
              "included.")

    # FR-008: permission exceptions as at a day.
    fr008 = fr_table(doc, "FR-008")
    append_to(fr008, "Current evidence",
              " The list reads as at a past day: ?as_at= returns the exceptions that stood at "
              "the end of that day, paginated, with as_at and history_starts, is_expired judged "
              "at that moment and granted_by_role null, because a role's permissions keep no "
              "history to compare against. A day before the person's account history, or before "
              "exceptions were tracked, is refused with 409 HISTORY_NOT_KEPT.")
    append_to(fr008, "Acceptance",
              " A past list shows what stood that day and refuses a day before its history "
              "(PermissionOverridesAsAtTests).")

    # FR-010: the session's branch reach, and the one list left.
    fr010 = fr_table(doc, "FR-010")
    append_to(fr010, "Current evidence",
              " branch_reach_payload renders the same answer for the sign-in response and /me as "
              "branch_reach: whole_tenant true for an unnarrowed caller, otherwise the exact "
              "branch ids, which are empty when every granted branch has left service. It is "
              "computed for the effective user, so under impersonation it describes the person "
              "proxied. A client builds a branch picker from it rather than from the branch list, "
              "which answers which branches the school runs, not which the reader may work in. "
              "The same reach decides which money a write may move: a bank account, or a ledger "
              "account behind one, that belongs to a branch outside the caller's reach is unknown "
              "to them on a write, and a branch's document is paid only from its own branch's "
              "account or a school-wide one (FR-032, Module 19).")
    edit_value(fr010, "Limit",
               "Two list surfaces are still unnarrowed and are recorded in Needs Attention: the "
               "financial statements, which aggregate ledger lines through several relation "
               "paths, and this module's own role and assignment lists. The source audit sees "
               "lookups by id, not list querysets, so it reports neither.",
               "One list surface is still unnarrowed and is recorded in Needs Attention: this "
               "module's own role list. The assignment list and detail show a branch-bound "
               "caller only holders posted at their branches or school-wide (FR-032), and the "
               "financial statements read a branch-bound reader's own journals (Module 19). The "
               "source audit sees lookups by id, not list querysets, so it would not report the "
               "role list.")

    # FR-019: the audit vocabulary.
    fr019 = fr_table(doc, "FR-019")
    append_to(fr019, "Current evidence",
              " record_rbac_audit refuses an action type the central trail does not register, "
              "raising before anything is written, because the mirror would otherwise drop the "
              "event without an error. FIELD_REGISTRY_SYNCED, FIELD_ACCESS_CHANGED, "
              "FIELD_ACCESS_RESET, FIELD_OVERRIDE_CREATED and FIELD_OVERRIDE_LIFTED are "
              "registered (vs_audit migration 0016), and permission-group creation, update and "
              "deletion are recorded as CREATE, UPDATE and DELETE on entity type "
              "permission_group.")
    append_to(fr019, "Acceptance",
              " An unregistered action type raises and writes nothing, and every literal type "
              "passed to emit_audit_event or record_rbac_audit in the source is registered "
              "(LiteralActionTypesAreRegisteredTests).")

    # FR-021: provisioning drops what the tenant may not hold.
    fr021 = fr_table(doc, "FR-021")
    append_to(fr021, "Current evidence",
              " The copy drops, and logs, a library default the new tenant may not hold rather "
              "than refusing it, because the caller asked for a role and not for that key, and a "
              "key the tenant may not hold confers nothing inside it; writing such a key "
              "directly is still refused. The prebuilt role's default field switches are copied "
              "in the same transaction, skipping a default on an inactive field or on a field "
              "the tenant may not hold, and keeping Write on a field no longer writable as Read "
              "only.")
    append_to(fr021, "Acceptance",
              " A stale library no longer takes school creation down: the key is dropped and "
              "Finance Admin provisioned, while the guard still refuses the key written directly "
              "(LibraryLeftoverBreaksTenantProvisioningTests), and the field defaults follow the "
              "same rules (ProvisioningCopiesFieldDefaultsTests).")
    edit_value(fr021, "Limit",
               "The provisioning copy is not filtered by scope, so a library default for a key "
               "later reclassified PLATFORM is refused at the grant and provisioning fails. "
               "School creation provisions Finance Admin, so a school cannot be created on a "
               "database whose library attached finance.entity.create while it was "
               "tenant-holdable; this is traced in code, and no test builds Finance Admin from "
               "seeded finance keys. See Needs Attention. ",
               "A default the copy drops is logged rather than reported to whoever created the "
               "school; audit_permission_scope names any such library row, and "
               "withdraw_from_tenants takes it back. ")

    # FR-022: the enforcement contract.
    fr022 = fr_table(doc, "FR-022")
    set_value(fr022, "Requirement", FR022_REQUIREMENT)
    set_value(fr022, "Current evidence", FR022_EVIDENCE)
    set_value(fr022, "Acceptance", FR022_ACCEPTANCE)
    set_value(fr022, "Limit", FR022_LIMIT)

    # FR-025, FR-027 and FR-028.
    edit_value(fr_table(doc, "FR-025"), "Limit",
               "a view gated solely by rbac_group_permission or by HasAnyModuleAccess passes "
               "untouched.",
               "a view gated solely by rbac_group_permission or by HasAnyModuleAccess passes "
               "untouched. Module-membership reads now open every finance and procurement "
               "dashboard and the finance reference lists (FR-031), and a dashboard block asks "
               "its key of the plan as well, through keys_within_plan, the bulk form of "
               "plan_refusal's verdict (FR-031).")
    edit_value(fr_table(doc, "FR-027"), "Limit",
               "Field exceptions are not kept in the record history.",
               "Field exceptions are kept in the record history and read as at a day (FR-030); a "
               "role's switches are not, so a past view carries no role comparison.")
    edit_value(fr_table(doc, "FR-028"), "Limit",
               "saying to grant it from the person's profile (Needs Attention).",
               "saying to grant it from the person's profile (Needs Attention)." + D13_SENTENCE)

    # New requirements.
    add_fr(doc, FR029_HEADING, "FR-029 | Implemented", FR029)
    set_status(doc, fr_table(doc, "FR-029"), "Implemented")
    add_fr(doc, FR030_HEADING, "FR-030 | Implemented with limits", FR030)
    set_status(doc, fr_table(doc, "FR-030"), "Implemented with limits")
    add_fr(doc, FR031_HEADING, "FR-031 | Implemented", FR031)
    set_status(doc, fr_table(doc, "FR-031"), "Implemented")
    add_fr(doc, FR032_HEADING, "FR-032 | Implemented with limits", FR032)
    set_status(doc, fr_table(doc, "FR-032"), "Implemented with limits")
    add_fr(doc, FR033_HEADING, "FR-033 | Implemented", FR033)
    set_status(doc, fr_table(doc, "FR-033"), "Implemented")

    # 5.2 Withdrawal.
    insert_row_after(row_labelled(doc, "Personal DENY override"), [
        "Close a field on a role, or deny it to one person",
        "Read or Write on one registered field, for every holder of the role or for one person.",
        "On the next request. Read off also takes Write, a personal DENY beats every role and "
        "ALLOW, and a reset returns the field to its default.",
    ])

    # 6 Data model.
    insert_row_after(row_labelled(doc, "Permission"), [
        "FieldDefinition",
        "One registered field of a resource that an administrator may restrict per role.",
        "Primary key module.resource.name, unique per resource. Code-owned: written only by "
        "sync_field_registry, deactivated rather than deleted. Carries label, group, api_names, "
        "sensitive (the default: closed when true), writable, open_on_create and a scope with no "
        "default and the guard a permission's has. Migrations 0022 and 0024.",
    ])
    after_overrides = insert_row_after(row_labelled(doc, "UserPermissionOverride"), [
        "RoleFieldAccess / PrebuiltRoleFieldAccess",
        "A role's, or a library role's, Read and Write on one field.",
        "Unique per role and field. A check constraint stores Write only with Read. A missing "
        "row is the field's default and a reset deletes the row. Guards on save, clean and "
        "bulk_create refuse a field the tenant may not hold and Write on a field that is not "
        "writable; the library refuses a PLATFORM field outright. Migration 0023.",
    ])
    insert_row_after(after_overrides, [
        "UserFieldAccessOverride",
        "One person's ALLOW or DENY on READ or WRITE of one field.",
        "Unique on (user, field, access), so a new exception replaces the old. Required reason, "
        "lazy expiry, tenant-pinned. Refuses a field the tenant may not hold and ALLOW WRITE on a "
        "field that is not writable. Tracked by vs_history, so the list reads as at a day. "
        "Migration 0023.",
    ])

    # 7.1 Registry routes.
    edit_cell(row_labelled(doc, "GET /rbac/vision/permission-modules/").cells[-1],
              "List backend-defined modules.",
              "List backend-defined modules with their readable labels.")
    edit_cell(row_labelled(doc, "GET /rbac/vision/permission-resources/").cells[-1],
              "List resources within modules.",
              "List resources within modules with their readable labels.")
    insert_row_after(row_labelled(doc, "GET /rbac/vision/permissions/{key}/"), [
        "GET /rbac/vision/fields/",
        "The field registry, active or not, filterable by module, resource, sensitive, is_active "
        "and search. platform.permissions.view; a write verb answers 405.",
    ])

    # 7.2 Tenant routes.
    assign_list = row_labelled(doc, "GET, POST /rbac/tenants/{slug}/role-assignments/")
    keep_format(assign_list.cells[-1], assign_list.cells[-1].text.rstrip() + (
        " A branch-bound caller lists only grants held by people posted at their branches or "
        "school-wide, and grants only within their branches to somebody posted only there "
        "(FR-032)."))
    insert_row_after(assign_list, [
        "GET, PATCH /rbac/tenants/{slug}/role-assignments/{id}/",
        "One grant. The view keys read it and the assign keys change it. A branch-bound caller "
        "reads only a grant held by somebody posted at their branches or school-wide, and "
        "changes one only within their branches (FR-032). Changing the grant's role to one "
        "carrying restricted keys the caller does not hold is refused, in one sentence naming "
        "the role.",
    ])
    for path in ("POST /rbac/tenants/{slug}/role-assignments/{id}/revoke/",
                 "POST /rbac/tenants/{slug}/role-assignments/{id}/replace/"):
        cell = row_labelled(doc, path).cells[-1]
        keep_format(cell, cell.text.rstrip() + " A branch-bound caller only within their "
                                               "branches (FR-032).")
    overrides = row_labelled(doc, "GET, POST /rbac/tenants/{slug}/users/{user_id}/permission-overrides/")
    keep_format(overrides.cells[-1], overrides.cells[-1].text.rstrip() + (
        " ?as_at=YYYY-MM-DD lists them as they stood at the end of that day, with "
        "granted_by_role null; a day before the history starts is 409 HISTORY_NOT_KEPT."))
    catalogue = row_labelled(doc, "GET /rbac/tenants/{slug}/permission-catalogue/")
    added = insert_row_after(catalogue, [
        "GET /rbac/tenants/{slug}/access-catalogue/",
        "The same vocabulary as a tree: module, resource, then the resource's permissions and "
        "restrictable fields, filterable by module, resource and search. Role view keys or "
        "platform.permissions.view, open before go-live. A platform operator holding a platform "
        "role-view key may assert a school, and entries follow that school's scope, so no "
        "PLATFORM entry appears.",
    ])
    added = insert_row_after(added, [
        "GET, PATCH /rbac/tenants/{slug}/roles/{key}/field-access/",
        "A role's Read and Write on every field the tenant may hold, filterable by module, "
        "resource, search and state (hidden, read_only, full). Reading takes a field access view "
        "or update key; PATCH takes school.field_access.update or platform.field_access.update, "
        "1 to 200 changes, each field once, atomic. Open before go-live; not reachable across "
        "tenants.",
    ])
    added = insert_row_after(added, [
        "GET, POST /rbac/tenants/{slug}/users/{user_id}/field-access-overrides/",
        "One person's field exceptions, filterable by mode and access, each with role_state. The "
        "permission-override keys; self-edit forbidden; a new exception on the same field and "
        "access replaces the old. ?as_at= lists them as at a day, with role_state null.",
    ])
    insert_row_after(added, [
        "DELETE /rbac/tenants/{slug}/users/{user_id}/field-access-overrides/{id}/",
        "Lift a field exception. The delete key, audited with the person's state before and "
        "after.",
    ])

    # 7.3 What a caller gets back.
    edit_cell(row_labelled(doc, "A key held through a group grant or HasAnyModuleAccess alone")
              .cells[-1],
              "both are rare and reach data that is meaningless without a module already reached "
              "through a gated route.",
              "a module-membership read opens a module's shared lists and dashboards (FR-031), "
              "and a dashboard block asks its key of the role and of the school's plan "
              "(FR-031).")
    added = row_labelled(doc, "An empty list")
    for condition, answer in (
        ("A submitted field the caller may not write, hidden or read-only",
         "403 field_write_denied, naming every offending field in one answer, and nothing is "
         "saved. An unchanged echo on update from a caller who may read the field, a blank on "
         "create and an open-on-create value on create are dropped instead."),
        ("A route whose whole response is a field the caller may not read, such as an import "
         "batch's uploaded file",
         "403 field_read_denied, naming the field, asked before the file is looked up, so it is "
         "the same whether or not the record holds one."),
        ("A field switch naming an unknown, inactive or platform field, or Write on a field that "
         "is not writable",
         "400. An unknown, inactive and platform field are refused alike, so a school cannot "
         "learn which platform fields exist; a field that is not writable is named."),
        ("A field switch change with no changes, more than 200, or one field twice",
         "400."),
        ("A field exception on yourself",
         "403, on create and on lift alike, under either identity of a proxy session."),
        ("An exceptions list read as at a day before its history",
         "409 HISTORY_NOT_KEPT."),
        ("A grant reaching past a branch-bound caller's branches, or held by somebody posted "
         "beyond them",
         "400 under branch or under user, saying a school-wide administrator can make it."),
        ("A change to a row shared beyond a branch-bound caller's branches",
         "403 SHARED_RECORD_READ_ONLY, naming who can make the change."),
        ("A dataset's own import key on delete, edit, issue resolution or the administrative "
         "feeds, or on rollback of anything but a bank statement's own batch",
         "403, the same for a batch that exists and one that does not."),
    ):
        added = insert_row_after(added, [condition, answer])

    # 8 Dependencies.
    dependencies = table_with_header(doc, "Dependency")
    append_to(dependencies, "Module 5, Audit and Activity Logging",
              " Every action type this module writes must be one Module 5 registers, and "
              "record_rbac_audit refuses one that is not.")
    append_to(dependencies, "Every engine app",
              " Each also declares its restrictable fields in its own field_access.py, and "
              "Field Access is only as complete as those declarations.")
    module12 = row_labelled(doc, "Module 12, Staff Management")
    keep_format(module12.cells[-1], module12.cells[-1].text.rstrip() + (
        " Its staff import takes the dataset key school.staff.import (FR-033), and its staff "
        "writes and bulk grant apply the branch-bound rules of FR-032."))
    added = insert_row_after(module12, [
        "Module 10, Bulk Data Import",
        "HasImportBatchRBACPermission builds on this module's gate and lets a registered dataset "
        "key stand in for the wizard's engine keys, and for an engine key past the wizard that "
        "its dataset declares, which only bank statements do, for rollback (FR-033).",
    ])
    insert_row_after(added, [
        "Module 19, Finance and Accounting",
        "Opens its shared reference reads and dashboards on module membership through "
        "HasAnyModuleAccess, and judges writes to a branch budget through caller_may_change "
        "(FR-031, FR-032).",
    ])

    # 8.1 Minimum verification.
    insert_after(doc, "•  The seeders and cleanup migration",
                 "•  Scope withdrawal: every tenant-side surface loses a reclassified key, an "
                 "organization tenant included; the platform tenant's roles, platform groups and "
                 "DENY overrides keep it; provisioning from a stale library drops the key and "
                 "still creates the role, while writing the key directly is refused.")
    insert_after(doc, "•  Durable role mutation failure",
                 "•  The audit vocabulary: an unregistered action type raises and writes nothing, "
                 "and every literal type in the source is registered.")
    insert_after(doc, "•  The branch source audit",
                 "•  Field Access: every guarded field declared, every declared name reaching a "
                 "serializer and a field's scope equal to its guarding key's; no serializer "
                 "reading or writing a registered name without enforcing it; every declared "
                 "surface, rendered for a caller with every field closed, carrying no registered "
                 "name at any depth; switch changes refused across tenants, exceptions refused on "
                 "yourself, and a platform field refused exactly as an unknown one; a route "
                 "serving one field refused field_read_denied exactly when its record's detail "
                 "would leave the field out; a failed "
                 "audit rolling a switch change back; and a cold evaluation costing three "
                 "queries.")
    insert_after(doc, "•  Field Access:",
                 "•  Branch-bound writes and dataset keys: a grant reaching past the caller's "
                 "branches, or held by somebody posted beyond them, refused; the holder list "
                 "narrowed to the caller's people and the school-wide; a school dataset key "
                 "admitted to the wizard steps and refused rollback, delete, edit, issue "
                 "resolution and the administrative reads alike for a batch that exists and one "
                 "that does not; the bank-statement key rolling back its own statement import "
                 "and nothing else past the wizard.")

    # 9 Needs Attention.
    attention = table_with_header(doc, "Pri.")
    school_creation = row_with_cell_starting(attention, 1, "School creation fails where")
    school_creation._tr.getparent().remove(school_creation._tr)
    lists = row_with_cell_starting(attention, 1, "Two surfaces still return rows")
    keep_format(lists.cells[1], NA_ROLE_LIST)
    keep_format(lists.cells[2], NA_ROLE_LIST_FIX)
    keep_format(lists.cells[3], "FR-010, FR-032")
    creation = row_with_cell_starting(attention, 1, "Creating a person with a restricted role")
    keep_format(creation.cells[1], creation.cells[1].text.rstrip() + D13_SENTENCE)
    keep_format(row_with_cell_starting(attention, 1, "Customer, vendor").cells[3],
                "FR-022, FR-027, FR-029")
    history = row_with_cell_starting(attention, 1, "Field exceptions keep no history")
    keep_format(history.cells[1], NA_ROLE_HISTORY)
    keep_format(history.cells[2], NA_ROLE_HISTORY_FIX)
    keep_format(history.cells[3], "FR-027, FR-030")
    edit_box(table_with_header(doc, "REMOVAL RULE").rows[0].cells[0], [
        ("replace", "• Against v1.27",
         "• Against v1.28 one item leaves: school creation no longer fails on a library default "
         "the tenant may not hold. Two are rewritten to the risk that remains: of the two "
         "unnarrowed lists, only the role list is left, and of the missing histories, only a "
         "role's switches and permissions."),
    ])

    # 10 Traceability.
    edit_paragraph(doc, "MRD v2.89 records Module 4", "MRD v2.89", f"MRD v{MRD_VERSION}")
    trace = table_with_header(doc, "MRD Module 4 capability")
    registry = row(trace, "Backend-owned permission registry by module, resource, and action")
    keep_format(registry.cells[1], "FR-001, FR-029")
    append_to(trace, "Backend-owned permission registry by module, resource, and action",
              " Modules and resources carry readable labels, and restrictable fields sit in the "
              "same tree.")
    edit_value(trace, "Role templates",
               "A library default for a key later reclassified PLATFORM breaks provisioning "
               "(Needs Attention P1).",
               "A library default the tenant may not hold is dropped and logged rather than "
               "breaking provisioning, and migration 0021 swept the ones left behind.")
    assign = row(trace, "Assign, revoke, and replace user roles")
    keep_format(assign.cells[1], "FR-009, FR-016, FR-028, FR-032")
    append_to(trace, "Assign, revoke, and replace user roles",
              " A branch-bound administrator grants, changes and revokes only within their "
              "branches, for people posted only there.")
    evaluation = row(trace, "Entity-aware permission evaluation")
    keep_format(evaluation.cells[1], "FR-007, FR-010, FR-012, FR-031, FR-032, FR-033")
    edit_value(trace, "Entity-aware permission evaluation",
               "Two unnarrowed list surfaces remain at P2.",
               "One unnarrowed list, the role list, remains at P2. A shared reference read opens "
               "on module membership, a branch-bound caller changes only rows wholly inside "
               "their branches, and a dataset's import key stands in for the wizard's engine "
               "keys, the bank-statement key for the rollback of its own statements as well.")
    append_to(trace, "Permission and assignment audit history",
              " An action type the central trail does not register is refused before anything "
              "is written.")
    append_to(trace, "Platform and tenant permission scope, enforced at grant",
              " A reclassified key is withdrawn from every tenant surface by migration, and the "
              "scope audit fails on a leftover grant.")
    # Two labels take the MRD's wording for their entries.
    for old, new in (("Permission groups", "Administrator-created permission groups"),
                     ("Assign, revoke, and replace user roles",
                      "Assign, revoke, and replace user roles with configured branch reach")):
        keep_format(row(trace, old).cells[0], new)
    field_access = row(trace, FIELD_ACCESS_OLD)
    keep_format(field_access.cells[0], FIELD_ACCESS_NEW)
    keep_format(field_access.cells[1], "FR-022, FR-027, FR-029, FR-030")
    keep_format(field_access.cells[2], (
        "Implemented. Every registered field sits under its resource in one tree with the "
        "permissions; each role has a Read and a Write switch on each, with one-person "
        "exceptions kept in the record history; every surface leaves out what the caller may "
        "not read and refuses what they may not write. Person records, finance banking and "
        "payroll, payments, vendors, imports and user accounts are registered."))
    update_reconciliation(doc, [
        ("replace", "• MRD v2.89 lists Module 4",
         f"• MRD v{MRD_VERSION} lists Module 4 as Roles & Permissions (RBAC), Phase V1, Backend "
         "Complete, In use Complete, code vs_rbac, with twenty-three capability entries. The "
         "module number, name, phase, states and ownership agree with this revision."),
        ("replace", "• Field Access over whole person records, with composite names rebuilt",
         "• Field Access, with composite names rebuilt from readable parts, is the entry mapped "
         "to FR-027 that took the count to twenty-three."),
        ("append", "• The field registry and its tree, per-role field switches with one-person "
         "exceptions, and their enforcement on every surface are read under the existing Field "
         "Access entry, which with FR-022, FR-029 and FR-030 covers every registered field and "
         "not only person records, so the count stays at twenty-three. MRD v2.94 words the "
         "entry the same way, as Field Access over every registered field, by role switch and "
         "one-person exception."),
        ("append", "• Shared reference reads on module membership, branch-bound writes and "
         "dataset import keys (FR-031 to FR-033) extend the existing evaluation and assignment "
         "entries rather than adding one."),
    ])

    keep_dependency_table_opening_together(doc)
    keep_table_rows_whole(doc)
    keep_requirement_tails(doc)
    log_change(doc, M04_TARGET, M04_SUMMARY)
    assert_absent_outside_log(
        doc, "MRD v2.89", "MRD v2.93", "FR-001 to FR-028", "FieldSecurityMixin", "_fls_permissions",
        "Field Access over whole person records", "names person records alone",
        "platform.staff_payroll.*", "Field-level security", "Two list surfaces",
        "Two surfaces still return", "Field exceptions keep no history",
        "Field exceptions are not kept", "Against v1.27", "not filtered by scope",
        "School creation fails where", "Needs Attention P1", "both are rare",
        "It is a mechanism, not a policy", "not of the plan", "answers to the role, not",
    )
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M04_DIR, M04_STEM, M04_TARGET),
           f"{M04_STEM.replace('_', ' ')} v{M04_TARGET}", M04_TARGET)


def main() -> None:
    patch_m04()


if __name__ == "__main__":
    main()
