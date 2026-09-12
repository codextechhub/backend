#!/usr/bin/env python3
"""Reconcile the Module 10 and Module 21 to 24 FRDs with the backend at 13675db.

Module 21 to 24 share vs_procurement, and three changes cut across them. Vendors
and stock locations are master data, and read a null branch as shared across the
tenant rather than as nobody's, with every vendor and location reference
resolved inside the caller's branches. Reversing an approval now asks the
document first: a purchase order already released to its vendor refuses, and
any other reversed document returns to PENDING. Selling by depth moved the AP
and GR/IR analytics onto procurement.analytics.view, which Module 23 had not
recorded. Module 24 also still listed two stock reports that were folded into
the item list, and Module 23 still carried an attachment gap that FR-014 had
closed.

Module 10 gains the six datasets a school imports for itself, which v1.2 named
in Needs Attention while FR-001 still said there were none, together with the
branch a batch is filed under, the branch-narrowed file download, the storage
fix that stopped every upload being stored empty, and the bell notice an import
run now sends through Module 8.

Each document takes one minor version. Source MRD references move to v2.76,
which is cut after these documents.

    python tools/patch_procurement_import_reconciliation_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.table import _Row

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
)

REVIEW_DATE = "11 September 2026"
SHORT_DATE = "11 Sep 2026"
MRD_VERSION = "2.76"
CODE_BASELINE = "Backend worktree at 13675db, 11 September 2026"
DEPLOYMENT = "Backend evidence only; nothing here is deployed."


# ── keep-format docx helpers ────────────────────────────────────────────────

def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError(f"Paragraph carries no run to inherit formatting from: {paragraph.text!r}")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    """Rewrite a one-paragraph cell, keeping the formatting it already has."""
    for paragraph in cell.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    set_run_text(cell.paragraphs[0], text)


def append_cell_text(cell, tail: str) -> None:
    set_cell_text(cell, cell.text.strip() + tail)


def rewrite_box(cell, lines: list[str]) -> None:
    """Rewrite a callout cell line by line.

    A callout is one bold coloured heading paragraph followed by regular bullet
    paragraphs, so line i reuses paragraph i, overflow clones the last bullet,
    and surplus paragraphs are removed. Cloning the heading for a new bullet
    would render every added line as another heading.
    """
    original = list(cell.paragraphs)
    if len(original) < 2:
        raise ValueError("A callout needs a heading and at least one bullet to reuse")
    for index, line in enumerate(lines):
        if index < len(original):
            set_run_text(original[index], line)
        else:
            clone = copy.deepcopy(original[-1]._p)
            cell.paragraphs[-1]._p.addnext(clone)
            set_run_text(cell.paragraphs[-1], line)
    for paragraph in original[len(lines):]:
        paragraph._p.getparent().remove(paragraph._p)


def unique_cells(row) -> list:
    seen, cells = set(), []
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        cells.append(cell)
    return cells


def find_table(doc, header: str, *, contains: str | None = None):
    """The first table whose first cell starts with ``header``.

    ``contains`` narrows to a table that also has a row whose first cell starts
    with that text, for documents where two tables share a header.
    """
    for table in doc.tables:
        if not table.rows[0].cells[0].text.strip().startswith(header):
            continue
        if contains is None or any(
            row.cells[0].text.strip().startswith(contains) for row in table.rows
        ):
            return table
    raise ValueError(f"Table not found: {header!r} containing {contains!r}")


def find_fr(doc, label: str):
    return find_table(doc, f"{label} |")


def find_row(table, prefix: str, *, column: int = 0):
    for row in table.rows:
        if row.cells[column].text.strip().startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix!r}")


def field(table, label: str):
    """The value cell of a requirement or control table row named ``label``."""
    return find_row(table, label).cells[1]


def clone_row(table, anchor, values: list[str], *, before: bool = False):
    new_tr = copy.deepcopy(anchor._tr)
    if before:
        anchor._tr.addprevious(new_tr)
    else:
        anchor._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)
    return row


def delete_row(row) -> None:
    row._tr.getparent().remove(row._tr)


def clone_styled_row(table, template, anchor, values: list[str]):
    """Copy ``template``'s formatting into a new row placed after ``anchor``."""
    new_tr = copy.deepcopy(template._tr)
    anchor._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    for cell, value in zip(unique_cells(row), values):
        set_cell_text(cell, value)
    return row


def hold_together(table, *, lead_rows: int = 2) -> None:
    """Keep every row whole, and the first ``lead_rows`` with the row after them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
    for row in table.rows[:lead_rows]:
        for cell in unique_cells(row):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def paginate(doc) -> None:
    """Stop a requirement's heading and banner, or a table's lead, being stranded.

    Each requirement table and the traceability table keep their rows whole and
    their first rows with the next, and the paragraph directly above each keeps
    with the table, so a heading never ends a page alone. The routes intro keeps
    with the sub-heading and table that follow it for the same reason.
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    for table in doc.tables:
        first = table.rows[0].cells[0].text.strip()
        if not ((first.startswith("FR-") and " | " in first) or first == "MRD capability"):
            continue
        hold_together(table)
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            Paragraph(previous, table._parent).paragraph_format.keep_with_next = True
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("All routes are mounted under /v1/procurement/"):
            paragraph.paragraph_format.keep_with_next = True


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def update_control(doc, *, version: str, module: int, supporting: str | None = None) -> None:
    control = find_table(doc, "Document", contains="Code baseline")
    set_cell_text(field(control, "Version"), version)
    set_cell_text(field(control, "Review date"), REVIEW_DATE)
    set_cell_text(field(control, "Code baseline"), CODE_BASELINE)
    set_cell_text(
        field(control, "Source MRD"),
        f"XVS Module Requirements Document v{MRD_VERSION} | Module {module}",
    )
    if supporting is not None:
        set_cell_text(field(control, "Supporting apps"), supporting)


def body_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Body paragraph not found: {prefix!r}")


def prepend_change_log(doc, version: str, summary: str) -> None:
    table = find_table(doc, "Version", contains="1.0")
    clone_row(table, table.rows[1], [version, SHORT_DATE, summary], before=True)


def finish(doc, output: Path, title: str, version: str) -> None:
    paginate(doc)
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


# ═════════════════════════════════════════════════════════════════════════════
# Module 21 - Vendor Management
# ═════════════════════════════════════════════════════════════════════════════

M21_DIR = "21-vendor-management"
M21_STEM = "XVS_M21_Vendor_Management_Functional_Requirements_Document"
M21_SOURCE, M21_TARGET = "1.3", "1.4"

M21_FR011_HEADING = "FR-011 Share Vendors Across Branches Without Exposing a Branch's Own"

M21_FR011 = {
    "header": "FR-011 | Implemented",
    "Requirement": (
        "Every branch must be able to buy from a vendor the tenant shares, and no "
        "branch may read, change or buy from a vendor another branch keeps to itself."
    ),
    "Current evidence": (
        "A vendor is master data rather than a document, so it reads a null branch "
        "inclusively, the reading vs_academics takes for a catalogue and vs_finance "
        "for a customer: a vendor with no branch is shared across the tenant and "
        "reachable from every branch, and a caller pinned to a branch also reaches "
        "that branch's own. Procurement documents read a null branch exclusively, "
        "and the two readings are one narrowing with one switch, chosen by the kind "
        "of row. The vendor list, the summary, the detail and update routes and the "
        "per-vendor insights are narrowed, and so is the resolver every vendor "
        "reference passes through: the vendor a purchase order, goods receipt, "
        "bill, payment, contract, RFQ invitation, quotation or assessment names, "
        "and a catalogue item's preferred vendor."
    ),
    "Acceptance": (
        "A storekeeper pinned to Ikeja can read and order from the stationer the "
        "tenant shares, and cannot read, rename, see the spend of, or raise a "
        "purchase against a vendor recorded as Lekki's own. A vendor another branch "
        "keeps to itself is reported exactly as one that does not exist, so an id "
        "cannot be used to learn that it is there. ProcurementCatalogueReadingTests "
        "covers each route and the resolver, and two of its tests assert that the "
        "shared vendor stays visible from every branch, because a narrowing that "
        "hides the tenant's own stationer is as much a defect as one that shows "
        "another branch's."
    ),
    "Current limit": (
        "The vendor API takes no branch input, so every vendor is created shared "
        "across the tenant and the narrowing has no effect on current data. It is "
        "in force for the day a tenant records a vendor as one branch's own, which "
        "the API cannot yet do."
    ),
}

M21_FR008_EVIDENCE = (
    "Assessments are point-in-time scorecards, list and create only, with no edit "
    "path. Listing them rides procurement.analytics.view, the key that also serves "
    "the vendor performance report they feed; recording one needs its own "
    "dedicated key."
)

M21_FR009_EVIDENCE = (
    "Vendor insight, category insight, catalogue usage insight, the vendor spend "
    "analysis and the vendor performance report are served as entity-scoped reads "
    "over the AP sub-ledger. The category and catalogue insights and the vendor "
    "summary ride procurement.report.view; per-vendor insight, the spend analysis "
    "and the performance report ride procurement.analytics.view, which is sold a "
    "depth deeper. The vendor summary and per-vendor insight reach only the "
    "vendors the caller's branches reach (FR-011)."
)

M21_BRANCH_ERROR = [
    "Naming a vendor another branch keeps to itself, by id or code",
    "Reported exactly as a vendor that does not exist: 404 on the vendor's own "
    "routes, a validation error on a document's vendor reference.",
]

M21_VENDOR_MODEL = [
    "Vendor",
    "Identity, status, contacts, tax, bank name and code, account, hold, KYC, and "
    "an optional branch.",
    "Sensitive fields gated on read and write; bank changes reset KYC. A null "
    "branch means shared across the tenant, and the API sets no other value yet.",
]

M21_MODULE4 = (
    "Supplies the keys, the field-level policy over contacts, tax and bank "
    "detail, and the branch narrowing, which vendors apply inclusively (FR-011)."
)

M21_TRACEABILITY_LEAD = (
    f"Module 21 carries 13 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Reading a null branch as shared strengthens the "
    "vendor profile entry without changing the count."
)

M21_CHANGE = (
    "Records that a vendor is master data and reads a null branch as shared across "
    "the tenant. The vendor list, summary, detail and update routes and per-vendor "
    "insights were not narrowed by branch at all, so a caller pinned to one branch "
    "could read another branch's vendor, rename it and read a year of its spend by "
    "naming an id. Procurement's exclusive reading of a document would have hidden "
    "the vendors a tenant shares and left a storekeeper nobody to order from, so "
    "vendors take the inclusive reading instead: the shared rows plus the caller's "
    "own, with a vendor another branch keeps to itself reported exactly as one "
    "that does not exist. New FR-011 records it, together with the resolver every "
    "vendor reference passes through, so a purchase, contract, invitation, "
    "quotation or assessment cannot name another branch's vendor. The vendor API "
    "takes no branch input, so every vendor is shared today and nothing changes on "
    "current data. FR-008 and FR-009 are corrected: both still named "
    "procurement.report.view for assessments and per-vendor insight after they "
    "moved to procurement.analytics.view. The typed errors, the Vendor data "
    "contract, the Module 4 dependency and MRD traceability are updated. Module 21 "
    f"stays at 13 capabilities. {DEPLOYMENT}"
)


def add_requirement_after(doc, previous_label: str, heading: str, fr: dict):
    """Clone the requirement block ``previous_label`` and fill the copy.

    A block is its Heading 2 paragraph, its table and the empty spacer after it.
    The spacer here is a plain paragraph, not the page break that ends a section,
    so copying it keeps the new block spaced like its neighbours.
    """
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    table = find_fr(doc, previous_label)
    heading_el = table._tbl.getprevious()
    spacer_el = table._tbl.getnext()
    previous_heading = Paragraph(heading_el, table._parent)
    spacer = Paragraph(spacer_el, table._parent)
    if not previous_heading.style.name.startswith("Heading"):
        raise ValueError("No heading directly above the requirement table")
    if spacer.text.strip() or 'w:type="page"' in spacer_el.xml:
        raise ValueError("The element after the requirement table is not a plain spacer")

    new_heading = copy.deepcopy(heading_el)
    new_table = copy.deepcopy(table._tbl)
    new_spacer = copy.deepcopy(spacer_el)
    spacer_el.addnext(new_heading)
    new_heading.addnext(new_table)
    new_table.addnext(new_spacer)

    set_run_text(Paragraph(new_heading, table._parent), heading)
    block = Table(new_table, table._parent)
    set_run_text(block.rows[0].cells[0].paragraphs[0], fr["header"])
    for label in ("Requirement", "Current evidence", "Acceptance", "Current limit"):
        set_cell_text(field(block, label), fr[label])
    return block


def patch_m21(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M21 Vendor Management Functional Requirements Document v{M21_TARGET}"

    # Bind every table before the requirement block is inserted.
    fr008, fr009 = find_fr(doc, "FR-008"), find_fr(doc, "FR-009")
    errors = find_table(doc, "Condition")
    model = find_table(doc, "Model")
    dependencies = find_table(doc, "Dependency")
    traceability = find_table(doc, "MRD capability")

    replace_cover_version(doc.tables[0], M21_SOURCE, M21_TARGET)
    update_control(doc, version=M21_TARGET, module=21)

    set_cell_text(field(fr008, "Current evidence"), M21_FR008_EVIDENCE)
    set_cell_text(field(fr009, "Current evidence"), M21_FR009_EVIDENCE)
    clone_row(errors, find_row(errors, "Creating a duplicate category"), M21_BRANCH_ERROR)
    # Exact match: "Vendor" is also the prefix of the VendorCategory row above it.
    vendor = next(row for row in model.rows if row.cells[0].text.strip() == "Vendor")
    for cell, value in zip(unique_cells(vendor), M21_VENDOR_MODEL):
        set_cell_text(cell, value)
    set_cell_text(find_row(dependencies, "Module 4").cells[1], M21_MODULE4)
    set_cell_text(find_row(traceability, "Vendor profile and status").cells[1], "FR-002, FR-011")
    set_run_text(body_paragraph(doc, "Module 21 carries"), M21_TRACEABILITY_LEAD)
    prepend_change_log(doc, M21_TARGET, M21_CHANGE)

    add_requirement_after(doc, "FR-010", M21_FR011_HEADING, M21_FR011)
    finish(doc, output, title, M21_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 22 - Procurement & Requisitions
# ═════════════════════════════════════════════════════════════════════════════

M22_DIR = "22-procurement-and-requisitions"
M22_STEM = "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document"
M22_SOURCE, M22_TARGET = "1.7", "1.8"

M22_BRANCH_SCOPE = (
    "Every tenant has at least one branch. A requisition may carry a branch or "
    "none, where none means tenant-wide, and which it carries follows the raiser's "
    "role grants rather than their staff record. Somebody granted at one branch "
    "stamps it; somebody granted at several who names none is refused and told to "
    "name one; somebody holding the role for the whole tenant files the request "
    "tenant-wide. A branch-scoped approval stage resolves people appointed at that "
    "branch plus those appointed tenant-wide. Reads follow the same answer and, "
    "for a document, read a null branch exclusively: a tenant-wide purchase is not "
    "shown to staff pinned to one branch. Master data reads it the other way. A "
    "vendor or a stock location with no branch is shared across the tenant and "
    "reachable from every branch, and only one another branch keeps to itself is "
    "withheld, so every vendor a requisition, RFQ invitation or quotation names is "
    "one the caller's branches reach."
)

M22_APPROVER_KEYS = (
    "Membership of the approver group or role the stage names, plus the queue's "
    "view key. No permission key confers approval; the stage decides."
)
M22_COMPETITION_KEYS = (
    "procurement.competition.override (CRITICAL), with a mandatory reason written "
    "into the audit record"
)
M22_ADMIN_DOES = (
    "Publish this tenant's approval ladders and read the coverage report. The "
    "approver groups those ladders name are composed in Module 7."
)

M22_FR003_EVIDENCE_TAIL = (
    " A reversal the engine records against the vote that decided a stage is "
    "handed back as well: on_action_reversed returns an APPROVED or REJECTED "
    "document to PENDING, and a requisition's status to PENDING_APPROVAL, so the "
    "overlay never reads APPROVED while the workflow shows the document back "
    "under review."
)
M22_FR003_ACCEPTANCE_TAIL = (
    " WorkflowApprovalTests.test_reversing_the_first_vote_returns_the_requisition_to_its_ladder "
    "proves a reversed first vote puts the requisition back to PENDING and "
    "PENDING_APPROVAL, and that both stages then accept a fresh vote."
)
M22_FR003_LIMIT = (
    "Before a reversal writes anything, the document is consulted only to refuse "
    "a purchase order already released to its vendor (Module 23). A requisition "
    "whose approval is reversed after a purchase order was raised from it returns "
    "to PENDING_APPROVAL while that order stands."
)

M22_FR010_EVIDENCE_TAIL = (
    " Each invited vendor is resolved inside the entity and the caller's "
    "branches, so a buyer pinned to one branch can invite a vendor the tenant "
    "shares and cannot invite one another branch keeps to itself; the same "
    "resolver serves a quotation's vendor."
)
M22_FR010_ACCEPTANCE_TAIL = (
    " ProcurementCatalogueReadingTests.test_a_purchase_cannot_be_raised_against_another_sites_vendor "
    "exercises that resolver, and test_the_unknown_and_the_unentitled_are_reported_alike "
    "proves a withheld vendor is refused exactly like one that does not exist."
)

M22_APPROVED_LEAVES = (
    "Terminal here unless an administrator reverses the deciding vote, which "
    "returns it to PENDING; a purchase order is a separate act in Module 23."
)
M22_REJECTED_LEAVES = "Terminal, unless the deciding vote is reversed, which returns it to PENDING."

M22_RFQ_MODEL_NOTE = (
    "Optionally raised off a requisition, whose branch it inherits. Creation does "
    "not check that the requisition is approved."
)

M22_MODULE7_TAIL = (
    " Its handlers also answer the engine's reversal contract: validate_reversal "
    "refuses a purchase order already released to its vendor, and "
    "on_action_reversed returns a reopened document to PENDING."
)
M22_MODULE21_TAIL = (
    " A vendor with no branch is shared across the tenant; one another branch "
    "keeps to itself cannot be invited."
)

M22_REVERSAL_GAP = (
    "• Reversing an approval consults the document only for a purchase order "
    "already released to its vendor. A requisition whose approval is reversed "
    "after a purchase order was raised from it goes back to PENDING_APPROVAL while "
    "that order stands, so the requisition reads unapproved beside a commitment "
    "made on the strength of its approval."
)

M22_TRACEABILITY_LEAD = (
    f"Module 22 carries 16 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Returning a reversed approval to PENDING strengthens "
    "workflow submission, and resolving vendors inside the caller's branches "
    "strengthens supplier invitations, without changing the count."
)

M22_CHANGE = (
    "Records what reversing an approval now does to a procurement document, and "
    "the inclusive reading of vendors. The engine's reversal reopened a workflow "
    "without telling the document, so a requisition kept reading APPROVED while "
    "the workflow showed it back under review, and an order could have its "
    "approval withdrawn after the vendor had been told. The handlers now answer "
    "the engine's reversal contract: a reversed vote that reopens its stage "
    "returns the document to PENDING, and a requisition to PENDING_APPROVAL, while "
    "an order already released to its vendor refuses the reversal (Module 23). "
    "FR-003, the requisition lifecycle and the Module 7 dependency are updated, "
    "and a new gap records that a requisition reversed after an order was raised "
    "from it goes back to pending while that order stands. The RFQ data contract "
    "is corrected: an RFQ may name any requisition in the entity, and nothing "
    "checks that the requisition is approved. Vendors "
    "are master data and read a null branch as shared, so an RFQ invitation or a "
    "quotation cannot name a vendor another branch keeps to itself (FR-010, "
    "Section 2.2, the Module 21 dependency). The actors table is corrected: the "
    "competition exception holder's row carried the approval administrator's "
    "rights in place of procurement.competition.override, and the approver row "
    "read as though approval came from holding a role, when the stage decides and "
    "names an approver group or a role. Module 22 stays at 16 capabilities. "
    f"{DEPLOYMENT}"
)


def patch_m22(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M22 Procurement and Requisitions Functional Requirements "
        f"Document v{M22_TARGET}"
    )

    actors = find_table(doc, "Actor")
    fr003, fr010 = find_fr(doc, "FR-003"), find_fr(doc, "FR-010")
    lifecycle = find_table(doc, "State", contains="DRAFT")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M22_SOURCE, M22_TARGET)
    update_control(doc, version=M22_TARGET, module=22)

    set_cell_text(find_row(actors, "Approver").cells[2], M22_APPROVER_KEYS)
    set_cell_text(find_row(actors, "Competition exception holder").cells[2], M22_COMPETITION_KEYS)
    set_cell_text(find_row(actors, "Approval administrator").cells[1], M22_ADMIN_DOES)

    set_run_text(body_paragraph(doc, "Every school has at least one branch."), M22_BRANCH_SCOPE)

    append_cell_text(field(fr003, "Current evidence"), M22_FR003_EVIDENCE_TAIL)
    append_cell_text(field(fr003, "Acceptance"), M22_FR003_ACCEPTANCE_TAIL)
    set_cell_text(field(fr003, "Current limit"), M22_FR003_LIMIT)
    append_cell_text(field(fr010, "Current evidence"), M22_FR010_EVIDENCE_TAIL)
    append_cell_text(field(fr010, "Acceptance"), M22_FR010_ACCEPTANCE_TAIL)

    set_cell_text(find_row(lifecycle, "APPROVED").cells[2], M22_APPROVED_LEAVES)
    set_cell_text(find_row(lifecycle, "REJECTED").cells[2], M22_REJECTED_LEAVES)
    set_cell_text(
        find_row(find_table(doc, "Model"), "RequestForQuotation").cells[2],
        M22_RFQ_MODEL_NOTE,
    )

    append_cell_text(find_row(dependencies, "Module 7").cells[1], M22_MODULE7_TAIL)
    append_cell_text(find_row(dependencies, "Module 21").cells[1], M22_MODULE21_TAIL)

    gap_cell = gaps.rows[0].cells[0]
    rewrite_box(gap_cell, [p.text for p in gap_cell.paragraphs] + [M22_REVERSAL_GAP])

    set_run_text(body_paragraph(doc, "Module 22 carries"), M22_TRACEABILITY_LEAD)
    prepend_change_log(doc, M22_TARGET, M22_CHANGE)
    finish(doc, output, title, M22_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 23 - Purchase Orders, Delivery & AP
# ═════════════════════════════════════════════════════════════════════════════

M23_DIR = "23-purchase-orders-delivery-and-ap"
M23_STEM = "XVS_M23_Purchase_Orders_Delivery_and_AP_Functional_Requirements_Document"
M23_SOURCE, M23_TARGET = "1.8", "1.9"

M23_READER = [
    "Finance reader",
    "Read AP aging, the reconciliation, GR/IR, cash requirements, spend analysis "
    "and cycle time.",
    "procurement.analytics.view, at Procurement Advanced",
]

M23_FR001_EVIDENCE_TAIL = (
    " The vendor an order names is resolved inside the entity and the caller's "
    "branches: a vendor with no branch is shared and orderable from every branch, "
    "and one another branch keeps to itself is refused exactly like one that does "
    "not exist. The same resolver serves a goods receipt, a bill, a payment, the "
    "duplicate-reference check, the eligible-invoice picker and the AP aging "
    "vendor drawer. Once approval has released the order's email to the vendor, "
    "with a delivery PENDING or SENT, reversing that approval is refused with "
    "REVERSAL_NOT_ALLOWED and the order has to be cancelled instead, where the "
    "vendor is told; before then a reversed vote returns the order to "
    "PENDING_APPROVAL."
)
M23_FR001_ACCEPTANCE_TAIL = (
    " ProcurementCatalogueReadingTests.test_a_purchase_cannot_be_raised_against_another_sites_vendor "
    "and test_the_school_wide_vendor_is_still_orderable_from_every_site cover the "
    "vendor resolver."
)
M23_FR001_LIMIT = (
    "No procurement test exercises the refusal to reverse an order already "
    "released to its vendor; the engine's reversal contract is tested with a "
    "stand-in handler."
)
M23_FR009_LIMIT = (
    "Reversing the approval behind a bill or a payment that has already posted is "
    "not refused: the approval overlay returns to PENDING while the ledger and "
    "the document's own status say it posted. See Section 9."
)
M23_FR012_EVIDENCE_TAIL = (
    " Every one of these reads rides procurement.analytics.view, sold at "
    "Procurement Advanced."
)

M23_REVERSAL_ERROR = [
    "Reversing the approval of an order already released to its vendor",
    "Refused with REVERSAL_NOT_ALLOWED; cancel the order instead.",
]
M23_BRANCH_ERROR = [
    "Naming a vendor another branch keeps to itself",
    "Refused exactly like a vendor that does not exist.",
]

M23_MODULE7_TAIL = (
    " Its handlers answer the engine's reversal contract: an order already "
    "released to its vendor refuses reversal, and any other reversed document "
    "returns to PENDING."
)
M23_MODULE21_TAIL = (
    " The vendor resolver reads a null branch as shared across the tenant."
)
M23_MODULE6 = [
    "Module 6, Configuration & Capability Management",
    "Holds the depth this module is sold at. Purchase orders and goods receipt "
    "answer to Procurement's Core band, vendor bills and vendor payments to Plus, "
    "and the AP and GR/IR analytics to Advanced.",
]

M23_POSTED_REVERSAL_GAP = (
    "• Reversing the approval behind a posted bill or payment is not refused. "
    "Only a purchase order already released to its vendor answers the engine's "
    "reversal check, so an administrator can reverse the vote behind a payment "
    "that has posted: its approval reads PENDING again while the ledger says it "
    "was paid, and a fresh rejection would leave a posted payment reading "
    "REJECTED. Required completion: refuse reversal once a bill or payment has "
    "posted."
)

M23_TRACEABILITY_LEAD = (
    f"Module 23 carries 23 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Refusing to reverse an order already released to "
    "its vendor strengthens workflow approval integration without changing the "
    "count."
)

M23_CHANGE = (
    "Records three changes and removes one stale gap. Reversing the approval of a "
    "purchase order already released to its vendor is refused with "
    "REVERSAL_NOT_ALLOWED, because a vendor who has read that the order is "
    "approved is already acting on it, and the order has to be cancelled "
    "instead; any other reversed approval returns the document to PENDING. A new "
    "gap records the part that is not covered: reversing the approval behind a "
    "posted bill or payment is not refused, so its approval reads PENDING again "
    "while the ledger says it was paid. The vendor an order, receipt, bill or "
    "payment names is resolved inside the caller's branches as well as the "
    "entity, reading a null branch as shared. AP aging, the reconciliation, GR/IR, "
    "cash requirements, spend analysis and cycle time are recorded under "
    "procurement.analytics.view at Procurement Advanced, and a Module 6 "
    "dependency states the bands: orders and receipts at Core, bills and payments "
    "at Plus. The attachment-URL bullet leaves Further Gaps, since FR-014 and v1.5 "
    "already record that each file is authorised per read. The traceability lead "
    "is corrected from 22 to the 23 entries its table and the MRD carry. "
    f"{DEPLOYMENT}"
)


def patch_m23(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M23 Purchase Orders Delivery and AP Functional Requirements "
        f"Document v{M23_TARGET}"
    )

    actors = find_table(doc, "Actor")
    fr001, fr009, fr012 = find_fr(doc, "FR-001"), find_fr(doc, "FR-009"), find_fr(doc, "FR-012")
    errors = find_table(doc, "Condition or route")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M23_SOURCE, M23_TARGET)
    update_control(doc, version=M23_TARGET, module=23)

    reader = find_row(actors, "Finance reader")
    for cell, value in zip(unique_cells(reader), M23_READER):
        set_cell_text(cell, value)

    append_cell_text(field(fr001, "Current evidence"), M23_FR001_EVIDENCE_TAIL)
    append_cell_text(field(fr001, "Acceptance"), M23_FR001_ACCEPTANCE_TAIL)
    set_cell_text(field(fr001, "Current limit"), M23_FR001_LIMIT)
    set_cell_text(field(fr009, "Current limit"), M23_FR009_LIMIT)
    append_cell_text(field(fr012, "Current evidence"), M23_FR012_EVIDENCE_TAIL)

    last = errors.rows[-1]
    clone_row(errors, last, M23_BRANCH_ERROR)
    clone_row(errors, last, M23_REVERSAL_ERROR)

    append_cell_text(find_row(dependencies, "Module 7").cells[1], M23_MODULE7_TAIL)
    append_cell_text(find_row(dependencies, "Module 21").cells[1], M23_MODULE21_TAIL)
    clone_row(dependencies, dependencies.rows[-1], M23_MODULE6)

    gap_cell = gaps.rows[0].cells[0]
    lines = [p.text for p in gap_cell.paragraphs]
    stale = [line for line in lines if line.startswith("• An attachment URL is a capability")]
    if len(stale) != 1:
        raise ValueError("Expected exactly one stale attachment-URL bullet")
    lines.remove(stale[0])
    lines.insert(1, M23_POSTED_REVERSAL_GAP)
    rewrite_box(gap_cell, lines)

    set_run_text(body_paragraph(doc, "Module 23 carries"), M23_TRACEABILITY_LEAD)
    prepend_change_log(doc, M23_TARGET, M23_CHANGE)
    finish(doc, output, title, M23_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 24 - Inventory & Stock Ledger
# ═════════════════════════════════════════════════════════════════════════════

M24_DIR = "24-inventory-and-stock-ledger"
M24_STEM = "XVS_M24_Inventory_and_Stock_Ledger_Functional_Requirements_Document"
M24_SOURCE, M24_TARGET = "1.1.1", "1.2"

M24_SUPPORTING = "vs_finance, vs_rbac, vs_audit, vs_tenants"

M24_READER = [
    "Stores reader",
    "Read items, stock locations, per-location balances, movements, and the "
    "valuation and reorder views of the item list.",
    "procurement.stock.view",
]
M24_MANAGER_DOES = "Maintain item master records and stock locations."

M24_FR008_EVIDENCE = (
    "Valuation and reorder are read from the item list and its summary rather "
    "than from separate reports. GET /stock-items/?needs_reorder=true lists the "
    "active items at or below their reorder level, and GET /stock-items/summary/ "
    "returns item counts, low-stock and out-of-stock counts and the carried value "
    "in kobo, alongside the movement ledger. Both accept a location, which makes "
    "every row, and the summary, report that store's quantity, value, unit cost "
    "and reorder state; without one they report the entity roll-up."
)
M24_FR008_ACCEPTANCE = (
    "Valuation is read from the maintained value rather than recomputed from the "
    "movement history, and the two agree because every movement maintains it. "
    "The entity roll-up is what reconciles to the inventory control account; a "
    "single location's total does not, and is not meant to. "
    "StockLocationTests.test_the_item_list_narrowed_to_a_store_reports_that_store "
    "and test_needs_reorder_is_measured_against_the_store_in_force cover the store "
    "filter, and StockConsoleAPITests.test_summary_counts_states_and_is_entity_scoped "
    "the summary."
)
M24_FR008_LIMIT = (
    "Without a location the list, the summary, the balance list and the movement "
    "ledger report every store the entity holds, whatever branch the caller is "
    "pinned to. See Section 9."
)

M24_FR009_EVIDENCE_TAIL = (
    " A stock location is master data and reads a null branch inclusively: a "
    "location with no branch is a store shared across the tenant and reachable "
    "from every branch, and a caller pinned to a branch also reaches that "
    "branch's own. The location list, detail and update routes, and the resolver "
    "every location reference passes through on an issue, an adjustment, the "
    "balance list, the item list, the summary and the movement ledger, withhold a "
    "store another branch keeps to itself and report it exactly like one that "
    "does not exist. Creating or moving a location takes its branch from the "
    "caller: one pinned to a single branch stamps it, one covering several who "
    "names none files a shared store, and naming a branch the caller does not "
    "work in is refused."
)
M24_FR009_ACCEPTANCE_TAIL = (
    " ProcurementCatalogueReadingTests.test_the_central_store_is_still_reachable_from_every_site, "
    "test_another_sites_store_is_not_readable_or_editable_by_id and "
    "test_both_lists_carry_the_shared_row_and_this_site_and_no_other cover the "
    "reading."
)

M24_LOCATION_MODEL = [
    "StockLocation",
    "Where stock physically sits: entity, optional branch, code, name, and the "
    "default flag.",
    "One default per entity. A null branch is a store shared across the tenant.",
]
M24_BALANCE_MODEL = [
    "StockBalance",
    "One item's quantity and value at one location.",
    "Carries that location's weighted average; the item's own totals are the "
    "roll-up across locations.",
]

M24_ITEMS_ROUTE = (
    "The item master. ?location= reports one store's figures, and "
    "?needs_reorder=true lists what is running out."
)
M24_SUMMARY_ROUTE = (
    "Counts, low and out of stock, and carried value, for the entity or one store."
)
M24_LOCATION_ROUTE = [
    "GET, POST /stock-locations/, /stock-locations/{id}/",
    "The stores stock sits in, shared or one branch's own.",
]
M24_BALANCE_ROUTE = [
    "GET /stock-balances/",
    "What each item holds at each location.",
]

M24_LOCATION_ERROR = [
    "Naming a stock location another branch keeps to itself",
    "Reported exactly like a location that does not exist.",
]
M24_UNNAMED_ERROR = [
    "Moving stock without naming a location when the entity has more than one",
    "Refused rather than defaulted.",
]

M24_MODULE4 = (
    "Supplies the separately grantable view, manage, issue and adjust keys, and "
    "the branch narrowing stock locations apply inclusively."
)
M24_MODULE6 = [
    "Module 6, Configuration & Capability Management",
    "Holds the depth this module is sold at. Every stock key answers to "
    "Procurement's Plus band.",
]

M24_UNNARROWED_GAP = (
    "• The balance list, the movement ledger, the item list and its summary are "
    "narrowed only when a location is named. Asked without one they report every "
    "store the entity holds, so a storekeeper pinned to one branch, who cannot "
    "open another branch's store by id, can still read its balances and every "
    "movement at it by leaving the filter off."
)

M24_TRACEABILITY_LEAD = (
    f"Module 24 carries 13 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below."
)

M24_CHANGE = (
    "Reconciles the document with the stock code as it stands. The valuation and "
    "reorder reports it listed were folded into the item list on 15 August: "
    "GET /stock-items/?needs_reorder=true answers what is running out and "
    "GET /stock-items/summary/ what the stock is worth, both for the entity or "
    "one store, and FR-008, the actors table and the route table now say so. The "
    "stock location and balance routes, and the StockLocation and StockBalance "
    "records, are added to the contracts they were missing from. Stock locations "
    "read a null branch as shared across the tenant: the location list, detail "
    "and update routes and every location reference withhold a store another "
    "branch keeps to itself, and a new location takes its branch from the caller "
    "(FR-009). A new gap records what that does not reach: the balance list, the "
    "movement ledger and the item list and summary are narrowed only when a "
    "location is named, so a branch-pinned storekeeper can read another branch's "
    "store by leaving the filter off. A Module 6 dependency records that every "
    "stock key sits at Procurement Plus, and vs_schools, which this engine does "
    "not import, leaves the supporting apps. Module 24 stays at 13 capabilities. "
    f"{DEPLOYMENT}"
)


def patch_m24(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        f"XVS M24 Inventory and Stock Ledger Functional Requirements Document v{M24_TARGET}"
    )

    actors = find_table(doc, "Actor")
    fr008, fr009 = find_fr(doc, "FR-008"), find_fr(doc, "FR-009")
    model = find_table(doc, "Model")
    routes = find_table(doc, "Method and path")
    errors = find_table(doc, "Condition")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M24_SOURCE, M24_TARGET)
    update_control(doc, version=M24_TARGET, module=24, supporting=M24_SUPPORTING)

    reader = find_row(actors, "Stores reader")
    for cell, value in zip(unique_cells(reader), M24_READER):
        set_cell_text(cell, value)
    set_cell_text(find_row(actors, "Stores manager").cells[1], M24_MANAGER_DOES)

    set_cell_text(field(fr008, "Current evidence"), M24_FR008_EVIDENCE)
    set_cell_text(field(fr008, "Acceptance"), M24_FR008_ACCEPTANCE)
    set_cell_text(field(fr008, "Current limit"), M24_FR008_LIMIT)
    append_cell_text(field(fr009, "Current evidence"), M24_FR009_EVIDENCE_TAIL)
    append_cell_text(field(fr009, "Acceptance"), M24_FR009_ACCEPTANCE_TAIL)

    item = find_row(model, "StockItem")
    clone_row(model, item, M24_BALANCE_MODEL)
    clone_row(model, item, M24_LOCATION_MODEL)

    items_route = find_row(routes, "GET, POST /stock-items/")
    set_cell_text(items_route.cells[1], M24_ITEMS_ROUTE)
    set_cell_text(find_row(routes, "GET /stock-items/summary/").cells[1], M24_SUMMARY_ROUTE)
    delete_row(find_row(routes, "GET /reports/stock-valuation/"))
    delete_row(find_row(routes, "GET /reports/stock-reorder/"))
    clone_row(routes, items_route, M24_LOCATION_ROUTE, before=True)
    clone_row(routes, items_route, M24_BALANCE_ROUTE, before=True)

    last = errors.rows[-1]
    clone_row(errors, last, M24_UNNAMED_ERROR)
    clone_row(errors, last, M24_LOCATION_ERROR)

    set_cell_text(find_row(dependencies, "Module 4").cells[1], M24_MODULE4)
    clone_row(dependencies, dependencies.rows[-1], M24_MODULE6)

    gap_cell = gaps.rows[0].cells[0]
    rewrite_box(gap_cell, [p.text for p in gap_cell.paragraphs] + [M24_UNNARROWED_GAP])

    set_run_text(body_paragraph(doc, "Module 24 carries"), M24_TRACEABILITY_LEAD)
    prepend_change_log(doc, M24_TARGET, M24_CHANGE)
    finish(doc, output, title, M24_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 10 - Bulk Data Import
# ═════════════════════════════════════════════════════════════════════════════

M10_DIR = "10-bulk-data-import"
M10_STEM = "XVS_M10_Bulk_Data_Import_Functional_Requirements_Document"
M10_SOURCE, M10_TARGET = "1.2", "1.3"

M10_SUPPORTING = (
    "vs_tenants, vs_schools, vs_calendar, vs_students, vs_academics, vs_staff, "
    "vs_user, vs_rbac, vs_workflow, vs_audit, vs_finance, vs_notifications, core"
)

M10_BOUNDARY_BULLET = (
    "• This module is the platform's only bulk write path. For CodeX it creates "
    "schools, branches and CodeX staff accounts through the same serializers the "
    "single-record API uses; for a school it loads the school's own calendar, "
    "students, staff, guardians, classes and subjects through the owning module's "
    "import resolver and creation services. Either way its guards are the owning "
    "module's guards restated at a second door."
)

M10_PURPOSE = (
    "Module 10 is how records arrive in bulk. An operator picks an official "
    "template, uploads a CSV or XLSX file, has every column and row checked "
    "against that template before anything is written, then runs the import and "
    "watches it row by row. Its organising commitment is that a file is never a "
    "shortcut: a CodeX row runs through the same serializer the single-record API "
    "uses, and a school's row through the owning module's own resolver and "
    "creation services, which validation and execution both call, so a rule "
    "enforced at the front door is enforced again here and a spreadsheet cannot "
    "buy a permission a request would be refused."
)

M10_UPLOAD_SCOPE = (
    "CSV and XLSX intake, header extraction, a stored preview snapshot of the "
    "parsed rows that every later step reads, and the branch the batch belongs to."
)
M10_EXECUTION_SCOPE = (
    "A background job per run, and a stored result for every row saying what was "
    "created, skipped or refused. Most datasets commit each row in its own "
    "savepoint; academic structure, subjects and guardians publish the whole file "
    "in one transaction, because half of one is worse than none."
)
M10_OWNER_SCOPE = (
    "The owning module. This engine calls SchoolCreateSerializer, "
    "BranchCreateSerializer and UserCreateSerializer for CodeX's datasets, and the "
    "calendar, students, staff, academics and guardians modules' own import "
    "resolvers for a school's, and holds no domain rules of its own."
)

M10_BOUNDARIES = (
    "Two boundaries govern every request. The first is the dataset: platform_only "
    "classifies each dataset type and fails closed, so a dataset added to the "
    "choices and forgotten is withheld from schools rather than handed to them. "
    "The three layers that can act on a dataset all ask the same function, "
    "because the template list is a courtesy, batch creation is the rule, and the "
    "executor is what catches a batch built before the rule existed or reached by "
    "a path nobody has thought of. The second boundary is the tenant and the "
    "branch: every read resolves its batch through one lookup that filters on the "
    "caller's asserted tenant and their branch reach, and issues, jobs, rows, "
    "audit events, notifications, rollback records and the uploaded file itself "
    "are all reached by first resolving that batch. The gate that lets a module's "
    "own import key stand in for the generic ones narrows the same way before it "
    "admits anybody. Narrowing the single lookup therefore narrows all of them, "
    "and another tenant's or another branch's batch answers 404 rather than 403, "
    "which is the same answer as a batch that does not exist. A new batch is "
    "filed under the uploader's branch, so the narrowing has something to read: "
    "an uploader pinned to one branch stamps it, one holding the role for the "
    "whole school or covering several branches files it for the school as a "
    "whole, and naming a branch the uploader does not work in is refused. "
    "Platform staff are deliberately unscoped, because provisioning a school means "
    "writing outside every tenant."
)

M10_MODULE_KEY_ACTOR = [
    "Holder of a module's import key",
    "Load that module's dataset through the wizard without the generic import keys.",
    "school.students.import for students and guardians, academics.structure.import "
    "for academic structure and subjects, finance.bankaccount.import for bank "
    "statements. It opens only batches of that dataset, in the caller's tenant "
    "and branches.",
]

M10_FR001_EVIDENCE = (
    "datasets.py classifies each dataset type. PLATFORM_ONLY_DATASETS holds "
    "schools, cx_users, bank_statements and branches; TENANT_DATASETS holds the six "
    "a school arrives with: calendar events, students, staff, academic structure, "
    "subjects and guardians, each admitted on the same three counts (the school "
    "already creates these through keys of its own, the handler writes nothing "
    "but the uploading tenant's rows, and there is a real reason to do it in "
    "bulk). platform_only fails closed, so a dataset added to the choices and not "
    "classified is withheld from schools. may_import is asked by the template "
    "list, by batch creation, and again by the executor for every row, so a batch "
    "built before the rule existed is refused at the point of the write. A "
    "refused row is skipped with a reason rather than raised, so the rest of the "
    "batch still reports. The template list, a template's detail, its download, "
    "the batch list, upload, detail, cancel and file download, validation, the "
    "issue list and export, the start of a run and the job list are open to a "
    "school still being set up, because loading its data is a step on its own "
    "checklist; template authoring stays shut. Reaching this engine at all is a "
    "second question, asked of the plan rather than the role: the navigation "
    "door, the onboarding card and the upload control each declare the "
    "bulk_import capability alongside the permission, and it sits at Core, so "
    "every plan reaches this engine."
)
M10_FR001_ACCEPTANCE = (
    "A school administrator listing templates is offered the six school datasets "
    "and none of the four platform ones. Naming a platform template in a crafted "
    "batch-create request is refused. A row reaching a handler with a platform "
    "dataset is skipped with the standard refusal message, which says nothing "
    "about what the dataset does. "
    "DatasetOwnershipRuleTests.test_the_school_datasets_are_the_six_a_school_arrives_with "
    "and test_every_dataset_type_is_classified hold the classification, "
    "ExecutorRefusesPlatformDatasetsTests the executor gate, and "
    "OnboardingImportSurfaceTests the surface a school still being set up reaches."
)
M10_FR001_LIMIT = (
    "Historical records, fee structures and timetables have no dataset type and "
    "no handler. The onboarding checklist step this serves, Upload Initial "
    "Datasets, is optional, so readiness and go-live are not blocked by it."
)

M10_FR002_EVIDENCE = (
    "ImportTemplate carries the dataset type, file format and status; "
    "ImportTemplateColumn carries each column's display name, target field, data "
    "type, required flag, choices, ordering and default. Validation reads the "
    "columns to build the expected headers, the row rules and the uniqueness "
    "rules; the executor reads the same columns to map a raw row onto the payload "
    "it hands the handler. seed_import publishes the schools, branches, CX users, "
    "bank statement, calendar and staff templates, and data migrations publish the "
    "students, academic structure, subjects and guardians templates, the students "
    "one carrying in its guidance the rules the file is checked against. A re-run "
    "of seed_import warns when another ACTIVE template serves the same dataset, "
    "because it matches on code and a renamed code leaves the old template "
    "answering beside the new one. The schools template's Package Plan column "
    "decides how deep the school reaches into every module rather than which "
    "modules it gets, and neither the schools nor the branches template carries a "
    "branch type, which the branch record no longer has. Creating a template needs "
    "import.templates.create; the internal configuration fields are gated "
    "separately behind the restricted import.templates.manage."
)
M10_FR002_ACCEPTANCE = (
    "A template's columns are the single source of both the validation vocabulary "
    "and the executor's field mapping. A school holding only templates.view can "
    "read a template, including while it is still being set up, and download its "
    "file, and can change nothing about it."
)
M10_FR002_LIMIT_TAIL = (
    " The duplicate-template check warns and retires nothing, because a dataset "
    "may legitimately be offered through two templates; retiring a renamed one is "
    "left to an operator."
)

M10_FR003_EVIDENCE = (
    "ImportBatch.clean refuses any format outside CSV, XLSX and XLS. The parser "
    "extracts uploaded_headers and writes preview_rows onto the batch; validation "
    "reads preview_rows, and so does execute_import. The uploaded file is written "
    "through the platform's database storage, which reads the upload through its "
    "chunks, so a file the parser has already read to the end is still stored "
    "whole; a bare read had stored every uploaded file with no bytes while its "
    "recorded size looked right. A post_delete signal removes the stored file when "
    "the batch is deleted."
)
M10_FR003_ACCEPTANCE = (
    "A file of an unsupported format is refused at creation. The rows validated "
    "are byte-for-byte the rows imported, because both read the stored snapshot. "
    "DatabaseStorageRewindTests proves a file already read to the end still stores "
    "its bytes and that the recorded size matches what was written."
)

M10_FR004_EVIDENCE = (
    "validate_import_batch runs template presence, headers against the template, "
    "rows against each column's type and choices, uniqueness rules declared on "
    "the template, dataset-specific rules, and cross-reference checks that "
    "resolve named records. Dataset-specific rules for schools and branches live "
    "here; the rest belong to the dataset's module, and the calendar, students, "
    "staff, academic structure, subjects and guardians each validate through the "
    "same resolver their execution calls, with bank statements checked by "
    "finance. The student rules refuse everything the enrol form refuses and "
    "more: a birth year that makes the child under two or over twenty-five, an "
    "admission date in the future or before the child was born, a guardian email "
    "that is not an address, a phone with fewer than seven digits, and any value "
    "longer than its column. Before anything is written they warn about a class "
    "the file fills past its capacity, a guardian contact used under two names, "
    "and a digit inside a name. Findings are written as ImportValidationIssue rows "
    "carrying row number, column, severity, a code from a closed vocabulary of "
    "thirteen, and a message."
)
M10_FR004_ACCEPTANCE = (
    "Each finding names the row and column an operator has to fix. The issue list "
    "is readable through the API and downloadable as a file. "
    "ImportCatchesWhatTheFormCatchesTests covers the student rules, "
    "WholeFileFaultsTests and RowFaultsTests the academic structure, "
    "SubjectImportTests the subjects and GuardianImportTests the guardians; the "
    "calendar's ValidationTests and the staff import's RowRefusalTests cover "
    "theirs."
)
M10_FR004_LIMIT = (
    "A CX users file is checked for structure, types and uniqueness only; its "
    "business rules are enforced when the row executes."
)

M10_FR006_EVIDENCE = (
    "execute_import creates an ImportJob, then iterates the stored rows inside its "
    "own transaction per row, so a crash on row 400 does not roll back rows 1 to "
    "399. Academic structure, subjects and guardians are the exception: their "
    "module validates the file as a whole and publishes it in one transaction, "
    "because half a structure, or a primary contact moved before the row meant to "
    "replace it fails, is worse than nothing. Each row writes an "
    "ImportJobRowResult holding the action taken, the model it created and its "
    "primary key where there is one, the raw row, the normalised payload, and "
    "either a status message or structured error details. The job carries "
    "progress percent and counters for processed, succeeded, failed and skipped "
    "rows, and the run is dispatched to a tracked Celery task that names the batch "
    "as its subject, so the request returns immediately."
)

M10_FR007_EVIDENCE = (
    "For CodeX's datasets the three handlers call SchoolCreateSerializer, "
    "BranchCreateSerializer and UserCreateSerializer with a request-like context "
    "naming the acting user. An imported CodeX hire is created PENDING_APPROVAL "
    "through UserCreationService and submitted into the platform-user approval "
    "workflow, exactly like a single add, and the CX handler names the platform "
    "tenant explicitly rather than inheriting the queuer's, because the "
    "serializer's fallback would otherwise create a CodeX hire inside a school. A "
    "school's datasets take their tenant from the batch and nowhere else, and "
    "their templates carry no school column. Interpretation lives in the owning "
    "module: the calendar, students and staff handlers call the resolver their "
    "validation called and write through that module's services, so a student is "
    "placed and matched to a guardian by the rules an enrolment uses, and a "
    "member of staff is created through the same user creation and invitation a "
    "single add uses."
)
M10_FR007_ACCEPTANCE = (
    "A rule enforced by a create serializer or a module's own service is enforced "
    "for an imported row without being restated here. An imported hire is "
    "indistinguishable from a hire added one at a time, including its approval "
    "request, and an imported student from one enrolled on the form. The "
    "calendar's ExecutionTests.test_the_uploading_school_is_the_only_school_it_can_write_to "
    "proves a school dataset cannot write outside the uploading tenant."
)
M10_FR007_LIMIT = (
    "Duplicate detection is per dataset: schools match on slug then name, branches "
    "on name within the tenant, CX users on email within the platform tenant, "
    "calendar entries on name and date, staff on email within the school, and a "
    "student number repeated in one file is refused. A near-duplicate that "
    "differs in punctuation is imported as a new record."
)

M10_FR008_EVIDENCE = (
    "Reversal is dispatched through a registry keyed on the row's target_model, "
    "with reversers for School, Branch, User and CalendarEvent; a model absent "
    "from the registry is refused rather than guessed at, and a row whose action "
    "was not a creation is refused because there is nothing it made. Each "
    "reverser locks its target with select_for_update and then re-checks the "
    "natural key the row recorded, the school's slug, the branch's name, the "
    "user's email or the event's name, against the record the id resolves to. "
    "Branch ownership is checked against the school the row itself named, because "
    "the provisioning datasets are platform-only and the batch's tenant is CodeX "
    "rather than the school the row created. A calendar entry is refused once an "
    "exam timetable has been built against it, as the events API refuses the "
    "same delete."
)
M10_FR008_ACCEPTANCE = (
    "A branches run whose rows created branch ids 9 to 12 deletes those four "
    "branches and no school. A recorded id whose record has since been renamed, "
    "or which now belongs to another school, is refused and left alone. The "
    "calendar's RollbackTests prove an entry goes with its audience and is "
    "refused when an exam timetable hangs off it."
)
M10_FR008_LIMIT = (
    "Rows written before the school primary key moved off the slug hold a "
    "non-numeric reference and are refused rather than matched on a guess. A "
    "Student, StaffProfile, StudentGuardian, SchoolClass or Subject has no "
    "reverser, so a school's students, staff, guardians, academic structure and "
    "subjects cannot be rolled back through this engine: every such row is "
    "refused with the model named, and a file imported in error has to be "
    "corrected through the owning module."
)

M10_FR011_EVIDENCE = (
    "One mixin resolves the batch, filtering on the caller's asserted tenant and "
    "their branch reach, and returning 404 rather than 403 so a batch belonging to "
    "somebody else is indistinguishable from one that does not exist. Validation "
    "issues, jobs, job details, row results, audit events, notifications, rollback "
    "records and the uploaded file are all reached by first resolving that batch, "
    "so the single narrowing carries to all of them; the file download, which had "
    "a tenant-only lookup of its own, now resolves through the same one. The gate "
    "that lets a module's own import key stand in for the generic ones narrows by "
    "branch as well, so the gate and the lookup cannot disagree about one id. A "
    "batch with no branch belongs to the school as a whole and is readable from "
    "every branch. The write half matches: the upload serializer requires the "
    "branch in its context and will not guess, the batch endpoint supplies the "
    "uploader's, and finance's statement wizard supplies the branch of the bank "
    "account the statement continues. Platform callers are deliberately unscoped, "
    "and ?school= inputs no longer influence scoping: the asserted tenant is "
    "authoritative."
)
M10_FR011_ACCEPTANCE = (
    "A school administrator requesting another school's batch by id receives 404, "
    "and so does a branch administrator requesting another branch's, on every "
    "nested route and on the file itself. ImportBatchFileDownloadBranchScopeTests "
    "and ImportBatchModuleKeyBranchScopeTests cover the read half; "
    "ImportBatchUploadBranchTests covers the write half, including an Ikeja upload "
    "the Lekki administrator cannot download; and the statement wizard's "
    "test_the_batch_takes_the_branch_of_the_account_it_continues covers the "
    "finance path."
)
M10_FR011_LIMIT = (
    "Platform staff see every tenant's imports with no per-tenant narrowing "
    "available in the module itself, which is correct for provisioning and means "
    "the module cannot grant a regional operator sight of only their own region. "
    "Batches uploaded before the branch was recorded carry none, so they read as "
    "belonging to the whole school and stay visible from every branch; nothing "
    "back-fills them."
)

M10_FR012_LIMIT = (
    "Cancellation sets a status and clears readiness. It does not interrupt a "
    "validation run already in flight, so a run can complete after the batch was "
    "cancelled and move it back to a validated status."
)

M10_FR014_EVIDENCE = (
    "ImportNotification rows are written when validation completes, when an "
    "import job completes, and when a rollback completes, addressed to the "
    "uploader or to the person who initiated the rollback. The rollback notice "
    "distinguishes a complete reversal from a partial one and states how many "
    "rows could not be reversed. The notification feed for a batch is readable by "
    "platform staff. Separately, the tracked job that runs an import notifies its "
    "owner through Module 8 when it finishes, as task.completed or task.failed on "
    "the in-app channel, and the notice names the batch so it opens on that "
    "batch's page."
)
M10_FR014_ACCEPTANCE = (
    "The person who started a run receives an in-app notice naming the file and "
    "the outcome, which leads back to the batch."
)
M10_FR014_LIMIT = (
    "The module's own rows are not dispatched through Module 8: the task that "
    "sends them marks them sent and delivers nothing, and their feed requires a "
    "platform key, so a school administrator never sees the validation or "
    "rollback notice. A rollback run in the background notifies nobody through "
    "Module 8, since its job kind is not one that notifies by default, and there "
    "is no email for either."
)

M10_FR015_EVIDENCE = (
    "The batch file download route is gated on import.batches.view and resolves "
    "the batch through the same scoped lookup as every other read, tenant and "
    "branch both. It reads the bytes through the file's own storage, because the "
    "database storage has no filesystem path and asking for one had made every "
    "download answer 500. The stored file is bound to the school that uploaded it "
    "and to the batch it belongs to, so the bytes cannot be fetched by naming "
    "them: a read is refused unless the caller's asserted school matches the "
    "file's, and the URL that carries it is signed for one person and expires. A "
    "post_delete signal removes the stored file when the batch row is deleted. "
    "Template files are separately downloadable under import.templates.view so "
    "an operator can start from the official shape."
)
M10_FR015_ACCEPTANCE_TAIL = (
    " An Ikeja administrator asking for a Lekki batch's file receives 404."
)
M10_FR015_LIMIT_TAIL = (
    " A batch uploaded before the storage read its upload through its chunks "
    "holds an empty file, and those bytes cannot be recovered; the file has to be "
    "uploaded again."
)

M10_STEP_UPLOAD = (
    "The file is stored, its headers are extracted, its rows are parsed into the "
    "batch's preview snapshot, and the batch is filed under the uploader's branch, "
    "or the school as a whole where they have no single branch to name. Status "
    "becomes uploaded."
)
M10_STEP_EXECUTE = (
    "Each row is mapped to a payload, run through its dataset handler inside its "
    "own savepoint, and recorded as created, skipped or failed. Academic "
    "structure, subjects and guardians are published as one transaction instead."
)
M10_STEP_FINISH = (
    "The job's counters and status settle, an audit event and a notification are "
    "written, the uploader's bell is told through Module 8, and the batch reflects "
    "the outcome."
)

M10_ROLLBACK_TAIL = (
    " A calendar entry may be reversed unless an exam timetable has been built "
    "against it. A school's students, staff, guardians, classes and subjects have "
    "no reverser at all, so a rollback of those files refuses every row and names "
    "the model it cannot reverse."
)

M10_BATCH_MODEL = (
    "One uploaded file: tenant, optional branch (the uploader's, or none for the "
    "school as a whole), uploader, template, dataset type, the stored file, parsed "
    "headers and preview rows, validation summary and readiness flags, and its "
    "status."
)

M10_API_TAIL = (
    " A school still being set up reaches the template list, detail and download, "
    "the batch list, upload, detail, cancel and file download, validation, the "
    "issue list and export, the start of a run and the job list, because loading "
    "its data is a step on its own checklist; template authoring stays shut to it."
)

M10_CELERY = (
    "Runs execution, and any rollback past fifty rows, as tracked jobs that name "
    "their batch; validation runs inside the request. With eager execution "
    "configured, a task failure propagates to the caller and is returned as a "
    "real error rather than a broker message."
)
M10_NEW_DEPENDENCIES = [
    [
        "Module 8, Notifications & Delivery",
        "Delivers the in-app task.completed or task.failed notice when an import "
        "run finishes, linked to its batch. The module's own ImportNotification "
        "rows do not pass through it.",
    ],
    [
        "Module 11, Student Management",
        "Owns the students and guardians imports: the resolver both passes call, "
        "the creation that places a student and matches a guardian as an "
        "enrolment does, and the school.students.import key that opens the wizard "
        "for those two datasets.",
    ],
    [
        "Module 12, Staff Management",
        "Owns the staff import: the resolver both passes call, and the creation "
        "that invites the person through the same user creation a single add "
        "uses. It registers no key of its own, so the generic import keys govern "
        "it.",
    ],
    [
        "Module 13, Academic Structure",
        "Owns the academic structure and subjects imports, each published as one "
        "transaction, and the academics.structure.import key that opens the "
        "wizard for both.",
    ],
    [
        "Module 14, Timetable & Calendar",
        "Owns the calendar import, and the rule that a calendar entry with an exam "
        "timetable built against it cannot be rolled back.",
    ],
]
M10_EVIDENCE = (
    "Operational evidence: the app's own suite covers the dataset-ownership gates, "
    "the validation publish gate, the tenant and branch scoping of reads, uploads "
    "and the file download, the admin email scope, the surface open to a school "
    "still being set up, and the rollback behaviour described in FR-008 through "
    "FR-010 and FR-017, including the case where a school and a branch share a "
    "primary key and the queued path taken when a rollback is too large to answer "
    "inside the request. Each school dataset's rules and execution are covered in "
    "its owning module's suite, named against each requirement above. This "
    "revision was written from the code and the tests at the baseline named in "
    "Document Control, not from a fresh run."
)

M10_NEEDS_ATTENTION = [
    "NEEDS ATTENTION",
    "• Three things a school might reasonably bring have no dataset: historical "
    "records, fee structures and timetables. None has a dataset type or a "
    "handler, so the onboarding screen shows historical records greyed and does "
    "not mention the other two. The six that do exist - students, staff, "
    "guardians, academic structure, subjects and calendar events - are seeded and "
    "reachable on every plan. The onboarding checklist item it serves is "
    "optional, and is absent altogether for a school whose plan does not reach "
    "bulk import, so readiness and go-live are not blocked either way.",
    "• Rollback reverses what CodeX imports and calendar entries, and nothing "
    "else. A school's students, staff, guardians, classes and subjects have no "
    "reverser, so a roll imported in error cannot be unwound here: every row is "
    "refused with the model named, and the records have to be corrected through "
    "the owning module.",
    "• The module's own import notifications are not dispatched through Module 8, "
    "and their feed requires a platform key, so a school administrator never sees "
    "the validation or rollback notice. The run itself does reach the uploader's "
    "bell as an in-app task notice linked to the batch; a rollback run in the "
    "background reaches nobody's.",
    "• A run cannot be stopped once it has started. Cancellation covers only the "
    "statuses before execution, and the row loop checks nothing, so a long import "
    "of a wrong file has to be allowed to finish before anything can be done "
    "about it.",
    "• Cancellation does not interrupt a validation run already in flight, so the "
    "run can complete after a batch was cancelled and move it back to a "
    "validated status.",
    "• Execution reads the batch's stored preview rows, and there is no streaming "
    "path, so the practical file size is bounded by one JSON column and one "
    "in-memory list. No limit is stated or enforced.",
    "• The two rollback paths report in different shapes: the finance path stores "
    "a line count and always records success, while the general path stores "
    "per-row outcomes and can record a partial result.",
    "• An uploaded roster is the most sensitive file this platform holds, and "
    "reading one back is bound to the school, the branch and the batch. What is "
    "not covered is archiving: only deletion retires the file, so a batch retired "
    "by archiving keeps its spreadsheet readable to whoever import.batches.view "
    "still admits.",
]

M10_TRACEABILITY_LEAD = (
    f"Module 10 of XVS Module Requirements Document v{MRD_VERSION} lists sixteen "
    "capabilities. Each maps to the requirements above."
)
M10_TRACEABILITY_NOTE = (
    "The MRD carries five needs-attention items for Module 10: the datasets a "
    "school still cannot bring, rollback that cannot reverse a school's own "
    "datasets, notifications that do not reach Module 8, a run that cannot be "
    "stopped, and an unbounded file size. Section 9 states those five and adds "
    "three the MRD does not carry, because they are narrower than the roadmap "
    "tracks: a validation run in flight is not interrupted by cancelling its "
    "batch, the two rollback paths report in different shapes, and an archived "
    "batch keeps its file readable. Each is a boundary of a capability the module "
    "delivers."
)

M10_CHANGE = (
    "Reconciles the document with the six datasets a school may now import for "
    "itself, which v1.2 recorded in Section 9 while FR-001 still said there were "
    "none. Calendar events, students, staff, academic structure, subjects and "
    "guardians are classified as a school's own, each loaded through its owning "
    "module's resolver, which validation and execution both call; academic "
    "structure, subjects and guardians publish the whole file in one transaction. "
    "FR-001, FR-002, FR-004, FR-006, FR-007 and the scope, workflow and dependency "
    "sections are updated to match, with the student rules that refuse everything "
    "the enrol form refuses. Rollback reverses calendar entries as well as "
    "CodeX's datasets and cannot reverse a school's students, staff, guardians, "
    "classes or subjects, which FR-008 and a new needs-attention item record. A "
    "batch now records the branch it was uploaded for, and its file is refused "
    "from another branch, where the download had a tenant-only lookup of its own "
    "(FR-011). Uploaded files were stored with no bytes and every download "
    "answered 500; both are fixed (FR-003, FR-015). An import run's completion "
    "reaches the uploader's bell through Module 8 and opens on the batch, so the "
    "notifications item is rewritten (FR-014). A school still being set up can "
    "read a template's rules, and the seeder warns when a renamed template code "
    "leaves the old one answering. Module 10 stays at sixteen capabilities. "
    f"{DEPLOYMENT}"
)


def patch_m10(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS M10 Bulk Data Import Functional Requirements Document v{M10_TARGET}"

    boundary = find_table(doc, "EVIDENCE BOUNDARY")
    scope = find_table(doc, "Area")
    out_of_scope = find_table(doc, "Concern")
    actors = find_table(doc, "Actor")
    frs = {label: find_fr(doc, label) for label in (
        "FR-001", "FR-002", "FR-003", "FR-004", "FR-006", "FR-007", "FR-008",
        "FR-011", "FR-012", "FR-014", "FR-015",
    )}
    steps = find_table(doc, "Step")
    model = find_table(doc, "Model")
    dependencies = find_table(doc, "Module", contains="Celery")
    attention = find_table(doc, "NEEDS ATTENTION")
    traceability = find_table(doc, "MRD capability")

    replace_cover_version(doc.tables[0], M10_SOURCE, M10_TARGET)
    update_control(doc, version=M10_TARGET, module=10, supporting=M10_SUPPORTING)

    box = boundary.rows[0].cells[0]
    lines = [p.text for p in box.paragraphs]
    if not lines[-1].startswith("• This module is the platform's only bulk write path"):
        raise ValueError("Evidence boundary no longer ends with the bulk-write bullet")
    rewrite_box(box, lines[:-1] + [M10_BOUNDARY_BULLET])

    set_run_text(body_paragraph(doc, "Module 10 is how records arrive in bulk."), M10_PURPOSE)
    set_cell_text(find_row(scope, "Upload and parsing").cells[1], M10_UPLOAD_SCOPE)
    set_cell_text(find_row(scope, "Execution").cells[1], M10_EXECUTION_SCOPE)
    set_cell_text(find_row(out_of_scope, "What a created record means").cells[1], M10_OWNER_SCOPE)
    set_run_text(body_paragraph(doc, "Two boundaries govern every request."), M10_BOUNDARIES)

    correct_file = find_row(actors, "Correct a bad file", column=1)
    clone_row(actors, correct_file, M10_MODULE_KEY_ACTOR)

    fr = frs["FR-001"]
    set_cell_text(field(fr, "Current evidence"), M10_FR001_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR001_ACCEPTANCE)
    set_cell_text(field(fr, "Current limit"), M10_FR001_LIMIT)

    fr = frs["FR-002"]
    set_cell_text(field(fr, "Current evidence"), M10_FR002_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR002_ACCEPTANCE)
    append_cell_text(field(fr, "Current limit"), M10_FR002_LIMIT_TAIL)

    fr = frs["FR-003"]
    set_cell_text(field(fr, "Current evidence"), M10_FR003_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR003_ACCEPTANCE)

    fr = frs["FR-004"]
    set_cell_text(field(fr, "Current evidence"), M10_FR004_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR004_ACCEPTANCE)
    set_cell_text(field(fr, "Current limit"), M10_FR004_LIMIT)

    set_cell_text(field(frs["FR-006"], "Current evidence"), M10_FR006_EVIDENCE)

    fr = frs["FR-007"]
    set_cell_text(field(fr, "Current evidence"), M10_FR007_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR007_ACCEPTANCE)
    set_cell_text(field(fr, "Current limit"), M10_FR007_LIMIT)

    fr = frs["FR-008"]
    set_cell_text(field(fr, "Current evidence"), M10_FR008_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR008_ACCEPTANCE)
    set_cell_text(field(fr, "Current limit"), M10_FR008_LIMIT)

    fr = frs["FR-011"]
    set_cell_text(field(fr, "Current evidence"), M10_FR011_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR011_ACCEPTANCE)
    set_cell_text(field(fr, "Current limit"), M10_FR011_LIMIT)

    set_cell_text(field(frs["FR-012"], "Current limit"), M10_FR012_LIMIT)

    fr = frs["FR-014"]
    set_cell_text(field(fr, "Current evidence"), M10_FR014_EVIDENCE)
    set_cell_text(field(fr, "Acceptance"), M10_FR014_ACCEPTANCE)
    set_cell_text(field(fr, "Current limit"), M10_FR014_LIMIT)

    fr = frs["FR-015"]
    set_cell_text(field(fr, "Current evidence"), M10_FR015_EVIDENCE)
    append_cell_text(field(fr, "Acceptance"), M10_FR015_ACCEPTANCE_TAIL)
    append_cell_text(field(fr, "Current limit"), M10_FR015_LIMIT_TAIL)

    set_cell_text(find_row(steps, "Upload").cells[1], M10_STEP_UPLOAD)
    set_cell_text(find_row(steps, "Execute").cells[1], M10_STEP_EXECUTE)
    set_cell_text(find_row(steps, "Finish").cells[1], M10_STEP_FINISH)

    rollback = body_paragraph(doc, "Rollback is available for a job that succeeded")
    set_run_text(rollback, rollback.text.strip() + M10_ROLLBACK_TAIL)

    set_cell_text(find_row(model, "ImportBatch").cells[1], M10_BATCH_MODEL)

    refusals = body_paragraph(doc, "Refusals are consistent across the surface")
    set_run_text(refusals, refusals.text.strip() + M10_API_TAIL)

    set_cell_text(find_row(dependencies, "Celery").cells[1], M10_CELERY)
    # The Module 6 row arrived without the label styling every row above it
    # carries, so it and the rows added after it are rebuilt from a styled row.
    styled = find_row(dependencies, "Module 19")
    module6 = find_row(dependencies, "Module 6")
    anchor = clone_styled_row(
        dependencies, styled, module6,
        ["Module 6, Configuration & Capability Management", module6.cells[1].text.strip()],
    )
    delete_row(module6)
    for values in reversed(M10_NEW_DEPENDENCIES):
        clone_styled_row(dependencies, styled, anchor, values)
    set_run_text(body_paragraph(doc, "Operational evidence:"), M10_EVIDENCE)

    rewrite_box(attention.rows[0].cells[0], M10_NEEDS_ATTENTION)

    set_cell_text(find_row(traceability, "Batch rollback").cells[2], "Implemented with limits")
    set_run_text(body_paragraph(doc, "Module 10 of XVS Module Requirements Document"), M10_TRACEABILITY_LEAD)
    set_run_text(body_paragraph(doc, "The MRD carries"), M10_TRACEABILITY_NOTE)

    prepend_change_log(doc, M10_TARGET, M10_CHANGE)
    finish(doc, output, title, M10_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    parser.add_argument(
        "--only", nargs="*", choices=["m10", "m21", "m22", "m23", "m24"],
        help="Write only these modules; the default writes all five.",
    )
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for tag, folder, stem, source, target, patch in (
        ("m10", M10_DIR, M10_STEM, M10_SOURCE, M10_TARGET, patch_m10),
        ("m21", M21_DIR, M21_STEM, M21_SOURCE, M21_TARGET, patch_m21),
        ("m22", M22_DIR, M22_STEM, M22_SOURCE, M22_TARGET, patch_m22),
        ("m23", M23_DIR, M23_STEM, M23_SOURCE, M23_TARGET, patch_m23),
        ("m24", M24_DIR, M24_STEM, M24_SOURCE, M24_TARGET, patch_m24),
    ):
        if args.only and tag not in args.only:
            continue
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists():
            raise SystemExit(f"Refusing to overwrite {output.name}")
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
