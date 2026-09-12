#!/usr/bin/env python3
"""Version the Module 1, 4, 6 and 9 FRDs against the backend at 19a50b3.

Taking back the role grants a tenant can no longer reach was a step inside
``change_plan``, so only a tier move performed it. It now hangs off the reach
itself in vs_config, and the service that performs it lives in
``vs_rbac.plan_grants``. Every document that described the old arrangement is
corrected:

* Module 1 records that a tier change, an uplift written or withdrawn, an
  uplift that lapsed and any later settlement, and ``manage.py apply_plans``
  all reconcile; that only a band out of depth is taken back and a key denied
  on purpose stays denied; that a first plan application revokes nothing; and
  that a role whose own permission dependencies refuse the settlement is
  reported in ``data.roles_needing_attention`` and in the audit trail rather
  than failing the billing write. It also records that a School is refused
  rather than created short of the roles CodeX ships, and that a Branch role's
  name follows its Branch.
* Module 4 names the service, both audit sources, and the deny set the
  settlement now carries, and corrects the provisioning requirement for a
  library that cannot supply a template.
* Module 6 records that the settlement hangs off its own writes, and the one
  endpoint that reconciles with nowhere to report a role it could not settle.
* Module 9 records the widened creation refusal it receives from Module 1.

Each document is read from its latest version and written beside it as the
next minor version; no earlier version is touched. Traceability is set to MRD
v2.77, which is cut after these documents.

    python tools/patch_reach_settlement_docs.py [--only m01] [--overwrite]
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_requirements_documents import (
    VERY_PALE_BLUE,
    WHITE,
    assert_no_em_dash,
    set_cell_shading,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)

REVIEW_DATE = "12 September 2026"
SHORT_DATE = "12 Sep 2026"
MRD_VERSION = "2.77"
CODE_BASELINE = "Backend worktree at 19a50b3, 12 September 2026"


# ── shared docx helpers ──────────────────────────────────────────────────────

def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError("Paragraph carries no run to inherit formatting from")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def cell_size(cell, default: float) -> float:
    """The point size the cell's text is set in, so a rewrite keeps it."""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.text.strip() and run.font.size:
                return run.font.size.pt
    return default


def rewrite(cell, text: str, *, default: float = 8.0) -> None:
    replace_cell(cell, text, size=cell_size(cell, default))


def append_text(cell, tail: str, *, default: float = 8.0) -> None:
    rewrite(cell, cell.text.strip() + tail, default=default)


def swap_text(cell, old: str, new: str, *, default: float = 8.0) -> None:
    text = cell.text
    if old not in text:
        raise ValueError(f"Text not found in cell: {old[:70]}")
    rewrite(cell, text.replace(old, new).strip(), default=default)


def cut_from(cell, marker: str, tail: str, *, default: float = 8.0) -> None:
    """Replace everything from ``marker`` onwards with ``tail``."""
    text = cell.text
    cut = text.find(marker)
    if cut < 0:
        raise ValueError(f"Marker not found in cell: {marker[:70]}")
    rewrite(cell, (text[:cut] + tail).strip(), default=default)


def replace_cover_version(table, source: str, target: str) -> None:
    """Rewrite the version on the cover, which is one run of one title block."""
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def replace_control_value(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            replace_cell(row.cells[1], value, size=9)
            return
    raise ValueError(f"Control row not found: {label}")


def set_control(table, label: str, value: str) -> None:
    """Rewrite a control-page value in place, keeping its own formatting."""
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            rewrite(row.cells[1], value, default=9.0)
            return
    raise ValueError(f"Control row not found: {label}")


def find_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Body paragraph not found: {prefix}")


def replace_body_paragraph(doc, prefix: str, text: str) -> None:
    set_run_text(find_paragraph(doc, prefix), text)


def find_row(table, prefix: str, *, col: int = 0, exact: bool = False):
    for row in table.rows:
        text = row.cells[col].text.strip()
        if (text == prefix) if exact else text.startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix}")


def clone_row(table, anchor, values: list[str], *, before: bool = True):
    """Copy ``anchor`` beside itself and write ``values`` into the copy."""
    sizes = [cell_size(cell, 8.0) for cell in anchor.cells]
    clone = copy.deepcopy(anchor._tr)
    (anchor._tr.addprevious if before else anchor._tr.addnext)(clone)
    row = next(r for r in table.rows if r._tr is clone)
    for cell, value, size in zip(row.cells, values, sizes):
        replace_cell(cell, value, size=size)
    return row


def restripe(table) -> None:
    """Alternate white and pale-blue body rows, as the generator lays them out."""
    for index, row in enumerate(table.rows[1:]):
        for cell in row.cells:
            set_cell_shading(cell, VERY_PALE_BLUE if index % 2 else WHITE)


#: Row properties that must follow ``w:cantSplit`` in a ``w:trPr``.
_AFTER_CANT_SPLIT = ("trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden")


def forbid_row_split(row) -> None:
    """Keep one table row on a single page, at its place in the schema order."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is not None:
        return
    cant_split = OxmlElement("w:cantSplit")
    for name in _AFTER_CANT_SPLIT:
        later = tr_pr.find(qn(f"w:{name}"))
        if later is not None:
            later.addprevious(cant_split)
            return
    tr_pr.append(cant_split)


def prepend_change_log(table, version: str, summary: str) -> None:
    template = table.rows[1]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], SHORT_DATE, size=8)
    replace_cell(row.cells[2], summary, size=8)


def rewrite_box(table, edit) -> None:
    """Rewrite a one-paragraph callout whose lines are separated by breaks."""
    paragraph = table.rows[0].cells[0].paragraphs[0]
    lines = paragraph.text.split("\n")
    set_run_text(paragraph, "\n".join(edit(lines)))


def require_fr(table, label: str):
    if not table.rows[0].cells[0].text.strip().startswith(label):
        raise ValueError(f"Table is not {label}")
    return table


def finish(doc, output: Path, title: str, version: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


EVIDENCE_ROW, ACCEPTANCE_ROW, LIMIT_ROW = 2, 3, 4


# ═════════════════════════════════════════════════════════════════════════════
# Module 1 - School & Branch Management
# ═════════════════════════════════════════════════════════════════════════════

M01_DIR = "01-school-and-branch-management"
M01_STEM = "XVS_M01_School_and_Branch_Management_Functional_Requirements_Document"
M01_SOURCE, M01_TARGET = "1.23", "1.24"

#: Every table this revision edits, bound before any row is inserted.
M01_TABLES = {
    "decision": 6, "FR-001": 9, "FR-006": 14, "FR-007": 15,
    "sequence": 30, "school_api": 35, "validation": 38, "dependencies": 39,
    "traceability": 42, "change_log": 43,
}

M01_DECISION_COUNT = (
    "• Module 1 remains Backend Partial and In use Complete in MRD v2.77, with "
    "twenty-three capability entries."
)

M01_DECISION_REACH = (
    "• The role grants follow the School's reach, not the plan change. A tier "
    "moved down, an uplift withdrawn, an uplift that lapsed and any later "
    "settlement, and the command that re-applies a plan all take back the "
    "grants the School can no longer reach, because the settlement hangs off "
    "the reach in vs_config rather than off one caller. Only a band out of "
    "depth is taken back: a module closed because a subscription expired or an "
    "operator denied it leaves every grant standing, because that state is "
    "meant to be reversed and paying would not give the roles back. A key the "
    "School denied on purpose stays denied, a School given its plan for the "
    "first time keeps every role it arrived with, and permission groups the "
    "School composed are left as they are."
)

M01_DECISION_REPORT = (
    "• A role that cannot be settled is reported, never enforced. A kept key "
    "may declare a dependency on one being taken away, and that refusal used "
    "to abort the whole billing write, so a School could stay on Premium in "
    "the product while its invoice said Standard. Each role is settled on its "
    "own savepoint now; one that refuses is left exactly as it was, named in "
    "the plan response as roles_needing_attention and in the audit trail under "
    "plan_downgrade_blocked, and every other role in the School is settled "
    "regardless. Nothing retries it: the refusal is deterministic, this "
    "deployment runs no scheduler, and the next settlement of that School's "
    "reach tries the role again."
)

M01_DECISION_SHORT = (
    "• No School is created short of the roles CodeX ships. A template the "
    "prebuilt library cannot supply used to answer with nothing, which reads "
    "exactly like success, so a School could be created holding two of the "
    "five roles the product is built around with nobody being told which were "
    "skipped. School and Branch creation ask before they write anything and "
    "refuse with 503 ADMIN_PROVISIONING_FAILED, naming every missing template "
    "at once; the library itself now holds all five on a fresh installation."
)

M01_DECISION_BRANCH_NAME = (
    "• A renamed Branch takes its roles with it. The name of a Branch's own "
    "copy of Branch Admin is composed from the Branch's, and the composition "
    "happened once at provisioning, so a School that renamed Ikeja to Yaba "
    "kept Branch Admin - Ikeja on its roles screen with no way to correct it. "
    "The rule is one function re-applied from the Branch's own save, so the "
    "console, a command, an import and the shell all obey it. A name the "
    "School chose for itself is left alone."
)

M01_FR001_EVIDENCE = (
    "SchoolCreateSerializer and School.save() create the School and Tenant "
    "atomically, together with at least one Branch. Required School and Branch "
    "administrators are inside that same transaction, and so are the roles "
    "CodeX ships: School Admin, Teacher, Finance Admin and Procurement Admin "
    "for the whole School, and a Branch Admin for each Branch. vs_rbac "
    "provision_role_from_prebuilt makes each one from the prebuilt role "
    "library, get-or-create on its key so a School never holds two copies, and "
    "copies the library's permissions onto it as they stand at creation. The "
    "library is guaranteed: vs_schools 0013 seeds Teacher, Finance Admin and "
    "Procurement Admin beside the two administrator templates 0010 already "
    "seeded, preserving a row an operator has refined or deactivated, and tops "
    "up the Schools created while the library was short, matched on the role's "
    "key or its name and copying permissions from the library rather than "
    "listing them. Creation also asks before it writes: require_prebuilt_roles "
    "reads vs_rbac missing_prebuilt_keys for all five templates and raises "
    "AdminProvisioningError naming every missing one at once, and Branch "
    "creation asks the same question for branch_admin, so an installation that "
    "cannot give a School its roles refuses the School rather than creating it "
    "short. A School created before creation provisioned the whole set "
    "receives the missing roles from vs_rbac 0017, which leaves a role the "
    "School already has untouched. The set of books and Module 9 onboarding "
    "control room remain best effort inside their own savepoints because "
    "neither is the credential needed to operate the new School."
)

M01_FR001_ACCEPTANCE = (
    "A 201 response persists one School, one protected Tenant, at least one "
    "Branch with exactly one main, the four whole-School roles and a Branch "
    "Admin role for each Branch, and every required administrator account, "
    "scoped role assignment, invitation record, and admin link. A missing "
    "Branch is rejected before writes. Any required administrator provisioning "
    "failure returns 503 and leaves none of those creation records behind, and "
    "so does a library that cannot supply one of the five roles. "
    "ANewSchoolGetsTheRolesCodeXShipsTests proves the full set arrives without "
    "anybody running a command afterwards; ASchoolIsNeverCreatedShortOfRolesTests "
    "proves a missing template refuses the School and names what is missing, "
    "and that a whole library still creates the School with every role."
)

M01_FR001_LIMIT_OLD = (
    "Teacher, Finance Admin and Procurement Admin arrive only where the "
    "prebuilt library holds them: the school-app migration guarantees the "
    "School Admin and Branch Admin templates on a fresh installation, nothing "
    "guarantees the other three, and a missing one is skipped without a "
    "warning. "
)

M01_FR006_EVIDENCE_TAIL = (
    " A Branch role's name follows its Branch. vs_rbac branch_role_name states "
    "the rule once, and a post_save receiver on Branch re-applies it whenever a "
    "rename lands, so the console, a management command, an import and the "
    "shell obey the same rule rather than each holding a copy of it; vs_rbac "
    "0020 backfills the roles whose names had already gone stale. A name the "
    "School chose for itself is left alone, because renaming a Branch is not "
    "permission to rename what a School decided to call something, and a name "
    "another role in the Tenant already holds is left alone rather than failing "
    "the rename on a unique constraint."
)

M01_FR006_ACCEPTANCE_TAIL = (
    " ARenamedBranchTakesItsRoleWithItTests proves a rename through the Branch "
    "update endpoint renames that Branch's own role, and vs_rbac's "
    "BranchRolesFollowTheBranchTests proves the other Branches' roles are left "
    "alone, that a School with one Branch gains no suffix when it renames, "
    "that a name the School chose itself survives, that a name another role "
    "already holds does not fail the rename, and that an edit which is not a "
    "rename leaves the roles alone."
)

M01_FR006_LIMIT = (
    "Standalone Branch creation still declares primary admin input optional "
    "while the create path requires it. The previous partial-success and "
    "reconciliation gap is closed, and so is the stale-name gap: the naming "
    "rule is re-applied from the Branch's own save. A School that drops back to "
    "one Branch keeps the suffix on the role it has until a Branch of that "
    "School is saved again, because the rule is re-applied from a Branch write "
    "rather than from a Branch count."
)

M01_FR007_EVIDENCE_OLD = "change_plan then takes back every role grant"

M01_FR007_EVIDENCE_TAIL = (
    "The role grants follow the School's reach. vs_config settles them from the "
    "write that narrowed it: set_entitlement settles where the depth it wrote "
    "is shallower than the depth it replaced, and set_depth_grant and "
    "clear_depth_grant settle whatever the write did, because an uplift changes "
    "reach by the clock as well as by a write. A tier change, an uplift "
    "written or withdrawn, an uplift that lapsed months ago and any later "
    "settlement, and manage.py apply_plans therefore all reconcile without a "
    "caller remembering to ask. apply_plan_entitlements writes every module "
    "first and settles once, because a settlement run between the rows reads a "
    "half-moved School. vs_rbac revoke_grants_beyond_the_tenants_depth then "
    "asks what is out of reach now rather than what changed, and writes each "
    "affected role's remaining keys through set_role_access under source "
    "plan_downgrade, so every revocation takes the role's lock, bumps its "
    "version and leaves an audit entry. Two narrowings are deliberate: only a "
    "band out of depth is taken back, so a module closed because a "
    "subscription expired or an operator denied it leaves every grant "
    "standing, and the School's existing denies are named in the write, so a "
    "key denied on purpose stays denied. A School given its plan for the first "
    "time keeps every role it arrived with. A role whose kept key declares a "
    "dependency on a revoked one is reported rather than enforced: each role is "
    "settled on its own savepoint, a refused role is left exactly as it was, an "
    "RBACAuditLog row is written under source plan_downgrade_blocked at "
    "severity WARNING and status FAILED, and every plan endpoint answers with "
    "data.roles_needing_attention, a list of role_key, role_name, "
    "permission_keys and the refusal in words, with the message naming the "
    "count and the roles. Permission groups the School composed are left as "
    "they are, and their keys meet the same gate when they resolve. "
)

M01_FR007_ACCEPTANCE_TAIL = (
    " EveryWayTheReachShrinksTakesTheGrantsWithItTests proves the same of an "
    "uplift withdrawn through the endpoint, an uplift that lapsed months ago, a "
    "plan re-applied by command and a plan exception that still covers its "
    "keys, and proves that a first application leaves the roles it found alone, "
    "that a lapsed subscription does not take the roles apart, and that a key "
    "denied on purpose is still denied afterwards. "
    "ARoleThatCannotBeSettledIsReportedNotEnforcedTests proves the plan change "
    "completes, the role it could not settle is left exactly as it was, every "
    "other role in the School is still settled, the audit trail names it, the "
    "plan response names it, and an ordinary plan change still revokes and "
    "reports nothing."
)

M01_FR007_LIMIT_OLD = (
    " Revocation runs on a plan change only. Withdrawing an uplift, an uplift "
    "lapsing, and apply_plans re-applying a plan leave any role grant beyond "
    "the depth that remains in place: refused by the plan gate, hidden from the "
    "picker, and back in force if the depth returns."
)

M01_FR007_LIMIT_TAIL = (
    " Nothing retries a role the settlement could not save. The refusal comes "
    "from the role's own dependency graph rather than from a busy database, "
    "this deployment runs no scheduler, and the role waits for the next "
    "settlement of that School's reach or for somebody to edit it; until then "
    "its keys sit outside the depth the plan reaches, so the gate refuses them "
    "at the door exactly as if they had gone. A key carried by a permission "
    "group or a personal ALLOW override is not taken back at all, and an "
    "uplift that has lapsed leaves its role rows standing until the School's "
    "reach is next settled."
)

M01_SEQUENCE_VALIDATE_TAIL = (
    " Ask the prebuilt role library whether it can supply every role the School "
    "will be given."
)

M01_SEQUENCE_VALIDATE_OUTCOME_TAIL = (
    " Refuse a School the library cannot give its roles, naming every missing "
    "template in one message."
)

M01_PLAN_CHANGE_ROUTE = (
    "Move the school onto another plan and re-grant every module in the same "
    "transaction, taking back the role grants the new depth no longer reaches. "
    "The response carries roles_needing_attention, naming any role left exactly "
    "as it was because its own permission dependencies refused the change. The "
    "plan it is already on is refused."
)

M01_PLAN_UPLIFT_ROUTE = (
    "Record a time-boxed depth uplift for one module. A reason is required. The "
    "write settles the school's role grants afterwards, so an uplift that "
    "lapsed earlier is tidied up here, and the response carries "
    "roles_needing_attention."
)

M01_PLAN_UPLIFT_CLEAR_ROUTE = (
    "Withdraw an uplift, returning the module to the depth its plan pays for. "
    "The role grants that depth no longer reaches go back with it, and the "
    "response carries roles_needing_attention."
)

M01_VALIDATION_PROVISIONING_TAIL = (
    " A prebuilt role library that cannot supply one of the five roles is "
    "refused the same way, asked before anything is written, with every missing "
    "template named in one message."
)

M01_DEP_RBAC = (
    "Evaluate platform permissions; provision the roles CodeX ships with every "
    "School and the School and Branch administrator assignments, each role "
    "copying the prebuilt library's permissions when it is made; keep a Branch "
    "role's name in step with its Branch from the Branch's own save; and take "
    "back the role grants beyond the depth the Tenant reaches, through the "
    "audited role-access service, wherever that reach narrows. The library "
    "holds all five templates on a fresh installation through the school-app "
    "migrations, which preserve existing or deliberately deactivated rows, and "
    "a template the library cannot supply refuses School or Branch creation "
    "rather than leaving the School short."
)

M01_TRACE_LEAD = (
    "This table reconciles every Module 1 capability entry in MRD v2.77 to the "
    "controlling functional requirement in this FRD. All twenty-three entries "
    "are represented. Two are restated rather than added: the roles CodeX ships "
    "now records that a School is refused rather than created short and that a "
    "Branch role follows its Branch, and the role-grant entry covers every way "
    "a School's reach shrinks rather than a tier move alone."
)

M01_TRACE_ROLES = (
    "Implemented; School Admin, Teacher, Finance Admin and Procurement Admin "
    "for the whole School and a Branch Admin for each Branch, copied from the "
    "prebuilt library when each is made. A Branch role names its Branch only "
    "once the School has two, and follows the Branch when it is renamed. The "
    "library holds all five templates on a fresh installation, and a School the "
    "library cannot supply is refused with 503 rather than created short."
)

M01_TRACE_REVOCATION_NAME = (
    "Every way a school's reach shrinks takes back the role grants beyond it"
)

M01_TRACE_REVOCATION = (
    "Implemented; a tier change, an uplift written or withdrawn, an uplift that "
    "lapsed and any later settlement, and apply_plans all reconcile, because "
    "the settlement hangs off the reach in vs_config rather than off the plan "
    "change. Only a band out of depth is taken back, a key denied on purpose "
    "stays denied, a first application revokes nothing, and a role whose own "
    "dependencies refuse the change is reported in the response and in the "
    "audit trail rather than failing the write."
)

M01_CHANGE = (
    "Records the role-grant settlement as it now works, and two creation rules "
    "that were not holding. Taking back the grants a School could no longer "
    "reach was a step inside the plan change, so only a tier move did it: an "
    "uplift withdrawn, an uplift that simply lapsed, and apply_plans all left "
    "the keys in force, refused at the door by the plan gate, invisible to the "
    "administrator who would remove them, and back the moment the School moved "
    "up again for an unrelated reason. The settlement now hangs off the reach "
    "itself in vs_config, which every one of those paths already goes through, "
    "and the service is vs_rbac revoke_grants_beyond_the_tenants_depth. Two "
    "narrowings are deliberate: only a band out of depth is taken back, so a "
    "module closed because a subscription expired leaves every grant standing, "
    "and a key the School denied on purpose stays denied. A School given its "
    "plan for the first time keeps every role it arrived with. A role whose "
    "kept key depends on a revoked one used to make the validator refuse and "
    "abort the whole billing write, so a School could stay on Premium in the "
    "product while its invoice said Standard; each role is now settled on its "
    "own savepoint, a refused role is left exactly as it was and named in "
    "data.roles_needing_attention and in the audit trail under "
    "plan_downgrade_blocked, and every other role is still settled. Separately, "
    "provisioning a role the prebuilt library did not hold answered with "
    "nothing, which reads like success, so a School could be created holding "
    "two of the five roles CodeX ships; it raises now, School and Branch "
    "creation refuse before anything is written with 503 "
    "ADMIN_PROVISIONING_FAILED naming every missing template, and vs_schools "
    "0013 puts Teacher, Finance Admin and Procurement Admin in the library and "
    "tops up existing Schools. A Branch role's name is re-composed from the "
    "Branch's own save, with vs_rbac 0020 backfilling the rows already stale, "
    "skipping a name a School chose itself and skipping a collision. FR-001, "
    "FR-006 and FR-007, the module decision, the creation sequence, the plan "
    "routes, the validation rules, the vs_rbac dependency and MRD traceability "
    "are updated. Module 1 stays Backend Partial and In use Complete, with "
    "twenty-three capability entries against MRD v2.77. Verified by vs_config "
    "105, vs_rbac 547, schools.vs_schools 358, schools.vs_staff 225, vs_user "
    "381 and schools.vs_onboarding 178 tests. Backend evidence only; nothing "
    "here is deployed."
)


def patch_m01(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M01 School and Branch Management Functional Requirements "
        f"Document v{M01_TARGET}"
    )
    t = {name: doc.tables[index] for name, index in M01_TABLES.items()}
    for label in ("FR-001", "FR-006", "FR-007"):
        require_fr(t[label], label)

    replace_cover_version(doc.tables[0], M01_SOURCE, M01_TARGET)
    control = doc.tables[1]
    replace_control_value(control, "Version", M01_TARGET)
    replace_control_value(control, "Review date", REVIEW_DATE)
    replace_control_value(control, "Code baseline", CODE_BASELINE)
    replace_control_value(
        control, "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 1",
    )

    def edit_decision(lines):
        if not lines[1].startswith("• Module 1 remains Backend Partial"):
            raise ValueError("Decision box count line not found")
        reach = next(
            index for index, line in enumerate(lines)
            if line.startswith("• A lower plan tier takes the role grants")
        )
        return [
            lines[0], M01_DECISION_COUNT, *lines[2:reach], M01_DECISION_REACH,
            M01_DECISION_REPORT, *lines[reach + 1:],
            M01_DECISION_SHORT, M01_DECISION_BRANCH_NAME,
        ]

    rewrite_box(t["decision"], edit_decision)
    forbid_row_split(t["decision"].rows[0])

    fr001 = t["FR-001"]
    rewrite(fr001.rows[EVIDENCE_ROW].cells[1], M01_FR001_EVIDENCE)
    rewrite(fr001.rows[ACCEPTANCE_ROW].cells[1], M01_FR001_ACCEPTANCE)
    swap_text(fr001.rows[LIMIT_ROW].cells[1], M01_FR001_LIMIT_OLD, "")

    fr006 = t["FR-006"]
    append_text(fr006.rows[EVIDENCE_ROW].cells[1], M01_FR006_EVIDENCE_TAIL)
    append_text(fr006.rows[ACCEPTANCE_ROW].cells[1], M01_FR006_ACCEPTANCE_TAIL)
    rewrite(fr006.rows[LIMIT_ROW].cells[1], M01_FR006_LIMIT)

    fr007 = t["FR-007"]
    evidence = fr007.rows[EVIDENCE_ROW].cells[1]
    text = evidence.text
    start = text.find(M01_FR007_EVIDENCE_OLD)
    if start < 0:
        raise ValueError("FR-007 revocation passage not found")
    keep = text.find("A plan change on a School still onboarding", start)
    if keep < 0:
        raise ValueError("FR-007 onboarding sentence not found")
    rewrite(evidence, text[:start] + M01_FR007_EVIDENCE_TAIL + text[keep:])
    append_text(fr007.rows[ACCEPTANCE_ROW].cells[1], M01_FR007_ACCEPTANCE_TAIL)
    swap_text(
        fr007.rows[LIMIT_ROW].cells[1],
        M01_FR007_LIMIT_OLD, M01_FR007_LIMIT_TAIL,
    )

    validate = find_row(t["sequence"], "1. Validate")
    append_text(validate.cells[1], M01_SEQUENCE_VALIDATE_TAIL)
    append_text(validate.cells[2], M01_SEQUENCE_VALIDATE_OUTCOME_TAIL)

    api = t["school_api"]
    rewrite(find_row(api, "/v1/i/{slug}/plan/", col=1).cells[3], M01_PLAN_CHANGE_ROUTE)
    rewrite(
        find_row(api, "/v1/i/{slug}/plan/uplifts/", col=1, exact=True).cells[3],
        M01_PLAN_UPLIFT_ROUTE,
    )
    rewrite(
        find_row(api, "/v1/i/{slug}/plan/uplifts/{capability}/", col=1).cells[3],
        M01_PLAN_UPLIFT_CLEAR_ROUTE,
    )

    append_text(
        find_row(t["validation"], "Required administrator provisioning").cells[1],
        M01_VALIDATION_PROVISIONING_TAIL,
    )
    rewrite(find_row(t["dependencies"], "vs_rbac").cells[2], M01_DEP_RBAC)

    replace_body_paragraph(
        doc, "This table reconciles every Module 1 capability entry", M01_TRACE_LEAD,
    )
    trace = t["traceability"]
    rewrite(find_row(trace, "The roles CodeX ships").cells[2], M01_TRACE_ROLES)
    downgrade = find_row(trace, "A plan downgrade takes back the role grants")
    rewrite(downgrade.cells[0], M01_TRACE_REVOCATION_NAME)
    rewrite(downgrade.cells[2], M01_TRACE_REVOCATION)

    prepend_change_log(t["change_log"], M01_TARGET, M01_CHANGE)
    finish(doc, output, title, M01_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 4 - Roles & Permissions (RBAC)
# ═════════════════════════════════════════════════════════════════════════════

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.18", "1.19"

#: Table positions in v1.18, all bound before anything is inserted.
M04_T = {
    "cover": 0, "control": 1, "contents": 2, "in_scope": 3, "FR-019": 27, "FR-021": 29,
    "FR-025": 33, "withdrawn": 36, "model": 38, "dependencies": 43,
    "needs": 44, "removal": 45, "trace": 46, "reconcile": 47, "changes": 48,
}

M04_SOURCE_SCOPE = (
    "Backend changes hanging the role-grant settlement off a tenant's reach "
    "rather than off a plan change, reporting a role whose own permission "
    "dependencies refuse that settlement instead of failing the write, keeping "
    "a branch role's name in step with its branch, and refusing to provision a "
    "role the prebuilt library does not hold (12 September 2026)"
)

M04_CODE_INSPECTED = (
    f"{CODE_BASELINE}: apps/vs_rbac (plan_grants.py, services.py, signals.py, "
    "exceptions.py, plan_gate.py, migrations 0017 to 0020), vs_config's "
    "capability and depth services, and the vs_schools plan endpoints, plan "
    "command and school creation they meet"
)

M04_MRD_BASELINE = (
    f"XVS Module Requirements Document v{MRD_VERSION}, Module 4, twenty-two "
    "capability entries"
)

M04_SCOPE_AUDIT_TAIL = (
    " A settlement that could not save a role records that under source "
    "plan_downgrade_blocked at severity WARNING and status FAILED, so which "
    "roles a narrowing of reach left behind is a question the trail answers by "
    "filtering rather than by inference."
)

M04_FR019_SOURCE_OLD = (
    "prebuilt provisioning, role suggestions, plan-downgrade revocation and "
    "Super Admin permission reconciliation, and names each as its source: an "
    "approved role change is approved_change_request, or "
    "self_approved_change_request with a self_approved flag when its requester "
    "decided it, and a downgrade is plan_downgrade."
)

M04_FR019_SOURCE_NEW = (
    "prebuilt provisioning, role suggestions, the settlement that follows a "
    "narrowing of a tenant's reach and Super Admin permission reconciliation, "
    "and names each as its source: an approved role change is "
    "approved_change_request, or self_approved_change_request with a "
    "self_approved flag when its requester decided it, and a settlement is "
    "plan_downgrade. A role that settlement could not save is recorded "
    "separately, by vs_rbac.plan_grants and under source "
    "plan_downgrade_blocked, at severity WARNING and status FAILED, carrying "
    "the keys that stayed behind and the refusal in words; the row is written "
    "after set_role_access has rolled that role's own savepoint back, so it "
    "survives the refusal instead of vanishing with it."
)

M04_FR019_LIMIT_OLD = (
    "the direct update path, an approved role change and a plan downgrade all "
    "supply grants alone, so each deletes any explicit deny on the role it "
    "touches, and the audit's before-and-after shows it going."
)

M04_FR019_LIMIT_NEW = (
    "the direct update path and an approved role change supply grants alone, so "
    "each deletes any explicit deny on the role it touches, and the audit's "
    "before-and-after shows it going. The reach settlement is the exception: it "
    "names the role's existing denies in the same write, so a key a school "
    "denied on purpose is still denied after a revocation rather than becoming "
    "merely ungranted and returning through a group."
)

M04_FR021_EVIDENCE = (
    "PrebuiltRoleTemplate rows are read-only library definitions, five of them: "
    "school_admin, branch_admin, finance_admin, procurement_admin and teacher. "
    "Bursar is retired and finance_manager is renamed finance_admin, which "
    "claims every finance and payments key by prefix, as procurement_admin "
    "claims every procurement key, restricted ones included; the prefix "
    "attachment is additive and never removes a default. seed_workflow_permissions "
    "adds the workflow defaults to school_admin, finance_admin, "
    "procurement_admin and branch_admin and grants the same keys to the school "
    "roles already built from them, additively. provision_role_from_prebuilt "
    "copies a template into a tenant-owned, unlocked system role through "
    "set_role_access, and raises PrebuiltRoleMissing where the library holds no "
    "active template for the key; it answered None instead, which reads exactly "
    "like success at any call site that does not check, and that is how a "
    "school was created holding two of the five roles the product ships. "
    "missing_prebuilt_keys answers the same question for a whole set before "
    "provisioning starts, so a caller refusing over an incomplete library names "
    "every missing role at once rather than sending an operator round the loop. "
    "A branch-scoped copy's key always carries the branch; its name is composed "
    "by branch_role_name, which carries the branch only once the tenant has "
    "more than one, and provisioning a role for a later branch renames each "
    "same-named sibling to say which branch it runs. sync_branch_role_names "
    "re-applies that one function from a post_save receiver on Branch, so a "
    "rename made by the console, a command, an import or the shell reaches the "
    "roles named after that branch; it rewrites only a name this code composed, "
    "leaves a name the tenant chose for itself, and leaves a name another role "
    "in the tenant already holds rather than failing the rename on the unique "
    "constraint. Migration 0020 backfills the rows that had already gone stale "
    "under the same three rules. School creation provisions the four "
    "tenant-wide templates and a Branch Admin per branch, and Migration 0017 "
    "gave every existing school the tenant-wide roles it lacked, matched by key "
    "or by name and granting only keys a tenant may hold, leaving a role the "
    "school already had untouched."
)

M04_FR021_ACCEPTANCE = (
    "A school created today holds School Admin, Teacher, Finance Admin, "
    "Procurement Admin and a Branch Admin without an operator running a command "
    "(ANewSchoolGetsTheRolesCodeXShipsTests.test_the_full_set_is_provisioned_"
    "without_anybody_running_a_command). A one-branch school's Branch Admin "
    "names no branch, and a second branch makes every sibling say which "
    "(test_one_branch_means_the_role_does_not_name_it, "
    "test_a_second_branch_makes_every_sibling_say_which). A library missing a "
    "template refuses the school and names what is missing rather than creating "
    "it short (ASchoolIsNeverCreatedShortOfRolesTests). A renamed branch takes "
    "the roles named after it (BranchRolesFollowTheBranchTests), while the "
    "other branches' roles, a name the school chose itself, and a name another "
    "role already holds are each left alone. Adoption remains idempotent, "
    "tenant-owned, editable, undeletable, and durably audited."
)

M04_FR021_LIMIT = (
    "The provisioning copy is not filtered by scope, so a library default for a "
    "key later reclassified PLATFORM is refused at the grant and provisioning "
    "fails. School creation provisions Finance Admin, so a school cannot be "
    "created on a database whose library attached finance.entity.create while "
    "it was tenant-holdable; this is traced in code, and no test builds Finance "
    "Admin from seeded finance keys. See Needs Attention. "
    "PrebuiltRoleTemplate.scope is written and read by nothing: Teacher is "
    "declared branch-scoped and is provisioned whole-tenant in every new school. "
    "The naming rule is re-applied from a branch save, so a school that drops "
    "back to one branch keeps the suffix on the role it has until a branch of "
    "that tenant is saved again."
)

M04_FR025_EVIDENCE_OLD = (
    "Moving a school down a tier revokes the role grants the new tier does not "
    "reach: change_plan calls revoke_grants_the_plan_no_longer_reaches after "
    "the new entitlements are written, reads each granted key through "
    "plan_reader, and removes the lost ones through set_role_access under "
    "source plan_downgrade, so each revocation takes the role's lock, bumps its "
    "version and is audited. Permission groups are left as the school composed "
    "them."
)

M04_FR025_EVIDENCE_NEW = (
    "A narrowing of what a tenant reaches takes the role grants beyond it. "
    "revoke_grants_beyond_the_tenants_depth, in vs_rbac.plan_grants, asks what "
    "is out of reach now rather than what a write changed, and removes the lost "
    "keys through set_role_access under source plan_downgrade, so each "
    "revocation takes the role's lock, bumps its version and is audited. It "
    "hangs off the reach itself: vs_config's set_entitlement calls it where the "
    "depth it wrote is shallower than the one it replaced, and set_depth_grant "
    "and clear_depth_grant call it whatever the write did, so a tier change, an "
    "uplift written or withdrawn, an uplift that lapsed by the clock and any "
    "later settlement, and manage.py apply_plans all reconcile without a caller "
    "remembering to ask. Only a band out of depth is taken back: a module that "
    "is closed rather than shallow leaves every grant standing, because an "
    "expired subscription and an operator's denial are states meant to be "
    "reversed and the gate already refuses every key behind them. An "
    "unprovisioned tenant loses nothing, exactly as the gate refuses it "
    "nothing, and a school given its plan for the first time keeps every role "
    "it arrived with. The write names the role's existing denies, so a key "
    "denied on purpose stays denied. Each role is settled on its own savepoint: "
    "one whose own permission dependencies refuse the change is left exactly as "
    "it was, recorded under source plan_downgrade_blocked, and named to the "
    "operator in the calling endpoint's data.roles_needing_attention, while "
    "every other role in the tenant is still settled. Permission groups are "
    "left as the school composed them."
)

M04_FR025_ACCEPTANCE_TAIL = (
    " Every other way the reach shrinks revokes the same way, and a role that "
    "cannot be settled is reported rather than allowed to fail the write "
    "(EveryWayTheReachShrinksTakesTheGrantsWithItTests, "
    "ARoleThatCannotBeSettledIsReportedNotEnforcedTests, both in "
    "schools.vs_schools)."
)

M04_FR025_LIMIT = (
    "Only rbac_permission is gated: a view gated solely by "
    "rbac_group_permission or by HasAnyModuleAccess passes untouched. "
    "Enforcement is one platform switch rather than one per school, so a staged "
    "rollout would need a platform-writable per-tenant value. The catalogue "
    "deliberately does not read that switch: it reports what the school bought "
    "whether or not refusals are being issued. The settlement revokes direct "
    "role grants only: a key carried by a group or a personal ALLOW override "
    "stays and is refused at the door, and moving back up restores nothing. "
    "Nothing retries a role the settlement could not save. The refusal is "
    "deterministic, no scheduler runs in this deployment, and the role waits "
    "for the next settlement of that tenant's reach or for somebody to edit it; "
    "its keys are refused at the door in the meantime, so what is left is a "
    "stale row rather than live access. A settlement performed by the vs_config "
    "entitlement endpoint reaches the audit trail and the logs and nothing "
    "else, because that endpoint passes no collector (Needs Attention)."
)

M04_WITHDRAWN_ROUTE = "Narrow what the school's plan reaches"

M04_WITHDRAWN_WHAT = (
    "Every direct role grant beyond the depth the school now reaches, on every "
    "role in the school. A tier moved down, an uplift withdrawn or lapsed, and "
    "a plan re-applied all count."
)

M04_WITHDRAWN_HOW = (
    "Immediate, through set_role_access under source plan_downgrade, so each "
    "role is locked, versioned and audited. Groups and overrides are left "
    "alone, denies are preserved, and moving back up restores nothing. A role "
    "whose own dependencies refuse the change keeps its grants and is reported "
    "under plan_downgrade_blocked."
)

M04_MODEL_AUDIT_TAIL = (
    " A role a reach settlement could not save is recorded here too, under "
    "source plan_downgrade_blocked at WARNING and FAILED, with the keys that "
    "stayed behind."
)

M04_DEP_CONFIG = (
    "Supplies proxy_idle_timeout_minutes, resolved per tenant and branch, which "
    "decides when an open-ended proxy session lapses. It also supplies the "
    "capability catalogue and the depth a tenant reaches into each module, "
    "which the plan gate reads through effective_capability, and the "
    "enforcement switch that decides whether that gate answers at all. The gate "
    "reads that switch through get_config on every gated request, uncached. The "
    "dependency runs the other way once: every vs_config write that can narrow "
    "a tenant's depth calls this module's revoke_grants_beyond_the_tenants_depth "
    "afterwards, which is the single return journey and is imported inside the "
    "function for that reason."
)

M04_DEP_MODULE1 = (
    "Consumes school_admin and branch_admin during required administrator "
    "creation, and provisions every tenant-wide prebuilt role, Teacher, Finance "
    "Admin and Procurement Admin beside School Admin, when a school is created; "
    "a template this module's library cannot supply refuses that creation "
    "rather than leaving the school short. Its plan endpoints and its "
    "apply_plans command reach this module's settlement through vs_config, and "
    "carry the roles it could not settle into their own responses. Its "
    "school-app migrations guarantee the five template identities on fresh "
    "installations; the RBAC seed command owns the templates' access defaults."
)

M04_VERIFY_SETTLEMENT = (
    "•  The reach settlement: every way a tenant's reach shrinks takes back the "
    "grants beyond it and leaves the ones it still reaches; a closed module and "
    "a first plan application take nothing; a key denied on purpose is still "
    "denied; and a role whose own dependencies refuse the change is left whole, "
    "reported to the caller and recorded under plan_downgrade_blocked while "
    "every other role is settled."
)

M04_NEED_REPORT = [
    "P2",
    "A role the settlement could not save waits for a person, and one endpoint "
    "cannot say so. The refusal is deterministic and nothing re-runs it: this "
    "deployment has no scheduler, so the role keeps grants beyond the tenant's "
    "depth until the reach is next settled or somebody edits the role. The keys "
    "are refused at the door in the meantime, so this is a stale row rather "
    "than live access. The plan endpoints and the apply_plans command name the "
    "role to whoever made the change; the vs_config entitlement endpoint "
    "reconciles with no collector, so a role it could not settle reaches the "
    "audit trail and the logs and nobody reading the response.",
    "Give the entitlement endpoint the same collector the plan endpoints use, "
    "so an operator writing a grant there is told what it left behind, and "
    "decide whether a report of roles standing outside their tenant's depth is "
    "worth a screen of its own.",
    "FR-025",
]

M04_REMOVAL_DELTA = (
    "• Against v1.18 no item leaves. One is added: a role the reach settlement "
    "could not save waits for a person, and the vs_config entitlement endpoint "
    "cannot report one."
)

M04_TRACE_LEAD = (
    f"MRD v{MRD_VERSION} records Module 4 as Roles & Permissions (RBAC), Phase "
    "V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two "
    "capability entries. The module number, name, phase, states and ownership "
    "agree with this revision. Each entry maps below; the plan-gate and role-"
    "template entries are restated rather than added to."
)

M04_TRACE_AUDIT = (
    "Implemented. Durable and append-only, with actor, reason, source (approved, "
    "self-approved, reach settlements and the roles a settlement could not save "
    "among them), approval reference where applicable, and complete direct "
    "grant, direct deny, group and effective combined before-and-after sets. The "
    "central trail remains a best-effort mirror."
)

M04_TRACE_GATE = (
    "Implemented and in service. The catalogue registers the switch ON, a school "
    "is refused under PLAN_UPGRADE_REQUIRED once its plan is applied, and every "
    "narrowing of what the school reaches revokes the direct role grants beyond "
    "it, reporting any role whose own dependencies refuse the change rather "
    "than failing the write."
)

M04_TRACE_TEMPLATES = (
    "Implemented. Five templates, School Admin, Branch Admin, Finance Admin, "
    "Procurement Admin and Teacher, provisioned as tenant-owned copies; school "
    "creation provisions all of them and is refused where the library cannot "
    "supply one, and migration 0017 gave every existing school the tenant-wide "
    "roles it lacked. A branch role's name follows its branch, backfilled by "
    "migration 0020. A library default for a key later reclassified PLATFORM "
    "breaks provisioning (Needs Attention P1)."
)

M04_RECONCILE = (
    "MRD RECONCILIATION\n"
    f"• MRD v{MRD_VERSION} lists Module 4 as Roles & Permissions (RBAC), Phase "
    "V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two "
    "capability entries. The module number, name, phase, states and ownership "
    "agree with this revision.\n"
    "• The count does not change. Hanging the settlement off a tenant's reach, "
    "reporting a role it could not save, and refusing to provision a role the "
    "library does not hold all restate entries the tracker already carries: the "
    "plan depth gate, the permission and assignment audit history, and the role "
    "templates.\n"
    "• Two capability names are restated in the tracker to match: the plan-gate "
    "entry now covers every way a school's reach shrinks rather than a "
    "downgrade alone, and the role-template entry records that a school is "
    "refused rather than created short.\n"
    "• Backend evidence only. Neither document claims deployment, production "
    "adoption, or data migration completion."
)

M04_CHANGE = (
    "Corrects what this module does when a tenant's reach narrows, and two "
    "provisioning rules that were not holding. Taking back out-of-plan grants "
    "was a step inside the schools app's plan change, so only a tier move did "
    "it and this document said so; the service is now "
    "vs_rbac.plan_grants.revoke_grants_beyond_the_tenants_depth, called from "
    "every vs_config write that can narrow a tenant's depth, so a tier change, "
    "an uplift written or withdrawn, an uplift that lapsed and any later "
    "settlement, and the apply_plans command all reconcile. Only a band out of "
    "depth is taken back, because a closed module is a state meant to be "
    "reversed; the write names the role's existing denies, so a key denied on "
    "purpose stays denied, which also withdraws this document's claim that a "
    "downgrade clears a role's explicit denies. A role whose kept key depends "
    "on a revoked one is reported rather than enforced: each role is settled on "
    "its own savepoint, a refused role is left exactly as it was and recorded "
    "under the new source plan_downgrade_blocked at WARNING and FAILED, and the "
    "calling endpoint names it in roles_needing_attention. Provisioning a "
    "template the library does not hold raises PrebuiltRoleMissing rather than "
    "answering None, and missing_prebuilt_keys lets a caller name every missing "
    "role in one refusal. A branch role's name is one function, branch_role_name, "
    "re-applied by sync_branch_role_names from the branch's own save, with "
    "migration 0020 backfilling the stale rows, skipping a name the school "
    "chose itself and skipping a collision. FR-019, FR-021 and FR-025 are "
    "updated, with the in-scope summary, the withdrawal routes, the data model, "
    "the vs_config and Module 1 dependencies, the minimum verification list, "
    "Needs Attention and MRD traceability following. One Needs Attention item "
    "is added: nothing retries a role the settlement could not save, and the "
    "vs_config entitlement endpoint reconciles with nowhere to report one. "
    "Module 4 remains Backend Complete and In use Complete, with twenty-two "
    "capability entries against MRD v2.77. Verified by vs_rbac 547 tests, with "
    "vs_config 105, schools.vs_schools 358, schools.vs_staff 225, vs_user 381 "
    "and schools.vs_onboarding 178 beside them. Backend evidence only; nothing "
    "here is deployed."
)


def patch_m04(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{M04_TARGET}"
    )
    t = {name: doc.tables[index] for name, index in M04_T.items()}
    for label in ("FR-019", "FR-021", "FR-025"):
        require_fr(t[label], label)

    replace_cover_version(t["cover"], M04_SOURCE, M04_TARGET)
    set_control(t["control"], "Version", M04_TARGET)
    set_control(t["control"], "Review date", REVIEW_DATE)
    set_control(t["control"], "Source scope", M04_SOURCE_SCOPE)
    set_control(t["control"], "Code inspected", M04_CODE_INSPECTED)
    set_control(t["control"], "MRD baseline", M04_MRD_BASELINE)

    rewrite(
        find_row(t["contents"], "10. MRD Traceability").cells[1],
        f"Agreement with MRD v{MRD_VERSION}'s twenty-two capability entries",
    )

    append_text(
        find_row(t["in_scope"], "The durable RBAC audit").cells[1],
        M04_SCOPE_AUDIT_TAIL,
    )

    fr019 = t["FR-019"]
    swap_text(fr019.rows[EVIDENCE_ROW].cells[1],
              M04_FR019_SOURCE_OLD, M04_FR019_SOURCE_NEW)
    swap_text(fr019.rows[LIMIT_ROW].cells[1],
              M04_FR019_LIMIT_OLD, M04_FR019_LIMIT_NEW)

    fr021 = t["FR-021"]
    rewrite(fr021.rows[EVIDENCE_ROW].cells[1], M04_FR021_EVIDENCE)
    rewrite(fr021.rows[ACCEPTANCE_ROW].cells[1], M04_FR021_ACCEPTANCE)
    rewrite(fr021.rows[LIMIT_ROW].cells[1], M04_FR021_LIMIT)

    fr025 = t["FR-025"]
    swap_text(fr025.rows[EVIDENCE_ROW].cells[1],
              M04_FR025_EVIDENCE_OLD, M04_FR025_EVIDENCE_NEW)
    append_text(fr025.rows[ACCEPTANCE_ROW].cells[1], M04_FR025_ACCEPTANCE_TAIL)
    rewrite(fr025.rows[LIMIT_ROW].cells[1], M04_FR025_LIMIT)

    plan_route = find_row(t["withdrawn"], "Move the school down a plan tier")
    rewrite(plan_route.cells[0], M04_WITHDRAWN_ROUTE)
    rewrite(plan_route.cells[1], M04_WITHDRAWN_WHAT)
    rewrite(plan_route.cells[2], M04_WITHDRAWN_HOW)

    append_text(find_row(t["model"], "RBACAuditLog").cells[2], M04_MODEL_AUDIT_TAIL)
    rewrite(find_row(t["dependencies"], "vs_config").cells[1], M04_DEP_CONFIG)
    rewrite(
        find_row(t["dependencies"], "Module 1, School and Branch Management").cells[1],
        M04_DEP_MODULE1,
    )
    replace_body_paragraph(doc, "•  A plan downgrade:", M04_VERIFY_SETTLEMENT)

    needs = t["needs"]
    clone_row(needs, needs.rows[-1], M04_NEED_REPORT, before=False)
    restripe(needs)

    def edit_removal(lines):
        if not lines[-1].startswith("• Against v1.17"):
            raise ValueError("Removal rule delta line not found")
        return [*lines[:-1], M04_REMOVAL_DELTA]

    rewrite_box(t["removal"], edit_removal)

    replace_body_paragraph(doc, "MRD v2.76 records Module 4", M04_TRACE_LEAD)
    trace = t["trace"]
    rewrite(find_row(trace, "Permission and assignment audit history").cells[2],
            M04_TRACE_AUDIT)
    rewrite(find_row(trace, "Plan depth gate with its own refusal").cells[2],
            M04_TRACE_GATE)
    rewrite(find_row(trace, "Role templates").cells[2], M04_TRACE_TEMPLATES)

    rewrite_box(t["reconcile"], lambda lines: M04_RECONCILE.split("\n"))
    prepend_change_log(t["changes"], M04_TARGET, M04_CHANGE)
    finish(doc, output, title, M04_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 6 - Configuration & Capability Management
# ═════════════════════════════════════════════════════════════════════════════

M06_DIR = "06-configuration-and-capability"
M06_STEM = (
    "XVS_M06_Configuration_and_Capability_Management_Functional_Requirements_Document"
)
M06_SOURCE, M06_TARGET = "1.1", "1.2"

M06_T = {
    "cover": 0, "control": 1, "FR-010": 16, "FR-012": 18, "FR-018": 24,
    "dependencies": 36, "needs": 37, "trace": 38, "changes": 39,
}

M06_FR010_EVIDENCE_TAIL = (
    " A grant that leaves the tenant reaching less than it did takes the role "
    "grants beyond the new depth with it: set_entitlement calls vs_rbac's "
    "revoke_grants_beyond_the_tenants_depth where the depth it wrote is "
    "shallower than the one it replaced, so what a tenant may do and what its "
    "roles say it may do stay one fact. A first grant takes nothing, because "
    "there was no shallower answer before it. A caller writing a tenant's "
    "modules as a set turns the reconciliation off for each row and settles "
    "once when every row is in place, because a settlement run between the rows "
    "of one batch reads a half-moved tenant and acts on it."
)

M06_FR010_ACCEPTANCE_TAIL = (
    " Moving a tenant's module to a shallower depth leaves no role holding keys "
    "beyond it, while a module closed rather than shallowed leaves every grant "
    "standing."
)

M06_FR010_LIMIT = (
    "A null depth means no depth limit at all rather than the shallowest one. "
    "Grants written before depth existed therefore reach every band, which is "
    "deliberate and is discussed under Needs Attention. The entitlement "
    "endpoint settles the tenant's role grants and passes no collector, so a "
    "role it could not settle reaches the audit trail and the logs and not the "
    "response the operator is reading."
)

M06_FR012_EVIDENCE_TAIL = (
    " Both writes settle the tenant's role grants afterwards, whatever the "
    "write did. This table is the one place where reach changes by the clock as "
    "well as by a write, so a deal edited today may be tidying up after one "
    "that lapsed in March, and asking what the tenant reaches now rather than "
    "what this write changed is what catches it. The keys an uplift reached go "
    "back with it when it is withdrawn, on its last day or months after it "
    "lapsed. Either write fills a caller's list with any role whose grants "
    "could not be settled, so the endpoint can name it to the operator; such a "
    "role is left exactly as it was and the write still completes, because a "
    "commercial decision is never blocked by one role's internal wiring."
)

M06_FR012_ACCEPTANCE_TAIL = (
    " Withdrawing an uplift takes back the keys the tier below it cannot reach, "
    "and an uplift nobody removed until months after it lapsed is settled at "
    "the next write rather than silently kept."
)

M06_FR012_LIMIT = (
    "One uplift per capability per tenant. A second overlapping deal on the same "
    "module replaces the first rather than stacking. Nothing in this deployment "
    "runs on a schedule, so an uplift that lapses by the clock stops answering "
    "at once while the role rows it reached stand until the tenant's reach is "
    "next settled."
)

M06_FR018_EVIDENCE_TAIL = (
    " A role grant this module's writes cause to be taken back is recorded in "
    "the RBAC trail rather than here, because the row changed is a role's, and "
    "so is a role that could not be settled, under its own source."
)

M06_DEP_RBAC = (
    "Enforces the nineteen PLATFORM-scoped keys, and consumes this module's "
    "evaluation in the plan gate, reading platform.entitlements.enforce through "
    "get_config on every gated request. The permission-to-band map lives there, "
    "not here. It also receives the one call that runs the other way: every "
    "write here that can narrow a tenant's depth asks vs_rbac to take back the "
    "role grants beyond it, and vs_rbac owns that service, its audit trail and "
    "the report of any role it could not settle."
)

M06_VERIFICATION = (
    "The requirements above name the tests behind each behaviour; this revision "
    "was traced from the code at the baseline named in Document Control and did "
    "not re-run them. This module's own suite ran 105 tests OK at that "
    "baseline, and the consumers of the settlement these writes now perform ran "
    "green beside it: vs_rbac 547 and schools.vs_schools 358, the latter "
    "holding the tests that exercise every way a tenant's reach can shrink. "
    "Backend evidence only; nothing here is deployed."
)

M06_NEED_COLLECTOR = [
    "6",
    "An entitlement written from the console reports nothing it left behind",
    "A write that narrows a tenant's depth takes back the role grants beyond "
    "it, and a role whose own permission dependencies refuse that change is "
    "left as it was and reported to the caller. The plan endpoints in Module 1 "
    "pass a collector and name the role to the operator who made the change. "
    "The entitlement endpoint here does not, so an operator who shallows a "
    "module from the configuration console is told the grant was saved and "
    "nothing more; the role is in the RBAC audit trail and in the logs, where "
    "nobody is looking at that moment.",
    "Collect the settlement's report in the entitlement endpoint and return it "
    "the way the plan endpoints do.",
]

M06_TRACE_LEAD = (
    f"Module 6 of XVS Module Requirements Document v{MRD_VERSION} lists "
    "twenty-six capabilities. Each maps to the requirements above; the role "
    "grants that now follow a narrowing of a tenant's depth are a consequence "
    "this module's writes trigger and vs_rbac performs, so they strengthen "
    "FR-010 and FR-012 without adding a capability or changing the count."
)

M06_TRACE_ENTITLEMENTS = "FR-010"
M06_CHANGE = (
    "Records that a write narrowing a tenant's depth now settles that tenant's "
    "role grants. Entitlements decide what the product offers and role grants "
    "are what a customer handed its own people, and the two were kept in step "
    "only by the schools app's plan change, so an uplift withdrawn here, an "
    "uplift that lapsed, and a module shallowed from the console all left roles "
    "holding keys the tenant could no longer reach. set_entitlement settles "
    "where the depth it wrote is shallower than the one it replaced, "
    "set_depth_grant and clear_depth_grant settle whatever they did because "
    "this table also changes reach by the clock, and a caller writing a "
    "tenant's modules as a set turns the per-row settlement off and settles "
    "once when every row is in place. The service is vs_rbac's, which is the "
    "single return journey in a dependency that otherwise runs one way. A role "
    "whose own permission dependencies refuse the change is left exactly as it "
    "was, reported to the caller and recorded in the RBAC trail, and the write "
    "still completes. FR-010, FR-012 and FR-018, the vs_rbac dependency, the "
    "verification evidence and MRD traceability are updated, and one current "
    "gap is added: the entitlement endpoint settles with no collector, so an "
    "operator writing a grant there is not told which role it left behind. "
    "Module 6 remains Backend Complete and In use Complete with twenty-six "
    "capability entries against MRD v2.77. Backend evidence only; nothing here "
    "is deployed."
)


def patch_m06(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M06 Configuration and Capability Management Functional "
        f"Requirements Document v{M06_TARGET}"
    )
    t = {name: doc.tables[index] for name, index in M06_T.items()}
    for label in ("FR-010", "FR-012", "FR-018"):
        require_fr(t[label], label)

    replace_cover_version(t["cover"], M06_SOURCE, M06_TARGET)
    set_control(t["control"], "Version", M06_TARGET)
    set_control(t["control"], "Review date", REVIEW_DATE)
    set_control(t["control"], "Code baseline", CODE_BASELINE)
    set_control(
        t["control"], "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 6, "
        "twenty-six capability entries",
    )

    fr010 = t["FR-010"]
    append_text(fr010.rows[EVIDENCE_ROW].cells[1], M06_FR010_EVIDENCE_TAIL)
    append_text(fr010.rows[ACCEPTANCE_ROW].cells[1], M06_FR010_ACCEPTANCE_TAIL)
    rewrite(fr010.rows[LIMIT_ROW].cells[1], M06_FR010_LIMIT)

    fr012 = t["FR-012"]
    append_text(fr012.rows[EVIDENCE_ROW].cells[1], M06_FR012_EVIDENCE_TAIL)
    append_text(fr012.rows[ACCEPTANCE_ROW].cells[1], M06_FR012_ACCEPTANCE_TAIL)
    rewrite(fr012.rows[LIMIT_ROW].cells[1], M06_FR012_LIMIT)

    append_text(t["FR-018"].rows[EVIDENCE_ROW].cells[1], M06_FR018_EVIDENCE_TAIL)

    rewrite(find_row(t["dependencies"], "vs_rbac").cells[1], M06_DEP_RBAC)
    replace_body_paragraph(
        doc, "The requirements above name the tests behind each behaviour",
        M06_VERIFICATION,
    )

    needs = t["needs"]
    clone_row(needs, needs.rows[-1], M06_NEED_COLLECTOR, before=False)
    restripe(needs)

    replace_body_paragraph(doc, "Module 6 of XVS Module Requirements Document",
                           M06_TRACE_LEAD)
    prepend_change_log(t["changes"], M06_TARGET, M06_CHANGE)
    finish(doc, output, title, M06_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 9 - School Onboarding
# ═════════════════════════════════════════════════════════════════════════════

M09_DIR = "09-school-onboarding"
M09_STEM = "XVS_M09_School_Onboarding_Functional_Requirements_Document"
M09_SOURCE, M09_TARGET = "2.11", "2.12"

M09_TABLES = {
    "boundary": 4, "out_of_scope": 6, "FR-006": 15, "journey": 26,
    "refusals": 33, "dependencies": 34, "change_log": 41,
}

M09_PURPOSE_NOTE = (
    "Version 2.12 records the widened creation refusal this module receives "
    "from Module 1. A School is now refused, before anything is written, where "
    "the prebuilt role library cannot give it the roles CodeX ships, so the "
    "five roles the creation boundary describes are a guarantee rather than a "
    "best effort. No behaviour in this module changed."
)

M09_BOUNDARY_CREATION = (
    "• Module 1 returns 503 and rolls the whole new School transaction back "
    "when its required administrator cannot be provisioned, and refuses the "
    "School the same way, before anything is written, where the prebuilt role "
    "library cannot supply one of the roles it would be given. Module 9 "
    "receives no half-created School to repair and no School holding fewer "
    "roles than it should."
)

M09_BOUNDARY_ROLES = (
    "• Module 1 creates the roles CodeX ships with every School, School Admin, "
    "Teacher, Finance Admin and Procurement Admin for the whole School and a "
    "Branch Admin for each Branch, and a staff record for each administrator. "
    "The library holds all five templates on a fresh installation, so the set "
    "no longer depends on somebody having run a seeding command. DEFAULT_ROLES "
    "reads only the School Admin role and its whole-tenant assignment, so none "
    "of that changes what it checks."
)

M09_OUT_OF_SCOPE_OWNER = (
    "Module 1, School and Branch Management. Onboarding verifies the result and "
    "never re-creates it. The required administrator, the roles CodeX ships and "
    "each administrator's staff record are part of the School transaction, and "
    "creation is refused when administrator provisioning fails or when the role "
    "library cannot supply a role the School would be given."
)

M09_FR006_EVIDENCE_TAIL = (
    " That refusal now covers the role library as well as the administrator: "
    "creation asks whether the library holds School Admin, Branch Admin, "
    "Teacher, Finance Admin and Procurement Admin before it writes anything, "
    "and refuses with the same 503 naming every missing template, so a School "
    "cannot reach this module holding two of the five roles with nobody having "
    "been told which were skipped."
)

M09_JOURNEY_CREATION = (
    "School, tenant, at least one Branch, the roles CodeX ships, and the "
    "required administrator account, staff record, scoped role assignment, and "
    "invitation record exist, all committed together; a library that cannot "
    "supply one of those roles refuses the creation instead. The School and "
    "tenant are PENDING. Books and the onboarding control room are best effort "
    "inside their own savepoints."
)

M09_REFUSAL_CONDITION = (
    "Required administrator provisioning fails during School creation, or the "
    "prebuilt role library cannot supply a role the School would be given"
)

M09_REFUSAL_ANSWER = (
    "ADMIN_PROVISIONING_FAILED from Module 1; no School or onboarding state "
    "remains, and the message names every missing template"
)

M09_DEP_MODULE1 = (
    "Creates the School, Tenant, at least one Branch, the roles CodeX ships, "
    "the required administrator and a staff record for each administrator, and "
    "package, and refuses the creation outright where the role library cannot "
    "supply one of those roles. Supplies SUSPENDED and routes the paired Tenant "
    "lifecycle through the shared transition service, so expiry written through "
    "School.status carries the same transactional proxy shutdown."
)

M09_TRACE_LEAD = (
    f"Module 9 carries 21 capability entries in MRD v{MRD_VERSION}. Each maps "
    "to the requirements below. The wider creation refusal Module 1 now raises "
    "changes nothing this module checks, so no entry or count changes."
)

M09_CHANGE = (
    "Records the wider creation refusal this module receives. Module 1 "
    "provisions the roles CodeX ships with every School, and a template the "
    "prebuilt library did not hold used to be skipped in silence, so a School "
    "could arrive here holding two of the five roles the product is built "
    "around. Creation now asks before it writes and refuses with the same 503 "
    "ADMIN_PROVISIONING_FAILED, naming every missing template, and the library "
    "holds all five on a fresh installation, so the roles named in this "
    "document's creation boundary are a guarantee rather than a best effort. "
    "The purpose note, the creation boundary, section 1.2, FR-006, the creation "
    "workflow, the typed refusal for a failed creation, the Module 1 dependency "
    "and the MRD traceability lead are updated. Module 9 stays Backend Complete "
    "and In use Partial with twenty-one capability entries against MRD v2.77. "
    "No code in this module changed, and no route, permission key, refusal "
    "code, audit action type or notification event moved. Backend evidence "
    "only; nothing here is deployed."
)


def patch_m09(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M09 School Onboarding Functional Requirements "
        f"Document v{M09_TARGET}"
    )
    t = {name: doc.tables[index] for name, index in M09_TABLES.items()}
    require_fr(t["FR-006"], "FR-006")

    replace_cover_version(doc.tables[0], M09_SOURCE, M09_TARGET)
    control = doc.tables[1]
    replace_control_value(control, "Version", M09_TARGET)
    replace_control_value(control, "Review date", REVIEW_DATE)
    replace_control_value(control, "Code baseline", CODE_BASELINE)
    replace_control_value(
        control, "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 9",
    )
    replace_control_value(
        control, "Supersedes", f"v{M09_SOURCE} and all earlier versions, retained unchanged",
    )

    replace_body_paragraph(doc, "Version 2.11 corrects the checklist",
                           M09_PURPOSE_NOTE)

    def edit_boundary(lines):
        if not lines[1].startswith("• Module 1 returns 503"):
            raise ValueError("Boundary creation line not found")
        if not lines[2].startswith("• Module 1 creates the roles CodeX ships"):
            raise ValueError("Boundary roles line not found")
        return [
            lines[0], M09_BOUNDARY_CREATION, M09_BOUNDARY_ROLES, *lines[3:],
        ]

    rewrite_box(t["boundary"], edit_boundary)

    rewrite(
        find_row(t["out_of_scope"], "Creating the school, its tenant").cells[1],
        M09_OUT_OF_SCOPE_OWNER,
    )
    append_text(t["FR-006"].rows[EVIDENCE_ROW].cells[1], M09_FR006_EVIDENCE_TAIL)
    rewrite(find_row(t["journey"], "1", exact=True).cells[2], M09_JOURNEY_CREATION)

    refusal = find_row(t["refusals"], "Required administrator provisioning fails")
    rewrite(refusal.cells[0], M09_REFUSAL_CONDITION)
    rewrite(refusal.cells[2], M09_REFUSAL_ANSWER)

    rewrite(
        find_row(t["dependencies"], "Module 1, School and Branch Management").cells[1],
        M09_DEP_MODULE1,
    )
    replace_body_paragraph(doc, "Module 9 carries 21 capability entries",
                           M09_TRACE_LEAD)
    prepend_change_log(t["change_log"], M09_TARGET, M09_CHANGE)
    finish(doc, output, title, M09_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

PATCHES = {
    "m01": (M01_DIR, M01_STEM, M01_SOURCE, M01_TARGET, patch_m01),
    "m04": (M04_DIR, M04_STEM, M04_SOURCE, M04_TARGET, patch_m04),
    "m06": (M06_DIR, M06_STEM, M06_SOURCE, M06_TARGET, patch_m06),
    "m09": (M09_DIR, M09_STEM, M09_SOURCE, M09_TARGET, patch_m09),
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    parser.add_argument("--only", nargs="*", choices=sorted(PATCHES))
    parser.add_argument(
        "--overwrite", action="store_true",
        help="Replace a target this script already wrote, when re-rendering a fix.",
    )
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for name in args.only or sorted(PATCHES):
        folder, stem, source, target, patch = PATCHES[name]
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists() and not args.overwrite:
            raise SystemExit(
                f"{output.name} already exists; pass --overwrite to replace it"
            )
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
