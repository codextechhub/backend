#!/usr/bin/env python3
"""Version the MRD and RBAC FRD for administrator-owned permission groups."""

from __future__ import annotations

import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)


REVIEW_DATE = "23 September 2026"
CHANGE_DATE = "23 Sep 2026"
MRD_SOURCE = "2.82"
MRD_TARGET = "2.83"
FRD_SOURCE = "1.22"
FRD_TARGET = "1.23"

MRD_SOURCE_SCOPE = (
    "Backend and console worktrees restoring administrator permission-group CRUD, "
    "seeding dependency rules, and retiring backend-created groups, 23 September 2026"
)
FRD_SOURCE_SCOPE = (
    "Backend and console worktrees separating backend-owned permission definitions "
    "from administrator-owned permission groups, 23 September 2026"
)

MRD_CHANGE_SUMMARY = (
    "Corrected Module 4 ownership after the backend-owned registry change. Permission "
    "definitions, dependency rules and prebuilt role defaults remain backend-owned. "
    "Permission groups are administrator-created role-assignment shortcuts with full "
    "console and API create, read, update and delete support under separate create, "
    "update and delete permissions. Fresh databases receive no backend-created groups. "
    "The seed flow reconciles prerequisite edges, backfills existing and prebuilt roles, "
    "then converts legacy system-group access to direct grants before deleting those "
    "groups. Group writes reject restricted, inactive and dependency-incomplete sets, "
    "infer tenant or platform scope, and write durable RBAC audit events. The complete "
    "834-test RBAC run, focused seeder tests, guide checks and frontend production build "
    "pass. Module 4 remains Backend Complete and In use Complete with twenty-two "
    "capability entries. Nothing here claims deployment."
)

FRD_CHANGE_SUMMARY = (
    "Separates permission-group ownership from the backend-owned permission registry. "
    "Modules, resources, actions, permissions, dependency definitions and prebuilt role "
    "defaults remain code-owned and readable through the console. Administrators may "
    "create, read, update and delete permission groups under separate operation keys. "
    "The API hides legacy system groups and refuses restricted, inactive or dependency-"
    "incomplete membership. Backend seeding now reconciles prerequisite edges and safely "
    "retires old system groups by copying their effective access to direct role grants "
    "before deletion. Fresh installations contain no seeded groups. The complete 834-"
    "test RBAC run, focused seeder tests, guide checks and frontend production build pass. "
    "The capability count remains twenty-two. Nothing here claims deployment."
)

MRD_DELTA_ROWS = [
    [
        "Permission definitions and defaults",
        "Backend owned",
        "Modules, resources, actions, permissions, dependency rules and prebuilt role "
        "defaults remain reconciled by backend code and exposed with readable labels.",
    ],
    [
        "Permission groups",
        "Administrator owned",
        "Administrators create, read, update and delete reusable role-assignment bundles. "
        "No backend group catalogue is seeded.",
    ],
    [
        "Group write authority",
        "Split by operation",
        "Create, update and delete use distinct platform.permission_groups keys. Reading "
        "uses platform.permissions.view.",
    ],
    [
        "Group safety",
        "Enforced",
        "Restricted and inactive keys are refused, every prerequisite must be present, "
        "scope is inferred from members, and each mutation is recorded in RBAC audit.",
    ],
    [
        "Permission dependencies",
        "Backend reconciled",
        "Every non-view action requires its resource view key where one exists. Explicit "
        "rules cover cross-resource needs, and existing roles are backfilled safely.",
    ],
    [
        "Legacy system groups",
        "Retired safely",
        "Effective group grants become direct role grants before legacy backend groups "
        "are deleted. Existing direct denies remain authoritative.",
    ],
]


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def update_control_table(table, *, version: str, source_scope: str, mrd_baseline=None):
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label in {"Source scope", "Code inspected"}:
            replace_cell(row.cells[1], source_scope, size=9)
        elif label == "MRD baseline" and mrd_baseline:
            replace_cell(row.cells[1], mrd_baseline, size=9)


def prepend_change_log(table, version: str, date: str, summary: str) -> None:
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], date, size=8)
    replace_cell(row.cells[2], summary, size=8)


def keep_rows_whole(table) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.core_properties.title = title
    doc.core_properties.version = MRD_TARGET

    replace_cover_version(doc.tables[0], MRD_SOURCE, MRD_TARGET)
    update_control_table(
        doc.tables[1], version=MRD_TARGET, source_scope=MRD_SOURCE_SCOPE,
    )

    for row in doc.tables[2].rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(
                row.cells[0], f"5. v{MRD_TARGET} Capability Delta",
                size=9, bold=True, color=BLUE,
            )
            replace_cell(
                row.cells[1],
                "Permission groups are administrator-owned role shortcuts",
                size=9,
            )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == f"5. v{MRD_SOURCE} Capability Delta":
            write_paragraph(
                paragraph, f"5. v{MRD_TARGET} Capability Delta",
                size=17, bold=True, space_before=15, space_after=8,
            )
        elif text.startswith("Two-layer platform and tenant authorization"):
            write_paragraph(
                paragraph,
                "Two-layer platform and tenant authorization with reusable, backend-"
                "owned permission definitions, dependency rules, role defaults, "
                "assignments and controlled overrides. Modules, resources, actions and "
                "permissions are reconciled in backend code and exposed read-only with "
                "readable labels. Administrators build permission groups and custom roles "
                "from that catalogue, while personal grants or denies remain controlled "
                "allocation tools. Platform-only scope is enforced independently from "
                "restricted sensitivity. Each permission records the product capability "
                "that governs it, and a second gate refuses a request the caller's role "
                "allows but the tenant's plan does not reach. Branch reach remains grant-"
                "derived, and every supported role, permission-group or personal override "
                "change records its actor and before-and-after access state.",
                size=9, space_after=5,
            )

    replace_cell(
        doc.tables[14].rows[1].cells[0],
        "▸  Administrator-created permission groups",
        size=8.5,
    )
    replace_cell(
        doc.tables[14].rows[7].cells[1],
        "▸  Administrator-created permission groups with validated scope and dependencies\n\n"
        "▸  Reserved approval-role keys refused at role creation",
        size=8.5,
    )
    decision_cell = doc.tables[14].rows[-1].cells[0]
    decision = decision_cell.text.rstrip()
    old = (
        "• Global permission definitions and shipped defaults have one write boundary: "
        "backend seeders. The console reads the ordered catalogue, builds custom roles "
        "from it and applies personal exceptions, but it cannot create, rename or delete "
        "the definitions themselves. Legacy registry-write keys are deactivated without "
        "deleting their history."
    )
    new = (
        "• Permission definitions, dependency rules and prebuilt role defaults have one "
        "write boundary: backend seeders. Permission groups are administrator-created "
        "shortcuts with separate create, update and delete authority. Fresh installations "
        "contain no seeded groups, and legacy backend groups are deleted only after their "
        "effective grants are preserved directly on roles."
    )
    replace_cell(decision_cell, decision.replace(old, new), size=8.2)

    rebuild_table(
        doc.tables[76],
        [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
        MRD_DELTA_ROWS,
        [1.85, 1.15, 4.05],
    )
    keep_rows_whole(doc.tables[76])
    prepend_change_log(doc.tables[78], MRD_TARGET, CHANGE_DATE, MRD_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_frd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{FRD_TARGET}"
    )
    doc.core_properties.title = title
    doc.core_properties.version = FRD_TARGET

    replace_cover_version(doc.tables[0], FRD_SOURCE, FRD_TARGET)
    update_control_table(
        doc.tables[1], version=FRD_TARGET, source_scope=FRD_SOURCE_SCOPE,
        mrd_baseline=f"XVS Module Requirements Document v{MRD_TARGET}",
    )

    for row in doc.tables[2].rows:
        if row.cells[0].text.strip() == "10. MRD Traceability":
            replace_cell(
                row.cells[1],
                f"Agreement with MRD v{MRD_TARGET}'s twenty-two capability entries",
                size=9,
            )

    replace_cell(
        doc.tables[3].rows[4].cells[1],
        "Administrator-created reusable bundles attachable to custom roles. The backend "
        "infers tenant or platform scope from membership, requires the complete dependency "
        "closure, and refuses inactive or restricted permissions. No groups are seeded.",
        size=8.5,
    )
    replace_cell(
        doc.tables[7].rows[3].cells[1],
        "Read the global permission catalogue, create and maintain permission groups, "
        "administer platform and tenant roles, and manage permitted personal exceptions. "
        "Permission definitions, dependencies and prebuilt defaults change through backend "
        "code and seeding; permission groups change through administrator CRUD.",
        size=8.2,
    )
    replace_cell(
        doc.tables[7].rows[3].cells[2],
        "platform.permissions.view for catalogue and group reads; "
        "platform.permission_groups.create, .update and .delete for group writes; "
        "platform.roles.* and the relevant override keys for allocation.",
        size=8.2,
    )

    replace_cell(
        doc.tables[9].rows[2].cells[1],
        "Permission.save composes key as f\"{module_id}.{resource.name}.{action_id}\". "
        "The master and module seeders reconcile the vocabulary, dependency rules and "
        "prebuilt role defaults. Registry collection and detail views use read-only "
        "generic views under platform.permissions.view, so POST, PUT, PATCH and DELETE "
        "return 405. show_access_registry prints the hierarchy and default role library "
        "in one place. Permission groups use their separate administrator CRUD routes.",
        size=8.5,
    )

    fr005 = doc.tables[13]
    replace_cell(
        fr005.rows[1].cells[1],
        "A permission group is an administrator-created grant path: attach it to a role "
        "and every key inside lands in the effective set. Its membership must be safe, "
        "complete and understandable before it can be reused.",
        size=8.5,
    )
    replace_cell(
        fr005.rows[2].cells[1],
        "PermissionGroupListCreateView and PermissionGroupDetailView provide CRUD under "
        "separate create, update and delete keys while reads use platform.permissions.view. "
        "The serializer infers scope, refuses inactive and restricted keys, validates the "
        "full dependency closure, and hides the legacy is_system marker. GroupPermission "
        "and TenantRoleGroup retain model-level scope and restricted-key backstops. Every "
        "create, update and delete writes a durable RBAC audit event.",
        size=8.5,
    )
    replace_cell(
        fr005.rows[3].cells[1],
        "An administrator can create, read, update and delete a group in the console. A "
        "tenant-only member set produces a tenant group; any platform key produces a "
        "platform group. Restricted, inactive or dependency-incomplete sets are refused.",
        size=8.5,
    )
    replace_cell(
        fr005.rows[4].cells[1],
        "A platform group may mix platform and tenant-scoped ordinary keys for CodeX use, "
        "but restricted keys remain excluded because group attachment carries no per-key "
        "approval provenance.",
        size=8.5,
    )

    fr015 = doc.tables[23]
    replace_cell(
        fr015.rows[2].cells[1],
        "PermissionDependency stores the graph, and PermissionDependencyValidator resolves "
        "transitive prerequisites for roles and groups. seed_permission_dependencies "
        "reconciles active definitions: each non-view action requires its resource's view "
        "key where one exists, explicit cross-resource rules are declared in code, and "
        "existing and prebuilt roles receive missing prerequisites without replacing "
        "direct denies. The dependencies endpoint can return the complete connected graph "
        "for one permission so the console shows more than the current page.",
        size=8.5,
    )
    replace_cell(
        fr015.rows[3].cells[1],
        "A role or group containing finance.invoice.update without finance.invoice.view is "
        "refused with a message naming the missing key. Re-seeding restores the declared "
        "graph and preserves working access on existing roles.",
        size=8.5,
    )
    replace_cell(
        fr015.rows[4].cells[1],
        "The general view prerequisite is structural. Domain-specific dependencies still "
        "need an explicit backend rule when the required permission sits on another "
        "resource or cannot be inferred from the action name.",
        size=8.5,
    )

    fr020 = doc.tables[28]
    replace_cell(
        fr020.rows[1].cells[1],
        "Administrators need reusable permission bundles for jobs that recur, without "
        "turning those local allocation choices into backend product definitions.",
        size=8.5,
    )
    replace_cell(
        fr020.rows[2].cells[1],
        "Fresh installations seed no PermissionGroup rows. The console exposes a readable "
        "module, resource and action picker for administrators to create and maintain "
        "groups. Restricted permissions are disabled in the picker and refused on the "
        "server; dependency-incomplete groups are refused. retire_system_permission_groups "
        "copies effective grants from legacy backend groups into direct role grants, leaves "
        "existing direct denies untouched, deletes the legacy groups and refreshes affected "
        "role versions.",
        size=8.5,
    )
    replace_cell(
        fr020.rows[3].cells[1],
        "The group list contains only administrator-created rows. An administrator may add, "
        "edit, deactivate or delete a group and attach it to custom roles. Running the full "
        "permission seed leaves zero backend groups while preserving access previously "
        "supplied through a legacy system group.",
        size=8.5,
    )
    replace_cell(
        fr020.rows[4].cells[1],
        "Group scope distinguishes tenant permissions from platform permissions. Branch "
        "reach still belongs to the role assignment and permission behavior rather than to "
        "the group row itself.",
        size=8.5,
    )

    fr023 = doc.tables[31]
    replace_cell(
        fr023.rows[2].cells[1],
        "seed_all_permissions runs definition sources in dependency order, synchronizes "
        "field definitions, reconciles the permission dependency graph, backfills missing "
        "prerequisites on existing and prebuilt roles, and retires legacy backend-created "
        "groups after preserving their access as direct grants. Prebuilt templates remain "
        "the backend-owned role defaults. Permission-group rows are never seeded.",
        size=8.5,
    )
    replace_cell(
        fr023.rows[3].cells[1],
        "A fresh environment receives the same backend-owned hierarchy, dependency rules "
        "and prebuilt role defaults, with no permission groups. Re-running the seeders is "
        "idempotent, preserves custom roles and administrator-created groups, and removes "
        "any reintroduced legacy system group without removing effective role access.",
        size=8.5,
    )

    replace_cell(
        doc.tables[38].rows[3].cells[2],
        "Unique per (permission, depends_on). Backend seeding reconciles structural and "
        "explicit edges; validators resolve them transitively and raise on cycles.",
        size=8.2,
    )
    replace_cell(
        doc.tables[38].rows[4].cells[2],
        "Administrator-owned. Scope is inferred from membership. GroupPermission refuses "
        "cross-scope and restricted membership; the serializer also refuses inactive and "
        "dependency-incomplete sets. A group grants nothing until attached.",
        size=8.2,
    )

    rebuild_table(
        doc.tables[40],
        ["Method and path", "Purpose and permission"],
        [
            ["GET /rbac/vision/permission-modules/", "List backend-defined modules. platform.permissions.view."],
            ["GET /rbac/vision/permission-modules/{name}/", "Read one backend-defined module. platform.permissions.view."],
            ["GET /rbac/vision/permission-resources/", "List resources within modules. platform.permissions.view."],
            ["GET /rbac/vision/permission-resources/{id}/", "Read one resource. platform.permissions.view."],
            ["GET /rbac/vision/permission-actions/", "List reusable actions. platform.permissions.view."],
            ["GET /rbac/vision/permission-actions/{name}/", "Read one action. platform.permissions.view."],
            ["GET /rbac/vision/permissions/", "List permissions with readable labels. platform.permissions.view."],
            ["GET /rbac/vision/permissions/{key}/", "Read one permission and its relationships. platform.permissions.view."],
            ["GET /rbac/vision/permission-dependencies/", "List or search prerequisite edges, or return one connected graph. platform.permissions.view."],
            ["GET /rbac/vision/permission-dependencies/{id}/", "Read one prerequisite edge. platform.permissions.view."],
            ["GET, POST /rbac/vision/permission-groups/", "List administrator groups with platform.permissions.view; create with platform.permission_groups.create."],
            ["GET, PATCH, PUT, DELETE /rbac/vision/permission-groups/{id}/", "Read with platform.permissions.view; edit with platform.permission_groups.update; delete with platform.permission_groups.delete."],
        ],
        [2.55, 4.5],
        font_size=8.2,
    )
    keep_rows_whole(doc.tables[40])

    for row in doc.tables[42].rows:
        if row.cells[0].text.strip() == "A permission set missing a prerequisite":
            replace_cell(
                row.cells[0],
                "A role or permission group missing a prerequisite",
                size=8.2,
            )
            replace_cell(
                row.cells[1],
                "400 listing each key and what it also requires.",
                size=8.2,
            )
        elif row.cells[0].text.strip() == "A definition mutation on a global registry route":
            replace_cell(
                row.cells[1],
                "405 for modules, resources, actions, permissions and dependencies. "
                "Permission groups are the exception and expose administrator CRUD under "
                "their operation-specific keys.",
                size=8.2,
            )

    for row in doc.tables[44].rows:
        if row.cells[1].text.strip().startswith("The school catalogue's branch reach is advisory"):
            replace_cell(
                row.cells[1],
                "Permission groups infer tenant or platform scope, but carry no branch-"
                "reach declaration. Branch reach belongs to the role assignment, so the "
                "group itself cannot explain whether a key has useful meaning when that "
                "role is pinned to one branch.",
                size=8.2,
            )
            replace_cell(
                row.cells[2],
                "Declare branch applicability on Permission and validate it when a role "
                "assignment names a branch, so every direct key and group-derived key "
                "follows the same rule.",
                size=8.2,
            )

    trace = doc.tables[46]
    replace_cell(
        trace.rows[1].cells[2],
        "Implemented. Seeders own permission definitions, dependency rules and prebuilt "
        "role defaults; definition APIs are read-only and backend labels let the console "
        "present permissions without exposing dotted keys. Groups and custom roles remain "
        "administrator-owned allocation records.",
        size=8.2,
    )
    replace_cell(
        trace.rows[2].cells[2],
        "Implemented. Backend seeding reconciles the graph, structural view prerequisites "
        "and explicit cross-resource rules, backfills roles, and the validators enforce "
        "the transitive closure on roles and groups.",
        size=8.2,
    )
    replace_cell(
        trace.rows[3].cells[2],
        "Implemented. Administrators have CRUD under separate operation keys. Scope and "
        "dependency closure are inferred and validated, restricted and inactive keys are "
        "refused, mutations are audited, and no backend groups are seeded.",
        size=8.2,
    )
    replace_cell(
        trace.rows[14].cells[2],
        "Implemented with limits. Seeders provide one reviewed definition source, "
        "reconcile dependencies, preserve existing role access while retiring legacy "
        "system groups, and leave administrator-created groups untouched; aggregate scope "
        "classification remains open.",
        size=8.2,
    )
    replace_cell(
        trace.rows[17].cells[0],
        "Administrator-created permission groups with validated scope and dependencies",
        size=8.2,
    )
    replace_cell(
        trace.rows[17].cells[2],
        "Implemented with limits. Administrators create the bundles they need, the server "
        "validates membership and audits every change, and legacy backend groups are "
        "retired safely. Branch applicability is not declared on Permission (P2).",
        size=8.2,
    )

    replace_cell(
        doc.tables[47].rows[0].cells[0],
        f"MRD RECONCILIATION\n"
        f"• MRD v{MRD_TARGET} lists Module 4 as Roles & Permissions (RBAC), Phase V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two capability entries. The module number, name, phase, states and ownership agree with this revision.\n"
        "• The count does not change. Administrator-owned group CRUD, backend dependency reconciliation and safe legacy-group retirement restate the existing permission-group, dependency and seeding capabilities.\n"
        "• Permission definitions and prebuilt role defaults remain backend-owned. Custom roles, administrator groups and personal exceptions remain supported allocation surfaces.\n"
        "• Backend and console verification are recorded separately from deployment and production adoption.",
        size=8.5,
    )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Module 4 is the platform's authorization boundary"):
            write_paragraph(
                paragraph,
                "Module 4 is the platform's authorization boundary. Every other module "
                "asks it whether a person may act and whose rows they may see. Permission "
                "definitions, dependency rules and prebuilt role defaults are owned in "
                "backend code and read through one global catalogue. Administrators still "
                "compose permission groups and custom roles from that catalogue and apply "
                "controlled personal grants or denies. Scope, restricted state, plan depth, "
                "branch reach and audit continue to govern those allocations.",
                size=9, space_after=5,
            )
        elif text.startswith("•  Backend ownership:"):
            write_paragraph(
                paragraph,
                "•  Ownership boundaries: global definition collections refuse writes; "
                "GET returns ordered definitions and readable labels; administrators create, "
                "update and delete permission groups under separate keys; custom roles and "
                "personal exceptions remain available through their allocation routes.",
                size=9, space_after=3,
            )
        elif text.startswith("FR-020  Publish the School Catalogue"):
            paragraph.text = (
                "FR-020  Let Administrators Maintain Reusable Permission Groups"
            )
        elif text.startswith(f"MRD v{MRD_SOURCE} records Module 4"):
            write_paragraph(
                paragraph,
                f"MRD v{MRD_TARGET} records Module 4 as Roles & Permissions (RBAC), Phase V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two capability entries. The module number, name, phase, states and ownership agree with this revision. Each entry maps below.",
                size=9, space_after=5,
            )

    prepend_change_log(doc.tables[48], FRD_TARGET, CHANGE_DATE, FRD_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    mrd_dir = root / "module-requirements"
    frd_dir = root / "functional-requirements" / "04-roles-and-permissions-rbac"

    mrd_source = mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"
    mrd_output = mrd_dir / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx"
    frd_stem = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
    frd_source = frd_dir / f"{frd_stem}_v{FRD_SOURCE}.docx"
    frd_output = frd_dir / f"{frd_stem}_v{FRD_TARGET}.docx"

    patch_mrd(mrd_source, mrd_output)
    patch_frd(frd_source, frd_output)
    print(mrd_output)
    print(frd_output)


if __name__ == "__main__":
    main()
