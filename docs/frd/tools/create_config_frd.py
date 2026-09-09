#!/usr/bin/env python3
"""Create the Module 6 Configuration & Capability Management FRD.

Module 6 has carried the depth model, the entitlement chain and, most
recently, a school-facing endpoint, and every one of those decisions has been
recorded in the MRD because the module had no document of its own. This is
that document: the first code-aligned baseline, written from vs_config as it
stands rather than from the MRD's summary of it.

    python tools/create_config_frd.py

Writes the FRD only. The MRD's Module 6 entry already names all twenty-six
capabilities and is reconciled here by traceability rather than by editing it
again in the same change.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_requirements_documents import (
    BLUE,
    GREY,
    add_body,
    add_callout as add_reference_callout,
    add_cover,
    add_heading,
    add_metadata_table,
    add_page_break,
    add_requirement,
    add_status_key,
    add_table,
    assert_no_em_dash,
    remove_body_content,
    set_headers,
    shrink_inherited_media,
    update_extended_title,
    write_paragraph,
)

FRD_VERSION = "1.0"
REVIEW_DATE = "9 September 2026"
MRD_VERSION = "2.71"
CODE_BASELINE = (
    "apps/vs_config on main at the data-export rebanding, with vs_rbac's plan "
    "gate and the school plan page as its principal consumers, reviewed on "
    "9 September 2026"
)


REQUIREMENTS = [
    {
        "id": "FR-001",
        "title": "Declare a Setting Before Anything May Write It",
        "status": "Implemented",
        "requirement": (
            "No configuration value exists without a declaration that states its "
            "type, its permitted values, and the scopes it may be written at. A key "
            "nobody declared is not a key with no rules; it is not a key."
        ),
        "evidence": (
            "ConfigurationDefinition holds key, label, description, value_type, "
            "default_value, validation_rules and allowed_scopes. seed_config_catalogue "
            "declares the platform and school-scoped definitions the product ships "
            "with and widens an existing platform-only row when a definition becomes "
            "school-writable, because get_or_create would otherwise leave a row that "
            "refuses every school write."
        ),
        "acceptance": (
            "Writing a value for an undeclared key is refused. Re-seeding is "
            "idempotent and never narrows an existing definition's scopes. A "
            "definition carries its own default, so a value that was never written "
            "still resolves."
        ),
        "limit": (
            "Definitions are archived rather than deleted, so a retired key stays "
            "readable to the audit trail that references it."
        ),
    },
    {
        "id": "FR-002",
        "title": "Place Every Scoped Row at Exactly One of Three Levels",
        "status": "Implemented",
        "requirement": (
            "A configuration value, a capability override and an audit event each "
            "belong to the platform, to one tenant, or to one branch, and the level "
            "must be unambiguous to both a constraint and a lookup."
        ),
        "evidence": (
            "ScopedModel gives all three the same two nullable keys and derives "
            "scope_key on save: 'platform', 'tenant:<id>' or 'branch:<id>'. A "
            "branch's tenant is read from the branch and must agree when both are "
            "supplied. The denormalised key exists because SQL treats NULL as "
            "unequal to NULL, so a unique constraint over the two nullable columns "
            "would admit duplicate platform rows."
        ),
        "acceptance": (
            "One row per definition per scope is enforced at the database. A scope "
            "chain is fetched in one query by scope_key. A branch row whose tenant "
            "contradicts its branch is refused."
        ),
        "limit": (
            "Platform scope is a placement and not an ownership: a row with both "
            "keys null belongs to no tenant and applies everywhere."
        ),
    },
    {
        "id": "FR-003",
        "title": "Resolve a Value by Precedence, Nearest Scope First",
        "status": "Implemented",
        "requirement": (
            "Reading a setting for a branch must answer with the branch's own value "
            "if it has one, then the school's, then the platform's, then the "
            "definition's default, and must give the same answer to every caller."
        ),
        "evidence": (
            "The resolution service walks the scope chain in that order and falls "
            "through to default_value. Callers reach it through conf.get_config and "
            "the effective-values endpoint rather than reading ConfigurationValue "
            "rows, so a screen and a service cannot disagree about what is in force."
        ),
        "acceptance": (
            "A branch value shadows its school's; removing it exposes the school's "
            "again without a write. A key never written anywhere resolves to its "
            "declared default rather than to nothing."
        ),
        "limit": (
            "Precedence is fixed. There is no per-definition policy that would let "
            "one setting resolve platform-first."
        ),
    },
    {
        "id": "FR-004",
        "title": "Refuse a Value the Declaration Does Not Permit",
        "status": "Implemented",
        "requirement": (
            "A write is checked against the declaration before it is stored: the "
            "type, the choice list, the numeric bounds, and whether the scope being "
            "written at is one the definition allows."
        ),
        "evidence": (
            "validate_value checks type, choices and bounds from validation_rules, "
            "and the scope guard refuses a level absent from allowed_scopes. The "
            "middle level is labelled 'school' in the declaration while the stored "
            "key reads 'tenant:<id>', which is deliberate: a school is a tenant, and "
            "the label is part of the definition's public shape."
        ),
        "acceptance": (
            "A wrongly typed value, a value outside its choices or bounds, and a "
            "value written at a disallowed scope are each refused with the setting "
            "named. A refused write leaves no row and no audit event."
        ),
        "limit": (
            "Validation describes one value in isolation. A rule spanning two "
            "settings has to be expressed as a guard under FR-005."
        ),
    },
    {
        "id": "FR-005",
        "title": "Let the Owning Module Judge Its Own Setting",
        "status": "Implemented",
        "requirement": (
            "Some settings are only legal given facts this module has no business "
            "knowing. Those must be judged by the module that owns the fact, without "
            "configuration importing it."
        ),
        "evidence": (
            "An owning app registers a guard for its own key from AppConfig.ready "
            "and the resolution service calls it without importing that app. A guard "
            "refuses by raising ConfigurationError or a subclass, which every write "
            "path already renders as the setting being rejected. The dependency "
            "points one way: everything may depend on configuration and "
            "configuration depends on nothing."
        ),
        "acceptance": (
            "'Run payroll per branch' is refused for a school with unposted staff, "
            "and vs_config never imports the staff app to decide it. A guard raising "
            "any other exception is not swallowed."
        ),
        "limit": (
            "A guard is registered at startup, so a deployment that fails to load "
            "the owning app silently loses that rule rather than failing closed."
        ),
    },
    {
        "id": "FR-006",
        "title": "Serve the Grouped Settings Screens as One Read Each",
        "status": "Implemented",
        "requirement": (
            "Platform, security and integration settings are screens rather than key "
            "lists, and each must resolve in a single call carrying every value in "
            "force for the caller's scope."
        ),
        "evidence": (
            "PlatformSettingsView, SecuritySettingsView and IntegrationSettingsView "
            "each read and write their own group through the same resolution and "
            "validation path as an individual value, so a grouped write is not a "
            "second way to bypass a declaration."
        ),
        "acceptance": (
            "A grouped read returns resolved values, not raw rows. A grouped write "
            "that fails validation on one field writes none of them."
        ),
        "limit": (
            "The groups are defined in code. Adding a setting to a screen is a "
            "deploy rather than a catalogue edit."
        ),
    },
    {
        "id": "FR-007",
        "title": "Test an Integration Without Being Able to Change It",
        "status": "Implemented",
        "requirement": (
            "An operator must be able to confirm that a deployment-owned connection "
            "works, without the platform storing, echoing or altering the "
            "credentials behind it."
        ),
        "evidence": (
            "The connections service performs bounded, read-only checks and is rate "
            "limited by a thirty-second cooldown per connection, raising "
            "ConnectionTestCoolingDown rather than queuing. Credentials remain "
            "deployment-owned and are never written through this module."
        ),
        "acceptance": (
            "A test reports reachable or not without returning the secret it used. "
            "A second test inside the cooldown is refused rather than run."
        ),
        "limit": (
            "A failure says the connection did not work. It does not distinguish a "
            "wrong credential from an unreachable host."
        ),
    },
    {
        "id": "FR-008",
        "title": "Hold the Capability Catalogue Exactly Two Levels Deep",
        "status": "Implemented",
        "requirement": (
            "A capability is either a module, which a school is sold, or a band, "
            "which is that module cut at an ordered depth. There is no third level "
            "and no band without a module."
        ),
        "evidence": (
            "Capability carries a self-referencing parent and a depth from an "
            "ordered IntegerChoices spaced by ten (Core 10, Plus 20, Advanced 30). "
            "A check constraint requires a band to have both a parent and a depth, "
            "and clean() refuses a band of a band and a module carrying a depth."
        ),
        "acceptance": (
            "A module with a depth is refused. A band without a parent is refused. "
            "A band whose parent is itself a band is refused. Spacing by ten leaves "
            "room to insert a depth between two without renumbering the rest."
        ),
        "limit": (
            "The three depths are fixed in code. Adding a fourth is a migration and "
            "a repricing, not a catalogue edit."
        ),
    },
    {
        "id": "FR-009",
        "title": "Refuse a Capability Graph That Cannot Resolve",
        "status": "Implemented",
        "requirement": (
            "A capability may require another, and the graph must never contain a "
            "self-requirement or a cycle that evaluation could not terminate on."
        ),
        "evidence": (
            "CapabilityDependency is unique on the pair and carries a check "
            "constraint that a capability cannot require itself. The evaluator "
            "tracks the path it is walking and raises CapabilityDependencyError on "
            "a cycle rather than recursing without bound. Procurement requires "
            "Finance and the parent portal requires the student portal."
        ),
        "acceptance": (
            "A duplicate dependency is refused by the database. A self-requirement "
            "is refused by the constraint. A cycle introduced by two separate "
            "writes is detected at evaluation and named."
        ),
        "limit": (
            "A cycle is caught when something evaluates it, not when it is written, "
            "so a bad pair can sit in the catalogue until a read walks it."
        ),
    },
    {
        "id": "FR-010",
        "title": "Grant a Module to a Tenant at a Stated Depth",
        "status": "Implemented",
        "requirement": (
            "An entitlement records that a tenant holds a module, how deep into it "
            "the tenant reaches, where the grant came from, and the window it "
            "covers."
        ),
        "evidence": (
            "CapabilityEntitlement carries state, source, depth and a starts_at / "
            "ends_at window, unique per capability and scope. The plan writes rows "
            "with source PACKAGE through vs_schools' apply_plan_entitlements, "
            "carrying the depth the plan reaches for that module and the "
            "subscription's own expiry as ends_at."
        ),
        "acceptance": (
            "A grant outside its window does not answer. Re-applying a plan rewrites "
            "the same rows rather than accumulating them. A school's grant is scoped "
            "to its tenant, so one school cannot read another's."
        ),
        "limit": (
            "A null depth means no depth limit at all rather than the shallowest "
            "one. Grants written before depth existed therefore reach every band, "
            "which is deliberate and is discussed under Needs Attention."
        ),
    },
    {
        "id": "FR-011",
        "title": "Open a Band When Its Module's Depth Reaches It",
        "status": "Implemented",
        "requirement": (
            "Nobody buys a band. A band answers when the depth the tenant holds on "
            "its parent module is at least the band's own depth."
        ),
        "evidence": (
            "The depth service resolves a tenant's depth for a module through a "
            "fixed chain - a live uplift, then the entitlement's depth, then a "
            "platform grant, then unlimited - and depth_allows compares the band's "
            "depth against it. A band declares requires_entitlement False, because "
            "its module's grant is what opens it."
        ),
        "acceptance": (
            "A tenant at Core reaches Core bands and is refused Plus and Advanced. "
            "Moving the module's depth to Plus opens every Plus band across the "
            "module at once, with no per-band write."
        ),
        "limit": (
            "Depth is per module. A school cannot hold Plus of one half of a module "
            "and Core of the other; that would be two modules."
        ),
    },
    {
        "id": "FR-012",
        "title": "Lift One Tenant's Depth for a Fixed Term",
        "status": "Implemented",
        "requirement": (
            "A deal made with one school must be able to deepen one module for a "
            "stated period without changing the school's plan, and must expire on "
            "its own."
        ),
        "evidence": (
            "CapabilityDepthGrant is unique on capability and tenant, carries its "
            "own window, and refuses a band because depth is a property of a module. "
            "Resolution takes the greater of the entitlement's depth and a live "
            "uplift, so an uplift raises and never lowers. A plan change does not "
            "touch it: it lives in its own table precisely so a tier moving "
            "underneath leaves it alone, and it stops mattering once the tier passes "
            "it."
        ),
        "acceptance": (
            "An uplift on Basic survives a move to Standard. An expired uplift "
            "returns the module to the depth the plan pays for without a write. An "
            "uplift naming a band is refused, naming the module to use instead."
        ),
        "limit": (
            "One uplift per capability per tenant. A second overlapping deal on the "
            "same module replaces the first rather than stacking."
        ),
    },
    {
        "id": "FR-013",
        "title": "Override a Capability at a Scope, in Both Directions",
        "status": "Implemented",
        "requirement": (
            "An operator must be able to switch a capability on or off for one "
            "platform, tenant or branch without editing what that tenant was "
            "granted, and must be able to stop overriding it."
        ),
        "evidence": (
            "CapabilityOverride is a ScopedModel unique per capability and scope, "
            "with states ENABLED, DISABLED and INHERIT. Evaluation reads the scope "
            "chain nearest-first and the first non-INHERIT state wins, after "
            "entitlement and dependencies have been settled."
        ),
        "acceptance": (
            "A branch override beats a tenant override, which beats a platform one. "
            "Setting INHERIT restores the underlying answer rather than storing the "
            "value it happened to have."
        ),
        "limit": (
            "An override is absolute at its scope: it cannot express 'on, but only "
            "to this depth'. Depth is changed by a grant or an uplift."
        ),
    },
    {
        "id": "FR-014",
        "title": "Answer 'Is This On?' the Same Way Everywhere",
        "status": "Implemented",
        "requirement": (
            "Whether a capability is on for a scope is computed from entitlement, "
            "dependencies, depth and overrides together, and the single-capability "
            "path and the bulk path must not be able to disagree."
        ),
        "evidence": (
            "effective_capability answers for one capability and "
            "BulkCapabilityEvaluator answers for the whole catalogue from a fixed "
            "set of preloaded queries; the bulk path mirrors the single path "
            "including the band rule, and a band whose module was left out of the "
            "evaluated set fails closed rather than reporting itself on. The API, "
            "the navigation and the plan gate all read the bulk path."
        ),
        "acceptance": (
            "The two paths return the same answer for every capability in the "
            "catalogue, which is asserted by test. A catalogue read costs a fixed "
            "handful of queries rather than one per capability."
        ),
        "limit": (
            "The evaluator is built per request. It is not a cache, and a change "
            "made mid-request is not seen by an evaluator already constructed."
        ),
    },
    {
        "id": "FR-015",
        "title": "Let a School Read What Its Own Plan Reaches",
        "status": "Implemented",
        "requirement": (
            "A school's own screens must be able to ask what the school bought, "
            "without holding a platform key and without being able to ask about "
            "anybody else."
        ),
        "evidence": (
            "GET /config/my-capabilities/ answers for the caller's asserted tenant "
            "with no RBAC key at all, and is open before go-live because a school's "
            "navigation is first drawn while it is still pending. It reports states "
            "only, never the entitlement rows, overrides or grants behind them. The "
            "tenant comes from the assertion the authentication layer has already "
            "validated against the caller, which is what makes the absent key safe."
        ),
        "acceptance": (
            "A teacher holding no configuration key is served. Naming a rival "
            "school's slug is refused before the view runs. The platform endpoint "
            "stays closed to a school. A row carries key and enabled and nothing "
            "else."
        ),
        "limit": (
            "It answers what is reachable, not which plan is in force. The plan "
            "itself is read from the console, which is Module 1's surface and "
            "platform-owned."
        ),
    },
    {
        "id": "FR-016",
        "title": "Distinguish a School Nobody Provisioned From One That Bought Nothing",
        "status": "Implemented",
        "requirement": (
            "A tenant holding no package grant at all has not bought nothing; it was "
            "never given a plan. Anything deciding what to show a school must not "
            "confuse the two."
        ),
        "evidence": (
            "tenant_is_provisioned asks whether any PACKAGE entitlement exists for "
            "the tenant, and self_effective_capabilities reports everything on when "
            "it does not. The rule lives beside the entitlements it reads and the "
            "plan gate re-exports it, so the endpoint a school's navigation reads "
            "and the gate that issues refusals cannot answer differently. The "
            "platform endpoint keeps the literal answer, because an operator asking "
            "what a school holds wants the rows."
        ),
        "acceptance": (
            "A school created before grants were written reliably reaches "
            "everything, on both the menu and the gate. One PACKAGE row is enough to "
            "make a school provisioned, so this cannot be used to reach past a real "
            "plan. Both halves are pinned by test."
        ),
        "limit": (
            "The rule retires itself rather than being removed: once every school "
            "has been given a plan it never fires, but nothing enforces that state."
        ),
    },
    {
        "id": "FR-017",
        "title": "Retire a Capability Without Deleting It",
        "status": "Implemented",
        "requirement": (
            "A capability that stops describing anything sellable must be able to "
            "leave the price list while the entitlements, overrides and audit "
            "history pointing at it stay readable."
        ),
        "evidence": (
            "seed_config_catalogue archives the keys named in RETIRED by setting "
            "is_active False rather than deleting the rows, and evaluation treats an "
            "inactive capability as off. Vendors left as a module and then as a "
            "Core band of Procurement; the platform module's three generic bands "
            "left because that module's keys live in named siblings at the same "
            "depth, so they held nothing and structurally never would."
        ),
        "acceptance": (
            "A retired capability answers off, its historical rows still resolve, "
            "and re-running the seeder does not resurrect it. An empty band is not "
            "by itself a reason to retire one."
        ),
        "limit": (
            "Retirement is a code list. An operator cannot retire a capability from "
            "the console, which is deliberate: it reprices every school at once."
        ),
    },
    {
        "id": "FR-018",
        "title": "Record Every Configuration Change Where It Can Be Found",
        "status": "Implemented",
        "requirement": (
            "Every write this module makes must be recorded locally with its actor, "
            "scope, reason and before-and-after, and mirrored into the platform "
            "audit trail so it can be found beside changes made elsewhere."
        ),
        "evidence": (
            "ConfigurationAuditEvent is a ScopedModel and remains locally "
            "authoritative. The audit service is the single coupling point to the "
            "platform trail, so a change to that service's interface touches one "
            "file in this app rather than every service in it."
        ),
        "acceptance": (
            "A value write, an entitlement change, an override and a depth grant "
            "each produce one local event and one mirrored one. A refused write "
            "produces neither."
        ),
        "limit": (
            "The mirror is best effort by design. A failure to write the platform "
            "copy does not roll back the configuration change it describes."
        ),
    },
    {
        "id": "FR-019",
        "title": "Filter, Facet and Save an Audit View Personally",
        "status": "Implemented",
        "requirement": (
            "An operator investigating a change must be able to narrow the trail, "
            "see what values the filters can take, and keep a narrowing they return "
            "to, without that narrowing becoming everybody's."
        ),
        "evidence": (
            "The audit list, facets and detail endpoints share one filter "
            "implementation with the export path, so a saved view and its export "
            "cannot drift. ConfigurationAuditSavedView is a ScopedModel owned by the "
            "user who created it."
        ),
        "acceptance": (
            "A saved view is visible to its owner and to nobody else. Facets are "
            "computed from the same scoped queryset the list uses, so an offered "
            "filter value always has rows behind it."
        ),
        "limit": (
            "Saved views are personal only. There is no shared or role-scoped view."
        ),
    },
    {
        "id": "FR-020",
        "title": "Export Audit Evidence Without Holding the Request Open",
        "status": "Implemented",
        "requirement": (
            "An export large enough to matter must not be produced inside the "
            "request that asked for it, and the resulting file must reach only the "
            "person who asked."
        ),
        "evidence": (
            "ConfigurationAuditExportJob records the request and the Celery task "
            "run_configuration_audit_export produces the file through the shared "
            "filter implementation. The download endpoint is gated on "
            "config.audit.export and serves the requesting user's own job."
        ),
        "acceptance": (
            "A request returns a job rather than a file. A completed job's file is "
            "refused to a different operator. The exported rows match the filters "
            "the list would have shown."
        ),
        "limit": (
            "Delivery depends on a worker. With no worker running a job stays "
            "pending rather than failing, and nothing on the screen distinguishes "
            "the two."
        ),
    },
    {
        "id": "FR-021",
        "title": "Show What Expires, and Reschedule in Bulk Without Half-Applying",
        "status": "Implemented",
        "requirement": (
            "An operator must be able to see which entitlements end when, and move "
            "many of them at once as a single decision that either lands completely "
            "or not at all."
        ),
        "evidence": (
            "The entitlement calendar reports renewal dates and expiry warnings from "
            "the windows already on the rows. The bulk schedule endpoint applies its "
            "changes atomically, and an explicit denial rejects the whole batch: "
            "existing grants keep their source, and only a target with no row is "
            "created as a manual grant."
        ),
        "acceptance": (
            "A batch containing one refused row changes nothing. A bulk write does "
            "not silently convert a package grant into a manual one."
        ),
        "limit": (
            "The calendar reads dates already stored. It does not model a renewal "
            "that has been agreed and not yet written."
        ),
    },
    {
        "id": "FR-022",
        "title": "Reserve Configuration to the Platform Tenant",
        "status": "Implemented with limits",
        "requirement": (
            "Configuration decides what every school reaches, so writing it is "
            "CodeX's and must not be reachable by a school however its roles are "
            "composed."
        ),
        "evidence": (
            "All nineteen keys in ConfigPermissions are PLATFORM-scoped, which the "
            "RBAC grant guard enforces independently of role composition, so a "
            "school role cannot hold one and cannot be given one. "
            "platform.entitlements.enforce, the switch the plan gate reads, is "
            "declared at platform scope for the same reason: a school-scoped switch "
            "would let a school with config.value.update disable its own plan gate. "
            "The one school-facing route, FR-015, carries no key and is read-only."
        ),
        "acceptance": (
            "A school administrator is refused every configuration endpoint except "
            "the self-scoped capability read. Attempting to grant a config key to a "
            "school role is refused at the grant path, not only at the door."
        ),
        "limit": (
            "The reservation is total. A setting a school ought to own for itself "
            "has no route today and would need a tenant-scoped key introduced "
            "deliberately."
        ),
    },
]


TRACEABILITY = [
    ["Typed configuration definitions and values", "FR-001, FR-004", "Implemented"],
    ["Platform, school, and branch scope resolution", "FR-002", "Implemented"],
    ["Precedence-aware effective-value resolution", "FR-003", "Implemented"],
    ["Configuration catalogue and validation", "FR-001, FR-004, FR-005", "Implemented"],
    ["Capability catalogue and dependencies", "FR-008, FR-009", "Implemented"],
    ["School capability entitlements", "FR-010", "Implemented"],
    ["Scoped capability overrides", "FR-013", "Implemented"],
    ["Effective-capability evaluation", "FR-014", "Implemented"],
    ["Platform settings endpoints", "FR-006", "Implemented"],
    ["Security settings endpoints", "FR-006", "Implemented"],
    ["Integration settings endpoints", "FR-006, FR-007", "Implemented"],
    ["Runtime authentication settings", "FR-003, FR-006", "Implemented"],
    ["Runtime mail and notification settings", "FR-003, FR-006", "Implemented"],
    ["School onboarding defaults", "FR-001, FR-003", "Implemented"],
    ["Configuration and capability audit records", "FR-018", "Implemented"],
    ["Configuration export and seeding commands", "FR-001, FR-017", "Implemented"],
    ["Entitlement renewal calendar and expiry warnings", "FR-021", "Implemented"],
    ["Atomic bulk entitlement scheduling with denial protection", "FR-021", "Implemented"],
    ["Personal configuration-audit saved views", "FR-019", "Implemented"],
    ["Asynchronous, user-owned audit exports", "FR-020", "Implemented"],
    ["Module and band catalogue with ordered depth", "FR-008", "Implemented"],
    ["Depth-limited capability entitlements", "FR-010, FR-011", "Implemented"],
    ["Time-boxed tenant depth grants", "FR-012", "Implemented"],
    ["Depth-aware effective-capability evaluation", "FR-011, FR-014", "Implemented"],
    ["Capability retirement without deletion", "FR-017", "Implemented"],
    ["A school reading the capabilities its own plan reaches", "FR-015, FR-016", "Implemented"],
]


NEEDS_ATTENTION = [
    ["1", "No setting a school may own",
     "Every one of the nineteen configuration keys is PLATFORM-scoped, so the "
     "only thing a school reaches on its own side is the read in FR-015, which "
     "carries no key at all. The moment a setting genuinely belongs to a school "
     "- a notification preference, a branding choice - it needs a tenant-scoped "
     "key and a write path, and neither exists.",
     "Introduce a tenant-scoped configuration key and the school-facing write "
     "path for it, deliberately and one setting at a time."],
    ["2", "The plan gate is one switch for the whole platform",
     "platform.entitlements.enforce is declared at platform scope, which is what "
     "stops a school disabling its own gate, and it also means enforcement is on "
     "for everybody or nobody. A staged rollout, or holding one school harmless "
     "while a mis-banding is corrected, is not expressible.",
     "A platform-writable per-tenant value, so the switch stays out of a "
     "school's hands while still being settable per school."],
    ["3", "A grant with no depth reaches everything",
     "Null depth means no limit rather than the shallowest limit, which is what "
     "keeps schools granted before depth existed from silently losing Advanced "
     "work. It also means any row written without a depth reaches every band. "
     "The rule retires itself as plans are applied and nothing enforces that it "
     "has.",
     "A report of PACKAGE entitlements still carrying a null depth, and a "
     "decision to backfill once it is empty."],
    ["4", "Five modules are priced but not built",
     "Attendance, Gradebook, Parents Management and both portals carry the full "
     "three-band shape and hold no permission keys, so their depths describe a "
     "price list rather than a product boundary. That is the intended shape for "
     "what is coming, and it means the catalogue cannot be read as a statement "
     "of what exists.",
     "No action while they are unbuilt. The distinction is recorded so an empty "
     "band is not mistaken for a retirement candidate."],
    ["5", "A pending export job and a lost one look the same",
     "Audit exports are produced by a Celery task. With no worker running, a job "
     "stays pending indefinitely and the screen shows the same state it shows "
     "one second after the request.",
     "A staleness threshold after which a pending job is reported as stalled."],
]


def add_real_bullets(doc, items, *, size=9):
    for item in items:
        paragraph = doc.add_paragraph()
        set_bullet_numbering(paragraph)
        write_paragraph(paragraph, item, size=size, space_after=2)


def set_bullet_numbering(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "1")
    num_pr.append(ilvl)
    num_pr.append(num_id)


def hold_requirement_together(doc):
    """Stop a requirement's heading and status banner stranding on their own.

    add_requirement builds a heading, a banner row and four content rows. Left
    to itself the table breaks wherever the page ends, and FR-022 landed with
    its heading and banner at the foot of one page and every word of its
    content on the next. Rows are made unsplittable, and the heading, the
    banner and the first content row are tied to what follows them, so the
    break can only fall between content rows.
    """
    table = doc.tables[-1]
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
    for row in table.rows[:2]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    for paragraph in reversed(doc.paragraphs):
        if paragraph.text.strip().startswith("FR-"):
            paragraph.paragraph_format.keep_with_next = True
            break


def add_callout(doc, title, lines, *, kind="info"):
    table = add_reference_callout(doc, title, lines, kind=kind)
    cell = table.cell(0, 0)
    for paragraph in cell.paragraphs[1:]:
        text = paragraph.text.removeprefix("• ")
        set_bullet_numbering(paragraph)
        write_paragraph(paragraph, text, size=8.6, space_after=2)
    return table


def build(reference: Path, output: Path) -> None:
    doc = Document(str(reference))
    remove_body_content(doc)
    title = (
        "XVS M06 Configuration and Capability Management Functional "
        f"Requirements Document v{FRD_VERSION}"
    )
    doc.core_properties.title = title
    doc.core_properties.subject = (
        "Code-aligned functional requirements for XVS Module 6"
    )
    doc.core_properties.author = "CodeX Team"
    doc.core_properties.version = FRD_VERSION
    set_headers(
        doc,
        "CodeX | Configuration & Capability Management | Functional "
        "Requirements Document (FRD)",
    )

    add_cover(
        doc,
        family="Functional Requirements Document",
        title="Module 6: Configuration & Capability Management",
        subtitle="XVision Systems | Code-aligned functional baseline",
        version=FRD_VERSION,
    )

    add_heading(doc, "Document Control", level=1)
    add_metadata_table(
        doc,
        [
            ("Document", "Functional Requirements Document (FRD)"),
            ("Module", "M06 | Configuration & Capability Management"),
            ("Version", FRD_VERSION),
            ("Review date", REVIEW_DATE),
            ("Code baseline", CODE_BASELINE),
            ("Source MRD",
             f"XVS Module Requirements Document v{MRD_VERSION} | Module 6, "
             "twenty-six capability entries"),
            ("Primary app", "vs_config"),
            ("Supporting apps",
             "core, vs_tenants, vs_rbac, vs_audit, schools/vs_schools"),
            ("Status", "Code-aligned baseline for Product and Engineering review"),
            ("Owner", "CodeX Team"),
        ],
    )
    add_callout(
        doc,
        "Evidence boundary",
        [
            "Implemented means the stated backend path is present in the "
            "inspected code. It does not prove frontend completion, deployment, "
            "production adoption or data migration.",
            "This module decides what every other module may be reached for, so "
            "a claim here is a claim about the whole platform's behaviour. "
            "Where a rule is enforced somewhere else, that place is named.",
            "The MRD and this FRD carry independent versions. A change can "
            "require one or both to advance.",
        ],
        kind="info",
    )
    add_page_break(doc)

    add_heading(doc, "Table of Contents", level=1)
    add_table(
        doc,
        ["Section", "Purpose"],
        [
            ["Document Control", "Version, ownership, baseline and evidence rules"],
            ["1. Purpose and Scope", "What this module owns and what it must not know"],
            ["2. Context and Status Model", "The two catalogues, the scope chain and depth"],
            ["3. Actors, Permissions, and Ownership", "Who may read and change configuration"],
            ["4. Functional Requirements", "Testable behaviour, evidence, acceptance and limits"],
            ["5. Workflows and Lifecycle Rules", "Resolution, evaluation and the plan chain"],
            ["6. Data Model and Relationships", "What is persisted, and the safety each row carries"],
            ["7. API and Validation Contracts", "Routes, keys, refusals and response rules"],
            ["8. Dependencies and Operational Evidence", "What this module needs, and what needs it"],
            ["9. Needs Attention", "Current gaps and required completion"],
            ["10. MRD Traceability", "All twenty-six Module 6 capability mappings"],
            ["11. Change Log", "Independent FRD revision history"],
        ],
        [2.35, 4.92],
        font_size=8.5,
    )
    add_body(
        doc,
        "Use the Word Navigation pane to jump between headings. This contents "
        "page is intentionally static for reliable headless rendering.",
        size=8.5,
        color=GREY,
    )
    add_page_break(doc)

    add_heading(doc, "1. Purpose and Scope", level=1)
    add_body(
        doc,
        "Module 6 answers two questions for the whole platform: what is this "
        "setting's value here, and may this school reach this thing at all. "
        "Everything else in XVS depends on it and it depends on nothing, which "
        "is the constraint that shapes the module: it holds typed declarations "
        "and a capability catalogue, and it never learns what a student, an "
        "invoice or a purchase order is.",
    )
    add_heading(doc, "1.1 In Scope", level=2)
    add_real_bullets(
        doc,
        [
            "Typed configuration definitions, their values at platform, school "
            "and branch scope, and precedence-aware resolution between them.",
            "The capability catalogue two levels deep: modules a school is sold, "
            "and bands cutting each module at Core, Plus or Advanced.",
            "Entitlements carrying a depth and a window, time-boxed depth "
            "uplifts, and scoped overrides in both directions.",
            "Effective evaluation for one capability and for the whole "
            "catalogue, including the self-scoped read a school's own screens "
            "use.",
            "The platform, security and integration settings screens, and a "
            "read-only connection test.",
            "The configuration audit trail, its facets and personal saved "
            "views, and asynchronous exports of it.",
        ],
    )
    add_heading(doc, "1.2 Out of Scope", level=2)
    add_real_bullets(
        doc,
        [
            "Refusing a request. This module answers whether a capability is "
            "on; vs_rbac's plan gate is what turns that answer into a 403, and "
            "the refusal's wording belongs there.",
            "Which permission key belongs to which band. That mapping lives in "
            "vs_rbac beside the keys it classifies.",
            "What a school pays. PackagePlan is Module 1's, and no rate, floor, "
            "billing unit or term exists anywhere yet.",
            "Deciding a plan. Translating a plan into grants is "
            "vs_schools' apply_plan_entitlements, which writes rows this module "
            "then evaluates.",
        ],
    )
    add_page_break(doc)

    add_heading(doc, "2. Context and Status Model", level=1)
    add_body(
        doc,
        "Two catalogues sit side by side. The configuration catalogue declares "
        "settings and holds their values; the capability catalogue declares "
        "what is sellable and holds who has it. They share one idea, which is "
        "the scope chain: a row belongs to the platform, to a school, or to a "
        "branch, and a read walks from the nearest scope outward.",
    )
    add_heading(doc, "2.1 Requirement Status", level=2)
    add_status_key(doc)
    add_heading(doc, "2.2 Depth Vocabulary", level=2)
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Module", "A capability a school is sold. Carries no depth of its "
                       "own and requires an entitlement."],
            ["Band", "That module cut at one depth. Requires no entitlement; it "
                     "opens when the module's depth reaches it."],
            ["Core / Plus / Advanced", "The three depths, stored as 10, 20 and "
                                       "30 so a fourth can be inserted without "
                                       "renumbering."],
            ["Unlimited", "A null depth. No limit rather than the shallowest "
                          "limit, so a grant written before depth existed still "
                          "reaches everything."],
            ["Uplift", "A time-boxed depth grant for one tenant and one module. "
                       "Raises, never lowers, and survives a plan change."],
            ["Provisioned", "The tenant holds at least one PACKAGE entitlement. "
                            "A tenant with none was never given a plan, which is "
                            "not the same as having bought nothing."],
        ],
        [1.75, 5.52],
        font_size=8.5,
    )
    add_page_break(doc)

    add_heading(doc, "3. Actors, Permissions, and Ownership", level=1)
    add_heading(doc, "3.1 Permission Matrix", level=2)
    add_table(
        doc,
        ["Actor", "May do", "Governed by"],
        [
            ["CodeX platform staff",
             "Declare settings, write values at any scope, manage the capability "
             "catalogue, grant and schedule entitlements, set overrides and "
             "uplifts, read and export the audit trail.",
             "The nineteen config.* keys, all PLATFORM-scoped."],
            ["A school's own users",
             "Read what their own school's plan reaches, and nothing else in "
             "this module.",
             "No key. The route is self-scoped and read-only (FR-015)."],
            ["Every other module",
             "Read a resolved value or an effective capability. Never write "
             "either.",
             "conf.get_config and the capability services."],
            ["The owning module of a setting",
             "Refuse a value its own domain makes illegal, without this module "
             "importing it.",
             "A guard registered from AppConfig.ready (FR-005)."],
        ],
        [1.55, 3.6, 2.12],
        font_size=8.3,
    )
    add_heading(doc, "3.2 Ownership Boundaries", level=2)
    add_real_bullets(
        doc,
        [
            "Configuration depends on nothing. Every rule needing another "
            "module's knowledge arrives as a registered guard, never an import.",
            "vs_rbac owns the permission-to-band map and the refusal. This "
            "module owns the catalogue those bands live in.",
            "vs_schools owns the plan and writes the grants. This module owns "
            "what a grant means once written.",
            "vs_audit receives a mirror of every configuration change. The local "
            "event stays authoritative.",
        ],
    )
    add_page_break(doc)

    add_heading(doc, "4. Functional Requirements", level=1)
    add_body(
        doc,
        "Each requirement records the required behaviour, the inspected "
        "evidence, the acceptance boundary and the current limit. Status is "
        "current state, not revision history.",
    )
    # No forced breaks between requirements. Each block already fills a page
    # naturally, so a break placed by index lands where the page has just
    # ended and leaves a blank sheet behind it.
    for requirement in REQUIREMENTS:
        add_requirement(doc, requirement)
        hold_requirement_together(doc)

    add_page_break(doc)
    add_heading(doc, "5. Workflows and Lifecycle Rules", level=1)
    add_heading(doc, "5.1 Resolving One Setting", level=2)
    add_table(
        doc,
        ["Step", "What happens", "Effect"],
        [
            ["1", "The caller names a key and a scope.",
             "An undeclared key is refused rather than defaulted."],
            ["2", "Candidate rows for the whole scope chain are fetched by "
                  "scope_key in one query.",
             "Branch, school and platform rows arrive together."],
            ["3", "The nearest scope holding a row wins.",
             "A branch value shadows its school's without deleting it."],
            ["4", "With no row anywhere, the definition's default answers.",
             "A key never written still resolves."],
        ],
        [0.6, 3.4, 3.27],
        font_size=8.4,
    )
    add_heading(doc, "5.2 Deciding Whether a Capability Is On", level=2)
    add_table(
        doc,
        ["Step", "What happens", "Effect"],
        [
            ["1", "An inactive capability answers off immediately.",
             "A retired row cannot be reached through history."],
            ["2", "For a band, its module is evaluated first.",
             "A band whose module is off is off, whatever its depth."],
            ["3", "For a band, the tenant's depth for that module is resolved: "
                  "live uplift, then entitlement depth, then platform grant, "
                  "then unlimited.",
             "One comparison decides every band of that module."],
            ["4", "For a module, its entitlement must be GRANTED and inside its "
                  "window.",
             "An expired subscription closes what it paid for."],
            ["5", "Every dependency is evaluated, cycles refused.",
             "Procurement is off if Finance is."],
            ["6", "The scope chain is read nearest-first for a non-INHERIT "
                  "override.",
             "A branch override beats a tenant one, which beats a platform one."],
        ],
        [0.6, 3.4, 3.27],
        font_size=8.4,
    )
    add_heading(doc, "5.3 From a Plan to a Refusal", level=2)
    add_body(
        doc,
        "The chain crosses three modules and is worth stating once. Module 1 "
        "applies a plan, writing one PACKAGE entitlement per module at the "
        "depth the plan reaches. Module 6 evaluates those rows into an "
        "on-or-off answer per capability. Module 4 maps each permission key to "
        "the band that governs it and turns a negative answer into "
        "PLAN_UPGRADE_REQUIRED. A school-facing screen asks Module 6 directly, "
        "through the self-scoped read, so what a menu shows and what the gate "
        "allows come from one evaluation.",
        size=9,
    )
    add_page_break(doc)

    add_heading(doc, "6. Data Model and Relationships", level=1)
    add_table(
        doc,
        ["Model", "Holds", "Safety contract"],
        [
            ["ConfigurationDefinition",
             "Key, type, default, validation rules and allowed scopes.",
             "Archived rather than deleted, so audit rows keep resolving."],
            ["ConfigurationValue",
             "One value per definition per scope.",
             "ScopedModel; unique on definition and scope_key."],
            ["Capability",
             "Modules and their bands, with depth and an active flag.",
             "Check constraint: a band has both a parent and a depth, and no "
             "band is a band of a band."],
            ["CapabilityDependency",
             "One capability requiring another.",
             "Unique on the pair; a capability cannot require itself."],
            ["CapabilityEntitlement",
             "What a tenant holds, at what depth, from what source, over what "
             "window.",
             "Unique on capability and scope. Null depth means unlimited."],
            ["CapabilityDepthGrant",
             "A time-boxed uplift for one tenant and one module.",
             "Unique on capability and tenant; refuses a band."],
            ["CapabilityOverride",
             "An operator switching one capability on or off at one scope.",
             "ScopedModel; unique on capability and scope_key."],
            ["ConfigurationAuditEvent",
             "Actor, scope, reason and before-and-after for every change.",
             "ScopedModel; locally authoritative, mirrored to vs_audit."],
            ["ConfigurationAuditSavedView",
             "A narrowing of the trail one operator returns to.",
             "Owned by its creator and visible to nobody else."],
            ["ConfigurationAuditExportJob",
             "A requested export and the file it produced.",
             "Served only to the operator who asked for it."],
        ],
        [1.7, 2.75, 2.82],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "7. API and Validation Contracts", level=1)
    add_body(
        doc,
        "Routes are mounted under /v1/config/ with the standard response "
        "envelope. Every route below is PLATFORM-scoped through its key except "
        "the self-scoped capability read, which carries no key at all.",
        size=9,
    )
    add_heading(doc, "7.1 Configuration", level=2)
    add_table(
        doc,
        ["Method and path", "Purpose and permission"],
        [
            ["GET, POST /config/definitions/",
             "The declaration catalogue. config.definition.view to read, "
             ".create to add."],
            ["GET, PATCH, DELETE /config/definitions/{key}/",
             "One declaration. .update to change, .archive to retire."],
            ["GET, POST /config/values/",
             "Values at a scope. config.value.view and .update."],
            ["DELETE /config/values/{key}/",
             "Clear a value at a scope, exposing the one behind it. "
             ".value.update."],
            ["GET /config/effective-values/ and /{key}/",
             "What is in force for a scope after precedence. .value.view."],
            ["GET, PATCH /config/platform-settings/",
             "The platform settings screen as one read and one write."],
            ["GET, PATCH /config/security-settings/",
             "The security screen. config.security.view and .manage."],
            ["GET, PATCH /config/integration-settings/",
             "The integration screen. config.integration.view and .manage."],
            ["POST /config/integration-settings/test/",
             "A bounded, read-only connection check. .integration.manage, and "
             "rate limited by a thirty-second cooldown."],
            ["GET /config/export/", "Configuration export. config.export.create."],
        ],
        [2.7, 4.57],
        font_size=8.3,
    )
    add_heading(doc, "7.2 Capabilities and Entitlements", level=2)
    add_table(
        doc,
        ["Method and path", "Purpose and permission"],
        [
            ["GET, POST /config/capabilities/",
             "The capability catalogue. config.capability.view and .manage."],
            ["GET, PATCH, DELETE /config/capabilities/{key}/",
             "One capability, including retirement by archiving."],
            ["GET, POST /config/entitlements/",
             "What a tenant holds. config.entitlement.view and .manage."],
            ["DELETE /config/entitlements/{capability}/",
             "Withdraw one entitlement. .entitlement.manage."],
            ["GET /config/entitlements/calendar/",
             "Renewal dates and expiry warnings. .entitlement.view."],
            ["POST /config/entitlements/bulk-schedule/",
             "Reschedule many at once, atomically. .entitlement.manage."],
            ["GET, POST /config/overrides/",
             "Scoped overrides. config.override.view and .manage."],
            ["GET /config/effective-capabilities/",
             "The literal evaluated state for a scope. config.capability.view, "
             "which no school role holds."],
            ["GET /config/my-capabilities/",
             "What the caller's own tenant reaches. No permission key; open "
             "before go-live; states only."],
        ],
        [2.7, 4.57],
        font_size=8.3,
    )
    add_heading(doc, "7.3 Audit", level=2)
    add_table(
        doc,
        ["Method and path", "Purpose and permission"],
        [
            ["GET /config/audit-events/ and /{event_id}/",
             "The trail and one event. config.audit.view."],
            ["GET /config/audit-events/facets/",
             "Filter values with rows behind them. .audit.view."],
            ["GET /config/audit-events/export/",
             "Direct export. config.audit.export."],
            ["GET, POST /config/audit-events/saved-views/",
             "Personal saved narrowings. .audit.view."],
            ["DELETE /config/audit-events/saved-views/{view_id}/",
             "Remove one of the caller's own. .audit.view."],
            ["GET, POST /config/audit-events/export-jobs/",
             "Request an asynchronous export and list your own."],
            ["GET /config/audit-events/export-jobs/{job_id}/download/",
             "Take the produced file. .audit.export, requester only."],
        ],
        [2.7, 4.57],
        font_size=8.3,
    )
    add_page_break(doc)
    add_heading(doc, "7.4 Refusals", level=2)
    add_table(
        doc,
        ["Condition", "Answer"],
        [
            ["A value for an undeclared key", "Refused. There is no such setting."],
            ["A value of the wrong type, outside its choices or bounds",
             "Refused, naming the setting. No row and no audit event."],
            ["A value written at a scope the definition does not allow",
             "Refused. allowed_scopes is part of the declaration."],
            ["A value the owning module's guard rejects",
             "Refused as the setting being rejected, carrying the guard's own "
             "reason."],
            ["A module carrying a depth, or a band without one",
             "Refused by check constraint and by clean()."],
            ["A depth uplift naming a band",
             "Refused, naming the module to use instead."],
            ["A bulk schedule containing one denied row",
             "The whole batch is rejected. Nothing is half-applied."],
            ["A second connection test inside the cooldown",
             "Refused rather than queued."],
            ["A school asking /config/effective-capabilities/",
             "403. The key is PLATFORM-scoped and no school role holds it."],
            ["A school asking /config/my-capabilities/ for another tenant",
             "404, before the view runs. The tenant comes from the validated "
             "assertion."],
            ["An export job downloaded by a different operator",
             "Refused. A job serves the person who asked for it."],
        ],
        [3.0, 4.27],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "8. Dependencies and Operational Evidence", level=1)
    add_table(
        doc,
        ["Dependency", "Contract"],
        [
            ["vs_tenants",
             "Supplies the tenant and branch a scope is expressed in, and "
             "resolves a branch inside a tenant. This module reads them and "
             "writes neither."],
            ["vs_rbac",
             "Enforces the nineteen PLATFORM-scoped keys, and consumes this "
             "module's evaluation in the plan gate. The permission-to-band map "
             "lives there, not here."],
            ["vs_audit",
             "Receives the mirror of every configuration change through one "
             "coupling point in this app."],
            ["Module 1, School and Branch Management",
             "Owns PackagePlan and writes PACKAGE entitlements through "
             "apply_plan_entitlements. This module evaluates what it wrote."],
            ["Celery and a worker",
             "Produces asynchronous audit export files. Nothing else in this "
             "module needs a worker."],
            ["Every other module",
             "Reads resolved values and effective capabilities, and may "
             "register a guard for its own setting. None writes configuration."],
        ],
        [2.1, 5.17],
        font_size=8.3,
    )
    add_heading(doc, "8.1 Verification Evidence", level=2)
    add_body(
        doc,
        "The module's own suite runs green at the code baseline named in "
        "Document Control: 105 configuration tests across the catalogue, the "
        "depth model and the self-scoped capability read. Its principal "
        "consumer, vs_rbac, runs 514 tests including the plan gate and the "
        "agreement between the role builder and that gate. Backend evidence "
        "only; nothing here is deployed.",
        size=9,
    )
    add_page_break(doc)

    add_heading(doc, "9. Needs Attention", level=1)
    add_body(
        doc,
        "Current state, not history. An item leaves this section when "
        "implementation and verification resolve it, and is rewritten when the "
        "risk changes shape.",
        size=9,
        color=GREY,
    )
    add_table(
        doc,
        ["Pri.", "Current gap", "Detail", "Required completion"],
        NEEDS_ATTENTION,
        [0.45, 1.55, 3.3, 1.97],
        font_size=8.2,
    )
    add_page_break(doc)

    add_heading(doc, "10. MRD Traceability", level=1)
    add_body(
        doc,
        f"Module 6 of XVS Module Requirements Document v{MRD_VERSION} lists "
        "twenty-six capabilities. Each maps to the requirements above.",
        size=9,
    )
    add_table(
        doc,
        ["MRD capability", "FRD requirement", "State"],
        TRACEABILITY,
        [3.3, 1.97, 2.0],
        font_size=8.3,
    )
    add_page_break(doc)

    add_heading(doc, "11. Change Log", level=1)
    add_table(
        doc,
        ["Version", "Date", "Change"],
        [[
            FRD_VERSION, REVIEW_DATE,
            "First code-aligned baseline for Module 6. Records the two "
            "catalogues and the one scope chain they share; declaration, "
            "validation and the registered guard that lets an owning module "
            "judge its own setting without this module importing it; the "
            "capability catalogue two levels deep, with depth spaced by ten and "
            "a null depth meaning unlimited rather than shallowest; "
            "entitlements, time-boxed uplifts and scoped overrides, and the "
            "single evaluation order the API, the navigation and the plan gate "
            "all read; the self-scoped capability read a school's own screens "
            "use, and the provisioned-versus-bought distinction that keeps a "
            "school created before grants existed from losing a product it "
            "pays for; retirement by archiving; and the audit trail, its "
            "personal saved views and its asynchronous exports. Five current "
            "gaps are recorded, the first being that every configuration key is "
            "platform-scoped, so a setting a school ought to own has no route. "
            "Backend evidence only; nothing here is deployed.",
        ]],
        [0.85, 1.25, 5.17],
        font_size=8.2,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root)
    reference = (
        root / "functional-requirements" / "10-bulk-data-import"
        / "XVS_M10_Bulk_Data_Import_Functional_Requirements_Document_v1.2.docx"
    )
    folder = root / "functional-requirements" / "06-configuration-and-capability"
    output = folder / (
        "XVS_M06_Configuration_and_Capability_Management_Functional_"
        f"Requirements_Document_v{FRD_VERSION}.docx"
    )
    build(reference, output)
    print(f"Wrote {output}")


if __name__ == "__main__":
    main()
