#!/usr/bin/env python3
"""Cut MRD v2.87 and three module FRDs for staff ID sign-in and the one identity trail.

What changed in the backend, and therefore in the documents:

* Sign-in and the self-service password reset take an ``identifier``: an email
  address, or a school's staff number when it holds no ``@``. A staff number
  resolves only inside the school the request names, whatever the tenant switch
  says, matches without regard to case, and signs nobody in when two numbers at
  one school differ only in case. A reset still goes to the email on file.
* Changing a sign-in email reissues the invitation of an account that has not
  been activated, so the link sent to the old address stops working. The school
  staff email route asks Field Access on ``email`` and takes an optional note.
* The staff history reads its account half from the identity AuditEvents that
  log_auth_event writes. It read AuthEventLog, which nothing outside the dev
  seeder wrote, so every school's account history was empty.
* AuthEventLog is dropped by vs_user migration 0015. Its event names live on as
  vs_user.auth_events.AuthEvent, and the /auth-events/ viewset is renamed
  AuthEventViewSet to match what it already read.

M03 also said the sign-in tenant switch was off. It has been on since 8f3b28b6
(20 August 2026); the claims that depended on it are corrected here.

    python tools/patch_staff_id_and_auth_events_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_record_history_docs import (
    ROOT,
    add_fr,
    append_rows,
    finish,
    frd_path,
    keep_format,
    log_change,
    replace_cell,
    set_control,
    set_cover_version,
    table_with_header,
    update_reconciliation,
)

REVIEW_DATE, SHORT_DATE = "26 September 2026", "26 Sep 2026"
MRD_SOURCE, MRD_TARGET = "2.86", "2.87"
CODE_BASELINE = (
    "Backend main at c9031ad7, with school-fe main at 0b3cf1b and console-fe main at "
    "bb3f3d9, 26 September 2026"
)
TEST_EVIDENCE = (
    "Verified by the vs_user suite (417 tests) and the schools.vs_staff suite (316 tests), "
    "each run on its own and all passing, 20 of them this change's own; and driven against "
    "the running school app."
)

# ── editing helpers ──────────────────────────────────────────────────────────


def fr_table(doc, fr: str):
    hits = [t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith(f"{fr} ")]
    if len(hits) != 1:
        raise ValueError(f"{fr} heads {len(hits)} tables")
    return hits[0]


def set_status(doc, table, status: str) -> None:
    """Rewrite a status FRD's merged 'FR-0NN | State' header, colour included.

    The header row's fill and text colour encode the state, so the row is taken
    from a requirement that already holds ``status`` and only its label is
    rewritten. Rewriting the text alone leaves an Implemented requirement drawn
    in the amber of Implemented with limits.
    """
    fr = table.rows[0].cells[0].text.split("|")[0].strip()
    templates = [t for t in doc.tables
                 if t is not table and "|" in t.rows[0].cells[0].text
                 and t.rows[0].cells[0].text.split("|", 1)[1].strip() == status]
    if not templates:
        raise ValueError(f"No requirement holds status {status!r} to copy its colour from")
    header = copy.deepcopy(templates[0].rows[0]._tr)
    table.rows[0]._tr.addprevious(header)
    table._tbl.remove(table.rows[1]._tr)
    for cell in table.rows[0].cells:
        keep_format(cell, f"{fr} | {status}")


def row(table, label: str):
    hits = [r for r in table.rows if r.cells[0].text.strip() == label]
    if len(hits) != 1:
        raise ValueError(f"{label!r} labels {len(hits)} rows")
    return hits[0]


def row_starting(table, start: str):
    hits = [r for r in table.rows if r.cells[0].text.strip().startswith(start)]
    if len(hits) != 1:
        raise ValueError(f"{start!r} starts {len(hits)} rows")
    return hits[0]


def set_value(table, label: str, text: str) -> None:
    keep_format(row(table, label).cells[-1], text)


def edit_cell(cell, old: str, new: str) -> None:
    text = cell.text
    if text.count(old) != 1:
        raise ValueError(f"{old[:60]!r} occurs {text.count(old)} times in the cell")
    keep_format(cell, text.replace(old, new))


def edit_value(table, label: str, old: str, new: str) -> None:
    edit_cell(row(table, label).cells[-1], old, new)


def remove_row(table, start: str) -> None:
    target = row_starting(table, start)
    target._tr.getparent().remove(target._tr)


def paragraph_starting(doc, start: str) -> Paragraph:
    hits = [p for p in doc.paragraphs if p.text.strip().startswith(start)]
    if len(hits) != 1:
        raise ValueError(f"{start!r} starts {len(hits)} paragraphs")
    return hits[0]


def edit_paragraph(doc, start: str, old: str, new: str) -> None:
    paragraph = paragraph_starting(doc, start)
    if paragraph.text.count(old) != 1:
        raise ValueError(f"{old[:60]!r} occurs {paragraph.text.count(old)} times")
    mrd_tools.retitle(paragraph, paragraph.text.replace(old, new))


def insert_after(doc, start: str, text: str) -> None:
    """Add a paragraph after the one starting with ``start``, formatted as it is."""
    paragraph = paragraph_starting(doc, start)
    clone = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(clone)
    mrd_tools.retitle(Paragraph(clone, paragraph._parent), text)


def assert_absent(doc, *needles: str) -> None:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for r in table.rows:
            texts.extend(c.text for c in r.cells)
    blob = "\n".join(texts)
    for needle in needles:
        at = blob.find(needle)
        if at >= 0:
            context = blob[max(0, at - 160):at + 80].replace("\n", " / ")
            raise ValueError(f"{needle!r} is still in the document: ...{context}...")


# ── structural repair ────────────────────────────────────────────────────────

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
#: CT_TcPr's child order in the transitional schema Word validates against.
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns",
    "cellDel", "cellMerge", "tcPrChange",
]
TCMAR_ORDER = ["top", "left", "bottom", "right"]


def repair_ooxml(doc) -> None:
    """Put back the schema order earlier revisions' table edits broke.

    Cloned and rebuilt tables left three defects that Word stops on when it
    opens the file: table-cell properties out of order, cell margins written as
    start and end where the transitional schema names left and right, and a
    table grid placed after the rows it describes. Each is reordered or
    renamed, a start or end margin repeating a left or right one is dropped, and
    a shading that names only a fill gains the required val="clear", which is
    how Word draws a fill with no pattern.
    """
    body = doc.element.body
    rank = {name: i for i, name in enumerate(TCPR_ORDER)}
    for tcpr in body.iter(W + "tcPr"):
        for child in sorted(tcpr, key=lambda c: rank.get(c.tag.replace(W, ""), len(rank))):
            tcpr.append(child)
    for mar in body.iter(W + "tcMar"):
        for bidi, ltr in (("start", "left"), ("end", "right")):
            for child in mar.findall(W + bidi):
                if mar.find(W + ltr) is not None:
                    mar.remove(child)
                else:
                    child.tag = W + ltr
        for child in sorted(mar, key=lambda c: TCMAR_ORDER.index(c.tag.replace(W, ""))
                            if c.tag.replace(W, "") in TCMAR_ORDER else len(TCMAR_ORDER)):
            mar.append(child)
    for shd in body.iter(W + "shd"):
        if shd.get(W + "val") is None:
            shd.set(W + "val", "clear")
    for tbl in body.iter(W + "tbl"):
        grid, props = tbl.find(W + "tblGrid"), tbl.find(W + "tblPr")
        if grid is not None and props is not None and props.getnext() is not grid:
            props.addnext(grid)



def normalise_change_log(doc) -> None:
    """Give every version row the change log's ordinary row format.

    An earlier revision restored a missing header by copying the header row,
    and the version rows cloned from it since inherited its fill, and sometimes
    its repeat-as-header flag, so Word drew them in the header's colours and
    printed some again at the top of every page of the log. Each such row is
    rebuilt from the first ordinary version row, keeping its text.
    """
    from patch_record_history_docs import change_log_table, is_header_styled

    table = change_log_table(doc)

    def is_header(r):
        return is_header_styled(table, r)

    ordinary = next((r for r in table.rows[1:] if not is_header(r)), None)
    if ordinary is None:
        raise ValueError("The change log has no ordinary version row to copy")
    for r in list(table.rows[1:]):
        if not is_header(r):
            continue
        texts = [c.text for c in r.cells]
        clone = copy.deepcopy(ordinary._tr)
        r._tr.addprevious(clone)
        table._tbl.remove(r._tr)
        rebuilt = next(x for x in table.rows if x._tr is clone)
        for cell, text in zip(rebuilt.cells, texts):
            keep_format(cell, text)


# ── M03 Identity, Team & Organogram ──────────────────────────────────────────

M03_DIR = "03-identity-team-and-organogram"
M03_STEM = "XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document"
M03_SOURCE, M03_TARGET = "1.16", "1.17"

M03_FR024 = [
    ("Requirement", (
        "A school's staff member must be able to sign in, and ask for a password reset, with "
        "the staff ID their school issued, in the same box that takes an address. A staff ID is "
        "the school's own label and is unique only inside that school, so it must never resolve "
        "anywhere else, and a refusal must look exactly like a wrong password."
    )),
    ("Current evidence", (
        "LoginRequestSerializer and PasswordResetRequestSerializer take identifier: a value "
        "holding an @ is read as an address and anything else as a staff number, and exactly one "
        "of an address, a staff number and an ID-card key is accepted. "
        "services.sign_in_scope.resolve_staff_number_account resolves the asserted tenant, then "
        "reads StaffProfile rows of that tenant only, through all_objects with the tenant named "
        "explicitly because no ambient tenant exists before sign-in. It refuses as "
        "TENANT_REQUIRED when no tenant is named, whatever REQUIRE_TENANT_ON_SIGN_IN says, "
        "because no platform-wide answer exists to who STF/0012 is. The match ignores case; "
        "stored uniqueness is case-sensitive, so two numbers differing only in case are "
        "ambiguous and sign nobody in. An unknown slug runs the same query against the sentinel "
        "key, as the address path does. LoginService.login records the number as typed in "
        "AuthAttempt.email_entered. PasswordService.request_reset resolves the same way and "
        "sends the link to the address on the account, never to what was typed."
    )),
    ("Acceptance", (
        "StaffIdSignInTests: Brightfield's BFS/STF/0012 signs its owner in, in any case and "
        "with surrounding spaces; the same number at Sunrise, with no tenant, or with a wrong "
        "password is INVALID_CREDENTIALS; two Brightfield numbers differing only in case sign "
        "nobody in; the attempt log holds the number as typed; an address in the same box still "
        "signs in. StaffIdPasswordResetTests: a staff number issues a reset for its owner at "
        "their school only, an unknown number and one at the wrong school answer byte-for-byte "
        "as a known one does, an address in the same box still works, and an empty box is 400."
    )),
    ("Current limit", (
        "A staff number is usually guessable (a school prefix and a counter), so it narrows a "
        "password-guessing attack to one school's numbering scheme. The per-account lockout "
        "(FR-010) and the IP throttle bound it; there is no MFA (gap 9.1). The attempt log's "
        "email_entered column holds a staff number when one was typed."
    )),
]

M03_FR025 = [
    ("Requirement", (
        "Correcting the address of an account that has not been activated must kill the "
        "invitation link already sent. The usual reason for the correction is a mistyped "
        "invitation, and whoever holds that link could otherwise set the password and sign in "
        "as the new member of staff."
    )),
    ("Current evidence", (
        "EmailChangeService.change_email, which both the platform route and the school staff "
        "route call, reissues the invitation through InvitationService.resend when the account "
        "is PENDING and its invitation is unused: the token is rotated, so the old link dies at "
        "once, and a fresh invitation is queued to the new address after commit and recorded as "
        "INVITATION_SENT. It also ends every session and records EMAIL_CHANGED with both "
        "addresses and, when given, the administrator's note."
    )),
    ("Acceptance", (
        "StaffEmailCorrectionTests: correcting funke@gmial.test to funke@gmail.test rotates "
        "the invitation's token and leaves it unused; an activated account's used invitation "
        "is left alone; a role whose Field Access leaves email read-only is refused 403 and "
        "the address is unchanged; the history shows the change and its note and neither "
        "address."
    )),
    ("Current limit", (
        "An activated account is signed out and signs in with the new address at once; nothing "
        "asks the new address to confirm first, so a mistyped correction to an active account "
        "stands until somebody notices. A school's staff can still sign in with their staff ID "
        "(FR-024) in that case."
    )),
]

M03_ATTENTION_EDITS = {
    "P2": [
        ("9.5 An actor that is not a User row", "AuthEventLog.actor and .subject, and vs_audit's "
         "actor_user, are foreign keys", "vs_audit's actor_user is a foreign key"),
        ("9.8 Two response shapes and one dead serializer", " AuthEventLogReadSerializer has no "
         "caller at all.", ""),
    ],
}

M03_TRACE_EDITS = [
    ("Tenant-aware email and password login", "FR-006, FR-009, FR-024", (
        "Implemented. The tenant is resolved before the account and the lookup scoped to it, "
        "and a sign-in naming no tenant is refused. A school's staff may sign in with their "
        "staff ID, resolved inside that school only.")),
    ("Password reset and password change", "FR-013, FR-014, FR-024", (
        "Implemented with limits. Each reset owns a hashed pr_ token, exact-row locked "
        "consumption prevents revival and replay, and mail names the account, tenant workspace "
        "and exact row expiry without claiming a password already changed. A school's staff may "
        "ask by staff ID; the link goes to the address on file. Used or expired rows are not "
        "swept.")),
    ("Authentication and account audit events", "FR-017, FR-025", (
        "Implemented with limits. Seventeen identity events reach the shared tenant-scoped "
        "trail behind platform.audit.view, sixteen under typed actions and the card-key rotation "
        "as a custom one; the trail is their only record. Correcting a pending account's "
        "address reissues its invitation.")),
]

M03_TRACE_NEW = [
    ["School staff sign-in and password reset by staff ID", "FR-024, FR-006",
     "Implemented. A staff number signs its owner in, or issues a reset to the address on "
     "file, only inside the school the request names, and every refusal is a wrong password."],
]

M03_RECONCILIATION = [
    ("replace", "• MRD v2.86 lists Module 3",
     "• MRD v2.87 lists Module 3 as Backend Partial and In use Complete with twenty "
     "capabilities. Sign-in and password reset by a school's staff ID is a new entry, mapped "
     "to FR-024. The ID-card sign-in entry maps to FR-008 and the platform staff profile read "
     "as at an earlier date to FR-023."),
]

M03_SUMMARY = (
    "Minor revision. Adds FR-024: a school's staff sign in, and ask for a password reset, with "
    "the staff ID their school issued, resolved only inside the school the request names and "
    "refused exactly as a wrong password is. Adds FR-025: correcting a pending account's "
    "address reissues its invitation, so the link sent to the old address stops working. "
    "Records that the AuthEventLog table is dropped (vs_user 0015), its event names kept as "
    "vs_user.auth_events.AuthEvent, which resolves gap 9.6; gaps 9.7 and 9.8 are renumbered "
    "9.6 and 9.7. Corrects FR-006, FR-007 and section 1.1, which said REQUIRE_TENANT_ON_SIGN_IN "
    "was off: it has been on since 8f3b28b6 on 20 August 2026. " + TEST_EVIDENCE
)


def patch_m03() -> None:
    doc = Document(str(frd_path(M03_DIR, M03_STEM, M03_SOURCE)))
    set_cover_version(doc, M03_SOURCE, M03_TARGET)
    set_control(doc, "Version", M03_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")

    scope = table_with_header(doc, "Area")
    edit_value(scope, "Sign-in", "Platform staff may present the random key printed on their ID "
               "card in place of an address;",
               "A school's staff may present the staff ID their school issued in place of an "
               "address, matched only inside that school. Platform staff may present the random "
               "key printed on their ID card in place of an address;")
    edit_value(scope, "The tenant assertion", "It is off; see FR-006.",
               "It is on: a sign-in or reset that names no tenant is refused; see FR-006.")

    fr006 = fr_table(doc, "FR-006")
    set_status(doc, fr006, "Implemented")
    edit_value(fr006, "Current evidence", "PasswordService.request_reset goes through the same "
               "function.", "PasswordService.request_reset goes through the same function, and "
               "a staff number goes through resolve_staff_number_account, which is always "
               "scoped (FR-024).")
    set_value(fr006, "Current limit", (
        "None at the tenant boundary. REQUIRE_TENANT_ON_SIGN_IN has been True since 8f3b28b6 on "
        "20 August 2026: a sign-in or reset naming no tenant is refused as TENANT_REQUIRED with "
        "the wrong-password payload, the unscoped branch is unreachable, and FR-007's guard "
        "stands down. The school product sends the slug it reads off its subdomain and the "
        "Console sends codex. SignInTenantRequiredSwitchTests exercises this state."))

    fr007 = fr_table(doc, "FR-007")
    edit_value(fr007, "Current limit", "bulk_create and bulk_update go round save()",
               "The guard stands down while REQUIRE_TENANT_ON_SIGN_IN is True, which it has been "
               "since 20 August 2026, so one address may now be an account at more than one "
               "tenant; it applies again only if the switch is turned off. bulk_create and "
               "bulk_update go round save()")

    edit_value(fr_table(doc, "FR-013"), "Current evidence",
               "Every reset answers silently on an unknown, out-of-scope or ineligible account",
               "A school's staff may name themselves by staff ID instead of an address (FR-024), "
               "and the link still goes to the address on the account. Every reset answers "
               "silently on an unknown, out-of-scope or ineligible account")

    edit_value(fr_table(doc, "FR-017"), "Current limit",
               "The AuthEventLog table this module defines is written only by the dev-data "
               "seeder; production identity evidence lives in vs_audit.",
               "The vs_audit trail is the only record of identity events: the AuthEventLog table "
               "that shadowed it was dropped by vs_user migration 0015, and its event names are "
               "vs_user.auth_events.AuthEvent.")
    edit_value(fr_table(doc, "FR-017"), "Current limit", "See gaps 9.5 and 9.6.", "See gap 9.5.")

    add_fr(doc, "FR-024  Sign In, or Ask for a Reset, With the Staff ID a School Issued",
           "FR-024 | Implemented", M03_FR024)
    add_fr(doc, "FR-025  Reissue a Pending Invitation When Its Address Is Corrected",
           "FR-025 | Implemented", M03_FR025)
    set_status(doc, fr_table(doc, "FR-024"), "Implemented")
    set_status(doc, fr_table(doc, "FR-025"), "Implemented")

    edit_value(fr_table(doc, "FR-004"), "Current limit",
               "The second tenant's copy cannot actually be created yet - FR-007 refuses it "
               "while the sign-in switch is off - so the constraint is in place ahead of the "
               "behaviour it enables.",
               "None. REQUIRE_TENANT_ON_SIGN_IN is on, so FR-007's guard stands down and one "
               "address may be an account at two tenants; this constraint is what keeps it to "
               "one account per tenant.")
    edit_value(fr007, "Requirement",
               "The database now permits ada@gmail.com at Bright Star and at Greenfield, and "
               "sign-in cannot yet tell them apart. The pair must therefore be refused at "
               "creation for exactly as long as that is true, and the refusal must lift by "
               "itself when it stops being true.",
               "The database permits ada@gmail.com at Bright Star and at Greenfield. For as "
               "long as sign-in cannot tell the two apart, the pair must be refused at "
               "creation, and the refusal must lift by itself when sign-in can.")

    model = table_with_header(doc, "Model")
    old = row_starting(model, "AuthEventLog")
    keep_format(old.cells[0], "AuthEvent")
    keep_format(old.cells[1], "The names of identity events. Not a table.")
    keep_format(old.cells[2], (
        "vs_user.auth_events.AuthEvent, seventeen TextChoices. log_auth_event records each "
        "event as an identity AuditEvent with the name in metadata['auth_event'], and every "
        "reader, the auth-event endpoint and the staff history among them, queries those rows. "
        "The AuthEventLog table that once held the enum was written only by the dev seeder and "
        "is dropped by migration 0015."))

    routes = table_with_header(doc, "Method and path")
    edit_value(routes, "POST /auth/login/",
               "Sign in with an address, or with an ID-card key for platform staff, and the "
               "password: exactly one of the two, or 400 on credentials. Optional tenant slug "
               "in the body, not read on the card path.",
               "Sign in with an address, a school staff number, or an ID-card key for platform "
               "staff, and the password: exactly one of the three, or 400 on credentials. "
               "identifier carries an address or a staff number, told apart by the @. The "
               "tenant slug in the body is required, and not read on the card path.")
    edit_value(routes, "POST /auth/password/reset/request/",
               "Self-service reset. Optional tenant slug in the body.",
               "Self-service reset, by address or, through identifier, a school staff number; "
               "the link goes to the address on the account. The tenant slug in the body is "
               "required.")
    account_routes = [t for t in doc.tables
                      if t.rows[0].cells[0].text.strip() == "Method and path"][1]
    edit_value(account_routes, "PATCH /{id}/email/change/",
               "Change an address and end every session.",
               "Change an address and end every session; a pending account's invitation is "
               "reissued to the new address (FR-025).")

    answers = table_with_header(doc, "Condition")
    keep_format(row(answers, "A sign-in naming both an address and a card key, or neither").cells[0],
                "A sign-in naming more than one of an address, a staff number and a card key, "
                "or none")
    append_rows(answers, [[
        "A staff number the named school does not hold, a staff number with no tenant, or two "
        "numbers at that school differing only in case",
        "401 with the identical INVALID_CREDENTIALS payload. A reset request answers 200 with "
        "the same sentence as a real one.",
    ]])

    edit_paragraph(doc, "vs_user ran 381 tests OK",
                   "; the suite was not run again for this revision.", ".")
    edit_paragraph(doc, "vs_user ran 381 tests OK",
                   paragraph_starting(doc, "vs_user ran 381 tests OK").text.split(".")[0],
                   "vs_user ran 417 tests OK and schools.vs_staff 316 OK at c9031ad7, each suite "
                   "run on its own for this revision")
    edit_paragraph(doc, "vs_user ran 417 tests OK", " Deployment is not inferred from these checks.",
                   " For this revision the school client's typecheck and its auth and staff "
                   "tests passed, and its sign-in, forgot-password and Change email address "
                   "screens were driven against the running backend with a staff ID and an "
                   "address. Deployment is not inferred from these checks.")
    edit_paragraph(doc, "•  Permission-denied and cross-tenant reads on the auth-event endpoint",
                   "AuthEventLogTenantIsolationTests", "AuthEventTenantIsolationTests")
    insert_after(doc, "•  Permission-denied and cross-tenant reads on the auth-event endpoint",
                 "•  Staff ID sign-in and reset at the naming school only, case-insensitive, "
                 "ambiguous numbers refused, and a corrected invitation killing its old link. "
                 "Covered by StaffIdSignInTests, StaffIdPasswordResetTests and "
                 "StaffEmailCorrectionTests.")

    attention = table_with_header(doc, "Pri.")
    remove_row_containing(attention, "9.6 AuthEventLog is a table nothing writes")
    for gap, old_text, new_text in (
        ("9.5 An actor that is not a User row", "AuthEventLog.actor and .subject, and vs_audit's "
         "actor_user, are foreign keys", "vs_audit's actor_user is a foreign key"),
        ("9.8 Two response shapes and one dead serializer", " AuthEventLogReadSerializer has no "
         "caller at all.", ""),
    ):
        target = row_with_cell_starting(attention, 1, gap)
        edit_cell(target.cells[2], old_text, new_text)
    renumber(attention, "9.7 is_staff", "9.6 is_staff")
    renumber(attention, "9.8 Two response shapes and one dead serializer on the identity "
             "surface", "9.7 Three spellings of one fact on the identity surface")
    edit_paragraph(doc, "Three structural questions remain explicit.",
                   "which gap 9.7 records", "which gap 9.6 records")

    trace = next(t for t in doc.tables
                 if t.rows[0].cells[0].text.strip().startswith("MRD Module"))
    for capability, coverage, state in M03_TRACE_EDITS:
        target = row(trace, capability)
        keep_format(target.cells[1], coverage)
        keep_format(target.cells[2], state)
    append_rows(trace, M03_TRACE_NEW)
    edit_paragraph(doc, "MRD v2.85 records Module 3", "MRD v2.85", f"MRD v{MRD_TARGET}")
    edit_paragraph(doc, f"MRD v{MRD_TARGET} records Module 3", "with eighteen capability "
                   "entries", "with twenty capability entries")
    update_reconciliation(doc, M03_RECONCILIATION)

    log_change(doc, M03_TARGET, M03_SUMMARY)
    assert_absent(doc, "It is off", "AuthEventLogTenantIsolationTests", "not run again",
                  "cannot actually be created yet", "cannot yet tell them apart",
                  "AuthEventLogReadSerializer", "9.8 Two response", "See gap 9.8",
                  "the unused AuthEventLog table remains")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M03_DIR, M03_STEM, M03_TARGET),
           f"{M03_STEM.replace('_', ' ')} v{M03_TARGET}", M03_TARGET)


def row_with_cell_starting(table, column: int, start: str):
    hits = [r for r in table.rows if r.cells[column].text.strip().startswith(start)]
    if len(hits) != 1:
        raise ValueError(f"{start!r} starts column {column} of {len(hits)} rows")
    return hits[0]


def remove_row_containing(table, start: str) -> None:
    target = row_with_cell_starting(table, 1, start)
    target._tr.getparent().remove(target._tr)


def renumber(table, start: str, new_start: str) -> None:
    target = row_with_cell_starting(table, 1, start)
    cell = target.cells[1]
    keep_format(cell, new_start + cell.text.strip()[len(start):])


# ── M05 Audit & Activity Logging ─────────────────────────────────────────────

M05_DIR = "05-audit-and-activity-logging"
M05_STEM = "XVS_M05_Audit_and_Activity_Logging_Functional_Requirements_Document"
M05_SOURCE, M05_TARGET = "1.5", "1.6"

M05_FR018 = [
    ("Requirement", (
        "A school's staff profile shows one person's account events beside their employment "
        "events. No school role holds platform.audit.view, and none should, so that slice must "
        "be readable under the staff module's own key, and it must carry no more of the event "
        "than the history needs."
    )),
    ("Current evidence", (
        "schools.vs_staff.views.lifecycle._account_events reads AuditEvent rows filtered to "
        "the view's tenant, entity_type User, entity_id the person's account, and "
        "metadata.auth_event in a fixed list of account events (sign-ins are left out). It runs "
        "behind school.teachers.view and the staff record's own tenant and branch scoping, and "
        "an as_at read cuts the rows at the end of that day. Each entry carries the event, its "
        "label, its time, the actor's id and display name, and, on an email change, the "
        "administrator's note. The addresses, IP address and agent in the metadata are not "
        "returned."
    )),
    ("Acceptance", (
        "StaffEmailCorrectionTests: an email change shows on the history with its note and "
        "actor, and neither address appears anywhere in the response; a pending account's "
        "reissued invitation shows as INVITATION_SENT; another school's account events never "
        "reach this history."
    )),
    ("Limit", (
        "The slice is one module's reader with its own copy of which events count, not a "
        "vs_audit surface; the event list lives in the staff view. A school reads its own "
        "person's account events this way whether or not it holds any audit key, which is the "
        "point, and the Event Explorer's own gate is unchanged."
    )),
]

M05_TRACE = [
    ["Tenant and platform visibility controls", "FR-018",
     "Implemented. A school staff record's history reads that one person's account events "
     "under school.teachers.view, inside the tenant, returning the event, time, actor and an "
     "email change's note and no other metadata."],
]

M05_RECONCILIATION = [
    ("replace", "•  MRD v2.86 lists Module 5",
     "•  MRD v2.87 lists Module 5 as Backend Complete and In use Complete with fourteen "
     "capabilities and no material capability gap. This FRD does not dispute Complete for the "
     "module's stated backend scope: every listed capability has a backend path."),
]

M05_SUMMARY = (
    "Minor revision. Adds FR-018: a school staff record's history reads one person's account "
    "events from this trail under school.teachers.view, inside the tenant and a fixed event "
    "list, returning the event, time, actor and an email change's note and no other metadata. "
    "It read vs_user.AuthEventLog, which nothing outside the dev seeder wrote, so every "
    "school's account history was empty. Records that AuthEventLog is dropped (vs_user 0015), "
    "leaving this trail the only record of identity events. Section 2.2, the Module 3 "
    "emitter row and one traceability row are updated. " + TEST_EVIDENCE
)


def patch_m05() -> None:
    doc = Document(str(frd_path(M05_DIR, M05_STEM, M05_SOURCE)))
    set_cover_version(doc, M05_SOURCE, M05_TARGET)
    set_control(doc, "Version", M05_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")

    edit_paragraph(doc, "Reading the trail is governed by the key and bounded by the row.",
                   "The view and export keys are deliberately tenant-holdable,",
                   "One narrow reader sits outside those keys: a school staff record's history "
                   "reads that person's account events under school.teachers.view, returning "
                   "the event, its time, the actor and an email change's note and nothing else "
                   "(FR-018). The view and export keys are deliberately tenant-holdable,")

    emitters = next(t for t in doc.tables if any(
        r.cells[0].text.strip().startswith("M3 Identity & Auth") for r in t.rows))
    edit_cell(row_starting(emitters, "M3 Identity & Auth").cells[1], "No second log.",
              "No second log: the AuthEventLog table that once shadowed it is dropped by "
              "vs_user migration 0015, and the event names are vs_user.auth_events.AuthEvent. "
              "An email change carries the administrator's note when one is given.")

    add_fr(doc, "FR-018  Let a Staff Record's History Read One Person's Account Events Under "
           "Its Own Key", "FR-018 | Implemented", M05_FR018)
    set_status(doc, fr_table(doc, "FR-018"), "Implemented")
    append_rows(next(t for t in doc.tables
                     if t.rows[0].cells[0].text.strip().startswith("MRD Module")), M05_TRACE)
    edit_paragraph(doc, "MRD v2.76 records Module 5",
                   "MRD v2.76 records Module 5 as Phase V1, Backend Complete, In use Complete, "
                   "with thirteen capability entries",
                   f"MRD v{MRD_TARGET} records Module 5 as Phase V1, Backend Complete, In use "
                   "Complete, with fourteen capability entries")
    update_reconciliation(doc, M05_RECONCILIATION)
    log_change(doc, M05_TARGET, M05_SUMMARY)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M05_DIR, M05_STEM, M05_TARGET),
           f"{M05_STEM.replace('_', ' ')} v{M05_TARGET}", M05_TARGET)


# ── M12 Staff Management ─────────────────────────────────────────────────────

M12_DIR, M12_STEM = "12-staff-management", "XVS_M12_Staff_Management_Functional_Requirements_Document"
M12_SOURCE, M12_TARGET = "2.9", "2.10"

M12_SUMMARY = (
    "Minor revision. FR-012 corrected: the account half of the history is read from the "
    "identity events vs_user records in the audit trail, narrowed to this tenant, this person "
    "and a fixed list of account events under this module's own key. Versions 2.1 to 2.9 said "
    "it read vs_user.AuthEventLog, which nothing outside the dev seeder wrote, so every "
    "school's account history was empty; that table is now dropped. FR-009's email change asks "
    "Field Access on email, takes an optional note shown on the history, and reissues a "
    "pending invitation so a link sent to a mistyped address stops working. The staff number "
    "is recorded as a sign-in identifier (M03 FR-024). The duplicated traceability row for "
    "audit history is merged. " + TEST_EVIDENCE
)


def patch_m12() -> None:
    doc = Document(str(frd_path(M12_DIR, M12_STEM, M12_SOURCE)))
    set_control(doc, "Version", M12_TARGET)
    set_control(doc, "Date", REVIEW_DATE)
    set_control(doc, "Supersedes", f"v{M12_SOURCE}")
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")
    set_control(doc, "Verified against", CODE_BASELINE)

    fr009 = fr_table(doc, "FR-009")
    edit_value(fr009, "Business rules",
               "(4) An email change is refused for an address already in use at this school, "
               "and reveals nothing about any other school.",
               "(4) An email change is refused for an address already in use at this school, "
               "and reveals nothing about any other school. The address is the email field of "
               "school.teachers, so a role whose Field Access leaves it read-only is refused "
               "even with the account key. An optional note of up to 200 characters is recorded "
               "with the event and shown on the history. An account still waiting to be "
               "activated has its invitation reissued to the new address, so the link sent to "
               "the old one stops working (M03 FR-025).")
    edit_value(fr009, "Refusals", "422 with the field named for an email already in use here;",
               "422 with the field named for an email already in use here; 403 "
               "field_write_denied for an email the caller's Field Access may not write;")

    fr012 = fr_table(doc, "FR-012")
    edit_value(fr012, "Permission",
               "The account half is read from vs_user.AuthEventLog, which carries the same "
               "events, is scoped to this tenant and this subject, and is the identity layer’s "
               "own record of them.",
               "The account half is read from the identity events vs_user records in the audit "
               "trail (the IDENTITY module, the person's account as the entity, the event named "
               "in metadata.auth_event), narrowed in the view to this tenant, this person and a "
               "fixed list of account events, under this module's own key (M05 FR-018). Of an "
               "event's metadata only an email change's note is returned; the addresses, IP "
               "address and agent are not.")
    rules = row(fr012, "Business rules").cells[-1]
    if not rules.text.rstrip().endswith("where it belongs."):
        raise ValueError("FR-012's business rules no longer end with rule (6)")
    keep_format(rules, rules.text.rstrip() + (
        " (7) An email change shows the administrator's note, when one was given, and never "
        "either address: the email is a Field Access field and the history is read more widely "
        "than it is."))

    fields = table_with_header(doc, "Field")
    edit_cell(row_starting(fields, "staff_number").cells[-1],
              "because a school that does not number its staff must not be forced to invent "
              "one.",
              "because a school that does not number its staff must not be forced to invent "
              "one. It is also a sign-in identifier: the person may sign in or ask for a "
              "password reset with it, matched without regard to case and only inside their own "
              "school, and two numbers at one school differing only in case sign nobody in "
              "(M03 FR-024).")

    endpoints = next(t for t in doc.tables if any(
        r.cells[0].text.strip() == "PATCH /v1/i/me/staff/<id>/account/email/" for r in t.rows))
    keep_format(row(endpoints, "PATCH /v1/i/me/staff/<id>/account/email/").cells[-1],
                "FR-009. Optional note; Field Access on email; a pending invitation is reissued.")

    insert_after(doc, "The three account keys are enforced:",
                 "Correcting an invited person's email kills the link sent to the old address "
                 "and sends a fresh one; a role with email read-only under Field Access is "
                 "refused and nothing changes; the history shows the change, its note and the "
                 "actor, and neither address (StaffEmailCorrectionTests).")

    edit_paragraph(doc, "MRD v2.85 records Module 12", "MRD v2.85", f"MRD v{MRD_TARGET}")
    edit_paragraph(doc, f"MRD v{MRD_TARGET} records Module 12", "with ten capability entries",
                   "with eleven capability entries")
    trace = table_with_header(doc, "MRD capability")
    for cell in trace.rows[0].cells[2:]:
        keep_format(cell, "State at c9031ad7")
    audit_rows = [r for r in trace.rows
                  if r.cells[0].text.strip() == "Audit history and field-level access"]
    if len(audit_rows) != 2:
        raise ValueError(f"Expected two audit-history trace rows, found {len(audit_rows)}")
    keep_format(audit_rows[0].cells[1], "FR-012, FR-022, FR-023")
    keep_format(audit_rows[0].cells[2], (
        "Implemented. Every write emits a STAFF audit event, and the profile history reads "
        "employment events and the person's account events from the identity trail, with an "
        "email change's note; a reversed leave decision leaves only the engine's entry. The "
        "profile and every tab read as at an earlier day, and the whole record's fields are "
        "switchable per role."))
    audit_rows[1]._tr.getparent().remove(audit_rows[1]._tr)

    log_change(doc, M12_TARGET, M12_SUMMARY)
    assert_absent(doc, "vs_user.AuthEventLog, which carries the same events")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M12_DIR, M12_STEM, M12_TARGET),
           f"{M12_STEM.replace('_', ' ')} v{M12_TARGET}", M12_TARGET)


# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_CONTENTS_NOTE = "Staff ID sign-in and one identity trail"
MRD_INTRO = (
    "This revision records sign-in and password reset by a school's staff ID, the invitation "
    "reissued when a pending account's address is corrected, the staff history reading its "
    "account half from the identity trail, and the retirement of the AuthEventLog table that "
    "trail made redundant."
)
MRD_CHANGE_SUMMARY = (
    "Adds sign-in and self-service password reset by a school's staff ID, resolved only inside "
    "the school the request names and refused exactly as a wrong password is; the reset link "
    "still goes to the address on file. Correcting a pending account's address reissues its "
    "invitation, so a link sent to a mistyped address stops working, and the school staff "
    "email route asks Field Access and takes a note. The staff history's account half reads the "
    "identity trail; it read AuthEventLog, which nothing outside the dev seeder wrote, so it was "
    "empty for every school. AuthEventLog is dropped (vs_user 0015). One capability entry joins "
    "(M03 +1), 510 across 31 modules; statuses do not move. M03 v1.17, M05 v1.6 and M12 v2.10. "
    "Backend evidence and local school-app verification; no deployment claim."
)
MRD_DELTA_ROWS = [
    ["Staff ID sign-in", "Scoped to the naming school",
     "A staff number signs its owner in, or issues a reset to the address on file, only inside "
     "the school the request names; unknown, foreign and ambiguous numbers answer as a wrong "
     "password."],
    ["Email correction", "Kills the old invitation link",
     "Correcting a pending account's address rotates its invitation and sends a fresh one; the "
     "school route asks Field Access on email and records a note."],
    ["Staff history", "Account half from the identity trail",
     "The history reads one person's account events under school.teachers.view, with the "
     "note and without the addresses; it had read a table nothing wrote."],
    ["AuthEventLog", "Retired",
     "The unused table is dropped; its event names live on as vs_user.auth_events.AuthEvent."],
    ["Tenant switch", "Recorded as on",
     "M03 said REQUIRE_TENANT_ON_SIGN_IN was off; it has been on since 20 August 2026."],
    ["Module FRDs", "Three revised", "M03 v1.17, M05 v1.6, M12 v2.10."],
]
MRD_MODULES = {
    3: {"add": [("▸  Tenant-aware email and password login",
                 "▸  School staff sign-in and password reset by staff ID")], "count": 20,
        "blurb_edits": [(
            "every other key is refused with one identical answer.",
            "every other key is refused with one identical answer. A school's staff may sign "
            "in, or ask for a password reset, with the staff ID their school issued, matched "
            "only inside that school.",
        )]},
    12: {"count": 11, "blurb_edits": [(
        "Documented by M12 FRD v2.9.",
        "The history's account half reads the person's account events from the identity "
        "trail, and correcting a sign-in email reissues a pending invitation so a link sent to "
        "a mistyped address stops working. Documented by M12 FRD v2.10.",
    )]},
}


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    doc = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta, log = tables[76], tables[78]
    grids = {m: tables[i] for m, i in mrd_tools.CAPABILITIES.items()}
    assert delta.rows[0].cells[0].text.strip().endswith("capability delta")
    assert index.rows[0].cells[5].text.strip() == "Entries"

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for r in control.rows:
        label = r.cells[0].text.strip()
        if label == "Version":
            replace_cell(r.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(r.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(r.cells[1], CODE_BASELINE, size=9)
    for r in contents.rows:
        if r.cells[0].text.strip().startswith("5."):
            keep_format(r.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(r.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (text.startswith("5. ") and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    blurbs = mrd_tools.module_blurbs(doc)
    for module, changes in MRD_MODULES.items():
        grid = grids[module]
        for anchor, new in changes.get("add", []):
            mrd_tools.add_bullet(grid, anchor, new)
        for old, new in changes.get("blurb_edits", []):
            text = blurbs[module].text
            if text.count(old) != 1:
                raise ValueError(f"Module {module} description holds {old!r} {text.count(old)} times")
            mrd_tools.retitle(blurbs[module], text.replace(old, new))
        keep_format(mrd_tools.row_keyed(index, 0, str(module)).cells[5], str(changes["count"]))

    total = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    for r in control.rows:
        if r.cells[0].text.strip() == "Capability entries":
            replace_cell(r.cells[1], str(total), size=9)
    if total != 510:
        raise ValueError(f"Capability total is {total}, expected 510")

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)
    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if intro:
        mrd_tools.retitle(intro[0], MRD_INTRO)
    else:
        section = paragraph_starting(doc, f"5. v{MRD_TARGET} Capability Delta")
        following = Paragraph(section._p.getnext(), section._parent)
        mrd_tools.retitle(following, MRD_INTRO)

    mrd_tools.prepend_change_log(log, MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m03()
    patch_m05()
    patch_m12()
    patch_mrd()


if __name__ == "__main__":
    main()
