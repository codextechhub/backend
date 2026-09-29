#!/usr/bin/env python3
"""Cut M25 v1.2: a dashboard block needs the school's plan, and pending means one thing.

What changed in the backend, and therefore in the document:

* The finance and procurement dashboards decide each block through
  DashboardReader, whose for_user keeps a key only when the reader's role holds
  it and the school's plan reaches it, through vs_rbac.plan_gate's
  keys_within_plan (4af5d331). A block below the school's band is absent from
  the payload, not refused. The routes open on module access, which the plan
  gate at the door does not examine, so this is where the plan applies.
* vs_finance.constants.PENDING_STATUSES is the one definition of pending for
  the refunds and write-offs list and the receivables dashboard (63b3ce27), and
  the approvals waiting on you block counts credit and debit notes (d6260eec),
  held by a test to every registered finance and payments workflow type.

v1.1 still recorded the finance overview and the procurement dashboard as
gated whole on finance.report.view and procurement.analytics.view, with no
per-card permission. Both have opened to any key of their domain, block by
block, since their per-block readers were built, so FR-002, Section 2.2, the
permission matrix and the routes are brought to the code with the new rule.

The MRD was checked and left at v2.93: its Module 25 entries already name the
finance and procurement dashboards, and no capability is added.

    python tools/patch_plan_aware_dashboard_blocks_docs.py
"""
from __future__ import annotations

from docx import Document

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_procurement_dashboards_and_routes_docs import spread_cover_lines
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    assert_absent_outside_log,
    insert_row_after,
    row_labelled,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_paragraph,
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
    set_value,
)

import patch_record_history_docs

REVIEW_DATE = "28 September 2026"
MRD_VERSION = "2.93"
CODE_BASELINE = (
    "Backend main at 27f1da0f, 28 September 2026, holding the finance pass merged at "
    "32194627: plan-aware dashboard blocks (4af5d331) and one definition of pending "
    "(63b3ce27, d6260eec)"
)
TEST_EVIDENCE = (
    "Traced from the code at 27f1da0f; FinanceDashboardPlanTests, "
    "ProcurementDashboardPlanTests, KeysWithinPlanTests, AdjustmentPendingCountTests and "
    "ApprovalsWaitingOnYouTests cover the new rules and were not re-run for this revision. "
    "Backend evidence only; nothing here is deployed."
)

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE

M25_DIR = "25-dashboards-and-analytics"
M25_STEM = "XVS_M25_Dashboards_and_Operational_Analytics_Functional_Requirements_Document"
M25_SOURCE, M25_TARGET = "1.1", "1.2"

FR002_EVIDENCE = (
    "A dashboard route opens to anybody holding a key of its domain, and each block inside "
    "it is computed only for a reader holding the key of the list it summarises. The finance "
    "overview opens to any finance or payments key and its Receivables & collections and "
    "Cash, spend & compliance tabs to any finance key; the procurement overview and its Spend "
    "& suppliers and Stock & receiving tabs open to any procurement key. The audit summary "
    "requires platform.audit.view and the team approval-load view the workflow instance key. "
    "Ticket and task dashboards reuse their own modules' permission sets."
)
FR002_ACCEPTANCE = (
    "A reader holding no key of a domain is refused its dashboard; one holding some of its "
    "keys is served the blocks behind those keys and null for the rest. Granting a dashboard "
    "does not require inventing a key that exists only for it."
)
FR002_LIMIT = (
    "A block is the smallest unit: a reader holding its key sees all of it. Outside finance "
    "and procurement a dashboard is still gated whole on its domain's key."
)

FR011_HEADING = "FR-011  Leave Out a Block the School's Plan Does Not Reach"
FR011 = [
    ("Requirement",
     "A dashboard must not show a figure from a part of the product the school has not "
     "bought, and must not refuse the whole screen because one of its figures is beyond the "
     "plan."),
    ("Current evidence",
     "The finance and procurement dashboards decide every block through DashboardReader, "
     "built once per request by DashboardReader.for_user. It keeps a key only when the "
     "reader's role holds it and the school's plan reaches it, asking Module 4's plan gate "
     "through keys_within_plan, the bulk form of the door's plan_refusal (Module 6, Section "
     "5.3), so DashboardReader.can answers both questions for each block at a fixed handful "
     "of queries. A block whose key fails either question is null in the payload: absent, "
     "not refused, and never a zero. That covers the finance overview and its receivables and "
     "spend tabs, the procurement overview and its Spend & suppliers and Stock & receiving "
     "tabs, and the restock draft's stock check, which refuses 403 where stock is beyond the "
     "plan. With the plan gate switched off, at a school no plan was ever applied to, and for "
     "a platform super admin, every key the role holds counts. The dashboard routes open on "
     "module access, which the plan gate at the door does not examine, so the block decision "
     "is where the plan applies to them."),
    ("Acceptance",
     "A Procurement Admin holding procurement.analytics.view at a school whose plan stops "
     "short of analytics gets no spend blocks, the answer the analytics screen itself gives, "
     "and gets them once the plan reaches that band. ProcurementDashboardPlanTests and "
     "FinanceDashboardPlanTests prove a block absent below its band and present at it; "
     "KeysWithinPlanTests proves every key survives with the gate off or at an unprovisioned "
     "school, and that a school at Core loses only the key banded at Plus."),
    ("Current limit",
     "The payload does not say why a block is null, so a screen cannot tell a reader who "
     "lacks the key from a school below the band, and cannot show the upgrade message the "
     "door gives."),
]

FR012_HEADING = "FR-012  Count Pending Work the Way Its List Does"
FR012 = [
    ("Requirement",
     "A count on a dashboard and the list it summarises must use one definition, so the two "
     "cannot disagree about what is pending or what is waiting on the reader."),
    ("Current evidence",
     "vs_finance.constants.PENDING_STATUSES, draft and pending approval, is the one "
     "definition of pending for the refunds and write-offs list and the receivables "
     "dashboard: a posted, voided or cancelled refund or write-off is finished and is not "
     "pending. The approvals waiting on you block counts the finance workflow items whose "
     "current stage waits on the reader, grouped by type through APPROVAL_TYPES: write-offs, "
     "refunds, concessions, credit and debit notes, expense claims, journals and payouts. "
     "Credit and debit notes share one document type and one route, so they are counted and "
     "named together."),
    ("Acceptance",
     "AdjustmentPendingCountTests proves voided and cancelled documents are not pending and "
     "that the list and the dashboard count the same documents. ApprovalsWaitingOnYouTests "
     "proves a credit note waiting on the reader is counted with the rest, and holds "
     "APPROVAL_TYPES to every registered finance and payments workflow type, so a type added "
     "later cannot be left out of the count unnoticed."),
    ("Current limit", "None."),
]

M25_SUMMARY = (
    "Minor revision. A dashboard block needs the school's plan as well as the reader's key: "
    "the finance and procurement dashboards keep a key only when the role holds it and the "
    "school's plan reaches it, through the plan gate's keys_within_plan, so a block below "
    "the school's band is absent from the payload rather than refused (new FR-011). Pending "
    "means one thing: PENDING_STATUSES is the one definition for the refunds and write-offs "
    "list and the receivables dashboard, and the approvals waiting on you block counts credit "
    "and debit notes, held by a test to every registered finance and payments workflow type "
    "(new FR-012). FR-002, Section 2.2, the permission matrix, the processing steps, the "
    "routes and refusals are brought to the code: the finance and procurement dashboards "
    "open to any key of their domain and decide block by block, which v1.1 recorded as "
    "gated whole on finance.report.view and procurement.analytics.view with no per-card "
    f"permission. Traceability is reconciled to MRD v{MRD_VERSION} at twelve capabilities. "
    + TEST_EVIDENCE
)


def set_requirements(trace, capability: str, requirements: str) -> None:
    """Rewrite the requirements column of one traceability row, leaving its state."""
    keep_format(row(trace, capability).cells[1], requirements)


def patch_m25() -> None:
    require_newest(str(ROOT / "functional-requirements" / M25_DIR / f"{M25_STEM}_v*.docx"),
                   M25_SOURCE)
    doc = Document(str(frd_path(M25_DIR, M25_STEM, M25_SOURCE)))
    set_cover_version(doc, M25_SOURCE, M25_TARGET)
    spread_cover_lines(doc)
    if f"Version: {M25_TARGET}" not in doc.tables[0].rows[0].cells[0].text:
        raise ValueError("The cover does not carry the new version")
    set_control(doc, "Version", M25_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD",
                f"XVS Module Requirements Document v{MRD_VERSION} | Module 25, twelve "
                "capability entries")
    supporting = row_labelled(doc, "Supporting apps").cells[-1]
    keep_format(supporting, supporting.text.rstrip() + ", and vs_rbac for the plan gate")

    # 2.2 Where each dashboard lives
    surfaces = table_headed(doc, "Surface", "Served by", "Gated on")
    set_value(surfaces, "Finance overview and statements",
              "Opening: any finance or payments key. Each block: the key of what it "
              "summarises, within the school's plan (FR-011).")
    set_value(surfaces, "Procurement dashboard and insights",
              "Opening the overview and its tabs: any procurement key. Each block: the key of "
              "the list it summarises, within the school's plan (FR-011). The insights: their "
              "own report or analytics keys.")

    # 3. Actors
    reader = row_labelled(doc, "A domain's reader")
    keep_format(reader.cells[-1],
                "Any key of the domain opens its dashboard; each block needs the key that "
                "lists what it summarises, within the school's plan.")
    edit_paragraph(doc, "A dashboard key is never invented.",
                   "requires the right to see the thing summarised.",
                   "requires the right to see the thing summarised, at a school whose plan "
                   "reaches it.")

    # 4. Requirements
    fr002 = fr_table(doc, "FR-002")
    set_value(fr002, "Current evidence", FR002_EVIDENCE)
    set_value(fr002, "Acceptance", FR002_ACCEPTANCE)
    set_value(fr002, "Current limit", FR002_LIMIT)
    add_fr(doc, FR011_HEADING, "FR-011 | Implemented", FR011)
    add_fr(doc, FR012_HEADING, "FR-012 | Implemented", FR012)

    # 5. Producing one figure
    steps = table_headed(doc, "Step", "What happens", "Effect")
    keep_format(row(steps, "2").cells[1],
                "The domain's own permission is checked, and then each block's key against "
                "the reader's role and the school's plan.")
    keep_format(row(steps, "2").cells[2],
                "No dashboard-only key exists to be granted by mistake, and no block shows "
                "what the school has not bought.")

    # 7. Routes and refusals
    routes = table_headed(doc, "Method and path", "Purpose and permission")
    finance = row(routes, "GET /finance/reports/dashboard/")
    keep_format(finance.cells[-1],
                "Every finance overview block in one payload. Any finance or payments key; each "
                "block needs its own key within the school's plan (FR-011).")
    insert_row_after(finance, [
        "GET /finance/reports/dashboard/receivables/ and /spend/",
        "The Receivables & collections and Cash, spend & compliance tabs. Any finance key; "
        "blocks as the overview.",
    ])
    procurement = row(routes, "GET /procurement/reports/dashboard/")
    keep_format(procurement.cells[-1],
                "Procurement overview. Any procurement key; each block needs its own key within "
                "the school's plan (FR-011).")
    insert_row_after(procurement, [
        "GET /procurement/reports/dashboard/suppliers/ and /stock/",
        "The Spend & suppliers and Stock & receiving tabs, on the overview's terms.",
    ])
    refusals = table_headed(doc, "Condition", "Answer")
    anchor = row(refusals, "A reader without the domain's key")
    added = insert_row_after(anchor, [
        "A block whose key the reader's role does not hold",
        "Null in the payload; the rest of the dashboard is served.",
    ])
    insert_row_after(added, [
        "A block whose key the school's plan does not reach",
        "Null in the payload, exactly as a key the role lacks: absent, not refused, and never "
        "a zero.",
    ])

    # 8. Dependencies and evidence
    rbac = row_labelled(doc, "vs_rbac, vs_tenants").cells[-1]
    keep_format(rbac, rbac.text.rstrip() + (
        " The plan gate's keys_within_plan decides which of a reader's keys the school's plan "
        "reaches, and so which blocks a finance or procurement dashboard computes."))
    evidence = paragraph_starting(doc, "This module has no suite of its own")
    mrd_tools.retitle(evidence, evidence.text.replace(
        " Backend evidence only; nothing here is deployed.",
        " The block rule is covered by FinanceDashboardPlanTests, ProcurementDashboardPlanTests "
        "and KeysWithinPlanTests, and the pending rule by AdjustmentPendingCountTests and "
        "ApprovalsWaitingOnYouTests. Backend evidence only; nothing here is deployed."))

    # 10. Traceability
    edit_paragraph(doc, "Module 25 of XVS Module Requirements Document", "v2.72",
                   f"v{MRD_VERSION}")
    trace = table_headed(doc, "MRD capability", "FRD requirement", "State")
    set_requirements(trace, "Finance dashboard and statements", "FR-002, FR-004, FR-011, FR-012")
    set_requirements(trace, "Accounts-receivable summaries", "FR-004, FR-005, FR-012")
    set_requirements(trace, "Procurement summaries and insights", "FR-002, FR-004, FR-011")

    log_change(doc, M25_TARGET, M25_SUMMARY)
    retired_site_word = "cam" + "pus"
    assert_absent_outside_log(doc, retired_site_word, retired_site_word.capitalize(),
                              "There is no per-card permission", "Document v2.72",
                              "procurement dashboard procurement.analytics.view")
    for table in doc.tables[1:]:
        mrd_tools.keep_rows_whole(table)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M25_DIR, M25_STEM, M25_TARGET),
           f"{M25_STEM.replace('_', ' ')} v{M25_TARGET}", M25_TARGET)


def main() -> None:
    patch_m25()


if __name__ == "__main__":
    main()
