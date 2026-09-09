#!/usr/bin/env python3
"""Version the Module 18 FRD for the payout recovery sweep's selection rule.

The sweep looked for a batch whose own status was DRAFT. That status is an
aggregate over the batch's instructions, so one confirmed child turned the batch
PROCESSING and dropped it out of every later sweep while its unsent siblings sat
PENDING. FR-021's evidence, acceptance and limit are revised to the rule the code
now applies, section 5.2 records that PROCESSING does not mean everything was
sent, and section 9's first gap stops saying a delayed payout waits in draft.

Traceability is reconciled to MRD v2.74: the lead line named 22 entries in
v2.54 while the table lists 23, and one capability is renamed to the tracker's
own wording. The MRD needs no new version, because Module 18's status, capability
count and control box all stay true.

    python tools/patch_payout_sweep_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)

REVIEW_DATE = "9 September 2026"
SHORT_DATE = "9 Sep 2026"

FRD_SOURCE_VERSION = "1.8"
FRD_TARGET_VERSION = "1.9"
MRD_VERSION = "2.74"
CODE_BASELINE = "Commit 6451925, 9 September 2026"

FRD_DIR = "18-payments-and-collections"
FRD_STEM = "XVS_M18_Payments_and_Collections_Functional_Requirements_Document"

#: Index of the FR-021 requirement table, and of the row each field sits on.
FR021_TABLE = 22
EVIDENCE_ROW, ACCEPTANCE_ROW, LIMIT_ROW = 2, 3, 4

#: Index of section 9's two boxes and of section 10's traceability table.
CONTROL_BOX_TABLE, GAPS_TABLE = 37, 38
TRACEABILITY_TABLE, CHANGE_LOG_TABLE = 39, 40

FR021_EVIDENCE = (
    "The workflow handler records the approval and registers the dispatch with "
    "transaction.on_commit, so the provider is called from a worker once the "
    "approval is durable. _assert_no_open_transaction refuses any provider call "
    "made with a database transaction open, at the single point every dispatch "
    "passes through. Each instruction moves PENDING to PROCESSING in its own "
    "committed transaction before its transfer is requested, and only a clean "
    "provider rejection marks it FAILED. dispatch_undispatched_payout_batches "
    "re-dispatches approved batches whose hand-off never ran. It selects on the "
    "instructions rather than on the batch: any approved batch still holding a "
    "PENDING instruction is swept, whether the batch reads DRAFT or PROCESSING, "
    "because the batch status is an aggregate over its children and one child "
    "reaching the provider turns it PROCESSING while its siblings sit unsent. "
    "The batch status decides only whether money may still move at all, through "
    "PAYOUT_BATCH_DISPATCHABLE, the named complement of the terminal states."
)

FR021_ACCEPTANCE = (
    "A rollback during approval leaves nothing sent and the batch unapproved. A "
    "dispatch attempted with a transaction open raises before the provider is "
    "called. An instruction already in flight is never sent again. A timeout or "
    "an unreachable provider leaves the payout in flight rather than failed. The "
    "sweep dispatches an approved batch whose enqueue was lost, and refuses one "
    "that was never approved. A batch whose first beneficiary has already been "
    "paid is still swept, and only its unsent instructions reach the provider. A "
    "batch whose status is terminal is left alone, whatever state a child row is "
    "in."
)

FR021_LIMIT = (
    "Recovery depends on a worker consuming the enqueue, or on beat running the "
    "ten-minute sweep. Where neither runs, an approved payout is delayed rather "
    "than lost, and nothing announces the delay. A batch that was sent halfway "
    "through reads as PROCESSING, which is also how a batch whose transfers are "
    "all genuinely in flight reads, so the two cannot be told apart on screen "
    "until the sweep finishes the first."
)

LIFECYCLE_PARAGRAPH = (
    "A payout instruction is created pending inside a batch and cannot reach "
    "dispatch until the batch has its required human approvals. The batch "
    "remains DRAFT during review. The approval commits first and the transfer is "
    "sent afterwards, outside that transaction, one instruction at a time; each "
    "instruction is claimed as PROCESSING before its own send, so a run that "
    "dies mid-flight leaves nothing for a retry to pay twice. The batch moves to "
    "PROCESSING once at least one instruction is in flight and derives its "
    "terminal state from its children. PROCESSING therefore says that some money "
    "has moved, not that all of it has, and a batch in it may still hold "
    "instructions nobody has sent, which is why recovery reads the instructions "
    "rather than the batch. Only a clean provider rejection marks an instruction "
    "FAILED; an uncertain failure stays in flight for the webhook or "
    "verification to settle. Confirmation remains idempotent and books the "
    "vendor payment once."
)

CONTROL_BOX = (
    "CURRENT PAYOUT CONTROL - FAIL CLOSED\n"
    "• Approval is not optional by route. Every single or bulk payout "
    "enters Module 7 and only the terminal callback may ask the provider to "
    "transfer.\n"
    "• New books receive tenant policy, whose stages name approver groups "
    "created empty. The shared platform row carries no steps, so an entity that "
    "has published no ladder of its own is refused rather than approved, and "
    "payouts forbid continuing without approval outright.\n"
    "• payout_approval_health names any active entity that still cannot "
    "resolve a standard ladder and fails the release check. Missing approver "
    "appointments are a separate parked-state health concern.\n"
    "• The default threshold is N500,000: one checker below it, and a "
    "distinct senior checker at or above it.\n"
    "• Approval and the transfer are separate transactions. The approving "
    "vote commits before the provider is called, so a rollback during approval "
    "cannot leave money sent, and a guard refuses any provider call made with a "
    "transaction open.\n"
    "• An instruction is claimed before its own send and only a clean "
    "provider rejection marks it failed, so no retry re-sends it and no timeout "
    "records moved money as a failure.\n"
    "• The recovery sweep finishes any approved batch that still holds an "
    "unsent instruction, whether or not some of its transfers have already gone, "
    "and moves no money against a batch whose status is terminal."
)

WORKER_GAP = (
    "• Webhook processing and post-approval payout dispatch both depend on "
    "a worker. With none running, events are stored and acknowledged but never "
    "booked, and an approved payout waits for the ten-minute sweep, which needs "
    "beat. A batch that was sent halfway through waits in processing rather than "
    "in draft, so the delay reaches no screen as an unsent batch."
)

TRACEABILITY_LEAD = (
    f"Module 18 carries 23 capability entries in MRD v{MRD_VERSION}. Each maps "
    "to the requirements below."
)

#: The tracker's own wording for the capability on row 9, which this table quotes.
PAYOUT_LIFECYCLE_CAPABILITY = "Payout creation and lifecycle"

CHANGE_SUMMARY = (
    "Corrected the recovery sweep's selection so a batch that was sent halfway "
    "through is still finished. The sweep looked for a batch whose own status "
    "was DRAFT, but that status is an aggregate over the batch's instructions "
    "and one confirmed child turns it PROCESSING, so a run that paid the first "
    "beneficiary and died before the second dropped out of every later sweep and "
    "left the rest unpaid for good. Selection now reads the instructions, any "
    "approved batch still holding a PENDING one, and the batch status decides "
    "only whether money may move at all, as the named complement of the terminal "
    "states. FR-021's evidence, acceptance and limit are revised, the fail-closed "
    "control box records what the sweep will and will not touch, section 5.2 "
    "records that PROCESSING does not mean everything was sent, and section 9's "
    "first gap stops saying a delayed payout waits in draft. Traceability is "
    f"reconciled to MRD v{MRD_VERSION}: the lead line named 22 entries in v2.54 "
    "while the table lists 23, and one capability takes the tracker's wording. "
    "No capability was added and no status changed, so the MRD needs no new "
    "version."
)


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_control_value(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            replace_cell(row.cells[1], value, size=9)
            return
    raise ValueError(f"Control row not found: {label}")


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)
                return
    raise ValueError(f"Cover version not found: {source}")


def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError("Paragraph carries no run to inherit formatting from")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_body_paragraph(doc, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_run_text(paragraph, text)
            return
    raise ValueError(f"Body paragraph not found: {prefix}")


def replace_boxed_line(cell, prefix: str, text: str) -> None:
    """Rewrite one bullet of a callout box, leaving its siblings alone."""
    for paragraph in cell.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_run_text(paragraph, text)
            return
    raise ValueError(f"Callout line not found: {prefix}")


def prepend_change_log(table, version: str, summary: str) -> None:
    template = table.rows[1]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], SHORT_DATE, size=8)
    replace_cell(row.cells[2], summary, size=8)


def patch_frd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M18 Payments and Collections Functional Requirements "
        f"Document v{FRD_TARGET_VERSION}"
    )
    doc.core_properties.title = title
    doc.core_properties.version = FRD_TARGET_VERSION

    replace_cover_version(doc.tables[0], FRD_SOURCE_VERSION, FRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Version", FRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Review date", REVIEW_DATE)
    replace_control_value(doc.tables[1], "Code baseline", CODE_BASELINE)
    replace_control_value(
        doc.tables[1],
        "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 18",
    )

    fr021 = doc.tables[FR021_TABLE]
    if not fr021.rows[0].cells[0].text.strip().startswith("FR-021"):
        raise ValueError("Table 22 is not FR-021")
    replace_cell(fr021.rows[EVIDENCE_ROW].cells[1], FR021_EVIDENCE, size=8.5)
    replace_cell(fr021.rows[ACCEPTANCE_ROW].cells[1], FR021_ACCEPTANCE, size=8.5)
    replace_cell(fr021.rows[LIMIT_ROW].cells[1], FR021_LIMIT, size=8.5)

    replace_body_paragraph(
        doc, "A payout instruction is created pending", LIFECYCLE_PARAGRAPH,
    )

    control_box = doc.tables[CONTROL_BOX_TABLE].rows[0].cells[0]
    set_run_text(control_box.paragraphs[0], CONTROL_BOX)
    replace_boxed_line(
        doc.tables[GAPS_TABLE].rows[0].cells[0],
        "• Webhook processing",
        WORKER_GAP,
    )

    replace_body_paragraph(doc, "Module 18 carries", TRACEABILITY_LEAD)
    traceability = doc.tables[TRACEABILITY_TABLE]
    for row in traceability.rows:
        if row.cells[0].text.strip().startswith("Payout creation"):
            replace_cell(row.cells[0], PAYOUT_LIFECYCLE_CAPABILITY, size=8)
            break
    else:
        raise ValueError("Payout creation capability row not found")

    prepend_change_log(
        doc.tables[CHANGE_LOG_TABLE], FRD_TARGET_VERSION, CHANGE_SUMMARY,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    folder = Path(args.root) / "functional-requirements" / FRD_DIR
    patch_frd(
        folder / f"{FRD_STEM}_v{FRD_SOURCE_VERSION}.docx",
        folder / f"{FRD_STEM}_v{FRD_TARGET_VERSION}.docx",
    )
    print(f"Wrote M18 FRD v{FRD_TARGET_VERSION}")


if __name__ == "__main__":
    main()
