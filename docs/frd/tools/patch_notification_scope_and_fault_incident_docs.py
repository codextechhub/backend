#!/usr/bin/env python3
"""Cut M08 v1.13 and M30 v1.3: a branch's own notification emails, and fault incidents.

What changed in the backend, and therefore in the documents:

* M08 (98e98798). A branch switches its own notification emails.
  NotificationSetting carries an optional branch (migration 0019) and
  NotificationEventType a branch_scoped flag, true for the six billing and four
  workflow events whose every sender names the branch the event is about.
  Resolution is branch, then tenant, then platform, then the event default. The
  settings routes take ?branch=, rows gain branch_scoped and can_edit, and three
  refusals are added. Which emails fire is the school's; the wording of every
  message is the platform's (the template key is platform-scoped).
* M08 (570410e1). Every settings change and every template create or edit
  writes a CONFIG_CHANGED audit record through one service.
* M08, checked and recorded: the school-wide Approval emails switch is Module 7's
  (workflow.notifications.enabled, 78e0ceb7), so M08 names it as out of scope.
* M30 (fa23e06d). A configuration fault found outside the alert engine opens
  one incident, deduplicated on Incident.fault_key (migration 0004). The school
  finance layer reports a school with two sets of books through it. Such an
  incident has no alert or rule behind it, sends no notification and does not
  resolve itself. The fired-alarm recipient is corrected to platform.health.update
  holders wherever v1.2 still read "create or update".

The MRD was not revised here; both documents cite v2.93, whose Module 8 and
Module 30 entries still carry 22 and 13 capabilities and which v2.94 leaves as
they were. M08's background delivery row takes the MRD's label for that entry,
background delivery, retries, and transactional operational alerts, and maps
FR-004, the must-send transactional events, which no row traced.

    python tools/patch_notification_scope_and_fault_incident_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    insert_row_after,
    row_labelled,
    table_headed,
)
from patch_staff_id_and_auth_events_docs import (
    edit_cell,
    edit_paragraph,
    edit_value,
    fr_table,
    insert_after,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
    set_status,
)

import patch_record_history_docs

REVIEW_DATE = "27 September 2026"
SOURCE_MRD = "XVS Module Requirements Document v2.93"

patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


def box_table(doc, heading: str):
    """The one-row boxed note whose text starts with ``heading``."""
    hits = [t for t in doc.tables
            if len(t.rows) == 1 and t.rows[0].cells[0].text.strip().startswith(heading)]
    if len(hits) != 1:
        raise ValueError(f"{heading!r} heads {len(hits)} boxes")
    return hits[0]


def box_cell(doc, heading: str):
    return box_table(doc, heading).rows[0].cells[0]


def remove_box(doc, heading: str) -> None:
    table = box_table(doc, heading)
    table._tbl.getparent().remove(table._tbl)


def _box_paragraph(cell, start: str) -> Paragraph:
    hits = [p for p in cell.paragraphs if p.text.strip().startswith(start)]
    if len(hits) != 1:
        raise ValueError(f"{start!r} starts {len(hits)} box paragraphs")
    return hits[0]


def box_sub(cell, old: str, new: str) -> None:
    """Replace text inside the one run that holds it, leaving the box's layout alone."""
    runs = [r for p in cell.paragraphs for r in p.runs if old in r.text]
    if len(runs) != 1 or runs[0].text.count(old) != 1:
        raise ValueError(f"{old[:60]!r} is not in exactly one run of the box")
    runs[0].text = runs[0].text.replace(old, new)


def box_replace(cell, start: str, text: str) -> None:
    mrd_tools.retitle(_box_paragraph(cell, start), text)


def box_add_after(cell, start: str, text: str) -> None:
    """Add a line after the paragraph starting ``start``, formatted as that paragraph."""
    anchor = _box_paragraph(cell, start)
    clone = copy.deepcopy(anchor._p)
    anchor._p.addnext(clone)
    mrd_tools.retitle(Paragraph(clone, anchor._parent), text)


def box_append(cell, text: str) -> None:
    last = [p for p in cell.paragraphs if p.text.strip()][-1]
    clone = copy.deepcopy(last._p)
    last._p.addnext(clone)
    mrd_tools.retitle(Paragraph(clone, last._parent), text)


def keep_together(table) -> None:
    """Hold a short table on one page: no row splits, and each row is tied to the next.

    A requirement table left to break wherever the page ends can strand its last
    row alone at the top of the following page, away from the requirement it
    belongs to.
    """
    for r in table.rows:
        props = r._tr.get_or_add_trPr()
        if props.find(qn("w:cantSplit")) is None:
            props.append(OxmlElement("w:cantSplit"))
    for r in table.rows[:-1]:
        for cell in r.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def drop_empty_paragraphs(cell) -> None:
    for paragraph in list(cell.paragraphs):
        if not paragraph.text.strip() and len(cell.paragraphs) > 1:
            paragraph._p.getparent().remove(paragraph._p)


# ── M08 Notifications & Delivery ─────────────────────────────────────────────

M08_DIR = "08-notifications-and-delivery"
M08_STEM = "XVS_M08_Notifications_and_Delivery_Functional_Requirements_Document"
M08_SOURCE, M08_TARGET = "1.12", "1.13"
M08_BASELINE = (
    "Backend main at 7f3756e2, 27 September 2026: settings and template changes audited "
    "at 570410e1, a branch's own settings at 98e98798"
)

M08_FR016_HEADING = "FR-016 Let Each Branch Set Its Own Emails"
M08_FR016 = [
    ("Requirement",
     "One branch switching off an email must not switch it off for another. When the only "
     "row a branch administrator could write was the whole school's, Ikeja muting fee emails "
     "muted Lekki's parents as well."),
    ("Current evidence",
     "NotificationSetting carries an optional branch (migration 0019): one row per tenant, "
     "branch, event and channel, the whole-tenant row still unique on its own, and a check "
     "that a branch row names its tenant. NotificationEventType.branch_scoped is true for "
     "exactly the events whose every sender names the branch the event is about: "
     "billing.invoice_issued, billing.statement_issued, billing.debit_note_issued, "
     "billing.credit_note_issued, billing.payment_received and billing.invoice_overdue, and "
     "workflow.stage_activated, workflow.rejected, workflow.returned and "
     "workflow.final_approved. An invoice, receipt, note or overdue notice passes the branch "
     "it is filed under, a statement the customer's branch, and a workflow notice the "
     "instance's. A branch row is consulted only for a branch_scoped event and only for "
     "recipients owned by that branch's tenant, so a school's branch never silences platform "
     "staff. GET and PATCH /notify/settings/ take ?branch=<id>; each row carries source "
     "(branch, tenant, platform or default), branch_scoped and can_edit. At a branch, true or "
     "false writes the branch's own row and null removes it, so the branch follows the whole "
     "school again."),
    ("Acceptance",
     "The branch must belong to the asserted tenant and lie inside the caller's reach; an "
     "unknown, foreign or unreached branch, and any branch at the platform layer, answer the "
     "same 404. A caller whose reach is limited to some branches reads the whole school with "
     "can_edit false on every row, and a whole-school PATCH from them is refused 403 "
     "BRANCH_SCOPE_REQUIRED, even in a school with one branch. At a branch, an event that is "
     "not branch_scoped is refused per item with BRANCH_NOT_CONFIGURABLE; null without a "
     "branch is refused with RESET_NEEDS_BRANCH. One refused item writes nothing. Ikeja "
     "switching an email off silences Ikeja alone, a branch may keep an email the school "
     "switched off, and the matrix still costs two queries at a branch "
     "(tests_branch_settings, 35 tests)."),
    ("Current limit",
     "A branch may set only the ten branch-scoped events. An event that any sender raises "
     "without a branch stays a whole-school setting, because a branch choice it could not "
     "honour would be a switch that does nothing. The platform default layer has no "
     "branches."),
]

M08_FR017_HEADING = "FR-017 Record Who Changed a Setting or a Template"
M08_FR017 = [
    ("Requirement",
     "A bursar who stops receiving payment emails, or a school whose invoice email suddenly "
     "reads differently, must be able to find out who changed it and when."),
    ("Current evidence",
     "Both administrative write paths call one service, vs_notifications/services/audit.py, "
     "which writes CONFIG / CONFIG_CHANGED to the platform audit trail, the pair Module 6 "
     "uses for its own value changes. A settings write records one event per event and "
     "channel whose stored value at that layer changed, with the value before and after, "
     "the layer (platform, tenant or branch, the branch named in the entity id, summary and "
     "metadata), and whether the row was created or removed. A template create or edit "
     "records the columns that changed, before and after: subject, body, action label and "
     "address, markup, markup ownership and active state."),
    ("Acceptance",
     "Switching a channel off and back on writes two records. Re-sending a stored value, an "
     "edit that changes nothing, and a refused PATCH record nothing. A platform-layer change "
     "is recorded as the platform's, and a branch change and its reset name the branch "
     "(NotificationChangeAuditTests, BranchSettingAuditTests)."),
    ("Current limit",
     "Recording is best effort by the audit trail's own contract: a failure to record never "
     "undoes the change."),
]

M08_APPROVAL_SWITCH_ROW = [
    "Whether a school's approvals notify anybody at all",
    "Module 7, Workflow & Approval Engine. Its one whole-school switch, "
    "workflow.notifications.enabled in the Module 6 configuration catalogue, is read before "
    "an approval event is raised: off, the approval events are never raised, so nothing "
    "reaches this module on any channel, the in-app one included. The school app shows it in "
    "Settings > Notifications above this module's settings (school-fe 78cd90d), which is not "
    "proof of deployment.",
]

M08_SENDERS_ROW = [
    "Modules 7 and 17, branch senders",
    "Name the branch an event is about when they raise it: the workflow instance's branch, "
    "the branch an invoice, receipt, credit or debit note or overdue invoice is filed under, "
    "and the customer's branch for a statement. That branch's own settings decide the "
    "branch-scoped events (FR-016).",
]

M08_SUMMARY = (
    "Minor revision. A branch sets its own notification emails. NotificationSetting gains an "
    "optional branch (migration 0019) and NotificationEventType a branch_scoped flag, true "
    "for the six billing and four workflow events whose every sender names the branch the "
    "event is about. Resolution is branch, then school, then platform, then the event "
    "default, and a branch speaks only for recipients its own school owns. The settings "
    "routes take ?branch=, rows carry source, branch_scoped and can_edit, and three refusals "
    "are added: BRANCH_SCOPE_REQUIRED, BRANCH_NOT_CONFIGURABLE and RESET_NEEDS_BRANCH. A "
    "branch administrator changes only their own branches' settings. Which emails fire is "
    "the school's; the wording of every message is the platform's. Every settings change "
    "and every template create or edit writes a CONFIG_CHANGED audit record. Adds FR-016 and "
    "FR-017 and updates scope, actors, FR-002, the dispatch workflow, the data model, the "
    "API contracts, dependencies, Needs Attention and traceability, and records that the "
    "school-wide Approval emails switch belongs to Module 7. The two resolved boxes leave "
    "Needs Attention, which holds current gaps only. Module 8 stays at 22 capability "
    "entries; the background delivery entry takes the MRD's wording, background delivery, "
    "retries, and transactional operational alerts, and maps FR-004, the transactional "
    "events no entry traced. Verified against backend 570410e1 and 98e98798, where vs_notifications ran "
    "224 tests, vs_workflow 430, vs_finance 873 and vs_rbac 892, all passing; not rerun for "
    "this revision. Backend evidence only; nothing here claims deployment."
)


#: The entry as the MRD words it; FR-004 is its transactional half.
M08_BACKGROUND_ENTRY = "Background delivery, retries, and transactional operational alerts"


def patch_m08() -> None:
    require_newest(str(ROOT / "functional-requirements" / M08_DIR / f"{M08_STEM}_v*.docx"),
                   M08_SOURCE)
    doc = Document(str(frd_path(M08_DIR, M08_STEM, M08_SOURCE)))
    set_cover_version(doc, M08_SOURCE, M08_TARGET)
    set_control(doc, "Version", M08_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", M08_BASELINE)
    set_control(doc, "Source MRD", f"{SOURCE_MRD} | Module 8")
    set_control(doc, "Supporting apps",
                "vs_tenants, vs_rbac, vs_user, vs_config, vs_audit, vs_health, "
                "vs_procurement, vs_finance, vs_workflow, schools/vs_onboarding, core")

    # 1. Scope
    scope = table_headed(doc, "Area", "Responsibility")
    edit_value(scope, "Settings",
               "The effective per-channel matrix, its platform defaults, and a tenant's "
               "overrides.",
               "The effective per-channel matrix, its platform defaults, a tenant's overrides "
               "for all of its branches, and one branch's own overrides for the emails sent "
               "about a branch. Every settings and template change is audited.")
    out_of_scope = table_headed(doc, "Not owned here", "Owner")
    insert_row_after(row(out_of_scope, "Mail transport"), M08_APPROVAL_SWITCH_ROW)

    # 2. Context
    edit_paragraph(
        doc, "Reading a feed needs no permission",
        "because they are three different kinds of authority.",
        "because they are three different kinds of authority. Which emails fire belongs to "
        "the school, and below it to each branch for the emails sent about a branch; the "
        "wording of every message belongs to the platform. A caller's branch reach decides "
        "which of the school's scopes they may configure.")

    # 3. Actors
    actors = table_headed(doc, "Actor", "May do", "Governed by")
    school_admin = row(actors, "Tenant administrator")
    keep_format(school_admin.cells[0], "School administrator")
    keep_format(school_admin.cells[1],
                "Switch an email on or off per event for the whole school, and for any of "
                "its branches.")
    keep_format(school_admin.cells[2],
                "communication.communication_permissions.enforce, a tenant key held by "
                "default by school_admin and branch_admin, inside their own tenant, with reach over the "
                "whole school")
    insert_row_after(school_admin, [
        "Branch administrator",
        "Switch an email on or off for their own branches only, for the events sent about a "
        "branch, and read the whole school's choices without changing them.",
        "The same key, with reach limited to some branches. A whole-school change is refused "
        "with BRANCH_SCOPE_REQUIRED.",
    ])
    platform = row(actors, "Platform staff")
    keep_format(platform.cells[1],
                "Create and edit templates, and preview any of them. Set the platform "
                "default layer every tenant inherits.")
    keep_format(platform.cells[2],
                "communication.notification_templates.configure, a platform key no school "
                "role can hold, so the wording of every message is the platform's; "
                "communication.communication_permissions.enforce in the platform tenant for "
                "the default layer")

    # 4. Requirements
    fr002 = fr_table(doc, "FR-002")
    edit_value(fr002, "Current evidence",
               "Resolution layers a tenant row over a platform row over the event's own "
               "default,",
               "Resolution layers a branch row, for an event sent about a branch (FR-016), "
               "over a tenant row over a platform row over the event's own default,")
    edit_value(fr002, "Acceptance",
               "A caller may read and write only their own tenant's rows.",
               "A caller may read and write only their own tenant's rows, and at a branch "
               "only the branches their reach covers.")
    edit_value(fr002, "Current limit",
               "Preferences are per tenant, not per person.",
               "Preferences are per tenant or per branch, not per person.")
    add_fr(doc, M08_FR016_HEADING, "FR-016 | Implemented", M08_FR016)
    add_fr(doc, M08_FR017_HEADING, "FR-017 | Implemented", M08_FR017)
    keep_together(fr_table(doc, "FR-016"))
    keep_together(fr_table(doc, "FR-017"))

    # 5. Workflow
    dispatch = table_headed(doc, "Step", "What happens", "Effect")
    edit_cell(row(dispatch, "1").cells[1],
              "A module raises an event with its context.",
              "A module raises an event with its context and, where it has one, the branch "
              "the event is about.")
    edit_cell(row(dispatch, "2").cells[2],
              "Owner row, then platform row, then the event default, once per owning tenant.",
              "The branch row for a branch-scoped event, counted only inside that branch's "
              "own tenant, then the owner row, then the platform row, then the event default, "
              "once per owning tenant.")

    # 6. Data model
    models = table_headed(doc, "Model", "Holds", "Notes")
    append_to(models, "NotificationEventType",
              " branch_scoped is true when every sender names the branch the event is about, "
              "so a branch may set it for itself.")
    edit_value(models, "NotificationSetting",
               "One row per tenant per pair, plus a platform row with no tenant. Constrained "
               "so each layer holds at most one.",
               "One row per tenant per pair for all of its branches, one per branch per pair "
               "for a branch's own choice, and a platform row with no tenant. Constrained so "
               "each layer holds at most one, a branch row names its tenant, and the platform "
               "layer has no branches. Deleting a branch removes its own rows.")

    # 7. API
    edit_paragraph(doc, "Routes are mounted under /v1/notify/.",
                   "settings and history are scoped to the caller's tenant.",
                   "settings are scoped to the caller's tenant or one of its branches, and "
                   "history to the caller's tenant.")
    settings_route = row_labelled(doc, "GET /notify/settings/, PATCH /notify/settings/update/")
    keep_format(settings_route.cells[-1],
                "The effective matrix for the caller's tenant, or with ?branch=<id> for one "
                "of its branches, and overrides by event and channel. Rows carry source, "
                "branch_scoped and can_edit. At a branch, null removes the branch's own "
                "value. Atomic, and every change is audited (FR-017).")
    errors = table_headed(doc, "Condition or route", "Answer")
    anchor = row(errors, "Configuring a transactional event")
    for values in reversed([
        ["A ?branch= that is unknown, another tenant's or outside the caller's reach, or any "
         "branch at the platform layer", "404, identical in every case."],
        ["A whole-school settings change by a caller whose reach is limited to some branches",
         "403, BRANCH_SCOPE_REQUIRED."],
        ["Setting, at a branch, an event that is not sent about a branch",
         "BRANCH_NOT_CONFIGURABLE, per item."],
        ["Resetting a value to null without naming a branch", "RESET_NEEDS_BRANCH, per item."],
    ]):
        insert_row_after(anchor, values)

    # 8. Dependencies
    deps = table_headed(doc, "Dependency", "Contract")
    append_to(deps, "Module 4, Roles & Permissions",
              " The template key is platform-scoped; the settings and history keys are tenant "
              "keys held by school_admin and branch_admin, and the caller's branch reach "
              "decides which of the school's scopes they may configure.")
    append_to(deps, "Module 5, Audit",
              " Every settings and template change is written to it as CONFIG_CHANGED "
              "(FR-017).")
    insert_row_after(row(deps, "Modules 7, 10, 26 and 29"), M08_SENDERS_ROW)

    # 9. Needs Attention: current gaps only
    remove_box(doc, "RESOLVED - A MESSAGE NOW BELONGS")
    remove_box(doc, "RESOLVED - THE DESIGN IS NO LONGER")
    intro = paragraph_starting(doc, "These are current risks and gaps, not history.")
    spacer = intro._p.getnext()
    if spacer.tag.endswith("}p") and not Paragraph(spacer, intro._parent).text.strip():
        spacer.getparent().remove(spacer)
    gaps = box_cell(doc, "FURTHER GAPS")
    drop_empty_paragraphs(gaps)
    box_sub(gaps, "FURTHER GAPS", "CURRENT GAPS")
    box_sub(gaps, "Delivery preferences are per tenant.",
            "Delivery preferences are per tenant or per branch.")

    # 10. Traceability
    mrd_tools.retitle(
        paragraph_starting(doc, "Module 8 carries 22 capability entries in MRD v2.80."),
        "Module 8 carries 22 capability entries in MRD v2.93. Each maps to the requirements "
        "below. A branch's own settings strengthen the effective matrix and the per-school "
        "overrides, and the record of settings and template changes strengthens notification "
        "audit events and templates, without changing the count.")
    trace = table_headed(doc, "MRD capability", "Requirements")
    background = row(trace, "Background delivery tasks and retries")
    keep_format(background.cells[0], M08_BACKGROUND_ENTRY)
    for capability, frs in (
        (M08_BACKGROUND_ENTRY, "FR-004, FR-012, FR-015"),
        ("Effective delivery-preference matrix", "FR-002, FR-016"),
        ("Per-school delivery overrides", "FR-002, FR-016"),
        ("Notification audit events", "FR-012, FR-017"),
        ("Notification templates", "FR-005, FR-009, FR-017"),
    ):
        keep_format(row(trace, capability).cells[-1], frs)

    log_change(doc, M08_TARGET, M08_SUMMARY)
    assert_absent_outside_log(doc, "MRD v2.80", "Tenant administrator", "RESOLVED - ",
                              "Background delivery tasks and retries",
                              "FURTHER GAPS", "Preferences are per tenant, not")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M08_DIR, M08_STEM, M08_TARGET),
           f"{M08_STEM.replace('_', ' ')} v{M08_TARGET}", M08_TARGET)


# ── M30 System Health & Monitoring ───────────────────────────────────────────

M30_DIR = "30-system-health-and-monitoring"
M30_STEM = "XVS_M30_System_Health_and_Monitoring_Functional_Requirements_Document"
M30_SOURCE, M30_TARGET = "1.2", "1.3"
M30_BASELINE = (
    "Backend main at 7f3756e2, 27 September 2026; configuration-fault incidents at fa23e06d"
)

M30_FR016_HEADING = "FR-016  Open One Incident per Configuration Fault"
M30_FR016 = [
    ("Requirement",
     "A configuration fault found outside the alert engine, such as a school whose books "
     "cannot name one primary entity, must reach an operator who can fix it, as exactly one "
     "incident however many requests meet it, rather than a logged 500 seen only by a bursar "
     "who cannot act on it."),
    ("Current evidence",
     "vs_health.faults.report_configuration_fault opens an automatic incident, "
     "INVESTIGATING, at the caller's severity, owner label Configuration check and team "
     "Platform, with the summary as its opening timeline note, unless an unresolved incident "
     "already carries the same Incident.fault_key. The key names one subject's fault, not the "
     "class of fault, and is matched exactly; it is indexed and returned by no serializer, so "
     "renaming or re-triaging the incident cannot break the match. The school finance layer "
     "is the first reporter: AmbiguousPrimaryEntity, raised when a school holds more than one "
     "active tenant-kind ledger entity, is reported at SEV2 under "
     "fal.ambiguous-primary-entity.tenant-<id>, naming the school and what to do, and the "
     "refusal still answers 500. A fault incident is listed, counted as active on the "
     "overview and included in reliability like any other incident."),
    ("Acceptance",
     "A read that raises three times opens exactly one incident, and two broken schools open "
     "one each. A repeat writes neither an incident nor a timeline entry. Resolving lets the "
     "next occurrence open a fresh incident. An operator or alert-driven incident, whose key "
     "is blank, never satisfies a fault's match, and a blank key is refused. Provisioning "
     "reports outside its own transaction, so the incident survives the rollback. A failing "
     "reporter changes neither the error nor its message (tests_faults, 8 tests; "
     "test_ambiguous_entity_reporting, 13 tests)."),
    ("Current limit",
     "A fault incident has no alert or rule behind it and sends no notification: it reaches "
     "the Health console only, because no configuration-fault event is registered in Module "
     "8. It does not resolve itself when the arrangement is corrected, unlike an alert-driven "
     "incident; an operator resolves it. Two requests arriving together can each open one."),
]

M30_FAULT_STEPS = [
    ("State", "Condition", "Effect"),
    ("First report", "No unresolved incident carries the fault key.",
     "Open one automatic incident, INVESTIGATING, at the caller's severity, with the opening "
     "note on its timeline."),
    ("Repeat", "An unresolved incident carries the key.",
     "Open nothing and append nothing; the caller receives no incident."),
    ("Resolved", "An operator resolves the incident.",
     "The key stops matching, so the next occurrence opens a fresh incident."),
    ("Reporter failure", "Reporting raises inside the detecting code.",
     "The failure is logged and swallowed; the original refusal reaches the request "
     "unchanged."),
]

M30_FAL_ROW = [
    "Module 19 Finance & Accounting, school finance layer",
    "Reports a school holding more than one active tenant-kind ledger entity "
    "(AmbiguousPrimaryEntity) through report_configuration_fault, from the envelope decorator "
    "both raise sites share and outside the provisioning transaction, so the incident "
    "survives its rollback. One incident per school; the request still answers 500.",
]

M30_SUMMARY = (
    "Minor revision. Adds FR-016: a configuration fault found outside the alert engine opens "
    "one incident. The school finance layer reports a school holding two active sets of books "
    "(AmbiguousPrimaryEntity) through vs_health.faults, from the decorator both raise sites "
    "share and outside the provisioning transaction, so the incident survives its rollback. "
    "Incident.fault_key (migration 0004) deduplicates: one unresolved incident per fault, a "
    "repeat writes nothing, not even a timeline entry, and resolving is how the next "
    "occurrence opens a fresh one. Such an incident has no alert or rule behind it, sends no "
    "notification and does not resolve itself. Scope, context, actors, workflow, data model, "
    "dependencies, Needs Attention and traceability follow, and the fired-alarm recipient is "
    "corrected to platform.health.update holders wherever the text still read create or "
    "update. Verified against backend fa23e06d, where vs_health ran 45 tests, schools.core "
    "219 and core 161, all passing; not rerun for this revision. Backend evidence only; "
    "nothing here claims deployment."
)

_BY_METHOD = "platform.health.create or platform.health.update, according to the method"


def add_fault_workflow(doc) -> None:
    """Section 5.4, cloned from 5.3's heading and table so it carries their look."""
    heading = paragraph_starting(doc, "5.3 Sustained Alert and Delivery")
    template = heading._p.getnext()
    if not template.tag.endswith("}tbl"):
        raise ValueError("5.3 is not followed by its table")
    new_heading = copy.deepcopy(heading._p)
    new_table = copy.deepcopy(template)
    template.addnext(new_heading)
    new_heading.addnext(new_table)
    mrd_tools.retitle(Paragraph(new_heading, heading._parent), "5.4 Configuration Fault")
    table = next(t for t in doc.tables if t._tbl is new_table)
    while len(table.rows) > len(M30_FAULT_STEPS):
        table._tbl.remove(table.rows[-1]._tr)
    for r, values in zip(table.rows, M30_FAULT_STEPS):
        for cell, value in zip(r.cells, values):
            keep_format(cell, value)
    keep_together(table)


def patch_m30() -> None:
    require_newest(str(ROOT / "functional-requirements" / M30_DIR / f"{M30_STEM}_v*.docx"),
                   M30_SOURCE)
    doc = Document(str(frd_path(M30_DIR, M30_STEM, M30_SOURCE)))
    set_cover_version(doc, M30_SOURCE, M30_TARGET)
    set_control(doc, "Version", M30_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", M30_BASELINE)
    set_control(doc, "Source MRD", f"{SOURCE_MRD} | Module 30")
    set_control(doc, "Supporting apps",
                "core, vs_tenants, vs_rbac, vs_user, vs_notifications, vs_config; "
                "configuration faults reported by schools.core.fal")

    # 1. Scope
    insert_after(doc, "Sustained and service-scoped breach evaluation",
                 "Configuration-fault incidents that another module reports when a data "
                 "arrangement is wrong: one unresolved incident per fault, with no alert rule, "
                 "metric or threshold behind it.")
    out_of_scope = table_headed(doc, "Not owned here", "Owner")
    insert_row_after(row(out_of_scope, "External uptime provider guarantees"), [
        "Deciding that a configuration is wrong",
        "The module whose rule it is. The school finance layer reports a school holding two "
        "sets of books; Health records the incident and never reads school data.",
    ])

    # 2. Context
    edit_paragraph(doc, "The module has three evidence paths.",
                   "The module has three evidence paths.",
                   "The module has four evidence paths.")
    edit_paragraph(doc, "The module has four evidence paths.",
                   "rather than ownership isolation.",
                   "rather than ownership isolation. Other modules report configuration "
                   "faults, which open incidents directly (FR-016).")
    decision = box_cell(doc, "CURRENT MODULE DECISION")
    box_sub(decision, "MRD v2.80", "MRD v2.93")
    box_append(decision, "A configuration fault opens an incident with no alert or rule "
                         "behind it. It reaches the console only, sends no notification, and "
                         "stays open until an operator resolves it.")

    # 3. Actors and permissions
    actors = table_headed(doc, "Actor", "May do", "Gate or contract")
    keep_format(row(actors, "Platform alarm recipient").cells[-1],
                "Active platform-tenant user whose effective grants include "
                "platform.health.update, the authority that maintains incidents and alert "
                "rules.")
    insert_row_after(row(actors, "Celery beat and workers"), [
        "Configuration-fault reporter",
        "Open one incident for a configuration fault its own module has found.",
        "Internal call to vs_health.faults.report_configuration_fault; no API identity.",
    ])
    matrix = table_headed(doc, "Operation", "Permission", "Scope")
    alarm = row(matrix, "Receive a fired Health alarm")
    keep_format(alarm.cells[1], "Effective platform.health.update")
    insert_row_after(alarm, [
        "Open a configuration-fault incident",
        "None; an internal service call, not reachable from the API",
        "Global incident record",
    ])
    edit_paragraph(doc, "The module imports no school application",
                   "as an engine dependency.",
                   "as an engine dependency. A configuration fault reaches Health through "
                   "vs_health.faults, which takes neutral strings and a fault key, never a "
                   "school.")

    # 4. Requirements
    add_fr(doc, M30_FR016_HEADING, "FR-016 | Implemented", M30_FR016)
    set_status(doc, fr_table(doc, "FR-016"), "Implemented with limits")
    keep_together(fr_table(doc, "FR-016"))

    # 5. Workflow
    lifecycle = table_headed(doc, "State", "Evaluation", "Effect")
    keep_format(row(lifecycle, "First-fire delivery").cells[1],
                "Resolve active platform.health.update holders.")
    add_fault_workflow(doc)

    # 6. Data model
    models = table_headed(doc, "Model", "Purpose", "Key contract")
    append_to(models, "Incident",
              "; fault_key, blank for operator and alert incidents, indexed and returned by "
              "no serializer")
    box_append(box_cell(doc, "OPERATIONAL EVIDENCE RULES"),
               "A configuration-fault incident's fault key is internal: no serializer returns "
               "it, so renaming or re-triaging the incident cannot break its deduplication.")

    # 7. API
    reads = table_headed(doc, "Endpoint", "Purpose")
    append_to(reads, "GET /v1/health/incidents/",
              "; configuration-fault incidents are listed like any other")

    # 8. Dependencies and evidence
    deps = table_headed(doc, "Dependency", "Contract")
    insert_row_after(row(deps, "core.BackgroundJob"), M30_FAL_ROW)
    edit_paragraph(doc, "Assign platform.health.create or platform.health.update",
                   _BY_METHOD, "platform.health.update")
    insert_after(doc, "Monitor the incident timeline",
                 "Watch the incident list for configuration faults: they send no "
                 "notification, and an operator resolves each once its arrangement is "
                 "corrected.")
    evidence = box_cell(doc, "INSPECTED EVIDENCE")
    box_sub(evidence, " The suite was not rerun for this documentation-only revision.", "")
    box_replace(evidence, "Module 30 suite: 37 tests passed",
         "Module 30 suite: 45 tests passed at fa23e06d, including collection, analytics, "
         "small-sample guards, sustained service-scoped alerting, delivery records, incident "
         "references, seeding, authentication, tenant filters and the configuration-fault "
         "rule (tests_faults, 8 tests).")
    box_add_after(evidence, "Module 30 suite:",
         "schools.core suite: 219 tests passed at fa23e06d, 13 of them "
         "test_ambiguous_entity_reporting: one incident per school, survival of the "
         "provisioning rollback, and a failing reporter that leaves the refusal unchanged. "
         "Neither suite was rerun for this revision.")

    # 9. Needs Attention
    attention = table_headed(doc, "Priority", "Current gap", "Required completion", "FR")
    insert_row_after(row(attention, "P1"), [
        "P2",
        "A configuration fault reaches nobody's inbox",
        "Register a configuration-fault event in Module 8 and send it to active "
        "platform.health.update holders when a fault incident opens, with a test that the "
        "first report notifies and a repeat does not.",
        "FR-016",
    ])
    removal = box_cell(doc, "REMOVAL RULE")
    box_sub(removal, "Remove this item only after", "Remove the P1 item only after")
    box_sub(removal, "the relevant monitor and SLO tests pass.",
            "the relevant monitor and SLO tests pass, and the P2 item only after a "
            "configuration-fault incident notifies its recipients under test.")

    # 10. Traceability
    edit_paragraph(
        doc, "Module 30 carries 13 capability entries",
        "Module 30 carries 13 capability entries in MRD v2.80. Each maps to the controlling "
        "requirements below.",
        "Module 30 carries 13 capability entries in MRD v2.93. Each maps to the controlling "
        "requirements below. Configuration-fault incidents strengthen incident records and "
        "events without changing the count.")
    trace = table_headed(doc, "MRD Module 30 capability", "Requirements", "Current state")
    incidents = row(trace, "Incident records and events")
    keep_format(incidents.cells[1], "FR-008, FR-011, FR-013, FR-016")
    keep_format(incidents.cells[2], "Implemented with FR-016 limit")

    log_change(doc, M30_TARGET, M30_SUMMARY)
    assert_absent_outside_log(doc, _BY_METHOD + " holders",
                              _BY_METHOD + " grant", "MRD v2.80", "three evidence paths")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M30_DIR, M30_STEM, M30_TARGET),
           f"{M30_STEM.replace('_', ' ')} v{M30_TARGET}", M30_TARGET)


def main() -> None:
    patch_m08()
    patch_m30()


if __name__ == "__main__":
    main()
