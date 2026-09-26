#!/usr/bin/env python3
"""Cut MRD v2.86 and five module FRDs for record history and person-record Field Access.

What changed in the backend, and therefore in the documents:

* A new engine app, vs_history, keeps a version of every tracked row whenever
  it is created, changed or deleted, including through QuerySet.update and the
  bulk writes, and answers "what did this row hold at the end of that day".
  A record's history starts at its first version and never earlier.
* Every profile read behind the student, guardian, staff and CX staff pages
  accepts ?as_at=YYYY-MM-DD and answers in the same shape as the live read,
  rendered through the live serializers so Field Access applies to the past.
* The name, personal, contact, employment and photograph fields of students,
  guardians, school staff and CX staff are registered Field Access fields, and
  a composite such as full_name is rebuilt from the parts the reader may read.
* A guardian's name is held in parts; names typed on one line are split and
  flagged for a person to confirm.

The script also repairs the change logs a previous revision damaged: the MRD
v2.85 and three FRDs wrote a version row over the header row. Each missing
header is restored from the last version that still carried it.

    python tools/patch_record_history_docs.py
"""
from __future__ import annotations

import copy
import glob
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)
import patch_mrd_v2_79_docs as mrd_tools

ROOT = Path(__file__).resolve().parents[1]
REVIEW_DATE, SHORT_DATE = "26 September 2026", "26 Sep 2026"
MRD_SOURCE, MRD_TARGET = "2.85", "2.86"
CODE_BASELINE = (
    "Backend worktree feat/record-history on 8b56c22e, with the school and "
    "console frontend worktrees, 26 September 2026"
)
#: The full backend suite, filled from the run that verified this revision.
TEST_EVIDENCE = (
    "Verified by the full backend suite, every app run on its own: 6207 tests, all passing, "
    "22 of them the history engine's own; and driven against the running app."
)

# ── generic helpers ──────────────────────────────────────────────────────────


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def keep_format(cell, text: str) -> None:
    """Rewrite a cell keeping the formatting of its first run."""
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    runs = cell.paragraphs[0].runs
    if not runs:
        cell.paragraphs[0].add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def control_table(doc):
    for table in doc.tables[:3]:
        if any(row.cells[0].text.strip() == "Version" for row in table.rows):
            return table
    raise ValueError("No control table carries a Version row")


def set_control(doc, label: str, value: str) -> None:
    table = control_table(doc)
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            keep_format(row.cells[-1], value)
            return
    raise ValueError(f"Control row not found: {label}")


def set_cover_version(doc, source: str, target: str) -> None:
    """The single-cell covers carry 'Version: x'; the grid covers carry none."""
    cell = doc.tables[0].rows[0].cells[0]
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return


def table_with_header(doc, first_cell: str):
    hits = [t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith(first_cell)]
    if not hits:
        raise ValueError(f"No table starts with {first_cell!r}")
    return hits[0]


def append_rows(table, rows: list[list[str]], *, size=8) -> None:
    for values in rows:
        template = table.rows[-1]
        template._tr.addnext(copy.deepcopy(template._tr))
        row = table.rows[-1]
        for cell, value in zip(row.cells, values):
            keep_format(cell, value)


def change_log_table(doc):
    """The table holding the version rows: the last one whose rows start with versions."""
    for table in reversed(doc.tables):
        first = table.rows[0].cells[0].text.strip()
        second = table.rows[1].cells[0].text.strip() if len(table.rows) > 1 else ""
        if first == "Version" and len(table.rows[0].cells) == 3:
            return table
        if first[:1].isdigit() and second[:1].isdigit() and len(table.rows[0].cells) == 3:
            return table
    raise ValueError("No change log table found")


def restore_header(table, header_source) -> None:
    """Put back a header row a previous revision wrote a version over."""
    if table.rows[0].cells[0].text.strip() == "Version":
        return
    header = copy.deepcopy(header_source.rows[0]._tr)
    table.rows[0]._tr.addprevious(header)
    # The row the header was overwritten by is a duplicate of a later row.
    first = table.rows[1]
    duplicates = [r for r in table.rows[2:] if r.cells[2].text == first.cells[2].text]
    if duplicates:
        first._tr.getparent().remove(first._tr)


def log_change(doc, version: str, summary: str, header_source=None) -> None:
    table = change_log_table(doc)
    if header_source is not None:
        restore_header(table, change_log_table(header_source))
    if table.rows[0].cells[0].text.strip() != "Version":
        raise ValueError("Change log has no header row and no source to restore it from")
    template = table.rows[1]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[1]
    keep_format(row.cells[0], version)
    keep_format(row.cells[1], REVIEW_DATE)
    keep_format(row.cells[2], summary)


def update_reconciliation(doc, edits) -> None:
    """Edit the MRD RECONCILIATION box line by line, keeping each line's formatting."""
    box = table_with_header(doc, "MRD RECONCILIATION").rows[0].cells[0]
    lines = [
        line.strip() for p in box.paragraphs for line in p.text.split("\n") if line.strip()
    ]
    lines = mrd_tools.edit_lines(lines, edits)
    if len(box.paragraphs) == 1 and len(box.paragraphs[0].runs) == 1:
        # One run whose lines are separated by breaks: keep it one run.
        box.paragraphs[0].runs[0].text = "\n".join(lines)
    else:
        mrd_tools.set_lines(box, lines)


def last_fr_table(doc):
    fr = [t for t in doc.tables if t.rows[0].cells[0].text.strip().startswith("FR-")]
    return fr[-1]


def add_fr(doc, heading: str, header_cell: str, rows: list[tuple[str, str]]) -> None:
    """Add one requirement: a heading cloned from the last one, and a table cloned from its table.

    The template's rows are reused in order and the labels rewritten, so a new
    requirement carries exactly the formatting of the ones beside it.
    """
    template = last_fr_table(doc)
    heading_p = template._tbl.getprevious()
    while heading_p is not None and heading_p.tag.split("}")[1] != "p":
        heading_p = heading_p.getprevious()
    if heading_p is None or not Paragraph(heading_p, doc).text.strip().startswith("FR-"):
        raise ValueError("The last requirement table has no FR heading above it")
    new_heading = copy.deepcopy(heading_p)
    new_table = copy.deepcopy(template._tbl)
    template._tbl.addnext(new_heading)
    new_heading.addnext(new_table)
    heading_par = Paragraph(new_heading, doc)
    mrd_tools.retitle(heading_par, heading)
    table = next(t for t in doc.tables if t._tbl is new_table)
    while len(table.rows) - 1 > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) - 1 < len(rows):
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    for cell in table.rows[0].cells:
        keep_format(cell, header_cell)
    for row, (label, value) in zip(table.rows[1:], rows):
        keep_format(row.cells[0], label)
        keep_format(row.cells[-1], value)


def finish(doc, output: Path, title: str, version: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise SystemExit(f"{output.name} already exists; never overwrite a version")
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)
    print(f"Wrote {output.name}")


def frd_path(folder: str, stem: str, version: str) -> Path:
    return ROOT / "functional-requirements" / folder / f"{stem}_v{version}.docx"


def oldest_with_header(folder: str, stem: str, newest: str) -> Document:
    """The newest earlier version whose change log still has its header."""
    for path in sorted(
        glob.glob(str(ROOT / "functional-requirements" / folder / f"{stem}_v*.docx")),
        key=lambda p: [int(x) for x in p.rsplit("_v", 1)[1][:-5].split(".")],
        reverse=True,
    ):
        doc = Document(path)
        try:
            if change_log_table(doc).rows[0].cells[0].text.strip() == "Version":
                return doc
        except ValueError:
            continue
    raise ValueError(f"No version of {stem} keeps its change log header")


# ── shared requirement text ──────────────────────────────────────────────────

AS_AT_ENGINE = (
    "vs_history.RecordVersion holds a copy of a tracked row's fields from recorded_at until "
    "the next version. The save, delete and many-to-many signals write one; so do "
    "QuerySet.update, bulk_create and bulk_update, through VersionedQuerySetMixin on every "
    "manager a tracked model declares, which vs_history.registry.track refuses a model "
    "without at start-up. A save that changes no tracked field writes nothing. "
    "manage.py baseline_record_history writes the first version of every row that has "
    "none, dated when it runs, and build.sh runs it on every deploy."
)

AS_AT_RULES = (
    "(1) ?as_at=YYYY-MM-DD is read at the end of that school day in Africa/Lagos: a change "
    "recorded before midnight counts and one recorded after does not. "
    "(2) No date, or today, is the live record. A future date or a malformed one answers 400 "
    "AS_AT_INVALID. "
    "(3) A day before the record's first version answers 409 HISTORY_NOT_KEPT with "
    "history_starts, and nothing is guessed for it. "
    "(4) The response has the live read's shape plus an as_at block (date, history_starts, "
    "photo_retired); the live read carries history_starts so a page knows the earliest day. "
    "(5) Rows are rebuilt from their versions and rendered through the live serializers, so "
    "Field Access and the owner rule apply to the past as to the present. "
    "(6) Which record a caller may open is decided by the live record and its branch, so a "
    "past view never widens what a caller sees. "
    "(7) Anything worked out from today (age, tenure, on leave, Completed leave) is worked "
    "out from the chosen day. "
    "(8) A document or photograph replaced since is reported held without a link, "
    "file_retired, because its file was retired when it was replaced."
)

# ── M11 Student Management ───────────────────────────────────────────────────

M11_DIR, M11_STEM = "11-student-management", "XVS_M11_Student_Management_Functional_Requirements_Document"
M11_SOURCE, M11_TARGET = "2.10", "2.11"

M11_FR026 = [
    ("Description", (
        "Read a student's profile, and a guardian's, as they stood at the end of an earlier "
        "day. An auditor asking what a child's surname, class or guardians were on the day of "
        "an exam gets the answer the record held then, or is told the history starts later."
    )),
    ("Permission", (
        "The live read's own key: school.students.view. No separate history permission: "
        "whoever may read a field reads its past values, and Field Access decides both."
    )),
    ("Trigger", (
        "?as_at=YYYY-MM-DD on GET /v1/students/<id>/, /guardians/, /class-history/, "
        "/status-history/, /history/, /subjects/ and /documents/, and on GET "
        "/v1/guardians/<id>/. " + AS_AT_ENGINE
    )),
    ("Business rules", AS_AT_RULES + (
        " (9) The student, guardian, guardian link, class enrolment and document rows are "
        "tracked; the status log and the audit timeline are dated rows already and are cut "
        "at the end of the day. (10) A class, subject or branch is named as it is named "
        "today, and the subjects are the ones the class's level offers today. "
        "(11) No status move is offered on a past record."
    )),
    ("Acceptance", (
        "(1) A phone and address changed on 10 March read the old values as at 5 March and "
        "the new ones as at 10 March. (2) A class moved on 12 March reads the old class and "
        "class history before it. (3) A guardian unlinked on 15 March is on the 10 March "
        "list with the phone they had then. (4) A document removed since reads attached with "
        "file_retired and no url. (5) A day before the history starts answers 409 with "
        "history_starts; a future day answers 400. (6) Another school's record answers 404 "
        "at any date, and so does another branch's to a branch-pinned head. (7) A field "
        "hidden from a role is absent from the past view exactly as from the live one."
    )),
]

M11_FR027 = [
    ("Description", (
        "Put every part of a student's and a guardian's record under Field Access, the "
        "name included, so a school decides per role who reads and who corrects each."
    )),
    ("Permission", "Field Access switches on school.students and school.guardians."),
    ("Trigger", (
        "Every serializer declared as a surface of those resources, the student type-ahead "
        "included, and the photograph routes."
    )),
    ("Business rules", (
        "(1) school.students registers first, middle and last name, date of birth, gender, "
        "nationality, state of origin, photo (as photo_url), address, phone, email, both "
        "emergency contact fields, previous school and student number beside the medical "
        "fields and the enrolment date. school.guardians registers the three name parts and "
        "photo beside the contact fields. "
        "(2) All are open by default, so no role loses anything on the day they ship. "
        "(3) The name parts, date of birth, gender and student number are open on create for "
        "a student, and a guardian's first and last name are open on create: enrolling is "
        "not correcting. "
        "(4) full_name is rebuilt from the name parts the reader may read, on the profile, "
        "the directory, the guardian surfaces and the type-ahead. A name printed on another "
        "module's document is that module's field. "
        "(5) The passport photograph's Write switch is asked when one is attached or removed, "
        "and its link is withheld from the checklist for a reader who may not read it; a "
        "guardian's photograph route asks its own switch."
    )),
    ("Acceptance", (
        "(1) A role with first name and phone closed reads neither on the profile, the list "
        "or a past view, and reads full_name as the last name alone. (2) A role that may "
        "read a guardian's last name but not change it is refused 403 field_write_denied. "
        "(3) Removing a guardian's photograph without its Write switch is refused 403. "
        "(4) The deep payload case renders every surface closed and finds no registered "
        "name at any depth."
    )),
]

M11_FR028 = [
    ("Description", (
        "Hold a guardian's name as a first, middle and last name, and ask a person to confirm "
        "any name the platform split from one line."
    )),
    ("Permission", "school.students.update to correct or confirm, as for the rest of the record."),
    ("Trigger", (
        "PATCH /v1/guardians/<id>/ with the parts; POST /v1/students/<id>/guardians/ and the "
        "enrol form with the parts for a new guardian; GET /v1/guardians/?name_review=true "
        "for the names awaiting review."
    )),
    ("Business rules", (
        "(1) Migration 0005 splits every existing full_name with schools.vs_students.names: "
        "leading honorifics are set aside, the first word is the first name, the last the "
        "last name, anything between the middle name. Every split row is flagged "
        "name_needs_review and its full_name kept exactly as typed. "
        "(2) A new guardian given only a one-line name, as the spreadsheet imports still do, "
        "is split and flagged the same way; one given its parts is composed and not flagged. "
        "(3) Correcting a name sends its parts; a one-line full_name on the correction route "
        "is refused 400 naming the parts to send, never silently ignored. "
        "(4) Sending the parts of a flagged name, changed or not, confirms it: the flag clears, "
        "full_name is composed from the parts, and the message says the name was confirmed."
    )),
    ("Acceptance", (
        "(1) 'Mrs Ngozi Adaeze Okafor' splits into Ngozi, Adaeze and Okafor, flagged, with "
        "full_name unchanged. (2) Sending those parts unchanged clears the flag and composes "
        "'Ngozi Adaeze Okafor'. (3) A blank last name is refused. (4) An enrolment naming a "
        "new guardian in parts creates it unflagged. (5) The directory filter lists only "
        "flagged guardians."
    )),
]

M11_CORRECTIONS = [
    [
        "v2.10 (section 7.2): a guardian's name is the single column full_name.",
        "The name is held in first_name, middle_name and last_name, with name_needs_review "
        "marking a name split from one line that nobody has confirmed. full_name stays the "
        "one-line form every list and search reads, kept as typed while flagged and composed "
        "from the parts once confirmed.",
        "apps/schools/vs_students/models.py, Guardian.save; names.py; "
        "migrations/0005_guardian_name_parts.py; FR-028.",
    ],
    [
        "v2.10 (FR-001, FR-018): a profile is read as it stands today.",
        "Every profile read accepts ?as_at= and answers as at the end of that day, from the "
        "record history vs_history keeps. Earlier days than the record's first version are "
        "refused, never guessed.",
        "apps/schools/vs_students/as_at.py and history.py; apps/vs_history; FR-026.",
    ],
    [
        "v2.10 (serializers docstring and section 3): the emergency contact carries no Field "
        "Access switch, and only the medical fields and the enrolment date are registered.",
        "The name, personal, contact, emergency, background and photograph fields are "
        "registered too, open by default. The emergency contact's switch starts open, so the "
        "reason it had none (a contact only an administrator can read is useless in an "
        "emergency) is honoured by the default rather than by leaving it unregistered.",
        "apps/schools/vs_students/field_access.py; FR-027.",
    ],
]

M11_SUMMARY = (
    "Minor revision. Three requirements join. FR-026 reads a student's profile and a "
    "guardian's as at the end of an earlier day, from versions the new vs_history engine "
    "keeps for the student, guardian, guardian link, enrolment and document rows; a day "
    "before a record's first version is refused rather than guessed. FR-027 puts the whole "
    "record under Field Access, the name included, with full_name rebuilt from the parts a "
    "reader may read. FR-028 holds a guardian's name in parts and flags every name split "
    "from one line for a person to confirm. Three rows join section 2.2, and the change "
    "log's order is restored. " + TEST_EVIDENCE
)

# ── M12 Staff Management ─────────────────────────────────────────────────────

M12_DIR, M12_STEM = "12-staff-management", "XVS_M12_Staff_Management_Functional_Requirements_Document"
M12_SOURCE, M12_TARGET = "2.8", "2.9"

M12_FR022 = [
    ("Description", (
        "Read a staff profile, and every tab on it, as it stood at the end of an earlier day: "
        "the record, the name on the account, the postings, the roles and their reach, the "
        "permission exceptions, the qualifications, the documents, the teaching assignments "
        "and the leave."
    )),
    ("Permission", "Each read's own key. No separate history permission."),
    ("Business rules", AS_AT_RULES + (
        " (9) StaffProfile (with its postings), StaffQualification, StaffDocument, "
        "TeachingAssignment and LeaveRequest are tracked here; the account, the role grants "
        "and the permission exceptions are tracked by vs_user and vs_rbac. (10) An account "
        "reads its stored status: a lockout is a running window, not a recorded state. "
        "(11) Field exceptions are not kept in history, so a past view does not show them."
    )),
    ("Acceptance", (
        "(1) A last name and job title changed on 10 March read the old values as at 5 March. "
        "(2) A role revoked on 12 March is held as at 5 March and revoked as at 12 March. "
        "(3) A qualification removed on 20 March is listed on 5 March. (4) Leave pending on "
        "2 March reads pending that day, and approved leave covering 5 March makes the "
        "person read on leave that day. (5) A day before the history starts answers 409. "
        "(6) Another school, or another branch to a pinned head, answers 404 at any date."
    )),
]

M12_FR023 = [
    ("Description", (
        "Put the whole staff record under Field Access: the name, the photograph and the "
        "employment details beside the personal details already registered."
    )),
    ("Permission", "Field Access switches on school.teachers."),
    ("Business rules", (
        "(1) school.teachers registers first, middle and last name, staff ID, job title, "
        "employment type, hire date, exit date and photo (sent as photo, read as photo_url), "
        "all open by default. (2) The first and last name live on the account and are open "
        "on create, since the Add form requires them. (3) The exit date is written only by "
        "the status change that ends employment, so it has a Read switch and no Write. "
        "(4) full_name is rebuilt from the name parts the reader may read. (5) The detail "
        "read carries first_name and last_name, so the edit form no longer splits full_name "
        "on its first space."
    )),
    ("Acceptance", (
        "(1) A role with the first name closed reads full_name as the last name alone. "
        "(2) The /me map lists exit_date read-only for every role. (3) The deep payload "
        "case renders every surface closed and finds no registered name at any depth."
    )),
]

M12_SUMMARY = (
    "Minor revision. FR-022 reads a staff profile and every tab on it as at the end of an "
    "earlier day, from versions vs_history keeps for the record, its postings, the account, "
    "the role grants, the permission exceptions, the qualifications, the documents, the "
    "teaching assignments and the leave. FR-023 registers the name, photograph and "
    "employment details under Field Access, with the exit date read-only for everybody and "
    "full_name rebuilt from the parts a reader may read. The change log's header, lost in "
    "an earlier revision, is restored. " + TEST_EVIDENCE
)

M12_TRACE = [
    ["Audit history and field-level access", "FR-012, FR-022, FR-023",
     "Implemented. The profile and every tab read as at an earlier day, and the whole "
     "record's fields are switchable per role."],
]

# ── M03 Identity, Team & Organogram ──────────────────────────────────────────

M03_DIR = "03-identity-team-and-organogram"
M03_STEM = "XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document"
M03_SOURCE, M03_TARGET = "1.15", "1.16"

M03_FR023 = [
    ("Requirement", (
        "A CX staff profile must be readable as it stood at the end of an earlier day, and "
        "every field on it, the name included, must follow Field Access on every surface that "
        "renders the profile: the full profile, the brief colleague view and the staff list."
    )),
    ("Current evidence", (
        "GET /platform-staff-profiles/<id>/?as_at= rebuilds the profile and its account from "
        "vs_history (vs_user.as_at, tracked by vs_user.history) and renders it through the "
        "full or brief serializer the live read would choose. platform.staff_profile registers "
        "the name, personal, contact, next-of-kin and employment fields beside the payroll "
        "bank fields; the account block is StaffProfileAccountSerializer, and the list and "
        "brief serializers are declared surfaces carrying the owner rule. UserUpdateSerializer "
        "asks the name switches when a CX staff member's name is corrected on somebody else's "
        "account. An account is tracked by an allow-list of identity fields, never by "
        "exclusion, so no credential reaches the history."
    )),
    ("Acceptance", (
        "A job title and last name changed on 10 March read the old values as at 5 March; "
        "a day before the history starts answers 409 with history_starts; a colleague's brief "
        "view and the staff list withhold a field their role may not read; a person always "
        "reads their own."
    )),
    ("Limit", (
        "The seat is the one the profile named that day, while the unit, department, division "
        "and line manager around it are read from today's organisation, which keeps its own "
        "effective-dated history. A photograph replaced since is not shown. Correcting a "
        "school account's name is governed by no platform switch."
    )),
]

M03_ATTENTION = [
    ["P3", "A past CX staff view reads the organisation around the seat from today",
     "Compose the unit, department, division and line manager from PositionAssignment's "
     "effective dates, so a restructured organisation reads as it stood.", "FR-023"],
]

M03_TRACE = [
    ["Platform staff profile read as at an earlier date", "FR-018, FR-023",
     "Implemented with limits. The profile and account read as they stood; the organisation "
     "around the seat reads as it stands."],
]

M03_SUMMARY = (
    "Adds FR-023: a CX staff profile reads as at the end of an earlier day, and every field on "
    "it, the name included, follows Field Access on the full profile, the brief colleague view "
    "and the staff list, with the account block switched as its own surface and name "
    "corrections on another person's account asking the same switch. One Needs Attention row "
    "and one traceability row join. The change log's header is restored. " + TEST_EVIDENCE
)

M03_RECONCILIATION = [
    ("replace", "• MRD v2.85 lists Module 3",
     "• MRD v2.86 lists Module 3 as Backend Partial and In use Complete with nineteen "
     "capabilities. The ID-card sign-in entry maps to FR-008; the flow existed without a "
     "tracker entry, and the random card key and card sign-in make it a capability of its "
     "own. Reading a platform staff profile as at an earlier date is a new entry, mapped to "
     "FR-023."),
]

# ── M04 Roles & Permissions ──────────────────────────────────────────────────

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.25", "1.26"

M04_FR027 = [
    ("Requirement", (
        "Field Access must cover a person's whole record, the name included, and a value "
        "built from registered fields must never carry one the reader may not read."
    )),
    ("Current evidence", (
        "FieldAccessMixin.field_composites names a composite (full_name) and the attribute "
        "path of each part; when a part is unreadable the composite is rebuilt from the "
        "readable parts in order. The student, guardian, school staff and CX staff "
        "declarations register the name, personal, contact, employment and photograph "
        "fields, open by default; the type-ahead and the CX staff list and brief views are "
        "declared surfaces. A photograph written through its own route asks its Write switch "
        "there (assert_writable). The registry guards list, with a reason, every serializer "
        "that emits or writes a registered name without being a surface: a person named "
        "beside something else, the Team account screen, account creation, and fields of the "
        "same name that belong to another resource."
    )),
    ("Acceptance", (
        "A role with the first name closed reads full_name as the last name alone on every "
        "surface; the deep payload cases find no registered name at any depth; the read and "
        "write path guards pass with every exception named."
    )),
    ("Limit", (
        "A name printed on another module's document (an invoice, a class list, an audit line) "
        "is that module's field and is not switched. Field exceptions are not kept in the "
        "record history."
    )),
]

M04_ATTENTION = [
    ["P2", "Customer, vendor, vendor contact, bank account, school profile and branch fields "
     "carry no switches yet",
     "Register the approved fields (customer name and billing details; vendor name, category, "
     "payment terms, KYC status, risk and hold; vendor contact name, email and phone; bank "
     "account and bank name; school name, address and registration ID; branch name, address "
     "and email) with the FinPro screens and the school and console settings screens handling "
     "an absent field, released together.", "FR-022, FR-027"],
    ["P3", "Field exceptions keep no history",
     "Track UserFieldAccessOverride in vs_history so a past staff view can show the exceptions "
     "in force that day.", "FR-027"],
]

M04_TRACE = [
    ["Field Access over whole person records, with composite names", "FR-022, FR-027",
     "Implemented. Student, guardian, school staff and CX staff records."],
]

M04_RECONCILIATION = [
    ("replace", "• MRD v2.85 lists Module 4",
     "• MRD v2.86 lists Module 4 as Roles & Permissions (RBAC), Phase V1, Backend Complete, "
     "In use Complete, code vs_rbac, with twenty-three capability entries. The module number, "
     "name, phase, states and ownership agree with this revision."),
    ("replace", "• The count does not change.",
     "• Administrator-owned group CRUD, backend dependency reconciliation and safe "
     "legacy-group retirement restate the existing permission-group, dependency and seeding "
     "capabilities rather than adding entries."),
    ("replace", "• Selected equal branch reach",
     "• Selected equal branch reach extends the existing role creation, assignment and "
     "entity-aware evaluation entries without adding one."),
    ("append",
     "• Field Access over whole person records, with composite names rebuilt from readable "
     "parts, is a new entry mapped to FR-027, which takes the count to twenty-three."),
]

M04_SUMMARY = (
    "Adds FR-027: Field Access covers a person's whole record. The mixin rebuilds a composite "
    "such as full_name from the parts a reader may read, the student, guardian, school staff "
    "and CX staff declarations register their name, personal, contact, employment and "
    "photograph fields open by default, photograph routes ask their Write switch, and the "
    "registry guards name every exception with its reason. Role grants and permission "
    "exceptions keep a record history (vs_rbac.history). Two Needs Attention rows join. The "
    "change log's header is restored. " + TEST_EVIDENCE
)

# ── M05 Audit & Activity Logging ─────────────────────────────────────────────

M05_DIR = "05-audit-and-activity-logging"
M05_STEM = "XVS_M05_Audit_and_Activity_Logging_Functional_Requirements_Document"
M05_SOURCE, M05_TARGET = "1.4", "1.5"

M05_FR017 = [
    ("Requirement", (
        "Beside the event trail, the platform must keep what a record held, so a record can be "
        "read as it stood at the end of an earlier day, and must never present a value it did "
        "not record as a past one."
    )),
    ("Current evidence", AS_AT_ENGINE + (
        " vs_history.as_at parses ?as_at=, answers version_at and instances_at (one indexed "
        "query per row or per owner, through the owners GIN index), and refuses a day before "
        "the first version with HISTORY_NOT_KEPT. A version records the audit identity's real "
        "actor, as the event trail does during a proxy session."
    )),
    ("Acceptance", (
        "A create writes a first version naming every field; an unchanged save writes nothing "
        "and a save of untracked columns reads nothing; update, bulk_create, bulk_update, "
        "delete and many-to-many changes from either side are recorded; a day reads changes up "
        "to its last minute in school time; the baseline command adds one version per unseen "
        "row and nothing on a rerun."
    )),
    ("Limit", (
        "Raw SQL and data migrations write no version and must call vs_history.recorder.record "
        "for the rows they touch. History starts on the day tracking reached a record."
    )),
]

M05_ATTENTION = [
    ["P2", "sync_field_registry writes an action type the audit vocabulary does not hold",
     "Register FIELD_REGISTRY_SYNCED in AuditActionType, or map the sync to an existing RBAC "
     "action, so the platform event it mirrors is not dropped with a validation error.",
     "FR-002"],
]

M05_TRACE = [
    ["Record history read as at an earlier day", "FR-017",
     "Implemented. Student, guardian, school staff, CX staff, account, role grant and "
     "permission exception rows."],
]

M05_RECONCILIATION = [
    ("replace", "•  MRD v2.76 lists Module 5",
     "•  MRD v2.86 lists Module 5 as Backend Complete and In use Complete with fourteen "
     "capabilities and no material capability gap. This FRD does not dispute Complete for the "
     "module's stated backend scope: every listed capability has a backend path."),
    ("replace", "•  No capability is added",
     "•  Record history, read as at an earlier day, is a new entry mapped to FR-017, taking "
     "the count to fourteen. The two documents agree on the trail's tenant confinement and "
     "remain versioned separately, each on its own decision."),
]

M05_SUMMARY = (
    "Adds FR-017: the vs_history engine keeps a version of every tracked row, through saves, "
    "deletes, many-to-many changes and the signal-free queryset writes, and answers what a "
    "row held at the end of an earlier day, refusing any day before its first version. One "
    "Needs Attention row records a defect found while verifying: sync_field_registry emits "
    "FIELD_REGISTRY_SYNCED, which the audit vocabulary refuses. " + TEST_EVIDENCE
)

# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_SOURCE_SCOPE = CODE_BASELINE
MRD_CONTENTS_NOTE = "Record history and person-record Field Access"
MRD_INTRO = (
    "This revision records two things a school will be asked about later. A new engine keeps "
    "what a person's record held, so a student's, guardian's, staff member's or CX staff "
    "member's profile can be read as it stood at the end of an earlier day, and it refuses a "
    "day before the record's history starts rather than showing today's values under an "
    "earlier date. And the whole of those records, the name included, is under Field Access, "
    "so a school decides per role who reads and who corrects each field. Six capability "
    "entries are added, taking the platform to 509."
)
MRD_CHANGE_SUMMARY = (
    "Adds record history: vs_history keeps a version of every tracked row and every profile "
    "read behind the student, guardian, staff and CX staff pages accepts a date, answering as "
    "at the end of that day through the live serializers so Field Access applies to the past; "
    "a day before a record's first version is refused. Puts the name, personal, contact, "
    "employment and photograph fields of those records under Field Access, open by default, "
    "with full_name rebuilt from the parts a reader may read. Holds a guardian's name in "
    "parts, with names split from one line flagged for review. Six capability entries join "
    "(M03 +1, M04 +1, M05 +1, M11 +2, M12 +1), 509 across 31 modules; statuses do not move. "
    "Restores the change log header v2.85 wrote over and retitles the delta section, which "
    "still named v2.83. M03 v1.16, M04 v1.26, M05 v1.5, M11 v2.11 and M12 v2.9. Backend "
    "evidence only."
)
MRD_DELTA_ROWS = [
    ["Record history", "Kept for person records",
     "Every tracked row gains a version on create, change and delete, including through "
     "update and the bulk writes; a baseline dates each existing record's history from the "
     "day it ran."],
    ["As at", "One date parameter on every profile read",
     "Student, guardian, staff and CX staff profiles and their tabs answer as at the end of a "
     "day, in the live shape, with Field Access applied; an earlier day than the history is "
     "refused, never guessed."],
    ["Person-record Field Access", "The whole record, the name included",
     "Name, personal, contact, employment and photograph fields are registered open by "
     "default; full_name is rebuilt from readable parts on every surface."],
    ["Guardian names", "Held in parts, reviewed when split",
     "Existing one-line names are split and flagged, kept as typed until a person confirms "
     "the parts."],
    ["Still to register", "Master data fields",
     "Customer, vendor, vendor contact, bank account, school profile and branch fields wait "
     "on a FinPro release and on the settings screens."],
    ["Module FRDs", "Five revised", "M03 v1.16, M04 v1.26, M05 v1.5, M11 v2.11, M12 v2.9."],
]

MRD_MODULES = {
    3: {"add": [("▸  Platform staff profiles",
                 "▸  Platform staff profile read as at an earlier date")], "count": 19},
    4: {"add": [("▸  Permission and assignment audit history",
                 "▸  Field Access over whole person records, with composite names")],
        "count": 23},
    5: {"add": [("▸  Entity-level audit trails",
                 "▸  Record history, read as at an earlier day")], "count": 14,
        "blurb_edits": [(
            "take the business change down with it.",
            "take the business change down with it. Beside the trail, vs_history keeps what a "
            "record held, so a person's record can be read as it stood at the end of an "
            "earlier day.",
        )]},
    11: {"add": [
        ("▸  Student profile and school-issued identifier",
         "▸  Student and guardian record read as at an earlier date"),
        ("▸  Guardian and emergency-contact relationships",
         "▸  Guardian names in parts, with review of names split from one line"),
    ], "count": 12, "blurb_edits": [(
        "Documented by M11 FRD v2.9",
        "Every profile read answers as at an earlier day, and the whole record is under "
        "Field Access. Documented by M11 FRD v2.11",
    )]},
    12: {"add": [("▸  Staff profile linked to identity",
                  "▸  Staff record read as at an earlier date")], "count": 11,
         "box_edits": [(
             "replace", "• Field-level access is not applied to the staff payload.",
             "• Field-level access covers the staff record: the name, personal, employment "
             "and photograph fields each have a Read and a Write switch per role, full_name is "
             "rebuilt from the parts a reader may read, and a past view applies the same "
             "switches. Field exceptions are not yet kept in the record history.",
         )], "blurb_edits": [("Documented by M12 FRD v2.8", "Documented by M12 FRD v2.9")]},
}


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    source = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    header_source = Document(str(folder / "XVS_Module_Requirements_Document_v2.84.docx"))
    doc = source
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta, log = tables[76], tables[78]
    grids = {m: tables[i] for m, i in mrd_tools.CAPABILITIES.items()}
    assert delta.rows[0].cells[0].text.strip().endswith("capability delta")
    assert index.rows[0].cells[5].text.strip() == "Entries"

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for row in control.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], MRD_SOURCE_SCOPE, size=9)
    for row in contents.rows:
        if row.cells[0].text.strip().startswith("5."):
            keep_format(row.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(row.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("5. ") and paragraph.style is not None and paragraph.style.name.startswith("Heading"):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    blurbs = mrd_tools.module_blurbs(doc)
    for module, changes in MRD_MODULES.items():
        grid = grids[module]
        for anchor, new in changes.get("add", []):
            mrd_tools.add_bullet(grid, anchor, new)
        if changes.get("box_edits"):
            box = grid.rows[-1].cells[0]
            mrd_tools.set_box(box, mrd_tools.edit_lines(mrd_tools.box_lines(box), changes["box_edits"]))
        for old, new in changes.get("blurb_edits", []):
            text = blurbs[module].text
            if text.count(old) != 1:
                raise ValueError(f"Module {module} description holds {old!r} {text.count(old)} times")
            mrd_tools.retitle(blurbs[module], text.replace(old, new))
        keep_format(mrd_tools.row_keyed(index, 0, str(module)).cells[5], str(changes["count"]))

    total = sum(int(row.cells[5].text.strip()) for row in index.rows[1:])
    for row in control.rows:
        if row.cells[0].text.strip() == "Capability entries":
            replace_cell(row.cells[1], str(total), size=9)
    if total != 509:
        raise ValueError(f"Capability total is {total}, expected 509")

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)
    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if intro:
        mrd_tools.retitle(intro[0], MRD_INTRO)

    restore_header(log, header_source.tables[78])
    mrd_tools.prepend_change_log(log, MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)

    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def patch_grid_frd(folder, stem, source, target, *, frs, corrections=None, trace=None,
                   summary, date_label="Date", supersedes=True):
    """M11 and M12: a grid cover, label-and-value requirement tables."""
    doc = Document(str(frd_path(folder, stem, source)))
    set_control(doc, "Version", target)
    set_control(doc, date_label, REVIEW_DATE)
    if supersedes:
        set_control(doc, "Supersedes", f"v{source}")
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET}")
    set_control(doc, "Verified against", CODE_BASELINE)
    for heading, header, rows in frs:
        add_fr(doc, heading, header, rows)
    if corrections:
        append_rows(table_with_header(doc, "What an earlier version said"), corrections)
    if trace:
        append_rows(table_with_header(doc, "MRD capability"), trace)
    log_change(doc, target, summary, header_source=oldest_with_header(folder, stem, source))
    title = f"{stem.replace('_', ' ')} v{target}"
    finish(doc, frd_path(folder, stem, target), title, target)


def patch_status_frd(folder, stem, source, target, *, frs, attention, trace, summary,
                     baseline_label, reconciliation):
    """M03, M04 and M05: a single-cell cover, status-headed requirement tables."""
    doc = Document(str(frd_path(folder, stem, source)))
    set_cover_version(doc, source, target)
    set_control(doc, "Version", target)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, baseline_label, CODE_BASELINE)
    for label, value in (("Code inspected", CODE_BASELINE),):
        try:
            set_control(doc, label, value)
        except ValueError:
            pass
    for label in ("Source MRD", "MRD baseline"):
        try:
            set_control(doc, label, f"XVS Module Requirements Document v{MRD_TARGET}")
        except ValueError:
            pass
    for heading, header, rows in frs:
        add_fr(doc, heading, header, rows)
    append_rows(table_with_header(doc, "Pri."), attention)
    update_reconciliation(doc, reconciliation)
    append_rows(next(t for t in doc.tables
                     if t.rows[0].cells[0].text.strip().startswith("MRD Module")), trace)
    log_change(doc, target, summary, header_source=oldest_with_header(folder, stem, source))
    title = f"{stem.replace('_', ' ')} v{target}"
    finish(doc, frd_path(folder, stem, target), title, target)


def main() -> None:
    if "PLACEHOLDER" in TEST_EVIDENCE:
        raise SystemExit("Refusing to cut: TEST_EVIDENCE is still a placeholder")
    patch_grid_frd(
        M11_DIR, M11_STEM, M11_SOURCE, M11_TARGET,
        frs=[
            ("FR-026  Read a record as at an earlier day", "FR-026  Read a Record As at an Earlier Day", M11_FR026),
            ("FR-027  Field Access over the whole record", "FR-027  Field Access Over the Whole Record", M11_FR027),
            ("FR-028  A guardian's name in parts", "FR-028  A Guardian's Name in Parts", M11_FR028),
        ],
        corrections=M11_CORRECTIONS, summary=M11_SUMMARY,
    )
    patch_grid_frd(
        M12_DIR, M12_STEM, M12_SOURCE, M12_TARGET,
        frs=[
            ("FR-022  Read a staff profile as at an earlier day", "FR-022  Read a Staff Profile As at an Earlier Day", M12_FR022),
            ("FR-023  Field Access over the whole staff record", "FR-023  Field Access Over the Whole Staff Record", M12_FR023),
        ],
        trace=M12_TRACE, summary=M12_SUMMARY,
    )
    patch_status_frd(
        M03_DIR, M03_STEM, M03_SOURCE, M03_TARGET,
        frs=[("FR-023  Read a CX Staff Profile As at an Earlier Day, With Every Field Switched",
              "FR-023 | Implemented with limits", M03_FR023)],
        attention=M03_ATTENTION, trace=M03_TRACE, summary=M03_SUMMARY,
        baseline_label="Code baseline", reconciliation=M03_RECONCILIATION,
    )
    patch_status_frd(
        M04_DIR, M04_STEM, M04_SOURCE, M04_TARGET,
        frs=[("FR-027  Switch a Person's Whole Record, and Never Leak a Part Through a Composite",
              "FR-027 | Implemented with limits", M04_FR027)],
        attention=M04_ATTENTION, trace=M04_TRACE, summary=M04_SUMMARY,
        baseline_label="Source scope", reconciliation=M04_RECONCILIATION,
    )
    patch_status_frd(
        M05_DIR, M05_STEM, M05_SOURCE, M05_TARGET,
        frs=[("FR-017  Keep What a Record Held, and Refuse a Day It Did Not Record",
              "FR-017 | Implemented with limits", M05_FR017)],
        attention=M05_ATTENTION, trace=M05_TRACE, summary=M05_SUMMARY,
        baseline_label="Code baseline", reconciliation=M05_RECONCILIATION,
    )
    patch_mrd()


if __name__ == "__main__":
    main()
