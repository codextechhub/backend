#!/usr/bin/env python3
"""Reconcile the MRD's Module 6 entry to the new M06 FRD.

Module 6 had no document of its own, so its decision box in the MRD had become
the place every decision about configuration and capability was written down:
ten bullets, each narrating one revision. The FRD now owns that detail, so the
box is rewritten as current state pointing at it, and keeps only the facts that
cross module boundaries and therefore belong in a cross-module tracker.

    python tools/patch_module6_reconciliation_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
    write_paragraph,
)

REVIEW_DATE = "9 September 2026"
CHANGE_DATE = "9 Sep 2026"

MRD_SOURCE, MRD_TARGET = "2.71", "2.72"
M06_FRD_VERSION = "1.0"

SOURCE_SCOPE = (
    "Documentation change reconciling Module 6 to its first Functional "
    "Requirements Document, and moving that module's accumulated detail out of "
    "the tracker and into it (9 September 2026)"
)

INTRO = (
    "This revision adds no capability and changes no status. It gives Module 6 "
    "the same shape every other documented module has: a tracker entry that "
    "says where the module stands, and an FRD that says how it behaves."
)

DELTA_ROWS = [
    ["Module 6's decision box", "Rewritten as current state",
     "Ten bullets, each narrating one past revision, replaced by what is true "
     "now plus the facts that cross module boundaries. The detail they carried "
     "is in M06 FRD v1.0, which states it as testable requirements rather than "
     "as history."],
    ["Module 6 traceability", f"Reconciled to M06 FRD v{M06_FRD_VERSION}",
     "All twenty-six capability entries map to one of the FRD's twenty-two "
     "requirements. The count, the module status and the integration status "
     "are unchanged."],
    ["Module 6's current gaps", "Recorded in the FRD",
     "Five, the first being that every configuration key is platform-scoped, "
     "so a setting a school ought to own for itself has no route. None is a "
     "defect, so none is added to Priority Gaps."],
]

M06_DECISION = [
    "Module 6 remains Backend Complete and In use Complete with twenty-six "
    f"capability entries, reconciled to M06 FRD v{M06_FRD_VERSION}.",
    "The detail now lives in the FRD. Twenty-two requirements, the resolution "
    "and evaluation orders, the API contracts and the refusal table are stated "
    "there as behaviour that can be tested. What stays here is what other "
    "modules have to know.",
    "Configuration is reserved to the platform tenant. All nineteen config.* "
    "keys are PLATFORM-scoped, which the RBAC grant guard enforces "
    "independently of how a school composes its roles, and the enforcement "
    "switch the plan gate reads is platform-scoped for the same reason. The "
    "one school-facing route carries no permission key at all and is "
    "read-only.",
    "The catalogue is two levels deep and depth is sold per module. A module "
    "is what a school holds; a band is that module cut at Core, Plus or "
    "Advanced, and it opens when the module's depth reaches it. A null depth "
    "means unlimited rather than shallowest, which is what keeps a school "
    "granted before depth existed from losing Advanced work.",
    "A school with no package grant was never given a plan, which is not the "
    "same as having bought nothing. The endpoint a school's navigation reads "
    "and the gate that issues refusals apply that rule from one place, so a "
    "menu cannot disagree with the product.",
    "Both platform bands sit at Core. Bulk import and data export are one "
    "promise read in two directions, and pricing either above the cheapest "
    "plan made starting or leaving a paid feature.",
    "Five current gaps are recorded in the FRD rather than here. The first is "
    "the shape of the module rather than an oversight: a setting a school "
    "ought to own for itself has no route, and one would have to be introduced "
    "deliberately.",
]

CHANGE_SUMMARY = (
    "Reconciled Module 6 to its first FRD and moved that module's accumulated "
    "detail into it. With no document of its own, Module 6's decision box had "
    "become where every configuration and capability decision was written "
    "down: ten bullets, each narrating one revision, several describing work "
    "finished many versions earlier. M06 FRD v1.0 states that ground as "
    "twenty-two testable requirements covering the two catalogues and the one "
    "scope chain they share, and all twenty-six capability entries map to one "
    "of them. The box now carries what a cross-module tracker should: that "
    "configuration is reserved to the platform tenant and why, that the "
    "catalogue is two levels deep with a null depth meaning unlimited, that a "
    "school with no package grant was never provisioned rather than having "
    "bought nothing, and that both platform bands sit at Core. No capability "
    "was added, no status changed, and the count stays at 492 across 31 "
    "modules. The FRD records five current gaps, none of them a defect, so "
    "Priority Gaps is unchanged. Documentation only; no code was touched."
)


def replace_cell(cell, text, **kwargs):
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def write_decision_cell(cell, title, lines, *, size=8.2):
    """Rewrite a decision box, keeping its banner styled as a banner."""
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_paragraph(
        cell.paragraphs[0], title.upper(), size=size, bold=True, color=BLUE,
        space_after=3,
    )
    for line in lines:
        paragraph = cell.add_paragraph()
        write_paragraph(paragraph, f"• {line}", size=size, space_after=2)


def retitle(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def replace_cover_version(table, source, target):
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def update_control_table(table, *, version, source_scope):
    for row in table.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], version, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], source_scope, size=9)


def prepend_change_log(table, version, date, summary):
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], date, size=8)
    replace_cell(row.cells[2], summary, size=8)


def keep_rows_whole(table):
    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.core_properties.version = MRD_TARGET

    replace_cover_version(doc.tables[0], MRD_SOURCE, MRD_TARGET)
    update_control_table(
        doc.tables[1], version=MRD_TARGET, source_scope=SOURCE_SCOPE,
    )
    for row in doc.tables[2].rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(
                row.cells[0], f"5. v{MRD_TARGET} Capability Delta",
                size=9, bold=True, color=BLUE,
            )
            replace_cell(
                row.cells[1],
                "Module 6 reconciled to its first FRD",
                size=9,
            )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == f"5. v{MRD_SOURCE} Capability Delta":
            retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")
        elif text.startswith("This revision") and len(text) > 80:
            retitle(paragraph, INTRO)

    write_decision_cell(
        doc.tables[18].rows[-1].cells[0], "Current decision", M06_DECISION,
    )

    rebuild_table(
        doc.tables[76],
        [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
        DELTA_ROWS,
        [1.85, 1.35, 3.85],
    )
    keep_rows_whole(doc.tables[76])

    prepend_change_log(doc.tables[78], MRD_TARGET, CHANGE_DATE, CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, f"XVS Module Requirements Document v{MRD_TARGET}")
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    mrd = Path(args.root) / "module-requirements"
    patch_mrd(
        mrd / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx",
        mrd / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
    )
    print(f"Wrote MRD v{MRD_TARGET}")


if __name__ == "__main__":
    main()
