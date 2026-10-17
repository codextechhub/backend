#!/usr/bin/env python3
"""Version the Module 3 FRD for session rows ending with every all-session revocation.

The MRD is read and left at its current version: its Module 3 entry lists no
gap or limitation this change resolves, and its capability list and states are
unchanged.
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    assert_no_em_dash,
    format_header_row,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)


REVIEW_DATE = "17 October 2026"
SHORT_DATE = "17 October 2026"
MRD_VERSION = "2.107"
FRD_SOURCE_VERSION = "1.23"
FRD_TARGET_VERSION = "1.24"
CODE_BASELINE = "Working tree on cf0ae409 (17 October 2026)"


def replace_cell(cell, text: str, **kwargs) -> None:
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_control_value(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            replace_cell(row.cells[1], value, size=9)
            return
    raise ValueError(f"Control row not found: {label}")


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def requirement_table(doc, fr: str):
    for table in doc.tables:
        if table.rows[0].cells[0].text.strip().startswith(f"{fr} |"):
            return table
    raise ValueError(f"Requirement table not found: {fr}")


def requirement_row(table, label: str):
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            return row
    raise ValueError(f"Requirement row not found: {label}")


def table_with_header(doc, header: str):
    for table in doc.tables:
        if table.rows[0].cells[0].text.strip() == header:
            return table
    raise ValueError(f"Table not found: {header}")


def set_paragraph_text(paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        paragraph._p.remove(run._r)


def repair_change_log(table) -> None:
    """Restore the change log's header row and the styling of the v1.23 entry.

    The v1.23 source carries its predecessor's entry in the header row, and its
    own entry in a copy of the header (blue fill, header type), so the table
    opens with a duplicate 1.22 entry and a styled 1.23 one. Row 0 becomes the
    header again and the 1.23 entry is rebuilt from an ordinary entry row.
    """
    for cell, label in zip(table.rows[0].cells, ["Version", "Date", "Summary"]):
        replace_cell(cell, label, size=8.5, bold=True)
    format_header_row(table.rows[0])

    styled, ordinary = table.rows[1], table.rows[2]
    values = [cell.text.strip() for cell in styled.cells]
    styled._tr.addprevious(copy.deepcopy(ordinary._tr))
    styled._tr.getparent().remove(styled._tr)
    for cell, value in zip(table.rows[1].cells, values):
        replace_cell(cell, value, size=8)


def prepend_change_log(table, version: str, summary: str) -> None:
    template = table.rows[1]
    template._tr.addprevious(copy.deepcopy(template._tr))
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=8)
    replace_cell(row.cells[1], SHORT_DATE, size=8)
    replace_cell(row.cells[2], summary, size=8)


def finish(doc, output: Path, title: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def patch_frd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M03 Identity Team and Organogram Functional Requirements "
        f"Document v{FRD_TARGET_VERSION}"
    )
    doc.core_properties.title = title
    doc.core_properties.version = FRD_TARGET_VERSION
    replace_cover_version(doc.tables[0], FRD_SOURCE_VERSION, FRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Version", FRD_TARGET_VERSION)
    replace_control_value(doc.tables[1], "Review date", REVIEW_DATE)
    replace_control_value(doc.tables[1], "Code baseline", CODE_BASELINE)
    replace_control_value(
        doc.tables[1],
        "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module M03",
    )

    fr013 = requirement_table(doc, "FR-013")
    replace_cell(
        requirement_row(fr013, "Current limit").cells[1],
        "Used and expired reset rows are retained indefinitely. Access tokens issued before a successful reset remain usable until their short expiry; refresh tokens are blacklisted and every session row of the account ends with PASSWORD_RESET.",
        size=8.5,
    )

    fr014 = requirement_table(doc, "FR-014")
    replace_cell(
        requirement_row(fr014, "Current evidence").cells[1],
        "PasswordService.change verifies the current password in the serializer, validates the new one against the shared policy, refuses a new password equal to the old, writes password_changed_at, and calls blacklist_all_user_tokens with the end reason PASSWORD_CHANGE. In one transaction that helper blacklists every outstanding refresh token, ends every active LoginSession row of the account with the reason and the time, the session that made the change included, and ends any impersonation session riding the account. PASSWORD_CHANGED is emitted. The helper is the one all-session revocation and takes a required reason from each caller: reset (PASSWORD_RESET), suspend (SUSPENDED), deactivate and a portal account closed by a product (DEACTIVATED), email change (EMAIL_CHANGE), and the self-service and administrative end-everything routes (FORCE_LOGOUT). A device session it ends releases the push subscription tied to that device, as signing out on the device does.",
        size=8.5,
    )
    replace_cell(
        requirement_row(fr014, "Acceptance").cells[1],
        "The endpoint requires only an active session, not a permission key, and is part of the pending-tenant surface so a school's first administrator can use it before go-live. EndingEverySessionTests proves a password change ends both of an account's sessions, the one that made it included, with PASSWORD_CHANGE and an end time, leaves another account's session live, and that signing out afterwards answers 200 and keeps the reason. DeviceSessionTests proves a device session ends at once with its reason and its device is released.",
        size=8.5,
    )
    replace_cell(
        requirement_row(fr014, "Current limit").cells[1],
        "Access tokens already issued remain valid until they expire; only refresh tokens are blacklisted, so a page holding an access token keeps working until its fifteen-minute expiry, while every session row reads ended at once.",
        size=8.5,
    )

    fr016 = requirement_table(doc, "FR-016")
    acceptance = requirement_row(fr016, "Acceptance").cells[1]
    replace_cell(
        acceptance,
        acceptance.text.strip()
        + " EndingEverySessionTests proves deactivation ends every session row with DEACTIVATED and blacklists each session's refresh token.",
        size=8.5,
    )

    reconciliation = next(
        table for table in doc.tables
        if table.rows[0].cells[0].text.strip().startswith("MRD RECONCILIATION")
    )
    cell = reconciliation.rows[0].cells[0]
    lines = cell.text.strip().split("\n")
    lines[1] = (
        f"• MRD v{MRD_VERSION} lists Module 3 as Backend Partial and In use Complete with twenty "
        "capabilities, each mapped above. This revision changes how session revocation is "
        "recorded inside an existing capability and adds none."
    )
    replace_cell(cell, "\n".join(lines), size=8.5)

    traceability = table_with_header(doc, "MRD Module 3 capability")
    for row in traceability.rows:
        if row.cells[0].text.strip() == "Session listing and revocation":
            replace_cell(
                row.cells[2],
                "Implemented with limits. Self-service and administrative targets are scoped. Every all-session revocation ends the session rows with its reason in the same transaction as the blacklist, so no list shows a revoked session as live; stale session rows are swept only when a list is opened.",
                size=8.2,
            )
            break
    else:
        raise ValueError("Traceability row not found: Session listing and revocation")

    rewritten = set()
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("For this revision the client-address tests ran"):
            rewritten.add("verification")
            set_paragraph_text(
                paragraph,
                "For this revision the session-ending change was verified on its own: EndingEverySessionTests and the device-session suite, then the vs_user app, all OK. Earlier evidence stands as recorded: the client-address tests at a5bae05e (core 20, vs_user 3) and the lockout, branch reach and as-at changes when committed (vs_user 405, 414 and 421). The client evidence is unchanged from v1.10: the migrated school client passed 46 focused auth tests and its production build, and its source and production bundle contain neither js-cookie nor refresh-token identifiers. At v1.17 the school client's typecheck and its auth and staff tests passed, and its sign-in, forgot-password and Change email address screens were driven against the running backend with a staff ID and an address. Deployment is not inferred from these checks.",
            )
        elif text.startswith("MRD v2.93 records Module 3"):
            rewritten.add("traceability")
            set_paragraph_text(
                paragraph,
                f"MRD v{MRD_VERSION} records Module 3 as Identity, Team & Organogram, Phase V1, Backend Partial, In use Complete, code ownership vs_user, with twenty capability entries. Each entry maps below.",
            )
        elif text.startswith("•  The empty-list response shape through the shared envelope"):
            rewritten.add("bullet")
            bullet = copy.deepcopy(paragraph._p)
            paragraph._p.addnext(bullet)
            set_paragraph_text(
                doc.paragraphs[index + 1],
                "•  Every all-session revocation ends the session rows: a password change ends the account's sessions, the changing one included, with PASSWORD_CHANGE and leaves another account's alone, a sign-out afterwards keeps the reason, deactivation ends them with DEACTIVATED, and a device session ended this way releases its push device. Covered by EndingEverySessionTests and DeviceSessionTests.",
            )
    if rewritten != {"verification", "traceability", "bullet"}:
        raise ValueError(f"Paragraphs not found; rewrote only {sorted(rewritten)}")

    change_log = doc.tables[-1]
    repair_change_log(change_log)
    prepend_change_log(
        change_log,
        FRD_TARGET_VERSION,
        "Closes the FR-014 limit that left LoginSession rows reading active after a password change. blacklist_all_user_tokens now ends every active session row of the account with a required end reason in the same transaction as the blacklist, and is the one all-session revocation: a password change (PASSWORD_CHANGE), a completed reset (PASSWORD_RESET) and a deactivation (DEACTIVATED), which ended no rows before, now do, alongside suspend, email change and force logout, which ended them separately. A device session it ends releases its push subscription. Signing out after a password change answers 200 and keeps the reason. FR-013, FR-014, FR-016, the session traceability row and minimum verification are updated, the MRD reconciliation names v2.107, and the change log's header row and v1.23 entry, which the v1.23 source carried in header form, are restored. MRD v2.107 checked; its Module 3 entry and capability count are unchanged. Backend evidence only; nothing here is deployed.",
    )
    finish(doc, output, title)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[3])
    args = parser.parse_args()
    root = args.root.resolve()
    frd_dir = root / "docs/frd/functional-requirements/03-identity-team-and-organogram"
    patch_frd(
        frd_dir / f"XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document_v{FRD_SOURCE_VERSION}.docx",
        frd_dir / f"XVS_M03_Identity_Team_and_Organogram_Functional_Requirements_Document_v{FRD_TARGET_VERSION}.docx",
    )


if __name__ == "__main__":
    main()
