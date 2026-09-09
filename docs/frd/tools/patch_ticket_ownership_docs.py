#!/usr/bin/env python3
"""Bring the M31 FRD up to the desk the code now runs.

Two things happened to Module 31 that its FRD never recorded. The first landed
earlier and was never written down: a school triages its own tickets and sends
up what it cannot solve, so the platform desk stopped spanning every tenant's
tickets and began spanning CodeX's own, the escalated ones, and whatever a
support user is personally on. v1.4 still says platform support may "read and
work any tenant's tickets", which has not been true since escalation shipped.

The second is the reason for this revision. Assignment never asked whose ticket
it was working on. Every assignee is CodeX support staff and an assignee is a
participant, so naming one admitted the desk to a school's thread and its
internal notes - with no escalation audit entry, no notification to the desk,
and on ``tickets.ticket.assign``, a key a school could grant itself, while
escalating takes ``tickets.ticket.manage``. The cheaper key did the bigger job.
Naming an owner is now the desk's alone, a platform owner may only be named on
a ticket the desk may have, and the key is platform-scoped.

    python tools/patch_ticket_ownership_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document

from generate_requirements_documents import (
    BLUE,
    assert_no_em_dash,
    rebuild_table,
    shrink_inherited_media,
    update_extended_title,
    write_cell,
)

REVIEW_DATE = "9 September 2026"
CODE_BASELINE = "a85fda4 (9 September 2026)"
MRD_SOURCE, MRD_TARGET = "2.74", "2.75"
MRD_VERSION = MRD_TARGET
M31_SOURCE, M31_TARGET = "1.4", "1.5"

STEM = "XVS_M31_Support_Tickets_Functional_Requirements_Document"
TITLE = f"XVS M31 Support Tickets Functional Requirements Document v{M31_TARGET}"

PURPOSE = (
    "Module 31 is the support desk and the privacy boundary for the how-to "
    "editorial loop: how a user asks for help, who triages it, how support "
    "works the request, who may read a conversation, and which dimensionless "
    "guide signals platform editors may review. Visibility remains the "
    "organising control, and escalation is what moves it: a school's ticket is "
    "the school's own until they send it up, so the platform desk spans what "
    "was escalated to it rather than everything every tenant files. Ticket "
    "participation never widens access, and guide events carry no tenant, "
    "actor, business-record, or arbitrary metadata dimension."
)

BOUNDARIES = (
    "Filing a ticket needs no permission at all, because a user locked out of "
    "support by a missing grant has no way to ask for the grant. Working one "
    "does: escalation, status changes and the audit trail each sit behind "
    "their own key. Naming an owner is not a key a school can hold at all, "
    "because every owner is CodeX support staff and choosing which of them "
    "works a ticket is CodeX's rota."
)

SCOPE_ROWS = [
    ["Ticket intake",
     "Filing, editing the editable fields, and the derived number, tenant, "
     "branch and source."],
    ["Visibility",
     "Who may list, read, and search - participants, tenant managers, and "
     "platform support on what was escalated to it."],
    ["Escalation",
     "The school's decision that a ticket is beyond them, which is what gives "
     "CodeX the ticket."],
    ["Conversation",
     "Public replies and internal notes, and the boundary between them."],
    ["Evidence",
     "Attachments on a ticket or a comment, validated before storage."],
    ["Ownership and lifecycle",
     "Ownership named by the platform desk from live RBAC eligibility, and "
     "guarded status transitions."],
    ["Accounting and reach",
     "The audit trail, the notifications, and the dashboard counters."],
    ["Guide editorial signals",
     "Closed how-to event intake, sanitized no-result searches, "
     "platform-health aggregates, scoped volume control, and disposable "
     "retention."],
]

ACTOR_ROWS = [
    ["Any signed-in user",
     "File a ticket, read and comment on their own threads.",
     "No key. Participation is the grant."],
    ["Tenant ticket manager",
     "Read, work and comment on the tickets of their own tenant that fall "
     "inside the branches their role grants reach, and send up what the "
     "school cannot solve.",
     "tickets.ticket.manage, inside their own tenant only"],
    ["Platform support",
     "The cross-tenant support console: read and work CodeX's own tickets, "
     "the tickets a school has escalated, and any ticket they are personally "
     "on.",
     "tickets.ticket.manage held on the PLATFORM tenant"],
    ["Platform assigner",
     "Set or clear a ticket's owner from the eligible set. A school never "
     "names an owner, whatever keys it holds.",
     "tickets.ticket.assign, a platform-scoped key no tenant role may hold"],
    ["History reader",
     "Read a ticket's operational audit trail.",
     "tickets.audit.view"],
    ["Active user",
     "Emit a validated guide, walkthrough, or settled no-result-search event.",
     "Active authentication, closed fields, and the guide_analytics throttle "
     "scope"],
    ["Platform guide editor",
     "Read aggregate guide engagement, search gaps, and walkthrough exits.",
     "platform.health.view held on the PLATFORM tenant"],
]

FR002_EVIDENCE = (
    "Listing resolves through one visibility queryset. Participants - "
    "requester and assignee - always see their own threads, whatever branch "
    "the ticket carries. A same-tenant holder of the manage key sees that "
    "tenant's tickets narrowed to the branches their role grants reach, and a "
    "ticket filed for the school as a whole is visible to all of them. "
    "Platform support spans three things rather than every tenant: CodeX's own "
    "tickets, the tickets a school has escalated, and anything a support user "
    "is personally on."
)

FR002_ACCEPTANCE = (
    "The tenant-kind check on that span is load-bearing: a school user granted "
    "manage administers their own tenant only and never inherits cross-tenant "
    "reach. The branch narrowing is applied to the manager arm alone and never "
    "to the participant arm, because a ticket carries the branch it was filed "
    "for rather than the branches of everyone on it. The object read runs the "
    "same arms as the list and answers 404 rather than 403, so a hidden ticket "
    "cannot be distinguished from a missing one and no reader can open by id "
    "what their own list omits."
)

FR005_REQUIREMENT = (
    "A queue where nothing is owned is a queue nobody works, and the person "
    "who owns a ticket has to be somebody who can work it."
)

FR005_EVIDENCE = (
    "Assignment sets the assignee and synchronises status, and clearing it "
    "releases the ticket. Eligible assignees are resolved from live RBAC "
    "rather than a list: active platform staff whose effective roles grant the "
    "manage key, with explicit direct denies winning over role and group "
    "grants. Naming an owner is the platform desk's own decision and is "
    "refused for any caller outside the platform tenant. An owner may only be "
    "named on a ticket the desk may have: CodeX's own, or a school's after "
    "that school escalated it."
)

FR005_ACCEPTANCE = (
    "The picker and the write answer through the same predicate, so a list can "
    "never offer a name the service would then refuse, and a school is never "
    "shown CodeX's staff on a ticket it has raised nothing with. Gating the "
    "write on the ticket's own state closes a route that reached further than "
    "escalation while asking less of the caller: an assignee is a participant, "
    "so naming one admitted the desk to the thread and its internal notes "
    "without the escalation's audit entry or the notification that tells the "
    "desk it has gained work. Clearing an owner is never gated on that state, "
    "because the owner already has the ticket and refusing would strand it "
    "with them."
)

FR005_LIMIT = (
    "An assignment made before this rule keeps its owner. Such a ticket is not "
    "retro-escalated, because an escalation nobody performed would be a false "
    "entry in the audit trail, so its owner keeps the thread and cannot hand "
    "it on until the school sends it up."
)

FR013_ROWS = [
    ["Requirement",
     "A school's own business - the projector in Room 3, a timetable clash - "
     "is not CodeX's to read, and a desk that receives all of it stops being "
     "read, taking the tickets that were meant for CodeX down with the rest. "
     "The school has to be able to decide what goes up."],
    ["Current evidence",
     "Escalation stamps the ticket rather than opening a second one: the "
     "escalation time and the person who decided it are recorded on the same "
     "row, so the reference the requester was given still resolves, the thread "
     "they have been reading still grows, and CodeX arrives to the whole "
     "history instead of a summary somebody retyped. It takes the manage key, "
     "because deciding the school is beaten by a thing is a triage decision. "
     "An optional note is posted as a public comment, so the person who raised "
     "the ticket can see that their school passed it on and why."],
    ["Acceptance",
     "The desk is notified at the moment of escalation rather than on the next "
     "comment, because the school escalated precisely because it was waiting. "
     "Before that moment the new-ticket notice goes to the school's own triage "
     "staff instead, which matters because recipients who cannot see a ticket "
     "are dropped: paging the platform desk about an unescalated ticket would "
     "send nothing at all and leave it unread with nobody aware it exists. "
     "Escalating twice is refused rather than repeated, and one of CodeX's own "
     "tickets cannot be escalated to CodeX."],
    ["Current limit",
     "There is no withdrawal. A ticket sent up stays up, and the visibility "
     "rule deliberately keeps a support user on any ticket they own so that a "
     "future withdrawal could not take a thread away mid-conversation."],
]

LIFECYCLE_ROWS = [
    ["1", "Filed.",
     "OPEN. Requester, tenant, branch and source derived from the actor. The "
     "school's own triage staff are told."],
    ["2", "Sent up to CodeX, if the school cannot solve it.",
     "Status unchanged. The escalation time and its decider are stamped on the "
     "same row, the desk is notified, and the thread is unbroken."],
    ["3", "Assigned.",
     "ASSIGNED. Owner named by the desk from the eligible set; clearing it "
     "releases the ticket."],
    ["4", "Picked up.", "IN_PROGRESS."],
    ["5", "Resolved.",
     "RESOLVED, with the timestamp recorded. Leaves every personal workload "
     "counter."],
    ["6", "Closed, or reopened.",
     "CLOSED, or back to IN_PROGRESS from either RESOLVED or CLOSED."],
]

DESK_ROUTE_ROWS = [
    ["GET, POST /support/tickets/", "List with filters, or file a new ticket."],
    ["GET, PATCH /support/tickets/{id}/",
     "Read the thread, or edit its editable fields."],
    ["POST /support/tickets/{id}/escalate/",
     "Send a school's ticket up to CodeX, with an optional public note. Needs "
     "the manage key."],
    ["POST /support/tickets/{id}/assign/",
     "Set or clear the owner. Platform staff only, on the assign key."],
    ["GET /support/tickets/{id}/eligible-assignees/",
     "Who may take it, resolved live from RBAC. Empty on a ticket the desk may "
     "not own."],
    ["POST /support/tickets/{id}/transition/",
     "Move the status. Needs the manage key."],
]

ERROR_ROWS = [
    ["Reading a ticket the caller may not see",
     "404, identical to a ticket that does not exist."],
    ["A status move outside the transition map", "Refused."],
    ["An unknown or malformed context field at creation",
     "Refused, naming the field."],
    ["A route pattern containing digits, a query string or a fragment",
     "Refused."],
    ["An unsupported, empty, oversized or mislabelled attachment",
     "Refused with a field error."],
    ["A school user setting or clearing an owner",
     "Refused. Naming the owner is the platform desk's."],
    ["Assigning a school's ticket before that school has escalated it",
     "Refused on the assignee field, naming escalation as what is missing."],
    ["Assigning to somebody outside the eligible set", "Refused."],
    ["Escalating a ticket already with CodeX, or one of CodeX's own",
     "Refused."],
    ["A comment id from another ticket", "Refused."],
    ["An unknown analytics field, event combination, id, or live route",
     "Refused before persistence."],
    ["More than the scoped event limit for one authenticated user",
     "429. No event row is created."],
    ["A summary window outside 1 to 365 days",
     "Refused with a typed days error."],
]

SEEDING_CONTRACT = (
    "seed_ticket_permissions registers the keys, grants them all to the "
    "platform roles, and attaches the request-side defaults to the school "
    "roles. tickets.ticket.assign is registered platform-scoped, so a school's "
    "roles screen never offers a key that would govern nothing there."
)

RESOLVED_BANNER = (
    "RESOLVED - ASSIGNMENT IS NO LONGER A QUIETER ESCALATION\n"
    "• Assignment asked who the assignee was and never whose ticket it "
    "was. Every assignee is CodeX support staff and an assignee is a "
    "participant, so naming one admitted the platform desk to a school's "
    "thread, its internal notes and its files - without the escalation audit "
    "entry, without the notification that tells the desk it has gained work, "
    "and on tickets.ticket.assign, which a school could grant itself, while "
    "escalating takes tickets.ticket.manage. Naming an owner is now the "
    "platform desk's alone, an owner may only be named on a ticket the desk "
    "may have, the assignee picker answers through the same rule as the write, "
    "and the key is platform-scoped with a migration taking back any grant a "
    "school was given. Two authority checks that answered for any support user "
    "on any ticket - managing, and reading or writing internal notes - now "
    "stop at the same visibility boundary, so the rule no longer rests on the "
    "viewset queryset alone."
)

TRACEABILITY_ROWS = [
    ["Ticket creation, retrieval, and update", "FR-001, FR-002"],
    ["Support queue and dashboard", "FR-007"],
    ["Search, filters, and pagination", "FR-002, FR-007"],
    ["Ticket comments", "FR-003"],
    ["Ticket attachments", "FR-004"],
    ["Ticket files stay on their own visibility-checked route", "FR-004"],
    ["Assignee and ownership controls", "FR-005"],
    ["Guarded status transitions", "FR-006"],
    ["Tenant and platform access scoping", "FR-002, FR-013"],
    ["Ticket audit history", "FR-009"],
    ["Ticket notifications", "FR-010, FR-013"],
    ["Privacy-safe guide and route context", "FR-008"],
    ["Strict creation-time context allowlist", "FR-008"],
    ["Module-registered ticket context vocabularies", "FR-008"],
    ["Privacy-safe guide engagement and search-gap event intake", "FR-011"],
    ["Platform guide-editor aggregates, scoped ingestion limits, and 180-day "
     "retention", "FR-012"],
    ["School triage and escalation to CodeX", "FR-013"],
]

TRACEABILITY_INTRO = (
    f"Module 31 carries 17 capability entries in MRD v{MRD_VERSION}. Each maps "
    "to the requirements below."
)

CHANGE_SUMMARY = (
    "Records the escalation boundary and takes assignment out of a school's "
    "hands. Escalation shipped without reaching this document: v1.4 still said "
    "platform support may read and work any tenant's tickets, when a school's "
    "ticket has been the school's own until they send it up, and the desk's "
    "span has been CodeX's own tickets, the escalated ones, and whatever a "
    "support user is personally on. FR-013 now states that rule, its guards "
    "and its notification, and FR-002 and the actor table are corrected to "
    "match. The revision's own change is FR-005. Assignment never asked whose "
    "ticket it was working on, and because every assignee is CodeX support "
    "staff and an assignee is a participant, naming one admitted the desk to a "
    "school's thread and its internal notes with none of escalation's audit "
    "entry, notification or triage grant - on a weaker key a school could "
    "grant itself. Naming an owner is now the platform desk's alone and only "
    "on a ticket the desk may have; tickets.ticket.assign is platform-scoped "
    "with a migration taking back any school grant; and the manage and "
    "internal-note authority checks stop at the same visibility boundary as "
    "the list. The lifecycle, the routes, the typed errors, the seeding "
    "contract and the traceability follow. No capability was added, and the "
    "traceability now carries all 17 MRD entries rather than 15: escalation is "
    "named as a capability of its own in MRD v2.75, and the entry for ticket "
    "files on their own visibility-checked route was missing."
)


def replace_cell(cell, text, **kwargs):
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    write_cell(cell, text, **kwargs)


def replace_cover_version(table, source, target):
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if source in run.text:
                run.text = run.text.replace(source, target)


def retitle(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def prepend_change_log(table, version, date, summary, *, size=8):
    template = table.rows[1]
    new_tr = copy.deepcopy(template._tr)
    template._tr.addprevious(new_tr)
    row = table.rows[1]
    replace_cell(row.cells[0], version, size=size)
    replace_cell(row.cells[1], date, size=size)
    replace_cell(row.cells[2], summary, size=size)


def fr_cells(table):
    """The label/value rows of an FR table, skipping its merged banner row."""
    return [(row.cells[0], row.cells[1]) for row in table.rows[1:]]


def set_fr_row(table, label, text, *, size=8.2):
    for name_cell, value_cell in fr_cells(table):
        if name_cell.text.strip() == label:
            replace_cell(value_cell, text, size=size)
            return True
    return False


def build_fr013(doc, template_table, heading_paragraph):
    """Add FR-013 by copying an existing FR table, so it inherits the styling.

    Placed at the end of Section 4 and numbered after the last requirement,
    because inserting it beside the ticket requirements it belongs with would
    renumber FR-006 through FR-012 and every traceability entry that names them.
    """
    new_heading = copy.deepcopy(heading_paragraph._p)
    heading_paragraph._p.addprevious(new_heading)
    from docx.text.paragraph import Paragraph

    paragraph = Paragraph(new_heading, heading_paragraph._parent)
    retitle(paragraph, "FR-013 Send a Ticket Up to CodeX")
    paragraph.style = doc.styles["Heading 2"]

    new_tbl = copy.deepcopy(template_table._tbl)
    new_heading.addnext(new_tbl)
    from docx.table import Table

    table = Table(new_tbl, template_table._parent)
    banner = table.rows[0].cells[0]
    replace_cell(banner, "FR-013 | Implemented", size=8.5, bold=True)
    for label, text in FR013_ROWS:
        if not set_fr_row(table, label, text):
            raise SystemExit(f"FR-013: no {label!r} row in the template")
    return table


def patch_frd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    doc.core_properties.title = TITLE
    doc.core_properties.version = M31_TARGET

    # Bound by reference before anything moves. Adding FR-013 and dropping a
    # resolved banner both shift every index after them, so indices read later
    # would address the wrong table and quietly rewrite it.
    control = doc.tables[1]
    scope_areas = doc.tables[3]
    actors = doc.tables[6]
    fr002, fr005, fr012 = doc.tables[8], doc.tables[11], doc.tables[18]
    lifecycle = doc.tables[19]
    desk_routes = doc.tables[22]
    typed_errors = doc.tables[24]
    dependencies = doc.tables[25]
    resolved_banner, stale_banner = doc.tables[26], doc.tables[27]
    traceability, change_log = doc.tables[29], doc.tables[30]

    replace_cover_version(doc.tables[0], M31_SOURCE, M31_TARGET)

    for row in control.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], M31_TARGET, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Code baseline":
            replace_cell(row.cells[1], CODE_BASELINE, size=9)
        elif label == "Source MRD":
            replace_cell(
                row.cells[1],
                f"XVS Module Requirements Document v{MRD_VERSION} | Module 31",
                size=9,
            )

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Module 31 is the support desk"):
            retitle(paragraph, PURPOSE)
        elif text.startswith("Filing a ticket needs no permission"):
            retitle(paragraph, BOUNDARIES)
        elif text.startswith("Module 31 carries"):
            retitle(paragraph, TRACEABILITY_INTRO)

    rebuild_table(scope_areas, ["Area", "Responsibility"], SCOPE_ROWS,
                  [1.6, 5.45])
    rebuild_table(actors, ["Actor", "May do", "Governed by"], ACTOR_ROWS,
                  [1.35, 3.3, 2.4])

    if not set_fr_row(fr002, "Current evidence", FR002_EVIDENCE):
        raise SystemExit("FR-002: evidence row not found")
    set_fr_row(fr002, "Acceptance", FR002_ACCEPTANCE)

    if not set_fr_row(fr005, "Requirement", FR005_REQUIREMENT):
        raise SystemExit("FR-005: requirement row not found")
    set_fr_row(fr005, "Current evidence", FR005_EVIDENCE)
    set_fr_row(fr005, "Acceptance", FR005_ACCEPTANCE)
    set_fr_row(fr005, "Current limit", FR005_LIMIT)

    workflows = next(
        p for p in doc.paragraphs
        if p.text.strip().startswith("5. Workflows and Lifecycle Rules")
    )
    build_fr013(doc, fr012, workflows)

    rebuild_table(lifecycle, ["Step", "Document", "Ledger effect"],
                  LIFECYCLE_ROWS, [0.55, 2.15, 4.35])
    rebuild_table(desk_routes, ["Method and path", "Purpose"],
                  DESK_ROUTE_ROWS, [3.05, 4.0])
    rebuild_table(typed_errors, ["Condition or route", "Answer"],
                  ERROR_ROWS, [3.85, 3.2])

    for row in dependencies.rows:
        if row.cells[0].text.strip() == "Seeding":
            replace_cell(row.cells[1], SEEDING_CONTRACT, size=8.2)

    # Section 9 is current state, not history. Both banners here record
    # defects closed in earlier revisions, and the change log holds them.
    replace_cell(resolved_banner.rows[0].cells[0], RESOLVED_BANNER, size=8.2)
    stale_banner._tbl.getparent().remove(stale_banner._tbl)

    rebuild_table(traceability, ["MRD capability", "Requirements"],
                  TRACEABILITY_ROWS, [4.95, 2.1])

    prepend_change_log(
        change_log, M31_TARGET, REVIEW_DATE, CHANGE_SUMMARY, size=8.2,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, TITLE)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


# ── the MRD ────────────────────────────────────────────────────────────────

MRD_SOURCE_SCOPE = (
    "Module 31's escalation boundary recorded, ticket ownership narrowed to "
    "the platform desk, and tickets.ticket.assign reclassified platform-scoped "
    "(9 September 2026)"
)

MRD_INTRO = (
    "This revision records a boundary Module 31 has enforced since escalation "
    "shipped and never wrote down, and narrows one it was not enforcing at "
    "all. A school's ticket is the school's own until they send it up, and "
    "naming its owner is CodeX's, on a key no school may hold."
)

MRD_CONTENTS_NOTE = "Module 31's escalation boundary and who names an owner"

M31_BLURB = (
    "Tenant-aware support queue with comments, files, school triage and "
    "escalation to CodeX, platform-owned assignment, controlled status "
    "changes, notifications, audit, privacy-safe guide-to-support context, and "
    "dimensionless how-to analytics for platform editorial review."
)

M31_NEW_CAPABILITY = "\u25b8  School triage and escalation to CodeX"

M31_DECISION = (
    "CURRENT DECISION\n"
    "\u2022 Module 31 remains Backend Complete and In use Complete with "
    "seventeen capability entries, reconciled to M31 FRD v1.5.\n"
    "\u2022 A school's ticket is the school's own until the school escalates "
    "it. The platform desk spans CodeX's own tickets, the ones a school sent "
    "up, and anything a support user is personally on, rather than everything "
    "every tenant files.\n"
    "\u2022 Naming a ticket's owner is the platform desk's. Every assignee is "
    "CodeX support staff, so choosing among them is CodeX's rota and "
    "tickets.ticket.assign is platform-scoped.\n"
    "\u2022 No material capability gap was identified within this module's "
    "stated backend scope."
)

MRD_DELTA_ROWS = [
    ["Module 31 capabilities", "16 to 17",
     "School triage and escalation to CodeX becomes an entry of its own. The "
     "behaviour shipped earlier and no capability described it, so this "
     "records what exists rather than counting anything new."],
    ["Module 31 ticket ownership", "Narrowed to the platform desk",
     "Assignment asked who the assignee was and never whose ticket it was. "
     "Every assignee is CodeX support staff and an assignee is a participant, "
     "so naming one admitted the desk to a school's thread and its internal "
     "notes with none of escalation's audit entry or notification, on a weaker "
     "key than escalating asks for."],
    ["tickets.ticket.assign", "Reclassified platform-scoped",
     "A school holding it can now act on nothing, so its roles screen no "
     "longer offers it. A migration takes back any grant a school was given, "
     "because the scope guard refuses a platform key on a tenant role and the "
     "rows would leave those roles unsaveable."],
    ["Module 31 status", "Unchanged at Complete and Complete",
     "The desk works as documented. This revision records a boundary and "
     "narrows an authority; it neither adds nor removes a capability."],
    ["M31 FRD", "Revised to v1.5",
     "Adds FR-013 for escalation, rewrites FR-005 for ownership, corrects the "
     "platform span in FR-002 and the actor table, and reconciles the "
     "traceability to all seventeen entries."],
]

MRD_CHANGE_SUMMARY = (
    "Records Module 31's escalation boundary and takes ticket ownership out of "
    "a school's hands. Escalation shipped without reaching either document: a "
    "school's ticket is the school's own until they send it up, and the "
    "platform desk spans CodeX's own tickets, the escalated ones, and whatever "
    "a support user is personally on rather than every ticket every tenant "
    "files. That is now a capability entry of its own, which takes Module 31 "
    "from 16 to 17 and the platform to 493 across 31 modules; nothing new was "
    "built for it. The change this revision carries is ownership. Assignment "
    "asked who the assignee was and never whose ticket it was, and because "
    "every assignee is CodeX support staff and an assignee is a participant, "
    "naming one admitted the desk to a school's thread and its internal notes "
    "with none of escalation's audit entry, notification or triage grant, on "
    "tickets.ticket.assign while escalating takes tickets.ticket.manage. "
    "Naming an owner is now the platform desk's alone and only on a ticket the "
    "desk may have; the key is platform-scoped with a migration taking back "
    "any school grant; and the manage and internal-note authority checks stop "
    "at the same visibility boundary as the ticket list. Module 31 stays "
    "Backend Complete and In use Complete, and M31 FRD advances to v1.5."
)


def keep_rows_whole(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))


def add_bullet(cell, text):
    """Append a bullet by cloning one already in the cell.

    The capability grid's bullets carry their own run formatting, and rewriting
    the cell to add a line would flatten it to the default. Cloning keeps the
    new entry indistinguishable from the sixteen beside it.
    """
    from docx.text.paragraph import Paragraph

    template = cell.paragraphs[0]
    clone = copy.deepcopy(template._p)
    template._p.addnext(clone)
    retitle(Paragraph(clone, cell), text)


def patch_mrd(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = f"XVS Module Requirements Document v{MRD_TARGET}"
    doc.core_properties.title = title
    doc.core_properties.version = MRD_TARGET

    control = doc.tables[1]
    contents = doc.tables[2]
    module_index = doc.tables[5]
    m31_capabilities = doc.tables[74]
    delta = doc.tables[76]
    change_log = doc.tables[78]

    replace_cover_version(doc.tables[0], MRD_SOURCE, MRD_TARGET)

    for row in control.rows:
        label = row.cells[0].text.strip()
        if label == "Version":
            replace_cell(row.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(row.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(row.cells[1], MRD_SOURCE_SCOPE, size=9)
        elif label == "Capability entries":
            replace_cell(row.cells[1], "493", size=9)

    for row in contents.rows:
        if row.cells[0].text.strip().startswith("5."):
            replace_cell(
                row.cells[0], f"5. v{MRD_TARGET} Capability Delta",
                size=9, bold=True, color=BLUE,
            )
            replace_cell(row.cells[1], MRD_CONTENTS_NOTE, size=9)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == f"5. v{MRD_SOURCE} Capability Delta":
            retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")
        elif text.startswith("This revision") and len(text) > 80:
            retitle(paragraph, MRD_INTRO)
        elif text.startswith("Tenant-aware support queue"):
            retitle(paragraph, M31_BLURB)

    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for row in module_index.rows:
        if row.cells[0].text.strip() == "31":
            replace_cell(
                row.cells[5], "17", size=7,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
            )

    # Escalation sits with the access rule it changes, not at the end.
    scoping = next(
        cell for row in m31_capabilities.rows for cell in row.cells
        if cell.text.strip() == "▸  Tenant and platform access scoping"
    )
    add_bullet(scoping, M31_NEW_CAPABILITY)
    replace_cell(
        m31_capabilities.rows[-1].cells[0], M31_DECISION, size=8.2,
    )

    rebuild_table(
        delta,
        [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
        MRD_DELTA_ROWS,
        [1.85, 1.35, 3.85],
    )
    keep_rows_whole(delta)

    prepend_change_log(change_log, MRD_TARGET, "9 Sep 2026", MRD_CHANGE_SUMMARY)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    root = Path(args.root)

    mrd = root / "module-requirements"
    patch_mrd(
        mrd / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx",
        mrd / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
    )
    print(f"Wrote MRD v{MRD_TARGET}")

    folder = root / "functional-requirements" / "31-support-tickets"
    patch_frd(
        folder / f"{STEM}_v{M31_SOURCE}.docx",
        folder / f"{STEM}_v{M31_TARGET}.docx",
    )
    print(f"Wrote M31 FRD v{M31_TARGET}")


if __name__ == "__main__":
    main()
