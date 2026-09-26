#!/usr/bin/env python3
"""Cut MRD v2.92 and M01 FRD v1.28: a field-level refusal says why in its message.

What changed in the backend, and therefore in the documents:

* The shared exception handler builds a DRF refusal's ``message`` from the
  refusal itself. One message is shown alone; several are each prefixed with
  their dotted field path, five at most and then "and N more"; a refusal that
  sets its own ``detail`` keeps it as the headline, so a sibling machine code
  never becomes copy. "An error occurred. Check the error details for more
  information." is left only for a body with no message in it. Before, every
  serializer and every upload refusal answered with that generic line and put
  the reason only under ``error.detail``, which the clients do not show.
* ``core.uploads.validate_upload`` words two refusals so they say what to do:
  a file whose bytes do not match its extension is told it is not really the
  type its name claims, and a ``.jfif`` (a JPEG under another name) is told to
  rename it to ``.jpg``.

The visible case was a school uploading its logo during onboarding: a JPEG
named crest.png was refused with the generic line and the admin never learned
why.

M01's traceability paragraph is corrected on the way: it named MRD v2.77 and
twenty-three entries while its table and the MRD carry twenty-four.

M09 names the logo only in its notification templates. M19 describes the
upload check's three tests and quotes no message. M03's activation refusal and
M11's refusal table name codes that this does not change. All checked and left.

Each document starts from the newest version in its folder and is written at
the next free number; the script refuses to run if a newer version has
appeared, and ``finish`` refuses to overwrite one.

    python tools/patch_field_error_message_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    finish,
    frd_path,
    keep_format,
    log_change,
    replace_cell,
    set_control,
    set_cover_version,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    row_labelled,
)
from patch_staff_id_and_auth_events_docs import (
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
)

import patch_record_history_docs

REVIEW_DATE, SHORT_DATE = "26 September 2026", "26 Sep 2026"
MRD_SOURCE, MRD_TARGET = "2.91", "2.92"
CODE_BASELINE = (
    "Backend main at 65a31144, which carries the field-level error message change "
    "(ac42a93d), 26 September 2026. The school app's logo field reads the message and is "
    "committed; it is pending release"
)
TEST_EVIDENCE = (
    "Verified by the full core (171 tests), vs_config (117), vs_exports (185), vs_tickets "
    "(114), vs_procurement (592), schools.vs_schools (377), vs_user (421) and vs_finance (872) "
    "suites, slow tests included, each app run on its own and all passing, and by uploading a "
    "mislabelled logo and a .jfif on a pending school's onboarding profile in the running "
    "school app. Nothing here claims deployment."
)

# log_change writes the module's own review date, read from this module.
patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


# ── M01 School & Branch Management ───────────────────────────────────────────

M01_DIR = "01-school-and-branch-management"
M01_STEM = "XVS_M01_School_and_Branch_Management_Functional_Requirements_Document"
M01_SOURCE, M01_TARGET = "1.27", "1.28"

M01_TRACE = (
    f"This table reconciles every Module 1 capability entry in MRD v{MRD_TARGET} to the "
    "controlling functional requirement in this FRD. All twenty-four entries are represented. "
    "Two are restated rather than added: the roles CodeX ships now records that a School is "
    "refused rather than created short and that a Branch role follows its Branch, and the "
    "role-grant entry covers every way a School's reach shrinks rather than a tier move "
    "alone. A refused logo saying why in the response message strengthens the branding and "
    "logo entry without changing the count."
)

M01_SUMMARY = (
    "Minor revision. A refused logo upload says why in the response message, not only under "
    "error.detail: a JPEG named crest.png is told it is not really a PNG and to save it again "
    "as one, and a .jfif is told to rename it to .jpg. The cause was platform-wide: the shared "
    "exception handler answered every field-level refusal with \"An error occurred. Check the "
    "error details for more information.\", and it now builds the message from the refusal "
    "(MRD current architecture). FR-009's acceptance and the logo route are updated, and the "
    "traceability paragraph now names MRD v2.92 and twenty-four entries, where it had named "
    f"MRD v2.77 and twenty-three. MRD v{MRD_TARGET}. " + TEST_EVIDENCE
)


def patch_m01() -> None:
    require_newest(str(ROOT / "functional-requirements" / M01_DIR / f"{M01_STEM}_v*.docx"),
                   M01_SOURCE)
    doc = Document(str(frd_path(M01_DIR, M01_STEM, M01_SOURCE)))
    set_cover_version(doc, M01_SOURCE, M01_TARGET)
    set_control(doc, "Version", M01_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET} | Module 1")

    append_to(fr_table(doc, "FR-009"), "Acceptance",
              " A refusal's reason is the response message itself, which the School's screen "
              "shows: a JPEG named crest.png is told it is not really a PNG and to save it "
              "again as one, and a .jfif is told to rename it to .jpg "
              "(SchoolLogoEndpointTests.test_a_mislabelled_image_says_why_in_the_message and "
              "test_a_jfif_is_told_to_rename_rather_than_refused_as_the_wrong_type).")

    route = row_labelled(doc, "POST, DELETE")
    cells = list({id(c._tc): c for c in route.cells}.values())
    if cells[1].text.strip() != "/v1/i/me/profile/logo/":
        raise ValueError("The POST, DELETE row is not the logo route")
    keep_format(cells[-1], cells[-1].text.rstrip() + (
        ". A refused file answers 400 with the reason as the message: not an image, over 2MB, "
        "not really the type its name claims, or a .jfif to rename"))

    mrd_tools.retitle(paragraph_starting(doc, "This table reconciles every Module 1"),
                      M01_TRACE)

    log_change(doc, M01_TARGET, M01_SUMMARY)
    assert_absent_outside_log(doc, "MRD v2.77", "twenty-three entries")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(M01_DIR, M01_STEM, M01_TARGET),
           f"{M01_STEM.replace('_', ' ')} v{M01_TARGET}", M01_TARGET)


# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_ARCHITECTURE = (
    "• A refusal says why in its message. When the server turns a request down field by "
    "field, the message is that field's own sentence, or each failing field named in turn, "
    "rather than a generic line pointing at the error detail; a refusal that sets its own "
    "detail keeps it. The generic line remains only for a refusal that carries no sentence at "
    "all."
)
MRD_CONTENTS_NOTE = "A field-level refusal says why in its message"
MRD_INTRO = (
    "This revision records that a refusal from any module says why in its message, and that "
    "a refused upload says what to do about the file."
)
MRD_CHANGE_SUMMARY = (
    "A field-level refusal says why in its message. The shared exception handler had answered "
    "every serializer and upload refusal with \"An error occurred. Check the error details for "
    "more information.\" and kept the reason under the error detail, which clients do not "
    "show; a school uploading a JPEG named crest.png during onboarding never learned why it "
    "was refused. The message is now built from the refusal, and the upload check tells a "
    "mislabelled file it is not really its claimed type and a .jfif to rename it. Current "
    "architecture gains the rule. Capability entries unchanged; statuses do not move. M01 "
    "v1.28. Backend evidence; the school app's logo field is committed and pending release; "
    "no deployment claim."
)
MRD_DELTA_ROWS = [
    ["Refusal messages", "Say why, on every module",
     "A field-level refusal's message is the field's own sentence, or each failing field named "
     "in turn, five at most; a refusal with its own detail keeps it."],
    ["Generic error line", "Kept for empty refusals only",
     "\"An error occurred. Check the error details\" answers only a refusal that carries no "
     "sentence at all."],
    ["Upload refusals", "Say what to do",
     "A file whose bytes do not match its name is told it is not really that type; a .jfif is "
     "told to rename it to .jpg."],
    ["Module FRDs", "One revised",
     "M01 v1.28, whose traceability paragraph also moves to twenty-four entries. M09, M19, "
     "M03 and M11 checked and left."],
]


def insert_bullet_after(doc, start: str, text: str) -> None:
    """Add a line to a boxed bullet list, below the one starting with ``start``, as it is set."""
    hits = [(p, c) for t in doc.tables for r in t.rows for c in r.cells
            for p in c.paragraphs if p.text.strip().startswith(start)]
    unique = {id(p._p): (p, c) for p, c in hits}
    if len(unique) != 1:
        raise ValueError(f"{start!r} starts {len(unique)} boxed lines")
    paragraph, cell = next(iter(unique.values()))
    clone = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(clone)
    mrd_tools.retitle(Paragraph(clone, cell), text)


def table_whose_first_cell(doc, predicate):
    hits = [t for t in doc.tables if predicate(t.rows[0].cells[0].text.strip())]
    if len(hits) != 1:
        raise ValueError(f"{len(hits)} tables match")
    return hits[0]


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    require_newest(str(folder / "XVS_Module_Requirements_Document_v*.docx"), MRD_SOURCE)
    doc = Document(str(folder / f"XVS_Module_Requirements_Document_v{MRD_SOURCE}.docx"))
    tables = doc.tables
    cover, control, contents, index = tables[0], tables[1], tables[2], tables[5]
    delta = table_whose_first_cell(doc, lambda h: h.endswith("capability delta"))
    log = table_whose_first_cell(doc, lambda h: h == "Version")
    assert index.rows[0].cells[5].text.strip() == "Entries"
    total_before = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])

    mrd_tools.replace_cover_version(cover, MRD_SOURCE, MRD_TARGET)
    for r in control.rows:
        label = r.cells[0].text.strip()
        if label == "Version":
            replace_cell(r.cells[1], MRD_TARGET, size=9)
        elif label == "Review date":
            replace_cell(r.cells[1], REVIEW_DATE, size=9)
        elif label == "Source scope":
            replace_cell(r.cells[1], CODE_BASELINE, size=9)
    for r in contents.rows:
        if r.cells[0].text.strip().startswith("5."):
            keep_format(r.cells[0], f"5. v{MRD_TARGET} Capability Delta")
            keep_format(r.cells[1], MRD_CONTENTS_NOTE)
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if (text.startswith("5. ") and paragraph.style is not None
                and paragraph.style.name.startswith("Heading")):
            mrd_tools.retitle(paragraph, f"5. v{MRD_TARGET} Capability Delta")

    insert_bullet_after(doc, "• The shared success envelope returns an empty list", MRD_ARCHITECTURE)

    blurb1 = mrd_tools.module_blurbs(doc)[1]
    if "Documented by M01" in blurb1.text:
        raise ValueError("Module 1's description already names an FRD version")
    mrd_tools.retitle(blurb1, blurb1.text.rstrip() + (
        " A refused logo says why in the response message, so a school is told its file is "
        "not really the type its name claims. Documented by M01 FRD v1.28."))

    total_after = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total_after != total_before:
        raise ValueError(f"Capability total moved from {total_before} to {total_after}")

    mrd_tools.rebuild_table(delta, [f"v{MRD_TARGET} capability delta", "Decision", "Evidence"],
                            MRD_DELTA_ROWS, mrd_tools.DELTA_WIDTHS)
    mrd_tools.keep_rows_whole(delta)
    intro = [p for p in doc.paragraphs
             if p.text.strip().startswith("This revision") and len(p.text) > 80]
    if len(intro) != 1:
        raise ValueError(f"{len(intro)} delta introductions found")
    mrd_tools.retitle(intro[0], MRD_INTRO)

    # Copied from the latest version row, so the new row reads at the same size.
    log.rows[1]._tr.addprevious(copy.deepcopy(log.rows[1]._tr))
    for cell, text in zip(log.rows[1].cells, (MRD_TARGET, SHORT_DATE, MRD_CHANGE_SUMMARY)):
        keep_format(cell, text)
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, folder / f"XVS_Module_Requirements_Document_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m01()
    patch_mrd()


if __name__ == "__main__":
    main()
