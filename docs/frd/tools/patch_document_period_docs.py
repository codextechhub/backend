#!/usr/bin/env python3
"""Cut MRD v2.107, M22 FRD v1.19 and M23 FRD v1.19: a procurement list answers for one period.

What the backend does, and therefore what the documents record:

* Every procurement list over a finance document takes ``?fiscal_year=<year>``,
  one of the entity's own fiscal years named by its label, or an inclusive
  ``?date_from=`` and ``?date_to=`` with either end left open. The narrowing is
  ``vs_finance.period_filter.narrow_document_list``, applied in
  ``_ProcBase.paginate``, so requisitions, RFQs, quotations, purchase orders,
  goods receipts, vendor bills, vendor credit notes and vendor payments all take
  it, and a list of vendors, contracts or stock ignores it.
* Each document is placed by its own date (``document_date_field``), and a
  fiscal year by its own start and end, so a September-to-August year is not
  read as a calendar year. The period is ANDed with ``?branch=`` and every
  other filter.
* A year the books do not have, a year with a range, a malformed date and a
  range ending before it starts are refused with 400 naming the field. The year
  is looked up only in the entity the caller is already authorised for.
* A named period replaces the default that hides archived fiscal years, so an
  archived year can be listed by naming it; without a period a list is as it
  was, archived years hidden unless ``?include_archived=true``.
* The requisition, RFQ, purchase-order and bill summaries count the same period
  and branch as the list beneath them.
* The screen exports follow neither the period nor the branch: they report the
  parameters as filters they do not apply and cover every branch and year in
  the reader's reach. That is recorded as a current gap in both FRDs and in the
  MRD's Needs Attention for Modules 22 and 23.

The MRD advances because two modules' documented capability and known limits
change; no capability entry is added and no status moves. M24 was checked and
left: a goods return has no list of its own and stock lists are not narrowed.
M19 was checked and left: the period helper lives in vs_finance but no finance
route takes it. M25 and M26 were checked and left: neither records a
procurement list's parameters, and no export code changed.

v1.18 of both FRDs and MRD v2.106 wrote a version row over their change log's
header and kept a duplicate of the row beneath it. Each header is restored
from the newest earlier version that still carries one, and the duplicate is
removed. The MRD's newest version rows had also lost their cell borders, so
they are rebuilt from the log's ordinary ruled row.

Each document starts from the newest version in its folder and is written at
the next number; the script refuses to run if a newer version has appeared, and
``finish`` refuses to overwrite one.

    /var/tmp/cxvenv/bin/python docs/frd/tools/patch_document_period_docs.py
"""
from __future__ import annotations

import copy

from docx import Document
from docx.text.paragraph import Paragraph

import patch_mrd_v2_79_docs as mrd_tools
from patch_document_type_labels_docs import require_newest
from patch_record_history_docs import (
    ROOT,
    add_fr,
    finish,
    frd_path,
    keep_format,
    log_change,
    oldest_with_header,
    set_control,
    table_with_header,
)
from patch_restricted_grant_ladder_docs import (
    append_to,
    assert_absent_outside_log,
    edit_blurb,
    edit_box,
    insert_row_after,
    row_labelled,
)
from patch_staff_id_and_auth_events_docs import (
    edit_paragraph,
    fr_table,
    normalise_change_log,
    paragraph_starting,
    repair_ooxml,
    row,
)

import patch_record_history_docs

REVIEW_DATE = "17 October 2026"
MRD_SOURCE, MRD_TARGET = "2.106", "2.107"
CODE_BASELINE = (
    "Backend main at 35e234f with the document period change in its working tree, "
    "17 October 2026. The shared screens' integration is pending"
)
TEST_EVIDENCE = (
    "Verified by vs_procurement.tests_document_period (20 tests) and the whole vs_procurement "
    "app (999 tests, OK); the full repository suite was not run. Backend evidence only; the "
    "shared screens' integration is pending and nothing here claims deployment."
)

# log_change writes the module's own review date, read from this module.
patch_record_history_docs.REVIEW_DATE = REVIEW_DATE


# ── shared text ──────────────────────────────────────────────────────────────

PARAMETERS = (
    "?fiscal_year=<year>, one of the entity's own fiscal years named by its label (2015), or "
    "?date_from=YYYY-MM-DD and ?date_to=YYYY-MM-DD, an inclusive range with either end left "
    "open"
)

NARROWING = (
    "The narrowing is vs_finance.period_filter.narrow_document_list, applied by "
    "_ProcBase.paginate, so every procurement list over a finance document takes it and a list "
    "of anything else (vendors, contracts, stock) ignores the parameters. Each document is "
    "placed by its own date, read through vs_finance.retention.document_date_field. A fiscal "
    "year is bounded by its own start and end dates, so where Bright Star School runs "
    "September to August, fiscal_year=2015 covers 1 September 2014 to 31 August 2015 and not "
    "the calendar year. The period is ANDed with every other filter the list takes, ?branch= "
    "included: Mrs Bello asking for Ikeja Branch and 2015 sees Ikeja's 2015 documents and "
    "only those. A named period replaces the default that leaves documents dated in an "
    "archived fiscal year out (vs_finance.archive.hide_archived), because asking for an "
    "archived 2015 is asking to see it and the period itself bounds the list. With no period "
    "a list is unchanged: archived years stay hidden unless ?include_archived=true, and every "
    "year that is not archived is listed."
)

REFUSALS = (
    "Four conditions are refused with 400, each naming its field: a fiscal year the entity "
    "does not have (fiscal_year: There is no financial year 2009 in these books.), a year and "
    "a date range together (fiscal_year: Choose a financial year or a date range, not both.), "
    "a date not written YYYY-MM-DD (date_from or date_to: Give the date as YYYY-MM-DD.), and "
    "a range ending before it starts (date_to: The end date must be on or after the start "
    "date.). The year is looked up only in the entity the caller is already authorised for, "
    "so the parameter cannot reveal another school's years."
)

TYPED_ERRORS = [
    ["Naming a fiscal year these books do not have",
     "400 on fiscal_year: There is no financial year 2009 in these books. Another entity's "
     "years are never consulted."],
    ["Naming a fiscal year and a date range together",
     "400 on fiscal_year: Choose a financial year or a date range, not both."],
    ["Giving date_from or date_to in any form but YYYY-MM-DD",
     "400 on that field: Give the date as YYYY-MM-DD."],
    ["Giving a date range that ends before it starts",
     "400 on date_to: The end date must be on or after the start date."],
]

MODULE19_TAIL = (
    " Its fiscal years bound a list narrowed by ?fiscal_year=, read only in the caller's own "
    "entity, and its archive rule decides which years a list leaves out when no period is "
    "named ({fr})."
)

REQUIREMENT = (
    "A reader must be able to list one year's documents, or one stretch of dates, without "
    "paging through every year the books hold, and the figures above a list must count the "
    "rows the list shows."
)


def insert_rows_after(anchor, rows: list[list[str]]):
    """Clone ``anchor`` once per row, in order, each directly below the last."""
    for values in rows:
        anchor = insert_row_after(anchor, values)
    return anchor


def append_box_line(doc, heading: str, line: str) -> None:
    edit_box(table_with_header(doc, heading).rows[0].cells[0], [("append", line)])


def set_trace(doc, capability: str, requirements: str) -> None:
    trace = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "MRD capability")
    keep_format(row(trace, capability).cells[-1], requirements)


def start_frd(folder: str, stem: str, source: str, target: str, module: int):
    require_newest(str(ROOT / "functional-requirements" / folder / f"{stem}_v*.docx"), source)
    doc = Document(str(frd_path(folder, stem, source)))
    mrd_tools.replace_cover_version(doc.tables[0], source, target)
    set_control(doc, "Version", target)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Code baseline", CODE_BASELINE)
    set_control(doc, "Source MRD", f"XVS Module Requirements Document v{MRD_TARGET} | Module M{module}")
    return doc


def keep_attention_with_its_boxes(doc) -> None:
    """Keep Needs Attention's heading and lead-in on the page its first box starts on.

    The lead-in is a single line followed by a one-row table, so without this
    a page break can fall between them and leave the heading stranded at the
    foot of a page with its boxes on the next.
    """
    heading = paragraph_starting(doc, "9. Needs Attention")
    lead = paragraph_starting(doc, "These are current risks and gaps")
    heading.paragraph_format.keep_with_next = True
    lead.paragraph_format.keep_with_next = True


def rebuild_borderless_rows(doc) -> None:
    """Give every change-log version row the bordered look of the log's ordinary rows.

    Some version rows were written without cell borders, so they print as loose
    text between a ruled header and ruled rows. Each one is rebuilt from the
    first ordinary row that carries borders, keeping its text. A log whose rows
    all carry borders, or none do, is left alone.
    """
    from patch_record_history_docs import change_log_table, is_header_styled

    table = change_log_table(doc)
    rows = [r for r in table.rows[1:] if not is_header_styled(table, r)]
    bordered = next((r for r in rows if "tcBorders" in r._tr.xml), None)
    if bordered is None:
        return
    for r in rows:
        if "tcBorders" in r._tr.xml:
            continue
        texts = [c.text for c in r.cells]
        clone = copy.deepcopy(bordered._tr)
        r._tr.addprevious(clone)
        table._tbl.remove(r._tr)
        rebuilt = next(x for x in table.rows if x._tr is clone)
        for cell, text in zip(rebuilt.cells, texts):
            keep_format(cell, text)


def close_frd(doc, folder: str, stem: str, source: str, target: str, summary: str) -> None:
    keep_attention_with_its_boxes(doc)
    log_change(doc, target, summary, header_source=oldest_with_header(folder, stem, source))
    assert_absent_outside_log(doc, "MRD v2.93", "MRD v2.106")
    repair_ooxml(doc)
    normalise_change_log(doc)
    finish(doc, frd_path(folder, stem, target), f"{stem.replace('_', ' ')} v{target}", target)


# ── M22 Procurement & Requisitions ───────────────────────────────────────────

M22_DIR = "22-procurement-and-requisitions"
M22_STEM = "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document"
M22_SOURCE, M22_TARGET = "1.18", "1.19"

M22_FR018_HEADING = "FR-018 Narrow a Document List to a Fiscal Year or a Date Range"
M22_FR018 = [
    ("Requirement", REQUIREMENT),
    ("Current evidence",
     f"The requisition, RFQ and quotation lists take {PARAMETERS}. A requisition is placed by "
     "its request date, an RFQ by its issue date and a quotation by its quote date. "
     f"{NARROWING} The purchase orders, receipts, bills, credit notes and payments of Module "
     "23 take the same period through the same choke point (Module 23 FR-017). The "
     "requisition summary and the RFQ summary count the same period and branch as the list "
     "beneath them, so every tile counts the rows it sits above; the RFQ summary's responses "
     "in are the submitted quotations on issued RFQs, placed by their RFQ's issue date. A "
     "month-to-date figure under a past year reads zero, which is literally true. "
     f"{REFUSALS}"),
    ("Acceptance",
     "vs_procurement.tests_document_period drives the requisition list and summary, and "
     "OrderAndSourcingPeriodTests proves the RFQ summary counts the year asked for. "
     "DocumentPeriodListTests proves a fiscal year lists the documents its own dates cover and "
     "ends on its own last day rather than the calendar's, branch and year together list only "
     "that branch in that year, a range is inclusive at both ends and either end may be "
     "open, no period still lists every year, each of the four refusals, that a reader without "
     "the list permission is still refused, and that the summary counts the period the list "
     "shows. ArchivedYearPeriodTests proves the default list leaves an archived year out while "
     "naming that year, or a range across it, lists it. CrossTenantPeriodTests proves another "
     "tenant's year is unknown here and another tenant's documents never answer a range. The "
     "RFQ and quotation lists reach the same narrowing through _ProcBase.paginate and are not "
     "driven by a test of their own."),
    ("Current limit",
     "The requisition export follows neither the period nor the branch. A screen export "
     "reports fiscal_year, date_from, date_to and branch as filters it does not apply, and the "
     "file covers every branch and year in the reader's reach, so Mrs Bello exporting the "
     "list she narrowed to Ikeja Branch in 2015 receives every branch and year she can read. "
     "The shared screens' integration of the period is pending; this requirement records the "
     "server contract."),
]

M22_FR015_TAIL = (
    " Under ?fiscal_year= or a date range, beside ?branch=, it counts the period the list "
    "shows, so its month to date under a past year reads zero (FR-018)."
)
M22_FR015_ACCEPTANCE_TAIL = (
    " test_the_summary_counts_the_same_period_the_list_shows proves it for Ikeja Branch in "
    "fiscal year 2015."
)

M22_GAP = (
    "• A requisition list narrowed to a fiscal year, a date range or a branch does not export "
    "as narrowed. The screen export reports those parameters as filters it does not apply and "
    "writes every branch and year in the reader's reach, so the file is wider than the table "
    "it was asked from (FR-018)."
)

M22_TRACE_TAIL = (
    " Narrowing a document list and its summary to a fiscal year or a date range strengthens "
    "the purchase requisition, requisition summary, request for quotation and quotation "
    "capture entries without changing the count."
)

M22_SUMMARY = (
    "Minor revision. New FR-018 records that the requisition, RFQ and quotation lists take "
    "?fiscal_year=, one of the entity's own fiscal years by its label, or an inclusive "
    "?date_from= and ?date_to= with either end open. Each document is placed by its own date, "
    "a fiscal year by its own start and end, and the period is ANDed with ?branch= and every "
    "other filter. The narrowing sits in _ProcBase.paginate, so Module 23's document lists "
    "take it too, and lists of vendors, contracts and stock ignore it. Four refusals answer "
    "400 naming the field, and the year lookup never leaves the caller's entity. A named "
    "period replaces the default that hides archived years; without one a list is unchanged. "
    "FR-015 and the RFQ summary count the same period and branch. The requisition, RFQ and "
    "quotation routes, a row for the RFQ summary route, the typed errors, the Module 19 "
    "dependency and traceability to MRD v2.107 follow, and a new gap records that the "
    "requisition export follows neither the period nor the branch. The change log's header "
    "row, which v1.18 wrote a version over, is restored and the duplicated 1.17 row removed. "
    "No Field Access field is added or changed: the period narrows rows, not fields. Module "
    f"22 stays at 16 capabilities. {TEST_EVIDENCE}"
)


def patch_m22() -> None:
    doc = start_frd(M22_DIR, M22_STEM, M22_SOURCE, M22_TARGET, 22)

    fr015 = fr_table(doc, "FR-015")
    append_to(fr015, "Current evidence", M22_FR015_TAIL)
    append_to(fr015, "Acceptance", M22_FR015_ACCEPTANCE_TAIL)
    add_fr(doc, M22_FR018_HEADING, "FR-018 | Implemented", M22_FR018)

    for label, text in (
        ("GET, POST /requisitions/",
         "List and raise requisitions. The list takes ?fiscal_year=, or ?date_from= and "
         "?date_to=, beside ?branch= (FR-018)."),
        ("GET /requisitions/summary/", "Pipeline state by status, over the list's branch and period."),
        ("GET, POST /rfqs/", "List and draft requests for quotation; the list takes the same period."),
        ("GET, POST /quotations/", "Capture and compare offers; the list takes the same period."),
    ):
        keep_format(row_labelled(doc, label).cells[-1], text)
    insert_row_after(row_labelled(doc, "GET, POST /rfqs/"), [
        "GET /rfqs/summary/",
        "Draft, open and closing-soon RFQs and the responses received, over the list's branch "
        "and period.",
    ])
    errors = table_with_header(doc, "Condition")
    insert_rows_after(row(errors, "Opening a dashboard with no procurement key"), TYPED_ERRORS)

    append_to(table_with_header(doc, "Dependency"), "Module 19, Finance & Accounting",
              MODULE19_TAIL.format(fr="FR-018"))
    append_box_line(doc, "FURTHER GAPS", M22_GAP)

    edit_paragraph(doc, "Module 22 carries 16 capability entries", "MRD v2.93",
                   f"MRD v{MRD_TARGET}")
    lead = paragraph_starting(doc, "Module 22 carries 16 capability entries")
    mrd_tools.retitle(lead, lead.text.rstrip() + M22_TRACE_TAIL)
    for capability, requirements in (
        ("Purchase requisitions", "FR-001, FR-018"),
        ("Requisition summaries", "FR-015, FR-016, FR-017, FR-018"),
        ("Requests for quotation", "FR-010, FR-018"),
        ("Quotation capture", "FR-012, FR-013, FR-018"),
    ):
        set_trace(doc, capability, requirements)

    close_frd(doc, M22_DIR, M22_STEM, M22_SOURCE, M22_TARGET, M22_SUMMARY)


# ── M23 Purchase Orders, Delivery & AP ───────────────────────────────────────

M23_DIR = "23-purchase-orders-delivery-and-ap"
M23_STEM = "XVS_M23_Purchase_Orders_Delivery_and_AP_Functional_Requirements_Document"
M23_SOURCE, M23_TARGET = "1.18", "1.19"

M23_FR017_HEADING = "FR-017 Narrow a Document List to a Fiscal Year or a Date Range"
M23_FR017 = [
    ("Requirement", REQUIREMENT),
    ("Current evidence",
     f"The purchase-order, goods-receipt, vendor-bill, vendor-credit-note and vendor-payment "
     f"lists take {PARAMETERS}. An order is placed by its order date, a receipt by its "
     "received date, a bill by its invoice date, a credit note by its note date and a payment "
     "by its payment date. A goods return has no list of its own and is read through its "
     f"receipt. {NARROWING} The requisition, RFQ and quotation lists of Module 22 take the "
     "same period (Module 22 FR-018). The purchase-order summary and the bill summary count "
     "the same period and branch as the list beneath them, so every tile counts the rows it "
     "sits above; the order summary's month to date under a past year reads zero, which is "
     "literally true, and a bill from that year still unpaid past its due date still counts "
     f"as overdue, judged on the school's day (Section 2.3). {REFUSALS}"),
    ("Acceptance",
     "vs_procurement.tests_document_period drives the requisition list and summary (Module 22 "
     "FR-018): a year bounded by its own dates, branch and year together, an inclusive and "
     "open-ended range, every year without a period, the four refusals, the list permission "
     "still enforced, an archived year listed when named and hidden by default, and another "
     "tenant's years and documents never answering. OrderAndSourcingPeriodTests proves the "
     "purchase-order list answers branch and year together and that the purchase-order summary "
     "counts the same orders as the list. The receipt, bill, credit-note and payment lists reach "
     "the same narrowing through _ProcBase.paginate, and the bill summary applies the same "
     "period through vs_finance.period_filter.period_q; no test names a period on those, so "
     "their coverage is the shared choke point's."),
    ("Current limit",
     "The purchase-order and bill exports follow neither the period nor the branch. A screen "
     "export reports fiscal_year, date_from, date_to and branch as filters it does not apply, "
     "and the file covers every branch and year in the reader's reach, so a bursar exporting "
     "Lekki Branch's 2015 bills receives every branch and year she can read. The shared "
     "screens' integration of the period is pending; this requirement records the server "
     "contract."),
]

M23_GAP = (
    "• An order or bill list narrowed to a fiscal year, a date range or a branch does not "
    "export as narrowed. The screen export reports those parameters as filters it does not "
    "apply and writes every branch and year in the reader's reach, so the file is wider than "
    "the table it was asked from (FR-017)."
)

M23_TRACE_TAIL = (
    " Narrowing a document list and its summary to a fiscal year or a date range strengthens "
    "the purchase-order, goods-received note, vendor invoice and vendor payment entries "
    "without changing the count."
)

M23_SUMMARY = (
    "Minor revision. New FR-017 records that the purchase-order, goods-receipt, vendor-bill, "
    "vendor-credit-note and vendor-payment lists take ?fiscal_year=, one of the entity's own "
    "fiscal years by its label, or an inclusive ?date_from= and ?date_to= with either end "
    "open, through the same _ProcBase.paginate choke point as Module 22's lists. Each "
    "document is placed by its own date, a fiscal year by its own start and end, and the "
    "period is ANDed with ?branch= and every other filter. Four refusals answer 400 naming "
    "the field, and the year lookup never leaves the caller's entity. A named period replaces "
    "the default that hides archived years; without one a list is unchanged. The "
    "purchase-order and bill summaries count the same period and branch. The order, receipt, "
    "bill and payment routes, a row for the bill summary route, the typed errors, the Module "
    "19 dependency and traceability to MRD v2.107 follow, and a new gap records that the "
    "order and bill exports follow neither the period nor the branch. The change log's header "
    "row, which v1.18 wrote a version over, is restored and the duplicated 1.17 row removed, "
    "and a table grid stored after its rows is put back in schema order. No Field Access "
    "field is added or changed: the period narrows rows, not fields. Module 23 stays at 25 "
    f"capabilities. {TEST_EVIDENCE}"
)


def patch_m23() -> None:
    doc = start_frd(M23_DIR, M23_STEM, M23_SOURCE, M23_TARGET, 23)

    add_fr(doc, M23_FR017_HEADING, "FR-017 | Implemented", M23_FR017)

    for label, text in (
        ("GET, POST /purchase-orders/, /purchase-orders/{id}/",
         "Commitments. The list takes ?fiscal_year=, or ?date_from= and ?date_to=, beside "
         "?branch= (FR-017)."),
        ("GET /purchase-orders/summary/", "Position across orders, over the list's branch and period."),
        ("GET, POST /goods-receipts/, /goods-receipts/{id}/",
         "Record what arrived; the list takes the same period."),
        ("GET, POST /vendor-invoices/, /vendor-invoices/{id}/",
         "Capture and read; the list takes the same period."),
        ("GET, POST /vendor-payments/, /vendor-payments/{id}/",
         "The settlement plan; the list takes the same period."),
    ):
        keep_format(row_labelled(doc, label).cells[-1], text)
    insert_row_after(row_labelled(doc, "GET, POST /vendor-invoices/, /vendor-invoices/{id}/"), [
        "GET /vendor-invoices/summary/",
        "The bill tabs' counts and the overdue balance, over the list's branch and period.",
    ])
    errors = table_with_header(doc, "Condition or route")
    insert_rows_after(row(errors, "Naming a vendor another branch keeps to itself"), TYPED_ERRORS)

    append_to(table_with_header(doc, "Dependency"), "Module 19, Finance & Accounting",
              MODULE19_TAIL.format(fr="FR-017"))
    append_box_line(doc, "FURTHER GAPS", M23_GAP)

    edit_paragraph(doc, "Module 23 carries 25 capability entries", "MRD v2.93",
                   f"MRD v{MRD_TARGET}")
    lead = paragraph_starting(doc, "Module 23 carries 25 capability entries")
    mrd_tools.retitle(lead, lead.text.rstrip() + M23_TRACE_TAIL)
    for capability, requirements in (
        ("Purchase-order creation and lifecycle", "FR-001, FR-017"),
        ("Goods-received notes", "FR-003, FR-017"),
        ("Vendor invoice capture", "FR-005, FR-017"),
        ("Vendor payment creation", "FR-007, FR-009, FR-016, FR-017"),
    ):
        set_trace(doc, capability, requirements)

    close_frd(doc, M23_DIR, M23_STEM, M23_SOURCE, M23_TARGET, M23_SUMMARY)


# ── MRD ──────────────────────────────────────────────────────────────────────

MRD_SOURCE_SCOPE = (
    "Backend main at 35e234f with the document period change in its working tree, 17 October "
    "2026. Backend evidence only; the shared screens' integration is pending and no "
    "deployment is claimed."
)
MRD_CONTENTS_NOTE = "Current contracts by area; procurement document periods, no capability added"
MRD_SECTION_HEADING = "5.2 17 October 2026 Capability Contracts"
MRD_SECTION_BULLETS = [
    "• Procurement document periods: every procurement list over a finance document takes one "
    "of the entity's own fiscal years or an inclusive date range, places each document by its "
    "own date, ANDs the period with the branch and every other filter, and lists an archived "
    "year when that year is named. The requisition, RFQ, purchase-order and bill summaries "
    "count the same period. Lists of vendors, contracts and stock are not narrowed.",
    "• Status and ownership: exports are not narrowed by the period or the branch, recorded "
    "under Needs Attention for Modules 22 and 23. Capability entries stay at 513, no module "
    "state moves, and M22 and M23 advance to FRD v1.19.",
]

M22_BLURB_OLD = "Documented by M22 FRD v1.16."
M22_BLURB_NEW = (
    "Every requisition, RFQ and quotation list narrows to one fiscal year or a date range "
    "beside the branch, and so do the summaries above them. Documented by M22 FRD v1.19."
)
M23_BLURB_OLD = "Documented by M23 FRD v1.16."
M23_BLURB_NEW = (
    "Every order, receipt, bill, credit note and payment list narrows to one fiscal year or a "
    "date range beside the branch, and so do the order and bill summaries. Documented by M23 "
    "FRD v1.19."
)
M22_ATTENTION = (
    "• A requisition list narrowed to a fiscal year, a date range or a branch exports every "
    "branch and year in the reader's reach. The export reports those filters as ones it does "
    "not apply rather than narrowing the file to what the screen shows."
)
M23_ATTENTION = (
    "• An order or bill list narrowed to a fiscal year, a date range or a branch exports every "
    "branch and year in the reader's reach. The export reports those filters as ones it does "
    "not apply rather than narrowing the file to what the screen shows."
)
MRD_CHANGE_SUMMARY = (
    "Every procurement list over a finance document, requisitions, RFQs, quotations, purchase "
    "orders, goods receipts, vendor bills, vendor credit notes and vendor payments, takes "
    "?fiscal_year=, the entity's own year by its label, or an inclusive ?date_from= and "
    "?date_to=. Each document is placed by its own date and the period is ANDed with ?branch= "
    "and every other filter. A year the books do not have, a year with a range, a malformed "
    "date and a reversed range are refused with 400 naming the field, and the year lookup "
    "stays inside the caller's entity. Naming a period lists an archived year's documents; "
    "without one the lists are unchanged. The requisition, RFQ, purchase-order and bill "
    "summaries count the same period. Exports are not narrowed by period or branch, recorded "
    "under Needs Attention for Modules 22 and 23. No Field Access field changes. Capability "
    "entries stay at 513 and no status moves. M22 and M23 advance to v1.19. The change log's "
    "header row, which v2.106 wrote a version over, is restored, and the version rows printed "
    "in the header's colours or without borders take the log's ordinary ruled format. "
    + TEST_EVIDENCE
)


def capability_box(doc, first_capability: str):
    """The Needs Attention cell of the module grid that lists ``first_capability``."""
    grids = [t for t in doc.tables
             if any(p.text.strip() == f"▸  {first_capability}"
                    for r in t.rows for c in r.cells for p in c.paragraphs)]
    if len(grids) != 1:
        raise ValueError(f"{first_capability!r} is listed in {len(grids)} grids")
    cells = list({id(c._tc): c for r in grids[0].rows for c in r.cells
                  if c.text.strip().startswith("NEEDS ATTENTION")}.values())
    if len(cells) != 1:
        raise ValueError(f"The {first_capability!r} grid has {len(cells)} Needs Attention boxes")
    return cells[0]


def add_section(doc, after_heading: str, before_heading: str, heading: str,
                bullets: list[str]) -> None:
    """Add a sub-section after ``after_heading``'s bullets, cloned from that sub-section."""
    paragraphs = doc.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == after_heading)
    end = next(i for i, p in enumerate(paragraphs)
               if i > start and p.text.strip() == before_heading)
    last = paragraphs[end - 1]
    if not last.text.strip().startswith("•"):
        raise ValueError(f"{after_heading!r} does not end on a bullet: {last.text[:60]!r}")
    anchor = last._p
    new_heading = copy.deepcopy(paragraphs[start]._p)
    anchor.addnext(new_heading)
    mrd_tools.retitle(Paragraph(new_heading, last._parent), heading)
    anchor = new_heading
    for text in bullets:
        clone = copy.deepcopy(last._p)
        anchor.addnext(clone)
        mrd_tools.retitle(Paragraph(clone, last._parent), text)
        anchor = clone


def patch_mrd() -> None:
    folder = ROOT / "module-requirements"
    stem = "XVS_Module_Requirements_Document"
    require_newest(str(folder / f"{stem}_v*.docx"), MRD_SOURCE)
    doc = Document(str(folder / f"{stem}_v{MRD_SOURCE}.docx"))
    header_source = Document(str(folder / f"{stem}_v2.105.docx"))
    index = doc.tables[5]
    if index.rows[0].cells[5].text.strip() != "Entries":
        raise ValueError("Table 5 is not the module index")
    total_before = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])

    mrd_tools.replace_cover_version(doc.tables[0], MRD_SOURCE, MRD_TARGET)
    set_control(doc, "Version", MRD_TARGET)
    set_control(doc, "Review date", REVIEW_DATE)
    set_control(doc, "Source scope", MRD_SOURCE_SCOPE)
    contents = doc.tables[2]
    keep_format(row(contents, "5. Current Capability Contracts").cells[1], MRD_CONTENTS_NOTE)

    edit_blurb(doc, 22, M22_BLURB_OLD, M22_BLURB_NEW)
    edit_blurb(doc, 23, M23_BLURB_OLD, M23_BLURB_NEW)
    edit_box(capability_box(doc, "Purchase requisitions"), [("append", M22_ATTENTION)])
    edit_box(capability_box(doc, "Purchase-order creation and lifecycle"),
             [("append", M23_ATTENTION)])

    add_section(doc, "5.1 14 October 2026 Capability Contracts",
                "6. Recommended Next Build Order", MRD_SECTION_HEADING, MRD_SECTION_BULLETS)

    total_after = sum(int(r.cells[5].text.strip()) for r in index.rows[1:])
    if total_after != total_before:
        raise ValueError(f"Capability total moved from {total_before} to {total_after}")

    log_change(doc, MRD_TARGET, MRD_CHANGE_SUMMARY, header_source=header_source)
    assert_absent_outside_log(doc, M22_BLURB_OLD, M23_BLURB_OLD)
    repair_ooxml(doc)
    normalise_change_log(doc)
    rebuild_borderless_rows(doc)
    finish(doc, folder / f"{stem}_v{MRD_TARGET}.docx",
           f"XVS Module Requirements Document v{MRD_TARGET}", MRD_TARGET)


def main() -> None:
    patch_m22()
    patch_m23()
    patch_mrd()


if __name__ == "__main__":
    main()
