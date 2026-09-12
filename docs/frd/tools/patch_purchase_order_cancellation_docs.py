#!/usr/bin/env python3
"""Record purchase-order cancellation, and the requisition rule that now depends on it.

A purchase order had no way to reach CANCELLED through the API, so the refusal a
buyer met on a requisition ("Cancel the order instead") named an action that did
not exist. Module 23 gains the route, its five refusals, the audit action and the
plain statement that no cancellation notice reaches the vendor. Module 22's
requisition rule is restated: a requisition refuses while an order raised from it
still stands or has taken delivery, and one whose every order was cancelled
without a delivery is reversible. Module 7 carries the same correction in its
account of the engine's reversal contract.

Each document takes one minor version. Source MRD references move to v2.78, which
carries the new Module 23 capability entry.

    python tools/patch_purchase_order_cancellation_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.table import _Row

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
)

REVIEW_DATE = "12 September 2026"
SHORT_DATE = "12 Sep 2026"
MRD_VERSION = "2.78"
CODE_BASELINE = "Backend worktree at b1b270d, 12 September 2026"
DEPLOYMENT = "Backend evidence only; nothing here is deployed."
VERIFIED = "Verified by 567 procurement tests and 759 finance tests."


# ── keep-format docx helpers ────────────────────────────────────────────────

def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError(f"Paragraph carries no run to inherit formatting from: {paragraph.text!r}")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    """Rewrite a one-paragraph cell, keeping the formatting it already has."""
    for paragraph in cell.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    set_run_text(cell.paragraphs[0], text)


def append_cell_text(cell, tail: str) -> None:
    set_cell_text(cell, cell.text.strip() + tail)


def rewrite_box(cell, lines: list[str]) -> None:
    """Rewrite a callout cell line by line.

    A callout is one bold coloured heading paragraph followed by regular bullet
    paragraphs, so line i reuses paragraph i, overflow clones the last bullet,
    and surplus paragraphs are removed. Cloning the heading for a new bullet
    would render every added line as another heading.
    """
    original = list(cell.paragraphs)
    if len(original) < 2:
        raise ValueError("A callout needs a heading and at least one bullet to reuse")
    for index, line in enumerate(lines):
        if index < len(original):
            set_run_text(original[index], line)
        else:
            clone = copy.deepcopy(original[-1]._p)
            cell.paragraphs[-1]._p.addnext(clone)
            set_run_text(cell.paragraphs[-1], line)
    for paragraph in original[len(lines):]:
        paragraph._p.getparent().remove(paragraph._p)


def replace_bullet(cell, prefix: str, text: str) -> None:
    """Rewrite exactly one callout bullet in place, identified by how it starts."""
    lines = [paragraph.text for paragraph in cell.paragraphs]
    matches = [index for index, line in enumerate(lines) if line.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one bullet starting {prefix!r}, found {len(matches)}")
    lines[matches[0]] = text
    rewrite_box(cell, lines)


def unique_cells(row) -> list:
    seen, cells = set(), []
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        cells.append(cell)
    return cells


def find_table(doc, header: str, *, contains: str | None = None):
    """The first table whose first cell starts with ``header``."""
    for table in doc.tables:
        if not table.rows[0].cells[0].text.strip().startswith(header):
            continue
        if contains is None or any(
            row.cells[0].text.strip().startswith(contains) for row in table.rows
        ):
            return table
    raise ValueError(f"Table not found: {header!r} containing {contains!r}")


def find_fr(doc, label: str):
    return find_table(doc, f"{label} |")


def find_row(table, prefix: str, *, column: int = 0):
    for row in table.rows:
        if row.cells[column].text.strip().startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix!r}")


def field(table, label: str):
    """The value cell of a requirement or control table row named ``label``."""
    return find_row(table, label).cells[1]


def clone_row(table, anchor, values: list[str], *, before: bool = False):
    new_tr = copy.deepcopy(anchor._tr)
    if before:
        anchor._tr.addprevious(new_tr)
    else:
        anchor._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)
    return row


def set_row_text(row, values: list[str]) -> None:
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)


def insert_requirement_after(doc, label: str, heading: str, fr: dict):
    """Add a requirement block below ``label``'s, spaced like its neighbours.

    A block is a Heading 2 paragraph, its table, and a plain spacer paragraph. The
    spacer is copied from above the anchor's own heading rather than from whatever
    follows the anchor's table, because the last requirement in a section is
    followed by the page break that ends the section, and cloning that would put
    the new block on a page of its own. Each clone is inserted directly after the
    anchor table, in reverse order, so the result reads spacer, heading, table.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    table = find_fr(doc, label)
    heading_el = table._tbl.getprevious()
    spacer_el = heading_el.getprevious()
    if not Paragraph(heading_el, table._parent).style.name.startswith("Heading"):
        raise ValueError("No heading directly above the requirement table")
    spacer = Paragraph(spacer_el, table._parent)
    if spacer.text.strip() or 'w:type="page"' in spacer_el.xml:
        raise ValueError("The element above the heading is not a plain spacer")

    new_table = copy.deepcopy(table._tbl)
    new_heading = copy.deepcopy(heading_el)
    new_spacer = copy.deepcopy(spacer_el)
    table._tbl.addnext(new_table)
    table._tbl.addnext(new_heading)
    table._tbl.addnext(new_spacer)

    set_run_text(Paragraph(new_heading, table._parent), heading)
    block = Table(new_table, table._parent)
    set_run_text(block.rows[0].cells[0].paragraphs[0], fr["header"])
    for name in ("Requirement", "Current evidence", "Acceptance", "Current limit"):
        set_cell_text(field(block, name), fr[name])
    return block


def hold_together(table, *, lead_rows: int = 2) -> None:
    """Keep every row whole, and the first ``lead_rows`` with the row after them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
    for row in table.rows[:lead_rows]:
        for cell in unique_cells(row):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def paginate(doc) -> None:
    """Stop a requirement's heading and banner, or a table's lead, being stranded."""
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    for table in doc.tables:
        first = table.rows[0].cells[0].text.strip()
        if not ((first.startswith("FR-") and " | " in first) or first == "MRD capability"):
            continue
        hold_together(table)
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            Paragraph(previous, table._parent).paragraph_format.keep_with_next = True
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("All routes are mounted under /v1/procurement/"):
            paragraph.paragraph_format.keep_with_next = True


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def update_control(doc, *, version: str, module: int) -> None:
    control = find_table(doc, "Document", contains="Code baseline")
    set_cell_text(field(control, "Version"), version)
    set_cell_text(field(control, "Review date"), REVIEW_DATE)
    set_cell_text(field(control, "Code baseline"), CODE_BASELINE)
    set_cell_text(
        field(control, "Source MRD"),
        f"XVS Module Requirements Document v{MRD_VERSION} | Module {module}",
    )


def body_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Body paragraph not found: {prefix!r}")


def prepend_change_log(doc, version: str, summary: str) -> None:
    table = find_table(doc, "Version", contains="1.0")
    clone_row(table, table.rows[1], [version, SHORT_DATE, summary], before=True)


def finish(doc, output: Path, title: str, version: str) -> None:
    paginate(doc)
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


# ═════════════════════════════════════════════════════════════════════════════
# Module 22 - Procurement & Requisitions
# ═════════════════════════════════════════════════════════════════════════════

M22_DIR = "22-procurement-and-requisitions"
M22_STEM = "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document"
M22_SOURCE, M22_TARGET = "1.9", "1.10"

M22_FR003_EVIDENCE = (
    "submit_for_approval hands the document to Module 7, which selects the "
    "template for the scope, activates stages, and calls the registered handler "
    "back on a terminal decision. The handler flips the procurement approval state "
    "and applies the document-type effect. A reversal the engine records against "
    "the vote that decided a stage is handed back as well: on_action_reversed "
    "returns an APPROVED or REJECTED document to PENDING, and a requisition's "
    "status to PENDING_APPROVAL, so the overlay never reads APPROVED while the "
    "workflow shows the document back under review. Whether a reversal is allowed "
    "at all is asked of the document before anything is written, and each document "
    "type answers for what its own approval released. A requisition asks what the "
    "orders raised from it did, not whether one was ever raised. Goods answer "
    "first and answer whatever state their order is in now: a posted receipt "
    "against any order raised from the requisition refuses the reversal, a "
    "cancelled order included, because the stock and the GR/IR liability it "
    "brought are not handed back by cancelling the order they arrived against. "
    "Failing that, an order that is still live, meaning neither cancelled nor "
    "reversed, refuses it too, because it is a commitment to a vendor that holds "
    "whether or not the vote behind it does. When every order raised has been "
    "cancelled and nothing arrived, the approval has released nothing that "
    "outlives it and the reversal proceeds. Purchase orders, vendor bills and "
    "vendor payments answer for their own in Module 23."
)
M22_FR003_ACCEPTANCE = (
    "The submit response carries the approval block, so the client learns in the "
    "same round trip when nobody can approve what was just submitted, and what "
    "would fix it. "
    "WorkflowApprovalTests.test_reversing_the_first_vote_returns_the_requisition_to_its_ladder "
    "proves a reversed first vote puts the requisition back to PENDING and "
    "PENDING_APPROVAL, and that both stages then accept a fresh vote. "
    "ProcurementApprovalReversalTests holds the three answers: "
    "test_a_requisition_an_order_was_raised_from_cannot_be_reversed for a live "
    "order, test_a_requisition_whose_cancelled_order_took_delivery_is_refused for "
    "goods that arrived against an order since cancelled, and "
    "test_a_requisition_whose_only_order_was_cancelled_is_reversible for the case "
    "that releases it. "
    "test_cancelling_the_only_order_frees_the_requisition_to_be_reversed runs the "
    "whole chain, and test_every_procurement_document_type_answers_the_reversal_check "
    "enumerates the registered document types rather than sampling them, so a "
    "fifth cannot be added without one."
)
M22_FR003_LIMIT = "None."

M22_APPROVED_LEAVES = (
    "Terminal here unless an administrator reverses the deciding vote, which "
    "returns it to PENDING. While an order raised from it still stands, or goods "
    "arrived against one, that reversal is refused; cancelling the order, which is "
    "a route a buyer can reach in Module 23, releases it again. Raising the order "
    "is in any case a separate act in Module 23."
)

M22_MODULE7 = (
    "Owns templates, approver resolution, routing, eligibility and the "
    "parked-approval repair. This module registers handlers, publishes its "
    "default ladders through the engine's publish service, and fences every call "
    "to the four procurement document types. Its handlers also answer the "
    "engine's reversal contract, each type for what its own approval released: a "
    "requisition while an order raised from it still stands or has taken "
    "delivery, an order that reached its vendor or has goods receipted against "
    "it, and a bill or a payment whose status has left draft or pending approval "
    "each refuse the reversal, while a document that released nothing returns to "
    "PENDING."
)

M22_TRACEABILITY_LEAD = (
    f"Module 22 carries 16 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Asking what a requisition's orders did, rather than "
    "whether one was raised, strengthens workflow submission without changing the "
    "count."
)

M22_CHANGE = (
    "Corrects the requisition's reversal rule and removes the limit it carried. "
    "v1.9 recorded that a requisition refuses once a purchase order has been "
    "raised from it, which asked whether a row existed rather than what it did: a "
    "cancelled order blocked the requisition for ever, and the refusal told a "
    "buyer to cancel an order that was already cancelled. The requisition now asks "
    "what its orders did. Goods answer first, whatever state their order is in "
    "now, because stock on the shelf and the GR/IR liability are not handed back "
    "by cancelling the order they arrived against; failing that a live order "
    "refuses, meaning one neither cancelled nor reversed; and a requisition whose "
    "every order was cancelled without a delivery is reversible. FR-003's "
    "evidence, acceptance and limit, the APPROVED row of the requisition "
    "lifecycle and the Module 7 dependency are restated. The limit is now None "
    "rather than the dead end v1.9 implied: cancelling a purchase order is a route "
    "a buyer can reach, recorded in Module 23 FR-015. Module 22 stays at 16 "
    f"capabilities. {VERIFIED} {DEPLOYMENT}"
)


def patch_m22(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M22 Procurement and Requisitions Functional Requirements "
        f"Document v{M22_TARGET}"
    )

    fr003 = find_fr(doc, "FR-003")
    lifecycle = find_table(doc, "State", contains="DRAFT")
    dependencies = find_table(doc, "Dependency")

    replace_cover_version(doc.tables[0], M22_SOURCE, M22_TARGET)
    update_control(doc, version=M22_TARGET, module=22)

    set_cell_text(field(fr003, "Current evidence"), M22_FR003_EVIDENCE)
    set_cell_text(field(fr003, "Acceptance"), M22_FR003_ACCEPTANCE)
    set_cell_text(field(fr003, "Current limit"), M22_FR003_LIMIT)

    set_cell_text(find_row(lifecycle, "APPROVED").cells[2], M22_APPROVED_LEAVES)
    set_cell_text(find_row(dependencies, "Module 7").cells[1], M22_MODULE7)

    set_run_text(body_paragraph(doc, "Module 22 carries"), M22_TRACEABILITY_LEAD)
    prepend_change_log(doc, M22_TARGET, M22_CHANGE)
    finish(doc, output, title, M22_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 23 - Purchase Orders, Delivery & AP
# ═════════════════════════════════════════════════════════════════════════════

M23_DIR = "23-purchase-orders-delivery-and-ap"
M23_STEM = "XVS_M23_Purchase_Orders_Delivery_and_AP_Functional_Requirements_Document"
M23_SOURCE, M23_TARGET = "1.11", "1.12"

M23_FR015_HEADING = "FR-015 Cancel a Commitment Nobody Will Fulfil"

M23_FR015 = {
    "header": "FR-015 | Implemented",
    "Requirement": (
        "An order that will never be delivered must be closable by the buyer who "
        "holds it, and refused wherever something downstream already depends on it."
    ),
    "Current evidence": (
        "POST /purchase-orders/{id}/cancel/ takes the order to CANCELLED under "
        "procurement.purchase_order.update, the state-change verb that already "
        "edits it, so no school has to grant a new key before a dead order can be "
        "closed. The order is resolved through the same branch-scoped resolver "
        "every detail and action endpoint uses and is locked for the duration, so "
        "an order at a branch the caller does not work in answers exactly as one "
        "that does not exist. The rules sit in the cancellation service rather "
        "than the view, so the school-facing layer meets the same refusals. Five "
        "conditions refuse, each with its own code: a missing reason "
        "(PURCHASE_ORDER_CANCEL_REASON_REQUIRED), goods already received "
        "(PURCHASE_ORDER_ALREADY_RECEIVED), a vendor bill that is not itself "
        "cancelled, matched at the header or through one of its lines "
        "(PURCHASE_ORDER_ALREADY_BILLED), an approval still in flight "
        "(PURCHASE_ORDER_UNDER_APPROVAL), and an order already cancelled or "
        "reversed (PURCHASE_ORDER_CANCEL_REFUSED). A payment needs no rule of its "
        "own, because it reaches an order only through a bill. Cancelling writes a "
        "PURCHASE_ORDER_CANCELLED audit record carrying the actor, the mandatory "
        "reason and the status before and after, and cancels a vendor email that "
        "approval had scheduled but nothing has sent."
    ),
    "Acceptance": (
        "A cancelled order leaves the pipeline KPIs, which already exclude a "
        "cancelled or reversed commitment. The approval overlay is left as it was: "
        "an order approved last week and cancelled today still reads APPROVED, "
        "because it was, and the ledger status carries the cancellation. The "
        "reason is mandatory, because it is what answers months later why an order "
        "to a supplier was withdrawn. PurchaseOrderCancellationTests covers the "
        "permitted issued order, the recorded reason, the approval that still "
        "reads approved, the order leaving the KPIs, and refusals for goods "
        "received, a live bill, a cancelled bill that does not block, an approval "
        "in flight, a second cancellation and a missing reason, with another "
        "branch's buyer answered exactly as for an order that does not exist."
    ),
    "Current limit": (
        "No cancellation notice reaches the vendor. This module has one "
        "vendor-facing order message, the order itself, so a supplier who already "
        "received an order learns of its withdrawal only when a person tells them. "
        "The route stops a scheduled email that has not gone out; anything queued "
        "or sent cannot be recalled."
    ),
}

M23_BUYER = [
    "Buyer",
    "Create and issue purchase orders, email them, retry a failed delivery, and "
    "cancel one that nothing stands against.",
    "procurement.purchase_order.* (create and update SENSITIVE; cancelling rides "
    "the update verb)",
]

M23_CANCEL_ROUTE = [
    "POST /purchase-orders/{id}/cancel/",
    "Withdraw a commitment, with a mandatory reason.",
]

M23_CANCEL_ERRORS = [
    ["Cancelling a purchase order without a reason",
     "Refused with PURCHASE_ORDER_CANCEL_REASON_REQUIRED."],
    ["Cancelling an order goods were received against",
     "Refused with PURCHASE_ORDER_ALREADY_RECEIVED; return the goods instead."],
    ["Cancelling an order a vendor bill stands against",
     "Refused with PURCHASE_ORDER_ALREADY_BILLED; credit or reverse the bill instead."],
    ["Cancelling an order whose approval is still being decided",
     "Refused with PURCHASE_ORDER_UNDER_APPROVAL; withdraw the approval first."],
    ["Cancelling an order already cancelled or reversed",
     "Refused with PURCHASE_ORDER_CANCEL_REFUSED."],
]

M23_MODULE7_TAIL = (
    " An order whose approval is still being decided refuses cancellation, "
    "because the engine owns that decision while it runs; withdrawing the approval "
    "through the engine returns the order to NOT_SUBMITTED, which cancellation "
    "then closes."
)

M23_VENDOR_GAP = (
    "• Cancelling a purchase order tells the vendor nothing. The only "
    "vendor-facing order message this module has is the order itself, so a "
    "supplier who already received one and then has it cancelled learns of the "
    "withdrawal when somebody rings them and not before. The audit row carries the "
    "reason, which is the record that the call is a person's job."
)

M23_TRACEABILITY_LEAD = (
    f"Module 23 carries 25 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Cancelling a purchase order with a recorded reason is "
    "the entry added in that version."
)

M23_NEW_CAPABILITY = [
    "Purchase-order cancellation with a recorded reason",
    "FR-015",
]

M23_CHANGE = (
    "Records purchase-order cancellation, which this module had no route for. An "
    "order that would never be delivered could not be closed through the API, so "
    "the refusal a buyer met on a requisition, to cancel the order instead, named "
    "an action that did not exist, and the requisition behind a dead order stayed "
    "approved for ever. New FR-015 records the route, the update verb it rides "
    "rather than a key no school has granted, the branch-scoped resolver and row "
    "lock, and the five refusals with their codes: no reason, goods received, a "
    "vendor bill that is not itself cancelled, an approval still in flight, and an "
    "order already closed. A payment needs no rule of its own because it reaches "
    "an order only through a bill. The route table, the typed errors, the buyer's "
    "row in the actors table and the Module 7 dependency are updated, and the "
    "PURCHASE_ORDER_CANCELLED audit record with its actor, reason and "
    "before-and-after is recorded. A new gap states plainly that no cancellation "
    "notice reaches the vendor: this module has one vendor-facing order message, "
    "the order itself, so a supplier who received an order learns of its "
    "withdrawal only when a person tells them. Module 23 moves from 24 to 25 "
    f"capability entries. {VERIFIED} {DEPLOYMENT}"
)


def patch_m23(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M23 Purchase Orders Delivery and AP Functional Requirements "
        f"Document v{M23_TARGET}"
    )

    actors = find_table(doc, "Actor")
    routes = find_table(doc, "Method and path", contains="POST /purchase-orders/{id}/submit/")
    errors = find_table(doc, "Condition or route")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")
    traceability = find_table(doc, "MRD capability")

    replace_cover_version(doc.tables[0], M23_SOURCE, M23_TARGET)
    update_control(doc, version=M23_TARGET, module=23)

    set_row_text(find_row(actors, "Buyer"), M23_BUYER)
    clone_row(routes, find_row(routes, "POST /purchase-orders/{id}/submit/"), M23_CANCEL_ROUTE)

    anchor = errors.rows[-1]
    for values in reversed(M23_CANCEL_ERRORS):
        anchor = clone_row(errors, anchor, values) if anchor is errors.rows[-1] \
            else clone_row(errors, anchor, values, before=True)
    append_cell_text(find_row(dependencies, "Module 7").cells[1], M23_MODULE7_TAIL)

    gap_cell = gaps.rows[0].cells[0]
    rewrite_box(gap_cell, [p.text for p in gap_cell.paragraphs] + [M23_VENDOR_GAP])

    clone_row(
        traceability,
        find_row(traceability, "Approval reversal refused once an order"),
        M23_NEW_CAPABILITY,
    )
    set_run_text(body_paragraph(doc, "Module 23 carries"), M23_TRACEABILITY_LEAD)

    insert_requirement_after(doc, "FR-014", M23_FR015_HEADING, M23_FR015)
    prepend_change_log(doc, M23_TARGET, M23_CHANGE)
    finish(doc, output, title, M23_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 7 - Workflow & Approval Engine
# ═════════════════════════════════════════════════════════════════════════════

M07_DIR = "07-workflow-and-approval-engine"
M07_STEM = "XVS_M07_Workflow_and_Approval_Engine_Functional_Requirements_Document"
M07_SOURCE, M07_TARGET = "1.10", "1.11"

M07_FR019_EVIDENCE = (
    "reverse_action asks the document's handler through validate_reversal before "
    "writing anything, then voids the vote: the original gains reversed_at, "
    "reversed_by, and a reason, and a reversal row is appended. Where the reversed "
    "vote is what decided its stage, the stage reopens, every stage activated "
    "after it returns to PENDING with its live votes voided and its snapshot "
    "cleared, a decided instance returns to IN_PROGRESS, and the handler's "
    "on_action_reversed puts the document back inside the same transaction. A vote "
    "the stage did not need is voided and nothing else moves. What each module "
    "refuses is its own. Payouts refuse once any instruction is claimed for the "
    "provider, and dispatch re-checks the approval under the same row lock; "
    "finance refuses once the document has moved past draft or pending approval; "
    "procurement answers for all four of its types through an overridable hook on "
    "its shared handler, refusing a requisition while an order raised from it "
    "still stands or has taken delivery, a purchase order once its vendor email is "
    "pending or sent or goods have been received against it, and a vendor invoice "
    "or vendor payment once its own status has left draft or pending approval, "
    "which is what posting does to it; user creation refuses once the account has "
    "left pending approval; role changes refuse once the request is decided. Leave "
    "returns a decided request to pending. A document type with no registered "
    "handler cannot be reversed. The view returns 404 for an action belonging to "
    "another tenant."
)

M07_REVERSAL_BULLET_PREFIX = "• Reversal is answered by the module that owns the document"
M07_REVERSAL_BULLET = (
    "• Reversal is answered by the module that owns the document, and every "
    "registered finance, procurement, payment, user-creation and role-change type "
    "now answers it: procurement refuses a requisition while an order raised from "
    "it still stands or took delivery, an order that reached its vendor or was "
    "received against, and a bill or a payment that has posted. The base handler "
    "still allows every reversal, so a leave request, and any document type "
    "registered later, is guarded by nothing until its own handler answers."
)

M07_DOMAIN_HANDLERS = (
    "Register document handlers and consume the approval outcome. Each handler "
    "also says whether a decision may still be reversed and puts its document back "
    "after one: payouts refuse once an instruction is claimed for the provider, "
    "finance once the document has posted, and procurement once what its own "
    "approval released has happened, which is an order still standing against a "
    "requisition or goods received against one, a purchase order that reached its "
    "vendor or was received against, and a posted vendor bill or vendor payment. "
    "Procurement's central ladder is published by this engine's publish service."
)

M07_TRACEABILITY_LEAD = (
    f"Module 7 carries 27 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Every registered procurement type answering the "
    "reversal check strengthens action reversal and the procurement workflow "
    "handlers without changing the count."
)

M07_CHANGE = (
    "Corrects one sentence of the procurement answer this document records, in the "
    "three places it appears. v1.10 said procurement refuses a requisition once a "
    "purchase order has been raised from it. It now refuses while an order raised "
    "from it still stands or has taken delivery, and allows the reversal once "
    "every order has been cancelled without a delivery, which is the case that "
    "releases a requisition whose supplier never delivered. FR-019's evidence, the "
    "reversal bullet in Needs Attention and the domain-handler dependency are "
    "restated. Nothing the engine does changes: it still asks the owning module "
    "before it writes, and still cannot withdraw what a vote released. Module 7 "
    f"stays at 27 capability entries. {DEPLOYMENT}"
)


def patch_m07(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M07 Workflow and Approval Engine Functional Requirements "
        f"Document v{M07_TARGET}"
    )

    fr019 = find_fr(doc, "FR-019")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M07_SOURCE, M07_TARGET)
    update_control(doc, version=M07_TARGET, module=7)

    set_cell_text(field(fr019, "Current evidence"), M07_FR019_EVIDENCE)
    set_cell_text(find_row(dependencies, "Modules 17 to 24").cells[1], M07_DOMAIN_HANDLERS)
    replace_bullet(gaps.rows[0].cells[0], M07_REVERSAL_BULLET_PREFIX, M07_REVERSAL_BULLET)

    set_run_text(body_paragraph(doc, "Module 7 carries"), M07_TRACEABILITY_LEAD)
    prepend_change_log(doc, M07_TARGET, M07_CHANGE)
    finish(doc, output, title, M07_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    parser.add_argument("--only", nargs="*", choices=["m07", "m22", "m23"])
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for tag, folder, stem, source, target, patch in (
        ("m07", M07_DIR, M07_STEM, M07_SOURCE, M07_TARGET, patch_m07),
        ("m22", M22_DIR, M22_STEM, M22_SOURCE, M22_TARGET, patch_m22),
        ("m23", M23_DIR, M23_STEM, M23_SOURCE, M23_TARGET, patch_m23),
    ):
        if args.only and tag not in args.only:
            continue
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists():
            raise SystemExit(f"Refusing to overwrite {output.name}")
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
