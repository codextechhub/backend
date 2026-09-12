#!/usr/bin/env python3
"""Version the Module 4 and Module 6 FRDs against the backend at 13675db.

Module 4 (Roles & Permissions) moves from v1.17 to v1.18. A role permission
change is decided by the rbac.role_change workflow ladder rather than by a
reviewer grant ceiling, and its requester may decide it. A restricted addition
is refused only where it lands on a role the editor holds, and the assignment
ceiling still refuses giving that role to yourself. Retiring the last way back
into role administration is refused, which becomes FR-026. A branch role names
its branch only once a school has more than one; school creation provisions
every tenant-wide library role; five keys no school can act on are PLATFORM; a
plan downgrade revokes out-of-plan role grants; and the enforcement switch is
registered ON. Three Needs Attention items are added.

Module 6 (Configuration & Capability) moves from v1.0 to v1.1. The catalogue
registers platform.entitlements.enforce ON, get_config still answers the plan
gate's own False where the key is not registered, the flag read is uncached,
and the switch's description promises a per-school excusal that its allowed
scopes refuse.

Both documents take MRD v2.76 as their source. Every existing cell is rewritten
through its first run, so the cell keeps the formatting it had; new rows and
the FR-026 block are deep copies of rows and blocks already in the state they
need.

    python tools/patch_role_change_ladder_docs.py [--overwrite]
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
)

REVIEW_DATE = "11 September 2026"
SHORT_DATE = "11 Sep 2026"
MRD_VERSION = "2.76"
CODE_BASELINE = "Backend worktree at 13675db, 11 September 2026"


# ── shared docx helpers ──────────────────────────────────────────────────────

def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError("Paragraph carries no run to inherit formatting from")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def write_kept(cell, text: str) -> None:
    """Replace a cell's text through its first paragraph's first run."""
    while len(cell.paragraphs) > 1:
        paragraph = cell.paragraphs[-1]
        paragraph._p.getparent().remove(paragraph._p)
    set_run_text(cell.paragraphs[0], text)


def edit_kept(cell, old: str, new: str) -> None:
    """Swap one exact passage inside a cell, refusing if it is not there."""
    text = cell.text
    if old not in text:
        raise ValueError(f"Passage not found: {old[:60]}")
    write_kept(cell, text.replace(old, new, 1))


def append_kept(cell, tail: str) -> None:
    write_kept(cell, cell.text.strip() + tail)


def replace_cover_version(table, source: str, target: str) -> None:
    """Rewrite the version on the cover, which is one run of one title block."""
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def find_row(table, prefix: str):
    for row in table.rows:
        if row.cells[0].text.strip().startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix}")


def set_control(table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip() == label:
            write_kept(row.cells[1], value)
            return
    raise ValueError(f"Control row not found: {label}")


def replace_body_paragraph(doc, prefix: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_run_text(paragraph, text)
            return
    raise ValueError(f"Body paragraph not found: {prefix}")


def find_body_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Body paragraph not found: {prefix}")


def clone_row(template, *, before=None, after=None):
    """Deep-copy ``template`` next to an anchor row and return the new row's XML."""
    clone = copy.deepcopy(template._tr)
    if before is not None:
        before._tr.addprevious(clone)
    elif after is not None:
        after._tr.addnext(clone)
    else:
        raise ValueError("Give an anchor row")
    return clone


def row_for(table, tr):
    for row in table.rows:
        if row._tr is tr:
            return row
    raise ValueError("Cloned row not found in its table")


def fill(row, values: dict[int, str]) -> None:
    for index, value in values.items():
        write_kept(row.cells[index], value)


def prepend_change_log(table, version: str, date: str, summary: str) -> None:
    tr = clone_row(table.rows[1], before=table.rows[1])
    row = row_for(table, tr)
    fill(row, {0: version, 1: date, 2: summary})


def insert_bullet_before(doc, anchor_prefix: str, text: str) -> None:
    anchor = find_body_paragraph(doc, anchor_prefix)
    clone = copy.deepcopy(anchor._p)
    anchor._p.addprevious(clone)
    for paragraph in doc.paragraphs:
        if paragraph._p is clone:
            set_run_text(paragraph, text)
            return
    raise ValueError("Inserted bullet not found")


def require_fr(table, label: str):
    if not table.rows[0].cells[0].text.strip().startswith(label):
        raise ValueError(f"Table is not {label}")
    return table


def _is_heading(element) -> bool:
    if element.tag != qn("w:p"):
        return False
    style = element.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return style is not None and style.get(qn("w:val"), "").startswith("Heading")


def keep_headings_with_their_tables(doc) -> None:
    """Hold a heading, any short lead-in, and a table's first two rows together.

    A heading style keeps itself with the next paragraph, but nothing keeps a
    lead-in paragraph with the table below it, or a table's first row with its
    second, so a long page could end on a heading and a lone header row or
    status bar. Keep-with-next on the lead-in paragraphs and the first row's
    paragraphs carries them over as one. Single-row callout boxes are left
    alone, because keeping one with the next paragraph would chain it onto the
    section that follows.
    """
    tables = {table._tbl: table for table in doc.tables}
    children = list(doc.element.body.iterchildren())
    for index, child in enumerate(children):
        if child.tag != qn("w:tbl") or len(tables[child].rows) < 2:
            continue
        lead_ins = []
        cursor = index - 1
        while cursor >= 0 and len(lead_ins) <= 2:
            element = children[cursor]
            if _is_heading(element):
                break
            if element.tag != qn("w:p"):
                lead_ins = None
                break
            lead_ins.append(element)
            cursor -= 1
        else:
            lead_ins = None
        if lead_ins is None or len(lead_ins) > 2 or cursor < 0:
            continue
        for paragraph in doc.paragraphs:
            if paragraph._p in lead_ins:
                paragraph.paragraph_format.keep_with_next = True
        for cell in tables[child].rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def finish(doc, output: Path, title: str, version: str) -> None:
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


EVIDENCE_ROW, ACCEPTANCE_ROW, LIMIT_ROW = 2, 3, 4


# ═════════════════════════════════════════════════════════════════════════════
# Module 4 - Roles & Permissions (RBAC)
# ═════════════════════════════════════════════════════════════════════════════

M04_DIR = "04-roles-and-permissions-rbac"
M04_STEM = "XVS_M04_Roles_and_Permissions_RBAC_Functional_Requirements_Document"
M04_SOURCE, M04_TARGET = "1.17", "1.18"

#: Table positions in v1.17, all bound before anything is inserted.
M04_T = {
    "cover": 0, "control": 1, "contents": 2, "in_scope": 3, "actors": 7,
    "FR-002": 10, "FR-003": 11, "FR-010": 18, "FR-017": 25, "FR-019": 27,
    "FR-021": 29, "FR-023": 31, "FR-025": 33,
    "withdrawn": 35, "model": 37, "api_routes": 40, "api_answers": 41,
    "dependencies": 42, "needs": 43, "removal": 44, "trace": 45,
    "reconcile": 46, "changes": 47,
}

M04_SOURCE_SCOPE = (
    "Backend changes deciding a role permission change through a workflow "
    "ladder, refusing a restricted addition only where it lands on the "
    "editor's own role, refusing to retire the last way back into role "
    "administration, taking out-of-plan grants back on a plan downgrade, and "
    "reclassifying five keys no school can act on (11 September 2026)"
)

M04_CODE_INSPECTED = (
    f"{CODE_BASELINE}: apps/vs_rbac (serializers/tenant.py, services.py, "
    "views.py, validators.py, workflow_handlers.py, plan_gate.py, scoping.py, "
    "migrations 0017 to 0019), the permission and prebuilt-role seeders, and "
    "the vs_workflow role-change ladder and vs_schools plan change they meet"
)

M04_MRD_BASELINE = (
    f"XVS Module Requirements Document v{MRD_VERSION}, Module 4, twenty-two "
    "capability entries"
)

M04_PURPOSE = (
    "Module 4 is the platform's authorization boundary. Every other module asks "
    "it the same two questions: may this person do this, and whose rows may "
    "they see. It answers from a governed permission registry, roles, "
    "assignments, and per-person exceptions. Two independent facts control a "
    "grant. Scope says whether a tenant may ever hold a key. The restricted "
    "flag says whether the key may be added without review, and review is "
    "needed where the addition lands on a role the person saving it holds: that "
    "is raised as a role change and decided through a workflow ladder. Added "
    "to anybody else's role it saves directly, because assignment refuses a "
    "role whose restricted keys the assigner does not already hold, so the "
    "editor cannot then hand that role to themselves. Groups and personal ALLOW "
    "overrides cannot carry a restricted key at all."
)

M04_BOUNDARIES = (
    "Five boundaries hold this module up, and they are deliberately different "
    "in kind. The tenant binds each request to one asserted tenant. Scope says "
    "whether that tenant may ever hold a key. Restriction says a sensitive key "
    "cannot land on its editor's own role without a role change decided by the "
    "role-change workflow ladder, and cannot be handed to anybody by an "
    "assigner who does not already hold it. A platform decision remains "
    "reserved to the platform tenant however its ordinary key was obtained. "
    "Finally, a tenant that has not gone live reaches only surfaces that opt in "
    "by name, and no role grant can reopen the rest."
)

M04_SCOPE_PREBUILT = (
    "PrebuiltRoleTemplate and its defaults, a platform-owned catalogue of five "
    "templates provisioned into a tenant's own roles rather than shared with "
    "it. School creation provisions the four tenant-wide templates and a Branch "
    "Admin per branch."
)

M04_SCOPE_ASSIGNMENTS = (
    "TenantUserRoleAssignment, whole-tenant or pinned to one branch, with "
    "revoke and replace and a per-branch uniqueness rule. Assign, replace, "
    "account creation, and draft submission all refuse a restricted role unless "
    "the actor already holds every restricted key in it, whoever the role is "
    "being given to, the actor included."
)

M04_SCOPE_ROLE_CHANGE = (
    "TenantRoleChangeRequest and its delta items, dependency validation, and "
    "atomic apply through set_role_access. Raising a request submits it to the "
    "rbac.role_change workflow ladder in the same transaction, and the handler "
    "this module registers with the engine applies the delta when the ladder "
    "approves. It is the route a restricted key takes onto a role its editor "
    "holds."
)

M04_ACTOR_SCHOOL_DOES = (
    "Build the tenant's role catalogue, including restricted additions to roles "
    "they do not hold, assign and revoke roles within their own grant ceiling, "
    "raise a role change for a restricted addition to a role they hold, and "
    "decide role changes where they hold the School Admin role the ladder "
    "names, their own included."
)

M04_ACTOR_SCHOOL_KEYS = (
    "school.roles.view / .create / .update / .delete / .assign for ordinary "
    "work; school.roles.approve to reach the decision endpoint, and a place on "
    "the ladder's frozen approver list for the vote to count; "
    "school.user_overrides.view / .manage for non-restricted exceptions."
)

M04_ACTOR_PLATFORM_DOES = (
    "Administer the platform registry and roles, and decide a platform role "
    "change where they hold the XVS Platform Admin role the platform ladder "
    "names. A platform reviewer is not an approver inside a school, because "
    "the engine resolves approvers from the school's own people. The Vision "
    "super administrator remains the bootstrap exception to the assignment "
    "ceiling."
)

M04_FR002_EVIDENCE_TAIL = (
    " Scope follows what a key's endpoint scopes to, not the module its name "
    "sits in. todo.task.view, .manage and .assign govern CodeX's own task "
    "tracker, whose queryset has no tenant column and scopes by the CX "
    "organogram; finance.entity.create creates a set of books, which CodeX "
    "gives a school when it is created; tickets.ticket.assign names which CodeX "
    "support person works a ticket. All five are PLATFORM. Migrations 0018 and "
    "0019 deleted every school role's grant of them first, because the grant "
    "models refuse a platform key on a tenant role and rows left behind would "
    "make those roles unsaveable, and 0019 also removed tickets.ticket.assign "
    "from the prebuilt defaults; 0018 did not do the same for "
    "finance.entity.create (see Needs Attention). platform.team.* and the "
    "platform.audit view and export keys stay tenant-holdable because their "
    "endpoints are tenant-scoped."
)

M04_FR002_ACCEPTANCE_TAIL = (
    " The same refusal now meets the five reclassified keys, and the tenant "
    "permission catalogue no longer offers them to a school. Only "
    "tickets.ticket.assign has a test pinning its scope "
    "(test_assignment_is_a_key_no_school_is_offered)."
)

M04_FR003_EVIDENCE = (
    "Scope remains enforced on every grant model, including bulk writes. "
    "Restricted authority has a separate path rule, and it asks whose role is "
    "being changed. Direct role creation refuses a restricted key, because it "
    "calls set_role_access without allow_restricted. A direct update refuses a "
    "restricted addition with 409 RESTRICTED_NEEDS_APPROVAL, carrying "
    "restricted_additions, only where the actor holds an active assignment to "
    "that role (_actor_holds_this_role); an addition to a role the actor does "
    "not hold is ordinary administration and saves, and the serializer passes "
    "that answer to set_role_access rather than letting the service decide "
    "again. GroupPermission and TenantRoleGroup reject restricted membership; "
    "UserPermissionOverride rejects restricted ALLOW; and assignment, replace, "
    "account creation, and draft submission refuse a role whose restricted keys "
    "the assigner does not already hold (missing_restricted_grant_authority), "
    "whoever receives it."
)

M04_FR003_ACCEPTANCE = (
    "The role path, group path, prebuilt path, assignment path, override path, "
    "account creation path, and draft-submission path all enforce scope. A "
    "restricted key reaches a custom role by two routes: an edit by somebody "
    "who does not hold that role, or a role change decided through the workflow "
    "ladder. Editing a role you do not hold and then giving it to yourself is "
    "closed by the assignment ceiling. "
    "test_giving_yourself_that_role_is_still_refused_by_the_grant_ceiling and "
    "test_giving_it_to_somebody_else_is_refused_by_the_same_ceiling pin that, "
    "beside test_a_restricted_addition_to_a_role_the_actor_holds_is_refused, "
    "test_the_same_addition_to_somebody_elses_role_goes_through and "
    "test_create_role_rejects_restricted_permission."
)

M04_FR003_LIMIT = (
    "Explicit DENY rows remain exempt from both grant controls because removing "
    "authority cannot escalate it. Seeded and provisioned system roles are the "
    "trusted bootstrap and carry direct restricted grants: Finance Admin every "
    "finance and payments key and Procurement Admin every procurement key, "
    "restricted ones included, and School Admin the restricted "
    "workflow.template.manage and workflow.group.manage. An addition to a role "
    "the editor does not hold reaches that role's current holders at once with "
    "no second person involved; the assignment ceiling governs who can be given "
    "the role, not what its existing holders receive."
)

M04_FR010_EVIDENCE_TAIL = (
    " visible_branch_ids_for answers the same question for many people in one "
    "query, reading each person's rows through the same _scope_from_rows "
    "function the single version uses, so a screen listing other people cannot "
    "read a grant differently from the gate. "
    "EveryBranchCarryingLookupIsAccountedForTests reads every request handler's "
    "source for a lookup by an outside id on a model carrying a branch, and "
    "fails on any that neither narrows to the caller's branches nor is listed "
    "with its reason in SETTLED_ELSEWHERE; its UNNARROWED registry of known "
    "holes is empty."
)

M04_FR010_LIMIT = (
    "Two list surfaces are still unnarrowed and are recorded in Needs "
    "Attention: the financial statements, which aggregate ledger lines through "
    "several relation paths, and this module's own role and assignment lists. "
    "The source audit sees lookups by id, not list querysets, so it reports "
    "neither. Whether a row with no branch is shared or is a scope of its own "
    "is a per-module reading rather than a platform rule: finance reads it "
    "inclusively, and procurement reads it exclusively on documents and "
    "inclusively on master data such as vendors and stock locations. It is "
    "argued at each call site and cannot be inferred from this module alone."
)

M04_FR017_REQUIREMENT = (
    "Editing what a role may do changes every holder. Ordinary additions and "
    "removals may be written directly by a role administrator, and so may a "
    "restricted addition to a role the editor does not hold. A restricted "
    "addition to a role the editor holds must be raised as a request with a "
    "justification and decided through an approval ladder that records who "
    "decided it before anything is granted."
)

M04_FR017_EVIDENCE = (
    "TenantRoleChangeRequest carries the requester, target role, required "
    "justification, PENDING / APPROVED / DENIED / APPLY_FAILED status, reviewer "
    "metadata, and normalised ADD or REMOVE deltas. raise_role_change_request "
    "creates the row and its deltas and submits it to the rbac.role_change "
    "ladder in one transaction, so a request the engine cannot route leaves no "
    "row behind. RoleChangeWorkflowHandler picks the tenant-less role-change "
    "template for a school, whose one stage names the school_admin system role "
    "with advance rule ANY and no skipping, or role-change-platform, owned by "
    "the platform tenant and naming xvs_platform_admin. The decision endpoint "
    "records a vote through record_action, which refuses anybody not on the "
    "stage's frozen approver list whatever keys they hold; a denial requires "
    "notes. The handler's on_approved callback is the only place a request's "
    "grants are written: apply_role_change_request replays the delta and calls "
    "set_role_access with the justification and request ID, under source "
    "approved_change_request, or self_approved_change_request with a "
    "self_approved flag when the requester decided it. Rejection, withdrawal "
    "and cancellation close the request as DENIED, and a reversal is refused "
    "once the grants have been applied. The handler declares "
    "allows_requester_self_approval, so the requester stays on the approver "
    "list, and refuses continue-without-approval, so an unstaffed stage parks "
    "rather than being released. Each request returns an approval block "
    "naming the ladder's status, the current stage, whether the reader can act "
    "and whether they raised it."
)

M04_FR017_ACCEPTANCE = (
    "Raising a request starts a ladder and grants nothing until it completes; "
    "approving applies the delta; rejecting closes the request and grants "
    "nothing; the approver of record is whoever cast the deciding vote. A head "
    "teacher who is the only School Admin holder can decide their own request, "
    "and that approval is recorded under its own audit source. A second School "
    "Admin holder is placed on the same stage. The exemption applies to this "
    "document type alone, and a look-alike role named School Admin without "
    "is_system_role confers no approval. RoleChangeLadderTests pins each case: "
    "test_raising_a_request_starts_a_ladder, "
    "test_nothing_is_granted_before_the_ladder_finishes, "
    "test_approving_applies_the_delta, "
    "test_rejecting_closes_the_request_and_grants_nothing, "
    "test_the_approver_of_record_is_whoever_voted, "
    "test_the_head_teacher_may_decide_her_own_request, "
    "test_self_approval_is_recorded_as_self_approval, "
    "test_a_second_administrator_gets_the_ordinary_two_person_review, "
    "test_the_exemption_does_not_leak_to_other_document_types, "
    "test_a_requester_on_an_ordinary_document_is_still_refused and "
    "test_a_look_alike_role_confers_no_approval."
)

M04_FR017_LIMIT = (
    "Separation of duties is not enforced for this document type. The "
    "requester is on the stage in every school and one approval completes it, "
    "so a head teacher can approve their own restricted addition even where a "
    "second administrator sits on the same stage; see Needs Attention. The "
    "reviewer grant ceiling that used to apply is gone, so a reviewer may "
    "approve a restricted key they do not hold. A request with no workflow "
    "instance is refused with 409 APPROVAL_MISSING and must be raised again. "
    "The approver list is frozen when the stage activates, so an administrator "
    "appointed afterwards cannot decide a request already waiting. Restricted "
    "removals may still be made directly because they reduce authority."
)

M04_FR019_SOURCE_OLD = (
    "set_role_access is the single supported boundary for direct role edits, "
    "approved requests, prebuilt provisioning, role suggestions and Super Admin "
    "permission reconciliation."
)

M04_FR019_SOURCE_NEW = (
    "set_role_access is the single supported boundary for direct role edits, "
    "approved role changes, prebuilt provisioning, role suggestions, "
    "plan-downgrade revocation and Super Admin permission reconciliation, and "
    "names each as its source: an approved role change is "
    "approved_change_request, or self_approved_change_request with a "
    "self_approved flag when its requester decided it, and a downgrade is "
    "plan_downgrade."
)

M04_FR019_LIMIT = (
    "school_id is a loose slug string rather than a foreign key, so the trail "
    "survives tenant deletion but is not joinable. The mirror into Module 5 can "
    "fail silently by design. Supplying grants to set_role_access without a "
    "deny set clears the role's explicit denies, which is the service's stated "
    "backward-compatible rule; the direct update path, an approved role change "
    "and a plan downgrade all supply grants alone, so each deletes any explicit "
    "deny on the role it touches, and the audit's before-and-after shows it "
    "going. The API exposes no deny field, so the denies at risk are those a "
    "migration or seeder wrote. Grants written by a data migration or by a "
    "seeder's sync into existing school roles, such as migration 0017's "
    "back-fill and seed_workflow_permissions, are direct row writes and carry "
    "no RBAC audit entry."
)

M04_FR021_EVIDENCE = (
    "PrebuiltRoleTemplate rows are read-only library definitions, five of "
    "them: school_admin, branch_admin, finance_admin, procurement_admin and "
    "teacher. Bursar is retired and finance_manager is renamed finance_admin, "
    "which claims every finance and payments key by prefix, as procurement_admin "
    "claims every procurement key, restricted ones included; the prefix "
    "attachment is additive and never removes a default. "
    "seed_workflow_permissions adds the workflow defaults to school_admin, "
    "finance_admin, procurement_admin and branch_admin and grants the same keys "
    "to the school roles already built from them, additively. "
    "provision_role_from_prebuilt copies a template into a tenant-owned, "
    "unlocked system role through set_role_access. A branch-scoped copy's key "
    "always carries the branch; its name carries the branch only once the "
    "school has more than one, and provisioning a role for a later branch "
    "renames each same-named sibling to say which branch it runs. School "
    "creation provisions the four tenant-wide templates and a Branch Admin per "
    "branch. Migration 0017 gave every existing school the tenant-wide roles it "
    "lacked, matched by key or by name and granting only keys a tenant may "
    "hold, and left a role the school already had untouched. vs_workflow "
    "migration 0013 removed the provisioned approver roles nobody held, found "
    "from the keys the published steps name rather than from a list."
)

M04_FR021_ACCEPTANCE = (
    "A school created today holds School Admin, Teacher, Finance Admin, "
    "Procurement Admin and a Branch Admin without an operator running a "
    "command "
    "(ANewSchoolGetsTheRolesCodeXShipsTests.test_the_full_set_is_provisioned_without_anybody_running_a_command). "
    "A one-branch school's Branch Admin names no branch, and a second branch "
    "makes every sibling say which (test_one_branch_means_the_role_does_not_name_it, "
    "test_a_second_branch_makes_every_sibling_say_which). Adoption remains "
    "idempotent, tenant-owned, editable, undeletable, and durably audited."
)

M04_FR021_LIMIT = (
    "The provisioning copy is not filtered by scope, so a library default for "
    "a key later reclassified PLATFORM is refused at the grant and provisioning "
    "fails. School creation provisions Finance Admin, so a school cannot be "
    "created on a database whose library attached finance.entity.create while "
    "it was tenant-holdable; this is traced in code, and no test builds Finance "
    "Admin from seeded finance keys. See Needs Attention. "
    "PrebuiltRoleTemplate.scope is written and read by nothing: Teacher is "
    "declared branch-scoped and is provisioned whole-tenant in every new "
    "school. A renamed sibling is not renamed back if a school returns to one "
    "branch."
)

M04_FR023_EVIDENCE_TAIL = (
    " Reachability is checked in aggregate rather than by a list of "
    "exceptions: UnenforcedKeysAreWithheldTests.test_no_registered_key_gates_nothing "
    "fails if any active registered key is gated by no view, and the ten "
    "approve keys that gated nothing, for journals, refunds, write-offs, "
    "payouts and procurement approval and their high-value variants, were "
    "deleted with their grants rather than hidden."
)

M04_FR023_LIMIT = (
    "Seeders remain the boundary's other authors. Reachability is checked in "
    "aggregate and scope classification still is not: no test or "
    "seed_all_permissions check proves that every active registered key "
    "carries a scope. Migration 0007 classified the original 344 / 40 split; "
    "keys added later depend on each seeder being correct, and a key "
    "reclassified later leaves any library default for it in place unless its "
    "migration removes that too."
)

M04_FR025_ENFORCE_OLD = (
    "platform.entitlements.enforce is on, at platform scope only so a school "
    "holding config.value.update cannot switch off its own gate, and a tenant "
    "holding no PACKAGE entitlement is treated as unprovisioned rather than "
    "unentitled and is never refused."
)

M04_FR025_ENFORCE_NEW = (
    "platform.entitlements.enforce is on: the configuration catalogue "
    "registers it ON, at platform scope only so a school holding "
    "config.value.update cannot switch off its own gate, and "
    "enforcement_enabled falls back to False only where the key is not "
    "registered at all, because a database with no catalogue refusing a "
    "paying school is the worse failure. The flag is read through get_config, "
    "which does not cache, so every gated request pays one query for it, "
    "named PLAN_GATE_FLAG_READ in the academics query budgets. A tenant "
    "holding no PACKAGE entitlement is treated as unprovisioned rather than "
    "unentitled and is never refused; that rule lives in vs_config beside the "
    "entitlements it reads, and the gate re-exports it."
)

M04_FR025_EVIDENCE_TAIL = (
    " Moving a school down a tier revokes the role grants the new tier does "
    "not reach: change_plan calls revoke_grants_the_plan_no_longer_reaches "
    "after the new entitlements are written, reads each granted key through "
    "plan_reader, and removes the lost ones through set_role_access under "
    "source plan_downgrade, so each revocation takes the role's lock, bumps "
    "its version and is audited. Permission groups are left as the school "
    "composed them."
)

M04_FR025_ACCEPTANCE_TAIL = (
    " A downgrade revokes what the new tier cannot reach, leaves what it still "
    "reaches, and records the revocation against the plan change "
    "(MovingDownATierTakesTheGrantsWithItTests). A seeded environment enforces "
    "and the switch stays a platform decision (TheCatalogueShipsEnforcementOnTests)."
)

M04_FR025_LIMIT = (
    "Only rbac_permission is gated: a view gated solely by "
    "rbac_group_permission or by HasAnyModuleAccess passes untouched. "
    "Enforcement is one platform switch rather than one per school, so a "
    "staged rollout would need a platform-writable per-tenant value. The "
    "catalogue deliberately does not read that switch: it reports what the "
    "school bought whether or not refusals are being issued. A downgrade "
    "revokes direct role grants only: a key carried by a group or a personal "
    "ALLOW override stays and is refused at the door, moving back up restores "
    "nothing, and because the revocation supplies grants without a deny set it "
    "also deletes the explicit denies on each role it touches (FR-019)."
)

M04_FR026_HEADING = "FR-026  Keep a Way Back Into Role Administration"
M04_FR026_HEADER = "FR-026 | Implemented with limits"

M04_FR026_REQUIREMENT = (
    "A school must not be able to switch off the last role through which "
    "anybody could administer roles, because every screen that could undo it "
    "sits behind the role just switched off, and recovering from it takes "
    "CodeX editing the database."
)

M04_FR026_EVIDENCE = (
    "_reject_last_way_in runs in the role detail serializer's validation "
    "whenever a save moves a role's status away from ACTIVE. It asks the "
    "invariant rather than the symptom: after the save, some other ACTIVE role "
    "in the tenant must carry a direct grant of school.roles.update or "
    "platform.roles.update and be held by at least one ACTIVE assignment. A "
    "role nobody holds is not counted as a way in. Where no such role "
    "survives, the save is refused with 400 under status, telling the caller "
    "to give somebody else a role that can manage roles first."
)

M04_FR026_ACCEPTANCE = (
    "Holy Cross has one head teacher, holding School Admin, who cannot take "
    "School Admin out of use; neither can an administrator retiring the "
    "only other role that could let anybody back in. Retiring a role while "
    "another still lets somebody in succeeds, so the guard does not become a "
    "ban. test_a_school_cannot_retire_its_only_way_back_in, "
    "test_a_role_can_be_retired_while_another_still_lets_somebody_in and "
    "test_a_role_nobody_holds_is_not_a_way_back_in pin the three cases."
)

M04_FR026_LIMIT = (
    "Only a status change is checked. Removing the role-update key through a "
    "permission edit, revoking the last assignment, and deleting a custom role "
    "that carries the key reach the same lockout unchecked. A holder is "
    "counted by assignment status alone, so a suspended account or a grant at "
    "a branch out of service still counts as a way in, and a key carried only "
    "through a group is not counted, which refuses more than it needs to. See "
    "Needs Attention."
)

M04_WITHDRAW_DEACTIVATE = (
    "Immediate, and invisible in the assignment rows, which still read ACTIVE. "
    "Refused where it would leave no active assignment to an active role "
    "carrying a role-update key (FR-026)."
)

M04_WITHDRAW_DOWNGRADE = {
    0: "Move the school down a plan tier",
    1: "Every direct role grant the new tier does not reach, on every role in "
       "the school.",
    2: "Immediate, through set_role_access under source plan_downgrade, so "
       "each role is locked, versioned and audited. Groups and overrides are "
       "left alone, and moving back up restores nothing.",
}

M04_MODEL_PREBUILT = (
    "Read-only to tenants. Five templates. scope is written and read by "
    "nothing; tier is read only for ordering. Prefix defaults are attached "
    "additively and never pruned."
)

M04_MODEL_CHANGE_REQUEST = (
    "PENDING, APPROVED, DENIED and APPLY_FAILED. Justification required; "
    "target role tenant-pinned. Routed as workflow document type "
    "rbac.role_change: the engine's instance holds the stage, the eligible "
    "approvers and who acted, and these statuses summarise it for the roles "
    "screen."
)

M04_API_ROLES = (
    "The tenant's role catalogue. Reading accepts school.roles.view, "
    "platform.roles.view or workflow.template.manage; creating takes "
    "school.roles.create or platform.roles.create. permission_keys or "
    "group_ids requires reason. A restricted permission is refused at "
    "creation; the creator, who does not hold the new role, may add it by "
    "editing the role afterwards."
)

M04_API_ROLE_DETAIL = (
    "One role, including permission_keys, group_ids and held_by_me, which says "
    "whether the reader holds an active assignment to it, so the screen can "
    "show before saving whether a restricted addition will save or need "
    "approval. Supplying either access field requires reason and writes one "
    "transactional before-and-after audit. A restricted addition to a role the "
    "reader holds is refused with 409 RESTRICTED_NEEDS_APPROVAL listing "
    "restricted_additions; to anybody else's role it saves. Retiring the last "
    "way back into role administration is refused. Delete is blocked for "
    "system or locked roles."
)

M04_API_REQUESTS = (
    "Raise or list change requests. The update keys raise; view keys list. "
    "Raising submits the request to the rbac.role_change ladder in the same "
    "transaction. Each request carries an approval block: the ladder's status, "
    "the current stage label, whether the reader can act, and whether they "
    "raised it."
)

M04_API_QUEUE = (
    "The tenant's role change requests, filterable by status and target role. "
    "school.roles.view or platform.roles.view. Platform callers may assert the "
    "target tenant."
)

M04_API_DECIDE = (
    "Approve or deny, recorded as a vote on the ladder's active stage. "
    "school.roles.approve or platform.roles.approve reaches the endpoint, and a "
    "place on the stage's frozen approver list makes the vote count. A denial "
    "requires notes. A request with no workflow instance is refused with 409 "
    "APPROVAL_MISSING."
)

M04_ANSWER_RESTRICTED = (
    "A group member or ALLOW override: 400 naming the restricted key. A role "
    "addition: 409 RESTRICTED_NEEDS_APPROVAL with restricted_additions, only "
    "where it lands on a role the caller holds; at creation, 400."
)

M04_ANSWER_CEILING = (
    "403, whoever the recipient is, the assigner included. The same ceiling "
    "applies to assign, replace, account creation, and draft submission."
)

M04_ANSWER_REQUESTER = {
    0: "A requester deciding their own role-change request",
    1: "Accepted: this document type lets the requester decide, and the "
       "approval is recorded as self_approved_change_request. A caller who is "
       "not on the stage's frozen approver list is refused by the engine under "
       "its own code, whatever keys they hold.",
}

M04_ANSWER_NEW = [
    {0: "Retiring the last active role, held by somebody, that carries a "
        "role-update key",
     1: "400 under status, naming what to do first. Retiring any other role is "
        "unaffected."},
    {0: "A decision on a request that has no workflow instance",
     1: "409 APPROVAL_MISSING, asking for the request to be raised again."},
    {0: "A request raised where the role-change ladder cannot be resolved",
     1: "Refused by the engine under its own typed code, and no request row is "
        "left behind."},
]

M04_DEP_WORKFLOW = (
    "Consumes resolve_users_with_permission to name approvers, holds "
    "workflow.template.manage, which this module accepts on the role list "
    "endpoint only, and decides role changes: a TenantRoleChangeRequest is "
    "routed as document type rbac.role_change through the role-change and "
    "role-change-platform templates, and the handler this module registers "
    "applies the delta on approval. The engine's requester self-approval "
    "exemption is declared for this document type alone."
)

M04_DEP_CONFIG_TAIL = (
    " The gate reads that switch through get_config on every gated request, "
    "uncached."
)

M04_DEP_SCHOOLS = (
    "Consumes school_admin and branch_admin during required administrator "
    "creation, and provisions every tenant-wide prebuilt role, Teacher, "
    "Finance Admin and Procurement Admin beside School Admin, when a school is "
    "created. Its plan change calls this module's set_role_access to revoke "
    "grants a lower tier does not reach. Its school-app migration still "
    "guarantees the two administrator template identities on fresh "
    "installations; the RBAC seed command owns the templates' access defaults."
)

M04_VERIFY_REWRITES = {
    "•  Both escalations end to end": (
        "•  Both escalations end to end: scope refuses a platform-only key on "
        "role and override paths; restriction refuses direct "
        "payments.payout.create at creation, in a group and as an ALLOW "
        "override, refuses a restricted addition to a role the editor holds "
        "with 409 RESTRICTED_NEEDS_APPROVAL, and lets the same addition to "
        "somebody else's role save while the assignment ceiling still refuses "
        "giving that role to the editor."
    ),
    "•  Bootstrap and ceiling": (
        "•  Bootstrap and ceiling: a platform actor may still grant a "
        "platform-scoped key inside CodeX; the Vision super administrator may "
        "assign the first restricted holder; every other assigner is limited "
        "to restricted keys they already hold, whoever receives the role."
    ),
    "•  Self-service refused on both controls": (
        "•  Self-service on the two controls: override create and lift reject "
        "both the actor and the effective user; a role change may be decided "
        "by its requester and is then recorded under "
        "self_approved_change_request, while the same requester on any other "
        "document type is refused."
    ),
}

M04_VERIFY_NEW = [
    "•  The last way back in: retiring the only active, held role carrying a "
    "role-update key is refused, retiring one while another still lets "
    "somebody in succeeds, and a role nobody holds does not count as a way in.",
    "•  The role-change ladder: nothing granted before it completes, the delta "
    "applied on approval, rejection granting nothing, and the requester "
    "exemption confined to rbac.role_change.",
    "•  A plan downgrade: it revokes what the new tier cannot reach, leaves "
    "what it can, and records the revocation against the plan change.",
    "•  The branch source audit: every lookup by an outside id on a "
    "branch-carrying model inside a request handler narrows to the caller's "
    "branches or is listed with its reason.",
]

M04_NEED_PREBUILT_SCOPE = (
    "PrebuiltRoleTemplate.scope is populated and read by nothing. It offers "
    "institution, branch, class and portal, the seeder sets it on all five "
    "templates, and no code reads it: Teacher is declared branch-scoped and "
    "school creation provisions it whole-tenant in every new school, which is "
    "the case the field looks as if it prevents. A field that looks like a "
    "control and is not one is how the last two escalations were possible."
)

M04_NEED_SEED_SCOPE = (
    "The platform registry API requires scope, and an aggregate test proves "
    "every active key gates some view, but nothing proves every active seeded "
    "key is classified. A forgotten seeder value remains safe at grant time, "
    "where it is refused, but the error appears after deployment rather than "
    "during seeding."
)

M04_NEED_SCHOOL_CREATION = {
    1: "School creation fails where the role library still gives Finance "
       "Admin finance.entity.create. Migration 0018 made that key PLATFORM and "
       "removed school roles' grants of it, but not the prebuilt default: the "
       "prebuilt seeder attaches every finance key by prefix and never removes "
       "one, and provisioning copies the defaults without filtering by scope. "
       "On a database whose library was seeded while the key was "
       "tenant-holdable, provisioning Finance Admin is refused at the grant, "
       "and because creation provisions Finance Admin, creating a school fails "
       "and rolls back. Traced in code at the baseline and not observed; no "
       "test builds Finance Admin from seeded finance keys, so nothing "
       "catches it.",
    2: "Delete the PrebuiltRolePermission rows for keys a tenant may not hold, "
       "as migration 0019 does for tickets.ticket.assign, and filter the "
       "provisioning copy to tenant-holdable keys so a later reclassification "
       "cannot do this again.",
    3: "FR-021",
}

M04_NEED_SELF_APPROVAL = {
    1: "A restricted addition to your own role needs only your own approval. "
       "The role-change ladder keeps the requester on its one stage and "
       "completes on one approval in every school, so a head teacher can raise "
       "a restricted addition to School Admin and approve it themselves even "
       "where a second administrator sits on the same stage. The approval is "
       "recorded as self_approved_change_request, which makes it visible "
       "afterwards, not reviewed beforehand.",
    2: "Exclude the requester whenever the stage resolves another eligible "
       "approver, keeping the exemption for the school with one "
       "administrator, or record here that one-person approval is the "
       "accepted rule.",
    3: "FR-017",
}

M04_NEED_LAST_WAY = {
    1: "The last-way-in guard covers one door. Retiring a role is refused "
       "where it would leave no active assignment to an active role carrying "
       "a role-update key, but removing that key through a permission edit, "
       "revoking the last such assignment, and deleting a custom role that "
       "carries it are not checked, and each locks the school out of role "
       "administration the same way. The guard also counts a holder whose "
       "account is suspended, or whose only grant sits at a branch out of "
       "service, as a way in.",
    2: "State the invariant once and ask it on every path that can remove the "
       "last holder or the last key: permission edits, revocation, "
       "replacement and deletion, counting only holders who can sign in.",
    3: "FR-026",
}

M04_REMOVAL = (
    "REMOVAL RULE\n"
    "• An item leaves Needs Attention only when implementation and relevant "
    "verification resolve it fully.\n"
    "• An item that is partly resolved is rewritten to the risk that remains, "
    "not deleted.\n"
    "• Against v1.17 no item leaves. Three are added: school creation failing "
    "on a stale library default, self-approval of a restricted addition to "
    "your own role, and the last-way-in guard covering one door. The "
    "prebuilt-scope and seeder-classification items are rewritten to what "
    "now holds."
)

M04_TRACE_LEAD = (
    f"MRD v{MRD_VERSION} records Module 4 as Roles & Permissions (RBAC), Phase "
    "V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two "
    "capability entries. The module number, name, phase, states and ownership "
    "agree with this revision. Each entry maps below; the last-way-in guard is "
    "the entry this revision adds."
)

M04_TRACE_STATES = {
    "Role templates": (
        "Implemented. Five templates, School Admin, Branch Admin, Finance "
        "Admin, Procurement Admin and Teacher, provisioned as tenant-owned "
        "copies; school creation provisions all of them, and migration 0017 "
        "gave every existing school the tenant-wide roles it lacked. A library "
        "default for a key later reclassified PLATFORM breaks provisioning "
        "(Needs Attention P1)."
    ),
    "Create and update roles": (
        "Implemented. Ordinary permission and group changes require a reason "
        "and cross one locked, transactional mutation service. A restricted "
        "addition saves directly to a role the editor does not hold and is "
        "refused with 409 RESTRICTED_NEEDS_APPROVAL on one they hold; the role "
        "detail says which through held_by_me."
    ),
    "Assign, revoke, and replace user roles": (
        "Implemented. Every assignment path refuses a role whose restricted "
        "keys the assigner does not hold, whoever receives it, which closes "
        "editing a role and then giving it to yourself."
    ),
    "Role-change request workflow": (
        "Implemented. A workflow ladder decides it: one School Admin stage, "
        "ANY, never skipped, with the requester allowed to decide their own "
        "and recorded as such. The reviewer grant ceiling is gone. "
        "Self-approval where a second administrator exists is at P2."
    ),
    "Permission and assignment audit history": (
        "Implemented. Durable and append-only, with actor, reason, source "
        "(approved, self-approved and plan-downgrade changes among them), "
        "approval reference where applicable, and complete direct grant, "
        "direct deny, group and effective combined before-and-after sets. The "
        "central trail remains a best-effort mirror."
    ),
    "Permission seeding commands": (
        "Implemented with limits. API-created keys require explicit scope, "
        "approval keys are restricted, system roles provide bootstrap, groups "
        "exclude restricted members, and every active key must gate some "
        "view; aggregate scope classification remains open."
    ),
    "Permission mapped to the capability that governs it": (
        "Implemented. Every registered key carries a verdict. Five keys no "
        "school can act on are platform-only, todo.task.view, .manage and "
        ".assign, finance.entity.create and tickets.ticket.assign, and the "
        "data export band sits at Core, on every plan."
    ),
    "Plan depth gate with its own refusal": (
        "Implemented and in service. The catalogue registers the switch ON, a "
        "school is refused under PLAN_UPGRADE_REQUIRED once its plan is "
        "applied, and a downgrade revokes the direct role grants the new tier "
        "does not reach."
    ),
}

M04_TRACE_ENTITY_TAIL = (
    " A source audit keeps every lookup by id narrowed or accounted for."
)

M04_TRACE_RENAMES = {
    "School permission group catalogue with declared branch reach": (
        "School permission group catalogue with declared branch reach, "
        "reaching finance, procurement and payments"
    ),
    "Role builder offers only what the gate allows": (
        "A role builder that offers only what the plan allows"
    ),
}

M04_TRACE_NEW = {
    0: "A school cannot retire the last way back into role administration",
    1: "FR-026",
    2: "Implemented with limits. Retiring a role is refused where nobody would "
       "still hold an active role carrying a role-update key; removing the "
       "key, revoking the last holder and deleting a role are not checked "
       "(P2).",
}

M04_RECONCILE = (
    "MRD RECONCILIATION\n"
    f"• MRD v{MRD_VERSION} lists Module 4 as Roles & Permissions (RBAC), Phase "
    "V1, Backend Complete, In use Complete, code vs_rbac, with twenty-two "
    "capability entries. The module number, name, phase, states and ownership "
    "agree with this revision.\n"
    "• The count rises by one. Refusing to retire the last way back into role "
    "administration is a rule a school meets on the roles screen, with a "
    "refusal of its own, in the same way the reserved approval-role keys "
    "are.\n"
    "• The count is otherwise unchanged. Deciding a role change through a "
    "workflow ladder restates the role-change request entry; refusing a "
    "restricted addition only on the editor's own role restates the create "
    "and update entry; revoking out-of-plan grants on a downgrade and "
    "registering enforcement ON restate the plan gate entry.\n"
    "• Backend evidence only. Neither document claims deployment, production "
    "adoption, or data migration completion."
)

M04_CHANGE_SUMMARY = (
    "Records how a role change is decided now, with four other RBAC changes "
    "the previous revisions had not reached. A restricted addition to a role "
    "was refused for everybody and could be approved only by a reviewer who "
    "already held the key, so a school whose one approver was the head teacher "
    "could raise a request nobody in the building could close. A role change "
    "is now a workflow ladder: raising it submits it in the same transaction, "
    "one School Admin stage decides it, the engine's approval callback applies "
    "the delta, and the requester may decide their own, recorded under "
    "self_approved_change_request. A restricted addition is refused, with 409 "
    "RESTRICTED_NEEDS_APPROVAL and the keys named, only where it lands on a "
    "role the editor holds, and the role detail says whether the reader holds "
    "it; the assignment ceiling still refuses giving that role to yourself, "
    "which closes the two-step route. Retiring the last role that could let "
    "anybody back into role administration is refused. A branch role names its "
    "branch only once a school has more than one, school creation provisions "
    "every tenant-wide library role, and migration 0017 gave existing schools "
    "the roles they lacked. Five keys no school can act on are PLATFORM, with "
    "their school grants removed first. A plan downgrade revokes out-of-plan "
    "role grants under source plan_downgrade, and the enforcement switch is "
    "registered ON. FR-002, FR-003, FR-010, FR-017, FR-019, FR-021, FR-023 and "
    "FR-025 are updated and FR-026 is added, with the actors, workflows, data "
    "model, API contracts, dependencies, minimum verification, Needs Attention "
    "and MRD traceability following; the approval queue's key is corrected to "
    "the view keys it has always read with. Three Needs Attention items are "
    "added, the first being that creating a school fails where the role "
    "library still gives Finance Admin finance.entity.create. Module 4 remains "
    "Backend Complete and In use Complete, with twenty-two capability entries "
    "against MRD v2.76. Verified by the vs_rbac suite, 541 tests OK at "
    "6611856, including RoleChangeLadderTests and the role detail tests named "
    "in FR-003 and FR-026; the provisioning and downgrade tests named in "
    "FR-021 and FR-025 sit in schools.vs_schools and were read rather than "
    "re-run for this revision. Backend evidence only; nothing here is deployed."
)


def insert_fr026(doc, fr023_table, fr025_table) -> None:
    """Deep-copy FR-023's heading and amber table, and place them after FR-025."""
    heading = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("FR-023  "):
            heading = paragraph
            break
    if heading is None:
        raise ValueError("FR-023 heading not found")

    new_heading = copy.deepcopy(heading._p)
    new_table = copy.deepcopy(fr023_table._tbl)
    fr025_table._tbl.addnext(new_heading)
    new_heading.addnext(new_table)

    for paragraph in doc.paragraphs:
        if paragraph._p is new_heading:
            set_run_text(paragraph, M04_FR026_HEADING)
            break
    else:
        raise ValueError("FR-026 heading not placed")

    for table in doc.tables:
        if table._tbl is new_table:
            set_run_text(table.rows[0].cells[0].paragraphs[0], M04_FR026_HEADER)
            write_kept(table.rows[1].cells[1], M04_FR026_REQUIREMENT)
            write_kept(table.rows[EVIDENCE_ROW].cells[1], M04_FR026_EVIDENCE)
            write_kept(table.rows[ACCEPTANCE_ROW].cells[1], M04_FR026_ACCEPTANCE)
            write_kept(table.rows[LIMIT_ROW].cells[1], M04_FR026_LIMIT)
            return
    raise ValueError("FR-026 table not placed")


def patch_m04(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M04 Roles and Permissions RBAC Functional Requirements "
        f"Document v{M04_TARGET}"
    )
    t = {name: doc.tables[index] for name, index in M04_T.items()}
    for label in ("FR-002", "FR-003", "FR-010", "FR-017", "FR-019", "FR-021",
                  "FR-023", "FR-025"):
        require_fr(t[label], label)

    replace_cover_version(t["cover"], M04_SOURCE, M04_TARGET)
    set_control(t["control"], "Version", M04_TARGET)
    set_control(t["control"], "Review date", REVIEW_DATE)
    set_control(t["control"], "Source scope", M04_SOURCE_SCOPE)
    set_control(t["control"], "Code inspected", M04_CODE_INSPECTED)
    set_control(t["control"], "MRD baseline", M04_MRD_BASELINE)

    write_kept(find_row(t["contents"], "4. Functional Requirements").cells[1],
               "FR-001 to FR-026, each with the code evidence behind it")
    write_kept(find_row(t["contents"], "10. MRD Traceability").cells[1],
               f"Agreement with MRD v{MRD_VERSION}'s twenty-two capability entries")

    replace_body_paragraph(doc, "Module 4 is the platform's authorization boundary",
                           M04_PURPOSE)
    replace_body_paragraph(doc, "Five boundaries hold this module up", M04_BOUNDARIES)

    write_kept(find_row(t["in_scope"], "Prebuilt role library").cells[1],
               M04_SCOPE_PREBUILT)
    write_kept(find_row(t["in_scope"], "Assignments").cells[1], M04_SCOPE_ASSIGNMENTS)
    write_kept(find_row(t["in_scope"], "Role-change approval").cells[1],
               M04_SCOPE_ROLE_CHANGE)

    school = find_row(t["actors"], "School administrator")
    write_kept(school.cells[1], M04_ACTOR_SCHOOL_DOES)
    write_kept(school.cells[2], M04_ACTOR_SCHOOL_KEYS)
    write_kept(find_row(t["actors"], "CodeX platform staff").cells[1],
               M04_ACTOR_PLATFORM_DOES)

    append_kept(t["FR-002"].rows[EVIDENCE_ROW].cells[1], M04_FR002_EVIDENCE_TAIL)
    append_kept(t["FR-002"].rows[ACCEPTANCE_ROW].cells[1], M04_FR002_ACCEPTANCE_TAIL)

    write_kept(t["FR-003"].rows[EVIDENCE_ROW].cells[1], M04_FR003_EVIDENCE)
    write_kept(t["FR-003"].rows[ACCEPTANCE_ROW].cells[1], M04_FR003_ACCEPTANCE)
    write_kept(t["FR-003"].rows[LIMIT_ROW].cells[1], M04_FR003_LIMIT)

    append_kept(t["FR-010"].rows[EVIDENCE_ROW].cells[1], M04_FR010_EVIDENCE_TAIL)
    write_kept(t["FR-010"].rows[LIMIT_ROW].cells[1], M04_FR010_LIMIT)

    write_kept(t["FR-017"].rows[1].cells[1], M04_FR017_REQUIREMENT)
    write_kept(t["FR-017"].rows[EVIDENCE_ROW].cells[1], M04_FR017_EVIDENCE)
    write_kept(t["FR-017"].rows[ACCEPTANCE_ROW].cells[1], M04_FR017_ACCEPTANCE)
    write_kept(t["FR-017"].rows[LIMIT_ROW].cells[1], M04_FR017_LIMIT)

    edit_kept(t["FR-019"].rows[EVIDENCE_ROW].cells[1],
              M04_FR019_SOURCE_OLD, M04_FR019_SOURCE_NEW)
    write_kept(t["FR-019"].rows[LIMIT_ROW].cells[1], M04_FR019_LIMIT)

    write_kept(t["FR-021"].rows[EVIDENCE_ROW].cells[1], M04_FR021_EVIDENCE)
    write_kept(t["FR-021"].rows[ACCEPTANCE_ROW].cells[1], M04_FR021_ACCEPTANCE)
    write_kept(t["FR-021"].rows[LIMIT_ROW].cells[1], M04_FR021_LIMIT)

    append_kept(t["FR-023"].rows[EVIDENCE_ROW].cells[1], M04_FR023_EVIDENCE_TAIL)
    write_kept(t["FR-023"].rows[LIMIT_ROW].cells[1], M04_FR023_LIMIT)

    edit_kept(t["FR-025"].rows[EVIDENCE_ROW].cells[1],
              M04_FR025_ENFORCE_OLD, M04_FR025_ENFORCE_NEW)
    append_kept(t["FR-025"].rows[EVIDENCE_ROW].cells[1], M04_FR025_EVIDENCE_TAIL)
    append_kept(t["FR-025"].rows[ACCEPTANCE_ROW].cells[1], M04_FR025_ACCEPTANCE_TAIL)
    write_kept(t["FR-025"].rows[LIMIT_ROW].cells[1], M04_FR025_LIMIT)

    withdrawn = t["withdrawn"]
    write_kept(find_row(withdrawn, "Deactivate the role").cells[2],
               M04_WITHDRAW_DEACTIVATE)
    anchor = find_row(withdrawn, "Suspend or lock the account")
    fill(row_for(withdrawn, clone_row(anchor, before=anchor)), M04_WITHDRAW_DOWNGRADE)

    write_kept(find_row(t["model"], "PrebuiltRoleTemplate").cells[2], M04_MODEL_PREBUILT)
    write_kept(find_row(t["model"], "TenantRoleChangeRequest").cells[2],
               M04_MODEL_CHANGE_REQUEST)

    routes = t["api_routes"]
    write_kept(find_row(routes, "GET, POST /rbac/tenants/{slug}/roles/").cells[1],
               M04_API_ROLES)
    write_kept(find_row(routes, "GET, PATCH, DELETE /rbac/tenants/{slug}/roles/{key}/").cells[1],
               M04_API_ROLE_DETAIL)
    write_kept(find_row(routes, "GET, POST /rbac/tenants/{slug}/role-change-requests/").cells[1],
               M04_API_REQUESTS)
    write_kept(find_row(routes, "GET /rbac/tenants/{slug}/role-change-requests/approval/").cells[1],
               M04_API_QUEUE)
    write_kept(find_row(routes, "POST /rbac/tenants/{slug}/role-change-requests/{id}/decide/").cells[1],
               M04_API_DECIDE)

    answers = t["api_answers"]
    write_kept(find_row(answers, "A direct role addition, group member").cells[1],
               M04_ANSWER_RESTRICTED)
    write_kept(find_row(answers, "A restricted role assigned by somebody").cells[1],
               M04_ANSWER_CEILING)
    fill(find_row(answers, "A requester trying to approve"), M04_ANSWER_REQUESTER)
    for values in M04_ANSWER_NEW:
        last = answers.rows[-1]
        fill(row_for(answers, clone_row(last, after=last)), values)

    deps = t["dependencies"]
    write_kept(find_row(deps, "Module 7, Workflow and Approval").cells[1], M04_DEP_WORKFLOW)
    append_kept(find_row(deps, "vs_config").cells[1], M04_DEP_CONFIG_TAIL)
    write_kept(find_row(deps, "Module 1, School and Branch Management").cells[1],
               M04_DEP_SCHOOLS)

    for prefix, text in M04_VERIFY_REWRITES.items():
        replace_body_paragraph(doc, prefix, text)
    for text in M04_VERIFY_NEW:
        insert_bullet_before(doc, "•  The empty-list response shape", text)

    needs = t["needs"]
    for row in needs.rows:
        if row.cells[1].text.strip().startswith("PrebuiltRoleTemplate.scope"):
            write_kept(row.cells[1], M04_NEED_PREBUILT_SCOPE)
            break
    else:
        raise ValueError("Prebuilt scope need not found")
    for row in needs.rows:
        if row.cells[1].text.strip().startswith("The platform registry API now requires scope"):
            write_kept(row.cells[1], M04_NEED_SEED_SCOPE)
            break
    else:
        raise ValueError("Seeder scope need not found")
    first_p1 = needs.rows[1]
    if first_p1.cells[0].text.strip() != "P1":
        raise ValueError("First need is not P1")
    fill(row_for(needs, clone_row(first_p1, after=first_p1)), M04_NEED_SCHOOL_CREATION)
    for values in (M04_NEED_SELF_APPROVAL, M04_NEED_LAST_WAY):
        last = needs.rows[-1]
        if last.cells[0].text.strip() != "P2":
            raise ValueError("Last need is not P2")
        fill(row_for(needs, clone_row(last, after=last)), values)

    set_run_text(t["removal"].rows[0].cells[0].paragraphs[0], M04_REMOVAL)

    trace = t["trace"]
    for capability, state in M04_TRACE_STATES.items():
        write_kept(find_row(trace, capability).cells[2], state)
    append_kept(find_row(trace, "Entity-aware permission evaluation").cells[2],
                M04_TRACE_ENTITY_TAIL)
    for old, new in M04_TRACE_RENAMES.items():
        write_kept(find_row(trace, old).cells[0], new)
    last = trace.rows[-1]
    fill(row_for(trace, clone_row(last, after=last)), M04_TRACE_NEW)
    replace_body_paragraph(doc, "MRD v2.53 records Module 4", M04_TRACE_LEAD)
    set_run_text(t["reconcile"].rows[0].cells[0].paragraphs[0], M04_RECONCILE)

    prepend_change_log(t["changes"], M04_TARGET, SHORT_DATE, M04_CHANGE_SUMMARY)

    # Last, because it adds a table and every index above is already bound.
    insert_fr026(doc, t["FR-023"], t["FR-025"])
    keep_headings_with_their_tables(doc)

    finish(doc, output, title, M04_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 6 - Configuration & Capability Management
# ═════════════════════════════════════════════════════════════════════════════

M06_DIR = "06-configuration-and-capability"
M06_STEM = (
    "XVS_M06_Configuration_and_Capability_Management_Functional_Requirements_Document"
)
M06_SOURCE, M06_TARGET = "1.0", "1.1"

M06_T = {"cover": 0, "control": 1, "FR-003": 9, "FR-022": 28,
         "dependencies": 36, "needs": 37, "changes": 39}

M06_FR003_LIMIT = (
    "Precedence is fixed. There is no per-definition policy that would let one "
    "setting resolve platform-first. get_config does not cache, so a caller "
    "reading a setting on every request pays for the read each time; the plan "
    "gate's enforcement flag is that cost on every gated endpoint, named "
    "PLAN_GATE_FLAG_READ in the academics query budgets."
)

M06_FR022_EVIDENCE = (
    "All nineteen keys in ConfigPermissions are PLATFORM-scoped, which the "
    "RBAC grant guard enforces independently of role composition, so a school "
    "role cannot hold one and cannot be given one. "
    "platform.entitlements.enforce, the switch the plan gate reads, is "
    "declared at platform scope for the same reason: a school-scoped switch "
    "would let a school with config.value.update disable its own plan gate. "
    "seed_config_catalogue registers it ON, so an environment built from the "
    "catalogue refuses what a school's plan does not reach. get_config returns "
    "the caller's own default where a key is not registered at all, and the "
    "plan gate's default is False, so a database with no configuration "
    "catalogue does not refuse a paying school. The one school-facing route, "
    "FR-015, carries no key and is read-only."
)

M06_FR022_ACCEPTANCE_TAIL = (
    " A seeded environment enforces, and the switch's allowed scopes are "
    "platform alone (vs_rbac's TheCatalogueShipsEnforcementOnTests: "
    "test_a_seeded_environment_enforces and test_it_stays_a_platform_decision)."
)

M06_FR022_LIMIT = (
    "The reservation is total. A setting a school ought to own for itself has "
    "no route today and would need a tenant-scoped key introduced "
    "deliberately. The enforcement switch's own description says a school can "
    "be excused from enforcement, which its platform-only allowed scope does "
    "not permit; see Needs Attention."
)

M06_DEP_RBAC = (
    "Enforces the nineteen PLATFORM-scoped keys, and consumes this module's "
    "evaluation in the plan gate, reading platform.entitlements.enforce "
    "through get_config on every gated request. The permission-to-band map "
    "lives there, not here."
)

M06_VERIFICATION = (
    "The requirements above name the tests behind each behaviour; this "
    "revision was traced from the code at the baseline named in Document "
    "Control and did not re-run them. The principal consumer, vs_rbac, ran 541 "
    "tests OK at 6611856, including the plan gate, the agreement between the "
    "role builder and that gate, and TheCatalogueShipsEnforcementOnTests, "
    "which reads the seeded catalogue rather than a test's own definition. "
    "Backend evidence only; nothing here is deployed."
)

M06_NEED_SWITCH = (
    "platform.entitlements.enforce is declared at platform scope, which is "
    "what stops a school disabling its own gate, and it also means "
    "enforcement is on for everybody or nobody. A staged rollout, or holding "
    "one school harmless while a mis-banding is corrected, is not "
    "expressible. The definition's own description, which an operator reads "
    "in the console, says the opposite: that the switch stays readable per "
    "school so a school can be excused without switching it off for the "
    "platform. A school-scoped value for it is refused by the definition's "
    "allowed scopes."
)

M06_NEED_SWITCH_COMPLETION = (
    "A platform-writable per-tenant value, so the switch stays out of a "
    "school's hands while still being settable per school. Until then, "
    "correct the description so it does not promise an excusal the "
    "declaration refuses."
)

M06_TRACE_LEAD = (
    f"Module 6 of XVS Module Requirements Document v{MRD_VERSION} lists "
    "twenty-six capabilities. Each maps to the requirements above; shipping the "
    "enforcement switch ON strengthens the plan-gate reservation in FR-022 "
    "without adding a capability or changing the count."
)

M06_CHANGE_SUMMARY = (
    "Records that the configuration catalogue ships the plan gate's "
    "enforcement switch ON. platform.entitlements.enforce was registered off, "
    "so an environment built from the catalogue behaved unlike the ones "
    "already enforcing; seed_config_catalogue now registers it true, while "
    "get_config still returns the plan gate's own False where the key is not "
    "registered at all, so a database with no catalogue does not refuse a "
    "paying school. The switch stays platform-scoped, pinned by "
    "TheCatalogueShipsEnforcementOnTests. get_config does not cache, so the "
    "gate pays one read per gated request, which the academics query budgets "
    "now name. FR-003's limit, FR-022, the vs_rbac dependency, the "
    "verification evidence and Needs Attention item 2 are updated; the last "
    "records that the switch's own description promises a per-school "
    "excusal its allowed scopes refuse. Module 6 remains Backend Complete and "
    "In use Complete with twenty-six capability entries against MRD v2.76. "
    "Backend evidence only; nothing here is deployed."
)


def patch_m06(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M06 Configuration and Capability Management Functional "
        f"Requirements Document v{M06_TARGET}"
    )
    t = {name: doc.tables[index] for name, index in M06_T.items()}
    require_fr(t["FR-003"], "FR-003")
    require_fr(t["FR-022"], "FR-022")

    replace_cover_version(t["cover"], M06_SOURCE, M06_TARGET)
    set_control(t["control"], "Version", M06_TARGET)
    set_control(t["control"], "Review date", REVIEW_DATE)
    set_control(t["control"], "Code baseline", CODE_BASELINE)
    set_control(
        t["control"], "Source MRD",
        f"XVS Module Requirements Document v{MRD_VERSION} | Module 6, "
        "twenty-six capability entries",
    )

    write_kept(t["FR-003"].rows[LIMIT_ROW].cells[1], M06_FR003_LIMIT)

    write_kept(t["FR-022"].rows[EVIDENCE_ROW].cells[1], M06_FR022_EVIDENCE)
    append_kept(t["FR-022"].rows[ACCEPTANCE_ROW].cells[1], M06_FR022_ACCEPTANCE_TAIL)
    write_kept(t["FR-022"].rows[LIMIT_ROW].cells[1], M06_FR022_LIMIT)

    write_kept(find_row(t["dependencies"], "vs_rbac").cells[1], M06_DEP_RBAC)
    replace_body_paragraph(doc, "The module's own suite runs green", M06_VERIFICATION)

    switch = find_row(t["needs"], "2")
    if not switch.cells[1].text.strip().startswith("The plan gate is one switch"):
        raise ValueError("Need 2 is not the enforcement switch")
    write_kept(switch.cells[2], M06_NEED_SWITCH)
    write_kept(switch.cells[3], M06_NEED_SWITCH_COMPLETION)

    replace_body_paragraph(doc, "Module 6 of XVS Module Requirements Document",
                           M06_TRACE_LEAD)

    prepend_change_log(t["changes"], M06_TARGET, REVIEW_DATE, M06_CHANGE_SUMMARY)
    finish(doc, output, title, M06_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    parser.add_argument(
        "--overwrite", action="store_true",
        help="Replace a target this script already wrote, when re-rendering a fix.",
    )
    parser.add_argument(
        "--only", choices=("m04", "m06"),
        help="Write one document rather than both.",
    )
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for name, folder, stem, source, target, patch in (
        ("m04", M04_DIR, M04_STEM, M04_SOURCE, M04_TARGET, patch_m04),
        ("m06", M06_DIR, M06_STEM, M06_SOURCE, M06_TARGET, patch_m06),
    ):
        if args.only and args.only != name:
            continue
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists() and not args.overwrite:
            raise SystemExit(f"{output.name} already exists; pass --overwrite to replace it")
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
