#!/usr/bin/env python3
"""Record the per-type reversal check and branch-scoped stock in Modules 22 to 24.

Reversing an approval used to ask whether the document was a purchase order, so
three of the four procurement types always allowed one. Each type now answers for
what its own approval released: a requisition once an order has been raised from
it, an order once it reached the vendor or goods were received against it, and a
bill or a payment once its own status has left draft or pending approval. Module
22 records the requisition's answer and drops the gap it carried, and Module 23
records the two payable answers, the order's new second ground, and drops the gap
that required them.

Module 24 is narrowed by branch on both sides. The balance list, the movement
ledger, the item list and its summary read the caller's own stores and the
school's shared ones rather than every store the entity holds, and issuing and
adjusting resolve the store the same way. The gap v1.2 opened is closed with it,
and filing or re-branching a store works, which it did not.

Each document takes one minor version. Source MRD references move to v2.77, which
is cut after these documents.

    python tools/patch_procurement_reversal_stock_scope_docs.py
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.table import _Row

from generate_requirements_documents import (
    assert_no_em_dash,
    shrink_inherited_media,
    update_extended_title,
)

REVIEW_DATE = "12 September 2026"
SHORT_DATE = "12 Sep 2026"
MRD_VERSION = "2.77"
CODE_BASELINE = "Backend worktree at 19a50b3, 12 September 2026"
DEPLOYMENT = "Backend evidence only; nothing here is deployed."
VERIFIED = "Verified by 553 procurement tests."


# ── keep-format docx helpers ────────────────────────────────────────────────

def set_run_text(paragraph, text: str) -> None:
    """Rewrite a paragraph in place, keeping the formatting of its first run."""
    if not paragraph.runs:
        raise ValueError(f"Paragraph carries no run to inherit formatting from: {paragraph.text!r}")
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_cell_text(cell, text: str) -> None:
    """Rewrite a one-paragraph cell, keeping the formatting it already has."""
    for paragraph in cell.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    set_run_text(cell.paragraphs[0], text)


def append_cell_text(cell, tail: str) -> None:
    set_cell_text(cell, cell.text.strip() + tail)


def rewrite_box(cell, lines: list[str]) -> None:
    """Rewrite a callout cell line by line.

    A callout is one bold coloured heading paragraph followed by regular bullet
    paragraphs, so line i reuses paragraph i, overflow clones the last bullet,
    and surplus paragraphs are removed. Cloning the heading for a new bullet
    would render every added line as another heading.
    """
    original = list(cell.paragraphs)
    if len(original) < 2:
        raise ValueError("A callout needs a heading and at least one bullet to reuse")
    for index, line in enumerate(lines):
        if index < len(original):
            set_run_text(original[index], line)
        else:
            clone = copy.deepcopy(original[-1]._p)
            cell.paragraphs[-1]._p.addnext(clone)
            set_run_text(cell.paragraphs[-1], line)
    for paragraph in original[len(lines):]:
        paragraph._p.getparent().remove(paragraph._p)


def drop_bullet(cell, prefix: str) -> None:
    """Remove exactly one callout bullet, identified by how it starts."""
    lines = [paragraph.text for paragraph in cell.paragraphs]
    matches = [line for line in lines if line.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one bullet starting {prefix!r}, found {len(matches)}")
    lines.remove(matches[0])
    rewrite_box(cell, lines)


def unique_cells(row) -> list:
    seen, cells = set(), []
    for cell in row.cells:
        if cell._tc in seen:
            continue
        seen.add(cell._tc)
        cells.append(cell)
    return cells


def find_table(doc, header: str, *, contains: str | None = None):
    """The first table whose first cell starts with ``header``.

    ``contains`` narrows to a table that also has a row whose first cell starts
    with that text, for documents where two tables share a header.
    """
    for table in doc.tables:
        if not table.rows[0].cells[0].text.strip().startswith(header):
            continue
        if contains is None or any(
            row.cells[0].text.strip().startswith(contains) for row in table.rows
        ):
            return table
    raise ValueError(f"Table not found: {header!r} containing {contains!r}")


def find_fr(doc, label: str):
    return find_table(doc, f"{label} |")


def find_row(table, prefix: str, *, column: int = 0):
    for row in table.rows:
        if row.cells[column].text.strip().startswith(prefix):
            return row
    raise ValueError(f"Row not found: {prefix!r}")


def field(table, label: str):
    """The value cell of a requirement or control table row named ``label``."""
    return find_row(table, label).cells[1]


def clone_row(table, anchor, values: list[str], *, before: bool = False):
    new_tr = copy.deepcopy(anchor._tr)
    if before:
        anchor._tr.addprevious(new_tr)
    else:
        anchor._tr.addnext(new_tr)
    row = _Row(new_tr, table)
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)
    return row


def set_row_text(row, values: list[str]) -> None:
    cells = unique_cells(row)
    if len(cells) != len(values):
        raise ValueError(f"Row has {len(cells)} cells, {len(values)} values given")
    for cell, value in zip(cells, values):
        set_cell_text(cell, value)


def hold_together(table, *, lead_rows: int = 2) -> None:
    """Keep every row whole, and the first ``lead_rows`` with the row after them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
    for row in table.rows[:lead_rows]:
        for cell in unique_cells(row):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True


def paginate(doc) -> None:
    """Stop a requirement's heading and banner, or a table's lead, being stranded.

    Each requirement table and the traceability table keep their rows whole and
    their first rows with the next, and the paragraph directly above each keeps
    with the table, so a heading never ends a page alone. The routes intro keeps
    with the sub-heading and table that follow it for the same reason.
    """
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    for table in doc.tables:
        first = table.rows[0].cells[0].text.strip()
        if not ((first.startswith("FR-") and " | " in first) or first == "MRD capability"):
            continue
        hold_together(table)
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            Paragraph(previous, table._parent).paragraph_format.keep_with_next = True
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("All routes are mounted under /v1/procurement/"):
            paragraph.paragraph_format.keep_with_next = True


def replace_cover_version(table, source: str, target: str) -> None:
    for paragraph in table.rows[0].cells[0].paragraphs:
        for run in paragraph.runs:
            if f"Version: {source}" in run.text:
                run.text = run.text.replace(f"Version: {source}", f"Version: {target}")
                return
    raise ValueError(f"Cover version not found: {source}")


def update_control(doc, *, version: str, module: int) -> None:
    control = find_table(doc, "Document", contains="Code baseline")
    set_cell_text(field(control, "Version"), version)
    set_cell_text(field(control, "Review date"), REVIEW_DATE)
    set_cell_text(field(control, "Code baseline"), CODE_BASELINE)
    set_cell_text(
        field(control, "Source MRD"),
        f"XVS Module Requirements Document v{MRD_VERSION} | Module {module}",
    )


def body_paragraph(doc, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Body paragraph not found: {prefix!r}")


def prepend_change_log(doc, version: str, summary: str) -> None:
    table = find_table(doc, "Version", contains="1.0")
    clone_row(table, table.rows[1], [version, SHORT_DATE, summary], before=True)


def finish(doc, output: Path, title: str, version: str) -> None:
    paginate(doc)
    doc.core_properties.title = title
    doc.core_properties.version = version
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    update_extended_title(output, title)
    shrink_inherited_media(output)
    assert_no_em_dash(output)


# ═════════════════════════════════════════════════════════════════════════════
# Module 22 - Procurement & Requisitions
# ═════════════════════════════════════════════════════════════════════════════

M22_DIR = "22-procurement-and-requisitions"
M22_STEM = "XVS_M22_Procurement_and_Requisitions_Functional_Requirements_Document"
M22_SOURCE, M22_TARGET = "1.8", "1.9"

M22_FR003_EVIDENCE_TAIL = (
    " Whether a reversal is allowed at all is asked of the document before "
    "anything is written, and each document type answers for what its own "
    "approval released. A requisition refuses once a purchase order has been "
    "raised from it: that order is a commitment to a vendor which stands whether "
    "or not the vote behind it does, so it is stopped by cancelling the order "
    "rather than by withdrawing the approval it was raised on. Purchase orders, "
    "vendor bills and vendor payments answer for their own in Module 23."
)
M22_FR003_ACCEPTANCE_TAIL = (
    " ProcurementApprovalReversalTests.test_a_requisition_an_order_was_raised_from_cannot_be_reversed "
    "and test_a_requisition_with_no_order_yet_is_still_reversible hold the two "
    "answers, and test_every_procurement_document_type_answers_the_reversal_check "
    "enumerates the registered document types rather than sampling them, so a "
    "fifth cannot be added without one."
)
M22_FR003_LIMIT = (
    "The refusal reads whether any purchase order was raised from the "
    "requisition, not whether one still stands, so a requisition whose only order "
    "was cancelled keeps its approval rather than returning to the queue."
)

M22_APPROVED_LEAVES = (
    "Terminal here unless an administrator reverses the deciding vote, which "
    "returns it to PENDING. Once a purchase order has been raised from it that "
    "reversal is refused and the order is cancelled instead. Raising the order is "
    "in any case a separate act in Module 23."
)

M22_MODULE7 = (
    "Owns templates, approver resolution, routing, eligibility and the "
    "parked-approval repair. This module registers handlers, publishes its "
    "default ladders through the engine's publish service, and fences every call "
    "to the four procurement document types. Its handlers also answer the "
    "engine's reversal contract, each type for what its own approval released: a "
    "requisition an order has been raised from, an order that reached its vendor "
    "or has goods receipted against it, and a bill or a payment whose status has "
    "left draft or pending approval each refuse the reversal, while a document "
    "that released nothing returns to PENDING."
)

M22_STALE_GAP = "• Reversing an approval consults the document only for a purchase order"

M22_TRACEABILITY_LEAD = (
    f"Module 22 carries 16 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Refusing to reverse the approval of a requisition an "
    "order was raised from strengthens workflow submission without changing the "
    "count."
)

M22_CHANGE = (
    "Records that each procurement document now answers for its own reversal, and "
    "closes the gap v1.8 opened. The shared check asked whether the document was a "
    "purchase order, so three of the four types always allowed a reversal and a "
    "requisition could be put back under review with an order already placed "
    "against it. Each type is now asked what its own approval released: a "
    "requisition refuses once a purchase order has been raised from it, and "
    "Module 23's orders, bills and payments answer for theirs. FR-003, the "
    "approved state in the requisition lifecycle and the Module 7 dependency are "
    "updated, and the reversal bullet leaves Further Gaps because the gap is "
    "closed. One limit takes its place inside FR-003: the refusal reads whether "
    "any order was raised, not whether one still stands, so a cancelled order "
    "keeps the requisition approved. ProcurementApprovalReversalTests covers the "
    "requisition's two answers, and its last test enumerates the registered "
    "document types so a fifth cannot be added without one. Module 22 stays at 16 "
    f"capabilities. {VERIFIED} {DEPLOYMENT}"
)


def patch_m22(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M22 Procurement and Requisitions Functional Requirements "
        f"Document v{M22_TARGET}"
    )

    fr003 = find_fr(doc, "FR-003")
    lifecycle = find_table(doc, "State", contains="DRAFT")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M22_SOURCE, M22_TARGET)
    update_control(doc, version=M22_TARGET, module=22)

    append_cell_text(field(fr003, "Current evidence"), M22_FR003_EVIDENCE_TAIL)
    append_cell_text(field(fr003, "Acceptance"), M22_FR003_ACCEPTANCE_TAIL)
    set_cell_text(field(fr003, "Current limit"), M22_FR003_LIMIT)

    set_cell_text(find_row(lifecycle, "APPROVED").cells[2], M22_APPROVED_LEAVES)
    set_cell_text(find_row(dependencies, "Module 7").cells[1], M22_MODULE7)

    drop_bullet(gaps.rows[0].cells[0], M22_STALE_GAP)

    set_run_text(body_paragraph(doc, "Module 22 carries"), M22_TRACEABILITY_LEAD)
    prepend_change_log(doc, M22_TARGET, M22_CHANGE)
    finish(doc, output, title, M22_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 23 - Purchase Orders, Delivery & AP
# ═════════════════════════════════════════════════════════════════════════════

M23_DIR = "23-purchase-orders-delivery-and-ap"
M23_STEM = "XVS_M23_Purchase_Orders_Delivery_and_AP_Functional_Requirements_Document"
M23_SOURCE, M23_TARGET = "1.9", "1.10"

M23_FR001_EVIDENCE_TAIL = (
    " A posted goods receipt refuses the reversal on the same ground, whatever "
    "became of the email: an order whose email never went out can still have been "
    "placed by telephone, and once the goods are in, the stock and the GR/IR "
    "liability both exist."
)
M23_FR001_ACCEPTANCE_TAIL = (
    " ProcurementApprovalReversalTests.test_a_purchase_order_already_sent_to_its_vendor_cannot_be_reversed, "
    "test_a_purchase_order_the_goods_arrived_against_cannot_be_reversed and "
    "test_a_purchase_order_that_released_nothing_is_still_reversible cover the "
    "three answers."
)
M23_FR001_LIMIT = "None."

M23_FR009_EVIDENCE_TAIL = (
    " The separation holds when a decision is undone as well. A bill or a payment "
    "refuses to have its approval reversed once its own status has left draft or "
    "pending approval, which is what posting does to it, because approval is the "
    "permission to post and the journal outlives the vote. A posted payment can "
    "therefore never read as though it were still waiting on one."
)
M23_FR009_ACCEPTANCE_TAIL = (
    " ProcurementApprovalReversalTests.test_a_posted_bill_cannot_be_reversed and "
    "test_a_posted_payment_cannot_be_reversed cover the refusal, while "
    "test_a_bill_still_in_draft_is_reversible and "
    "test_a_payment_still_in_draft_is_reversible prove a document that has not "
    "posted is still reversible."
)
M23_FR009_LIMIT = (
    "The check reads the document's own status rather than the ledger, so a "
    "cancelled or reversed payment refuses the reversal too. A bill the ladder "
    "approved but nobody has posted stays reversible, which is the intended "
    "reading: approval is the permission to post, and until it is used nothing "
    "has happened."
)

M23_ORDER_REVERSAL_ERROR = [
    "Reversing the approval of an order already sent to its vendor, or received against",
    "Refused with REVERSAL_NOT_ALLOWED; cancel the order instead.",
]
M23_POSTED_REVERSAL_ERROR = [
    "Reversing the approval behind a posted bill or payment",
    "Refused with REVERSAL_NOT_ALLOWED, naming the status the document is now in.",
]

M23_MODULE7 = (
    "Runs approval over purchase orders, vendor invoices and vendor payments "
    "through the ladders Module 22 publishes. Its handlers answer the engine's "
    "reversal contract one type at a time: an order that reached its vendor or has "
    "goods receipted against it refuses the reversal, so does a bill or a payment "
    "whose status has left draft or pending approval, and a document that released "
    "nothing returns to PENDING."
)

M23_STALE_GAP = "• Reversing the approval behind a posted bill or payment is not refused."

M23_TRACEABILITY_LEAD = (
    f"Module 23 carries 23 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Refusing to reverse the approval behind a released "
    "order, or behind a posted bill or payment, strengthens workflow approval "
    "integration without changing the count."
)

M23_CHANGE = (
    "Closes the reversal gap v1.9 recorded, and widens the order's own refusal. "
    "Only a purchase order answered the engine's reversal check, so the approval "
    "behind a bill or a payment that had already posted could be sent back to "
    "pending while the ledger said it was paid, and a fresh rejection would have "
    "left a posted payment reading rejected. Each type now answers for what its "
    "own approval released: a bill or a payment refuses once its status has left "
    "draft or pending approval, which is what posting does to it. An order refuses "
    "on a second ground as well, which is new: goods received against it. An order "
    "whose email never went out can still have been placed by telephone, and once "
    "the goods are in, the stock and the GR/IR liability both exist. FR-001 and "
    "FR-009 record it, the typed-error table gains the posted bill and payment and "
    "restates the order's row, and the Module 7 dependency states the contract per "
    "type. FR-001's limit that no procurement test exercised the refusal goes with "
    "the gap: ProcurementApprovalReversalTests covers all four types and "
    "enumerates the registered ones. FR-009's limit is restated as current state, "
    "because a bill the ladder approved but nobody posted is still reversible on "
    "purpose. Module 23 stays at 23 capabilities. "
    f"{VERIFIED} {DEPLOYMENT}"
)


def patch_m23(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        "XVS M23 Purchase Orders Delivery and AP Functional Requirements "
        f"Document v{M23_TARGET}"
    )

    fr001, fr009 = find_fr(doc, "FR-001"), find_fr(doc, "FR-009")
    errors = find_table(doc, "Condition or route")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M23_SOURCE, M23_TARGET)
    update_control(doc, version=M23_TARGET, module=23)

    append_cell_text(field(fr001, "Current evidence"), M23_FR001_EVIDENCE_TAIL)
    append_cell_text(field(fr001, "Acceptance"), M23_FR001_ACCEPTANCE_TAIL)
    set_cell_text(field(fr001, "Current limit"), M23_FR001_LIMIT)

    append_cell_text(field(fr009, "Current evidence"), M23_FR009_EVIDENCE_TAIL)
    append_cell_text(field(fr009, "Acceptance"), M23_FR009_ACCEPTANCE_TAIL)
    set_cell_text(field(fr009, "Current limit"), M23_FR009_LIMIT)

    order_row = find_row(errors, "Reversing the approval of an order")
    set_row_text(order_row, M23_ORDER_REVERSAL_ERROR)
    clone_row(errors, order_row, M23_POSTED_REVERSAL_ERROR)

    set_cell_text(find_row(dependencies, "Module 7").cells[1], M23_MODULE7)

    drop_bullet(gaps.rows[0].cells[0], M23_STALE_GAP)

    set_run_text(body_paragraph(doc, "Module 23 carries"), M23_TRACEABILITY_LEAD)
    prepend_change_log(doc, M23_TARGET, M23_CHANGE)
    finish(doc, output, title, M23_TARGET)


# ═════════════════════════════════════════════════════════════════════════════
# Module 24 - Inventory & Stock Ledger
# ═════════════════════════════════════════════════════════════════════════════

M24_DIR = "24-inventory-and-stock-ledger"
M24_STEM = "XVS_M24_Inventory_and_Stock_Ledger_Functional_Requirements_Document"
M24_SOURCE, M24_TARGET = "1.2", "1.3"

M24_READER_DOES = (
    "Read items, stock locations, per-location balances, movements, and the "
    "valuation and reorder views of the item list, for the stores they work in."
)

M24_FR008_EVIDENCE = (
    "Valuation and reorder are read from the item list and its summary rather "
    "than from separate reports. GET /stock-items/?needs_reorder=true lists the "
    "active items at or below their reorder level, and GET /stock-items/summary/ "
    "returns item counts, low-stock and out-of-stock counts and the carried value "
    "in kobo, alongside the movement ledger. Both accept a location, which makes "
    "every row, and the summary, report that store's quantity, value, unit cost "
    "and reorder state. Without one they answer for the stores the caller works "
    "in, summed across them, and only a caller entitled to every store reads the "
    "entity roll-up. The catalogue itself is not narrowed: every item is listed, "
    "and one held only where this caller cannot see reports nothing rather than "
    "falling back to the entity's total. Reorder is measured against the total in "
    "scope, because a school holding four hundred at one branch and none at "
    "another is not short of them."
)
M24_FR008_ACCEPTANCE = (
    "Valuation is read from the maintained value rather than recomputed from the "
    "movement history, and the two agree because every movement maintains it. The "
    "entity roll-up is what reconciles to the inventory control account; the "
    "stores one caller can see do not, and are not meant to. "
    "StockBranchScopeTests.test_the_item_list_reports_what_this_callers_stores_hold, "
    "test_an_item_held_only_at_another_branch_reports_nothing and "
    "test_the_summary_values_only_the_stores_in_scope cover the narrowing; "
    "test_head_offices_item_list_is_the_school_total and "
    "test_head_offices_summary_is_the_school_total cover the caller who reads "
    "every store; and SingleBranchStockIsUnchangedTests proves a school with one "
    "branch reads exactly what it read before."
)
M24_FR008_LIMIT = (
    "The counts describe the whole catalogue while the figures describe the stores "
    "in scope, so an item held only at another branch is counted as out of stock "
    "for a caller pinned to one, rather than left out of their counts."
)

M24_FR009_EVIDENCE_TAIL = (
    " Reads of that stock are narrowed the same way, through the store each row "
    "sits at: the balance list, the movement ledger, the item list and its summary "
    "all answer for the caller's own stores and the school's shared ones rather "
    "than for every store the entity holds. So does a movement. A caller pinned to "
    "branches who names no store is answered from the ones they work in, where a "
    "single visible store needs no naming, none is refused with a message saying "
    "to ask for one, and a choice between several is refused asking which, without "
    "naming a store they may not see. A caller entitled to every store keeps the "
    "stock service's own defaulting."
)
M24_FR009_ACCEPTANCE_TAIL = (
    " StockBranchScopeTests.test_a_storekeeper_reads_their_own_store_and_the_shared_one "
    "and test_a_storekeeper_never_reads_another_branchs_movements cover the reads; "
    "StockMovementComesFromTheCallersStoreTests.test_issuing_without_a_store_cannot_draw_from_another_branchs, "
    "test_adjusting_without_a_store_cannot_touch_another_branchs and "
    "test_a_central_store_is_still_everybodys_to_issue_from cover the writes; and "
    "StockLocationBranchWriteTests.test_a_storekeeper_files_a_new_store_at_their_own_branch, "
    "test_somebody_over_two_branches_files_it_school_wide and "
    "test_a_storekeeper_cannot_file_a_store_at_another_branch cover the branch a "
    "filed store takes."
)

M24_FR011_EVIDENCE_TAIL = (
    " Where the caller is pinned to branches, the store a movement is answered "
    "from is one of their own rather than the entity's default, so a school whose "
    "only book store stands at another branch asks which store rather than issuing "
    "from a shelf its storekeeper cannot even open."
)
M24_FR011_ACCEPTANCE = (
    "A single-store entity's calls and responses are unchanged, and so are those "
    "of a school with one branch, whose storekeepers are not narrowed at all. Once "
    "more than one store is within the caller's reach, a movement that names none "
    "is refused rather than defaulted, because silently drawing from the main "
    "store when somebody meant the annex is a quieter version of the bug locations "
    "exist to fix. A storekeeper who has exactly one store of their own is not "
    "asked which, however many the school has."
)

M24_MOVEMENT_ERROR = [
    "Moving stock without naming a location when the caller can reach more than one",
    "Refused rather than defaulted, and refused as well when they can reach none.",
]

M24_BALANCE_ROUTE = "What each item holds at each store the caller works in."
M24_ITEMS_ROUTE = (
    "The item master. Figures describe the caller's stores; ?location= narrows "
    "them to one, and ?needs_reorder=true lists what is running out."
)
M24_SUMMARY_ROUTE = (
    "Counts, low and out of stock, and carried value, for the caller's stores or "
    "one named store."
)
M24_MOVEMENTS_ROUTE = "The immutable movement ledger, for the caller's stores."

M24_MODULE4 = (
    "Supplies the separately grantable view, manage, issue and adjust keys, and "
    "the branch narrowing stock locations apply inclusively. Balances, movements "
    "and item figures reach that same narrowing through the location they sit at, "
    "so what a caller may read and what they may move from are one rule."
)

M24_STALE_GAP = "• The balance list, the movement ledger, the item list and its summary are narrowed only"

M24_TRACEABILITY_LEAD = (
    f"Module 24 carries 13 capability entries in MRD v{MRD_VERSION}. Each maps to "
    "the requirements below. Narrowing balances, movements and item figures to the "
    "caller's own stores strengthens stock locations and per-location balances "
    "without changing the count."
)

M24_CHANGE = (
    "Closes the gap v1.2 opened: stock is narrowed by branch on the way out and on "
    "the way in. The balance list, the movement ledger, the item list and its "
    "summary narrowed only when a location was named, so a storekeeper pinned to "
    "one branch read another's balances and every movement at them by leaving the "
    "filter off, and issuing without naming a store drew from another branch's "
    "shelf. All four reads now go through the catalogue rule the rest of "
    "procurement uses, reached through the store a row sits at. Item figures are "
    "summed over the stores in scope and every listed item is answered, so one "
    "held only at another branch reports nothing rather than falling back to the "
    "school total. Issuing and adjusting resolve the store the same way: a caller "
    "pinned to branches who names none is answered from their own, where one needs "
    "no naming, none is refused, and several is refused asking which. A null branch "
    "still means shared across the school, so a central store stays readable and "
    "usable from every branch, and a school with one branch or one store reads and "
    "moves exactly what it did before. Filing a store, or moving one to another "
    "branch, now works at all: the wrapper behind those routes dropped the keyword "
    "that tells master data to take the shared reading, and creation and patching "
    "raised a server error on five routes nothing covered. FR-008, FR-009 and "
    "FR-011, the actors table, the route purposes, the typed errors, the Module 4 "
    "dependency and MRD traceability are updated. Module 24 stays at 13 "
    f"capabilities. {VERIFIED} {DEPLOYMENT}"
)


def patch_m24(source: Path, output: Path) -> None:
    doc = Document(str(source))
    title = (
        f"XVS M24 Inventory and Stock Ledger Functional Requirements Document v{M24_TARGET}"
    )

    actors = find_table(doc, "Actor")
    fr008, fr009, fr011 = find_fr(doc, "FR-008"), find_fr(doc, "FR-009"), find_fr(doc, "FR-011")
    routes = find_table(doc, "Method and path")
    errors = find_table(doc, "Condition")
    dependencies = find_table(doc, "Dependency")
    gaps = find_table(doc, "FURTHER GAPS")

    replace_cover_version(doc.tables[0], M24_SOURCE, M24_TARGET)
    update_control(doc, version=M24_TARGET, module=24)

    set_cell_text(find_row(actors, "Stores reader").cells[1], M24_READER_DOES)

    set_cell_text(field(fr008, "Current evidence"), M24_FR008_EVIDENCE)
    set_cell_text(field(fr008, "Acceptance"), M24_FR008_ACCEPTANCE)
    set_cell_text(field(fr008, "Current limit"), M24_FR008_LIMIT)

    append_cell_text(field(fr009, "Current evidence"), M24_FR009_EVIDENCE_TAIL)
    append_cell_text(field(fr009, "Acceptance"), M24_FR009_ACCEPTANCE_TAIL)

    append_cell_text(field(fr011, "Current evidence"), M24_FR011_EVIDENCE_TAIL)
    set_cell_text(field(fr011, "Acceptance"), M24_FR011_ACCEPTANCE)

    set_cell_text(find_row(routes, "GET /stock-balances/").cells[1], M24_BALANCE_ROUTE)
    set_cell_text(find_row(routes, "GET, POST /stock-items/").cells[1], M24_ITEMS_ROUTE)
    set_cell_text(find_row(routes, "GET /stock-items/summary/").cells[1], M24_SUMMARY_ROUTE)
    set_cell_text(find_row(routes, "GET /stock-movements/").cells[1], M24_MOVEMENTS_ROUTE)

    set_row_text(find_row(errors, "Moving stock without naming a location"), M24_MOVEMENT_ERROR)
    set_cell_text(find_row(dependencies, "Module 4").cells[1], M24_MODULE4)

    drop_bullet(gaps.rows[0].cells[0], M24_STALE_GAP)

    set_run_text(body_paragraph(doc, "Module 24 carries"), M24_TRACEABILITY_LEAD)
    prepend_change_log(doc, M24_TARGET, M24_CHANGE)
    finish(doc, output, title, M24_TARGET)


# ═════════════════════════════════════════════════════════════════════════════

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parents[1])
    parser.add_argument(
        "--only", nargs="*", choices=["m22", "m23", "m24"],
        help="Write only these modules; the default writes all three.",
    )
    args = parser.parse_args()
    root = Path(args.root) / "functional-requirements"

    for tag, folder, stem, source, target, patch in (
        ("m22", M22_DIR, M22_STEM, M22_SOURCE, M22_TARGET, patch_m22),
        ("m23", M23_DIR, M23_STEM, M23_SOURCE, M23_TARGET, patch_m23),
        ("m24", M24_DIR, M24_STEM, M24_SOURCE, M24_TARGET, patch_m24),
    ):
        if args.only and tag not in args.only:
            continue
        directory = root / folder
        output = directory / f"{stem}_v{target}.docx"
        if output.exists():
            raise SystemExit(f"Refusing to overwrite {output.name}")
        patch(directory / f"{stem}_v{source}.docx", output)
        print(f"Wrote {folder} v{target}")


if __name__ == "__main__":
    main()
